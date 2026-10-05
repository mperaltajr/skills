"""Build Slide-Lab-Install-Guide.docx (2 pages: page 1 the install, page 2 everything after).

  py -3 docs/make_install_guide.py

The paste-in request is copied from README.md "Step 2" on every run, so the two
cannot drift; the build fails if it is not one plain ASCII line. Layout follows
the 2026-10-02 design ruling: 11 pt body on 15 pt leading, about 88-character
lines, bold accent headers over thin rules (prints cleanly), two callouts.
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parents[1] / "Slide-Lab-Install-Guide.docx"
OUT.parent.mkdir(exist_ok=True)
AS_OF = "Accurate as of 5 October 2026"
REPO_URL = "https://github.com/mperaltajr/skills"
ACC = "0B3C49"; ACCENT = RGBColor(0x0B, 0x3C, 0x49)
TEXT = RGBColor(0x22, 0x22, 0x22); MUTED = RGBColor(0x5A, 0x5A, 0x5A)
RULE = "C9D1D4"; TINT = "EEF2F3"; PASTE_FILL = "F5F7F8"
FONT, MONO = "Calibri", "Consolas"

def req():
    t = (REPO / "README.md").read_text(encoding="utf-8")
    m = re.search(r"^### Step 2\..*?^```[^\n]*\n(.*?)\n```", t, re.S | re.M)
    r = m.group(1).strip("\n")
    assert "\n" not in r and "\t" not in r and all(32 <= ord(c) <= 126 for c in r)
    return r

def exact(pf, pts):
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(pts)

def pborder(p, sides, sz, color, space):
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    for s in ("top", "left", "bottom", "right"):
        if s in sides:
            e = OxmlElement(f"w:{s}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sides[s] if isinstance(sides, dict) else sz))
            e.set(qn("w:space"), str(space[s] if isinstance(space, dict) else space)); e.set(qn("w:color"), color)
            b.append(e)
    ppr.append(b)

def pshade(p, fill):
    ppr = p._p.get_or_add_pPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill)
    ppr.append(s)

def link(p, url, text, size):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts"); f.set(qn("w:ascii"), FONT); f.set(qn("w:hAnsi"), FONT); rpr.append(f)
    c = OxmlElement("w:color"); c.set(qn("w:val"), ACC); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); p._p.append(h)

def rich(p, parts, size, color=TEXT):
    for part in parts:
        if isinstance(part, str): part = (part, "")
        txt, st = part
        if st.startswith("link:"):
            link(p, st[5:], txt, size); continue
        r = p.add_run(txt); r.font.size = Pt(size); r.font.color.rgb = color
        r.bold = "b" in st; r.italic = "i" in st

def para(c, parts, size=11, lead=15, before=0, after=6, color=TEXT, kwn=False):
    p = c.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); exact(pf, lead)
    pf.keep_with_next = kwn; pf.widow_control = True
    rich(p, parts if isinstance(parts, list) else [parts], size, color)
    return p

def heading(doc, text, before=12):
    p = doc.add_paragraph(style="Heading 1"); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(6); exact(pf, 18); pf.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = ACCENT; r.font.name = FONT
    rpr = r._element.get_or_add_rPr(); rpr.rFonts.set(qn("w:eastAsia"), FONT)
    return p

def callout(doc, parts, after=8):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for sd in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{sd}")
        if sd == "left":
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "24"); e.set(qn("w:color"), ACC)
        else:
            e.set(qn("w:val"), "nil")
        e.set(qn("w:space"), "0"); b.append(e)
    pr.append(b)
    c = t.rows[0].cells[0]
    tcpr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), TINT); tcpr.append(sh)
    cmar(c, t=80, b=80, l=180, r=144)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); exact(p.paragraph_format, 15)
    rich(p, parts, 11)
    widths(t, [5.7])
    t._tbl.tblPr.find(qn("w:tblInd")).set(qn("w:w"), "180")
    g = doc.add_paragraph(); g.paragraph_format.space_after = Pt(0); exact(g.paragraph_format, after)
    return p

def pborder_extra(p):
    b = p._p.pPr.find(qn("w:pBdr"))
    for s, sp in (("top", 4), ("bottom", 4)):
        e = OxmlElement(f"w:{s}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), str(sp)); e.set(qn("w:color"), TINT)
        b.insert(0 if s == "top" else len(b), e)
    # reorder to schema order top,left,bottom,right
    order = {"top": 0, "left": 1, "bottom": 2, "right": 3}
    kids = sorted(list(b), key=lambda e: order[e.tag.split("}")[1]])
    for k in list(b): b.remove(k)
    for k in kids: b.append(k)
    e = OxmlElement("w:right"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
    e.set(qn("w:space"), "4"); e.set(qn("w:color"), TINT); b.append(e)

def paste(doc, text):
    lab = para(doc, [("Copy all of this into Claude (it is one long line)", "b")], size=10, lead=13, after=3, color=ACCENT, kwn=True)
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(8); exact(pf, 14)
    pf.left_indent = Pt(9); pf.right_indent = Pt(9); pf.keep_together = True
    pborder(p, ("top", "left", "bottom", "right"), 6, ACC, {"top": 6, "bottom": 6, "left": 8, "right": 8})
    pshade(p, PASTE_FILL)
    r = p.add_run(text); r.font.name = MONO; r.font.size = Pt(10); r.font.color.rgb = TEXT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), MONO); r._element.rPr.rFonts.set(qn("w:hAnsi"), MONO)

def tborders(t):
    pr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{s}")
        if s in ("bottom", "insideH"):
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), RULE)
        else:
            e.set(qn("w:val"), "nil")
        e.set(qn("w:space"), "0"); b.append(e)
    pr.append(b)

def cell_bottom(cell, sz, color):
    tcpr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color); e.set(qn("w:space"), "0")
    b.append(e); tcpr.append(b)

def cmar(cell, t=50, b=50, l=0, r=144):
    tcpr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for s, v in (("top", t), ("start", l), ("bottom", b), ("end", r)):
        e = OxmlElement(f"w:{s}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcpr.append(m)

def widths(t, ws):
    t.autofit = False; pr = t._tbl.tblPr
    w = pr.find(qn("w:tblW"))
    w.set(qn("w:w"), str(int(sum(ws) * 1440))); w.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); pr.append(lay)
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), "0"); ind.set(qn("w:type"), "dxa"); pr.append(ind)
    for gc, x in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), ws): gc.set(qn("w:w"), str(int(x * 1440)))
    for row in t.rows:
        for c, x in zip(row.cells, ws): c.width = Inches(x)

def table(doc, header, rows, ws, after_gap=True):
    n = len(rows) + (1 if header else 0)
    t = doc.add_table(rows=n, cols=len(ws)); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    tborders(t)
    ri0 = 0
    if header:
        tr = t.rows[0]._tr.get_or_add_trPr(); tr.append(OxmlElement("w:tblHeader")); tr.append(OxmlElement("w:cantSplit"))
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; cmar(c, t=40, b=60); cell_bottom(c, 8, ACC)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); exact(p.paragraph_format, 13)
            p.paragraph_format.keep_with_next = True
            rich(p, [(h, "b")], 10, ACCENT)
        ri0 = 1
    for ri, row in enumerate(rows, start=ri0):
        t.rows[ri]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]; cmar(c)
            paras = val if (isinstance(val, list) and val and isinstance(val[0], list)) else [val]
            for k, pv in enumerate(paras):
                p = c.paragraphs[0] if k == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(4 if k < len(paras) - 1 else 0); exact(p.paragraph_format, 13)
                parts = pv if isinstance(pv, list) else [pv]
                if ci == 0: parts = [(x, "b") if isinstance(x, str) else x for x in parts]
                rich(p, parts, 10)
    widths(t, ws)
    g = doc.add_paragraph(); g.paragraph_format.space_after = Pt(0); exact(g.paragraph_format, 6)
    return t

def page_field(p, instr):
    for kind, val in (("begin", None), ("instr", f" {instr} "), ("separate", None), ("text", "1"), ("end", None)):
        r = p.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
        if kind == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = val
        elif kind == "text":
            e = OxmlElement("w:t"); e.text = val
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
        r._r.append(e)

def build():
    request = req()
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1.4)
    sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7)
    sec.header_distance = Inches(0.4); sec.footer_distance = Inches(0.35)
    n = doc.styles["Normal"]; n.font.name = FONT; n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    n.font.size = Pt(11); n.font.color.rgb = TEXT
    h1 = doc.styles["Heading 1"]; h1.font.name = FONT; h1.font.size = Pt(14); h1.font.color.rgb = ACCENT; h1.font.bold = True
    rf = h1.element.rPr.rFonts
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    rf.set(qn("w:ascii"), FONT); rf.set(qn("w:hAnsi"), FONT); rf.set(qn("w:eastAsia"), FONT)
    doc.core_properties.title = "Slide Lab install guide"; doc.core_properties.author = "Slide Lab"; doc.core_properties.comments = AS_OF

    fp = sec.footer.paragraphs[0]
    r = fp.add_run(f"Slide Lab install guide  |  {AS_OF}  |  Page "); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    page_field(fp, "PAGE")
    r = fp.add_run(" of "); r.font.size = Pt(8.5); r.font.color.rgb = MUTED
    page_field(fp, "NUMPAGES")

    # Title block
    p = doc.add_paragraph(style="Title") if False else doc.add_paragraph()
    pf = p.paragraph_format; pf.space_after = Pt(4); exact(pf, 28)
    r = p.add_run("Slide Lab install guide"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = ACCENT
    p = para(doc, ["Three steps, about 20 to 45 minutes the first time, plus however long your "
                   "software portal takes to approve and install LibreOffice."], size=12, lead=16, after=10, color=MUTED)
    p = para(doc, ["Slide Lab is an add-on for Claude that turns your message, notes and numbers "
                   "into a finished PowerPoint deck on your client's own template, with their "
                   "colors, fonts and layouts. It helps you shape the storyline, shows you three "
                   "designs for each slide to pick from, then builds the deck and checks every "
                   "slide before you open it."], after=0)
    pborder(p, {"bottom": 8}, 8, ACC, {"bottom": 10})

    heading(doc, "Step 1. Install three programs yourself", before=10)
    callout(doc, [("Request LibreOffice on day one. ", "b"), "On company computers it can need approval."])
    para(doc, ["Install these yourself from your company's software portal (Software Center, "
               "Self Service or similar). If you are not on a company computer, use the official sites."], after=6, kwn=True)
    table(doc, ["Program", "Why Slide Lab needs it", "Where to get it"], [
        ["Claude (desktop app or Claude Code)", "Runs Slide Lab",
         ["However your company provides Claude; otherwise ", ("claude.ai/code", "link:https://claude.ai/code")]],
        ["Python 3.10 or newer", "Runs the build steps behind the scenes",
         ["Software portal, or ", ("python.org", "link:https://www.python.org/downloads/")]],
        ["LibreOffice", "Draws slide previews. Required on a Mac; on Windows, PowerPoint can stand in",
         ["Software portal, or ", ("libreoffice.org", "link:https://www.libreoffice.org/download/")]],
    ], [1.45, 2.15, 2.1])

    heading(doc, "Step 2. Let Claude do the rest")
    para(doc, ["Open Claude (desktop app or Claude Code) and paste this request exactly as written."], kwn=True)
    paste(doc, request)
    para(doc, ["Claude asks before each download or install (your company's settings may "
               "require that). Click ", ("Allow", "b"), " each time."])
    para(doc, ["At the end Claude shows a setup table. If every piece is OK, the last line "
               "reads \"Slide Lab is ready.\" Any row that is not OK says exactly what to do."], after=0)

    heading(doc, "Step 3. Restart Claude")
    para(doc, ["Close and reopen Claude so it loads Slide Lab. Then type ",
               ("/slide-lab", "b"), " or just ask for a deck."], after=0)

    # Page 2
    h = heading(doc, "Your first deck", before=0); h.paragraph_format.page_break_before = True
    cp = callout(doc, [("Keep your work out of OneDrive. ", "b"), "Work in this folder on your own drive:"])
    cp.add_run().add_break()
    rr = cp.add_run("C:\\Users\\<you>\\Slide Lab\\sessions\\<Client>\\"); rr.bold = True; rr.font.size = Pt(11); rr.font.color.rgb = TEXT
    cp.add_run().add_break()
    rich(cp, ["The setup creates it, with a Desktop shortcut. Syncing hundreds of working files slowed builds from about 4 minutes per page to 9 to 20 and made previews time out. Copy finished decks to OneDrive or SharePoint."], 11)
    table(doc, None, [
        ["Your client's template", ["It can live anywhere, but must be 16:9 at 13.333 x 7.5 inches (PowerPoint's standard widescreen). Register each template once, about 5 to 10 minutes. Claude asks which colors are the main, "
                                     "highlight and cover colors and which layouts to use, then you open one sample slide in "
                                     "PowerPoint and confirm it looks right. Your original file is never changed."]],
        ["Start small", ["Start with 5 to 8 slides. Bring the main message, the audience, and your numbers or notes."]],
        ["How long a deck takes", ["A typical 10 to 15 slide deck takes 1 to 2 hours end to end, including your review time."]],
        ["When you are done", ["Say ", ("the deck is good", "b"), ". Slide Lab deletes the working files it no longer needs "
                               "and keeps the deck and everything needed to edit it later."]],
    ], [1.45, 4.25])

    heading(doc, "Getting updates")
    para(doc, ["Slide Lab checks for updates once a day when you start using it, and asks before "
               "installing. Say yes, and restart Claude if it asks you to. To check any time, say ",
               ("update Slide Lab", "b"), "."])

    heading(doc, "If something goes wrong")
    table(doc, ["If this happens", "Do this"], [
        ["Something in the setup is not working", ["Tell Claude ", ("check my Slide Lab setup", "b"),
                                                   ". It runs a checker, changes nothing, and shows what to fix."]],
        ["The setup check says Git is missing", ["Git is a download tool Slide Lab uses. It is usually already there, because "
                                                 "Claude Code on Windows uses it. Get it from your software portal or ",
                                                 ("git-scm.com", "link:https://git-scm.com"), "."]],
        ["A \"certificate\" error during setup", ["Try again first; it usually clears on its own. If it keeps happening, "
                                                  "tell Claude, or see the Install section of README.md."]],
        ["Anything else, or a deck did not come out right", ["Type ", ("/slidelab-log", "b"),
                                                             " in Claude. It writes the problem report for you; send the file to Mario Peralta on Teams or by email, never anywhere public."]],
    ], [1.9, 3.8])

    heading(doc, "Where to learn more")
    para(doc, ["Open ", ("Slide-Lab-Tutorial.html", "b"), " in the Slide Lab folder for a step-by-step "
               "walkthrough with timings. The example files next to it (Slide-Lab-Example-Storyline.html, "
               "Slide-Lab-Example-Review.html and Slide-Lab-Example-Deck.pptx) show what each stage looks like."])
    para(doc, [("Prefer to type the install commands yourself? ", "b"),
               "See the Install section of README.md in the Slide Lab folder or at ",
               (REPO_URL, "link:" + REPO_URL), "."], size=10, lead=13, color=MUTED, after=0)
    doc.save(OUT); return OUT

if __name__ == "__main__":
    print(build())
