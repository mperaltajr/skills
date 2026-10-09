"""html_emit.py — draw a translation plan into a native, editable PowerPoint slide.

translate_html.py reads a worker's HTML design in the browser that drew it and
turns it into a plan: a list of shapes and text boxes, in paint order, with
every geometry and style decision already made. This module only draws that
plan. Keeping the drawing here (rather than in each generated script) means a
fix to how something is drawn applies to every slide on the next finalize.

Each op kind maps to plain python-pptx: rectangles, rounded rectangles, ovals,
triangles, freeform polygons and polylines, straight connectors, text boxes
with styled runs (a whole list is one text box with real bullets), native
tables (graphicFrame), and library icons (inserted as their vector shapes
through icon_helper). No charts, no embedded objects, no pictures of text.

The rules that used to live as prose in the translator agent's instructions are
code here, applied the same way every time:
  - no shadow is effectRef idx="0", never a deleted <p:style> child
  - letter spacing is the raw spc attribute at px x 75
  - line spacing is exact points from the computed line-height
  - font sizes snap to PowerPoint's grid
  - one-line text never wraps
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from twins.helpers import set_letter_spacing, snap_font_pt

CANVAS_W, CANVAS_H = 1280, 720

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
          "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
_KIND = {"rect": MSO_SHAPE.RECTANGLE, "rounded": MSO_SHAPE.ROUNDED_RECTANGLE,
         "oval": MSO_SHAPE.OVAL, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE}
_DASH = {"dashed": MSO_LINE.DASH, "dotted": MSO_LINE.ROUND_DOT}


def E(px: float) -> Emu:
    return Emu(int(round(px * 9525)))


def _no_shadow(shape) -> None:
    """LibreOffice draws the theme's preset shadow on any autoshape; the designs
    are flat. effectRef idx="0" says "no effect" while keeping <p:style> whole:
    deleting the child instead is what made PowerPoint refuse a deck."""
    eff = shape._element.find(".//" + qn("a:effectRef"))
    if eff is not None:
        eff.set("idx", "0")


def _fill(shape, color, alpha: float = 1.0) -> None:
    if not color:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    if alpha < 0.999:
        srgb = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(round(alpha * 100000)))}))


def _gradient(shape, grad: dict, alpha: float = 1.0) -> None:
    """A CSS linear gradient as a native gradient fill.

    CSS measures the angle clockwise from "up" (180deg = top to bottom);
    DrawingML measures it clockwise from "left to right", so subtract 90.
    Each stop keeps its own transparency.
    """
    shape.fill.gradient()
    grad_fill = shape._element.spPr.find(qn("a:gradFill"))
    gs_lst = grad_fill.find(qn("a:gsLst"))
    for child in list(gs_lst):
        gs_lst.remove(child)
    for pos, color, a in grad["stops"]:
        gs = gs_lst.makeelement(qn("a:gs"), {"pos": str(int(round(pos * 100000)))})
        clr = gs.makeelement(qn("a:srgbClr"), {"val": color})
        a = a * alpha
        if a < 0.999:
            clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(round(a * 100000)))}))
        gs.append(clr)
        gs_lst.append(gs)
    for old in grad_fill.findall(qn("a:lin")) + grad_fill.findall(qn("a:path")):
        grad_fill.remove(old)
    ang = (grad["angle"] - 90) % 360
    lin = grad_fill.makeelement(qn("a:lin"), {"ang": str(int(round(ang * 60000))), "scaled": "0"})
    gs_lst.addnext(lin)


def _line(shape, line) -> None:
    if not line:
        shape.line.fill.background()
        return
    shape.line.color.rgb = RGBColor.from_string(line["color"])
    shape.line.width = E(line["w"])
    if line.get("dash") in _DASH:
        shape.line.dash_style = _DASH[line["dash"]]


def _shape(slide, op) -> None:
    sp = slide.shapes.add_shape(_KIND[op["kind"]], E(op["x"]), E(op["y"]),
                                E(max(op["w"], 0.5)), E(max(op["h"], 0.5)))
    if op["kind"] == "rounded" and op.get("radius_ratio") is not None:
        sp.adjustments[0] = op["radius_ratio"]
    if op.get("rotation"):
        sp.rotation = op["rotation"]
    if op.get("gradient"):
        _gradient(sp, op["gradient"], op.get("alpha", 1.0))
    else:
        _fill(sp, op.get("fill"), op.get("alpha", 1.0))
    _line(sp, op.get("line"))
    _no_shadow(sp)
    if op.get("name"):
        sp.name = op["name"]


def _freeform(slide, op) -> None:
    pts = op["points"]
    if len(pts) < 2:
        return
    builder = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    builder.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=bool(op.get("closed")))
    sp = builder.convert_to_shape()
    _fill(sp, op.get("fill") if op.get("closed") else None, op.get("alpha", 1.0))
    _line(sp, op.get("line"))
    _no_shadow(sp)
    if op.get("name"):
        sp.name = op["name"]


def _connector(slide, op) -> None:
    (x1, y1), (x2, y2) = op["points"]
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _line(c, op.get("line"))
    # Connectors carry the same theme style, shadow included.
    _no_shadow(c)
    if op.get("name"):
        c.name = op["name"]


def _bullet(p, spec) -> None:
    """A real PowerPoint bullet on paragraph p: a character (buChar) or a
    number (buAutoNum), in its own color and size. spec None = no bullet."""
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct", "a:buSzPts", "a:buFontTx",
                "a:buFont", "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    # bullets go after lnSpc / spcBef / spcAft and before tabLst / defRPr
    anchor = next((pPr.find(qn(t)) for t in ("a:tabLst", "a:defRPr", "a:extLst")
                   if pPr.find(qn(t)) is not None), None)

    def put(el):
        if anchor is not None:
            anchor.addprevious(el)
        else:
            pPr.append(el)
        return el

    if not spec:
        put(pPr.makeelement(qn("a:buNone"), {}))
        return
    if spec.get("color"):
        clr = put(pPr.makeelement(qn("a:buClr"), {}))
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": spec["color"]}))
    pct = int(round(max(25.0, min(400.0, float(spec.get("pct") or 100))) * 1000))
    put(pPr.makeelement(qn("a:buSzPct"), {"val": str(pct)}))
    if spec.get("autonum"):
        put(pPr.makeelement(qn("a:buAutoNum"), {"type": spec["autonum"],
                                                 "startAt": str(int(spec.get("start") or 1))}))
    else:
        put(pPr.makeelement(qn("a:buFont"), {"typeface": spec.get("font") or "Arial"}))
        put(pPr.makeelement(qn("a:buChar"), {"char": spec["char"]}))


def _paragraphs(tf, op, default_lh_px: float, end_size: bool = False) -> None:
    """Write op's paragraphs (runs) into text frame tf. op["para"], when there,
    holds per-paragraph settings: list level, left/right indent, hanging
    indent, space before, line height and the bullet (a one-box list)."""
    props = op.get("para") or []
    for i, para in enumerate(op["paragraphs"]):
        pr = props[i] if i < len(props) else {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN.get(pr.get("align") or op.get("align", "left"), PP_ALIGN.LEFT)
        lh_px = pr.get("line_height_px") or default_lh_px
        if lh_px:
            p.line_spacing = Pt(lh_px * 0.75)
        if props:
            if pr.get("space_before_px"):
                p.space_before = Pt(round(pr["space_before_px"] * 0.75, 2))
            pPr = p._p.get_or_add_pPr()
            if pr.get("lvl"):
                pPr.set("lvl", str(min(8, int(pr["lvl"]))))
            pPr.set("marL", str(int(round(max(0.0, pr.get("marL_px") or 0) * 9525))))
            if pr.get("marR_px"):
                pPr.set("marR", str(int(round(max(0.0, pr["marR_px"]) * 9525))))
            pPr.set("indent", str(int(round((pr.get("indent_px") or 0) * 9525))))
            _bullet(p, pr.get("bullet"))
        size_pt = None
        for r in para:
            if not r["t"]:
                continue
            run = p.add_run()
            run.text = r["t"]
            f = run.font
            if r.get("font"):
                f.name = r["font"]
            f.size = Pt(snap_font_pt(r["size_px"] * 0.75))
            size_pt = size_pt or f.size
            f.bold = bool(r.get("bold"))
            f.italic = bool(r.get("italic"))
            if r.get("color"):
                f.color.rgb = RGBColor.from_string(r["color"])
            if r.get("letter_spacing_px"):
                set_letter_spacing(run, r["letter_spacing_px"])
        if end_size and size_pt is not None:
            # the paragraph mark at the text's size, so an empty line or a
            # table row is not held open at the 18pt default
            end = p._p.get_or_add_endParaRPr()
            end.set("sz", str(int(round(size_pt.pt * 100))))


def _text(slide, op) -> None:
    tb = slide.shapes.add_textbox(E(op["x"]), E(op["y"]), E(max(op["w"], 1)), E(max(op["h"], 1)))
    if op.get("name"):
        tb.name = op["name"]
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.auto_size = MSO_AUTO_SIZE.NONE
    # One-line text never wraps. PowerPoint and LibreOffice set a font slightly
    # wider than Chrome, so a box that fits in Chrome by a pixel would break a
    # short word mid-word ("Taa S"). Letting it overhang is the lesser error.
    tf.word_wrap = bool(op.get("wrap"))
    _paragraphs(tf, op, op["line_height_px"], end_size=bool(op.get("para")))


# "No Style, No Grid": the table carries no theme banding or grid of its own;
# every fill and border below is set cell by cell, as the design drew it.
_NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
_EDGES = (("L", "a:lnL"), ("R", "a:lnR"), ("T", "a:lnT"), ("B", "a:lnB"))


def _edge(tcPr, tag, line) -> None:
    ln = tcPr.makeelement(qn(tag), {"w": str(int(round((line or {}).get("w", 0) * 9525))) if line else "0"})
    if line:
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": line["color"]}))
        ln.append(sf)
        dash = {"dashed": "dash", "dotted": "sysDot"}.get(line.get("dash") or "", "solid")
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    else:
        ln.append(ln.makeelement(qn("a:noFill"), {}))
    tcPr.append(ln)


def _table(slide, op) -> None:
    """A native PowerPoint table: column widths, row heights, merged cells,
    per-cell fill, per-edge borders, inner margins and styled text, as the
    plan decided (translate_html._table_group)."""
    nr, nc = len(op["rows"]), len(op["cols"])
    gf = slide.shapes.add_table(nr, nc, E(op["x"]), E(op["y"]), E(op["w"]), E(op["h"]))
    if op.get("name"):
        gf.name = op["name"]
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        if tblPr.get(attr) is not None:
            tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = _NO_STYLE
    for j, w in enumerate(op["cols"]):
        tbl.columns[j].width = E(w)
    for i, h in enumerate(op["rows"]):
        tbl.rows[i].height = E(h)
    for c in op["cells"]:
        if c["rs"] > 1 or c["cs"] > 1:
            tbl.cell(c["r"], c["c"]).merge(tbl.cell(c["r"] + c["rs"] - 1, c["c"] + c["cs"] - 1))
    by_pos = {(c["r"], c["c"]): c for c in op["cells"]}
    for i in range(nr):
        for j in range(nc):
            cell = tbl.cell(i, j)
            slot = op["slots"][i][j]
            c = by_pos.get((i, j))
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for child in list(tcPr):
                tcPr.remove(child)
            mar = (c or {}).get("mar") or [0, 0, 0, 0]
            for attr, v in zip(("marL", "marT", "marR", "marB"), mar):
                tcPr.set(attr, str(int(round(max(0.0, v) * 9525))))
            tcPr.set("anchor", "t")
            for key, tag in _EDGES:
                _edge(tcPr, tag, slot.get(key))
            if slot.get("fill"):
                sf = tcPr.makeelement(qn("a:solidFill"), {})
                clr = sf.makeelement(qn("a:srgbClr"), {"val": slot["fill"]})
                if (slot.get("alpha") or 1.0) < 0.999:
                    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(round(slot["alpha"] * 100000)))}))
                sf.append(clr)
                tcPr.append(sf)
            else:
                tcPr.append(tcPr.makeelement(qn("a:noFill"), {}))
            tf = cell.text_frame
            if c and c.get("paragraphs"):
                _paragraphs(tf, c, c.get("line_height_px") or 0, end_size=True)
            else:
                # an empty cell must not hold its row open at the 18pt default
                p = tf.paragraphs[0]
                end = p._p.get_or_add_endParaRPr()
                end.set("sz", "100")
                p.line_spacing = Pt(1)


def _icon(slide, op) -> None:
    """A library icon, inserted as its real vector shapes (icon_helper), tinted
    with the sketch's color. An unknown name, or an icon with no drawable
    picture, gets icon_helper's labeled placeholder: never a silent gap."""
    import sys
    scripts = str(Path(__file__).resolve().parent.parent / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    import icon_helper
    args = (E(op["x"]), E(op["y"]), E(max(op["w"], 1)), E(max(op["h"], 1)))
    name = op.get("name") or "?"
    if op.get("placeholder"):
        icon_helper._insert_placeholder(slide, *args, icon_name=name)
        slide.shapes[-1].name = f"icon-{name}"
        return
    icon_helper.insert_icon(name, slide, *args, accent_color="#" + (op.get("color") or "000000"))


_DRAW = {"shape": _shape, "freeform": _freeform, "connector": _connector, "text": _text,
         "icon": _icon, "table": _table}


def build(plan: dict):
    """Draw a plan onto a fresh 1280x720 slide. Returns (prs, slide)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(CANVAS_W), E(CANVAS_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for op in plan["ops"]:
        _DRAW[op["op"]](slide, op)
    return prs, slide


def build_and_save(plan: dict, out_path) -> None:
    prs, _ = build(plan)
    prs.save(str(Path(out_path)))
