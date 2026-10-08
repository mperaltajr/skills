"""Chrome normalization: the takeaway, footnotes and source line in fixed places.

Owner's rule (2026-10-08): on every finished slide Slide Lab places the
takeaway, the footnote(s) and the source line itself, in one fixed position
and size, whatever the designer drew:

  takeaway   directly under the title, at the title's text left edge, in the
             template's own Subtitle-slot size (else 16 pt), not bold, in the
             template's main text color; the same on every slide of the deck
  source     in the bottom band at the template's source position (its
             registered Source slot, else its footer slot, else Slide Lab's
             standard source line), at the title's text left edge, 9 pt, even
             where the template centers its footer box
  footnotes  directly above the source line, stacked upward when there are
             several, same left edge, 9 pt

normalize_chrome() runs per slide in finalize (scripts/finalize_deck.py), so
it holds for every path (direct, sketch, redesign, QC fix), and on the
registration mock slide (scripts/register_template.py). Roles are found from
the template's fields (registered takeaway / source slot ids, SUBTITLE and
FOOTER placeholders) and from shape names (subtitle*, takeaway*, source*,
footnote*).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx.oxml.ns import qn

_SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))
import _chrome_schema as CS  # noqa: E402

_DEFAULT_INSET_EMU = 91440            # PowerPoint's default inner margin, 0.1 in
_EMU_PX = CS.EMU_PER_PX_AT_1280
_TITLE_TYPES = (1, 3)
_SUBTITLE_TYPES = (4,)
_FOOTER_TYPES = (15,)
_SLIDE_NUMBER_TYPES = (13,)


def _ph_type(shape):
    try:
        if not shape.is_placeholder:
            return None
        return int(shape.placeholder_format.type or 0)
    except Exception:
        return None


def _ph_idx(shape):
    try:
        return int(shape.placeholder_format.idx) if shape.is_placeholder else None
    except Exception:
        return None


def _inset(shape, attr: str) -> int:
    """Effective inner margin (EMU): the shape's own, else its layout's and
    master's placeholder, else PowerPoint's default."""
    chain = [shape._element]
    try:
        base = shape._base_placeholder if shape.is_placeholder else None
        while base is not None:
            chain.append(base._element)
            base = getattr(base, "_base_placeholder", None)
    except Exception:
        pass
    for el in chain:
        bp = el.find(".//" + qn("a:bodyPr"))
        if bp is not None and bp.get(attr) is not None:
            try:
                return int(bp.get(attr))
            except ValueError:
                pass
    return _DEFAULT_INSET_EMU


def _text_of(shape) -> str:
    try:
        return (shape.text_frame.text or "").strip() if shape.has_text_frame else ""
    except Exception:
        return ""


def _name(shape) -> str:
    try:
        return (shape.name or "").strip().lower()
    except Exception:
        return ""


def _size_from(el) -> float | None:
    """A text size (pt) written on a placeholder: its list style's first
    level, else any run, end-of-paragraph or default run size."""
    if el is None:
        return None
    lst = el.find(".//" + qn("a:lstStyle"))
    if lst is not None:
        node = lst.find(qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
        if node is not None and node.get("sz"):
            return int(node.get("sz")) / 100
    for tag in ("a:rPr", "a:endParaRPr", "a:defRPr"):
        for node in el.iter(qn(tag)):
            if node.get("sz"):
                return int(node.get("sz")) / 100
    return None


def _slot_size_pt(layout, idx: int) -> float | None:
    """The takeaway slot's own text size on `layout` (placeholder idx), read
    through the layout and then the master, as PowerPoint inherits it."""
    ph = next((p for p in layout.placeholders if _ph_idx(p) == idx), None)
    if ph is None:
        return None
    size = _size_from(ph._element)
    if size:
        return size
    master = layout.slide_master
    t = _ph_type(ph) or 2
    for mp in master.placeholders:
        mt = _ph_type(mp)
        if (t in _TITLE_TYPES and mt in _TITLE_TYPES) or (t not in _TITLE_TYPES and mt == 2):
            size = _size_from(mp._element)
            if size:
                return size
    ts = master._element.find(qn("p:txStyles"))
    if ts is not None:
        node = ts.find(qn("p:bodyStyle") + "/" + qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
        if node is not None and node.get("sz"):
            return int(node.get("sz")) / 100
    return None


def template_takeaway_pt(prs, chrome_spec=None, layout_chrome=None) -> float:
    """The deck's one takeaway size: the template's own Subtitle-slot size
    when a registered layout has a takeaway slot, else 16 pt."""
    from twins.composer import _find_named_layout
    cands = []
    if layout_chrome is not None and getattr(layout_chrome, "subtitle_placeholder_idx", None) is not None:
        cands.append((layout_chrome.name, layout_chrome.subtitle_placeholder_idx))
    for name, lc in (getattr(chrome_spec, "layouts", None) or {}).items():
        if getattr(lc, "subtitle_placeholder_idx", None) is not None:
            cands.append((name, lc.subtitle_placeholder_idx))
    for name, idx in cands:
        lay = _find_named_layout(prs, name)
        if lay is None:
            continue
        size = _slot_size_pt(lay, int(idx))
        if size:
            return float(size)
    return float(CS.CHROME_TAKEAWAY_FONT_PT)


def _title_geometry(slide, layout_chrome) -> tuple[int, int, int]:
    """(text left, text right, box bottom) of the slide's title, in EMU."""
    t_idx = getattr(layout_chrome, "title_placeholder_idx", None) if layout_chrome else None
    title = next((sh for sh in slide.shapes if _ph_type(sh) in _TITLE_TYPES
                  or (t_idx is not None and _ph_idx(sh) == t_idx)), None)
    if title is None:
        title = next((sh for sh in slide.shapes if _name(sh).startswith("title")
                      and getattr(sh, "has_text_frame", False)), None)
    if title is None:
        try:
            title = next((ph for ph in slide.slide_layout.placeholders
                          if _ph_type(ph) in _TITLE_TYPES), None)
        except Exception:
            title = None
    if title is not None and title.left is not None and title.width:
        left, width = int(title.left), int(title.width)
        top, height = int(title.top or 0), int(title.height or 0)
        return (left + _inset(title, "lIns"), left + width - _inset(title, "rIns"),
                top + height)
    tx = getattr(layout_chrome, "title_box_x_px", None) if layout_chrome else None
    tw = getattr(layout_chrome, "title_box_width_px", None) if layout_chrome else None
    ty = getattr(layout_chrome, "title_box_y_px", None) if layout_chrome else None
    th = getattr(layout_chrome, "title_box_height_px", None) if layout_chrome else None
    if None in (tx, tw, ty, th):
        b = CS.canonical_title_box()
        tx, tw, ty, th = b.x_px, b.w_px, b.y_px, b.h_px
    return (tx * _EMU_PX + _DEFAULT_INSET_EMU, (tx + tw) * _EMU_PX - _DEFAULT_INSET_EMU,
            (ty + th) * _EMU_PX)


def source_top_emu(slide, layout_chrome) -> int:
    """The template's source position (EMU from the top of the slide): the
    registered Source slot on the layout, else the layout's footer slot,
    else the layout's recorded source box / Slide Lab's standard one."""
    s_idx = getattr(layout_chrome, "source_placeholder_idx", None) if layout_chrome else None
    try:
        layout_phs = list(slide.slide_layout.placeholders)
    except Exception:
        layout_phs = []
    if s_idx is not None:
        ph = next((p for p in layout_phs if _ph_idx(p) == s_idx), None)
        if ph is not None and ph.top is not None:
            return int(ph.top)
    ph = next((p for p in layout_phs if _ph_type(p) in _FOOTER_TYPES), None)
    if ph is not None and ph.top is not None:
        return int(ph.top)
    box = None
    if layout_chrome is not None:
        try:
            from twins.helpers import _chrome_box_for
            box = _chrome_box_for(layout_chrome, "source")
        except Exception:
            box = None
    box = box or CS.canonical_source_box()
    return int(box.y_px) * _EMU_PX


def _set_runs(shape, *, size_pt=None, bold=None, color_scheme=None,
              clear_italic=False) -> None:
    for p in shape.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("algn", "l")
        for a in ("indent", "marL"):
            if pPr.get(a) is not None:
                del pPr.attrib[a]
        holders = [r._r.get_or_add_rPr() for r in p.runs]
        end = p._p.find(qn("a:endParaRPr"))
        if end is not None:
            holders.append(end)
        for rpr in holders:
            if size_pt is not None:
                rpr.set("sz", str(int(round(size_pt * 100))))
            if bold is not None:
                rpr.set("b", "1" if bold else "0")
            if clear_italic and rpr.get("i") is not None:
                del rpr.attrib["i"]
            if color_scheme:
                for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:pattFill"):
                    for f in rpr.findall(qn(tag)):
                        rpr.remove(f)
                fill = rpr.makeelement(qn("a:solidFill"), {})
                fill.append(fill.makeelement(qn("a:schemeClr"), {"val": color_scheme}))
                # solidFill comes before the font and link children of rPr
                later = {qn(t) for t in ("a:highlight", "a:uLnTx", "a:uLn", "a:uFillTx",
                                          "a:uFill", "a:latin", "a:ea", "a:cs", "a:sym",
                                          "a:hlinkClick", "a:hlinkMouseOver", "a:rtl",
                                          "a:extLst")}
                anchor = next((c for c in rpr if c.tag in later), None)
                if anchor is not None:
                    anchor.addprevious(fill)
                else:
                    rpr.append(fill)


def _set_frame(shape, *, left: int, top: int, width: int, height: int | None) -> None:
    shape.left, shape.top = int(left), int(top)
    shape.width = int(max(width, _EMU_PX))
    if height is not None:
        shape.height = int(height)
    bp = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if bp is not None:
        # no inner margin left, top or bottom: the text itself sits at the
        # placed edge, whether the box is the template's slot or a drawn one
        for ins in ("lIns", "tIns", "bIns"):
            bp.set(ins, "0")
        bp.set("wrap", "square")
        bp.set("anchor", "t")


def _lines_estimate(text: str, width_emu: int, size_pt: float) -> int:
    """Rough line count for a short note line (an average letter is half an em)."""
    width_px = max(1.0, width_emu / _EMU_PX)
    per_line = max(8, int(width_px / (size_pt * 96 / 72 * 0.5)))
    return max(1, sum(max(1, -(-len(para) // per_line)) for para in (text or "").split("\n")))


def normalize_chrome(slide, layout_chrome=None, *, takeaway_pt: float | None = None,
                     on_dark: bool | None = None) -> dict:
    """Place the slide's takeaway, footnotes and source line by the owner's
    rule (module docstring). Returns {"takeaway", "source", "footnotes"}:
    how many of each were placed."""
    done = {"takeaway": 0, "source": 0, "footnotes": 0}
    if layout_chrome is not None and getattr(layout_chrome, "layout_class", "") == "bespoke" \
            and not getattr(layout_chrome, "has_page_number", True):
        return done          # a cover or section divider keeps its own art-directed lines
    if on_dark is None:
        on_dark = getattr(layout_chrome, "text_role", "dark_on_light") == "light_on_dark"
    t_pt = float(takeaway_pt or CS.CHROME_TAKEAWAY_FONT_PT)
    note_pt = float(CS.CHROME_NOTE_FONT_PT)
    text_left, text_right, title_bottom = _title_geometry(slide, layout_chrome)
    width = max(_EMU_PX * 40, text_right - text_left)
    sub_idx = getattr(layout_chrome, "subtitle_placeholder_idx", None) if layout_chrome else None
    src_idx = getattr(layout_chrome, "source_placeholder_idx", None) if layout_chrome else None

    takeaways, sources, notes, page_nums = [], [], [], []
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        t, idx, name = _ph_type(sh), _ph_idx(sh), _name(sh)
        if t in _SLIDE_NUMBER_TYPES or name.startswith("page-number"):
            page_nums.append(sh)
            continue
        if not _text_of(sh):
            continue
        if (sub_idx is not None and idx == sub_idx) or t in _SUBTITLE_TYPES or \
                (t is None and (name.startswith("subtitle") or name.startswith("takeaway"))):
            takeaways.append(sh)
        elif (src_idx is not None and idx == src_idx) or t in _FOOTER_TYPES or \
                (t is None and name.startswith("source")):
            sources.append(sh)
        elif t is None and name.startswith("footnote"):
            notes.append(sh)

    # takeaway: the layout's slot keeps its own height position; a drawn line
    # goes just under the title
    takeaways.sort(key=lambda s: (0 if s.is_placeholder else 1))
    for extra in takeaways[1:]:
        # a designer's second copy of the same takeaway line
        if _text_of(extra) == _text_of(takeaways[0]):
            extra._element.getparent().remove(extra._element)
    for sh in takeaways[:1]:
        if sh.is_placeholder and sh.top is not None:
            top, height = int(sh.top), int(sh.height or CS.CANONICAL_SUBTITLE_H * _EMU_PX)
        else:
            top = title_bottom + CS.CHROME_TAKEAWAY_GAP_PX * _EMU_PX
            height = max(int(sh.height or 0), CS.CANONICAL_SUBTITLE_H * _EMU_PX)
        _set_frame(sh, left=text_left, top=top, width=width, height=height)
        _set_runs(sh, size_pt=t_pt, bold=False, clear_italic=True,
                  color_scheme="bg1" if on_dark else "tx1")
        done["takeaway"] += 1

    # source line: the template slot first; any further source lines stack
    # above it like footnotes
    sources.sort(key=lambda s: (0 if s.is_placeholder else 1, -int(s.top or 0)))
    src_top = source_top_emu(slide, layout_chrome)
    right = text_right
    for pn in page_nums:
        try:
            if pn.left is not None and int(pn.left) > text_left and \
                    abs(int(pn.top or 0) - src_top) < 40 * _EMU_PX:
                right = min(right, int(pn.left) - CS.CHROME_PAGE_NUMBER_GAP_PX * _EMU_PX)
        except Exception:
            pass
    src_w = max(_EMU_PX * 40, right - text_left)
    line_h = int(round(note_pt * 1.2 * 96 / 72 * _EMU_PX))
    stack_top = src_top
    if sources:
        sh = sources[0]
        h = _lines_estimate(_text_of(sh), src_w, note_pt) * line_h
        _set_frame(sh, left=text_left, top=src_top, width=src_w, height=h)
        _set_runs(sh, size_pt=note_pt)
        done["source"] += 1
        notes = sources[1:] + notes

    def _order(s):
        m = re.search(r"(\d+)", _name(s))
        return (int(m.group(1)) if m else 999, int(s.top or 0))
    for sh in sorted(notes, key=_order, reverse=True):
        h = _lines_estimate(_text_of(sh), src_w, note_pt) * line_h
        top = stack_top - CS.CHROME_NOTE_GAP_PX * _EMU_PX - h
        _set_frame(sh, left=text_left, top=top, width=src_w, height=h)
        _set_runs(sh, size_pt=note_pt)
        stack_top = top
        done["footnotes"] += 1
    return done
