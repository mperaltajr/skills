#!/usr/bin/env python3
"""Smoke test: body text keeps the body font when heading and body are one family.

A template can register one family for both heading and body while the
theme's heading face is that family's bold face (heading "Fict Sans Bold",
body "Fict Sans"). Registration strips the style word, so both come out as
"Fict Sans". The theme pass then bound every "Fict Sans" run to the heading
font (+mj-lt, the bold face), and every body box on every slide came out bold
(2026-10-08).

Checks, on a made-up template:
  - the theme's two font names, as Slide Lab reads them, are one family
  - a body run in that family binds to the body font (+mn-lt) and stays
    regular; a run the designer set bold keeps bold
  - two distinct families still bind as before (heading name -> +mj-lt,
    body name and any other font -> +mn-lt)

Run:  py -3 slide-builder/tests/run_body_font_binding_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import io
import re
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from twins.client_theme import _extract_raw_theme, apply_theme_to_shape_xml  # noqa: E402


def _template_with_fonts(path: Path, major: str, minor: str) -> None:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    for part in prs.part.package.iter_parts():
        if "/theme/theme" in str(part.partname).lower():
            xml = part.blob.decode("utf-8")
            xml = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', r"\g<1>" + major, xml)
            xml = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', r"\g<1>" + minor, xml)
            part._blob = xml.encode("utf-8")
    prs.save(str(path))


def _runs(shape):
    out = []
    for r in shape.element.iter(qn("a:r")):
        rpr = r.find(qn("a:rPr"))
        latin = rpr.find(qn("a:latin")) if rpr is not None else None
        out.append((latin.get("typeface") if latin is not None else None,
                    rpr.get("b") if rpr is not None else None))
    return out


def _slide_with(faces: dict) -> tuple:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, (name, (face, bold)) in enumerate(faces.items()):
        tb = s.shapes.add_textbox(Inches(1), Inches(0.5 + i * 0.8), Inches(6), Inches(0.6))
        tb.name = name
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = f"Text {name}"
        run.font.name = face
        run.font.bold = bold
        run.font.size = Pt(14)
    return prs, s


def main() -> int:
    print("[1] one family for heading and body, heading face bold")
    with tempfile.TemporaryDirectory() as td:
        tpl = Path(td) / "one_family.pptx"
        _template_with_fonts(tpl, "Fict Sans Bold", "Fict Sans")
        raw = _extract_raw_theme(tpl)
        major, minor = raw["major_font"], raw["minor_font"]
    assert major == minor == "Fict Sans", (major, minor)
    print(f"    theme fonts as read: heading={major!r} body={minor!r}")
    _prs, s = _slide_with({"body": ("Fict Sans", False), "strong": ("Fict Sans", True),
                           "other": ("Some Other Face", False)})
    for sh in list(s.shapes):
        apply_theme_to_shape_xml(sh.element, {}, major_font=major, minor_font=minor)
    by = {sh.name: _runs(sh)[0] for sh in s.shapes}
    assert by["body"] == ("+mn-lt", "0"), f"body text bound to {by['body']}, wanted body font, regular"
    assert by["strong"] == ("+mn-lt", "1"), f"bold run lost its bold or font: {by['strong']}"
    assert by["other"][0] == "+mn-lt", by["other"]
    print("    ok: body text -> body font, regular; bold stays bold where set")

    print("[2] two distinct families bind as before")
    _prs, s = _slide_with({"head": ("Georgia", False), "body": ("Calibri", False),
                           "other": ("Some Other Face", False)})
    for sh in list(s.shapes):
        apply_theme_to_shape_xml(sh.element, {}, major_font="Georgia", minor_font="Calibri")
    by = {sh.name: _runs(sh)[0][0] for sh in s.shapes}
    assert by == {"head": "+mj-lt", "body": "+mn-lt", "other": "+mn-lt"}, by
    print("    ok: heading name -> heading font; body and other names -> body font")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
