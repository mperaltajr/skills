#!/usr/bin/env python3
"""Smoke test: a QC fix the user approved in chat compiles without the two pages.

  1. before any deck is compiled, apply_qc_fix.py refuses (a first build goes
     through REVIEW.html and FINAL-CHECK.html)
  2. a normal first build compiles through the pages
  3. after a slide is rebuilt with one option, apply_qc_fix.py keeps the other
     picks, records the user's words, finalizes, writes the final-check record
     and compiles (exit 0)
  4. it refuses a fixed slide that has more than one option (the user picks)

Run:  py -3 slide-builder/tests/run_qc_fix_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))
import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402


def _picks_line(out: Path, picks: dict) -> str:
    body = _state.canonical_picks(picks)
    tok = _state.read_state(out)["review"]["token"]
    return f"PICKS {body} CHECK {_state.approval_check(tok, body)}"


def _agent_part_survives() -> None:
    """[5] a QC fix on a sketch slide the translator agent finished does not
    convert it again (2026-10-06: every run re-converted, erased the agent's
    drawing and exited 3 again, for ever)."""
    import os
    from pptx import Presentation
    print("[5] a fix on a slide the translator agent finished compiles, keeping its drawing")
    saved = os.environ.get("SLIDE_LAB_TRANSLATOR")
    os.environ["SLIDE_LAB_TRANSLATOR"] = "script"
    tmp, out = H.new_build(2)
    try:
        sd = out / "slide_02"
        native = sd / "option_A_native.py"
        H.write_option(out, 1, "A")
        H.write_sketch(out, 2, "A")
        assert H.finalize(out).returncode == 0
        assert H.run("build_review.py", "--out", out).returncode == 0
        r = H.run("record_picks.py", "--out", out, "--approved",
                  _picks_line(out, {"slide_01": "A", "slide_02": "A"}))
        assert r.returncode == 0 and native.exists(), r.stdout[-1200:]
        H.agent_finishes(native)
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        assert H.run("build_review.py", "--out", out, "--final").returncode == 0
        tok = _state.read_state(out)["final_check"]["token"]
        assert H.run("compile_picks.py", "--out", out, "--final-token", tok).returncode == 0

        # the QC fix: a new design for slide 2, with a part only the agent draws
        H.write_sketch(out, 2, "A", text="The fixed note.")
        r = H.run("apply_qc_fix.py", "--out", out, "--slides", "2", "--approved", "fix slide 2")
        assert r.returncode == 3 and "option_A_native.py" in r.stdout, r.stdout[-1500:]
        finished = H.agent_finishes(native)
        r = H.run("apply_qc_fix.py", "--out", out, "--slides", "2", "--approved", "fix slide 2")
        assert r.returncode == 0, (
            f"exit {r.returncode}: the finished slide was converted again "
            f"(the loop)\n{r.stdout[-1500:]}")
        assert native.read_text(encoding="utf-8") == finished, \
            "the agent's drawing was rewritten"
        assert "kept as is" in r.stdout, r.stdout[-1500:]
        deck = Presentation(str(out / "final_deck.pptx"))
        names = {sh.name for sh in deck.slides[1].shapes}
        assert "agent-drawn-mark" in names, f"the agent's shape is not in the deck: {names}"
        print("    ok: second run exit 0, the agent's part byte-identical and in the deck")

        print("[6] a design changed after the conversion is converted again")
        H.write_sketch(out, 2, "A", text="A second fix to the note.")
        r = H.run("apply_qc_fix.py", "--out", out, "--slides", "2", "--approved", "fix it again")
        assert r.returncode == 3, r.stdout[-1500:]
        assert "# FALLBACK_PENDING:" in native.read_text(encoding="utf-8")
        assert "The fixed note" not in (sd / "option_A_native.plan.json").read_text(
            encoding="utf-8"), "the plan was not redrawn from the new design"
        print("    ok: re-converted, waiting on the agent for the new design")
    finally:
        if saved is None:
            os.environ.pop("SLIDE_LAB_TRANSLATOR", None)
        else:
            os.environ["SLIDE_LAB_TRANSLATOR"] = saved
        H.cleanup(tmp)


def main() -> int:
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1, "A")
        H.write_option(out, 2, "A")
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-1500:]

        print("[1] refused before any deck is compiled")
        r = H.run("apply_qc_fix.py", "--out", out, "--slides", "2", "--approved", "fix it")
        assert r.returncode == 5 and "no deck has been compiled" in r.stdout, r.stdout
        print("    ok")

        print("[2] the first build compiles through the pages")
        assert H.run("build_review.py", "--out", out).returncode == 0
        r = H.run("record_picks.py", "--out", out, "--approved",
                  _picks_line(out, {"slide_01": "A", "slide_02": "A"}))
        assert r.returncode == 0, r.stdout
        assert H.finalize(out).returncode == 0
        assert H.run("build_review.py", "--out", out, "--final").returncode == 0
        tok = _state.read_state(out)["final_check"]["token"]
        r = H.run("compile_picks.py", "--out", out, "--final-token", tok)
        assert r.returncode == 0, r.stdout[-1500:]
        print("    ok")

        print("[3] a rebuilt slide compiles from the chat approval")
        r = H.run("build_deck.py", "--slide", "2", "--out", out, "--template", H.TEMPLATE,
                  "--pattern", "direct")
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        H.write_option(out, 2, "A")
        assert H.finalize(out, "--slide", "2").returncode == 0

        print("[4] refused when the fixed slide has two options")
        H.write_option(out, 2, "B")
        r = H.run("apply_qc_fix.py", "--out", out, "--slides", "2", "--approved", "fix it")
        assert r.returncode == 5 and "has 2 options" in r.stdout, r.stdout
        for p in (out / "slide_02").glob("option_B*"):
            p.unlink()
        print("    ok")

        r = H.run("apply_qc_fix.py", "--out", out, "--slides", "2",
                  "--approved", "fix the chart on slide 2")
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-800:]
        st = _state.read_state(out)
        assert st["review"]["picks"] == {"slide_01": "A", "slide_02": "A"}, st["review"]
        ov = [o for o in st.get("overrides", []) if o["override"] == "qc_fix_approved_in_chat"]
        assert ov and "fix the chart on slide 2" in ov[-1]["detail"], st.get("overrides")
        assert (out / "final_deck.pptx").exists()
        assert json.loads((out / "picks.json").read_text(encoding="utf-8"))["slide_02"] == "A"
        print("    ok: other picks kept, user's words recorded, deck compiled")
    finally:
        H.cleanup(tmp)
    _agent_part_survives()
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
