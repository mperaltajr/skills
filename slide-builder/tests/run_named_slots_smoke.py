#!/usr/bin/env python3
"""Smoke test: template setup learns text slots NAMED Subtitle / Takeaway / Source.

Many templates make the takeaway line and the source line ordinary text (BODY)
placeholders and say what they are only in the name. Registration counted only
PowerPoint's own SUBTITLE / FOOTER types, so on those templates the takeaway
was drawn as a loose shape and the source line had no home (2026-10-06).

On a COPY of Slide Lab's own test template (the fixture is never changed), one
layout gets a BODY slot named "Subtitle" under the title, a BODY slot named
"Takeaway" at the foot of the page (not the line under the title), and its
footer slot turned into a BODY slot named "Source". Then:

  1. names: slot_role_from_name reads the names (and ignores others)
  2. registration: chrome.yml records subtitle_placeholder_idx = the Subtitle
     slot (not the Takeaway box at the foot), source_placeholder_idx = the
     Source slot, the body starts 12 px below the subtitle slot and ends above
     the source slot; an untouched layout is unchanged
  3. an older chrome.yml without the new fields still loads (fields None)
  4. the registration self-test writes the takeaway and the source line into
     the template's own slots, with no loose copies
  5. the composer fills the source slot by its idx; the build's finishing step
     puts the takeaway in the Subtitle slot and draws no loose 'subtitle' shape
  6. the build's finishing step puts the source line (template field
     'footer') in the Source slot by its idx, with no fallback box and no
     "no footer slot" warning

Run:  py -3 slide-builder/tests/run_named_slots_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import contextlib
import copy
import io
import json
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

import _paths as _p  # noqa: E402
import run_layout_inheritance_smoke as fixture  # noqa: E402
from _chrome_schema import load_chrome_yml, slot_role_from_name, validate_chrome_dict  # noqa: E402

LAYOUT = "body_canonical_light"
OTHER = "body_canonical_dark"
SUB_IDX, TAKEAWAY_IDX, SOURCE_IDX = 13, 14, 20
TITLE = "A test title for the named slots"
SUB = "A test takeaway line that belongs in the template's own slot"
SOURCE = "Source: sample data for the test, not real"
EMU = 9525


def _add_body_slot(layout, name: str, idx: int, x: int, y: int, w: int, h: int) -> None:
    """Clone the layout's content placeholder as a BODY slot with this name/box (px)."""
    src = next(ph for ph in layout.placeholders if int(ph.placeholder_format.type) == 7)
    sp = copy.deepcopy(src._element)
    ph = sp.find(".//" + qn("p:nvPr")).find(qn("p:ph"))
    ph.set("type", "body")
    ph.set("idx", str(idx))
    sp.find(".//" + qn("p:cNvPr")).set("name", name)
    sp.find(".//" + qn("p:cNvPr")).set("id", str(100 + idx))
    _set_box(sp, x * EMU, y * EMU, w * EMU, h * EMU)
    layout.shapes._spTree.append(sp)


def _set_box(sp, x: int, y: int, w: int, h: int) -> None:
    """Give a placeholder its own position (EMU), as a real template's slot has."""
    from lxml import etree
    sppr = sp.find(qn("p:spPr"))
    for old in sppr.findall(qn("a:xfrm")):
        sppr.remove(old)
    xfrm = etree.Element(qn("a:xfrm"))
    sppr.insert(0, xfrm)
    etree.SubElement(xfrm, qn("a:off"), x=str(x), y=str(y))
    etree.SubElement(xfrm, qn("a:ext"), cx=str(w), cy=str(h))


def _make_named_copy(dst: Path) -> None:
    prs = Presentation(str(fixture.FIXTURE_PPTX))
    lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
    title = next(ph for ph in lay.placeholders if int(ph.placeholder_format.type) == 1)
    t_x, t_bottom = int(title.left) // EMU, (int(title.top) + int(title.height)) // EMU
    _add_body_slot(lay, "Subtitle 3", SUB_IDX, t_x, t_bottom + 4, int(title.width) // EMU, 30)
    _add_body_slot(lay, "Takeaway", TAKEAWAY_IDX, t_x, 560, int(title.width) // EMU, 40)
    n = 0
    for ph in lay.placeholders:
        if int(ph.placeholder_format.type) == 15:
            # keep the footer's spot: as a BODY slot it would otherwise inherit
            # the master's body position
            _set_box(ph._element, int(ph.left), int(ph.top), int(ph.width), int(ph.height))
            el = ph._element.find(".//" + qn("p:nvPr")).find(qn("p:ph"))
            el.set("type", "body")
            el.set("idx", str(SOURCE_IDX))
            ph._element.find(".//" + qn("p:cNvPr")).set("name", "Source")
            n += 1
    assert n == 1, f"expected one footer slot on {LAYOUT}, found {n}"
    prs.save(str(dst))


def main() -> int:
    if not fixture.FIXTURE_PPTX.exists():
        fixture.build_fixture(fixture.FIXTURE_PPTX)
    tmp = Path(tempfile.mkdtemp(prefix="slidelab_named_slots_"))
    try:
        print("[1] names")
        assert slot_role_from_name("Subtitle 2") == ("subtitle", 0)
        assert slot_role_from_name("Subheading") == ("subtitle", 0)
        assert slot_role_from_name("Takeaway") == ("subtitle", 2)
        assert slot_role_from_name("Source") == ("source", 0)
        assert slot_role_from_name("Sources") == ("source", 0)
        assert slot_role_from_name("Footnote 4") == ("source", 1)
        for other in ("Content", "Text Placeholder 2", "Title 1", "Resources", ""):
            assert slot_role_from_name(other) is None, other
        print("    ok")

        print("[2] registration records the named slots")
        tpl = tmp / "named_slots.pptx"
        _make_named_copy(tpl)
        import register_template as rt
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            fixture.register_fixture(tpl)
        spec = load_chrome_yml(_p.chrome_yml(tpl))
        lc = spec.layouts[LAYOUT]
        assert lc.subtitle_placeholder_idx == SUB_IDX, lc.subtitle_placeholder_idx
        assert lc.source_placeholder_idx == SOURCE_IDX, lc.source_placeholder_idx
        prs = Presentation(str(tpl))
        lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
        by_idx = {ph.placeholder_format.idx: ph for ph in lay.placeholders}
        sub_bottom = (int(by_idx[SUB_IDX].top) + int(by_idx[SUB_IDX].height)) // EMU
        src_top = int(by_idx[SOURCE_IDX].top) // EMU
        assert lc.body_top_y_px == sub_bottom + 12, (lc.body_top_y_px, sub_bottom)
        assert lc.body_bottom_y_px <= src_top, (lc.body_bottom_y_px, src_top)
        untouched = spec.layouts[OTHER]
        assert untouched.subtitle_placeholder_idx is None
        assert untouched.source_placeholder_idx is None
        print(f"    ok: subtitle idx {lc.subtitle_placeholder_idx} (not the Takeaway box at "
              f"the foot), source idx {lc.source_placeholder_idx}, body {lc.body_top_y_px}-"
              f"{lc.body_bottom_y_px} px")

        print("[3] an older chrome.yml without the new fields still loads")
        import yaml
        raw = yaml.safe_load(_p.chrome_yml(tpl).read_text(encoding="utf-8"))
        for lay_d in raw["layouts"].values():
            for k in ("source_placeholder_idx", "background_hex", "background_kind",
                      "body_left_px", "body_right_px"):
                lay_d.pop(k, None)
        old = validate_chrome_dict(raw).layouts[LAYOUT]
        assert old.source_placeholder_idx is None and old.background_hex is None
        print("    ok")

        print("[4] registration self-test writes both lines into the template's slots")
        tj = _p.theme_json(tpl)
        d = json.loads(tj.read_text(encoding="utf-8"))
        d["default_content_layout"] = LAYOUT
        tj.write_text(json.dumps(d, indent=2), encoding="utf-8")
        fails, infos = rt._render_mock_page_selftest(tpl)
        assert fails == [], fails
        mock = Presentation(str(_p.selftest_pptx(tpl)))
        s = mock.slides[0]
        ph_text = {ph.placeholder_format.idx: ph.text_frame.text for ph in s.placeholders}
        assert "takeaway" in ph_text.get(SUB_IDX, "").lower(), ph_text
        assert "Slide Lab analysis" in ph_text.get(SOURCE_IDX, ""), ph_text
        loose = [sh.name for sh in s.shapes if not sh.is_placeholder
                 and (sh.name or "").lower().startswith(("subtitle", "source"))]
        assert not loose, f"loose copies drawn as well: {loose}"
        print("    ok: takeaway in idx %d, source in idx %d, no loose copies" % (SUB_IDX, SOURCE_IDX))

        print("[5] composer fills the source slot by idx; finishing uses the Subtitle slot")
        from twins.composer import _populate_layout_placeholders, clone_missing_chrome_placeholders
        prs = Presentation(str(tpl))
        lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
        slide = prs.slides.add_slide(lay)
        found = _populate_layout_placeholders(slide, title=TITLE, footer=SOURCE,
                                              title_idx=0, footer_idx=SOURCE_IDX)
        assert found["footer"], found
        got = {ph.placeholder_format.idx: ph.text_frame.text for ph in slide.placeholders}
        assert got.get(SOURCE_IDX) == SOURCE, got

        import finalize_deck as fd
        prs = Presentation(str(tpl))
        lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
        slide = prs.slides.add_slide(lay)
        clone_missing_chrome_placeholders(slide, lay)
        src = Presentation().slides.add_slide(Presentation().slide_layouts[6])
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            fd._apply_body_canonical_finishing(
                slide, prs, lc, src, 3,
                template_fields_override={"title": TITLE, "subtitle": SUB})
        got = {ph.placeholder_format.idx: ph.text_frame.text for ph in slide.placeholders}
        assert got.get(SUB_IDX) == SUB, got
        assert not [sh for sh in slide.shapes if not sh.is_placeholder
                    and (sh.name or "").lower().startswith("subtitle")], "loose subtitle drawn"
        print("    ok")

        print("[6] finishing puts the source line in the Source slot, with no warning")
        prs = Presentation(str(tpl))
        lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
        slide = prs.slides.add_slide(lay)
        clone_missing_chrome_placeholders(slide, lay)
        src = Presentation().slides.add_slide(Presentation().slide_layouts[6])
        said = io.StringIO()
        with contextlib.redirect_stdout(said), contextlib.redirect_stderr(said):
            fd._apply_body_canonical_finishing(
                slide, prs, lc, src, 3,
                template_fields_override={"title": TITLE, "subtitle": SUB, "footer": SOURCE})
        got = {ph.placeholder_format.idx: ph.text_frame.text for ph in slide.placeholders}
        assert got.get(SOURCE_IDX) == SOURCE, got
        assert "no footer slot" not in said.getvalue(), said.getvalue()
        assert not [sh for sh in slide.shapes if not sh.is_placeholder
                    and SOURCE in (sh.text_frame.text if sh.has_text_frame else "")], \
            "the source line was drawn as a loose text box as well"
        print(f"    ok: source line in idx {SOURCE_IDX}, no fallback box, no warning")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
