#!/usr/bin/env python3
"""Smoke test for the copied-page figure ledger.

Two real failures collide in this feature. A supplied one-pager was replaced by
an exec-summary nobody asked for; and a faithful replica shipped the source's
stale numbers ($300B / 6-12 months / 3x) when the brief already had the verified
values. Fixing the first naively causes the second.

The ledger's contract (since 2026-10-13, the lighter check), pinned down here:
  * rows only for figures that are on BOTH the copied page and the picked,
    finished design (a figure the design dropped is not asked about); years,
    dates, slide numbers and prose are never rows;
  * table cells and grouped shapes ARE read (that is where a one-pager keeps
    its numbers);
  * surfaces that cannot be read at all are reported, never silently passed;
  * an unanswered row BLOCKS the compile;
  * the machine never claims the figures match the brief, only that a human has
    or has not answered each one;
  * a PDF page or a picture (no shape tree) gets an honest section: no rows,
    the whole page recorded as checked by eye only, compile proceeds, and
    delivery says so.

Run:  py -3 slide-builder/tests/run_source_ledger_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import _state  # noqa: E402

TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c6360000002000100ffff03000006000557bfabd4000000"
    "0049454e44ae426082")


def _supplied_page(path: Path, png: Path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.6))
    tb.text_frame.text = "Market reaches $300B by 2027"
    note = s.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(6), Inches(0.4))
    note.text_frame.text = "Prepared for the steering committee"   # no figure
    t = s.shapes.add_table(2, 2, Inches(0.5), Inches(1.2), Inches(5), Inches(1)).table
    t.cell(0, 0).text = "Lead time"; t.cell(0, 1).text = "6-12 months"
    t.cell(1, 0).text = "Upside";    t.cell(1, 1).text = "3x"
    png.write_bytes(TINY_PNG)
    s.shapes.add_picture(str(png), Inches(6), Inches(2), Inches(1), Inches(1))
    prs.save(str(path))


def design(out: Path, n: int, texts: list, letter: str = "A") -> Path:
    """A finished option file, as finalize leaves it: slide_NN/option_X.pptx."""
    d = out / f"slide_{n:02d}"
    d.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, t in enumerate(texts):
        s.shapes.add_textbox(Inches(0.5), Inches(0.4 + 0.7 * i), Inches(8),
                             Inches(0.6)).text_frame.text = t
    f = d / f"option_{letter}.pptx"
    prs.save(str(f))
    return f


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="source_ledger_smoke_"))
    try:
        out = tmp / "build"; out.mkdir()
        page = tmp / "one_pager.pptx"
        _supplied_page(page, tmp / "x.png")
        _state.record_prep(out, "h", out)
        _state.record_review(out)
        # The rest of the approval chain, so the ledger is the only thing
        # standing between this build and a compile.
        _state.record_picks(out, {"slide_01": "A"})
        # The picked design kept $300B and 6-12 months, and dropped 3x.
        design(out, 1, ["The market reaches $300B", "Lead time is 6-12 months"])
        tok = _state.record_final_check(out, {})
        env = {**os.environ, "PYTHONPATH": str(SCRIPTS)}

        print("[1] rows only for figures on the copied page AND the picked design")
        r = subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "build",
                            "--out", str(out), "--deck", str(page), "--slide", "1"],
                           capture_output=True, text=True, env=env)
        assert r.returncode == 0, r.stderr[-400:]
        led = json.loads((out / "source_ledger.json").read_text(encoding="utf-8"))
        sec = led["pages"]["slide_01"]
        figs = sorted(row["figure"] for row in sec["rows"])
        assert figs == ["$300B", "6-12"], f"wrong rows: {figs}"
        texts = [row["source_text"] for row in sec["rows"]]
        assert "6-12 months" in texts, f"table cell missed: {texts}"
        assert sec["not_on_design"] == 1, "3x is not on the design and is not asked about"
        assert not any("2027" == row["figure"] for row in sec["rows"]), "a year is not a figure"
        assert sec["fingerprint"].startswith("A:"), sec["fingerprint"]
        print(f"    ok: {figs}; 3x left out (not on the design); the year skipped")

        print("[2] unreadable surfaces are reported, not silently passed")
        assert sec["unreachable"], "the picture should be reported as unreadable"
        assert "picture" in sec["unreachable"][0]["why"]
        print(f"    ok: {sec['unreachable'][0]['why']}")

        print("[3] unanswered rows BLOCK the compile")
        ok, why = _state.check_compile_allowed(out, tok)
        assert not ok and "not answered" in why and "slide 1" in why, why
        print(f"    ok: {why[:110]}...")

        print("[4] replace_with needs an actual replacement")
        for row in sec["rows"]:
            row["resolution"] = "replace_with"      # but no replacement value
        (out / "source_ledger.json").write_text(json.dumps(led), encoding="utf-8")
        subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "status",
                        "--out", str(out)], capture_output=True, text=True, env=env)
        ok, _ = _state.check_compile_allowed(out, tok)
        assert not ok, "an empty replacement must not count as resolved"
        print("    ok: a blank replacement is not an answer")

        print("[5] answering every row releases the compile")
        for row in sec["rows"]:
            if row["figure"] == "6-12":
                row["resolution"] = "replace_with"; row["replacement"] = "sold out"
            else:
                row["resolution"] = "keep_source"
        (out / "source_ledger.json").write_text(json.dumps(led), encoding="utf-8")
        subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "status",
                        "--out", str(out)], capture_output=True, text=True, env=env)
        ok, why = _state.check_compile_allowed(out, tok)
        assert ok, f"compile still refused after every row was answered: {why}"
        print("    ok: compile released once every figure is answered by a human")

        print("[6] the banner never reports a green 'no conflicts'; counts per slide")
        qc = json.loads((out / "brief_qc.json").read_text(encoding="utf-8"))
        assert "1 kept" in qc["summary"] and "unreadable" in qc["summary"], qc["summary"]
        assert any("NOT checked" in w for w in qc["warnings"]), qc["warnings"]
        st = _state.read_state(out)["source_ledger"]
        assert st["keep_source"] == 1 and st["unreachable"] >= 1, st
        assert st["by_slide"]["slide_01"]["kept"] == 1, st
        print(f"    ok: {qc['summary'][:110]}...")

        print("[7] a section made from another design than the pick is refused")
        design(out, 1, ["The market reaches $450B"])          # the design changed
        ok, why = _state.check_compile_allowed(out, tok)
        assert ok, "no pin in _meta.json here: the stale check is for pinned slides"
        (out / "_meta.json").write_text(json.dumps({"slides": [
            {"n": 1, "pinned_source_page": f"{page} slide 1"}]}), encoding="utf-8")
        ok, why = _state.check_compile_allowed(out, tok)
        assert not ok and "older design" in why and "--all-picked" in why, why
        r = subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "build",
                            "--out", str(out), "--all-picked"],
                           capture_output=True, text=True, env=env)
        assert r.returncode == 0, r.stdout + r.stderr
        sec = json.loads((out / "source_ledger.json").read_text(encoding="utf-8"))["pages"]["slide_01"]
        assert sec["rows"] == [] and sec["not_on_design"] == 3, sec
        assert _state.check_compile_allowed(out, tok)[0]
        print("    ok: refused as older than the pick; --all-picked rebuilt it (no figure "
              "came through, so no row; the pin stays)")

        _visual_pages(tmp, env)
        _old_format(tmp)
        print("\nSMOKE PASSED.")
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _fictional_page(path: Path) -> None:
    """A made-up one-page PDF or picture (never a client file)."""
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (1280, 720), "white")
    d = ImageDraw.Draw(img)
    d.text((80, 80), "Sample page: output grows to 12 units by year 3", fill="black")
    d.rectangle((80, 200, 600, 400), outline="black")
    img.save(str(path))


def _visual_pages(tmp: Path, env: dict) -> None:
    import check_done
    print("[8] a PDF or a picture: an honest section, checked by eye only")
    for suffix in (".pdf", ".png"):
        out = tmp / f"build_{suffix[1:]}"; out.mkdir()
        page = tmp / f"sample_page{suffix}"
        _fictional_page(page)
        _state.record_prep(out, "h", out)
        _state.record_review(out)
        _state.record_picks(out, {"slide_01": "A"})
        design(out, 1, ["Output grows to 12 units"])
        tok = _state.record_final_check(out, {})
        r = subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "build",
                            "--out", str(out), "--deck", str(page), "--slide", "1"],
                           capture_output=True, text=True, env=env)
        assert r.returncode == 0, r.stdout[-400:] + r.stderr[-800:]
        assert "CHECKED BY EYE ONLY" in r.stdout, r.stdout
        sec = json.loads((out / "source_ledger.json").read_text(encoding="utf-8"))["pages"]["slide_01"]
        assert sec["rows"] == [], "nothing on a PDF or picture may be claimed as bound"
        assert sec["source_kind"] == "visual", sec
        assert len(sec["unreachable"]) == 1 and "by eye only" in sec["unreachable"][0]["why"]
        st = _state.read_state(out)["source_ledger"]
        assert st["visual_only"] and st["unresolved"] == 0 and st["unreachable"] == 1, st
        ok, why = _state.check_compile_allowed(out, tok)
        assert ok, f"compile still refused for a {suffix} page: {why}"
        qc = json.loads((out / "brief_qc.json").read_text(encoding="utf-8"))
        assert "by eye only" in qc["summary"], qc["summary"]
        line = check_done.supplied_page_line(st)
        assert "CHECKED BY EYE ONLY" in line, line
        subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "status",
                        "--out", str(out)], capture_output=True, text=True, env=env)
        assert _state.read_state(out)["source_ledger"]["visual_only"], "status lost the flag"
        print(f"    ok ({suffix}): {sec['unreachable'][0]['why'][:70]}...")
    r = subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "build",
                        "--out", str(tmp / "build_pdf"), "--deck", str(tmp / "sample_page.pdf"),
                        "--slide", "2", "--target-slide", "1"],
                       capture_output=True, text=True, env=env)
    assert r.returncode == 2 and "out of range" in r.stderr, r.stderr[-400:]
    other = tmp / "page.docx"; other.write_bytes(b"not a page")
    r = subprocess.run([sys.executable, str(SCRIPTS / "source_ledger.py"), "build",
                        "--out", str(tmp / "build_pdf"), "--deck", str(other)],
                       capture_output=True, text=True, env=env)
    assert r.returncode == 2 and "Traceback" not in r.stderr and "PDF" in r.stderr, r.stderr
    print("    ok: a page past the PDF's end and an unsupported file are refused plainly")


def _old_format(tmp: Path) -> None:
    """[9] a one-page ledger from before 2026-10-13 is still read: its open
    rows block, and as a pinned slide it counts as older than the pick."""
    print("[9] the old one-page format is read")
    out = tmp / "build_old"; out.mkdir()
    _state.record_prep(out, "h", out)
    _state.record_review(out)
    _state.record_picks(out, {"slide_04": "A"})
    tok = _state.record_final_check(out, {})
    (out / "source_ledger.json").write_text(json.dumps({
        "source_page": "x.pptx", "source_slide": 2, "generated_at": "2026-10-01T09:00:00",
        "rows": [{"slide": 4, "addr": "s2/sh1", "source_text": "$300B", "resolution": None}],
        "unreachable": []}), encoding="utf-8")
    ok, why = _state.check_compile_allowed(out, tok)
    assert not ok and "slide 4" in why, why
    import source_ledger
    assert list(source_ledger.load_ledger(out)["pages"]) == ["slide_04"]
    print("    ok: read as slide 4's section; its open row blocks")


if __name__ == "__main__":
    sys.exit(main())
