#!/usr/bin/env python3
"""translate_replay.py — replay real past translations through translate_html.py.

Not a smoke test (it needs real session folders and takes a while). This is the
go/no-go harness for replacing the slide-builder-translator agent with the
script: for every option a past session translated with the agent, it

  1. copies the design and the agent's script into a work folder (the session
     folders are only ever read),
  2. translates the design with the script,
  3. builds both into .pptx and checks each opens cleanly (pptx_openability),
  4. compares the text on the script's slide with the text the design shows,
  5. renders design / agent / script side by side, plus a blind copy where the
     agent and script panels are in random order (key kept separately).

Run:
  py -3 tests/translate_replay.py --work <scratch dir> <session dir> [<session dir> ...]
  add --limit N to try a few first.
Writes <work>/replay_summary.json and <work>/sheets/, <work>/blind/.
"""
from __future__ import annotations

import argparse
import json
import random
import re
import shutil
import subprocess
import sys
import traceback
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

import translate_html as T  # noqa: E402


def find_options(session: Path):
    """Every slide_NN/option_X.html that the agent translated in this session."""
    for npy in sorted(session.rglob("slide_*/option_?_native.py")):
        if "_prev" in npy.parts or "_raw" in npy.parts:
            continue
        letter = npy.stem.split("_")[1]
        html = npy.parent / f"option_{letter}.html"
        if html.exists():
            yield html, npy


def _norm_chars(s: str) -> Counter:
    return Counter(c for c in s.lower() if c.isalnum())


def pptx_text(pptx: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(pptx))
    out = []
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            # a native table's text is in its cells
            for row in sh.table.rows:
                for cell in row.cells:
                    out.append(cell.text_frame.text)
    return "\n".join(out)


def run_script(py: Path, timeout=90) -> bool:
    r = subprocess.run([sys.executable, str(py)], cwd=str(py.parent), capture_output=True,
                       text=True, timeout=timeout)
    return r.returncode == 0 and py.with_suffix(".pptx").exists()


def render_png(pptx: Path, outdir: Path):
    from render_slides import render_libre
    outdir.mkdir(parents=True, exist_ok=True)
    try:
        render_libre(pptx, outdir, dpi=96)
    except Exception:
        return None
    pngs = sorted(outdir.glob("*.png"))
    return pngs[0] if pngs else None


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", required=True, type=Path)
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--no-render", action="store_true")
    ap.add_argument("sessions", nargs="+", type=Path)
    args = ap.parse_args(argv)
    work: Path = args.work
    (work / "sheets").mkdir(parents=True, exist_ok=True)
    (work / "blind").mkdir(parents=True, exist_ok=True)

    # Dedupe identical designs (one session folder was copied wholesale).
    # The work folder is named after the option's whole path below the folder
    # given (session folder included): given a client folder ("Flux"), every
    # session has a slide_01/option_A, and naming by the given folder alone
    # made them all one work folder, so only 101 of 263 options were scored.
    seen, jobs, keys = set(), [], set()
    for sess in args.sessions:
        for html, npy in find_options(sess):
            data = html.read_bytes()
            if data in seen:
                continue
            seen.add(data)
            try:
                rel = html.relative_to(sess).with_suffix("")
                parts = [sess.name] + list(rel.parts)
            except ValueError:
                parts = list(html.with_suffix("").parts[-4:])
            key = "__".join(re.sub(r"[^A-Za-z0-9]+", "_", p).strip("_") for p in parts if p)
            base, n = key, 2
            while key.lower() in keys:          # still unique on a case-blind disk
                key, n = f"{base}_{n}", n + 1
            keys.add(key.lower())
            jobs.append((key, html, npy))
    if args.limit:
        jobs = jobs[:args.limit]
    print(f"{len(jobs)} unique translated options")

    from playwright.sync_api import sync_playwright
    from pptx_openability import check_openability
    from pptx import Presentation
    from PIL import Image

    rng = random.Random(20261001)
    results, blind_key = [], {}
    with sync_playwright() as pw:
        from _browser import launch
        browser = launch(pw)
        page = browser.new_page(viewport={"width": 1280, "height": 720})
        for i, (key, html, npy) in enumerate(jobs, 1):
            wd = work / "opts" / key
            res = {"key": key, "html": str(html)}
            try:
                if wd.exists():
                    shutil.rmtree(wd)
                (wd / "agent").mkdir(parents=True)
                (wd / "script").mkdir(parents=True)
                # agent: run its own script beside a copy of its siblings
                for f in html.parent.iterdir():
                    if f.is_file() and f.suffix in (".py", ".html", ".png", ".json", ".css", ".svg"):
                        shutil.copy2(f, wd / "agent" / f.name)
                shutil.copy2(html, wd / "script" / html.name)
                res["agent_built"] = run_script(wd / "agent" / npy.name)

                data = T.extract(page, wd / "script" / html.name, True)
                plan, fields = T.plan_from_extract(data, True)
                rep = T.write_outputs(wd / "script" / html.name, wd / "script",
                                      html.stem.split("_")[1], plan, fields,
                                      data.get("visible_text", ""))
                res["script_built"] = run_script(wd / "script" / npy.name)
                res["fallback"] = rep["fallback"]
                res["warnings"] = [w["code"] for w in rep["warnings"]]
                res["counts"] = rep["counts"]

                for who in ("agent", "script"):
                    p = (wd / who / npy.name).with_suffix(".pptx")
                    if p.exists():
                        res[f"{who}_open_issues"] = [x["issue"][:120] for x in
                                                     check_openability(Presentation(str(p)))]

                # text: everything the design shows (minus the grafted chrome)
                shown = data.get("visible_text", "")
                for k, v in fields.items():
                    shown = shown.replace(v, " ", 1) if v else shown
                want = _norm_chars(shown)
                sp = (wd / "script" / npy.name).with_suffix(".pptx")
                if sp.exists():
                    got = _norm_chars(pptx_text(sp))
                    missing = want - got
                    res["text_chars"] = sum(want.values())
                    res["text_missing"] = sum(missing.values())
                    res["text_missing_sample"] = "".join(sorted(missing.elements()))[:60]

                if not args.no_render and res.get("script_built"):
                    design = wd / "design.png"
                    subprocess.run([sys.executable, str(SKILL / "scripts" / "render_html.py"),
                                    str(wd / "script" / html.name), str(design)],
                                   capture_output=True, timeout=120)
                    panels = {"design": design}
                    for who in ("agent", "script"):
                        p = (wd / who / npy.name).with_suffix(".pptx")
                        if p.exists():
                            png = render_png(p, wd / f"render_{who}")
                            if png:
                                panels[who] = png
                    if all(k in panels for k in ("design", "agent", "script")):
                        ims = {k: Image.open(v).convert("RGB").resize((1280, 720)) for k, v in panels.items()}
                        sheet = Image.new("RGB", (3840, 720), "white")
                        for j, k in enumerate(("design", "agent", "script")):
                            sheet.paste(ims[k], (1280 * j, 0))
                        sheet.resize((1920, 360)).save(work / "sheets" / f"{key}.png")
                        order = ["agent", "script"]
                        rng.shuffle(order)
                        blind = Image.new("RGB", (3840, 720), "white")
                        blind.paste(ims["design"], (0, 0))
                        blind.paste(ims[order[0]], (1280, 0))
                        blind.paste(ims[order[1]], (2560, 0))
                        bname = f"b{len(blind_key) + 1:03d}.png"
                        blind.resize((1920, 360)).save(work / "blind" / bname)
                        blind_key[bname] = {"key": key, "X": order[0], "Y": order[1]}
            except Exception as exc:
                res["error"] = f"{type(exc).__name__}: {exc}"
                res["trace"] = traceback.format_exc()[-800:]
            results.append(res)
            print(f"[{i}/{len(jobs)}] {key}: script_built={res.get('script_built')} "
                  f"fallback={len(res.get('fallback') or [])} "
                  f"missing_text={res.get('text_missing')}/{res.get('text_chars')} "
                  f"{'ERROR ' + res['error'] if 'error' in res else ''}", flush=True)
        browser.close()

    n = len(results)
    built = [r for r in results if r.get("script_built")]
    summary = {
        "options": n,
        "script_built": len(built),
        "errors": sum(1 for r in results if "error" in r),
        "script_open_failures": sum(1 for r in built if r.get("script_open_issues")),
        "agent_open_failures": sum(1 for r in results if r.get("agent_open_issues")),
        "needs_fallback": sum(1 for r in results if r.get("fallback")),
        "text_lost_options": sum(1 for r in built if (r.get("text_missing") or 0) > 0),
        "text_missing_chars": sum(r.get("text_missing") or 0 for r in built),
        "text_total_chars": sum(r.get("text_chars") or 0 for r in built),
        "warning_counts": dict(Counter(w for r in results for w in r.get("warnings", []))),
    }
    (work / "replay_summary.json").write_text(json.dumps({"summary": summary, "results": results}, indent=1),
                                              encoding="utf-8")
    (work / "blind_key.json").write_text(json.dumps(blind_key, indent=1), encoding="utf-8")
    print(json.dumps(summary, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
