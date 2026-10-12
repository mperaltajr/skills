#!/usr/bin/env python3
"""Smoke test: the whole-file checks of each quality-check round block "done"
(owner's decision 2026-10-09).

record_vision_qc.py runs the whole-file hygiene and type-scale check every
round. Its findings used to be printed and recorded, and nothing stopped a
deck with a Major from them being called done.

  1. a Major from the whole-file checks (scratch speaker notes) makes
     check_done refuse, naming the finding and its key
  2. keeping it needs the user's words: empty words and an unknown key are
     refused; with the words recorded check_done passes and lists the kept
     finding with those exact words at delivery
  3. a Critical from them (placeholder text) blocks and cannot be kept
  4. fixed (gone on the next round): nothing blocks; an Advisory alone never
     blocks
  5. a round recorded before this change (no keys stored) is still checked

No renderer is used: tests/_qc_helper.py draws stand-in page images.

Run:  py -3 slide-builder/tests/run_file_checks_block_smoke.py
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import os
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(HERE.parents[1] / "slide-qc" / "scripts"))
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import _state  # noqa: E402
import qc_pages  # noqa: E402
import _qc_helper as Q  # noqa: E402

TITLES = ["Northwind Ferries annual review for the board",
          "Evening crossings lifted ridership by fourteen percent",
          "Fares held flat for commuters through the year"]


def _deck(path: Path, notes: str = "", lorem: bool = False, small: bool = False) -> Path:
    p = Presentation()
    for i, t in enumerate(TITLES):
        s = p.slides.add_slide(p.slide_layouts[0 if i == 0 else 5])
        s.shapes.title.text = t
        if i == 1 and notes:
            s.notes_slide.notes_text_frame.text = notes
        if i == 2 and lorem:
            tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
            tb.text_frame.text = "Lorem ipsum dolor sit amet consectetur"
        if i == 2 and small:
            tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
            tb.text_frame.text = "Crossings ran on time through the winter months"
            for r in tb.text_frame.paragraphs[0].runs:
                r.font.size = Pt(10.5)
    p.save(str(path))
    return path


def _compile(d: Path, deck: Path) -> None:
    K = _state.option_key
    opts = [K(i, "A") for i in range(1, len(TITLES) + 1)]
    if not (_state.read_state(d).get("stages") or {}).get("prep"):
        _state.record_prep(d, "h", d)
    _state.begin_finalize(d)
    _state.record_option_qc(d, {k: {"blocks": 0, "reasons": []} for k in opts})
    _state.end_finalize(d, "ok")
    _state.record_compile(d, kind="picks", output=deck, slides=len(TITLES), options=opts)


def _round(d: Path, deck: Path) -> tuple[int, str]:
    Q.fake_render(d, deck)
    rc, out, _plan = Q.read_round(d)
    return rc, out


def _done(d: Path) -> tuple[int, str]:
    r = subprocess.run([sys.executable, str(SCRIPTS / "check_done.py"), "--out", str(d)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace",
                       env={**os.environ, "PYTHONPATH": str(SCRIPTS)})
    return r.returncode, r.stdout


def _keep(d: Path, *args) -> tuple[int, str]:
    return Q.record(d, *args)


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    tmp = Path(tempfile.mkdtemp(prefix="file_checks_"))

    print("[0] a clean deck is deliverable")
    d0 = tmp / "clean"
    d0.mkdir()
    deck0 = _deck(d0 / "Northwind Review.pptx")
    _compile(d0, deck0)
    rc, out = _round(d0, deck0)
    assert rc == 0, out
    assert not qc_pages.file_check_findings(_state.read_state(d0)), \
        qc_pages.file_check_findings(_state.read_state(d0))
    rc, out = _done(d0)
    assert rc == 0 and "DELIVERABLE" in out, out
    print("    ok")

    print("[1] a Major from the whole-file checks blocks done")
    d = tmp / "major"
    d.mkdir()
    deck = _deck(d / "Northwind Review.pptx", notes="TODO: check this number")
    _compile(d, deck)
    rc, out = _round(d, deck)
    assert rc == 0 and "whole-file checks BLOCK done" in out, out
    fs = qc_pages.file_check_findings(_state.read_state(d))
    assert len(fs) == 1 and fs[0]["severity"] == "Major" and fs[0]["category"] == "speaker-notes", fs
    key = fs[0]["key"]
    assert key.startswith("speaker-notes:s2:") and f"[{key}]" in out, (key, out)
    rc, out = _done(d)
    assert rc == 1 and "whole-file checks" in out and key in out and "--keep-file-check" in out, out
    print(f"    ok: check_done refuses, naming {key}")

    print("[2] kept with the user's exact words")
    rc, out = _keep(d, "--keep-file-check", key, "--user-words", "")
    assert rc == 3 and "exact words" in out, out
    rc, out = _keep(d, "--keep-file-check", "speaker-notes:s9:abcdef", "--user-words", "fine")
    assert rc == 3 and "no Critical or Major" in out, out
    words = "Keep the notes, they are for me when I present on Thursday"
    rc, out = _keep(d, "--keep-file-check", key, "--user-words", words)
    assert rc == 0 and words in out, out
    rc, out = _done(d)
    assert rc == 0 and "DELIVERABLE" in out, out
    assert "kept file checks: 1" in out and words in out and key in out, out
    # The keep stays valid on the next round while the same finding remains.
    rc, out = _round(d, deck)
    assert rc == 0 and "BLOCK done" not in out and "kept by the user" in out, out
    rc, out = _done(d)
    assert rc == 0, out
    print("    ok: empty words and an unknown key refused; kept -> deliverable, listed "
          "with the user's words; the keep carries to the next round")

    print("[3] a Critical blocks and cannot be kept")
    d3 = tmp / "critical"
    d3.mkdir()
    deck3 = _deck(d3 / "Northwind Review.pptx", lorem=True)
    _compile(d3, deck3)
    rc, out = _round(d3, deck3)
    assert rc == 0, out
    crit = [f for f in qc_pages.file_check_findings(_state.read_state(d3))
            if f["severity"] == "Critical"]
    assert crit, qc_pages.file_check_findings(_state.read_state(d3))
    rc, out = _keep(d3, "--keep-file-check", crit[0]["key"], "--user-words", "leave it")
    assert rc == 3 and "cannot be kept" in out, out
    rc, out = _done(d3)
    assert rc == 1 and "Critical and 0 Major still open" in out, out
    print("    ok")

    print("[4] fixed on the next round: nothing blocks; Advisory never blocks")
    deck3b = _deck(d3 / "Northwind Review.pptx", small=True)
    _compile(d3, deck3b)
    rc, out = _round(d3, deck3b)
    assert rc == 0 and "BLOCK done" not in out, out
    hyg = _state.read_state(d3)["vision_deck_look"]["hygiene"]
    assert hyg["advisories"] >= 1 and hyg["criticals"] == 0 and hyg["majors"] == 0, hyg
    rc, out = _done(d3)
    assert rc == 0 and "DELIVERABLE" in out, out
    print(f"    ok: {hyg['advisories']} Advisory, deliverable")

    print("[5] a round recorded before keys were stored is still checked")
    st = _state.read_state(d)
    st.pop("file_checks_kept", None)
    for v in st["vision_deck_look"]["hygiene"]["violations"]:
        v.pop("key", None)
    _state._write(d, st)
    rc, out = _done(d)
    assert rc == 1 and key in out, out
    print("    ok")
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
