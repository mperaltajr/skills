#!/usr/bin/env python3
"""End-to-end smoke: page numbers are real slide-number fields (owner's rule,
2026-10-12: page numbers are never hardcoded).

A fictional 3-slide build goes through the normal pipeline (prep, option
scripts, review + picks, finalize, final check, compile; _e2e_harness.py).
Then:
  1  every page-number box in the compiled deck holds a slide-number field
     (<a:fld type="slidenum">) whose stored number is the slide's position,
     and no typed number
  2  the slides are reordered in the compiled file (python-pptx sldIdLst):
     each page still holds a field (so PowerPoint renumbers it), and
     fix_page_numbers.py then sets the stored numbers to the new positions
     without changing anything else
  3  the hygiene check finds nothing on the compiled deck; on a copy with
     typed numbers it reports a Major tagged page_number_field, which
     qc_fix_policy.py fixes automatically, and fix_page_numbers.py clears it
  4  a placeholder whose template has no field gets a new one (valid id)
  5  when PowerPoint is installed: the reordered deck, opened read-only
     through ppt_safe, shows 1, 2, 3 (the user's open presentations are the
     same before and after). When LibreOffice is installed: its PDF of the
     reordered deck shows 1, 2, 3 too.

Run:  py -3 slide-builder/tests/run_page_number_field_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import json
import re
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
QC = SKILL.parent / "slide-qc" / "scripts"
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(QC))

import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from twins.composer import is_page_number_shape  # noqa: E402

GUID_RE = re.compile(r"^\{[0-9A-F]{8}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{12}\}$", re.I)


def _ok(r, what):
    assert r.returncode == 0, f"{what} failed:\n{r.stdout[-2500:]}\n{r.stderr[-2500:]}"
    return r


def page_numbers(deck: Path) -> list[list[dict]]:
    """Per slide: one record per page-number box: stored text, field count,
    typed bare-number runs, field ids."""
    out = []
    for slide in Presentation(str(deck)).slides:
        rows = []
        for sh in slide.shapes:
            if not (getattr(sh, "has_text_frame", False) and is_page_number_shape(sh)):
                continue
            body = sh.text_frame._txBody
            flds = [f for f in body.iter(qn("a:fld")) if f.get("type") == "slidenum"]
            typed = [r for r in body.iter(qn("a:r"))
                     if re.fullmatch(r"\s*\d+\s*", r.findtext(qn("a:t")) or "")]
            rows.append({"name": sh.name, "text": (sh.text_frame.text or "").strip(),
                         "fields": len(flds), "typed": len(typed),
                         "stored": [f.findtext(qn("a:t")) for f in flds],
                         "ids": [f.get("id") for f in flds]})
        out.append(rows)
    return out


def reorder(src: Path, dst: Path, order: list[int]) -> None:
    prs = Presentation(str(src))
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for i in order:
        lst.append(ids[i])
    prs.save(str(dst))


def non_pn_xml(deck: Path) -> list[bytes]:
    """Each slide's XML (canonical) without its page-number boxes."""
    from lxml import etree
    res = []
    for slide in Presentation(str(deck)).slides:
        root = etree.fromstring(etree.tostring(slide._element))
        for sp in list(root.iter(qn("p:sp"))):
            ph = sp.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}")
            name = sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").get("name", "")
            if (ph is not None and ph.get("type") == "sldNum") or re.match(
                    r"^\s*(page|slide)[\s_-]*(number|num|no)", name, re.I):
                sp.getparent().remove(sp)
        res.append(etree.tostring(root, method="c14n"))
    return res


def hygiene(deck: Path) -> list[dict]:
    r = H.run_path(QC / "check_pptx_hygiene.py", deck)
    assert r.returncode == 0, r.stderr[-1500:]
    return [v for v in json.loads(r.stdout)["violations"] if v["category"] == "page_number_field"]


def main() -> int:
    print("[0] a fictional 3-slide build through the normal pipeline")
    tmp, out = H.new_build(3)
    try:
        for i in (1, 2, 3):
            H.write_option(out, i)
        picks = {f"slide_{i:02d}": "A" for i in (1, 2, 3)}
        _ok(H.run("build_review.py", "--out", out), "review page")
        token = _state.read_state(out)["review"]["token"]
        body = _state.canonical_picks(picks)
        _ok(H.run("record_picks.py", "--out", out, "--approved",
                  f"PICKS {body} CHECK {_state.approval_check(token, body)}"), "picks")
        _ok(H.finalize(out), "finalize")
        r = _ok(H.run("build_review.py", "--out", out, "--final"), "final check")
        tok = _state.read_state(out)["final_check"]["token"]
        _ok(H.run("compile_picks.py", "--out", out, "--final-token", tok), "compile")
        deck = _state.compiled_deck(out)
        print(f"    ok: compiled {deck.name}")

        print("[1] every page-number box is a slidenum field holding its position")
        pns = page_numbers(deck)
        assert len(pns) == 3, pns
        numbered = [i for i, rows in enumerate(pns, 1) if rows]
        assert numbered, f"no slide carries a page number: {pns}"
        for i, rows in enumerate(pns, 1):
            for row in rows:
                assert row["fields"] == 1, f"slide {i}: {row}"
                assert row["typed"] == 0, f"slide {i}: typed number left: {row}"
                assert row["stored"] == [str(i)] and row["text"] == str(i), f"slide {i}: {row}"
                assert all(GUID_RE.match(x or "") for x in row["ids"]), row
        print(f"    ok: slides {numbered} hold a field with the right stored number")

        print("[2] reorder the compiled slides: each page still holds a field")
        moved = tmp / "reordered.pptx"
        reorder(deck, moved, [2, 0, 1])
        after = page_numbers(moved)
        for i, rows in enumerate(after, 1):
            for row in rows:
                assert row["fields"] == 1 and row["typed"] == 0, f"slide {i} after reorder: {row}"
        # the stored numbers are the OLD positions now (3, 1, 2): what a viewer
        # shows below is computed from the field, not read from the stored text
        assert [r[0]["stored"] for r in after if r] == [["3"], ["1"], ["2"]][:len(numbered)] \
            or len(numbered) < 3, after
        fixed = tmp / "reordered_fixed.pptx"
        _ok(H.run("fix_page_numbers.py", moved, "--out", fixed), "fix_page_numbers")
        assert moved.exists(), "the input was touched"
        for i, rows in enumerate(page_numbers(fixed), 1):
            for row in rows:
                assert row["stored"] == [str(i)] and row["fields"] == 1, f"slide {i}: {row}"
        assert non_pn_xml(moved) == non_pn_xml(fixed), "the converter changed other XML"
        print("    ok: fields survive the reorder; the converter re-stores 1, 2, 3 and nothing else")

        print("[3] hygiene: none on the compiled deck; a typed copy is a Major, fixed automatically")
        assert hygiene(deck) == [], hygiene(deck)
        typed = tmp / "typed.pptx"
        prs = Presentation(str(deck))
        for pos, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if getattr(sh, "has_text_frame", False) and is_page_number_shape(sh) \
                        and (sh.text_frame.text or "").strip():
                    sh.text_frame.text = str(pos)        # what the old pipeline wrote
        prs.save(str(typed))
        found = hygiene(typed)
        assert len(found) == len(numbered), found
        assert all(f["severity"] == "Major" and f.get("auto_fix") for f in found), found
        import qc_fix_policy
        d, cat = qc_fix_policy.decide(f"slide {found[0]['slide']} [page_number_field] Major: "
                                      f"{found[0]['issue']}")
        assert (d, cat) == ("auto", "page_number_field"), (d, cat)
        refixed = tmp / "typed_fixed.pptx"
        r = _ok(H.run("fix_page_numbers.py", typed, "--out", refixed), "fix_page_numbers")
        assert "typed number -> field" in r.stdout, r.stdout
        assert hygiene(refixed) == [], "the converter did not clear the finding"
        print(f"    ok: {len(found)} Major(s) -> auto; converter clears them")

        print("[4] a placeholder whose template has no field gets a new one")
        from twins.composer import (_populate_layout_placeholders,
                                    clone_missing_chrome_placeholders)
        p2 = Presentation()
        lay = next(l for l in p2.slide_layouts
                   if any(int(ph.placeholder_format.type) == 13 for ph in l.placeholders))
        s = p2.slides.add_slide(lay)
        clone_missing_chrome_placeholders(s, lay)
        ph = next(p for p in s.placeholders if int(p.placeholder_format.type) == 13)
        for f in list(ph._element.iter(qn("a:fld"))):
            f.getparent().remove(f)                     # a template with no field
        assert not list(ph._element.iter(qn("a:fld")))
        _populate_layout_placeholders(s, page_num="4")
        flds = [f for f in ph._element.iter(qn("a:fld")) if f.get("type") == "slidenum"]
        assert len(flds) == 1 and flds[0].findtext(qn("a:t")) == "4", ph.text_frame.text
        assert GUID_RE.match(flds[0].get("id") or ""), flds[0].get("id")
        print("    ok: field created, id", flds[0].get("id"))

        print("[5] what PowerPoint and LibreOffice show on the reordered deck")
        want = [str(i) if after[i - 1] else None for i in range(1, 4)]
        import render_slides as RS
        if sys.platform == "win32" and RS.powerpoint_available():
            import ppt_safe
            from _ppt_snapshot import snapshot
            before = snapshot()
            assert ppt_safe.check_opens(moved, 3)["ok"], "PowerPoint does not open the deck cleanly"
            shown = []
            with ppt_safe.session() as (app, _):
                with ppt_safe.opened(app, moved) as pres:
                    for i in range(1, pres.Slides.Count + 1):
                        t = None
                        for j in range(1, pres.Slides(i).Shapes.Count + 1):
                            sh = pres.Slides(i).Shapes(j)
                            try:
                                if sh.Type == 14 and sh.PlaceholderFormat.Type == 13:
                                    t = sh.TextFrame.TextRange.Text.strip()
                            except Exception:
                                pass
                        shown.append(t)
            assert snapshot() == before, "the user's PowerPoint changed"
            assert shown == want, f"PowerPoint shows {shown}, want {want}"
            print(f"    ok: PowerPoint shows {shown} (stored numbers were the old positions)")
        else:
            print("    skipped PowerPoint: not installed")
        if RS.libreoffice_available():
            import pypdfium2 as pdfium
            pdf, used = RS.to_pdf(moved, tmp / "lo_pdf", renderer="libreoffice")
            if used == "libreoffice":
                doc = pdfium.PdfDocument(str(pdf))
                got = []
                for i in range(len(doc)):
                    w, h = doc[i].get_size()
                    txt = doc[i].get_textpage().get_text_bounded(
                        left=w * 0.85, bottom=0, right=w, top=h * 0.10)
                    nums = re.findall(r"\d+", txt or "")
                    got.append(nums[-1] if nums else None)
                doc.close()
                assert got == want, f"LibreOffice shows {got}, want {want}"
                print(f"    ok: LibreOffice shows {got}")
            else:
                print(f"    skipped LibreOffice: the PDF came from {used}")
        else:
            print("    skipped LibreOffice: not installed")
    finally:
        H.cleanup(tmp)

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
