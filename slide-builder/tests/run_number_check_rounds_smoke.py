#!/usr/bin/env python3
"""Smoke test: number_check.py tags each round-2 finding with whether its
page changed since round 1 (owner's decision 2026-10-14, speed verdict Q2 D:
round 2 stays on every page; measure where its findings land). Made-up
content.

  1. start --round 1 keeps every page's fingerprint
  2. after one page's text changes and a page is inserted before it, start
     --round 2 and record: a finding on the edited page is "changed", one on
     an untouched page that moved down by one is "unchanged" (page numbers
     ignored) with its round-1 page, a finding with no page is not tagged
  3. the tally and the refusal without a round fingerprint

Run:  py -3 slide-builder/tests/run_number_check_rounds_smoke.py
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402
import contextlib
import io
import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import number_check as N  # noqa: E402


def _deck(path: Path, pages: list[str]) -> None:
    prs = Presentation()
    for i, t in enumerate(pages, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = t
        s.shapes.add_textbox(Inches(9), Inches(7), Inches(0.5), Inches(0.3)).text_frame.text = str(i)
    prs.save(str(path))


def _run(*args) -> tuple[int, str]:
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = N.main(list(args))
    return rc, buf.getvalue()


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        out = Path(td)
        deck = out / "d.pptx"
        _deck(deck, ["Revenue reaches $30.8B by Year 5", "Costs fall 12% a year",
                     "Margins reach 46%"])

        print("[1] round 1")
        rc, o = _run("start", "--out", str(out), "--round", "1", "--deck", str(deck))
        assert rc == 0 and "all 3 page(s)" in o and "Then:" in o, o
        f1 = out / "f1.txt"
        f1.write_text("slide 2: 12% is 11% in the model\n", encoding="utf-8")
        rc, o = _run("record", "--out", str(out), "--round", "1", "--findings", str(f1))
        assert rc == 0 and "round 1: 1 finding(s)" in o, o
        print("    ok")

        print("[2] round 2 after an edit and an insert")
        _deck(deck, ["Revenue reaches $30.8B by Year 5", "A new page about cash",
                     "Costs fall 11% a year", "Margins reach 46%"])
        rc, o = _run("start", "--out", str(out), "--round", "2", "--deck", str(deck))
        assert rc == 0, o
        f2 = out / "f2.json"
        f2.write_text(json.dumps([
            {"page": 3, "issue": "11% needs its source"},
            {"page": 4, "issue": "46% is EBITDA margin, say so"},
            {"issue": "deck-wide: years written two ways"}]), encoding="utf-8")
        rc, o = _run("record", "--out", str(out), "--round", "2", "--findings", str(f2))
        assert rc == 0, o
        rec = json.loads((out / "_qc" / "number_check" / "findings_round_2.json").read_text("utf-8"))
        fs = rec["findings"]
        assert fs[0]["page_changed_since_round_1"] is True, fs[0]
        assert fs[1]["page_changed_since_round_1"] is False and fs[1]["round_1_page"] == 3, fs[1]
        assert fs[2]["page_changed_since_round_1"] is None, fs[2]
        assert rec["tally"] == {"findings": 3, "on_changed_pages": 1, "on_unchanged_pages": 1,
                                "not_tagged": 1}, rec["tally"]
        assert "on an UNCHANGED page: slide 4" in o, o
        print("    ok:", o.splitlines()[0])

        print("[3] tally and refusal")
        rc, o = _run("tally", "--out", str(out))
        assert rc == 0 and "round 1:" in o and "round 2:" in o, o
        rc, o = _run("record", "--out", str(out), "--round", "3", "--findings", str(f2))
        assert rc == 3 and "start --out" in o, o
        print("    ok")
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
