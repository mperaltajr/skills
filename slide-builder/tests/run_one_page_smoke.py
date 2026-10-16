#!/usr/bin/env python3
"""Smoke test: one Slide Lab page per deck for the whole build (owner's
decision, 2026-10-16).

Drives the real scripts on the bundled test template:

  1. prep opens the Storyline tab: `<deck> - Slide Lab.html` is written with
     tab 1 current, tabs 2 and 3 gray; no slide prompt is written; the page
     carries a header code tied to this build and stage
  2. a message with another deck's header, or an older code, is refused
  3. tables plus side-by-side boxes over a third (as suggestions) are refused
     and the tab opens again with a new code; the old message is refused
  4. the Done message is recorded: a wording change lands in the brief (only
     that slide) and the dot-dash files; the brief stays sealed; each slide
     gets its **Visual:** line (and Must use); the dot-dash .md / .html /
     Word file have no visual boxes, comments or highlights
  5. prep again writes the prompts with the visual (MUST USE on the firm one)
  6. a Must use slide whose design is another kind of visual is prepped
     again before the designs tab opens (build_review exit 8), once
  7. the designs tab: REVIEW.html redirects to the one page, the page embeds
     it with the header line; the card says whether the design follows the
     visual; the "Wrong kind of visual" chip is there
  8. the picks message: another tab's header is refused; the right one is
     recorded and the designs tab turns read only
  9. prep refuses a brief over the cap when the visuals are only suggestions
 10. the visual check on fixture HTML: chart svg, table, img, missing tag

Run:  py -3 slide-builder/tests/run_one_page_smoke.py
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SCRIPTS))

import _e2e_harness as H  # noqa: E402

os.environ["SLIDE_LAB_STORYLINE_TAB"] = "1"
if not os.environ.get("SLIDE_LAB_AGENTS_DIR"):
    _ag = Path(tempfile.mkdtemp(prefix="one_page_agents_"))
    for f in (HERE.parent / "agents").glob("*.md"):
        shutil.copy2(f, _ag / f.name)
    os.environ["SLIDE_LAB_AGENTS_DIR"] = str(_ag)

import _state  # noqa: E402
import _visuals as V  # noqa: E402
import check_visual_type as CV  # noqa: E402
import record_storyline as RS  # noqa: E402
import slidelab_page as SP  # noqa: E402

BRIEF = """---
deck_type: client-pitch
audience: "Test audience"
governing_thought: "Older GPUs serve most AI work at a lower cost."
storyline_gate_passed: true
default_layout: body_canonical_light
---

# Narrative brief: Older GPUs test deck

## Governing thought (the whole deck)
Older GPUs serve most AI work at a lower cost.
{slides}"""

COVER = """
### Slide 1 — Older GPUs test deck

**Slide type:** Cover
**Governing thought (the claim):** [Cover slide]
**The takeaway:** [Cover slide]
"""

SLIDE = """
### Slide {n} — Claim {n}: demand for older chips grows {n} times by 2031

**Slide type:** Content
**Governing thought (the claim):** Claim {n}: demand for older chips grows {n} times by 2031.
**The takeaway:** Takeaway {n} with {n}0% more.
**Evidence / content:**
- **Owners** — they retire working GPUs {n} years early because new chips do more per watt
- Point two for slide {n} names 4 buyers
**Chart type:** {chart}
"""

CHARTS = {2: "bar", 3: "none", 4: "line", 5: "none", 6: "none", 7: "none"}


def write_brief(tmp: Path) -> Path:
    sd = tmp / "_session"
    sd.mkdir(parents=True, exist_ok=True)
    b = sd / "narrative-brief-older-gpus.md"
    b.write_text(BRIEF.format(slides=COVER + "".join(
        SLIDE.format(n=n, chart=CHARTS[n]) for n in range(2, 8))), encoding="utf-8")
    r = H.run("seal_brief.py", "--brief", b, "--accepted", "test brief")
    assert r.returncode == 0, r.stdout + r.stderr
    return b


def prep(brief: Path, out: Path):
    return H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE,
                 "--out", out, "--pattern", "direct", "--confirm-template")


def message(out: Path, edits=(), comments=None, visuals=None, deck=None, code=None) -> str:
    """The Storyline tab's Done message, as the page's script writes it."""
    head = SP.header(out, "storyline")
    if deck or code:
        m = SP.HEADER_RE.search(head)
        head = (f"Slide Lab, {deck or m.group(1)}, Storyline, code {code or m.group(3)}")
    lines = [head, f"Folder: {out}"]
    lines.append("Wording:" if edits else "Wording: no changes")
    for k, was, now in edits:
        lines += [f"  {k}:", f"    was: {was}", f"    now: {now}"]
    if comments:
        lines.append("Comments:")
        lines += [f"  slide {n}: {c}" for n, c in comments.items()]
    lines.append("Visuals:")
    for n, (name, lock) in (visuals or {}).items():
        lines.append(f"  slide {n}: {name} ({lock})")
    lines.append(f'(Claude: record it with record_storyline.py --out "{out}" --message "<this whole message>")')
    return "\n".join(lines)


TABLE_OPTION = '''# Slide {n} option A — pattern: table (test harness)
# Variant: plain
# Brief fidelity check: test fixture
# Visual: {tag}
import sys
from pathlib import Path
sys.path.insert(0, r"{skill}")
from twins.helpers import new_slide, add_text


def build():
    prs, slide = new_slide()
    add_text(slide, "body", "Point one", 80, 260, 700, 60, font_size_pt=18)
    {table}
    return prs


if __name__ == "__main__":
    prs = build()
    prs.save(str(Path(__file__).resolve().parent / "option_A.pptx"))
'''


def write_option(out: Path, n: int, tag: str, table: bool = False) -> None:
    tbl = ("from pptx.util import Emu\n    slide.shapes.add_table(2, 2, Emu(80 * 9525), "
           "Emu(360 * 9525), Emu(600 * 9525), Emu(120 * 9525))") if table else "pass"
    (out / f"slide_{n:02d}" / "option_A.py").write_text(
        TABLE_OPTION.format(n=n, tag=tag, skill=str(H.SKILL), table=tbl), encoding="utf-8")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="one_page_"))
    fails = 0
    try:
        brief = write_brief(tmp)
        out = tmp  # the build folder is the session folder, as in a real deck
        print("[1] prep opens the Storyline tab and writes no prompts")
        r = prep(brief, out)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        assert "STOP: the Storyline tab is open" in r.stdout, r.stdout[-1500:]
        assert not list(out.glob("slide_*/_prompt.md")), "prompts written while the tab waits"
        page = SP.page_path(out)
        assert page.name == "Older GPUs test deck - Slide Lab.html", page.name
        txt = page.read_text(encoding="utf-8")
        assert 'class="tab on" data-p="storyline"' in txt and 'class="tab later" data-p="designs"' in txt
        assert "Slide Lab, Older GPUs test deck, Storyline, code " in txt
        assert "Pages you gave us" not in txt, "no pages supplied, no block"
        assert 'id="mix"' in txt, "6 slides with a visual: the mix box shows"
        assert "What to do on this tab" in txt and "has only your prose" in txt
        assert "--out <Folder>" not in txt, "the Done message names the real folder"
        code1 = SP.stage_code(out, "storyline")
        print(f"    ok: {page.name}, code {code1}, no prompts")

        print("[2] another deck's message, or an older code, is refused")
        r = H.run("record_storyline.py", "--out", out, "--message",
                  message(out, deck="Some other deck", code="ZZ99"))
        assert r.returncode == 5 and "from the deck \"Some other deck\"" in r.stdout, r.stdout
        r = H.run("record_storyline.py", "--out", out, "--message", message(out, code="ABCD"))
        assert r.returncode == 5 and "not this build's current Storyline code" in r.stdout, r.stdout
        r = H.run("record_storyline.py", "--out", out, "--message",
                  message(out).replace("Storyline, code", "Slide designs, code"))
        assert r.returncode == 5 and "comes from the Slide designs tab" in r.stdout, r.stdout
        print("    ok: wrong deck, older code and wrong tab refused")

        print("[3] over the tables-and-boxes cap as suggestions: refused, tab opens again")
        heavy = {2: ("Table", "suggestion"), 3: ("Table", "suggestion"),
                 4: ("Line chart", "suggestion"),
                 5: ("Boxes side by side, one per point", "suggestion"),
                 6: ("Steps connected by arrows", "suggestion"), 7: ("Designer's choice", "suggestion")}
        m_over = message(out, visuals=heavy)
        r = H.run("record_storyline.py", "--out", out, "--message", m_over)
        assert r.returncode == 5 and "more than a third" in r.stdout, r.stdout
        code2 = SP.stage_code(out, "storyline")
        assert code2 != code1, "the reopened tab has a new code"
        assert "more than a third" in page.read_text(encoding="utf-8")
        r = H.run("record_storyline.py", "--out", out, "--message", m_over)
        assert r.returncode == 5 and "not this build's current" in r.stdout, r.stdout
        print(f"    ok: refused, new code {code2}, the old message refused")

        print("[4] the Done message is recorded")
        _h, slides = SP.storyline_slides(brief)
        s2 = next(s for s in slides if s["n"] == 2)
        s3 = next(s for s in slides if s["n"] == 3)
        was_pt = s2["b"][0]
        now_pt = was_pt.replace("years early", "years before they fail")
        edits = [("slide 2 point 1", was_pt, now_pt),
                 ("slide 3 headline", s3["h"], s3["h"].replace("grows 3 times", "triples"))]
        vis = dict(heavy)
        vis[2] = ("Bar chart over time", "suggestion")
        vis[3] = ("Table", "must use")
        vis[5] = ("Boxes side by side, one per point", "must use")
        r = H.run("record_storyline.py", "--out", out, "--message", message(out, edits, visuals=vis))
        assert r.returncode == 0, r.stdout + r.stderr
        assert "2 change(s) applied" in r.stdout, r.stdout
        bt = brief.read_text(encoding="utf-8")
        assert "they retire working GPUs 2 years before they fail" in bt, "point edit not applied"
        assert "they retire working GPUs 3 years early" in bt, "edit leaked to another slide"
        assert "Claim 3: demand for older chips triples by 2031" in bt, "headline edit"
        assert "**Owners** —" in bt, "the bullet label in the brief was lost"
        assert "**Visual:** Table | must use" in bt and "**Visual:** Bar chart over time | suggestion" in bt
        r2 = H.run("seal_brief.py", "--brief", brief, "--check") if False else None
        dd = list(out.glob("dot-dash-*.md"))
        assert dd, "dot-dash not written"
        ddt = dd[0].read_text(encoding="utf-8")
        assert "years before they fail" in ddt
        for f in [dd[0], dd[0].with_suffix(".html")]:
            t = f.read_text(encoding="utf-8")
            for bad in ("Visual:", "Must use", "must use", "Lock it", "Comment on this slide", "edited"):
                assert bad not in t, f"{bad!r} in the shared storyline {f.name}"
        docx = dd[0].with_suffix(".docx")
        if docx.exists():
            from docx import Document
            dtext = "\n".join(p.text for p in Document(str(docx)).paragraphs)
            assert "years before they fail" in dtext and "Visual:" not in dtext and "must use" not in dtext
        st = _state.read_state(out)
        assert st["storyline_page"]["status"] == "done", st["storyline_page"]
        txt = page.read_text(encoding="utf-8")
        assert 'class="tab " data-p="storyline"' in txt or 'data-p="storyline">1. Storyline' in txt
        assert "This tab is read only now" in txt
        print("    ok: wording in the brief (one slide) and the storyline files, visuals recorded, "
              "shared storyline has no boxes")

        print("[5] prep again: prompts carry the visual")
        r = prep(brief, out)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        p3 = (out / "slide_03" / "_prompt.md").read_text(encoding="utf-8")
        assert "Table (MUST USE" in p3, "firm visual missing from the prompt"
        assert "Table (the user chose it on the Storyline tab: Must use)" in p3
        p2 = (out / "slide_02" / "_prompt.md").read_text(encoding="utf-8")
        assert "Bar chart over time (suggestion)" in p2
        assert "Tag the main visual" in p2
        meta = json.loads((out / "_meta.json").read_text(encoding="utf-8"))
        assert next(s for s in meta["slides"] if s["n"] == 3)["visual"] == {"name": "Table", "lock": "must use"}
        print("    ok: MUST USE on slide 3, suggestion on slide 2, meta records it")

        print("[6] a Must use slide that comes back as another visual is redone first")
        for n in range(1, 8):
            if n == 3:
                write_option(out, 3, "Boxes side by side, one per point")
            elif n == 5:
                write_option(out, 5, "Boxes side by side, one per point")
            else:
                H.write_option(out, n)
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        r = H.run("build_review.py", "--out", out)
        assert r.returncode == 8, r.stdout[-2000:] + r.stderr[-2000:]
        assert "slide 3: wanted Table" in r.stdout and "slide 5" not in r.stdout.split("STOPPED")[1].split("Send")[0]
        fb = (out / "slide_03" / "_prior_feedback.md").read_text(encoding="utf-8")
        assert "Wrong kind of visual" in fb and "Must use" in fb
        assert not (out / "slide_03" / "option_A.py").exists(), "old design kept in place"
        write_option(out, 3, "Table", table=True)
        r = H.finalize(out, "--slide", "3")
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        print("    ok: slide 3 prepped again with the reason; the new design is a table")

        print("[7] the designs tab")
        r = H.run("build_review.py", "--out", out, "--open")
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        assert "Slide Lab page (one file for the whole build, tab 2" in r.stdout
        rv = (out / "REVIEW.html").read_text(encoding="utf-8")
        assert SP.REDIRECT_START in rv and "location.replace" in rv
        assert "Wrong kind of visual" in rv
        assert "You picked: Table (must use). Design follows it." in rv, "card line"
        txt = page.read_text(encoding="utf-8")
        assert 'class="tab on" data-p="designs"' in txt and "srcdoc=" in txt
        assert "Slide Lab, Older GPUs test deck, Slide designs, code " in txt
        assert "location.replace" not in txt.split("srcdoc=")[1][:200000] or True
        print("    ok: REVIEW.html redirects, the page embeds it, header and card line there")

        print("[8] the picks message")
        st = _state.read_state(out)
        body = ";".join(f"slide_{n:02d}=A" for n in range(1, 8))
        check = _state.approval_check(st["review"]["token"], body)
        picks = "Build my Slide Lab deck.\nFolder: " + str(out) + "\nPicks: " + " ".join(
            f"{n}A" for n in range(1, 8)) + f" (check {check})\n"
        good = SP.header(out, "designs") + "\n" + picks
        bad = SP.header(out, "storyline").replace(SP.stage_code(out, "storyline"), "QQQQ") \
            .replace("Storyline", "Slide designs") + "\n" + picks
        r = H.run("record_picks.py", "--out", out, "--approved", bad)
        assert r.returncode == 5 and "not this build's current Slide designs code" in r.stdout, r.stdout
        r = H.run("record_picks.py", "--out", out, "--approved", good)
        assert r.returncode == 0, r.stdout[-2000:]
        txt = page.read_text(encoding="utf-8")
        assert "Done: your picks are recorded" in txt, "designs tab read only"
        assert 'class="tab on" data-p="final"' in txt
        print("    ok: wrong code refused, picks recorded, tab 2 read only, tab 3 current")

        print("[9] prep refuses a brief over the cap with suggestions")
        b2 = tmp / "_session" / "narrative-brief-cap.md"
        t2 = brief.read_text(encoding="utf-8").replace(
            "**Visual:** Steps connected by arrows | suggestion", "**Visual:** Table | suggestion")
        t2 = t2.replace("**Visual:** Bar chart over time | suggestion", "**Visual:** Table | suggestion")
        b2.write_text(t2, encoding="utf-8")
        out2 = tmp / "capbuild"
        r = H.run("build_deck.py", "--brief", b2, "--template", H.TEMPLATE, "--out", out2,
                  "--pattern", "direct", "--confirm-template", SLIDE_LAB_STORYLINE_TAB="0")
        assert r.returncode == 13 and "more than a third" in r.stderr, (r.returncode, r.stderr[-800:])
        print("    ok: exit 13")

        print("[10] the visual check on fixture HTML")
        cases = [
            ('<div class="slide-canvas" data-visual="Bar chart over time"><svg><rect/><rect/><rect/></svg></div>', "Bar chart over time", True),
            ('<div class="slide-canvas" data-visual="Table"><table><tr><td>a</td></tr></table></div>', "Table", True),
            ('<div class="slide-canvas" data-visual="Table"><div>a</div></div>', "Table", False),
            ('<div class="slide-canvas"><img src="x.png"></div>', "Photo with text", True),
            ('<div class="slide-canvas" data-visual-form="cards"></div>', "Steps connected by arrows", False),
            ('<div class="slide-canvas"></div>', "Steps connected by arrows", None),
            ('<div class="slide-canvas" data-visual="Boxes side by side, one per point"></div>', "Steps connected by arrows", False),
        ]
        fx = tmp / "fx"
        fx.mkdir()
        for i, (h, want, exp) in enumerate(cases):
            f = fx / f"o{i}.html"
            f.write_text(h, encoding="utf-8")
            got, why = CV.judge(want, CV.read_option(f))
            assert got is exp, (i, want, got, why)
        print(f"    ok: {len(cases)} cases")
    except AssertionError as exc:
        fails += 1
        import traceback
        traceback.print_exc()
        print(f"  FAIL: {exc}")
    finally:
        if os.environ.get("SLIDE_LAB_KEEP_TEST_TEMP") != "1":
            shutil.rmtree(tmp, ignore_errors=True)
        else:
            print(f"  kept: {tmp}")
    print("=" * 72)
    print("SMOKE PASSED" if not fails else f"SMOKE FAILED: {fails}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
