#!/usr/bin/env python3
"""Smoke test: the allowed icon list, its one lookup, and icons from sketch to slide.

Before 2026-10-08 a sketch's `<i data-icon-name="gear">` rendered as a blank
gap on the review page, and the translator dropped it from the editable slide
without a word. On 2026-10-21 the owner approved one allowed list of about
240 icons (icons/checked-icons.json), one lookup (scripts/icon_names.py), and
icon rules. This checks:

  0. the list: every name opens exactly one existing file (XML and preview);
     no two names share a file; no two allowed files are the same drawing;
     every topic has icons; an unknown name is refused (the lookup and
     insert_icon both stop); old names map to the right plain name; the
     default color is the template dark text color (TEXT_DARK), not the old
     purple; the prompt's icon block offers 2 to 4 topics plus the status
     group, and only those; build_deck fills it;
  1. the shipped SVG previews cover the library and stay a reasonable size;
  2. render_html draws icon pixels in each icon's box, in the designer's
     color (white on the dark card), and a labeled placeholder plus a
     warning for the unknown name;
  3. the translator plans an icon step per icon (none silently dropped), at
     the right box and color, turns an old name into its plain name, and
     warns about the unknown name; the review card blocks it (Critical);
  4. the finished slide has the real vector icon (a group named icon-<name>,
     alt text "Icon: <name>") at that box, tinted, with unique shape ids, and
     a labeled placeholder for the unknown name;
  5. translate_html.py --emit creates a missing output folder;
  6. slide-qc's icon checks: an icon inside a chart (Major), more than 6 on a
     page, icons on more than a third of the pages, one icon on 3+ pages
     (Advisory).

Run:  py -3 slide-builder/tests/run_icons_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

import icon_names  # noqa: E402
import icon_svg  # noqa: E402

SKETCH = """<!DOCTYPE html><html><head><meta charset="utf-8"><style>
*{box-sizing:border-box}html,body{margin:0}
.slide-canvas{width:1280px;height:720px;position:relative;background:#fff;font-family:Arial}
.ic{width:48px;height:48px}
</style></head><body><div class="slide-canvas" data-visual-form="cards">
<div style="position:absolute;left:100px;top:200px;width:300px;height:200px;border:1px solid #ccc;padding:20px">
  <i class="ic" data-icon-name="gear-refresh" style="color:#1F3A93"></i>
  <p style="margin:10px 0 0;font-size:16px">Known icon</p></div>
<div style="position:absolute;left:480px;top:200px;width:300px;height:200px;background:#1F3A93;padding:20px;color:#FFFFFF">
  <img src="icons/server-86.svg" class="ic" alt="">
  <p style="margin:10px 0 0;font-size:16px">Old name on a dark card</p></div>
<div style="position:absolute;left:860px;top:200px;width:300px;height:200px;border:1px solid #ccc;padding:20px">
  <i class="ic" data-icon-name="no-such-icon" data-icon-color="#C00000"></i>
  <p style="margin:10px 0 0;font-size:16px">Unknown icon</p></div>
</div></body></html>
"""
# Each icon's 48 px box: card left/top + 20 px padding (+1 px border).
# Keyed by the name the slide gets: server-86 is the old name of server-rack.
BOX = {"gear-refresh": (121, 221), "server-rack": (500, 220), "no-such-icon": (881, 221)}


def run(*args):
    env = {**os.environ, "PYTHONIOENCODING": "utf-8", "SLIDE_LAB_NO_OPEN": "1"}
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True,
                          env=env, encoding="utf-8", errors="replace")


def check_list() -> None:
    print("[0] the allowed list and its lookup")
    ics = icon_names.icons()
    assert 230 <= len(ics) <= 250, len(ics)
    names = [ic["name"] for ic in ics]
    files = [ic["file"] for ic in ics]
    assert len(set(names)) == len(names), "two icons share a name"
    assert len(set(files)) == len(files), "two names share a file"
    # every name (and every old name) opens exactly one existing file
    opened: dict[str, set] = {}
    for ic in ics:
        assert (icon_svg.ICONS / f"{ic['file']}.xml").exists(), ic
        assert (icon_svg.SVG_DIR / f"{ic['file']}.svg").exists(), ic
        for key in [ic["name"], ic["file"]] + list(ic.get("old_names") or []):
            opened.setdefault(key, set()).add(ic["file"])
        assert icon_names.file_for(ic["name"]) == ic["file"], ic
        assert icon_svg.status(ic["name"]) == "checked", ic
        assert ic.get("min_px", 24) in (24, 32), ic
        assert ic["shows"].strip() and ic["topic"] in icon_names.topics(), ic
    clash = {k: v for k, v in opened.items() if len(v) > 1}
    assert not clash, f"a name opens more than one file: {clash}"
    # no two allowed files are the same drawing
    seen: dict[str, str] = {}
    for ic in ics:
        h = hashlib.md5((icon_svg.SVG_DIR / f"{ic['file']}.svg").read_bytes()).hexdigest()
        assert h not in seen, f"{ic['name']} is the same drawing as {seen[h]}"
        seen[h] = ic["name"]
    # every topic has icons; the status group is there
    for t in icon_names.topics():
        assert icon_names.by_topic(t), f"topic {t} has no icons"
    st = {ic["name"] for ic in icon_names.by_topic(icon_names.status_topic())}
    assert {"check-mark", "x-circle", "question", "info"} <= st, st
    # owner's decision B: rocket and lightbulb stay, marked use sparingly
    for n in ("rocket", "lightbulb"):
        assert "sparingly" in (icon_names.lookup(n) or {}).get("note", ""), n
    # unknown names are refused
    for bad in ("no-such-icon", "chip", "speech", "certified-75", ""):
        assert icon_names.lookup(bad) is None, bad
        assert icon_svg.status(bad) == "unknown", bad
        try:
            icon_names.resolve(bad, quiet=True)
            raise AssertionError(f"{bad!r} was not refused")
        except icon_names.UnknownIcon as exc:
            assert "checked-icons.json" in str(exc), exc
    # old names map to the right plain name
    for old, new in {"gear": "gear-refresh", "people": "org-chart", "team-35": "org-chart",
                     "mobile-100": "mobile-data", "search-31": "person-search",
                     "search-33": "person-search", "chart-bar": "bar-chart",
                     "bank-12": "gavel", "server-86": "server-rack",
                     "shield-warning": "vault", "globe": "globe"}.items():
        ic = icon_names.lookup(old)
        assert ic and ic["name"] == new, (old, ic)
    assert icon_svg.status("gear") == "old-name"
    assert icon_svg.warning_for("gear")["code"] == "ICON_OLD_NAME"
    assert icon_svg.warning_for("no-such-icon")["code"] == "ICON_UNKNOWN"
    # default color = the template dark text color (TEXT_DARK), not the old purple
    from twins.helpers import TEXT_DARK
    assert icon_names.DEFAULT_COLOR.lstrip("#").upper() == str(TEXT_DARK).upper(), \
        (icon_names.DEFAULT_COLOR, str(TEXT_DARK))
    import icon_helper
    import inspect
    assert "A100FF" not in inspect.getsource(icon_helper).upper()
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from patches import get_blank_layout
    prs = Presentation()
    slide = prs.slides.add_slide(get_blank_layout(prs))
    assert icon_helper.insert_icon("coins", slide, 0, 0, 457200, 457200)
    grp = slide.shapes[-1]
    fills = {c.get("val") for c in grp._element.iter(qn("a:srgbClr"))}
    assert fills == {str(TEXT_DARK).upper()}, fills
    assert grp.name == "icon-coins" and grp._element.xpath("./*[1]/p:cNvPr")[0].get("descr") == "Icon: coins"
    try:
        icon_helper.insert_icon("no-such-icon", slide, 0, 0, 457200, 457200)
        raise AssertionError("insert_icon took an unknown name")
    except icon_names.UnknownIcon:
        pass
    from twins.helpers import add_icon_from_library
    add_icon_from_library(slide, "x", 10, 10, 32, "gear")   # old name still builds
    fills = {c.get("val") for c in slide.shapes[-1]._element.iter(qn("a:srgbClr"))}
    assert fills == {str(TEXT_DARK).upper()}, fills
    # the prompt's icon block: 2 to 4 topics plus status, only those
    text = "GPU leasing revenue; power and cooling cost per MW; servers in the data center"
    picked = icon_names.pick_topics(text)
    assert 2 <= len(picked) <= 4 and "energy" in picked and "finance" in picked, picked
    block = icon_names.topic_block(text)
    offered = set(picked) | {icon_names.status_topic()}
    for ic in ics:
        shown = f"`{ic['name']}`" in block
        assert shown == (ic["topic"] in offered), (ic["name"], ic["topic"], offered)
    assert "`cooling`" in block and "`server-rack`" in block and "`check-mark`" in block
    assert len(icon_names.pick_topics("")) == 2       # nothing matches: two default topics
    assert len(icon_names.pick_topics(text * 5 + " people team customer data strategy time")) == 4
    tpl = (SKILL / "prompt.md").read_text(encoding="utf-8")
    assert "{{ICON_TOPICS}}" in tpl and "6 icons on a page" in tpl
    import build_deck
    seeds = {"content_hash": "0" * 32, "pattern_pick_seed": "1" * 32, "variant_seed_a": "a" * 32,
             "variant_seed_b": "b" * 32, "variant_seed_c": "c" * 32}
    ph = build_deck.build_placeholders(
        {"slide_n": 1, "title": "Cooling costs fall", "governing_thought": "power and cooling per MW"},
        1, "", Path("t.pptx"), Path("out"), seeds, "(none)")
    assert "`cooling`" in ph["ICON_TOPICS"] and "Status" in ph["ICON_TOPICS"], ph["ICON_TOPICS"]
    rendered = build_deck.render_prompt(tpl, ph)
    assert "{{ICON_TOPICS}}" not in rendered and "`cooling`" in rendered
    print(f"    ok: {len(ics)} icons in {len(icon_names.topics())} topics; lookup, old names, "
          f"refusal, default color and prompt block")


def check_qc(tmp: Path) -> None:
    print("[6] slide-qc's icon checks")
    import check_pptx_hygiene as H
    import icon_helper
    from pptx import Presentation
    from pptx.util import Emu
    from patches import get_blank_layout
    prs = Presentation()
    E = 9525   # EMU per px
    for n in range(1, 5):
        slide = prs.slides.add_slide(get_blank_layout(prs))
        icon_helper.insert_icon("coins", slide, 50 * E, 50 * E, 32 * E, 32 * E)
        if n == 1:
            # a drawn chart: bars named chart-bar-..., an icon on top of them
            for k in range(3):
                b = slide.shapes.add_shape(1, Emu((300 + 60 * k) * E), Emu(300 * E),
                                           Emu(40 * E), Emu(200 * E))
                b.name = f"chart-bar-{k}"
            icon_helper.insert_icon("growth-chart", slide, 370 * E, 350 * E, 24 * E, 24 * E)
        if n == 2:
            for k in range(6):
                icon_helper.insert_icon("check-mark", slide, (100 + 40 * k) * E, 600 * E, 24 * E, 24 * E)
    deck = tmp / "icons_qc.pptx"
    prs.save(str(deck))
    viol = H.run_all_checks(deck)["violations"]
    cats = [(v["severity"], v["category"], v["slide"]) for v in viol if v["category"].startswith("icon")]
    assert ("Major", "icon_in_chart", 1) in cats, cats
    assert sum(1 for c in cats if c[1] == "icon_in_chart") == 1, cats   # the coins icon is outside
    assert ("Advisory", "icon_use", 2) in cats, cats                     # 7 icons on page 2
    issues = " | ".join(v["issue"] for v in viol if v["category"] == "icon_use")
    assert "4 of 4 pages" in issues, issues                              # more than a third
    assert "Icon 'coins' is on 4 pages" in issues, issues                # one icon on 3+ pages
    import qc_fix_policy
    assert qc_fix_policy.decide("slide 1 [icon_in_chart] Major: icon inside a chart")[0] == "auto"
    assert qc_fix_policy.decide("slide 2 [icon_use] Advisory: 7 icons on one page")[0] == "none"
    print(f"    ok: {len(cats)} icon findings")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="icons_smoke_"))
    try:
        check_list()

        print("[1] previews ship for the library")
        xml = len(list(icon_svg.ICONS.glob("*.xml")))
        svgs = list(icon_svg.SVG_DIR.glob("*.svg"))
        assert len(svgs) >= xml - 3, (len(svgs), xml)
        size = sum(p.stat().st_size for p in svgs)
        assert size < 4_000_000, f"previews are {size / 1e6:.1f} MB"
        # The preview is the slide's geometry: converting again gives the same file.
        assert icon_svg.xml_to_svg(icon_svg.ICONS / "gear.xml") == icon_svg.svg_for("gear-refresh")
        print(f"    ok: {len(svgs)} previews, {size / 1e6:.1f} MB")

        print("[2] the sketch render draws the icons, in color, and marks the unknown one")
        html = tmp / "option_A.html"
        html.write_text(SKETCH, encoding="utf-8")
        png = tmp / "option_A.sketch.png"
        r = run(SKILL / "scripts" / "render_html.py", html, png)
        assert r.returncode == 0, r.stderr[-800:]
        assert "ICON_UNKNOWN" in r.stderr and "no-such-icon" in r.stderr, r.stderr
        from PIL import Image
        im = Image.open(png).convert("RGB")

        def pixels(name, test):
            x0, y0 = BOX[name]
            return sum(1 for x in range(x0, x0 + 48) for y in range(y0, y0 + 48)
                       if test(im.getpixel((x, y))))
        blue = pixels("gear-refresh", lambda p: p[2] > 100 and p[0] < 90 and p[1] < 110)
        white = pixels("server-rack", lambda p: min(p) > 200)
        grey = pixels("no-such-icon", lambda p: abs(p[0] - p[2]) < 20 and 100 < p[0] < 200)
        assert blue > 80, f"gear box has {blue} icon pixels"
        assert white > 80, f"server box on the dark card has {white} white pixels"
        assert grey > 20, f"unknown icon box has {grey} placeholder pixels"
        print(f"    ok: {blue} / {white} / {grey} icon pixels")

        print("[3] the translator plans every icon; [5] --emit creates its folder")
        emit = tmp / "not" / "yet" / "there"
        r = run(SKILL / "scripts" / "translate_html.py", "--html", html, "--emit", emit)
        assert r.returncode == 0, r.stdout[-800:] + r.stderr[-800:]
        plan = json.loads((emit / "option_A_native.plan.json").read_text(encoding="utf-8"))
        icons = {o["name"]: o for o in plan["ops"] if o["op"] == "icon"}
        assert set(icons) == set(BOX), icons.keys()
        assert icons["gear-refresh"]["color"] == "1F3A93" and icons["server-rack"]["color"] == "FFFFFF"
        assert icons["no-such-icon"]["placeholder"] and not icons["gear-refresh"]["placeholder"]
        for name, (x0, y0) in BOX.items():
            o = icons[name]
            assert x0 - 1 <= o["x"] and o["x"] + o["w"] <= x0 + 49, (name, o)
            assert y0 - 1 <= o["y"] and o["y"] + o["h"] <= y0 + 49, (name, o)
        rep = json.loads((emit / "option_A_translation_report.json").read_text(encoding="utf-8"))
        codes = {w["code"] for w in rep["warnings"]}
        assert {"ICON_UNKNOWN", "ICON_OLD_NAME"} <= codes, rep["warnings"]
        import build_review
        _form, notes = build_review.sketch_facts(html, None)
        sev = {n["code"]: n["severity"] for n in notes}
        assert sev.get("ICON_UNKNOWN") == "Critical" and sev.get("ICON_OLD_NAME") == "Advisory", notes
        print(f"    ok: {len(icons)} icon steps; old name made plain; the unknown name is Critical")

        print("[4] the finished slide has the vector icons at their boxes, tinted")
        r = run(emit / "option_A_native.py")
        assert r.returncode == 0, r.stderr[-800:]
        from pptx import Presentation
        from pptx.oxml.ns import qn
        prs = Presentation(str(emit / "option_A_native.pptx"))
        slide = prs.slides[0]
        by_name = {s.name: s for s in slide.shapes}
        for name in ("gear-refresh", "server-rack"):
            sh = by_name.get(f"icon-{name}")
            assert sh is not None, f"icon-{name} missing: {list(by_name)}"
            o = icons[name]
            assert abs(sh.left / 9525 - o["x"]) < 1 and abs(sh.width / 9525 - o["w"]) < 1, name
            fills = {c.get("val") for c in sh._element.iter(qn("a:srgbClr"))}
            assert ("#" + o["color"]).lstrip("#") in fills, (name, fills)
        ph = by_name.get("icon-no-such-icon")
        assert ph is not None and "[no-such-icon]" in ph.text_frame.text
        ids = [int(e.get("id")) for e in slide.shapes._spTree.iter(qn("p:cNvPr"))]
        assert len(ids) == len(set(ids)), "duplicate shape ids"
        print("    ok")

        check_qc(tmp)
        print("SMOKE PASSED.")
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
