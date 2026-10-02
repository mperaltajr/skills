#!/usr/bin/env python3
"""Smoke test: chat registration (`propose` then `commit-cli`) stores every
answer the user gives in chat.

  - --cover-layout is written to theme.json as cover_layout (it used to be
    asked for and then silently dropped)
  - --reference-slide N writes a reference_slide block to brand.yml
  - --default-content-layout is written as before
  - a cover layout name that is not in the template stops the commit (exit 2)

Run:  py -3 slide-builder/tests/run_commit_cli_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(HERE))

import run_layout_inheritance_smoke as fixture  # noqa: E402
import _paths as _p  # noqa: E402
from pptx import Presentation  # noqa: E402

REG = SCRIPTS / "register_template.py"


def _run(tmp: Path, *args) -> subprocess.CompletedProcess:
    env = dict(os.environ, SLIDE_LAB_REGISTRY=str(tmp / "registry.json"))
    return subprocess.run([sys.executable, str(REG), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8",
                          errors="replace", env=env, timeout=600)


def _template(tmp: Path) -> Path:
    """The layout fixture plus one content slide to point --reference-slide at."""
    tpl = tmp / "chat_reg_template.pptx"
    prs = Presentation(str(fixture.FIXTURE_PPTX))
    prs.slide_width, prs.slide_height = 12192000, 6858000   # 13.333 x 7.5 in
    layout = next(l for m in prs.slide_masters for l in m.slide_layouts
                  if l.name == "body_canonical_light")
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = "A typical content page"
    prs.save(str(tpl))
    return tpl


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="commit_cli_smoke_"))
    try:
        tpl = _template(tmp)
        r = _run(tmp, "propose", tpl)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        assert "commit-cli" in r.stdout and "register.html in the chat" not in r.stdout, \
            "propose still points the user at register.html"
        prop = json.loads(_p.register_proposal_json(tpl).read_text(encoding="utf-8"))
        colors = prop["colors"]
        slots = [s for s in ("dk2", "accent1", "accent2", "lt2") if s in colors]
        assert len(slots) >= 2, colors
        base = ["commit-cli", tpl,
                "--primary-slot", slots[0], "--primary-hex", colors[slots[0]],
                "--accent-slot", slots[1], "--accent-hex", colors[slots[1]],
                "--default-content-layout", "body_canonical_light"]

        print("[1] a cover layout that is not in the template stops the commit")
        r = _run(tmp, *base, "--cover-layout", "No Such Layout")
        assert r.returncode == 2 and "cover_layout" in r.stdout, r.stdout[-1500:]
        print("    ok: exit 2 with the list of real layouts")

        print("[2] cover layout + reference slide are stored")
        r = _run(tmp, *base, "--cover-layout", "cover_light", "--reference-slide", "1")
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-1500:]
        theme = json.loads(_p.theme_json(tpl).read_text(encoding="utf-8"))
        assert theme.get("cover_layout") == "cover_light", theme.get("cover_layout")
        assert theme.get("default_content_layout") == "body_canonical_light"
        brand = _p.brand_yml(tpl).read_text(encoding="utf-8")
        assert "reference_slide:" in brand, "reference slide not written to brand.yml"
        print("    ok: theme.json cover_layout + brand.yml reference_slide")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
