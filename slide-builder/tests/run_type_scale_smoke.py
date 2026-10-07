#!/usr/bin/env python3
"""Smoke test: the type scale (body 12pt default, 10.5pt floor, at most 3
body sizes, exceptions 9pt+).

  1. a slide that keeps to the scale passes: 12/14/16 body, a 40pt hero
     figure, 9pt chart labels named chart-..., a 9pt source line
  2. body text at 10pt is a finding; at 10.5pt it is only a note (below the
     12pt default but allowed)
  3. a fourth body size is flagged
  4. a chart label or source under 9pt is flagged
  5. slide-qc's hygiene check reports them as Major "text size" findings
  6. the FINAL-CHECK.html block lists the slide
  7. Slide Lab's own 9 pt option badge (chrome-option-badge) is skipped: not
     under the floor, not a 4th size. A designer's 9 pt shape named
     chrome-foo is still flagged (exact names only)
  8. a real all-options compile with --badge passes slide-qc's hygiene check
     with no text-size Major

Run:  py -3 slide-builder/tests/run_type_scale_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE.parent / "scripts"))
sys.path.insert(0, str(ROOT / "slide-qc" / "scripts"))
import type_scale  # noqa: E402
import check_pptx_hygiene  # noqa: E402


def _box(slide, name, text, pt, top=1.0):
    tb = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(6), Inches(0.5))
    tb.name = name
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(pt)
    return tb


def _deck(path, slides):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for boxes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for b in boxes:
            _box(s, *b)
    prs.save(str(path))


GOOD = [("heading-1", "Vietnam clears both bars", 16), ("key-1", "$1.2B by 2030", 14),
        ("detail-1", "Partners 4, distribution 4, rules 4", 12), ("hero", "60%", 40),
        ("chart-ylab-200", "200", 9), ("chart-legend-vn", "Vietnam", 9),
        ("src", "Source: illustrative data, not real", 9, 6.9)]


def _all_options_compile() -> None:
    """[8] the real pipeline: two options, all-options deck, --badge."""
    print("[8] an all-options compile with --badge has no text-size Major")
    sys.path.insert(0, str(HERE))
    import _e2e_harness as H
    import _state
    tmp, out = H.new_build(1)
    try:
        H.write_option(out, 1, "A")
        H.write_option(out, 1, "B")
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        assert H.run("build_review.py", "--out", out).returncode == 0
        tok = _state.read_state(out)["review"]["token"]
        line = f"Picks: all options (check {_state.approval_check(tok, 'ALL')})"
        r = H.run("record_picks.py", "--out", out, "--approved", line)
        assert r.returncode == 0, r.stdout[-1200:]
        assert H.finalize(out).returncode == 0
        assert H.run("build_review.py", "--out", out, "--final").returncode == 0
        ftok = _state.read_state(out)["final_check"]["token"]
        r = H.run("compile_picks.py", "--out", out, "--final-token", ftok,
                  "--all-variations", "--badge")
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        deck = out / "final_deck_all_variations.pptx"
        prs = Presentation(str(deck))
        badges = [sh for s in prs.slides for sh in s.shapes
                  if sh.name == "chrome-option-badge"]
        assert len(badges) == 2, f"expected a badge on each option, got {len(badges)}"
        qc = check_pptx_hygiene.run_all_checks(deck)["violations"]
        maj = [v for v in qc if v.get("category") == "text size"
               and v["severity"] == "Major"]
        assert not maj, f"the labeled all-options deck fails its own text-size check: {maj}"
        print(f"    ok: {len(prs.slides)} labeled slides, 0 text-size Majors")
    finally:
        H.cleanup(tmp)


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        deck = Path(td) / "t.pptx"
        _deck(deck, [
            GOOD,
            GOOD + [("cell-1", "Fails on size", 10)],
            GOOD + [("big-1", "Gate 2 at month 12", 18)],
            GOOD + [("chart-xlab-2024", "2024", 8), ("footnote-1", "1. Estimate", 8, 6.95)],
        ])
        res = dict(type_scale.check_pptx(deck))

        print("[1] a slide on the scale passes")
        assert res[1] == [], res[1]
        print("    ok")

        print("[2] body at 10pt is a finding; at 10.5pt only a note")
        assert any("under 10.5 pt" in p for p in res[2]), res[2]
        with tempfile.TemporaryDirectory() as td2:
            d2 = Path(td2) / "n.pptx"
            _deck(d2, [[g for g in GOOD if g[0] != "detail-1"]
                       + [("cell-1", "Fails on size", 10.5)]])
            (_, probs, nts), = type_scale.check_pptx(d2, with_notes=True)
            assert probs == [] and nts and "below the 12 pt default" in nts[0], (probs, nts)
            qc = check_pptx_hygiene.run_all_checks(d2)["violations"]
            sev = {v["severity"] for v in qc if v.get("category") == "text size"}
            assert sev == {"Advisory"}, qc
        print("    ok")

        print("[3] a fourth body size is flagged")
        assert any("4 body text sizes" in p for p in res[3]), res[3]
        print("    ok")

        print("[4] chart label or footnote under 9pt is flagged")
        assert any("under 9 pt" in p for p in res[4]), res[4]
        assert not any("under 10.5 pt" in p for p in res[4]), res[4]
        print("    ok")

        print("[5] slide-qc reports Major text-size findings")
        out = check_pptx_hygiene.run_all_checks(deck)
        ts = [v for v in out["violations"] if v.get("category") == "text size"
              and v["severity"] == "Major"]
        assert ts and all(v["severity"] == "Major" for v in ts), ts
        assert {v["slide"] for v in ts} == {2, 3, 4}, ts
        print("    ok")

        print("[6] the FINAL-CHECK block lists the slide")
        block = type_scale.html_block([("Slide 1", res[1]), ("Slide 2", res[2])])
        assert "Slide 2" in block and "Slide 1" not in block, block
        assert "ok" in type_scale.html_block([("Slide 1", [])])
        print("    ok")

        print("[7] the option badge is skipped; a designer's chrome-foo is not")
        from compile_picks import _stamp_option_badge
        d7 = Path(td) / "badge.pptx"
        _deck(d7, [GOOD, GOOD + [("chrome-foo", "Looks like chrome", 9)]])
        prs = Presentation(str(d7))
        for s in prs.slides:
            assert _stamp_option_badge(s, prs, "Option B"), "badge not placed"
        prs.save(str(d7))
        names = {sh.name for sh in prs.slides[0].shapes}
        assert names & type_scale.PIPELINE_CHROME_NAMES, (
            f"compile_picks stamps a name type_scale does not skip: {names}")
        r7 = dict(type_scale.check_pptx(d7))
        assert r7[1] == [], f"Slide Lab's own badge failed the type scale: {r7[1]}"
        assert any("under 10.5 pt" in p for p in r7[2]), (
            f"a worker shape named chrome-foo dodged the check: {r7[2]}")
        qc = check_pptx_hygiene.run_all_checks(d7)["violations"]
        maj = [v for v in qc if v.get("category") == "text size"
               and v["severity"] == "Major"]
        assert {v["slide"] for v in maj} == {2}, maj
        print("    ok")
    _all_options_compile()
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
