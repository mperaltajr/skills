#!/usr/bin/env python3
"""Smoke test: the sketch converter draws tables as native PowerPoint tables
and lists as one text box with real bullets (owner's decision, 2026-10-08).

Two fictional slides:
  table  header row in <thead>, 6 body rows, a highlighted row, a merged
         cell (rowspan), numbers right-aligned, borders under the rows
  list   a 5-item <ul> with bold lead-ins and one nested item, plus a
         <ul> whose markers are small squares drawn with ::before

Checks:
  - the table is ONE graphicFrame table: 7 rows x 4 columns, the merged
    cell, the highlighted row's fill, right-aligned numbers, borders,
    column widths and row heights as designed
  - each list is ONE text box: one paragraph per item, real bullets
    (buChar), the nested item at level 1, bold lead-ins kept, the square
    markers turned into a bullet character in their own color
  - every character of the design is on the slide; the deck opens
    (pptx_openability); the self-check found nothing; the type-scale check
    sees table text
  - renders close to the sketch in LibreOffice and, on Windows with
    PowerPoint, in PowerPoint too (the safe way: our file only, read-only,
    the user's presentations untouched)
  - a table is ALWAYS a native table (owner's decision, 2026-10-27): one
    with rounded corners is native (corners square, a TABLE_NOTE); one whose
    render differs (planted wrong fills) stays native with a TABLE_NOTE and
    is never handed to the agent; a list that renders differently steps
    down with a warning and keeps every character
  - groups (10/26 deck, page 2, 2026-10-26): cells spanning 2 and 3 rows with a
    black 1px line under each group, light gray 1px lines inside, come out
    native with no note; with one black group line deleted from the native
    table, the self-check still notes it (the table stays native)

Needs Playwright's Chromium and LibreOffice (PowerPoint part skipped
without it, or with SLIDE_LAB_SKIP_POWERPOINT=1).
Run:  py -3 slide-builder/tests/run_native_tables_lists_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import copy
import json
import os
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

import numpy as np  # noqa: E402
from lxml import etree  # noqa: E402
import translate_html as T  # noqa: E402
import translate_alarm as A  # noqa: E402

BASE = """body { margin: 0; }
.slide-canvas { position: relative; width: 1280px; height: 720px; background: #fff;
  font-family: Arial, sans-serif; color: #1a1a1a; overflow: hidden; }
.title { position: absolute; left: 53px; top: 40px; font-size: 32px; font-weight: 700; }
"""

TABLE = """<!doctype html><html><head><meta charset="utf-8"><style>""" + BASE + """
table.t { position: absolute; left: 53px; top: 140px; width: 900px; border-collapse: collapse;
  table-layout: fixed; font-size: 16px; line-height: 22px; }
.t col.c1 { width: 300px; } .t col.c2, .t col.c3, .t col.c4 { width: 200px; }
.t thead th { background: #1f3a5f; color: #fff; font-weight: 700; text-align: left; padding: 10px 12px; }
.t thead th.num { text-align: right; }
.t tbody td { padding: 9px 12px; border-bottom: 1px solid #c8ced6; }
.t td.num { text-align: right; }
.t tr.hl td { background: #fff3cd; font-weight: 700; }
.t td.group { background: #eef2f6; vertical-align: top; color: #1f3a5f; font-weight: 700; }
.note { position: absolute; left: 53px; top: 600px; font-size: 14px; color: #555; }
</style></head><body><div class="slide-canvas">
<div class="title" data-template-field="title">Regional unit volumes by quarter</div>
<table class="t" data-shape-id="volume-table">
<colgroup><col class="c1"><col class="c2"><col class="c3"><col class="c4"></colgroup>
<thead><tr><th>Region</th><th class="num">Q1 units</th><th class="num">Q2 units</th><th class="num">Change</th></tr></thead>
<tbody>
<tr><td class="group" rowspan="2">North and Lakes</td><td class="num">1,240</td><td class="num">1,410</td><td class="num">+13.7%</td></tr>
<tr><td class="num">980</td><td class="num">1,020</td><td class="num">+4.1%</td></tr>
<tr><td>Coastal South</td><td class="num">2,115</td><td class="num">2,060</td><td class="num">-2.6%</td></tr>
<tr class="hl"><td>Central Plains</td><td class="num">3,480</td><td class="num">3,905</td><td class="num">+12.2%</td></tr>
<tr><td>Mountain West</td><td class="num">760</td><td class="num">812</td><td class="num">+6.8%</td></tr>
<tr><td>Island Territories</td><td class="num">305</td><td class="num">298</td><td class="num">-2.3%</td></tr>
</tbody></table>
<div class="note" data-shape-id="note">Fictional figures for a test slide</div>
</div></body></html>"""

LIST = """<!doctype html><html><head><meta charset="utf-8"><style>""" + BASE + """
ul.pts { position: absolute; left: 53px; top: 140px; width: 760px; margin: 0; padding-left: 26px;
  font-size: 18px; line-height: 26px; }
ul.pts > li { margin-bottom: 14px; }
ul.pts li b { color: #1f3a5f; }
ul.pts ul { margin: 8px 0 0 0; padding-left: 24px; font-size: 16px; line-height: 22px; }
ul.sq { position: absolute; left: 860px; top: 140px; width: 360px; margin: 0; padding: 0;
  list-style: none; font-size: 16px; line-height: 22px; }
ul.sq li { position: relative; padding-left: 20px; margin-bottom: 10px; }
ul.sq li::before { content: ''; position: absolute; left: 0; top: 8px; width: 7px; height: 7px;
  background: #c0392b; }
</style></head><body><div class="slide-canvas">
<div class="title" data-template-field="title">Five moves for the next planning cycle</div>
<ul class="pts" data-shape-id="moves">
<li><b>Simplify intake:</b> one request form replaces the four used across the teams today</li>
<li><b>Share the queue:</b> every team sees the same work list, ordered by due date</li>
<li><b>Measure weekly:</b> cycle time and backlog are reviewed every Monday
  <ul><li>Each review takes fifteen minutes and has a single owner</li></ul></li>
<li><b>Fund in stages:</b> money is released when each stage clears its check</li>
<li><b>Close the loop:</b> every finished item is reviewed with the requester</li>
</ul>
<ul class="sq" data-shape-id="checks">
<li>Owners named for every stage</li>
<li>Dates agreed with each team lead</li>
<li>Budget released in three steps</li>
</ul>
</div></body></html>"""

GROUPS = """<!doctype html><html><head><meta charset="utf-8"><style>""" + BASE + """
*, *::before, *::after { box-sizing: border-box; }
table.g { position: absolute; left: 55px; top: 212px; width: 1171px; border-collapse: collapse;
  table-layout: fixed; font-family: "Segoe UI", Arial, sans-serif; font-size: 16px;
  line-height: 1.25; }
.g thead tr { background: #000; }
.g th { color: #fff; font-weight: 700; text-align: left; padding: 0 12px; height: 40px; }
.g td { padding: 4px 12px; height: 50px; border-bottom: 1px solid #D3D1C7; vertical-align: middle; }
.g tr.end td { border-bottom: 1.5px solid #000; }
.g td.grp { font-weight: 700; vertical-align: top; padding-top: 14px; background: #F2F2F2;
  border-bottom: 1.5px solid #000; }
.g td.nm { font-weight: 700; }
</style></head><body><div class="slide-canvas">
<div class="title" data-template-field="title">Six planning tools grouped by purpose</div>
<table class="g" data-shape-id="groups-table">
<colgroup><col style="width:150px"><col style="width:290px"><col style="width:511px"><col style="width:220px"></colgroup>
<thead><tr><th>Group</th><th>Name</th><th>What it is for</th><th>Owner</th></tr></thead>
<tbody>
<tr class="end"><td class="grp">Sizing</td><td class="nm">Market sizing (by region)</td><td>Size of the opportunity in each market*</td><td>Research and Insight</td></tr>
<tr><td class="grp" rowspan="2">Scenarios</td><td class="nm">Business scenarios and ranges</td><td>Tests the plan's most important drivers</td><td>Planning and Finance</td></tr>
<tr class="end"><td class="nm">Regional policy scenarios</td><td>Regional policy scenarios on the market sizing and plan inputs</td><td>Policy and Partners</td></tr>
<tr><td class="grp" rowspan="3">Core tools</td><td class="nm">Long-range financial model</td><td>Value, revenue and returns for the long-range plan</td><td>Planning and Finance</td></tr>
<tr><td class="nm">Operating financial model</td><td>Runs the business: P&amp;L, monthly forecast, impact of operating changes</td><td>Planning and Finance</td></tr>
<tr class="end"><td class="nm">Pricing tool</td><td>Price and margin for each deal</td><td>Planning and Finance</td></tr>
</tbody></table>
</div></body></html>"""

BLANK = """<!doctype html><html><head><meta charset="utf-8"><style>""" + BASE + """
table.b { position: absolute; left: 55px; top: 200px; width: 900px; border-collapse: collapse;
  table-layout: fixed; font-size: 16px; line-height: 24px; }
.b td { padding: 14px 16px; height: 52px; border-bottom: 1px solid #ddd; vertical-align: top; }
.b td.q { background: #f4f4f4; font-size: 18px; }
</style></head><body><div class="slide-canvas">
<table class="b" data-shape-id="bar-table"><colgroup><col style="width:450px"><col style="width:450px"></colgroup>
<tbody>
<tr><td>1&nbsp; Format</td><td class="q" rowspan="3">"Demand may fall"<br><br>"Risk went up"<br><br>"Watch this"</td></tr>
<tr><td>2&nbsp; Timing</td></tr>
<tr><td>3&nbsp; Owner</td></tr>
</tbody></table>
</div></body></html>"""

ROUNDED = TABLE.replace("border-collapse: collapse;", "border-collapse: separate; border-spacing: 0;"
                        " border-radius: 12px; overflow: hidden;").replace(
    ".t thead th {", ".t thead th { border-radius: 12px 12px 0 0;")


def chars(s: str) -> Counter:
    return Counter(c for c in s.lower() if c.isalnum())


def pptx_text(pptx: Path) -> str:
    from pptx import Presentation
    out = []
    for sh in Presentation(str(pptx)).slides[0].shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text_frame.text for r in sh.table.rows for c in r.cells]
    return "\n".join(out)


def build_native(d: Path) -> Path:
    py = d / "option_A_native.py"
    r = subprocess.run([sys.executable, str(py)], cwd=str(d), capture_output=True, text=True,
                       timeout=120)
    assert r.returncode == 0, r.stderr[-1500:]
    return py.with_suffix(".pptx")


def region_diff(design: np.ndarray, native: np.ndarray, box) -> float:
    x0, y0, x1, y1 = (int(v) for v in box)
    return float(np.abs(design[y0:y1, x0:x1] - native[y0:y1, x0:x1]).mean())


def ppt_names():
    """Open presentations in the user's PowerPoint (attach only), or None."""
    try:
        import ppt_safe
        if not ppt_safe.powerpoint_running():
            return None
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        app = win32com.client.GetActiveObject("PowerPoint.Application")
        return sorted(ppt_safe.open_names(app))
    except Exception:
        return None


def main() -> int:
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from pptx_openability import check_openability
    import type_scale as TS

    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        dirs = {}
        for name, html in (("table", TABLE), ("list", LIST), ("rounded", ROUNDED),
                           ("groups", GROUPS)):
            d = td / name
            d.mkdir()
            (d / "option_A.html").write_text(html, encoding="utf-8")
            dirs[name] = d
        names = ("table", "list", "rounded", "groups")
        reports = T.translate_many([(dirs[n] / "option_A.html", dirs[n], "A", True)
                                    for n in names])
        rep = dict(zip(names, reports))

        # ---- table -----------------------------------------------------------
        r = rep["table"]
        assert r["self_check"]["ran"], r["warnings"]
        assert r["counts"]["native_tables"] == 1, r["counts"]
        assert not r["fallback"], r["fallback"]
        assert not r["self_check"]["fired"], r["self_check"]
        plan = json.loads((dirs["table"] / "option_A_native.plan.json").read_text(encoding="utf-8"))
        tops = [o for o in plan["ops"] if o["op"] == "table"]
        assert len(tops) == 1, [o["op"] for o in plan["ops"]]
        tb = tops[0]
        assert len(tb["rows"]) == 7 and len(tb["cols"]) == 4, (len(tb["rows"]), len(tb["cols"]))
        assert [round(w) for w in tb["cols"]] == [300, 200, 200, 200], tb["cols"]
        merged = [c for c in tb["cells"] if c["rs"] == 2]
        assert len(merged) == 1 and merged[0]["c"] == 0, merged
        hl_row = 4                                   # header, 2 merged rows, coastal, central
        assert all(s["fill"] == "FFF3CD" for s in tb["slots"][hl_row]), tb["slots"][hl_row]
        assert tb["slots"][0][0]["fill"] == "1F3A5F"
        assert tb["slots"][3][1]["B"] and tb["slots"][3][1]["B"]["color"] == "C8CED6"
        nums = [c for c in tb["cells"] if c["c"] >= 1 and c["r"] >= 1]
        assert nums and all(c["align"] == "right" for c in nums), [c["align"] for c in nums]
        # no stray shapes or text boxes left for the table's cells
        assert not [o for o in plan["ops"] if o["op"] in ("text", "shape")
                    and 140 <= o.get("y", 0) <= 430 and o.get("x", 0) < 960], \
            [o for o in plan["ops"] if o["op"] != "table"]
        pp = build_native(dirs["table"])
        prs = Presentation(str(pp))
        frames = [s for s in prs.slides[0].shapes if getattr(s, "has_table", False) and s.has_table]
        assert len(frames) == 1
        t = frames[0].table
        assert t.cell(1, 0).is_merge_origin and t.cell(2, 0).is_spanned
        assert t.cell(4, 0).text_frame.paragraphs[0].runs[0].font.bold is True
        assert abs(t.rows[0].height / 9525 - tb["rows"][0]) < 1
        issues = check_openability(prs)
        assert not issues, issues
        design_text = chars("Region Q1 units Q2 units Change North and Lakes 1,240 1,410 +13.7% 980 "
                            "1,020 +4.1% Coastal South 2,115 2,060 -2.6% Central Plains 3,480 3,905 "
                            "+12.2% Mountain West 760 812 +6.8% Island Territories 305 298 -2.3% "
                            "Fictional figures for a test slide")
        missing = design_text - chars(pptx_text(pp))
        assert not missing, missing
        ts = TS.check_slide(prs.slides[0], prs.slide_height)
        assert 12.0 in ts["body_sizes"], ts            # 16px table text = 12pt, now seen
        assert not TS.problems(ts), TS.problems(ts)
        print("  ok: table is one native table (7x4, merged cell, highlighted row, "
              "right-aligned numbers, borders), opens, keeps every character")

        # ---- list ------------------------------------------------------------
        r = rep["list"]
        assert r["counts"]["single_box_lists"] == 2, (r["counts"], r["warnings"])
        assert not r["fallback"], r["fallback"]
        assert not r["self_check"]["fired"], r["self_check"]
        pp_l = build_native(dirs["list"])
        prs_l = Presentation(str(pp_l))
        boxes = {s.name: s for s in prs_l.slides[0].shapes if s.has_text_frame}
        moves = boxes["moves"].text_frame
        assert len(moves.paragraphs) == 6, [p.text for p in moves.paragraphs]
        xml = boxes["moves"]._element.xml
        assert xml.count("<a:buChar") == 6, xml.count("<a:buChar")
        lvls = [p._p.pPr.get("lvl") if p._p.pPr is not None else None for p in moves.paragraphs]
        assert lvls == [None, None, None, "1", None, None], lvls
        leads = [p.runs[0] for p in moves.paragraphs if p.runs and p.runs[0].text.endswith(":")]
        assert len(leads) == 5 and all(run.font.bold for run in leads)
        assert all(not p.runs[1].font.bold for p in moves.paragraphs if len(p.runs) > 1)
        checks = boxes["checks"]._element.xml
        assert checks.count("<a:buChar") == 3 and 'val="C0392B"' in checks, checks[:400]
        assert len([s for s in prs_l.slides[0].shapes if s.has_text_frame]) == 2
        assert not check_openability(prs_l)
        lt = ("Simplify intake: one request form replaces the four used across the teams today "
              "Share the queue: every team sees the same work list, ordered by due date "
              "Measure weekly: cycle time and backlog are reviewed every Monday Each review "
              "takes fifteen minutes and has a single owner Fund in stages: money is released "
              "when each stage clears its check Close the loop: every finished item is reviewed "
              "with the requester Owners named for every stage Dates agreed with each team lead "
              "Budget released in three steps")
        assert not (chars(lt) - chars(pptx_text(pp_l)))
        print("  ok: each list is one text box with real bullets (nested level, bold "
              "lead-ins, square markers as a colored bullet), keeps every character")

        # ---- rendered close to the sketch: LibreOffice and PowerPoint ----------
        with sync_playwright() as pw:
            from _browser import launch
            br = launch(pw)
            page = br.new_page(viewport={"width": 1280, "height": 720})
            data_t = T.extract(page, dirs["table"] / "option_A.html", True)
            data_l = T.extract(page, dirs["list"] / "option_A.html", True)
            br.close()
        design_t, design_l = A.load_rgb(data_t["_shot"]), A.load_rgb(data_l["_shot"])
        engines = ["libreoffice"]
        from render_slides import powerpoint_available
        if powerpoint_available() and os.environ.get("SLIDE_LAB_SKIP_POWERPOINT") != "1":
            engines.append("powerpoint")
        # today's output (separate shapes / one box per item), for comparison:
        # the native version must render at least as close to the sketch
        from twins.html_emit import build as _build
        legacy = []
        for nm, data in (("table", data_t), ("list", data_l)):
            lp, _ = T.plan_from_extract(data, True)
            for gid in list(lp.groups):
                g_ = lp.groups[gid]
                if g_["kind"] == "table":       # one version only: swap in its shapes by hand
                    cur = {id(o) for o in g_["chain"][0]}
                    lp.ops = [o for o in lp.ops if id(o) not in cur] + list(g_["shapes"])
                    lp.ops.sort(key=lambda o: o["_order"])
                    continue
                while g_["cur"] < len(g_["chain"]) - 1:
                    T._step_group(lp, gid, "comparison")
            prs_old, _ = _build(T._public_plan(lp))
            f = td / f"legacy_{nm}.pptx"
            prs_old.save(str(f))
            legacy.append(f)
        for eng in engines:
            before = ppt_names() if eng == "powerpoint" else None
            got_t, got_l, old_t, old_l = A.render_full([pp, pp_l] + legacy, engine=eng)
            after = ppt_names() if eng == "powerpoint" else None
            assert got_t and got_l and old_t and old_l, f"{eng} did not render"
            if eng == "powerpoint":
                assert before == after, "the user's open presentations changed"
            tb_box, l_box = (53, 140, 953, 430), (53, 140, 1230, 420)
            dt = region_diff(design_t, got_t[0], tb_box)
            dl = region_diff(design_l, got_l[0], l_box)
            ot = region_diff(design_t, old_t[0], tb_box)
            ol = region_diff(design_l, old_l[0], l_box)
            # every character is where the renderer put text
            for got, want in ((got_t, design_text), (got_l, chars(lt))):
                pdf_text = chars("".join(c["ch"] for c in got[1]))
                assert not (want - pdf_text), (eng, want - pdf_text)
            # (sizes snap to PowerPoint's list, e.g. 18px -> 14pt, so text is a
            # little wider than the sketch in both versions)
            assert dt <= ot + 0.5 and dl <= ol + 0.5, (eng, dt, ot, dl, ol)
            assert dt < 6.0 and dl < 15.0, (eng, dt, dl)
            print(f"  ok: renders close to the sketch in {eng} (mean pixel difference, "
                  f"native vs separate shapes: table {dt:.2f} vs {ot:.2f}, lists {dl:.2f} "
                  f"vs {ol:.2f} of 255)"
                  + ("; the user's PowerPoint presentations untouched" if before is not None else ""))

        # ---- always a native table ----------------------------------------------
        r = rep["rounded"]
        assert r["counts"]["native_tables"] == 1, (r["counts"], r["warnings"])
        notes = [w["detail"] for w in r["warnings"] if w["code"] == "TABLE_NOTE"]
        assert any("rounded corners are drawn square" in n for n in notes), r["warnings"]
        assert not any(w["code"] in ("TABLE_KEPT_AS_SHAPES", "TABLE_NOT_NATIVE")
                       for w in r["warnings"]), r["warnings"]
        assert not [f for f in r["fallback"] if (f.get("id") or "").startswith("volume-table")], \
            r["fallback"]
        pp_r = build_native(dirs["rounded"])
        prs_r = Presentation(str(pp_r))
        assert sum(1 for s_ in prs_r.slides[0].shapes if getattr(s_, "has_table", False)
                   and s_.has_table) == 1
        assert not (design_text - chars(pptx_text(pp_r)))
        print("  ok: a table with rounded corners is a native table (corners square, with a "
              "note), keeping every character")

        # ---- groups: merged cells with black group lines (10/26 deck, page 2) ----------
        r = rep["groups"]
        assert r["self_check"]["ran"], r["warnings"]
        assert r["counts"]["native_tables"] == 1, r["counts"]
        assert not [w for w in r["warnings"] if w["code"].startswith("TABLE")], r["warnings"]
        assert not r["fallback"], r["fallback"]
        assert not r["self_check"]["fired"], r["self_check"]
        plan_g = json.loads((dirs["groups"] / "option_A_native.plan.json").read_text(encoding="utf-8"))
        tg = next(o for o in plan_g["ops"] if o["op"] == "table")
        assert sorted(c["rs"] for c in tg["cells"] if c["rs"] > 1) == [2, 3], tg["cells"]
        pp_g = build_native(dirs["groups"])
        prs_g = Presentation(str(pp_g))
        tbl_g = next(s_ for s_ in prs_g.slides[0].shapes if getattr(s_, "has_table", False)
                     and s_.has_table).table
        # PowerPoint draws a merged cell's borders from its first cell: the
        # group's black bottom line is written there
        for r0 in (2, 4):
            tcPr = tbl_g.cell(r0, 0)._tc.tcPr
            lnB = tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}lnB")
            assert lnB is not None and 'val="000000"' in etree.tostring(lnB).decode(), r0
        print("  ok: cells spanning 2 and 3 rows with black 1px group lines and light gray "
              "1px lines inside are one native table, no note (merged cell carries its "
              "bottom line)")

        # a deleted group line is still noticed, and the table stays native
        with sync_playwright() as pw:
            from _browser import launch
            br = launch(pw)
            page = br.new_page(viewport={"width": 1280, "height": 720})
            data_g = T.extract(page, dirs["groups"] / "option_A.html", True)
            bdir = td / "blank"
            bdir.mkdir()
            (bdir / "option_A.html").write_text(BLANK, encoding="utf-8")
            data_b = T.extract(page, bdir / "option_A.html", True)
            br.close()
        # a blank line between lines (<br><br>) stays, as an empty paragraph,
        # in a table cell; a no-break space is not collapsed (2026-10-27)
        plan_b, _ = T.plan_from_extract(data_b, True)
        tb_b = next(o for o in plan_b.ops if o["op"] == "table")
        q = next(c for c in tb_b["cells"] if c["rs"] == 3)
        assert ["".join(r["t"] for r in p) for p in q["paragraphs"]] == \
            ['"Demand may fall"', "", '"Risk went up"', "", '"Watch this"'], q["paragraphs"]
        first = "".join(r["t"] for r in tb_b["cells"][0]["paragraphs"][0])
        assert first == "1  Format", repr(first)
        print("  ok: a blank line set with <br><br> in a cell is kept; a no-break space is kept")
        plan_d, _ = T.plan_from_extract(data_g, True)
        top_d = next(o for o in plan_d.ops if o["op"] == "table")
        for s_ in top_d["slots"][3]:                    # planted: the Scenarios group's
            s_["B"] = None                               # black line deleted
        for s_ in top_d["slots"][4]:
            s_["T"] = None
        T.self_check([(plan_d, data_g)])                 # hand_over=True, as in a build
        codes_d = [w["code"] for w in plan_d.warnings]
        assert "TABLE_NOTE" in codes_d, plan_d.warnings
        assert any(o["op"] == "table" for o in plan_d.ops), [o["op"] for o in plan_d.ops]
        assert not [f for f in plan_d.fallbacks if (f.get("id") or "").startswith("groups-table")], \
            plan_d.fallbacks
        assert "TABLE_NOT_NATIVE" not in codes_d and not plan_d.self_check["fired"], codes_d
        print("  ok: one black group line deleted: the self-check notes it, the table stays "
              "native and nothing is handed to the agent")

        # a native table and lists that render differently step down
        plan_t, _ = T.plan_from_extract(data_t, True)
        plan_l, _ = T.plan_from_extract(data_l, True)
        top = next(o for o in plan_t.ops if o["op"] == "table")
        for row in top["slots"]:
            for s in row:
                s["fill"] = "00AA00"                    # planted: wrong fills
        for o in plan_l.ops:
            for pr in o.get("para") or []:
                if pr.get("bullet"):
                    pr["bullet"]["color"] = "00CC00"     # planted: wrong bullet color
                    pr["bullet"]["pct"] = 300
        # hand_over=False: what the step-down leaves on the slide, before any
        # element is handed to the agent (whose part it then is to draw)
        T.self_check([(plan_t, data_t), (plan_l, data_l)], hand_over=False)
        codes_t = [w["code"] for w in plan_t.warnings]
        codes_l = [w["code"] for w in plan_l.warnings]
        # a table whose render can't match stays a native table, with a note
        assert "TABLE_NOTE" in codes_t and "TABLE_NOT_NATIVE" not in codes_t, codes_t
        assert any(o["op"] == "table" for o in plan_t.ops), [o["op"] for o in plan_t.ops]
        assert T.native_counts(plan_t)["native_tables"] == 1
        assert "LIST_NOT_SINGLE_BOX" in codes_l, codes_l          # disc list: straight to items
        assert "LIST_MARKERS_AS_SHAPES" in codes_l, codes_l       # square markers: shapes, one box
        with tempfile.TemporaryDirectory() as t2:
            for nm, pl in (("t", plan_t), ("l", plan_l)):
                from twins.html_emit import build
                prs2, _ = build(T._public_plan(pl))
                f = Path(t2) / f"{nm}.pptx"
                prs2.save(str(f))
                want = design_text if nm == "t" else chars(lt)
                assert not (want - chars(pptx_text(f))), (nm, want - chars(pptx_text(f)))
        # and, as in a build (hand_over=True), the table is never handed to the agent
        plan_t2, _ = T.plan_from_extract(data_t, True)
        for row in next(o for o in plan_t2.ops if o["op"] == "table")["slots"]:
            for s in row:
                s["fill"] = "00AA00"
        T.self_check([(plan_t2, data_t)])
        assert any(o["op"] == "table" for o in plan_t2.ops)
        assert not [f for f in plan_t2.fallbacks if (f.get("id") or "").startswith("volume-table")], \
            plan_t2.fallbacks
        print("  ok: a table whose render can't match stays a native table with a note "
              "(never shapes, never the agent); a list steps down with a warning "
              "(bullets -> marker shapes -> item boxes), no text lost")
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
