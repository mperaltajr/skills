#!/usr/bin/env python3
"""Smoke test: picks and the final check approved in chat (owner's decisions
D1 and D2, 2026-10-06).

Once REVIEW.html or FINAL-CHECK.html is built and opened for the current
files, the user's own words in chat count as that page's approval. They are
recorded verbatim, bound to the page's token and files, and listed by
check_done as chat approvals (not as gates passed over).

  1. picks in chat are refused before REVIEW.html was opened
  2. refused when a slide's options changed after the page was opened
  3. refused when the picks name a letter the user's words do not
  4. accepted with the page open: recorded with the user's words and the
     review token; the paste path (--approved) still works
  5. "build it" is refused before FINAL-CHECK.html was opened, and after a
     re-finalize cleared the page the user saw
  6. refused when a finished file changed after the page was opened
  7. accepted with the page open; --final-token and --approved-in-chat
     together is bad usage
  8. check_done lists both as chat approvals, not as overrides

Run:  py -3 slide-builder/tests/run_chat_approval_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))
import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402


def _chat_picks(out: Path, words: str, picks: str):
    return H.run("record_picks.py", "--out", out, "--approved-in-chat", words,
                 "--picks", picks)


def main() -> int:
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1, "A")
        H.write_option(out, 2, "A")
        assert H.finalize(out).returncode == 0

        print("[1] picks in chat are refused before REVIEW.html was opened")
        assert H.run("build_review.py", "--out", out).returncode == 0
        r = _chat_picks(out, "use A on both slides", "1A 2A")
        assert r.returncode == 5 and "has not been opened" in r.stdout, r.stdout
        assert not (_state.read_state(out)["review"].get("picks")), "picks were recorded"
        print("    ok")

        print("[2] refused when a slide's options changed after the page was opened")
        r = H.run("build_review.py", "--out", out, "--open")
        assert r.returncode == 0, r.stdout[-800:]
        assert _state.read_state(out)["review"]["opened"]["how"] == "test"
        opt = out / "slide_02" / "option_A.py"
        opt.write_text(opt.read_text(encoding="utf-8") + "\n# changed after the page opened\n",
                       encoding="utf-8")
        r = _chat_picks(out, "use A on both slides", "1A 2A")
        assert r.returncode == 5 and "changed after REVIEW.html was opened" in r.stdout, r.stdout
        assert "slide_02" in r.stdout, r.stdout
        print("    ok")

        print("[3] refused when the picks name a letter the user did not")
        assert H.run("build_review.py", "--out", out, "--open").returncode == 0
        r = _chat_picks(out, "looks good, go", "1A 2A")
        assert r.returncode == 5 and "do not mention" in r.stdout, r.stdout
        r = H.run("record_picks.py", "--out", out, "--approved-in-chat", "use A")
        assert r.returncode == 5 and "needs --picks" in r.stdout, r.stdout
        print("    ok")

        print("[4] accepted with the page open; the user's words are on record")
        words = "use A on both slides"
        r = _chat_picks(out, words, "1A 2A")
        assert r.returncode == 0, r.stdout
        assert "from the user's words in chat" in r.stdout, r.stdout
        st = _state.read_state(out)
        assert st["review"]["picks"] == {"slide_01": "A", "slide_02": "A"}, st["review"]
        assert st["review"]["approved_via"] == "chat"
        ca = st["chat_approvals"][-1]
        assert ca["kind"] == "picks" and ca["words"] == words, ca
        assert ca["bound"]["review_token"] == st["review"]["token"], ca
        assert set(ca["bound"]["stamps"]) == {"slide_01", "slide_02"}, ca
        # the paste path is unchanged
        body = "slide_01=A;slide_02=A"
        line = f"Picks: 1A 2A (check {_state.approval_check(st['review']['token'], body)})"
        r = H.run("record_picks.py", "--out", out, "--approved", line)
        assert r.returncode == 0, r.stdout
        assert _state.read_state(out)["review"]["approved_via"] == "page"
        r = H.run("record_picks.py", "--out", out, "--approved", line,
                  "--approved-in-chat", words)
        assert r.returncode == 2, r.stdout
        # chat-approve again, so the record the delivery lists is the chat one
        assert _chat_picks(out, words, "1A 2A").returncode == 0
        print("    ok: recorded with the review token and both slides' stamps; "
              "the pasted line still works")

        print("[5] 'build it' is refused before FINAL-CHECK.html was opened")
        assert H.finalize(out).returncode == 0
        assert H.run("build_review.py", "--out", out, "--final").returncode == 0
        r = H.run("compile_picks.py", "--out", out, "--approved-in-chat", "build it")
        assert r.returncode == 5 and "FINAL-CHECK.html has not been opened" in r.stdout, r.stdout
        assert H.run("build_review.py", "--out", out, "--final", "--open").returncode == 0
        assert _state.read_state(out)["final_check"]["opened"]["how"] == "test"
        assert H.finalize(out).returncode == 0      # clears the page the user saw
        assert H.run("build_review.py", "--out", out, "--final").returncode == 0
        r = H.run("compile_picks.py", "--out", out, "--approved-in-chat", "build it")
        assert r.returncode == 5 and "has not been opened" in r.stdout, r.stdout
        assert not (out / "final_deck.pptx").exists()
        print("    ok: before the page was opened, and after a re-finalize")

        print("[6] refused when a finished file changed after the page was opened")
        assert H.run("build_review.py", "--out", out, "--final", "--open").returncode == 0
        themed = out / "slide_01" / "option_A.pptx"
        original = themed.read_bytes()
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(str(themed))
        prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        prs.save(str(themed))
        r = H.run("compile_picks.py", "--out", out, "--approved-in-chat", "build it")
        assert r.returncode == 5 and "changed after the final" in r.stdout, r.stdout[-800:]
        themed.write_bytes(original)
        assert "chat_approvals" not in _state.read_state(out) or not [
            a for a in _state.read_state(out)["chat_approvals"] if a["kind"] == "final_check"]
        print("    ok")

        print("[7] accepted with the page open for the current files")
        tok = _state.read_state(out)["final_check"]["token"]
        r = H.run("compile_picks.py", "--out", out, "--approved-in-chat", "build it",
                  "--final-token", tok)
        assert r.returncode == 2, r.stdout
        r = H.run("compile_picks.py", "--out", out, "--approved-in-chat", "build it")
        assert r.returncode == 0, r.stdout[-1500:]
        deck = out / "final_deck.pptx"
        assert deck.exists()
        st = _state.read_state(out)
        fc = [a for a in st["chat_approvals"] if a["kind"] == "final_check"]
        assert len(fc) == 1 and fc[0]["words"] == "build it", fc
        assert fc[0]["bound"]["final_token"] == tok, fc
        assert set(fc[0]["bound"]["digests"]) == {"slide_01/A", "slide_02/A"}, fc
        print("    ok: compiled; the words are bound to the page's token and digests")

        print("[8] check_done lists them as chat approvals, not as overrides")
        _state.record_vision_qc(out, str(deck), slides_reviewed=2)
        r = H.run("check_done.py", "--out", out)
        assert r.returncode == 0, r.stdout
        assert "chat approvals: 2 recorded" in r.stdout, r.stdout
        assert '"use A on both slides"' in r.stdout and '"build it"' in r.stdout, r.stdout
        ov_part = r.stdout.split("chat approvals")[0]
        assert "build it" not in ov_part and "use A on both" not in ov_part, r.stdout
        print("    ok")
    finally:
        H.cleanup(tmp)
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
