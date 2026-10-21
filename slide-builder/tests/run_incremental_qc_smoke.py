#!/usr/bin/env python3
"""Smoke test: the owner's definition of done for the quality check (item 8,
session report, approved 2026-10-09), and proof of reading.

  1. one render, reused: a folder holding a render of these exact deck bytes
     by the same program at the same dpi is reused (nothing drawn again); a
     changed image, other deck bytes or another dpi are not reused
  2. per-page records: page, image hash, renderer, dpi, time, open findings
     with severity, beside the deck summary and the deck-wide look
  3. the script works out the changed pages by image hash; a round must
     cover every one of them, and cannot claim a page it was not given
  4. every round records one deck-wide look (thumbnail sheets, each changed
     page beside an unchanged page of the same layout) and runs the
     whole-file hygiene and type-scale check itself
  5. check_done: every page's reviewed image is the current render's image
     of this exact deck, a deck-wide look for this deck, no Critical or Major
     open; a finding on an unchanged page carries forward and stays open
  6. every page is read again when there is no earlier review, the page
     count, the renderer, the dpi or the template changed
  7. proof of reading: a wrong or old check code, a quote that is not on the
     page, a page with no text and no description, a count larger than the
     notes prove, all refused; the codes are in no text file

No renderer is used: tests/_qc_helper.py draws stand-in page images. The
renderers themselves were measured separately (render_manifest.py's notes).

Run:  py -3 slide-builder/tests/run_incremental_qc_smoke.py
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
QC_SCRIPTS = HERE.parents[1] / "slide-qc" / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(QC_SCRIPTS))
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import _state  # noqa: E402
import qc_pages  # noqa: E402
import render_manifest as RM  # noqa: E402
import _qc_helper as Q  # noqa: E402

TITLES = ["Northwind Ferries annual review for the board",
          "Evening crossings lifted ridership by fourteen percent",
          "Fares held flat for commuters through the year",
          None,                                   # a page with no text at all
          "Two vessels refitted ahead of the winter schedule"]


def _deck(path: Path, titles=TITLES, word: str = "") -> Path:
    p = Presentation()
    for i, t in enumerate(titles):
        if t is None:
            s = p.slides.add_slide(p.slide_layouts[6])
            s.shapes.add_shape(1, Inches(2), Inches(2), Inches(4), Inches(2))
            continue
        s = p.slides.add_slide(p.slide_layouts[0 if i == 0 else 5])
        s.shapes.title.text = t.replace("fourteen", word) if word else t
    p.save(str(path))
    return path


def _build(tmp: Path, name: str, titles=TITLES) -> tuple[Path, Path]:
    d = tmp / name
    d.mkdir()
    deck = _deck(d / "final_deck.pptx", titles)
    K = _state.option_key
    _state.record_prep(d, "h", d)
    _state.begin_finalize(d)
    opts = [K(i, "A") for i in range(1, len(titles) + 1)]
    _state.record_option_qc(d, {k: {"blocks": 0, "reasons": []} for k in opts})
    _state.end_finalize(d, "ok")
    _state.record_compile(d, kind="picks", output=deck, slides=len(titles), options=opts)
    return d, deck


def _recompile(d: Path, deck: Path) -> None:
    _state.record_compile(d, kind="picks", output=deck,
                          slides=len(Presentation(str(deck)).slides),
                          options=[_state.option_key(i, "A")
                                   for i in range(1, len(Presentation(str(deck)).slides) + 1)])


def _done(d: Path) -> tuple[int, str]:
    r = subprocess.run([sys.executable, str(SCRIPTS / "check_done.py"), "--out", str(d)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace",
                       env={**os.environ, "PYTHONPATH": str(SCRIPTS)})
    return r.returncode, r.stdout


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    tmp = Path(tempfile.mkdtemp(prefix="incremental_qc_"))
    try:
        # ------------------------------------------------------------ 1
        print("[1] a render of these exact bytes is reused; anything else is not")
        d, deck = _build(tmp, "reuse")
        qc = qc_pages.qc_dir(d)
        man = Q.fake_render(d, deck)
        assert man["deck_digest"] == RM.file_md5(deck) and man["dpi"] == 150 and \
            man["slides"] == 5 and all(p["hash"] for p in man["pages"]), man
        ok, why = RM.reusable(qc, deck, "libreoffice", 150)
        assert ok, why
        import render_slides as RS
        before = {p.name: p.stat().st_mtime_ns for p in qc.glob("slide_*.png")}
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            got = RS.render(deck, qc, dpi=150, renderer="libreoffice", reuse=True)
        assert got == "libreoffice" and "Reused" in buf.getvalue(), buf.getvalue()
        assert before == {p.name: p.stat().st_mtime_ns for p in qc.glob("slide_*.png")}, \
            "a reused render must not draw anything"
        assert not RM.reusable(qc, deck, "libreoffice", 200)[0], "another dpi"
        assert not RM.reusable(qc, deck, "powerpoint", 150)[0], "another program"
        other = _deck(tmp / "other.pptx", word="nine")
        assert not RM.reusable(qc, other, "libreoffice", 150)[0], "other deck bytes"
        from PIL import Image
        Image.new("RGB", (400, 225), (0, 0, 0)).save(qc / "slide_02.png")
        ok, why = RM.reusable(qc, deck, "libreoffice", 150)
        assert not ok and "slide_02.png" in why, why
        assert RM.verify_images(man, qc) == ["slide_02.png"]
        print("    ok: reused without drawing; other dpi / program / bytes / a changed "
              "image are not reused")

        # ------------------------------------------------------------ 2 + 4
        print("[2] a full first round writes per-page records and the deck-wide look")
        d, deck = _build(tmp, "rounds")
        qc = qc_pages.qc_dir(d)
        Q.fake_render(d, deck)
        maj = {"severity": "Major", "category": "chrome", "issue": "source line moved"}
        adv = {"severity": "Advisory", "issue": "tight spacing"}
        rc, o, plan = Q.read_round(d, findings={3: [maj], 5: [adv]})
        assert rc == 0, o
        assert plan["full"] and plan["read_pages"] == [1, 2, 3, 4, 5], plan
        assert "no earlier page-by-page review" in plan["reasons"][0], plan["reasons"]
        st = _state.read_state(d)
        man = RM.read(qc)
        pg = st["vision_pages"]["pages"]
        assert sorted(pg, key=int) == ["1", "2", "3", "4", "5"], pg.keys()
        for n in range(1, 6):
            r = pg[str(n)]
            assert r["hash"] == man["pages"][n - 1]["hash"] and r["renderer"] == "libreoffice" \
                and r["dpi"] == 150 and r["reviewed_at"], r
        assert pg["3"]["findings"] == [{"severity": "Major", "category": "chrome",
                                        "issue": "source line moved"}], pg["3"]
        look = st["vision_deck_look"]
        assert look["deck_digest"] == RM.file_md5(deck) and look["hygiene"]["ran"], look
        assert look["sheets"] and not look["pairs"], "a full read has no unchanged page to pair"
        assert st["vision_qc"]["slides_reviewed"] == 5 and st["vision_qc"]["majors"] == 1
        assert "vision_plan" not in st, "the round's codes are one-time"
        assert st["vision_qc"]["no_text_pages"] == [4], st["vision_qc"]
        rc, o = _done(d)
        assert rc == 1 and "1 Major" in o, o
        print("    ok: page, hash, program, dpi, time, findings; sheets; hygiene ran; "
              "the open Major refuses")

        # ------------------------------------------------------------ 3 + 4 + 5
        print("[3] only the pages whose image changed are read; the round must cover them")
        Q.fake_render(d, deck, vary={2: "edited"})
        assert qc_pages.unreviewed_pages(_state.read_state(d), RM.read(qc)) == [2]
        plan = qc_pages.make_plan(d, _return_codes=True)
        blank = (qc / qc_pages.NOTES).read_text(encoding="utf-8")       # as plan wrote it
        assert plan["read_pages"] == [2] and plan["carried"] == [1, 3, 4, 5], plan
        pair = plan["pairs"][0]
        assert pair["changed"] == 2 and pair["partner"] == 3 and pair["same_layout"], pair
        assert (qc / "to_read" / "pair_02.png").exists() and (qc / "to_read" / "page_02.png").exists()
        assert not (qc / "to_read" / "page_01.png").exists(), "an unchanged page is not handed out"
        Q.fill_notes(d, plan, skip_pages=(2,))
        rc, o = Q.record(d)
        assert rc == 3 and "no verified note: 2" in o and "images changed" in o, o
        notes = json.loads(blank)
        notes["pages"][0].update(code=plan["_codes"]["page:2"], quote=TITLES[1])
        notes["pages"].append({"page": 1, "code": "AAAAAA", "quote": TITLES[0]})
        (qc / qc_pages.NOTES).write_text(json.dumps(notes), encoding="utf-8")
        rc, o = Q.record(d)
        assert rc == 3 and "page 1 was not in this round's pages" in o, o
        notes["pages"].pop()
        (qc / qc_pages.NOTES).write_text(json.dumps(notes), encoding="utf-8")
        print("    ok: changed page 2 only, paired with page 3 (same layout); a skipped "
              "changed page and an unasked page are refused")

        print("[4] the deck-wide look is required every round")
        (qc / qc_pages.NOTES).write_text(blank, encoding="utf-8")   # as plan wrote it
        Q.fill_notes(d, plan, skip_look=True)
        rc, o = Q.record(d)
        assert rc == 3 and "deck-wide look: the check code for sheet_01.png is missing" in o, o
        print("    ok")

        print("[5] carried findings stay open; check_done follows the current render")
        Q.fill_notes(d, plan)
        rc, o = Q.record(d)
        assert rc == 0 and "1 page(s) read" in o and "4 carried" in o, o
        st = _state.read_state(d)
        assert st["vision_pages"]["pages"]["3"]["findings"][0]["severity"] == "Major"
        rc, o = _done(d)
        assert rc == 1 and "1 Major" in o and "carry forward" in o, o
        rc, o, _ = Q.read_round(d, also=[3], findings={3: []})      # page 3 fixed and re-read
        assert rc == 0, o
        rc, o = _done(d)
        assert rc == 0 and "1 Advisory open" in o, "the Advisory on page 5 carries: " + o
        Q.fake_render(d, deck, vary={2: "edited", 5: "moved"})      # the deck looks different
        rc, o = _done(d)
        assert rc == 1 and "never reviewed as they look in this deck: page 5" in o, o
        Q.fake_render(d, deck, vary={2: "edited"})
        assert _done(d)[0] == 0
        _deck(deck, word="nine")                                    # new bytes, no render
        _recompile(d, deck)
        rc, o = _done(d)
        assert rc == 1 and "no render of this exact deck" in o, o
        print("    ok: the carried Major refuses until page 3 is re-read; a page that "
              "looks different is unreviewed; new bytes need a new render")

        # ------------------------------------------------------------ 6
        print("[6] every page is read again when the count, program, dpi or template change")

        def _full_reason(dd, dk, **kw):
            Q.fake_render(dd, dk, **kw)
            a = qc_pages.assess(_state.read_state(dd), RM.read(qc_pages.qc_dir(dd)))
            return a["full"], " ".join(a["reasons"])

        d, deck = _build(tmp, "fallback")
        Q.fake_render(d, deck)
        assert Q.read_round(d)[0] == 0
        full, why = _full_reason(d, deck)
        assert not full, why
        full, why = _full_reason(d, deck, renderer="powerpoint")
        assert full and "renderer changed" in why, why
        full, why = _full_reason(d, deck, dpi=200)
        assert full and "dpi changed" in why, why
        _deck(deck, TITLES + ["A sixth page about the winter timetable"])
        _recompile(d, deck)
        full, why = _full_reason(d, deck)
        assert full and "page count changed (5 -> 6)" in why, why
        _deck(deck)
        _recompile(d, deck)
        tpl = tmp / "tpl.pptx"
        with zipfile.ZipFile(deck) as zi, zipfile.ZipFile(tpl, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in zi.infolist():
                data = zi.read(it.filename)
                if it.filename == "ppt/theme/theme1.xml":
                    data = data.replace(b'name="Office Theme"', b'name="Other Theme"')
                zo.writestr(it, data)
        shutil.move(str(tpl), str(deck))
        _recompile(d, deck)
        full, why = _full_reason(d, deck)
        assert full and "template" in why, why
        rc, o = _done(d)
        assert rc == 1 and "never reviewed" in o, o
        print("    ok: renderer / dpi / page count / template each force a full read")

        # ------------------------------------------------------------ 7
        print("[7] proof of reading")
        d, deck = _build(tmp, "proof")
        qc = qc_pages.qc_dir(d)
        Q.fake_render(d, deck)
        plan = qc_pages.make_plan(d, _return_codes=True)
        codes = set(plan["_codes"].values())
        for f in qc.rglob("*"):
            if f.is_file() and f.suffix in (".json", ".md", ".txt"):
                text = f.read_text(encoding="utf-8", errors="replace")
                assert not any(c in text for c in codes), f"a check code is readable in {f}"
        st_text = (d / "_state.json").read_text(encoding="utf-8")
        assert not any(c in st_text for c in codes), "the build record holds only hashes"

        Q.fill_notes(d, plan, wrong_code_pages=(2,))
        rc, o = Q.record(d)
        assert rc == 3 and "page 2: the check code is wrong" in o, o
        Q.fill_notes(d, plan, bad_quote_pages=(3,))
        rc, o = Q.record(d)
        assert rc == 3 and "is not text on page 3" in o, o
        Q.fill_notes(d, plan)
        notes = json.loads((qc / qc_pages.NOTES).read_text(encoding="utf-8"))
        for e in notes["pages"]:
            if e["page"] == 2:
                e["quote"] = "Fares held flat for commuters"     # page 3's own text
            if e["page"] == 4:
                e["seen"] = ""                                   # no text, no description
        (qc / qc_pages.NOTES).write_text(json.dumps(notes), encoding="utf-8")
        rc, o = Q.record(d)
        assert rc == 3 and "is not text on page 2" in o and "page 4 has no text" in o, o
        Q.fill_notes(d, plan)
        rc, o = Q.record(d, "--slides-reviewed", 71)
        assert rc == 3 and "claims more pages than the notes prove (5)" in o, o
        assert "vision_pages" not in _state.read_state(d), "a refused round records nothing"
        old = plan
        plan = qc_pages.make_plan(d, _return_codes=True)                  # a new round
        Q.fill_notes(d, old)                                              # last round's codes
        rc, o = Q.record(d)
        assert rc == 3 and "check code is wrong" in o, "codes are one-time: " + o
        Q.fill_notes(d, plan)
        rc, o = Q.record(d, "--slides-reviewed", 5)
        assert rc == 0 and "5 page(s) read with proof" in o, o
        print("    ok: wrong code, a quote from another page, a page with nothing "
              "written, an over-claim and last round's codes are all refused")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
