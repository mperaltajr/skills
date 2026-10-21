#!/usr/bin/env python3
"""Smoke test: pages that repeat are designed once and copied by script.

Owner's decision 2026-10-14 (speed verdict Q3, items 1, 2 and 4), on the real
pipeline with made-up pages:

  1. prep sends a designer only to the FIRST section divider; its _prompt.md
     carries the PAGE VALUES block, the other dividers' prompts say "no
     worker"; the copies are on the direct path even when the router would
     have picked the sketch path
  2. finalize writes every copy from the picked design (clone_pages.py); each
     copy matches the master everywhere except the words that differ, shows
     its own section number (never the master's), and passes check_fit
  3. a second finalize changes nothing; a hand edit on a copy is undone (and
     kept aside) with a warning naming the master
  4. a copy takes its master's pick; a comment on a copy goes to the master in
     edit mode, and after the edit every copy is made again from it
  5. a rebuild of a copy is refused, naming the master
  6. after compile, the cover, contents and executive summary are kept for the
     template; the next deck on it sends no designer for them and fills them
     from its own storyline, word for word (contents page ranges from its
     sections)
  7. the brief readers: bold labels, section lines, page ranges, refusals of a
     block a copy cannot fill

Run:  py -3 slide-builder/tests/run_clone_pages_smoke.py
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))

import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402
import clone_pages as CP  # noqa: E402

from pptx import Presentation  # noqa: E402

BRIEF = """---
deck_type: client-pitch
audience: "Test audience"
governing_thought: "The test deck proves repeated pages are copied, not redesigned."
storyline_gate_passed: true
default_layout: body_canonical_light
---

# Copy test

## Governing thought (the whole deck)
The test deck proves repeated pages are copied, not redesigned.
{slides}"""

COVER = """
### Slide {n} — {title}

**Slide type:** Cover / Title
**Layout:** cover_light
**Governing thought (the claim):** {title}
**Evidence / content:**
- **Title:** {title}
- **Subtitle:** {subtitle}
- **Presenter:** {presenter}
- **Date:** October 2026
"""

SUMMARY = """
### Slide {n} — Executive summary

**Slide type:** Executive Summary
**Governing thought (the claim):** The case in brief.
**The takeaway:** none, no line under the headline on this page
**Evidence / content:**
{rows}
"""

CONTENTS = """
### Slide {n} — What this deck covers

**Slide type:** Agenda / Contents
**Governing thought (the claim):** What this deck covers.
**The takeaway:** none
**Evidence / content:**
{rows}
"""

DIVIDER = """
### Slide {n} — {title}

**Slide type:** Section divider
**Layout:** section_divider
**Governing thought (the claim):** {title}
**The takeaway:** none, no line under the headline on this page
**Section description:** {desc}
**Evidence / content:**
- {title}
"""

CONTENT = H.SLIDE

HEAD = '''# Slide {n} option {L} - pattern: Full canvas (test master)
# Variant: test master for copies
# Brief fidelity check: every word comes from the PAGE VALUES block
{block}
import sys
from pathlib import Path
sys.path.insert(0, r"{skill}")
from twins.helpers import new_slide, add_text

'''

TAIL = '''

if __name__ == "__main__":
    prs = build()
    prs.save(str(Path(__file__).resolve().parent / "option_{L}.pptx"))
'''

DIVIDER_BODY = '''def build():
    prs, slide = new_slide()
    add_text(slide, "section-number", f"{SECTION_NUMBER:02d}", 80, 200, 300, 90,
             font_size_pt=40, bold=True)
    add_text(slide, "section-name", SECTION_TITLE, 80, 300, 900, 60, font_size_pt=28)
    add_text(slide, "section-line", SECTION_DESCRIPTION, 80, 370, 1000, 50, font_size_pt=16)
    for i, name in enumerate(SECTIONS, start=1):
        add_text(slide, f"rail-{i:02d}", f"{i:02d} {name}", 80 + (i - 1) * 300, 560, 290, 40,
                 font_size_pt=12, bold=(i == SECTION_NUMBER))
    return prs
'''

COVER_BODY = '''def build():
    prs, slide = new_slide()
    add_text(slide, "cover-name", COVER_TITLE, 80, 220, 1000, 80, font_size_pt=32, bold=True)
    add_text(slide, "cover-line", COVER_SUBTITLE, 80, 320, 1000, 40, font_size_pt=16)
    add_text(slide, "cover-who", COVER_PRESENTER, 80, 520, 600, 30, font_size_pt=12)
    add_text(slide, "cover-date", COVER_DATE, 80, 560, 600, 30, font_size_pt=12)
    return prs
'''

SUMMARY_BODY = '''def build():
    prs, slide = new_slide()
    for i, row in enumerate(SUMMARY_ROWS):
        y = 200 + i * 90
        add_text(slide, f"row-label-{i + 1:02d}", row["label"], 80, y, 300, 70, font_size_pt=12,
                 bold=True)
        add_text(slide, f"row-text-{i + 1:02d}", row["text"], 400, y, 780, 70, font_size_pt=12)
    return prs
'''

CONTENTS_BODY = '''def build():
    prs, slide = new_slide()
    for i, item in enumerate(CONTENTS):
        y = 200 + i * 80
        add_text(slide, f"item-name-{i + 1:02d}", f"{item['number']:02d}  {item['name']}",
                 80, y, 500, 36, font_size_pt=16, bold=True)
        add_text(slide, f"item-pages-{i + 1:02d}",
                 f"Pages {item['first_page']} to {item['last_page']}", 600, y, 200, 36,
                 font_size_pt=12)
        add_text(slide, f"item-line-{i + 1:02d}", item["description"], 80, y + 36, 1000, 36,
                 font_size_pt=12)
    return prs
'''


def _deck_one() -> str:
    s = [COVER.format(n=1, title="Older GPUs find a second market",
                      subtitle="For discussion with investors", presenter="A. Founder"),
         SUMMARY.format(n=2, rows="\n".join([
             "- **01 THE MARKET, PAGES 4 TO 5** — Demand is large and growing.",
             "- **02 THE PLAN, PAGES 6 TO 7** — The plan buys, tests and resells.",
             "- **03 THE ASK, PAGES 8 TO 9** — The ask funds the first two years."])),
         CONTENTS.format(n=3, rows="\n".join([
             "- **01 THE MARKET** — Pages 4 to 5: How big the demand is.",
             "- **02 THE PLAN** — Pages 6 to 7: What the business does.",
             "- **03 THE ASK** — Pages 8 to 9: What the money pays for."])),
         DIVIDER.format(n=4, title="The Market", desc="How big the demand is."),
         CONTENT.format(n=5),
         DIVIDER.format(n=6, title="The Plan", desc="What the business does."),
         CONTENT.format(n=7),
         DIVIDER.format(n=8, title="The Ask", desc="What the money pays for."),
         CONTENT.format(n=9)]
    return BRIEF.format(slides="".join(s))


def _deck_two() -> str:
    s = [COVER.format(n=1, title="A second deck with other words",
                      subtitle="Board update", presenter="B. Chair"),
         SUMMARY.format(n=2, rows="\n".join([
             "- **01 WHERE WE ARE** — Sales doubled this year.",
             "- **02 WHAT IS NEXT** — Two new sites open in spring."])),
         CONTENTS.format(n=3, rows="\n".join([
             "- **01 WHERE WE ARE** — The year in numbers.",
             "- **02 WHAT IS NEXT** — The sites and the hiring plan."])),
         DIVIDER.format(n=4, title="Where We Are", desc="The year in numbers."),
         CONTENT.format(n=5),
         CONTENT.format(n=6),
         DIVIDER.format(n=7, title="What Is Next", desc="The sites and the hiring plan."),
         CONTENT.format(n=8)]
    return BRIEF.format(slides="".join(s))


def _prep(root: Path, name: str, text: str) -> tuple[Path, object]:
    brief = root / f"{name}.md"
    brief.write_text(text, encoding="utf-8")
    assert H.run("seal_brief.py", "--brief", brief, "--accepted", "test").returncode == 0
    out = root / name
    r = H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE, "--out", out,
              "--pattern", "direct", "--confirm-template")
    assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
    return out, r


def _write_master(out: Path, n: int, body: str, L: str = "A") -> Path:
    """Stand in for the designer: the block exactly as its _prompt.md gives it."""
    prompt = (out / f"slide_{n:02d}" / "_prompt.md").read_text(encoding="utf-8")
    m = re.search(r"```python\n(# --- PAGE VALUES.*?# --- END PAGE VALUES ---)\n```", prompt, re.S)
    assert m, f"slide {n}: no PAGE VALUES block in its _prompt.md"
    p = out / f"slide_{n:02d}" / f"option_{L}.py"
    p.write_text(HEAD.format(n=n, L=L, block=m.group(1), skill=str(H.SKILL)) + body
                 + TAIL.format(L=L), encoding="utf-8")
    return p


def _shapes(pptx: Path) -> list[tuple]:
    s = Presentation(str(pptx)).slides[0]
    out = []
    for sh in s.shapes:
        txt = sh.text_frame.text if getattr(sh, "has_text_frame", False) else ""
        out.append((sh.name, sh.left, sh.top, sh.width, sh.height, txt))
    return out


def _texts(pptx: Path) -> str:
    return "\n".join(t for *_x, t in _shapes(pptx))


def _message(out: Path, picks: dict, feedback: str = "") -> str:
    body = _state.canonical_picks(picks)
    tok = _state.read_state(out)["review"]["token"]
    plain = " ".join(f"{int(k.split('_')[1])}{v}" for k, v in sorted(picks.items()))
    return (f"Build my Slide Lab deck.\nFolder: {out}\n"
            f"Picks: {plain} (check {_state.approval_check(tok, body)})\n") + feedback


def unit() -> None:
    print("[7] the brief readers and the block checks")
    items = CP.bullets("- **01 THE GPU ECONOMY, PAGES 4 TO 7** — AI demand is large.\n"
                       "  - first point.\n- plain line, wrapped\n  onto two lines.")
    assert items[0] == {"label": "01 THE GPU ECONOMY, PAGES 4 TO 7",
                        "text": "AI demand is large", "points": ["first point"]}, items
    assert items[1]["label"] == "" and items[1]["text"] == "plain line, wrapped onto two lines"
    # A client-style divider: the line is in its first bullet, after a bold label
    # and a repeat of the title.
    d = {"slide_n": 4, "title": "The GPU Economy", "archetype": "Section divider",
         "so_what": "none, no line under the headline on this page",
         "evidence_content": "- **SECTION 01** — The GPU Economy: Demand for older GPUs "
                             "outgrows used supply, and no buyer takes them.\n"}
    assert CP._section_description(d) == ("Demand for older GPUs outgrows used supply, and "
                                          "no buyer takes them"), CP._section_description(d)
    # The contents' own page range is checked against the sections'.
    slides = [{"slide_n": 1, "archetype": "Agenda / Contents",
               "evidence_content": "- **01 A** — Pages 2 to 9: one.\n"},
              {"slide_n": 2, "archetype": "Section divider", "title": "A"},
              {"slide_n": 5, "archetype": "Content"}]
    rows, warns = CP._contents(slides[0], CP.sections(slides))
    assert rows[0]["first_page"] == 2 and rows[0]["last_page"] == 5, rows
    assert rows[0]["description"] == "one" and warns and "2 to 9" in warns[0], warns
    # Blocks a copy cannot fill are refused with the reason.
    good = CP.render_block({"CONTENTS": []})
    assert CP.parse_block(good, CP.CONTENTS)[2] == ""
    assert "lacks SUMMARY_ROWS" in CP.parse_block(good, CP.SUMMARY)[2]
    bad = good.replace("CONTENTS = []", "CONTENTS = []\nW = 3")
    assert "cannot fill" in CP.parse_block(bad, CP.CONTENTS)[2]
    assert "no PAGE VALUES" in CP.parse_block("x = 1\n", CP.CONTENTS)[2]
    assert CP._type_problems({"SECTION_NUMBER": "01"}, {"SECTION_NUMBER": 1})
    # Picks and comments of copies.
    meta = {"slides": [{"n": 4, "page_copy": {"role": "master", "kind": "divider"}},
                       {"n": 6, "page_copy": {"role": "copy", "kind": "divider", "master": 4}}]}
    assert CP.check_copy_picks(meta, {"slide_04": "A", "slide_06": "A"}) == ""
    assert CP.check_copy_picks(meta, {"slide_04": "A", "slide_06": "-"}) == ""
    why = CP.check_copy_picks(meta, {"slide_04": "B", "slide_06": "A"})
    assert "pick B on slide 6" in why, why
    fb, notes, why = CP.route_copy_feedback(meta, {"slide_06": {"note": "bigger number"}},
                                            {"slide_04": "A", "slide_06": "A"})
    assert not why and "slide_06" not in fb and "bigger number" in fb["slide_04"]["note"], fb
    _fb, _n, why = CP.route_copy_feedback(meta, {"slide_06": {"note": "x"}},
                                          {"slide_04": "-", "slide_06": "A"})
    assert "left out" in why, why
    # The router's sketch path is overridden for copied pages.
    import build_deck as BD

    class _A:
        brief = Path(tempfile.gettempdir()) / "b.md"
        out = Path(tempfile.gettempdir()) / "o"
        template = H.TEMPLATE
    sl = [{"slide_n": 1, "archetype": "Executive Summary", "evidence_content": "- **A** — b"},
          {"slide_n": 2, "archetype": "Section divider", "title": "One"},
          {"slide_n": 3, "archetype": "Analytical"},
          {"slide_n": 4, "archetype": "Section divider", "title": "Two"}]
    pps = BD._classify_all_slides(sl, "auto")
    assert pps["1"] == "sketch", pps
    plan, rc = BD.plan_page_copies(_A, sl, pps, None, None)
    assert rc == 0 and pps["1"] == "direct" and pps["3"] == "sketch", pps
    assert plan[2]["role"] == "master" and plan[4] == {**plan[4], "role": "copy", "master": 2}
    print("    ok")


def main() -> int:
    root = Path(tempfile.mkdtemp(prefix="slidelab_clone_pages_"))
    agents = root / "agents"
    agents.mkdir()
    for f in (H.SKILL / "agents").glob("slide-builder-*.md"):
        shutil.copy(f, agents / f.name)
    os.environ["SLIDE_LAB_AGENTS_DIR"] = str(agents)
    try:
        unit()
        out, r = _prep(root, "deck_one", _deck_one())

        print("[1] prep: one designer for the dividers; copies get no worker")
        assert "Do NOT send a worker for slide(s) 6, 8" in r.stdout, r.stdout[-2500:]
        assert "slides 6, 8 are copies of it" in r.stdout, r.stdout[-2500:]
        meta = json.loads((out / "_meta.json").read_text(encoding="utf-8"))
        pcs = {s["n"]: s.get("page_copy") or {} for s in meta["slides"]}
        assert pcs[4]["role"] == "master" and pcs[4]["copies"] == [6, 8], pcs[4]
        assert pcs[6]["master"] == 4 and pcs[6]["values"]["SECTION_NUMBER"] == 2, pcs[6]
        assert pcs[8]["values"]["SECTION_DESCRIPTION"] == "What the money pays for"
        assert pcs[8]["values"]["SECTIONS"] == ["The Market", "The Plan", "The Ask"]
        assert {pcs[n]["role"] for n in (1, 2, 3)} == {"master"}, "no kept design yet"
        assert not pcs[5], "a content page is never copied"
        assert CP.no_worker_slides(meta) == [6, 8]
        p4 = (out / "slide_04" / "_prompt.md").read_text(encoding="utf-8")
        assert "## 0. Page values" in p4 and "SECTION_NUMBER = 1" in p4
        p6 = (out / "slide_06" / "_prompt.md").read_text(encoding="utf-8")
        assert p6.startswith("# NO WORKER FOR THIS SLIDE: slide 6 is a copy of slide 4")
        plan_md = (out / "dispatch_plan.md").read_text(encoding="utf-8")
        assert "## Pages made by script" in plan_md
        print("    ok")

        print("[2] finalize makes the copies; each matches the master but for its words")
        _write_master(out, 4, DIVIDER_BODY)
        _write_master(out, 1, COVER_BODY)
        _write_master(out, 2, SUMMARY_BODY)
        _write_master(out, 3, CONTENTS_BODY)
        for n in (5, 7, 9):
            H.write_option(out, n)
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-2000:]
        assert "[0b] Pages copied by script" in r.stdout, r.stdout[-3000:]
        assert "slide 6 option A" in r.stdout and "slide 8 option A" in r.stdout
        c6 = (out / "slide_06" / "option_A.py").read_text(encoding="utf-8")
        assert c6.startswith("# Slide 6 option A") and "# PAGE COPY: made by clone_pages.py " \
            "from slide 4 option A" in c6, c6[:300]
        master = _shapes(out / "slide_04" / "option_A.pptx")
        allowed = {"section-number", "section-name", "section-line"}
        for n, num, name in ((6, "02", "The Plan"), (8, "03", "The Ask")):
            copy = _shapes(out / f"slide_{n:02d}" / "option_A.pptx")
            assert len(copy) == len(master), (n, len(copy), len(master))
            differ = set()
            for a, b in zip(master, copy):
                assert a[:5] == b[:5], f"slide {n}: {b[0]} moved or resized ({a[:5]} -> {b[:5]})"
                if a[5] != b[5]:
                    differ.add(a[0])
            title_like = {x for x in differ if x.lower().startswith("title")}
            assert differ - title_like <= allowed, (n, differ)
            texts = {sh[0]: sh[5] for sh in copy}
            assert texts["section-number"] == num, (n, texts["section-number"])
            assert texts["section-name"] == name
            assert "01" != texts["section-number"]
            r = H.run("check_fit.py", out / f"slide_{n:02d}" / "option_A.pptx")
            assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        print("    ok: same shapes in the same places; only the number, name and line "
              "differ; 02 and 03 shown; check_fit passes")

        print("[3] a second finalize changes nothing; a hand edit on a copy is undone")
        stamp = _state.option_files_stamp(out / "slide_06")
        r = H.finalize(out)
        assert r.returncode == 0 and "already current" in r.stdout, r.stdout[-2000:]
        assert _state.option_files_stamp(out / "slide_06") == stamp
        f8 = out / "slide_08" / "option_A.py"
        f8.write_text(f8.read_text(encoding="utf-8").replace("font_size_pt=28",
                                                              "font_size_pt=20"),
                      encoding="utf-8")
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-2000:]
        assert "WARNING: slide 8 option A was changed by hand" in r.stdout, r.stdout[-2000:]
        assert "font_size_pt=28" in f8.read_text(encoding="utf-8")
        assert "font_size_pt=20" in (out / "slide_08" / "_prev_copy" / "option_A.py").read_text(
            encoding="utf-8")
        print("    ok")

        print("[4] picks: a copy takes its master's pick; its comment goes to the master")
        assert H.run("build_review.py", "--out", out).returncode == 0
        page = (out / "REVIEW.html").read_text(encoding="utf-8")
        assert "A copy of slide 4&#x27;s design" in page or "A copy of slide 4's design" in page
        picks = {f"slide_{n:02d}": "A" for n in range(1, 10)}
        msg = _message(out, picks, "Feedback:\n  slide_06:\n    other: make the number bigger\n")
        r = H.run("record_picks.py", "--out", out, "--approved", msg)
        assert r.returncode == 0, r.stdout[-2000:]
        assert "comment on slide 6 goes to slide 4" in r.stdout, r.stdout[-2000:]
        req = out / "slide_04" / "_edit_request.md"
        assert req.exists() and not (out / "slide_06" / "_edit_request.md").exists()
        assert "(from slide 6, a copy of this design) make the number bigger" in req.read_text(
            encoding="utf-8")
        m4 = out / "slide_04" / "option_A.py"
        m4.write_text(m4.read_text(encoding="utf-8").replace("font_size_pt=40",
                                                              "font_size_pt=44"),
                      encoding="utf-8")
        r = H.run("record_picks.py", "--out", out, "--edits-done")
        assert r.returncode == 0, r.stdout[-2000:]
        # A one-slide finalize of the master finalizes the copies it changed
        # (a QC fix on the master reaches every divider in one command).
        r = H.finalize(out, "--slide", "4")
        assert r.returncode == 0, r.stdout[-2000:]
        assert "[8] Finalizing the copies of slide 4: slide(s) 6, 8" in r.stdout, r.stdout[-2000:]
        for n in (6, 8):
            assert "font_size_pt=44" in (out / f"slide_{n:02d}" / "option_A.py").read_text(
                encoding="utf-8"), f"slide {n} was not made again from the edited master"
            done = _shapes(out / f"slide_{n:02d}" / "option_A.pptx")
            assert done and done[0], "copy not finalized"
        print("    ok: the edit request is on slide 4; after the edit, finalize --slide 4 "
              "makes and finalizes both copies")

        print("[5] a rebuild of a copy is refused, naming the master")
        r = H.run("build_deck.py", "--template", H.TEMPLATE, "--out", out, "--slide", "6")
        assert r.returncode == 2 and "rebuild slide 4" in r.stderr, r.stderr[-1500:]
        print("    ok")

        print("[6] after compile the cover, contents and summary are kept for the template")
        r = H.run("build_review.py", "--out", out, "--final")
        assert r.returncode == 0, r.stderr[-1500:]
        tok = _state.read_state(out)["final_check"]["token"]
        r = H.run("compile_picks.py", "--out", out, "--final-token", tok)
        assert r.returncode == 0, r.stdout[-3000:]
        kept = Path(H.PAGE_MASTERS) / H.TEMPLATE.stem
        assert sorted(p.name for p in kept.glob("*.py")) == [
            "contents.py", "cover.py", "exec_summary.py"], list(kept.iterdir())
        assert "kept the cover design of slide 1" in r.stdout, r.stdout[-2500:]

        out2, r = _prep(root, "deck_two", _deck_two())
        meta2 = json.loads((out2 / "_meta.json").read_text(encoding="utf-8"))
        pc2 = {s["n"]: s.get("page_copy") or {} for s in meta2["slides"]}
        assert all(pc2[n]["master"] == "template" for n in (1, 2, 3)), pc2
        assert pc2[7]["master"] == 4 and pc2[4]["role"] == "master"
        assert "Do NOT send a worker for slide(s) 1, 2, 3, 7" in r.stdout, r.stdout[-2500:]
        _write_master(out2, 4, DIVIDER_BODY)
        for n in (5, 6, 8):
            H.write_option(out2, n)
        r = H.finalize(out2)
        assert r.returncode == 0, r.stdout[-3000:] + r.stderr[-2000:]
        t1 = _texts(out2 / "slide_01" / "option_A.pptx")
        assert "A second deck with other words" in t1 and "B. Chair" in t1, t1
        assert "Older GPUs" not in t1
        t2 = _texts(out2 / "slide_02" / "option_A.pptx")
        assert "Sales doubled this year" in t2 and "Two new sites open in spring" in t2, t2
        assert "THE ASK" not in t2
        t3 = _texts(out2 / "slide_03" / "option_A.pptx")
        assert "Pages 4 to 6" in t3 and "Pages 7 to 8" in t3, t3
        assert "The sites and the hiring plan" in t3
        for n in (1, 2, 3):
            r = H.run("check_fit.py", out2 / f"slide_{n:02d}" / "option_A.pptx")
            assert r.returncode == 0, r.stdout[-1500:]
        print("    ok: kept, then filled for the second deck with its own words; "
              "contents pages 4 to 6 and 7 to 8 from its sections")
    finally:
        H.cleanup(root)
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
