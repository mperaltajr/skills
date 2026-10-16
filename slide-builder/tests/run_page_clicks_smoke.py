#!/usr/bin/env python3
"""Click test: every button of the one Slide Lab page, the way a user clicks it
(owner's report, 2026-10-16: "the Build it button and Send fixes doesn't work
either ... check all of the clicks for it so that it works").

Builds a real small deck through the scripts on the bundled test template
(Storyline tab, then the designs, then the final check), keeps the one page
as it is at each stage, then opens each copy from disk in headless Microsoft
Edge (its own browser, never the user's) and clicks every control on every
tab:

  Storyline     the tabs, edit a line, + Comment, every visual box,
                Suggestion / Must use, Done, Let the designers pick all
  Slide designs Pick A/B/C, the picture itself, Replace these, +3 designs,
                Leave out, every quick-feedback chip, the text boxes, Clear,
                All options in one deck (keep switch, Back), Build my deck
                (not ready, redesign, build), the copy box's Copy and Close,
                the copy box opening when copying is truly blocked, the
                storyline panel
  Final check   Build it, Send fixes (empty and with text), the copy box when
                copying is blocked; on a second build with a copied page,
                Keep / Replace in the copied-pages table
  Finished tabs read only (no live button, a note says so)
  On their own  REVIEW.html?stay and FINAL-CHECK.html?stay, and the plain
                REVIEW.html / FINAL-CHECK.html jumping to the one page's tab

Every copy is checked by pasting what is on the clipboard into the page
(Ctrl+V), so "Copied" on screen with nothing copied fails. Any script error
on a page fails. The browser's own notice that it refused the clipboard
inside the embedded tab is expected (the older copy command is used then)
and is not counted.

Run:  py -3 slide-builder/tests/run_page_clicks_smoke.py
      SLIDE_LAB_CLICKS_SHOTS=<folder> keeps a screenshot of each stage there.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import os
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))

import run_one_page_smoke as S  # noqa: E402  (sets the Storyline tab on, copies the agent files)
import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402
import slidelab_page as SP  # noqa: E402

OPT3 = {"SLIDE_LAB_OPTIONS_PER_SLIDE": "3"}
EXPECTED_NOISE = ("Permissions policy violation", "Images loaded lazily")
RESULTS: list[tuple[str, str, bool, str]] = []


def rec(tab: str, control: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((tab, control, bool(ok), detail))


def must(r, what: str) -> None:
    assert r.returncode == 0, f"{what} failed:\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}"


# ------------------------------------------------------------------ the builds

def build_main(root: Path) -> dict:
    """7 slides, 3 options each: snapshots of the one page at each stage."""
    out = root
    brief = S.write_brief(root)
    must(S.prep(brief, out), "prep")
    page = SP.page_path(out)
    snaps = {"storyline": out / "stage_storyline.html"}
    shutil.copy2(page, snaps["storyline"])
    vis = {2: ("Bar chart over time", "suggestion"), 3: ("Table", "suggestion"),
           4: ("Line chart", "suggestion"), 5: ("Steps connected by arrows", "suggestion"),
           6: ("Designer's choice", "suggestion"), 7: ("Designer's choice", "suggestion")}
    must(H.run("record_storyline.py", "--out", out, "--message", S.message(out, visuals=vis)),
         "record_storyline")
    must(H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE, "--out", out,
               "--pattern", "direct", "--confirm-template", **OPT3), "prep with visuals")
    for n in range(1, 8):
        for L in "ABC":
            H.write_option(out, n, L)
    must(H.run("finalize_deck.py", "--out", out, "--template", H.TEMPLATE, **OPT3), "finalize")
    must(H.run("build_review.py", "--out", out, "--open", **OPT3), "build_review")
    snaps["designs"] = out / "stage_designs.html"
    shutil.copy2(page, snaps["designs"])
    snaps["review"] = out / "REVIEW.html"
    st = _state.read_state(out)
    body = _state.canonical_picks({f"slide_{n:02d}": "A" for n in range(1, 8)})
    msg = (SP.header(out, "designs") + "\n"
           + f"PICKS {body} CHECK {_state.approval_check(st['review']['token'], body)}")
    must(H.run("record_picks.py", "--out", out, "--approved", msg, **OPT3), "record_picks")
    must(H.run("finalize_deck.py", "--out", out, "--template", H.TEMPLATE, **OPT3), "finalize 2")
    must(H.run("build_review.py", "--out", out, "--final", "--open"), "final check")
    snaps["final"] = out / "stage_final.html"
    shutil.copy2(page, snaps["final"])
    snaps["final_alone"] = out / "FINAL-CHECK.html"
    snaps["page"] = page
    return snaps


def build_copied(root: Path) -> Path:
    """A 2-slide build whose slide 1 copies a supplied page: the final check
    has the copied-pages table (Keep / Replace)."""
    from pptx import Presentation
    from pptx.util import Inches
    root.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, t in enumerate(["Market reaches $300B", "Lead time is 6-12 months"]):
        s.shapes.add_textbox(Inches(0.5), Inches(0.4 + 0.8 * i), Inches(8),
                             Inches(0.6)).text_frame.text = t
    prs.save(str(root / "deck_a.pptx"))
    s1 = H.SLIDE.format(n=1).replace(
        "**Slide type:** Content",
        "**Slide type:** Content\n**Pinned source page:** deck_a.pptx slide 1").replace(
        "- Point one for slide 1", "- The market is worth about $450B")
    brief = root / "brief.md"
    brief.write_text(H.BRIEF.format(slides=s1 + H.SLIDE.format(n=2)), encoding="utf-8")
    must(H.run("seal_brief.py", "--brief", brief, "--accepted", "test"), "seal")
    out = root / "out"
    off = {"SLIDE_LAB_STORYLINE_TAB": "0"}
    must(H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE, "--out", out,
               "--pattern", "direct", "--confirm-template", **off), "prep (copied page)")
    body1 = H.OPTION_SCRIPT.format(n=1, L="A", skill=str(H.SKILL)).replace(
        "Point one\\nPoint two\\nPoint three", "The market reaches $300B\\nLead time is 6-12 months")
    H.write_option(out, 1, body=body1)
    H.write_option(out, 2)
    must(H.run("finalize_deck.py", "--out", out, "--template", H.TEMPLATE, **off), "finalize")
    must(H.run("build_review.py", "--out", out, **off), "review")
    tok = _state.read_state(out)["review"]["token"]
    body = _state.canonical_picks({"slide_01": "A", "slide_02": "A"})
    must(H.run("record_picks.py", "--out", out, "--approved",
               f"PICKS {body} CHECK {_state.approval_check(tok, body)}", **off), "picks")
    must(H.run("finalize_deck.py", "--out", out, "--template", H.TEMPLATE, **off), "finalize 2")
    must(H.run("build_review.py", "--out", out, "--final", **off), "final check")
    return SP.page_path(out)


# ------------------------------------------------------------------ the browser

def _launch(pw):
    """Edge first (the owner's browser), as a double-clicked file opens: no
    extra switches. Chrome or the bundled Chromium only when Edge is missing."""
    for kw in ({"channel": "msedge"}, {"channel": "chrome"}, {}):
        try:
            return pw.chromium.launch(headless=True, **kw), kw.get("channel") or "chromium"
        except Exception:
            continue
    raise RuntimeError("no browser could start (Edge, Chrome or Playwright's Chromium)")


class Page:
    """One opened file: its console errors and the clipboard checks."""

    def __init__(self, browser, url: str, webdriver: bool = True):
        self.ctx = browser.new_context(viewport={"width": 1400, "height": 900})
        if not webdriver:
            self.ctx.add_init_script(
                "Object.defineProperty(navigator, 'webdriver', {get: () => false});")
        self.p = self.ctx.new_page()
        self.p.set_default_timeout(8000)
        self.errors: list[str] = []
        self.dialogs: list[str] = []
        self.p.on("pageerror", lambda e: self.errors.append(f"script error: {e}"))
        self.p.on("console", self._console)
        self.p.on("dialog", self._dialog)
        self.p.goto(url)
        self.p.wait_for_timeout(600)

    def _console(self, m) -> None:
        if m.type == "error" and not any(x in m.text for x in EXPECTED_NOISE):
            self.errors.append(f"console: {m.text[:200]}")

    def _dialog(self, d) -> None:
        self.dialogs.append(d.message[:80])
        d.accept()

    def frame(self, selector: str):
        for f in self.p.frames:
            try:
                if f.query_selector(selector):
                    return f
            except Exception:
                continue
        return None

    def sentinel(self) -> None:
        """Put a known text on the clipboard with a real Ctrl+C, so a later
        paste proves what the page copied (not something from before)."""
        self.p.evaluate(self._PROBE + "t.value = 'SENTINEL'; t.focus(); t.select(); }")
        self.p.keyboard.press("Control+C")

    # The probe box goes inside an open modal box when there is one (the rest
    # of the page cannot take focus then).
    _PROBE = """() => { let t = document.getElementById('__clk');
          if (!t) { t = document.createElement('textarea'); t.id = '__clk';
            t.style.cssText = 'position:fixed;left:0;bottom:0;width:40px;height:20px;z-index:2147483647'; }
          (document.querySelector('dialog[open]') || document.body).appendChild(t); """

    def clipboard(self) -> str:
        self.p.evaluate(self._PROBE + "t.value = ''; t.focus(); }")
        self.p.keyboard.press("Control+V")
        return self.p.evaluate("() => document.getElementById('__clk').value")

    def close(self) -> None:
        self.ctx.close()


def copied(pg: Page, act, *needles) -> tuple[bool, str]:
    """Run `act` (a click), then check the clipboard holds every needle."""
    pg.sentinel()
    act()
    pg.p.wait_for_timeout(350)
    got = pg.clipboard()
    miss = [n for n in needles if n not in got]
    if got == "SENTINEL":
        return False, "nothing was copied"
    return (not miss), ("copied: " + got.splitlines()[0][:70] if not miss
                        else f"copied text lacks {miss}: {got[:120]!r}")


# ------------------------------------------------------------------ storyline tab

def check_storyline(browser, snap: Path, shots: Path | None) -> None:
    T = "Storyline"
    pg = Page(browser, snap.as_uri())
    p = pg.p
    for tab in ("designs", "final", "storyline"):
        p.click(f'.tab[data-p="{tab}"]')
        on = p.evaluate(f"document.getElementById('{tab}').classList.contains('on')")
        rec("All tabs", f"tab switch to {tab}", on, "pane shown" if on else "pane did not show")
    el = p.locator('li[contenteditable="true"]').first
    el.click()
    p.keyboard.press("End")
    p.keyboard.type(" EDITED-BY-TEST")
    rec(T, "edit a line", "edited" in (el.get_attribute("class") or ""), "line turns yellow")
    p.locator(".cbtn").first.click()
    opened = p.evaluate("document.querySelector('.comment').classList.contains('open')")
    p.locator(".ctext").first.fill("merge this with slide 5")
    rec(T, "+ Comment on this slide", opened, "comment box opens")
    sels = p.locator(".vsel")
    n_sel = sels.count()
    ok_all, bad = True, []
    for i in range(n_sel):
        s = sels.nth(i)
        opts = s.locator("option").all_inner_texts()
        cur = s.input_value()
        want = next(o for o in opts if o != cur)
        before = p.evaluate(f"document.querySelectorAll('.skw')[{i}].innerHTML")
        s.select_option(label=want)
        after = p.evaluate(f"document.querySelectorAll('.skw')[{i}].innerHTML")
        if s.input_value() != want or (before == after and want != cur):
            ok_all = False
            bad.append(i + 1)
    mix = p.evaluate("(document.getElementById('mixmsg') || {}).textContent || ''")
    rec(T, f"visual select (all {n_sel})", ok_all and n_sel > 0 and bool(mix),
        f"sketch and mix update; mix says: {mix[:50]}" if ok_all else f"no update on box {bad}")
    must_lbl = p.locator(".lock").first.locator("label").nth(1)
    must_lbl.click()
    m1 = p.locator(".lock").first.locator("input").nth(1).is_checked()
    p.locator(".lock").first.locator("label").nth(0).click()
    s1 = p.locator(".lock").first.locator("input").nth(0).is_checked()
    must_lbl.click()
    rec(T, "Suggestion / Must use", m1 and s1, "both choices take")
    ok, why = copied(pg, lambda: p.click("#done"), "Storyline, code", "EDITED-BY-TEST",
                     "merge this with slide 5", "(must use)")
    hint = p.evaluate("document.getElementById('outhint').textContent")
    rec(T, "Done: copy my changes", ok and hint.startswith("Copied"), why)
    if shots:
        p.screenshot(path=str(shots / "M_clicks_storyline.png"))
    ok, why = copied(pg, lambda: p.click("#all"), "the designers pick every visual")
    rec(T, "Let the designers pick all the visuals", ok, why)
    rec(T, "no script errors", not pg.errors, "; ".join(pg.errors[:3]))
    pg.close()


# ------------------------------------------------------------------ designs tab

def check_designs(browser, url: str, where: str, shots: Path | None) -> None:
    T = f"Slide designs{where}"
    pg = Page(browser, url)
    p = pg.p
    if "#designs" in url:
        p.click('.tab[data-p="designs"]')
    f = pg.frame("#btn-build")
    if f is None:
        rec(T, "page loads", False, "no Build my deck button found")
        pg.close()
        return
    ev = f.evaluate
    for L in "ABC":
        f.click(f'#card-slide_01 button.pick[data-letter="{L}"]')
        rec(T, f"Pick {L}", ev("pickForSlide('slide_01')") == L,
            "slide 1 picked " + str(ev("pickForSlide('slide_01')")))
    f.click('#card-slide_01 button.pick[data-letter="C"]')
    rec(T, "Pick again clears", not ev("pickForSlide('slide_01')"), "pick cleared")
    f.click('#card-slide_01 .option[data-letter="B"] .option-frame')
    rec(T, "click the design picture", ev("pickForSlide('slide_01')") == "B", "picture picks it")
    f.click("#card-slide_02 button.none:not(.three)")
    rec(T, "Replace these", ev("isReplace('slide_02') && !isThree('slide_02')")
        and f.is_visible("#regen-slide_02"), "marked to replace, note shows")
    f.click("#card-slide_03 button.none.three")
    rec(T, "+3 designs", ev("isThree('slide_03')"), "three new designs asked")
    f.click('#card-slide_04 button.pick.leave')
    rec(T, "Leave out", ev("pickForSlide('slide_04')") == "-", "slide left out")
    chips = f.locator("#card-slide_05 .chip")
    n_ch = chips.count()
    for i in range(n_ch):
        chips.nth(i).click()
    quick = ev("fbFor('slide_05').quick || ''")
    rec(T, f"quick-feedback chips (all {n_ch})",
        n_ch > 0 and len(quick.split("; ")) == n_ch
        and f.locator("#card-slide_05 .chip.selected").count() == n_ch,
        f"{len(quick.split('; '))} of {n_ch} recorded")
    boxes = f.locator('#card-slide_05 textarea[data-field]')
    n_box = boxes.count()
    for i in range(n_box):
        boxes.nth(i).fill(f"note {i} from the test")
    fb = ev("fbFor('slide_05')")
    rec(T, f"text boxes (all {n_box})",
        n_box > 0 and all(fb.get(boxes.nth(i).get_attribute("data-field")) == f"note {i} from the test"
                          for i in range(n_box)), "each box recorded")
    pg.p.wait_for_timeout(1900)  # let an earlier note fade
    f.click("#btn-build")
    pg.p.wait_for_timeout(200)
    toast = ev("document.getElementById('toast').textContent")
    rec(T, "Build my deck (slides not decided)", "still need a pick" in toast, f"note: {toast[:60]}")
    for n in range(5, 8):
        f.click(f'#card-slide_0{n} button.pick[data-letter="A"]')
    ok, why = copied(pg, lambda: f.click("#btn-build"), "Slide designs, code",
                     "Redesign some slides", "New designs for: slide 2, slide 3",
                     "Three designs for: slide 3", "note 0 from the test")
    rec(T, "Build my deck (with Replace: redesign message)",
        ok and not ev("document.getElementById('picks-dialog').open"), why)
    f.click("#card-slide_02 button.pick[data-letter='A']")
    f.click("#card-slide_03 button.pick[data-letter='A']")
    ok, why = copied(pg, lambda: f.click("#btn-build"), "Slide designs, code",
                     "Build my Slide Lab deck.", "Picks: 1B 2A 3A 4- 5A 6A 7A (check ")
    dlg_open = ev("document.getElementById('picks-dialog').open")
    rec(T, "Build my deck", ok and not dlg_open,
        why + ("" if not dlg_open else "; the copy box opened anyway"))
    if shots and not where:
        p.screenshot(path=str(shots / "M_clicks_designs.png"))
    f.click("#btn-all")
    on = ev("document.body.classList.contains('all-mode')")
    f.click('#card-slide_01 .keep-toggle[data-letter="C"]')
    off = ev("leftOut('slide_01').indexOf('C') >= 0")
    rec(T, "All options in one deck (first click, keep switch)", on and off,
        "switches show; C left out")
    ok, why = copied(pg, lambda: f.click("#btn-all"), "Build all options in one Slide Lab deck",
                     "Picks: 1AB ")
    rec(T, "All options in one deck (copy)", ok and any("converts every option" in d
                                                        for d in pg.dialogs), why)
    f.click("#btn-all-cancel")
    rec(T, "Back to one pick per slide", not ev("document.body.classList.contains('all-mode')"),
        "normal mode again")
    # the copy box, when copying is truly blocked
    ev("""() => { window.__realExec = document.execCommand.bind(document);
      document.execCommand = () => false;
      if (navigator.clipboard) navigator.clipboard.writeText = () => Promise.reject(new Error('blocked')); }""")
    f.click("#btn-build")
    pg.p.wait_for_timeout(300)
    dlg = ev("document.getElementById('picks-dialog').open")
    body = ev("document.getElementById('picks-dlg-body').textContent")
    rec(T, "copy box when copying is blocked", dlg and "Picks: 1B" in body,
        "the box opens with the message" if dlg else "no box: the user gets nothing")
    ev("() => { document.execCommand = window.__realExec; }")
    ok, why = copied(pg, lambda: f.click("#btn-dlg-copy"), "Build my Slide Lab deck.", "Picks: 1B")
    rec(T, "copy box: Copy", ok, why)
    f.click("#btn-dlg-close")
    rec(T, "copy box: Close", not ev("document.getElementById('picks-dialog').open"), "box closed")
    f.click(".storyline-section summary")
    o1 = ev("document.querySelector('.storyline-section').open")
    f.click(".storyline-section summary")
    o2 = ev("document.querySelector('.storyline-section').open")
    rec(T, "storyline panel open / close", o1 and not o2, "opens and closes")
    f.click("#btn-clear")
    pg.p.wait_for_timeout(200)
    cleared = ev("SLIDE_IDS.every(s => !pickForSlide(s) && !isReplace(s))") and \
        f.locator(".chip.selected").count() == 0
    rec(T, "Clear", cleared, "picks, replace requests and notes cleared")
    rec(T, "no script errors", not pg.errors, "; ".join(pg.errors[:3]))
    pg.close()


# ------------------------------------------------------------------ final check

def check_final(browser, url: str, where: str, shots: Path | None) -> None:
    T = f"Final check{where}"
    pg = Page(browser, url)
    p = pg.p
    if "#final" in url:
        p.click('.tab[data-p="final"]')
    f = pg.frame("#btn-go")
    if f is None:
        rec(T, "page loads", False, "no Build it button found")
        pg.close()
        return
    ev = f.evaluate
    ok, why = copied(pg, lambda: f.click("#btn-go"), "Final check, code",
                     "Build it: I checked the finished slides.", "Final check: ")
    rec(T, "Build it", ok and not f.is_visible("#copybox"), why)
    pg.p.wait_for_timeout(1900)
    f.click("#btn-fix")
    pg.p.wait_for_timeout(200)
    toast = ev("document.getElementById('toast').textContent")
    rec(T, "Send fixes (box empty)", "Say what needs fixing first" in toast, f"note: {toast[:50]}")
    f.fill("#fix", "slide 2 title is cut off")
    ok, why = copied(pg, lambda: f.click("#btn-fix"), "Final check, code",
                     "Fix these before building", "slide 2 title is cut off")
    rec(T, "Send fixes", ok and not f.is_visible("#copybox"), why)
    if shots and not where:
        p.screenshot(path=str(shots / "M_clicks_final.png"))
    ev("""() => { document.execCommand = () => false;
      if (navigator.clipboard) navigator.clipboard.writeText = () => Promise.reject(new Error('blocked')); }""")
    f.click("#btn-go")
    pg.p.wait_for_timeout(300)
    shown = f.is_visible("#copybox") and "Build it:" in ev("document.getElementById('copytext').value")
    rec(T, "copy box when copying is blocked", shown, "box shows the message" if shown else "no box")
    rec(T, "no script errors", not pg.errors, "; ".join(pg.errors[:3]))
    pg.close()


def check_copied_table(browser, page: Path) -> None:
    T = "Final check (copied page)"
    pg = Page(browser, page.as_uri() + "#final")
    f = pg.frame("#btn-go")
    if f is None or not f.query_selector("tr.cp-row"):
        rec(T, "copied-pages table", False, "no table rows on the page")
        pg.close()
        return
    rows = f.locator("tr.cp-row")
    n = rows.count()
    pg.sentinel()
    f.click("#btn-go")
    pg.p.wait_for_timeout(250)
    toast = f.evaluate("document.getElementById('toast').textContent")
    rec(T, "Build it before every row is answered", "Answer every row" in toast
        and pg.clipboard() == "SENTINEL", f"note: {toast[:50]}")
    for i in range(n):
        rows.nth(i).locator("input[type=radio][value=keep]").check()
    ok, why = copied(pg, lambda: f.click("#btn-go"), "Build it:", f"Copied-page figures: all {n} kept")
    rec(T, f"Keep (all {n} rows) then Build it", ok, why)
    pg.p.wait_for_timeout(1900)
    rows.nth(0).locator("input[type=radio][value=replace]").check()
    w = rows.nth(0).locator("input.cp-with")
    if w.count():
        w.fill("$450B")
    ok, why = copied(pg, lambda: f.click("#btn-go"), "figures copied from the earlier page to replace",
                     "$450B" if w.count() else "slide 1")
    rec(T, "Replace a row then Build it", ok, why)
    rec(T, "no script errors", not pg.errors, "; ".join(pg.errors[:3]))
    pg.close()


def check_read_only(browser, snap: Path) -> None:
    pg = Page(browser, snap.as_uri())
    p = pg.p
    live = p.evaluate("document.querySelectorAll('#storyline [contenteditable=\"true\"], #storyline #done').length")
    rec("Storyline (finished)", "read only", live == 0, f"{live} live control(s)")
    p.click('.tab[data-p="designs"]')
    f = pg.frame("#btn-build")
    if f is None:
        rec("Slide designs (finished)", "read only", False, "tab missing")
    else:
        n_live = f.evaluate("document.querySelectorAll('button:not(:disabled), textarea:not(:disabled)').length")
        note = f.evaluate("(document.querySelector('.slidelab-ro-banner') || {}).textContent || ''")
        rec("Slide designs (finished)", "read only", n_live == 0 and "read only" in note,
            f"{n_live} live control(s); note: {note[:40]}")
    rec("Finished tabs", "no script errors", not pg.errors, "; ".join(pg.errors[:3]))
    pg.close()


def check_redirects(browser, review: Path, final: Path, page: Path) -> None:
    for f, tab in ((review, "designs"), (final, "final")):
        pg = Page(browser, f.as_uri(), webdriver=False)
        url = pg.p.url
        ok = url.endswith(page.name.replace(" ", "%20") + "#" + tab)
        rec("On their own", f"{f.name} opened plainly jumps to the one page", ok, url[-60:])
        pg.close()


# ------------------------------------------------------------------ main

def main() -> int:
    shots = os.environ.get("SLIDE_LAB_CLICKS_SHOTS")
    shots = Path(shots) if shots else None
    if shots:
        shots.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="page_clicks_"))
    fails = 0
    try:
        print("[1] build a real small deck to each stage")
        snaps = build_main(tmp / "deck")
        cp_page = build_copied(tmp / "cp")
        print("    ok: Storyline, Slide designs, Final check, and a build with a copied page")
        from playwright.sync_api import sync_playwright
        import _browser
        with _browser.scratch_env(), sync_playwright() as pw:
            browser, name = _launch(pw)
            print(f"[2] click every control ({name}, headless, its own browser)")
            try:
                check_storyline(browser, snaps["storyline"], shots)
                check_designs(browser, snaps["designs"].as_uri() + "#designs", "", shots)
                check_designs(browser, snaps["review"].as_uri() + "?stay", " (REVIEW.html?stay)", None)
                check_final(browser, snaps["final"].as_uri() + "#final", "", shots)
                check_final(browser, snaps["final_alone"].as_uri() + "?stay",
                            " (FINAL-CHECK.html?stay)", None)
                check_copied_table(browser, cp_page)
                check_read_only(browser, snaps["final"])
                check_redirects(browser, snaps["review"], snaps["final_alone"], snaps["page"])
            finally:
                browser.close()
    except AssertionError as exc:
        fails += 1
        import traceback
        traceback.print_exc()
        print(f"  FAIL: {exc}")
    finally:
        if os.environ.get("SLIDE_LAB_KEEP_TEST_TEMP") != "1":
            shutil.rmtree(tmp, ignore_errors=True)
        else:
            print(f"  kept: {tmp}")
    w = max([len(r[0]) for r in RESULTS] + [4])
    c = max([len(r[1]) for r in RESULTS] + [7])
    for tab, control, ok, detail in RESULTS:
        print(f"  {'ok  ' if ok else 'DEAD'} {tab:<{w}}  {control:<{c}}  {detail}")
    dead = [r for r in RESULTS if not r[2]]
    fails += len(dead)
    print("=" * 72)
    print(f"{len(RESULTS)} checks, {len(dead)} failed")
    print("SMOKE PASSED" if not fails else f"SMOKE FAILED: {fails}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
