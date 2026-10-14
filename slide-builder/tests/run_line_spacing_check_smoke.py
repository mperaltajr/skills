#!/usr/bin/env python3
"""Smoke test: the line-spacing file check (type_scale.line_spacing_slide).

2026-10-13 report item 10: the converter wrote exact line spacing below the
font size on 439 of 545 paragraphs, and nothing in the pictures showed it.
The check (owner-approved 2026-10-14, speed verdict Q2 item 6) flags exact
spacing tighter than the font needs, and is calibrated so the converter's
deliberate one-line spacing (0.8 x the size) with its top inset (2026-10-13)
does not fire. Made-up content.

  1. lines that meet: a wrapped paragraph at 9 pt spacing on 12 pt text is a
     Major; at 14 pt it passes; a one-line box at 0.8 x passes this rule
  2. paragraphs: 10 pt spacing on 12 pt text with 6 pt before each paragraph
     passes; with no space between them it is a Major
  3. capitals above the box: a one-line box at 0.8 x with no top inset is an
     Advisory; with translate_html.top_pad_px as its inset it passes
  4. slide-qc's hygiene check carries both (Major "line-spacing", fixed
     automatically as an overlap; Advisory), and qc_fix_policy maps the tag
  5. the converter's own output passes (needs Playwright's Chromium): a
     design with one-line labels, a hero number, wrapped body text and a list

Run:  py -3 slide-builder/tests/run_line_spacing_check_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

import type_scale as TS  # noqa: E402
import translate_html as T  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
LONG = ("Revenue grows in every year of the plan because each site adds capacity "
        "and every customer group buys more of it")


def _box(slide, x, y, w, h, paras, size=12, spacing=None, before=None, tins=0):
    tb = slide.shapes.add_textbox(Emu(x * 9525), Emu(y * 9525), Emu(w * 9525), Emu(h * 9525))
    tf = tb.text_frame
    tf.word_wrap = True
    body = tf._txBody.find(A + "bodyPr")
    for k in ("lIns", "rIns", "bIns"):
        body.set(k, "0")
    body.set("tIns", str(int(tins)))
    body.set("anchor", "t")
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.name = "Arial"
        ppr = p._p.get_or_add_pPr()
        if spacing is not None:
            ln = etree.SubElement(ppr, A + "lnSpc")
            etree.SubElement(ln, A + "spcPts").set("val", str(int(spacing * 100)))
        if before is not None:
            sb = etree.SubElement(ppr, A + "spcBef")
            etree.SubElement(sb, A + "spcPts").set("val", str(int(before * 100)))
    return tb


def _slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def main() -> int:
    print("[1] lines that meet")
    _p, s = _slide()
    _box(s, 50, 50, 200, 80, [LONG], 12, spacing=9, tins=7200)
    r = TS.line_spacing_slide(s)
    assert len(r["touch"]) == 1 and r["touch"][0][:2] == (9.0, 12.0), r
    assert TS.line_spacing_problems(r) and "meet the next" in TS.line_spacing_problems(r)[0]
    _p, s = _slide()
    _box(s, 50, 50, 200, 80, [LONG], 12, spacing=14)
    assert TS.line_spacing_slide(s)["touch"] == []
    _p, s = _slide()
    _box(s, 50, 50, 600, 30, ["One short label"], 12, spacing=10,
         tins=T.top_pad_px({"font": "Arial", "bold": False, "size_px": 16}, 10 / 0.75) * 9525)
    assert TS.line_spacing_slide(s) == {"touch": [], "caps_above": []}, TS.line_spacing_slide(s)
    print("    ok: 9 pt on 12 pt wrapped flagged; 14 pt and a one-line 0.8 x box pass")

    print("[2] paragraphs")
    _p, s = _slide()
    _box(s, 50, 50, 600, 120, ["First point", "Second point", "Third point"], 12,
         spacing=10, before=6)
    assert TS.line_spacing_slide(s)["touch"] == [], TS.line_spacing_slide(s)
    _p, s = _slide()
    _box(s, 50, 50, 600, 120, ["First point", "Second point", "Third point"], 12, spacing=10)
    assert len(TS.line_spacing_slide(s)["touch"]) == 2, TS.line_spacing_slide(s)
    print("    ok: space before keeps them apart; without it the 2 later lines are flagged")

    print("[3] capitals above the box")
    _p, s = _slide()
    _box(s, 50, 50, 600, 30, ["One short label"], 12, spacing=10, tins=0)
    r = TS.line_spacing_slide(s)
    assert r["touch"] == [] and len(r["caps_above"]) == 1, r
    assert TS.line_spacing_problems(r) == [] and "above the box" in TS.line_spacing_notes(r)[0]
    assert len(T.caps_above_frame(s)) == 1, "same rule as the converter's advisory"
    print(f"    ok: {r['caps_above'][0][0]} px above without the inset (Advisory)")

    print("[4] the hygiene check and the fix policy")
    import check_pptx_hygiene as H
    import qc_fix_policy as P
    with tempfile.TemporaryDirectory() as td:
        prs, s = _slide()
        _box(s, 50, 50, 200, 80, [LONG], 12, spacing=9, tins=7200)
        _box(s, 50, 300, 600, 30, ["One short label"], 12, spacing=10, tins=0)
        f = Path(td) / "deck.pptx"
        prs.save(str(f))
        v = [x for x in H.run_all_checks(f)["violations"] if x["category"] == "line-spacing"]
        sev = sorted(x["severity"] for x in v)
        assert sev == ["Advisory", "Major"], v
    assert P.normalize_tag("line-spacing") == "overlap"
    major = next(x for x in v if x["severity"] == "Major")
    assert P.decide(f"slide 1 [line-spacing] Major: {major['issue']}") == ("auto", "overlap")
    print("    ok: Major + Advisory, tag maps to overlap")

    print("[5] the converter's own output passes")
    design = """<!doctype html><html><head><meta charset="utf-8"><style>
body { margin: 0; }
.slide-canvas { position: relative; width: 1280px; height: 720px; background: #fff;
  font-family: Arial, sans-serif; color: #1a1a1a; overflow: hidden; }
.a { position: absolute; margin: 0; }
</style></head><body><div class="slide-canvas">
<p class="a" data-shape-id="hero" style="left:60px;top:80px;font-size:64px;font-weight:700">$4.8B</p>
<p class="a" data-shape-id="label" style="left:60px;top:200px;font-size:16px">Total by year five</p>
<p class="a" data-shape-id="small" style="left:60px;top:240px;font-size:12px">Source: made-up figures</p>
<p class="a" data-shape-id="body" style="left:400px;top:80px;width:300px;font-size:16px;line-height:22px">
Two lines of body text that wrap inside their box as designed</p>
<ul class="a" data-shape-id="pts" style="left:400px;top:220px;width:300px;font-size:16px;line-height:22px">
<li>First point</li><li>Second point that is long enough to wrap onto a second line</li></ul>
</div></body></html>"""
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        html = td / "option_A.html"
        html.write_text(design, encoding="utf-8")
        T.translate_many([(html, td, "A", True)], check=False)
        plan = json.loads((td / "option_A_native.plan.json").read_text(encoding="utf-8"))
        from twins.html_emit import build
        _prs, slide = build(plan)
        r = TS.line_spacing_slide(slide)
        assert r == {"touch": [], "caps_above": []}, r
        tight = [p for o in plan["ops"] if o["op"] == "text" for p in [o]
                 if o.get("line_height_px") and o["paragraphs"][0]
                 and o["line_height_px"] < T._snapped_px(o["paragraphs"][0][0]["size_px"])]
        assert tight, "the design should include converter one-line boxes set below the size"
    print(f"    ok: no finding; {len(tight)} box(es) set below their size on purpose pass")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
