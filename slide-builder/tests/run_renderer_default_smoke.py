#!/usr/bin/env python3
"""Smoke test: PowerPoint is the default renderer on Windows, LibreOffice the
backup, and each takes over when the other fails twice (owner's decision,
2026-10-09).

  [1] the choice: auto = PowerPoint on Windows with PowerPoint, LibreOffice
      without it (SLIDE_LAB_NO_POWERPOINT=1 stands in for a Mac);
      settings.json / SLIDE_LAB_RENDERER "libreoffice" or "powerpoint" start
      with that one; a program that is not installed gives way to the other
  [2] slide-qc's render_slides.py draws through PowerPoint by default, with
      LibreOffice installed and with it turned off (SLIDE_LAB_NO_LIBREOFFICE=1)
  [3] LibreOffice fails twice -> PowerPoint takes over (exactly two
      LibreOffice tries); PowerPoint fails twice -> LibreOffice takes over
      (exactly two PowerPoint tries); both fail -> one error naming both
  [4] the one-line PowerPoint notice is printed once per command, however
      many renders the command runs, and not again by its child processes

Fictional files only; the user's open PowerPoint presentations are recorded
before and after (attaching only) and must be unchanged.
Skips the PowerPoint parts (still prints SMOKE PASSED) where PowerPoint can't
be driven.
Run:  py -3 slide-builder/tests/run_renderer_default_smoke.py
"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
QC = SKILL.parent / "slide-qc" / "scripts"
for p in (str(SKILL), str(SKILL / "scripts"), str(QC)):
    sys.path.insert(0, p)

import render_slides as RS  # noqa: E402
from _ppt_snapshot import snapshot  # noqa: E402


def _deck(path: Path, n: int = 2) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        tb.text_frame.text = f"Fictional slide {i + 1}"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


class _Env:
    def __init__(self, **kv):
        self.kv = kv

    def __enter__(self):
        self.saved = {k: os.environ.get(k) for k in self.kv}
        for k, v in self.kv.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v

    def __exit__(self, *a):
        for k, v in self.saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def _run(args, **env):
    e = {**os.environ, "PYTHONIOENCODING": "utf-8", "SLIDE_LAB_NO_OPEN": "1"}
    e.pop("SLIDE_LAB_PPT_NOTICE_SHOWN", None)
    e.update(env)
    r = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True,
                       env=e, timeout=600, encoding="utf-8", errors="replace")
    assert r.returncode == 0, f"{args}\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}"
    return r


def main() -> int:
    before = snapshot()
    has_ppt, has_lo = RS.powerpoint_available(), RS.libreoffice_available()
    print(f"  PowerPoint installed: {has_ppt}; LibreOffice installed: {has_lo}")

    print("[1] which program starts")
    with _Env(SLIDE_LAB_RENDERER=None):
        auto = RS.pick_renderer()
    if sys.platform == "win32" and has_ppt:
        assert auto == "powerpoint", auto
    elif has_lo:
        assert auto == "libreoffice", auto
    with _Env(SLIDE_LAB_RENDERER=None, SLIDE_LAB_NO_POWERPOINT="1"):
        assert RS.pick_renderer() == "libreoffice"          # a Mac
    if has_lo:
        with _Env(SLIDE_LAB_RENDERER="libreoffice"):
            assert RS.pick_renderer() == "libreoffice"
    if has_ppt:
        with _Env(SLIDE_LAB_RENDERER="powerpoint"):
            assert RS.pick_renderer() == "powerpoint"
        with _Env(SLIDE_LAB_RENDERER="libreoffice", SLIDE_LAB_NO_LIBREOFFICE="1"):
            assert RS.pick_renderer() == "powerpoint", "a missing LibreOffice must give way"
    with _Env(SLIDE_LAB_RENDERER="powerpoint", SLIDE_LAB_NO_POWERPOINT="1"):
        assert RS.pick_renderer() == "libreoffice", "a missing PowerPoint must give way"
    # the settings file is read when no override is set
    import json
    assert json.loads(RS.SETTINGS_JSON.read_text(encoding="utf-8"))["renderer"] == "auto"
    real_settings = RS.SETTINGS_JSON
    with tempfile.TemporaryDirectory() as tdn:
        fake = Path(tdn) / "settings.json"
        fake.write_text('{"renderer": "libreoffice"}', encoding="utf-8")
        RS.SETTINGS_JSON = fake
        try:
            with _Env(SLIDE_LAB_RENDERER=None):
                assert RS.renderer_setting() == "libreoffice"
        finally:
            RS.SETTINGS_JSON = real_settings
    print(f"    ok: auto -> {auto}; overrides and give-way rules hold")

    with tempfile.TemporaryDirectory() as tdn:
        td = Path(tdn)
        deck = _deck(td / "fictional.pptx")

        if has_ppt:
            print("[2] slide-qc renders through PowerPoint by default")
            r = _run([QC / "render_slides.py", deck, td / "a"])
            assert "(PowerPoint" in r.stdout and len(list((td / "a").glob("slide_*.png"))) == 2, r.stdout
            r = _run([QC / "render_slides.py", deck, td / "b"], SLIDE_LAB_NO_LIBREOFFICE="1")
            assert "(PowerPoint" in r.stdout and len(list((td / "b").glob("slide_*.png"))) == 2, r.stdout
            if has_lo:
                r = _run([QC / "render_slides.py", deck, td / "c"], SLIDE_LAB_RENDERER="libreoffice")
                assert "(LibreOffice" in r.stdout, r.stdout
            print("    ok: PowerPoint with LibreOffice present and with it missing; "
                  "the setting switches to LibreOffice")

        if has_ppt and has_lo:
            print("[3] a program that fails twice hands over to the other")
            tries = []
            real_run = RS._run_soffice

            def lo_fails(cmd):
                tries.append(1)
                return subprocess.CompletedProcess(cmd, 3, "", "simulated LibreOffice failure")
            RS._run_soffice = lo_fails
            try:
                eng = RS.render(deck, td / "f1", 40, renderer="libreoffice", quiet=True)
            finally:
                RS._run_soffice = real_run
            assert eng == "powerpoint" and len(tries) == 2, (eng, tries)
            assert len(list((td / "f1").glob("slide_*.png"))) == 2
            import ppt_safe
            ptries = []
            real_exp, real_pdf = ppt_safe.export_pngs, ppt_safe.export_pdf

            def ppt_fails(*a, **k):
                ptries.append(1)
                raise RuntimeError("simulated PowerPoint failure")
            ppt_safe.export_pngs = ppt_fails
            ppt_safe.export_pdf = ppt_fails
            try:
                eng = RS.render(deck, td / "f2", 40, renderer="powerpoint", quiet=True)
                assert eng == "libreoffice" and len(ptries) == 2, (eng, ptries)
                ptries.clear()
                pdf, eng = RS.to_pdf(deck, td / "f3", renderer="powerpoint")
                assert eng == "libreoffice" and pdf.exists() and len(ptries) == 2
                RS._run_soffice = lo_fails
                try:
                    RS.render(deck, td / "f4", 40, renderer="powerpoint", quiet=True)
                    raise AssertionError("both failing raised nothing")
                except RuntimeError as exc:
                    msg = str(exc)
                finally:
                    RS._run_soffice = real_run
            finally:
                ppt_safe.export_pngs, ppt_safe.export_pdf = real_exp, real_pdf
            assert "simulated PowerPoint failure" in msg and "simulated LibreOffice failure" in msg, msg
            print("    ok: LibreOffice x2 -> PowerPoint; PowerPoint x2 -> LibreOffice; "
                  "both failing -> one error with both reasons")

        if has_ppt:
            print("[4] the PowerPoint notice, once per command")
            script = td / "three_renders.py"
            script.write_text(
                "import sys, subprocess\n"
                f"sys.path.insert(0, r'{QC}')\n"
                "import render_slides as RS\n"
                f"for k in range(3):\n"
                f"    RS.render(r'{deck}', r'{td}' + f'/n{{k}}', 40, quiet=True)\n"
                "subprocess.run([sys.executable, r'" + str(QC / "render_slides.py") + "', r'"
                + str(deck) + "', r'" + str(td / "child") + "'], check=True)\n",
                encoding="utf-8")
            r = _run([script])
            both = r.stdout + r.stderr
            assert both.count(RS.NOTICE) == 1, both
            assert "editing in PowerPoint" in RS.NOTICE
            print("    ok: 3 renders + a child render, the notice printed once")

    after = snapshot()
    assert before[1] == after[1], "the user's open presentations changed"
    if before[0]:
        assert after[0], "PowerPoint was running before and is not now"
    print(f"  ok: the user's PowerPoint: {len(before[1])} presentation(s) open before, "
          f"the same {len(after[1])} after")
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
