#!/usr/bin/env python3
"""Smoke test: no grafted slide may ship two shapes claiming the same id.

Background. A shape's id (<p:cNvPr id="...">) must be unique within its slide.
Both graft paths build a slide by deepcopying elements out of source slides that
each numbered their own shapes from 1, so an untouched graft routinely lands
three or four shapes all claiming id "2".

LibreOffice renders that file. python-pptx reads it. check_pptx_hygiene passes
it. PowerPoint says "found a problem with content", offers to repair, and then
refuses to open the deck. An 87-slide deck shipped twice in that state because
every offline check was green.

There are TWO graft sites, and a fix to one leaves the other shipping the defect:
  1. scripts/finalize_deck.py  (per-option themed pptx)
  2. scripts/compile_picks.py  (assembling picked options into the deck)
plus the package check that has to catch a hand-edit made after compile.

Run:  py -3 slide-builder/tests/run_shape_id_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import inspect
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from twins.composer import (  # noqa: E402
    find_duplicate_shape_ids,
    reassign_shape_ids,
)


def _ids(slide) -> list[str]:
    return [c.get("id") for c in slide.shapes._spTree.iter(qn("p:cNvPr"))]


def main() -> int:
    print("1. a slide carrying the graft's id collision")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(4):
        slide.shapes.add_textbox(Inches(0.5 + i), Inches(1), Inches(1), Inches(1))
    # Reproduce what four deepcopied source trees leave behind.
    for cNvPr in list(slide.shapes._spTree.iter(qn("p:cNvPr")))[1:]:
        cNvPr.set("id", "2")
    assert find_duplicate_shape_ids(prs), "the check failed to see a real collision"
    print(f"    ok: collision detected — {find_duplicate_shape_ids(prs)[0]}")

    print("2. reassign resolves it")
    changed = reassign_shape_ids(slide)
    after = _ids(slide)
    assert changed == 3, f"expected 3 shapes moved, got {changed}"
    assert len(set(after)) == len(after), f"ids still collide after reassign: {after}"
    assert find_duplicate_shape_ids(prs) == []
    assert after[0] == "1", "the root group's id must not be reassigned"
    print(f"    ok: {after}")

    print("3. reassign is idempotent and leaves unique ids where they were")
    snapshot = list(after)
    assert reassign_shape_ids(slide) == 0, "a second pass moved shapes that were already unique"
    assert _ids(slide) == snapshot, "a second pass renumbered unique ids"
    print("    ok: no churn on a clean slide")

    print("4. both graft sites call it")
    import finalize_deck
    import compile_picks
    fsrc = inspect.getsource(finalize_deck.graft_and_theme)
    assert "reassign_shape_ids" in fsrc, (
        "finalize_deck's graft no longer renumbers shape ids; every themed "
        "option would ship a deck PowerPoint refuses to open")
    csrc = inspect.getsource(compile_picks.copy_picked_slide_into)
    assert "reassign_shape_ids" in csrc, (
        "compile_picks's graft no longer renumbers shape ids")
    print("    ok: finalize_deck and compile_picks both renumber")

    print("5. the package check catches a post-compile hand-edit")
    isrc = inspect.getsource(compile_picks.assert_package_integrity)
    assert "find_duplicate_shape_ids" in isrc, (
        "assert_package_integrity no longer checks shape ids; badges or label "
        "patches applied to the compiled deck could re-introduce the defect")
    print("    ok: assert_package_integrity checks shape ids")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
