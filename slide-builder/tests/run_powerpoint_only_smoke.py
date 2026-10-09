#!/usr/bin/env python3
"""Smoke test: Slide Lab runs without LibreOffice on Windows, through
PowerPoint, without disturbing the user's PowerPoint (owner's decision,
2026-10-08).

Part 1, the quit rule (no PowerPoint needed, a stand-in app):
  - PowerPoint was not running and nothing is open afterwards -> quit
  - PowerPoint was not running, but the user opened a deck while the render
    ran -> NOT quit (the race this fixes)
  - PowerPoint was already running -> never quit
  - only the presentation we opened is closed, read-only and windowless

Part 2, end to end with LibreOffice turned off (SLIDE_LAB_NO_LIBREOFFICE=1),
on fictional files, while whatever the user has open in PowerPoint stays
open (the open presentations are recorded before and after, by attaching,
never opening or closing anything of theirs):
  - slide-qc render (render_slides.py, default engine -> PowerPoint fallback)
  - slide-qc --engine ppt and export_slides.py (no longer refuse)
  - the converter's self-check (translate_html.py -> PowerPoint render)
  - the translator agent's self-check render (selfcheck_render.py)
  - the registration self-test and finalize are proved by running their own
    smokes with the same switch (see the CHANGELOG); this test covers the
    renderers they call
  - render time per slide, both engines, printed

Skipped (prints SMOKE PASSED with a note) when PowerPoint can't be driven.
Run:  py -3 slide-builder/tests/run_powerpoint_only_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import time
import types
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
QC = SKILL.parent / "slide-qc" / "scripts"
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(QC))

import ppt_safe  # noqa: E402

DESIGN = """<!doctype html><html><head><meta charset="utf-8"><style>
body { margin: 0; }
.slide-canvas { position: relative; width: 1280px; height: 720px; background: #fff;
  font-family: Arial, sans-serif; color: #1a1a1a; overflow: hidden; }
.title { position: absolute; left: 53px; top: 40px; font-size: 32px; font-weight: 700; }
.card { position: absolute; left: 53px; top: 160px; width: 500px; height: 200px;
  background: #eef2f6; border-radius: 6px; }
.card p { position: absolute; left: 24px; top: 24px; width: 450px; margin: 0;
  font-size: 18px; line-height: 26px; }
</style></head><body><div class="slide-canvas">
<div class="title" data-template-field="title">A fictional test slide</div>
<div class="card" data-shape-id="card"><p data-shape-id="card-body">Three steps move every
request from intake to done in one shared queue.</p></div>
</div></body></html>"""


# ---------------------------------------------------------------------------
# Part 1: the quit rule, with a stand-in PowerPoint
# ---------------------------------------------------------------------------
class _Pres:
    def __init__(self, app, name):
        self.app, self.FullName, self.closed = app, name, False
        self.Slides = types.SimpleNamespace(Count=1)

    def Close(self):
        self.closed = True
        self.app.pres.remove(self)


class _Coll:
    def __init__(self, app):
        self.app = app

    @property
    def Count(self):
        return len(self.app.pres)

    def __call__(self, i):
        return self.app.pres[i - 1]

    def Open(self, path, ro, untitled, window):
        assert (ro, untitled, window) == (True, False, False), "must open read-only, no window"
        p = _Pres(self.app, path)
        self.app.pres.append(p)
        self.app.opened.append(p)
        if self.app.user_opens_during:
            self.app.pres.append(_Pres(self.app, "user-deck.pptx"))   # the race
        return p


class _App:
    def __init__(self, user_opens_during=False, already=()):
        self.pres = [_Pres(self, n) for n in already]
        self.opened, self.quit, self.user_opens_during = [], False, user_opens_during
        self.Presentations = _Coll(self)
        self.Windows = types.SimpleNamespace(Count=0)
        self.Visible = False

    def Quit(self):
        self.quit = True


def _fake_session(app, running):
    fake_client = types.SimpleNamespace(Dispatch=lambda name: app)
    mods = {"win32com": types.SimpleNamespace(client=fake_client), "win32com.client": fake_client,
            "pythoncom": types.SimpleNamespace(CoInitialize=lambda: None, CoUninitialize=lambda: None)}
    saved = {k: sys.modules.get(k) for k in mods}
    real_running = ppt_safe.powerpoint_running
    sys.modules.update(mods)
    ppt_safe.powerpoint_running = lambda: running
    try:
        with ppt_safe.session() as (a, started):
            with ppt_safe.opened(a, Path("fictional.pptx")):
                pass
    finally:
        ppt_safe.powerpoint_running = real_running
        for k, v in saved.items():
            if v is None:
                sys.modules.pop(k, None)
            else:
                sys.modules[k] = v
    return app


def part1():
    a = _fake_session(_App(), running=False)
    assert a.quit and all(p.closed for p in a.opened), "should quit a PowerPoint it started"
    a = _fake_session(_App(user_opens_during=True), running=False)
    assert not a.quit, "quit although the user opened a deck during the render"
    assert [p.FullName for p in a.pres] == ["user-deck.pptx"], "the user's deck was touched"
    a = _fake_session(_App(already=("their-deck.pptx",)), running=True)
    assert not a.quit and [p.FullName for p in a.pres] == ["their-deck.pptx"]
    print("  ok: quit only a PowerPoint we started with nothing open in it; the user "
          "opening a deck mid-render keeps it running; only our file is closed")


# ---------------------------------------------------------------------------
# Part 2: end to end without LibreOffice
# ---------------------------------------------------------------------------
def _snapshot():
    """(running, open presentation names) of the user's PowerPoint; attach only."""
    if not ppt_safe.powerpoint_running():
        return False, []
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    try:
        app = win32com.client.GetActiveObject("PowerPoint.Application")
        return True, sorted(ppt_safe.open_names(app))
    except Exception:
        return True, ["<could not attach>"]


def _deck(path: Path, n: int) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        tb.text_frame.text = f"Fictional slide {i + 1}"
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    prs.save(str(path))
    return path


def _run(args, env, timeout=600):
    r = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True,
                       env=env, timeout=timeout)
    assert r.returncode in (0,), f"{args[0]} failed:\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}"
    return r.stdout


def part2():
    from render_slides import powerpoint_available
    if not powerpoint_available():
        print("  skipped part 2: PowerPoint can't be driven on this computer")
        return
    env = dict(os.environ, SLIDE_LAB_NO_LIBREOFFICE="1", SLIDE_LAB_NO_OPEN="1",
               PYTHONIOENCODING="utf-8")
    before = _snapshot()
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        deck = _deck(td / "fictional.pptx", 4)
        # slide-qc render, default engine: falls back to PowerPoint
        t0 = time.monotonic()
        out = _run([QC / "render_slides.py", deck, td / "png_default"], env)
        t_ppt = (time.monotonic() - t0) / 4
        assert "PowerPoint" in out and len(list((td / "png_default").glob("slide_*.png"))) == 4, out
        # --engine ppt and export_slides.py work with PowerPoint open
        _run([QC / "render_slides.py", deck, td / "png_ppt", "--engine", "ppt"], env)
        assert len(list((td / "png_ppt").glob("slide_*.png"))) == 4
        _run([QC / "export_slides.py", deck, "--out", td / "png_export"], env)
        assert len(list((td / "png_export").glob("slide_*.png"))) == 4
        print("  ok: slide-qc renders through PowerPoint (default fallback, --engine ppt, "
              "export_slides.py)")
        # the converter's self-check
        d = td / "slide_01"
        d.mkdir()
        (d / "option_A.html").write_text(DESIGN, encoding="utf-8")
        _run([SKILL / "scripts" / "translate_html.py", "--html", d / "option_A.html"], env)
        rep = json.loads((d / "option_A_translation_report.json").read_text(encoding="utf-8"))
        assert rep["self_check"]["ran"], rep["warnings"]
        assert not any(w["code"] == "SELF_CHECK_SKIPPED" for w in rep["warnings"])
        print("  ok: the converter's self-check ran through PowerPoint")
        # the translator agent's self-check render
        r = subprocess.run([sys.executable, str(d / "option_A_native.py")], cwd=str(d),
                           capture_output=True, text=True, timeout=120)
        assert r.returncode == 0, r.stderr
        _run([SKILL / "scripts" / "selfcheck_render.py", d / "option_A_native.pptx"], env)
        assert (d / "option_A_native.png").exists()
        print("  ok: selfcheck_render.py renders through PowerPoint")
        # LibreOffice timing for comparison (when installed)
        env_lo = dict(env)
        env_lo.pop("SLIDE_LAB_NO_LIBREOFFICE")
        t_lo = None
        try:
            from render_slides import _resolve_soffice
            _resolve_soffice()
            t0 = time.monotonic()
            _run([QC / "render_slides.py", deck, td / "png_lo"], env_lo)
            t_lo = (time.monotonic() - t0) / 4
        except RuntimeError:
            pass
    after = _snapshot()
    assert before[1] == after[1], "the user's open presentations changed"
    if before[0]:
        assert after[0], "PowerPoint was running before and is not now"
    print(f"  ok: the user's PowerPoint: running before={before[0]} after={after[0]}, "
          f"{len(before[1])} presentation(s) open before and the same {len(after[1])} after")
    print(f"  timing per slide (4-slide deck, one process each): PowerPoint {t_ppt:.2f}s"
          + (f", LibreOffice {t_lo:.2f}s" if t_lo is not None else ""))


def main() -> int:
    part1()
    part2()
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
