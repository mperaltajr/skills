#!/usr/bin/env python3
"""Smoke test for the Tier-2 QC gates.

Three things that were decorative or half-implemented before:

  1. `severity: "block"` did nothing. finalize counted blocking findings, printed
     the tally, and returned 0; nothing downstream read it, so a deck with
     blocking defects reported DONE. Now finalize records the count and compile
     refuses on it.
  2. Package integrity was unchecked. A structural edit outside the pipeline can
     leave an ORPHANED slide part (in the zip, referenced by no <p:sldId>).
     LibreOffice reports that as "source file could not be loaded", which is how
     the symptom got misfiled as a stale LibreOffice profile. The shipped
     13-slide deck still has exactly this defect.
  3. `chrome_zone_overlap` only checked the TOP zone despite its own comment
     promising the bottom one, and tested shape.top only (never top+height), so
     it could never see content reaching DOWN into the footer/page-number band.
     That is the class that recurred across slides 9/10/11.

Run:  py -3 slide-builder/tests/run_qc_gates_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import inspect
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pptx import Presentation  # noqa: E402

import _state  # noqa: E402
from compile_picks import assert_package_integrity  # noqa: E402


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="qc_gates_smoke_"))
    try:
        print("[1] a recorded QC block refuses the compile")
        out = tmp / "build"
        out.mkdir()
        _state.record_prep(out, "h1", out)
        tok = _state.record_review(out)
        ok, _ = _state.check_compile_allowed(out, tok)
        assert ok, "clean build should be compilable"
        _state.record_qc(out, 2, "2 option(s) with blocking QC findings")
        ok, why = _state.check_compile_allowed(out, tok)
        assert not ok, "a recorded QC block must refuse the compile"
        assert "blocking" in why, why
        _state.record_qc(out, 0, "clean")          # the fix re-runs finalize
        ok, _ = _state.check_compile_allowed(out, tok)
        assert ok, "clearing the blocks must re-allow the compile"
        print("    ok: block -> refuse, and a clean re-finalize re-allows it")

        print("[2] an orphaned slide part is caught")
        good = tmp / "good.pptx"
        prs = Presentation()
        for _ in range(3):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(str(good))
        assert assert_package_integrity(good) == [], "a clean deck must pass"

        orphan = tmp / "orphan.pptx"
        shutil.copy2(good, orphan)
        prs2 = Presentation(str(orphan))
        lst = prs2.slides._sldIdLst
        lst.remove(list(lst)[1])      # unlist the slide but keep its part + rel
        prs2.save(str(orphan))
        problems = assert_package_integrity(orphan)
        assert problems, "an orphaned slide part must be detected"
        assert "orphaned" in problems[0], problems
        print(f"    ok: {problems[0]}")

        print("[3] chrome-zone check covers the BOTTOM zone and exempts placeholders")
        import finalize_deck
        src = inspect.getsource(finalize_deck.run_option_qc)
        assert "BOTTOM_ZONE_PX" in src, "bottom invariant zone still not implemented"
        assert "bottom_px > BOTTOM_ZONE_PX" in src, "bottom zone is not actually tested"
        assert "bottom_px = top_px + h_px" in src, "still testing shape.top only, not the bbox"
        assert "is_placeholder" in src, (
            "inherited template placeholders are not exempt; the check fired on "
            "every option of every slide and therefore carried no signal")
        print("    ok: bottom zone tested via the bounding box; template chrome exempt")

        print("\nSMOKE PASSED.")
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
