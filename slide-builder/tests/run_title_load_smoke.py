#!/usr/bin/env python3
"""Smoke test: title / subtitle / takeaway band load (owner's rule, 2026-10-08).

A two-line title is allowed; a two-line title + a takeaway under it + a filled
bottom takeaway band is too much; each line must add something the others do
not say.

  - title_load.repeat_share: a line that reuses the title's content words
    scores >= 0.5, a line with new facts scores low, a too-short line is not
    judged
  - finished deck (slide-qc's check_pptx_hygiene): a band repeating the title
    is a content Major (content true, auto_fix false); a two-line title with a
    subtitle and a filled band is an Advisory; a one-line title with a
    subtitle that adds a number gets nothing; a footnote at the bottom is not
    taken for a band
  - brief side (brief_check): the repeat and two-line-title rows are
    recommendations, and recommendations never block a seal

Run:  py -3 slide-builder/tests/run_title_load_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import title_load as T  # noqa: E402

LONG_TITLE = ("Demand for AI compute outruns new supply, and every NVIDIA refresh widens "
              "the gap that retired GPUs can fill for locked-out buyers")
REPEAT_BAND = "Every NVIDIA refresh widens the gap retired GPUs can fill for locked-out buyers"
NEW_SUB = "About 8.4M GPUs retire over five years while 30% of demand goes unmet"


def _slide(prs, title, subtitle=None, band=None, footnote=None, band_name="takeaway-txt"):
    s = prs.slides.add_slide(prs.slide_layouts[5])          # Title Only
    s.shapes.title.text = title
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.5), Inches(0.3), Inches(12.3), Inches(1.0)
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size = Pt(28)
        r.font.name = "Arial"
    if subtitle:
        b = s.shapes.add_textbox(Inches(0.5), Inches(1.35), Inches(12.3), Inches(0.35))
        b.text_frame.text = subtitle
    if band:
        b = s.shapes.add_textbox(Inches(0.6), Inches(6.45), Inches(11.8), Inches(0.3))
        b.name = band_name
        b.text_frame.text = band
    if footnote:
        b = s.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.25))
        b.name = "TextBox 9"
        b.text_frame.text = footnote
    return s


def main() -> int:
    print("[1] repeat_share")
    assert T.repeat_share(LONG_TITLE, REPEAT_BAND) >= 0.5, T.repeat_share(LONG_TITLE, REPEAT_BAND)
    assert T.repeat_share(LONG_TITLE, NEW_SUB) < 0.5, T.repeat_share(LONG_TITLE, NEW_SUB)
    assert T.repeat_share("Claim number 1 is stated plainly here", "Takeaway 1.") == 0.0
    print("    ok")

    print("[2] finished deck findings")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _slide(prs, LONG_TITLE, subtitle=NEW_SUB, band=REPEAT_BAND)             # 1: Major + Advisory
    _slide(prs, "Kinfield is profitable from Year 1", subtitle=NEW_SUB)      # 2: nothing
    _slide(prs, LONG_TITLE, subtitle=NEW_SUB,                                # 3: footnote, not a band
           footnote="* Every NVIDIA refresh widens the gap retired GPUs can fill for buyers")
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "deck.pptx"
        prs.save(str(p))
        import check_pptx_hygiene as hy
        res = hy.run_all_checks(p)
    found = [(v["slide"], v["severity"], v["category"]) for v in res["violations"]
             if v["category"] in ("repeats-title", "title-load")]
    assert (1, "Major", "repeats-title") in found, found
    assert (1, "Advisory", "title-load") in found, found
    assert not [f for f in found if f[0] in (2, 3)], found
    major = next(v for v in res["violations"] if v["category"] == "repeats-title")
    assert major["content"] is True and major["auto_fix"] is False, major
    assert "bottom takeaway band" in major["issue"], major["issue"]
    print("    ok:", found)

    print("[3] brief side: recommendation rows never block a seal")
    import brief_check
    rows = [{"slide": 1, "field": "takeaway", "problem": "recommendation: x", "text": "t",
             "level": "recommendation"},
            {"slide": 1, "field": "title", "problem": "buzzword", "text": "t", "level": "issue"}]
    assert brief_check.blocking(rows) == rows[1:], brief_check.blocking(rows)
    assert brief_check.blocking(rows[:1]) == []
    brief = Path(tempfile.mkdtemp(prefix="slidelab_tl_")) / "brief.md"
    try:
        brief.write_text(
            "---\ndeck_type: client-pitch\naudience: \"Test\"\n"
            "governing_thought: \"Test.\"\nstoryline_gate_passed: true\n---\n\n# Title load test\n\n"
            "## Governing thought (the whole deck)\nTest.\n\n"
            "### Slide 1 — Every NVIDIA refresh widens the gap retired GPUs can fill\n\n"
            "**Slide type:** Content\n"
            "**Governing thought (the claim):** Every NVIDIA refresh widens the gap retired GPUs can fill.\n"
            "**The takeaway:** Each NVIDIA refresh widens the gap that retired GPUs fill, 30% of demand.\n"
            "**Evidence / content:**\n- Point one\n- Point two\n", encoding="utf-8")
        issues = brief_check.check(brief)
    finally:
        brief.unlink(missing_ok=True)
        brief.parent.rmdir()
    recs = [i for i in issues if i.get("level") == "recommendation" and "repeats the title" in i["problem"]]
    assert recs, issues
    assert all(i not in brief_check.blocking(issues) for i in recs)
    print("    ok:", recs[0]["problem"][:70])
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
