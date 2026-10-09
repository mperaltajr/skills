#!/usr/bin/env python3
"""Smoke test: pictures and charts survive finalize and assembly.

Both copy steps copied a slide's shapes without the files they point at.
finalize re-attached only top-level pictures; compile re-attached nothing. Any
picked page with a photo, a grouped photo, a logo or a native chart failed
the package-integrity check, and compile refused the whole deck with a
message blaming an edit "outside the pipeline" (2026-10-08).

Through the real pipeline (prep, finalize, final check, compile) on the
bundled test template, with one page carrying a top-level photo, a photo
inside a group, a small logo and a native chart:
  - compile exits 0 and the finished deck passes the integrity and
    PowerPoint-openability checks
  - the finished slide still has all three pictures (with their image files)
    and the chart (with its chart part and embedded workbook)
  - the rendered page shows the photo's and the chart's colors
  - copy_shape_with_parts on its own: a picture referenced twice is carried
    once, and a link to another slide is dropped, not left dangling

Run:  py -3 slide-builder/tests/run_picture_graft_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402

PHOTO_RGB = (200, 30, 40)        # the photo is a solid red block: easy to find on the render
CHART_RGB = "1F7A3A"             # the chart's bars are green

PICTURE_SCRIPT = '''# Slide 1 option A -- pattern: photo and chart (test harness)
# Variant: plain
# Brief fidelity check: test fixture
import sys
from pathlib import Path
sys.path.insert(0, r"{skill}")
from twins.helpers import new_slide, add_text
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
PX = 9525


def build():
    prs, slide = new_slide()
    photo = HERE / "_photo.png"
    Image.new("RGB", (300, 200), {photo_rgb}).save(photo)
    logo = HERE / "_logo.png"
    Image.new("RGB", (60, 60), (20, 40, 160)).save(logo)
    pic = slide.shapes.add_picture(str(photo), Emu(80 * PX), Emu(200 * PX), Emu(300 * PX), Emu(200 * PX))
    pic.name = "photo"
    grp = slide.shapes.add_group_shape()
    grp.name = "photo-group"
    g = grp.shapes.add_picture(str(photo), Emu(80 * PX), Emu(430 * PX), Emu(150 * PX), Emu(100 * PX))
    g.name = "grouped-photo"
    lg = slide.shapes.add_picture(str(logo), Emu(1150 * PX), Emu(140 * PX), Emu(40 * PX), Emu(40 * PX))
    lg.name = "logo"
    data = CategoryChartData()
    data.categories = ["One", "Two", "Three"]
    data.add_series("Values", (4, 7, 5))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(560 * PX), Emu(200 * PX),
                                Emu(560 * PX), Emu(380 * PX), data)
    gf.name = "chart"
    plot = gf.chart.plots[0]
    for pt in range(3):
        f = plot.series[0].points[pt].format.fill
        f.solid()
        f.fore_color.rgb = RGBColor.from_string("{chart_rgb}")
    add_text(slide, "body", "A page with pictures and a chart", 80, 150, 460, 30, font_size_pt=14)
    return prs


if __name__ == "__main__":
    build().save(str(HERE / "option_A.pptx"))
'''


def _review(out: Path) -> str:
    r = H.run("build_review.py", "--out", out)
    assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
    return _state.read_state(out)["review"]["token"]


def _line(token: str, picks) -> str:
    body = _state.canonical_picks(picks)
    return f"PICKS {body} CHECK {_state.approval_check(token, body)}"


def _near(png: Path, xy_px: tuple, rgb: tuple, tol: int = 40) -> bool:
    """Is the color at (x, y) of the 1280-wide page close to rgb?"""
    from PIL import Image
    im = Image.open(png).convert("RGB")
    k = im.width / 1280
    got = im.getpixel((int(xy_px[0] * k), int(xy_px[1] * k)))
    return all(abs(a - b) <= tol for a, b in zip(got, rgb))


def _unit_checks() -> None:
    import io
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from twins.composer import copy_shape_with_parts
    from pptx_openability import check_openability
    src = Presentation()
    s1 = src.slides.add_slide(src.slide_layouts[6])
    s2 = src.slides.add_slide(src.slide_layouts[6])
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), (10, 200, 10)).save(buf, "PNG")
    a = s1.shapes.add_picture(io.BytesIO(buf.getvalue()), Inches(1), Inches(1))
    b = s1.shapes.add_picture(io.BytesIO(buf.getvalue()), Inches(3), Inches(1))
    tb = s1.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "go to slide 2"
    run.hyperlink.address = None
    rid = s1.part.relate_to(s2.part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide")
    rpr = run._r.get_or_add_rPr()
    h = rpr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}hlinkClick",
                        {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": rid,
                         "action": "ppaction://hlinksldjump"})
    rpr.append(h)
    dst = Presentation()
    d = dst.slides.add_slide(dst.slide_layouts[6])
    memo: dict = {}
    for sh in (a, b, tb):
        d.shapes._spTree.append(copy_shape_with_parts(sh.element, s1.part, d.part, memo))
    images = [r for r in d.part.rels.values() if r.reltype.endswith("/image")]
    assert len(images) == 1, f"one picture used twice should be carried once, got {len(images)}"
    assert not [r for r in d.part.rels.values() if r.reltype.endswith("/slide")], \
        "a link to another slide was carried"
    out = io.BytesIO()
    dst.save(out)
    out.seek(0)
    issues = check_openability(Presentation(out))
    assert not issues, issues
    print("    ok: one copy of a twice-used picture; the cross-slide link dropped cleanly")


def main() -> int:
    from pptx import Presentation
    from compile_picks import assert_package_integrity

    print("[1] a page with a photo, a grouped photo, a logo and a chart, through the pipeline")
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1, body=PICTURE_SCRIPT.format(
            skill=str(SKILL), photo_rgb=PHOTO_RGB, chart_rgb=CHART_RGB))
        H.write_option(out, 2)
        picks = {"slide_01": "A", "slide_02": "A"}
        r = H.run("record_picks.py", "--out", out, "--approved", _line(_review(out), picks))
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-2500:] + r.stderr[-2500:]
        themed = sorted((out / "slide_01").glob("*themed*.pptx")) or \
            sorted((out / "slide_01").glob("option_A*.pptx"))
        r = H.run("build_review.py", "--out", out, "--final")
        tok = (_state.read_state(out).get("final_check") or {}).get("token")
        assert r.returncode == 0 and tok, r.stdout[-1500:] + r.stderr[-1500:]
        r = H.run("compile_picks.py", "--out", out, "--final-token", tok)
        assert r.returncode == 0, r.stdout[-2500:] + r.stderr[-2500:]
        assert "REFUSED" not in r.stdout, r.stdout[-1500:]
        deck = _state.compiled_deck(out)
        assert deck.exists(), "no final deck"
        print("    ok: compile exit 0")

        print("[2] the finished deck opens and keeps every picture and the chart")
        problems = assert_package_integrity(deck)
        assert problems == [], problems
        prs = Presentation(str(deck))
        s = prs.slides[0]
        rels = list(s.part.rels.values())
        n_img = sum(1 for x in rels if x.reltype.endswith("/image"))
        charts = [x.target_part for x in rels if x.reltype.endswith("/chart")]
        assert n_img >= 2, f"pictures lost their files ({n_img} image link(s))"
        assert len(charts) == 1, f"the chart lost its part ({len(charts)})"
        assert any(x.reltype.endswith("/package") for x in charts[0].rels.values()), \
            "the chart lost its embedded workbook"
        names = {sh.name for sh in s.shapes}
        assert {"photo", "photo-group", "logo", "chart"} <= names, names
        grp = next(sh for sh in s.shapes if sh.name == "photo-group")
        inner = list(grp.shapes)[0]
        assert inner.image.blob, "the grouped photo has no image"
        assert s.shapes[[sh.name for sh in s.shapes].index("chart")].chart.plots[0], "chart unreadable"
        print(f"    ok: {n_img} image file(s), 1 chart with its workbook; integrity clean")

        print("[3] the rendered page shows the photo and the chart")
        pngs = sorted(out.rglob("*.png"))
        cand = [p for p in pngs if "slide_01" in p.as_posix() or p.name.endswith("01.png")]
        page = next((p for p in cand if _near(p, (230, 300), PHOTO_RGB)), None)
        assert page is not None, f"no rendered page shows the photo: {[p.name for p in cand][:8]}"
        from PIL import Image
        im = Image.open(page).convert("RGB")
        k = im.width / 1280
        green = sum(1 for x in range(560, 1120, 8) for y in range(200, 580, 8)
                    if all(abs(a - b) <= 50 for a, b in zip(
                        im.getpixel((int(x * k), int(y * k))),
                        tuple(int(CHART_RGB[i:i + 2], 16) for i in (0, 2, 4)))))
        assert green > 20, f"the chart's bars are not on the render ({green} samples)"
        assert _near(page, (155, 480), PHOTO_RGB), "the grouped photo is not on the render"
        print(f"    ok: photo, grouped photo and chart drawn ({page.name})")
    finally:
        H.cleanup(tmp)

    print("[4] copy_shape_with_parts on its own")
    _unit_checks()
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
