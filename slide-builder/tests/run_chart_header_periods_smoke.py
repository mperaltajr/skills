#!/usr/bin/env python3
"""Smoke test: chart header line, legend place and no final periods (owner's
rules, 2026-10-09; slide-builder/scripts/chart_rules.py).

A fictional deck ("Northwind Bakeries", made-up numbers):

  1. strip_final_periods removes one final period from text boxes, table
     cells, group children and a period in a run of its own, and records each
     change (slide, shape, before / after); it keeps periods between
     sentences, ellipses, abbreviations (Co., e.g., U.S., a single capital
     "J."), list numbers ("3.") and text that never had one
  2. after the strip nothing is left for the QC check to flag
  3. a correct drawn chart (chart-title "Title, [Unit]" top left, legend one
     row top right) passes
  4. a legend under the chart is flagged; a stacked legend at the top right
     is flagged; a header with no unit over a value axis is flagged; a chart
     with no header line is flagged
  5. native charts: a legend at the bottom with no title is flagged (header,
     legend); a titled chart with its unit and the legend at the top passes
  6. slide-qc's hygiene check reports them as layout Majors with the keys
     final_period, chart_header, legend_position, and qc_fix_policy classes
     all three "auto"
  7. finalize carries the records: RESULT.md lists each removed period

Run:  py -3 slide-builder/tests/run_chart_header_periods_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE.parent / "scripts"))
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(ROOT / "slide-qc" / "scripts"))
import chart_rules  # noqa: E402
import check_pptx_hygiene  # noqa: E402
import qc_fix_policy  # noqa: E402


def _box(shapes, name, text, x=1.0, y=1.0, w=4.0, h=0.4, pt=12):
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tb.text_frame.text = text
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)
    return tb


def _rect(shapes, name, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    sh = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    return sh


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _drawn_chart(slide, *, header=None, legend="top-right", stacked=False):
    """A two-series bar chart drawn from shapes, plot area 1.0-7.0 in x,
    2.5-6.0 in y; value ticks at the left."""
    sh = slide.shapes
    if header is not None:
        _box(sh, "chart-title", header, x=1.0, y=1.8, w=4.0, h=0.3)
    for i, v in enumerate((0, 20, 40, 60)):
        _box(sh, f"chart-ylab-{v}", str(v), x=0.5, y=5.8 - i * 1.0, w=0.4, h=0.25, pt=9)
    for i in range(4):
        _rect(sh, f"chart-bar-aa-{i}", 1.2 + i * 1.4, 4.0, 0.5, 2.0)
        _rect(sh, f"chart-bar-bb-{i}", 1.75 + i * 1.4, 3.0, 0.5, 3.0)
    if legend == "top-right":
        ys = (1.8, 2.1) if stacked else (1.8, 1.8)
        xs = (5.6, 6.3) if not stacked else (6.3, 6.3)
        for (key, x, y) in (("aa", xs[0], ys[0]), ("bb", xs[1], ys[1])):
            _rect(sh, f"chart-legend-swatch-{key}", x, y + 0.08, 0.12, 0.12)
            _box(sh, f"chart-legend-{key}", key.upper(), x=x + 0.15, y=y, w=0.5, h=0.3, pt=9)
    elif legend == "bottom":
        for key, x in (("aa", 3.0), ("bb", 3.8)):
            _rect(sh, f"chart-legend-swatch-{key}", x, 6.5, 0.12, 0.12)
            _box(sh, f"chart-legend-{key}", key.upper(), x=x + 0.15, y=6.42, w=0.5, h=0.3, pt=9)


def _native_chart(slide, *, title=None, legend_pos=None):
    data = CategoryChartData()
    data.categories = ["FY1", "FY2", "FY3"]
    data.add_series("AA", (1.0, 2.0, 3.0))
    data.add_series("BB", (1.5, 2.5, 3.5))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2.5),
                                Inches(6), Inches(4), data)
    gf.name = "Chart 1"
    ch = gf.chart
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
    else:
        ch.has_title = False
    if legend_pos is not None:
        ch.has_legend = True
        ch.legend.position = legend_pos
        ch.legend.include_in_layout = False
    return gf


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

        # ---- slide 1: text items, periods to remove and to keep ----
        s1 = _blank(prs)
        sh = s1.shapes
        _box(sh, "title-ish", "Northwind grows bakery sales 12%.")
        _box(sh, "takeaway", 'Concise text on "so what". Explains takeaway.', y=1.5)
        _box(sh, "panel-text", 'Concise text on "so what". Explains takeaway', y=2.0)
        _box(sh, "partner", "Supplied by Smith & Co.", y=2.5)
        _box(sh, "ellipsis", "More to come...", y=3.0)
        _box(sh, "ellipsis-char", "More to come…", y=3.3)
        _box(sh, "example", "Pastries, breads, e.g.", y=3.6)
        _box(sh, "country", "Sold in the U.S.", y=3.9)
        _box(sh, "initial", "Founded by Anna J.", y=4.2)
        _box(sh, "step-number", "3.", x=6, y=1.0, w=0.5)
        bul = _box(sh, "bullets", "First point.", x=6, y=1.5)
        bul.text_frame.add_paragraph().text = "Second point"
        bul.text_frame.add_paragraph().text = "Third point."
        split = _box(sh, "split-run", "Ends in its own run", x=6, y=2.5)
        split.text_frame.paragraphs[0].add_run().text = "."
        _box(sh, "footnote-1", "1. Volumes are loaves sold.", x=1, y=6.6, pt=9)
        _box(sh, "source", "Source: Northwind sales ledger, 2026.", x=1, y=6.9, pt=9)
        tbl = sh.add_table(2, 2, Inches(6), Inches(3.2), Inches(4), Inches(1)).table
        tbl.cell(0, 0).text = "Region"
        tbl.cell(0, 1).text = "Loaves."
        tbl.cell(1, 0).text = "North"
        tbl.cell(1, 1).text = "1.2M."
        grp = sh.add_group_shape()
        _box(grp.shapes, "group-label", "Grouped label.", x=6, y=4.5)

        print("[1] strip_final_periods removes final periods and keeps the rest")
        changes = chart_rules.strip_final_periods(s1, slide_n=1)
        changed = {c["shape"] for c in changes}
        texts = {}
        for name, tf in chart_rules.iter_text_frames(s1.shapes):
            texts[name] = [p.text for p in tf.paragraphs]
        assert texts["title-ish"] == ["Northwind grows bakery sales 12%"], texts["title-ish"]
        assert texts["takeaway"] == ['Concise text on "so what". Explains takeaway'], texts["takeaway"]
        assert texts["panel-text"] == ['Concise text on "so what". Explains takeaway']
        assert texts["partner"] == ["Supplied by Smith & Co."]
        assert texts["ellipsis"] == ["More to come..."]
        assert texts["ellipsis-char"] == ["More to come…"]
        assert texts["example"] == ["Pastries, breads, e.g."]
        assert texts["country"] == ["Sold in the U.S."]
        assert texts["initial"] == ["Founded by Anna J."]
        assert texts["step-number"] == ["3."]
        assert texts["bullets"] == ["First point", "Second point", "Third point"], texts["bullets"]
        assert texts["split-run"] == ["Ends in its own run"], texts["split-run"]
        assert texts["footnote-1"] == ["1. Volumes are loaves sold"]
        assert texts["source"] == ["Source: Northwind sales ledger, 2026"]
        cells = {k: v for k, v in texts.items() if " cell " in k}
        assert sorted(sum(cells.values(), [])) == sorted(["Region", "Loaves", "North", "1.2M"]), cells
        grp_key = next(k for k in texts if k.endswith("/group-label"))
        assert texts[grp_key] == ["Grouped label"]
        # one record per changed paragraph: title, takeaway, 2 bullets,
        # split run, footnote, source, 2 cells, group label = 10
        assert len(changes) == 10, (len(changes), changes)
        for c in changes:
            assert c["slide"] == 1 and c["before"].endswith(".") and not c["after"].endswith("."), c
        assert {"title-ish", "takeaway", "bullets", "split-run", "footnote-1",
                "source"} <= changed, changed
        print(f"    ok ({len(changes)} removed)")

        print("[2] nothing left for the check after the strip")
        assert chart_rules.leftover_periods(s1) == [], chart_rules.leftover_periods(s1)
        print("    ok")

        # ---- slide 2: a correct drawn chart ----
        s2 = _blank(prs)
        _drawn_chart(s2, header="Loaves sold¹, [millions]")
        # ---- slide 3: legend under the chart, header without a unit ----
        s3 = _blank(prs)
        _drawn_chart(s3, header="Loaves sold by region", legend="bottom")
        # ---- slide 4: a leftover period, a stacked legend, no header ----
        s4 = _blank(prs)
        _drawn_chart(s4, header=None, stacked=True)
        _box(s4.shapes, "callout", "Sales doubled.", x=8, y=2.0)
        # ---- slide 5: native chart, legend at the bottom, no title ----
        s5 = _blank(prs)
        _native_chart(s5, title=None, legend_pos=XL_LEGEND_POSITION.BOTTOM)
        # ---- slide 6: native chart done right ----
        s6 = _blank(prs)
        _native_chart(s6, title="Loaves sold, [millions]", legend_pos=XL_LEGEND_POSITION.TOP)
        deck = td / "northwind.pptx"
        prs.save(str(deck))
        prs = Presentation(str(deck))

        print("[3] a correct drawn chart passes")
        p2 = chart_rules.chart_problems(prs.slides[1])
        assert p2 == [], p2
        print("    ok")

        print("[4] misplaced legend, missing unit, stacked legend, missing header")
        p3 = {p["kind"] for p in chart_rules.chart_problems(prs.slides[2])}
        assert p3 == {"legend", "unit"}, p3
        p4 = chart_rules.chart_problems(prs.slides[3])
        kinds4 = sorted(p["kind"] for p in p4)
        assert kinds4 == ["header", "legend", "unit"], p4
        assert any("stacked" in p["detail"] for p in p4), p4
        print("    ok")

        print("[5] native charts: bottom legend and no title flagged; top legend + title passes")
        p5 = {p["kind"] for p in chart_rules.chart_problems(prs.slides[4])}
        assert {"header", "legend"} <= p5, p5
        p6 = chart_rules.chart_problems(prs.slides[5])
        assert p6 == [], p6
        print("    ok")

        print("[6] slide-qc hygiene: layout Majors with policy keys, classed automatic")
        viol = check_pptx_hygiene.run_all_checks(deck)["violations"]
        mine = [v for v in viol if v["category"] in ("final_period", "chart_header",
                                                     "legend_position")]
        by = {(v["slide"], v["category"]) for v in mine}
        assert (4, "final_period") in by, by
        assert not any(s == 1 for s, _ in by), f"slide 1 was already cleaned: {by}"
        assert not any(s in (2, 6) for s, _ in by), f"correct charts flagged: {by}"
        assert {(3, "legend_position"), (3, "chart_header"), (4, "legend_position"),
                (4, "chart_header"), (5, "chart_header"), (5, "legend_position")} <= by, by
        assert all(v["severity"] == "Major" and v.get("auto_fix") for v in mine), mine
        for key in ("final_period", "chart_header", "legend_position"):
            assert qc_fix_policy.CATEGORIES[key][0] == "auto", key
            d, cat = qc_fix_policy.decide(f"slide 4 [{key}] Major: x")
            assert (d, cat) == ("auto", key), (d, cat)
        print(f"    ok ({len(mine)} findings)")

        print("[7] finalize lists each removed period in RESULT.md")
        import finalize_deck as fd
        st = fd.OptionStatus(slide_n=1, letter="A", py_path=td / "a.py",
                             pptx_path=td / "a.pptx", raw_archive_path=td / "raw.pptx",
                             themed_pptx_path=td / "a.pptx", themed_png_path=td / "a.png")
        st.period_fixes = changes
        res = fd.write_result(td, td / "template.pptx", [st])
        text = res.read_text(encoding="utf-8")
        assert "## Final periods removed (automatic)" in text, text
        assert "Northwind grows bakery sales 12%.' -> 'Northwind grows bakery sales 12%'" in text, text
        print("    ok")
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
