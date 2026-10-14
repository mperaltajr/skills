#!/usr/bin/env python3
"""Smoke test: several slides copy supplied pages, and every one is checked.

Leo narrative session report, item 3 (owner's decision 2026-10-13, the
lighter check). A build pinned 15 slides to pages of two earlier decks. Each
`source_ledger.py build` overwrote the one ledger file, and the gates read one
count from the last run: resolving the last page's rows would have shipped the
other 14 unchecked (a silent false pass). Now, on the real pipeline:

  1. prep prints a note naming the pinned slides (no stop)
  2. after the picks, finalize checks each pinned slide against its picked,
     finished design: one section per slide, from two different files; rows
     only for figures on BOTH the copied page and the design; a figure that is
     in the brief word for word is counted, not listed
  3. FINAL-CHECK.html shows the open rows in one table (keep / replace) and
     its record is bound to those rows
  4. a pinned slide with no section stops compile, naming the slide and the
     one command (--all-picked); after it, compile passes and the final-check
     approval answers every row the table showed ("kept on the final check");
     COMPILED.md lists each copied page
  5. a pin that cannot be read is a plain error, never a crash

Made-up pages only. Run:  py -3 slide-builder/tests/run_copied_pages_smoke.py
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402
import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))

import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402


def _page(path: Path, slides: list) -> None:
    prs = Presentation()
    for texts in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, t in enumerate(texts):
            s.shapes.add_textbox(Inches(0.5), Inches(0.4 + 0.8 * i), Inches(8),
                                 Inches(0.6)).text_frame.text = t
    prs.save(str(path))


def _option(out: Path, n: int, body_text: str) -> None:
    body = H.OPTION_SCRIPT.format(n=n, L="A", skill=str(H.SKILL)).replace(
        "Point one\\nPoint two\\nPoint three", body_text)
    H.write_option(out, n, body=body)


def main() -> int:
    root = Path(tempfile.mkdtemp(prefix="slidelab_copied_pages_"))
    # This checkout's own agent files stand in for the installed ones, so the
    # test never depends on (or touches) ~/.claude/agents.
    import os
    import shutil
    agents = root / "agents"
    agents.mkdir()
    for f in (H.SKILL / "agents").glob("slide-builder-*.md"):
        shutil.copy(f, agents / f.name)
    os.environ["SLIDE_LAB_AGENTS_DIR"] = str(agents)
    try:
        _page(root / "deck_a.pptx", [["Market reaches $300B", "Lead time is 6-12 months",
                                      "Upside of 3x"]])
        _page(root / "deck_b.pptx", [["Unrelated first page"],
                                     ["Costs fall by 40%", "Ten sites in total"]])
        s1 = H.SLIDE.format(n=1).replace(
            "**Slide type:** Content",
            "**Slide type:** Content\n**Pinned source page:** deck_a.pptx slide 1").replace(
            "- Point one for slide 1", "- The market is worth about $450B")
        s2 = H.SLIDE.format(n=2).replace(
            "**Slide type:** Content",
            "**Slide type:** Content\n**Pinned source page:** `deck_b.pptx` slide 2").replace(
            "- Point one for slide 2", "- Costs fall by 40%")
        s3 = H.SLIDE.format(n=3)
        brief = root / "brief.md"
        brief.write_text(H.BRIEF.format(slides=s1 + s2 + s3), encoding="utf-8")
        assert H.run("seal_brief.py", "--brief", brief, "--accepted", "test").returncode == 0
        out = root / "out"

        print("[1] prep names the pinned slides, and does not stop")
        r = H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE,
                  "--out", out, "--pattern", "direct", "--confirm-template")
        assert r.returncode == 0, r.stderr[-1500:]
        assert "2 slide(s) copy a supplied page: slide 1, 2." in r.stdout, r.stdout[-1500:]
        print("    ok")

        print("[2] after the picks, finalize checks each pinned slide on its pick")
        _option(out, 1, "The market reaches $300B\\nLead time is 6-12 months")
        _option(out, 2, "Costs fall by 40%")
        _option(out, 3, "Plain words here")
        assert H.finalize(out).returncode == 0
        assert not (out / "source_ledger.json").exists(), "no check before the picks"
        assert H.run("build_review.py", "--out", out).returncode == 0
        tok = _state.read_state(out)["review"]["token"]
        body = _state.canonical_picks({"slide_01": "A", "slide_02": "A", "slide_03": "A"})
        r = H.run("record_picks.py", "--out", out, "--approved",
                  f"PICKS {body} CHECK {_state.approval_check(tok, body)}")
        assert r.returncode == 0, r.stdout[-800:]
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-1500:]
        assert "Copied pages" in r.stdout and "slide 1" in r.stdout, r.stdout[-1500:]
        led = json.loads((out / "source_ledger.json").read_text(encoding="utf-8"))
        assert sorted(led["pages"]) == ["slide_01", "slide_02"], led["pages"].keys()
        a, b = led["pages"]["slide_01"], led["pages"]["slide_02"]
        assert Path(a["source_page"]).name == "deck_a.pptx" and a["source_slide"] == 1
        assert Path(b["source_page"]).name == "deck_b.pptx" and b["source_slide"] == 2
        assert sorted(r_["figure"] for r_ in a["rows"]) == ["$300B", "6-12"], a["rows"]
        assert a["not_on_design"] >= 1, "3x is not on the picked design"
        assert b["rows"] == [] and b["brief_matches"] == 1, b
        print("    ok: two sections from two files; $300B and 6-12 to answer; 40% "
              "matches the brief (counted)")

        print("[3] FINAL-CHECK.html shows one table, bound to its rows")
        r = H.run("build_review.py", "--out", out, "--final")
        assert r.returncode == 0, r.stderr[-1200:]
        page = (out / "FINAL-CHECK.html").read_text(encoding="utf-8")
        assert "Numbers copied from your pages" in page and "$300B" in page, "no table"
        assert "cpAnswers" in page and 'class="cp-row"' in page
        assert "1 more figure(s)" in page, "the brief match is counted on the page"
        fc = _state.read_state(out)["final_check"]
        assert sorted(fc["ledger"]["rows"]) == sorted(r_["id"] for r_ in a["rows"]), fc
        ftok = fc["token"]
        print("    ok")

        print("[4] a pinned slide with no check stops compile, naming it")
        led["pages"].pop("slide_02")
        (out / "source_ledger.json").write_text(json.dumps(led), encoding="utf-8")
        r = H.run("compile_picks.py", "--out", out, "--final-token", ftok)
        assert r.returncode == 5, r.stdout[-1500:]
        assert "slide 2 (no copied-page check yet)" in r.stdout, r.stdout[-1500:]
        assert "--all-picked" in r.stdout and "Never remove a pin" in r.stdout
        r = H.run("source_ledger.py", "build", "--out", out, "--all-picked")
        assert r.returncode == 0, r.stdout + r.stderr
        r = H.run("compile_picks.py", "--out", out, "--final-token", ftok)
        assert r.returncode == 0, r.stdout[-2000:]
        led = json.loads((out / "source_ledger.json").read_text(encoding="utf-8"))
        rows = led["pages"]["slide_01"]["rows"]
        assert rows and all(r_["resolution"] == "keep_source" and "FINAL-CHECK" in r_["note"]
                            for r_ in rows), rows
        comp = (out / "COMPILED.md").read_text(encoding="utf-8")
        assert "## Copied pages" in comp and "slide 1: copied from deck_a.pptx slide 1" in comp
        st = _state.read_state(out)["source_ledger"]["by_slide"]
        assert st["slide_01"]["kept"] == 2 and st["slide_02"]["brief_matches"] == 1, st
        print("    ok: refused naming slide 2 and --all-picked; then compiled, the two "
              "figures kept on the final check, COMPILED.md lists both pages")

        print("[5] a pin that cannot be read is a plain error")
        import source_ledger
        for bad, word in (("deck_a.pptx", "which slide"), ("missing.pptx slide 3", "not found")):
            try:
                source_ledger.parse_pin(bad, [root])
                raise AssertionError(f"accepted {bad!r}")
            except source_ledger.PinError as exc:
                assert word in str(exc), exc
        p, n = source_ledger.parse_pin(f'"{root}\\deck_b.pptx", slide 2', [])
        assert p.name == "deck_b.pptx" and n == 2
        meta = json.loads((out / "_meta.json").read_text(encoding="utf-8"))
        meta["slides"][0]["pinned_source_page"] = "nowhere.pptx slide 1"
        (out / "_meta.json").write_text(json.dumps(meta), encoding="utf-8")
        r = H.run("source_ledger.py", "build", "--out", out, "--all-picked")
        assert r.returncode == 2 and "not found" in r.stderr and "Traceback" not in r.stderr, \
            r.stderr
        print("    ok")
    finally:
        H.cleanup(root)
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
