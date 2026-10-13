#!/usr/bin/env python3
"""Smoke test: page numbers are slide-number FIELDS that keep the template's
character formatting (size, color, typeface), at every place that writes them.

Background. On many templates the slide-number field's own <a:rPr> is the ONLY
carrier of size/color/typeface: the layout AND master can both ship an empty
<a:lstStyle/>. Dropping it used to discard all three, so the number fell back to
the presentation default (~13pt black): the "page numbers render large" defect.
Owner's rule (2026-10-12): page numbers are never typed text. The pipeline used
to replace the field with a literal run, so a reordered deck kept stale numbers.
Now the field is kept (or created) and only its stored number is set.

The places that write a page number:
  1. twins/composer.py::write_slide_number_field (finalize's graft, via
     _populate_layout_placeholders)
  2. scripts/compile_picks.py::_restamp_page_number (re-stamps position at compile)
  3. twins/composer.py::stamp_page_number (the converter fix_page_numbers.py and
     helpers.add_footer's page-number box)
A fix to only one lets another re-introduce the defect. This test covers all.

Run:  py -3 slide-builder/tests/run_page_number_format_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.oxml import parse_xml  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from twins.composer import (  # noqa: E402
    stamp_page_number, write_literal_run_preserving_field, write_slide_number_field,
)

FIELD_ID = "{4D2A7B31-0000-0000-0000-000000000001}"
FIELD_XML = (
    '<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    f'id="{FIELD_ID}" type="slidenum">'
    '<a:rPr lang="en-US" sz="800">'
    '<a:solidFill><a:srgbClr val="8A93A0"/></a:solidFill>'
    '<a:latin typeface="Arial"/></a:rPr>'
    '<a:t>2</a:t></a:fld>'
)
RUN_XML = (
    '<a:r xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:rPr lang="en-US" sz="800">'
    '<a:solidFill><a:srgbClr val="8A93A0"/></a:solidFill>'
    '<a:latin typeface="Arial"/></a:rPr>'
    '<a:t>5</a:t></a:r>'
)
GUID_RE = re.compile(r"^\{[0-9A-F]{8}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{4}-[0-9A-F]{12}\}$")


def _para_with(xml):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.3))
    tb.name = "Slide Number 4"
    para = tb.text_frame.paragraphs[0]
    para._p.append(parse_xml(xml))
    return prs, slide, tb, para


def _assert_formatting(rpr, where: str):
    assert rpr is not None, f"{where}: NO rPr, formatting was lost"
    assert rpr.get("sz") == "800", f"{where}: size lost (sz={rpr.get('sz')})"
    latin = rpr.find(qn("a:latin"))
    assert latin is not None and latin.get("typeface") == "Arial", f"{where}: typeface lost"
    assert rpr.find(qn("a:solidFill")) is not None, f"{where}: color lost"


def _only_field(para, want: str, where: str):
    flds = para._p.findall(qn("a:fld"))
    assert len(flds) == 1 and flds[0].get("type") == "slidenum", f"{where}: not one slidenum field"
    assert flds[0].findtext(qn("a:t")) == want, f"{where}: stored number {flds[0].findtext(qn('a:t'))!r}"
    assert not para._p.findall(qn("a:r")), f"{where}: a typed run was left"
    return flds[0]


def main() -> int:
    print("[1] composer graft keeps the template's field: id, formatting, new number")
    _prs, _s, _tb, para = _para_with(FIELD_XML)
    write_slide_number_field(para, 7)
    fld = _only_field(para, "7", "composer")
    assert fld.get("id") == FIELD_ID, "the template's field id was replaced"
    _assert_formatting(fld.find(qn("a:rPr")), "composer")
    print("    ok: same field, 8pt / 8A93A0 / Arial kept, stored number 7")

    print("[2] a typed number becomes a field with the run's formatting and a fresh id")
    _prs2, _s2, _tb2, para2 = _para_with(RUN_XML)
    write_slide_number_field(para2, 5)
    fld2 = _only_field(para2, "5", "convert")
    assert GUID_RE.match(fld2.get("id") or ""), f"bad field id {fld2.get('id')!r}"
    _assert_formatting(fld2.find(qn("a:rPr")), "convert")
    # rPr is the first child of a:fld, a:t the last (schema order)
    assert fld2[0].tag == qn("a:rPr") and fld2[-1].tag == qn("a:t"), "field child order wrong"
    print("    ok: run -> field, formatting kept, id", fld2.get("id"))

    print("[3] compile-time re-stamp keeps the field and sets the position")
    import compile_picks
    _prs3, slide3, tb3, para3 = _para_with(FIELD_XML)
    compile_picks._restamp_page_number(slide3, 12)
    fld3 = _only_field(tb3.text_frame.paragraphs[0], "12", "compile")
    _assert_formatting(fld3.find(qn("a:rPr")), "compile")
    _prs4, slide4, tb4, _p4 = _para_with(RUN_XML)
    compile_picks._restamp_page_number(slide4, 3)
    _only_field(tb4.text_frame.paragraphs[0], "3", "compile (typed)")
    print("    ok: field kept at compile; a typed number is converted at compile")

    print("[4] stamp_page_number leaves an empty box alone; other text in the line stays")
    prs5 = Presentation()
    s5 = prs5.slides.add_slide(prs5.slide_layouts[6])
    tb5 = s5.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.3))
    tb5.name = "page-number"
    assert stamp_page_number(tb5, 4) == "", "an empty page-number box must be left alone"
    _prs6, _s6, tb6, para6 = _para_with(FIELD_XML)
    pre = parse_xml('<a:r xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                           '<a:t>Page </a:t></a:r>')
    para6._p.insert(0, pre)
    assert stamp_page_number(tb6, 9) == "kept"
    assert tb6.text_frame.text == "Page 9", tb6.text_frame.text
    print("    ok: empty box untouched; 'Page <field>' -> 'Page 9'")

    print("[5] the plain-text writer is still used for non-page-number text")
    _prs7, _s7, _tb7, para7 = _para_with(FIELD_XML)
    run = write_literal_run_preserving_field(para7, "Source: x")
    _assert_formatting(run._r.find(qn("a:rPr")), "literal writer")
    import inspect
    from twins import composer
    src = inspect.getsource(composer._populate_layout_placeholders)
    assert "page_number=True" in src, "page numbers no longer go through the field writer"
    print("    ok")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
