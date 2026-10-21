#!/usr/bin/env python3
"""Smoke test: finalize's chrome roles, section tag, palette codes and checks.

The validated fixes from a session report (2026-10-09):

  1. One "page takeaway" rule (_chrome_schema.is_page_takeaway): a SUBTITLE
     placeholder, a box named exactly "subtitle" anywhere, "subtitle-..." only
     in the top 35% of the page, never a "takeaway..." name. Source and
     footnote boxes count only in the bottom 25% or when their text starts like
     one. normalize_chrome leaves a bottom band named "takeaway-bar" and a body
     row named "source-1" where they are, and deletes only a second copy of
     the real takeaway.
  2. The section tag: drawn by chrome_rules.place_section_tag from the brief,
     one per slide (still one after a second finalize), none for "(none)",
     none on a cover or divider, 8 px from the top for a title box at y=34,
     in the template's accent color, and a plain stop when there is no room.
  3. Palette codes: no two Slide Lab palette roles share a color code, and on
     a Meridian-like theme the soft accent C2A3FF becomes the light blue
     B0C5DB while the card border C2A3FE stays a gray.
  4. Checks: the template's own master and layout colors are on-palette; the
     repeated-subtitle finding is a content Major (asked, never fixed); type
     scale keeps a loose "takeaway" box in its count.
  5. Text fit (_chrome_schema.fit_findings): "FY2 H1" at 14 pt bold in a 68 px
     cell is flagged, at 12 pt it is not; a word wider than its box and a
     table growing into a shape below it stop; a wrapped two-word label is
     flagged. check_fit.py reports and exits 1 on a stop.
  6. Through the real pipeline: finalize draws one section tag (and still one
     on a second run), none for "(none)", and refuses --skip-render once the
     user has picked.

No renderer needed (finalize runs with --skip-render).
Run:  py -3 slide-builder/tests/run_finalize_chrome_checks_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402  (temp folders under the scratch root, deleted at exit)
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))

import _chrome_schema as CS  # noqa: E402
import _paths as _p  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402
from twins import chrome_rules as CR  # noqa: E402
from twins import client_theme as CT  # noqa: E402
from twins import helpers as HP  # noqa: E402

EMU = 9525
TEMPLATE = HERE / "fixtures" / "layout_diverse_template.pptx"
ARIAL = Path(r"C:\Windows\Fonts\arial.ttf")
ARIAL_BD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def ok(msg: str) -> None:
    print(f"  ok: {msg}")


def box(slide, name, text, x, y, w, h, pt=12, bold=False):
    tb = slide.shapes.add_textbox(Emu(x * EMU), Emu(y * EMU), Emu(w * EMU), Emu(h * EMU))
    tb.name = name
    tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.name = "Arial"
    return tb


def geom(sh):
    return (int(sh.left), int(sh.top), int(sh.width), int(sh.height))


# ---------------------------------------------------------------------------
def phase_roles() -> None:
    print("\n[1] page takeaway / source / footnote roles")
    prs, slide = HP.new_slide()
    h = CS.slide_height_emu(slide)
    assert abs(h - 720 * EMU) < EMU, h
    sub = box(slide, "subtitle", "The real takeaway", 60, 600, 500, 30)
    sub_top = box(slide, "subtitle-sowhat", "A takeaway variant", 60, 100, 500, 30)
    sub_low = box(slide, "subtitle-note", "Low variant", 60, 500, 500, 30)
    band = box(slide, "takeaway-bar", "Bottom band text", 60, 600, 900, 40)
    band_top = box(slide, "takeaway", "Card named takeaway", 60, 110, 300, 40)
    src_row = box(slide, "source-1", "Hyperscalers 42%", 60, 300, 300, 24)
    src_low = box(slide, "source-line", "Made-up figures", 60, 640, 600, 20)
    src_mid_text = box(slide, "source-2", "Source: a report", 60, 320, 300, 20)
    fn_mid = box(slide, "footnote-1", "1. A note in the body", 60, 340, 300, 20)
    fn_body = box(slide, "footnote-row", "Not a note", 60, 360, 300, 20)
    assert CS.is_page_takeaway(sub, h)
    assert CS.is_page_takeaway(sub_top, h)
    assert not CS.is_page_takeaway(sub_low, h)
    assert not CS.is_page_takeaway(band, h) and not CS.is_page_takeaway(band_top, h)
    ok('"subtitle" anywhere and "subtitle-..." in the top 35% are the takeaway; '
       '"takeaway..." names never are')
    assert CS.note_role(src_row, h) is None
    assert CS.note_role(src_low, h) == "source"
    assert CS.note_role(src_mid_text, h) == "source"
    assert CS.note_role(fn_mid, h) == "footnote"
    assert CS.note_role(fn_body, h) is None
    ok("source/footnote names count only in the bottom 25% or when the text starts like one")

    # normalize_chrome: bands and body rows untouched, real takeaway moved
    prs, slide = HP.new_slide()
    real = box(slide, "subtitle", "Costs fall 12% in Year 2", 300, 400, 500, 30, pt=20, bold=True)
    copy = box(slide, "subtitle-copy", "Costs fall 12% in Year 2", 60, 120, 500, 30)
    band = box(slide, "takeaway-bar", "Costs fall 12% in Year 2", 60, 600, 900, 40)
    row = box(slide, "source-1", "Hyperscalers 42%", 60, 300, 300, 24)
    before_band, before_row = geom(band), geom(row)
    CR.normalize_chrome(slide, None)
    names = [s.name for s in slide.shapes]
    assert "takeaway-bar" in names and geom(band) == before_band, "bottom band was moved or deleted"
    assert "source-1" in names and geom(row) == before_row, "body row was pulled into the source stack"
    assert names.count("subtitle") == 1 and "subtitle-copy" not in names, names
    assert geom(real) != (300 * EMU, 400 * EMU, 500 * EMU, 30 * EMU), "the real takeaway was not placed"
    ok("normalize_chrome leaves takeaway-bar and source-1 alone, deletes only the takeaway's second copy")

    prs, slide = HP.new_slide()
    loose = box(slide, "takeaway-box", "Only a card", 60, 300, 300, 40)
    g0 = geom(loose)
    done = CR.normalize_chrome(slide, None)
    assert geom(loose) == g0 and done.get("takeaway", 0) in (0, None, False), done
    ok("nothing is moved when no shape passes the rule")


# ---------------------------------------------------------------------------
def _layout_slide(layout_name: str):
    prs = Presentation(str(TEMPLATE))
    spec = CS.load_chrome_yml(_p.chrome_yml(TEMPLATE))
    lc = spec.layouts[layout_name]
    layout = next(l for l in prs.slide_layouts if l.name == layout_name)
    slide = prs.slides.add_slide(layout)
    for sh in slide.placeholders:
        if sh.placeholder_format.type in (1, 3):
            sh.text_frame.text = "A title that states the claim"
    return prs, slide, lc


def _tags(slide):
    return [s for s in slide.shapes if s.name == CS.SECTION_TAG_NAME]


def phase_section_tag() -> None:
    print("\n[2] section tag")
    prs, slide, lc = _layout_slide("body_canonical_light")
    assert CR.place_section_tag(slide, lc, "The market", accent_hex="3A6EA5")
    tags = _tags(slide)
    assert len(tags) == 1, len(tags)
    tag = tags[0]
    title = next(s for s in slide.placeholders if s.placeholder_format.type in (1, 3))
    gap_px = (int(title.top) - int(tag.top) - int(tag.height)) / EMU
    assert abs(gap_px - CS.SECTION_TAG_GAP_PX) < 0.6, gap_px
    rpr = tag.text_frame.paragraphs[0].runs[0]._r.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
    assert rpr.get("sz") == "1050" and rpr.get("b") == "1" and rpr.get("cap") == "small", rpr.attrib
    srgb = rpr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    assert srgb is not None and srgb.get("val") == "3A6EA5"
    ok("one tag, just above the title box, 10.5 pt bold small caps, accent 3A6EA5")
    assert CR.place_section_tag(slide, lc, "The market", accent_hex="3A6EA5")
    assert len(_tags(slide)) == 1
    ok("still one tag after a second run")
    assert not CR.place_section_tag(slide, lc, "(none)", accent_hex="3A6EA5")
    assert not _tags(slide)
    assert not CR.place_section_tag(slide, lc, "", accent_hex="3A6EA5")
    ok('no tag for "(none)" or an empty tag (an earlier one is removed)')

    prs, slide, lc = _layout_slide("body_canonical_dark")
    CR.place_section_tag(slide, lc, "Financials", accent_hex="3A6EA5")
    srgb = _tags(slide)[0]._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    assert srgb.get("val") == "FFFFFF"
    ok("white on a dark layout")

    for name in ("cover_light", "section_divider", "cover_dark"):
        prs, slide, lc = _layout_slide(name)
        assert not CR.place_section_tag(slide, lc, "The market", accent_hex="3A6EA5")
        assert not _tags(slide)
    prs, slide, lc = _layout_slide("body_canonical_light")
    assert not CR.place_section_tag(slide, lc, "The market", accent_hex="3A6EA5", page_type="cover")
    ok("no tag on a cover, a divider or a page typed cover")

    # A client template's title box sits at x=53, y=34: the tag lands ~8 px from the top
    prs, slide = HP.new_slide()
    kin = SimpleNamespace(name="Meridian Content", layout_class="body-canonical",
                          text_role="dark_on_light", has_page_number=True,
                          title_placeholder_idx=None, title_box_x_px=53,
                          title_box_width_px=1100, title_box_y_px=34, title_box_height_px=60)
    CR.place_section_tag(slide, kin, "The market", accent_hex="3A6EA5")
    t = _tags(slide)[0]
    assert abs(int(t.top) / EMU - 8) < 0.6, int(t.top) / EMU
    ok(f"Meridian-like title box at y=34: tag top at {int(t.top) / EMU:.0f} px")
    tight = SimpleNamespace(**{**kin.__dict__, "title_box_y_px": 10})
    prs, slide = HP.new_slide()
    try:
        CR.place_section_tag(slide, tight, "The market", accent_hex="3A6EA5")
    except CR.SectionTagNoRoomError as e:
        assert "no room" in str(e) and "Ask the user" in str(e)
    else:
        raise AssertionError("no room above the title did not stop")
    ok("no room above the title stops with a plain message")


# ---------------------------------------------------------------------------
def phase_palette() -> None:
    print("\n[3] palette codes")
    codes = {k: v for k, v in vars(CT).items()
             if k.startswith("SLIDE_LAB_") and isinstance(v, str) and len(v) == 6
             and all(c in "0123456789ABCDEFabcdef" for c in v)}
    by_hex: dict = {}
    for k, v in codes.items():
        by_hex.setdefault(v.upper(), []).append(k)
    shared = {h: ks for h, ks in by_hex.items() if len(ks) > 1}
    assert not shared, f"palette roles share a code: {shared}"
    helper_cols = {n: str(getattr(HP, n)) for n in
                   ("BRAND_PRIMARY", "BRAND_PRIMARY_MID", "BRAND_ACCENT", "BRAND_ACCENT_SOFT",
                    "CARD_BG", "CARD_BORDER")}
    assert len(set(helper_cols.values())) == len(helper_cols), helper_cols
    assert str(HP.CARD_BORDER).upper() == "C2A3FE" and str(HP.BRAND_ACCENT_SOFT).upper() == "C2A3FF"
    ok("no two palette roles share a color code (CARD_BORDER is C2A3FE)")

    th = CT.ClientTheme(dk1="000000", lt1="FFFFFF", dk2="0B1F33", lt2="EEF1F5",
                        accent1="3A6EA5", accent2="0B1F33", accent3="8A93A0",
                        accent4="DADDE0", accent5="B0C5DB", accent6="5F5E5A",
                        major_font="Arial", minor_font="Arial",
                        brand_primary="0B1F33", brand_accent="3A6EA5")
    cm = th.color_map()
    assert cm["C2A3FF"] == "B0C5DB", cm["C2A3FF"]
    assert cm["C2A3FE"] == CT._mix_hex(th.dk1, th.lt1, 0.85), cm["C2A3FE"]
    assert cm["C2A3FF"] != cm["C2A3FE"]
    assert cm["460073"] == "0B1F33" and cm["A100FF"] == "3A6EA5"
    ok(f"Meridian-like theme: soft accent C2A3FF -> {cm['C2A3FF']}, card border C2A3FE -> {cm['C2A3FE']}")


# ---------------------------------------------------------------------------
def phase_checks() -> None:
    print("\n[4] palette, repeated subtitle, type scale")
    import finalize_deck as FD
    own = FD._template_own_colors(TEMPLATE)
    assert own, "no colors read from the template's masters and layouts"
    th = CT.load_client_theme(str(TEMPLATE))
    pal = FD._expected_palette_for_theme(th, TEMPLATE)
    assert own <= pal
    ok(f"{len(own)} template master/layout colors count as on-palette")

    st = SimpleNamespace(page_title="Costs fall 12% in Year 2 as volumes grow",
                         page_subtitle="Costs fall 12% in Year 2")
    f = FD._repeat_finding(st)
    assert f and f["major"] and f["content"] and f["auto_fix"] is False and f["slide_level"]
    assert f["category"] == "rewrite" and "Your wording choice" in f["plain"][0]
    st2 = SimpleNamespace(page_title="Costs fall 12% in Year 2",
                          page_subtitle="Savings fund the second site by FY4")
    assert FD._repeat_finding(st2) is None
    ok("a takeaway that repeats the title is a content Major (asked, never fixed)")

    import type_scale as TS
    prs, slide = HP.new_slide()
    box(slide, "takeaway", "Loose card text", 60, 300, 400, 30, pt=7)
    res = TS.check_slide(slide, prs.slide_height)
    assert [t for _pt, t in res["small_body"] if t == "Loose card text"], res
    ok("type scale checks a loose box named takeaway")

    import translate_html as TH
    plan, _ = TH.plan_from_extract({"fields": {"title": "A claim", "section_tag": "THE MARKET"},
                                    "canvas": {"w": 1280, "h": 720}, "boxes": [], "texts": [],
                                    "svgs": [], "fallbacks": []},
                                   subtitle_as_shape=False)
    codes = [w["code"] for w in plan.warnings]
    assert "SECTION_TAG_FIELD" in codes, plan.warnings
    ok("the converter warns when a sketch marks the section tag as a template field")


# ---------------------------------------------------------------------------
def _fonts():
    assert ARIAL.exists() and ARIAL_BD.exists(), "Arial not installed"
    return CS.fit_fonts(heading="Arial", body="Arial", heading_ttf=str(ARIAL),
                        heading_bold_ttf=str(ARIAL_BD), body_ttf=str(ARIAL))


def _fy_table(slide, pt):
    gf = slide.shapes.add_table(2, 2, Emu(100 * EMU), Emu(200 * EMU), Emu(278 * EMU), Emu(60 * EMU))
    gf.name = "roadmap-table"
    t = gf.table
    t.columns[0].width = Emu(200 * EMU)
    t.columns[1].width = Emu(78 * EMU)
    for r in range(2):
        t.rows[r].height = Emu(30 * EMU)
        for c in range(2):
            cell = t.cell(r, c)
            cell.margin_left, cell.margin_right = Emu(10 * EMU), Emu(0)
            cell.margin_top = cell.margin_bottom = Emu(0)
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = "FY2 H1" if c == 1 else "GPU Leasing"
            run.font.size = Pt(pt)
            run.font.bold = c == 1
            run.font.name = "Arial"
    return gf


def phase_fit() -> None:
    print("\n[5] text fit")
    fonts = _fonts()
    prs, slide = HP.new_slide()
    _fy_table(slide, 14)
    res = CS.fit_findings(slide, fonts)
    fy = [f for f in res["findings"] if f["text"] == "FY2 H1"]
    assert fy and all(f["kind"] == "near_limit" and not f["stop"] for f in fy), res["findings"]
    ok(f'"FY2 H1" 14 pt bold in a 68 px room is flagged ({fy[0]["detail"]})')
    prs, slide = HP.new_slide()
    _fy_table(slide, 12)
    res = CS.fit_findings(slide, fonts)
    assert not [f for f in res["findings"] if f["text"] == "FY2 H1"], res["findings"]
    ok('"FY2 H1" at 12 pt in the same room is not')

    prs, slide = HP.new_slide()
    box(slide, "kpi", "Securitization", 100, 100, 60, 40, pt=16)
    box(slide, "label", "Closing cash", 300, 100, 100, 40, pt=14)
    gf = slide.shapes.add_table(3, 1, Emu(600 * EMU), Emu(200 * EMU), Emu(200 * EMU), Emu(60 * EMU))
    gf.name = "tall-table"
    for r in range(3):
        gf.table.rows[r].height = Emu(20 * EMU)
        cell = gf.table.cell(r, 0)
        cell.text_frame.text = "A long cell sentence that wraps onto several lines in this narrow column"
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(12)
    under = slide.shapes.add_shape(1, Emu(600 * EMU), Emu(290 * EMU), Emu(200 * EMU), Emu(40 * EMU))
    under.name = "card-below"
    res = CS.fit_findings(slide, fonts, body_bottom_px=660)
    kinds = {(f["kind"], f["shape"].split(" row")[0]) for f in res["findings"]}
    assert ("broken_word", "kpi") in kinds, kinds
    assert ("short_label_wraps", "label") in kinds, kinds
    assert ("table_overflow", "tall-table") in kinds, kinds
    assert all(f["stop"] for f in res["findings"] if f["kind"] in ("broken_word", "table_overflow"))
    ok("broken word and table running into a shape stop; a wrapped two-word label is flagged")

    tmp = Path(tempfile.mkdtemp(prefix="slidelab_fit_"))
    (tmp / "slide_01").mkdir()
    prs.save(str(tmp / "slide_01" / "option_A.pptx"))
    r = subprocess.run([sys.executable, str(SKILL / "scripts" / "check_fit.py"),
                        str(tmp / "slide_01" / "option_A.pptx"), "--template", str(TEMPLATE)],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    assert r.returncode == 1 and "STOP" in r.stdout, r.stdout + r.stderr
    ok("check_fit.py lists the stops and exits 1")


# ---------------------------------------------------------------------------
def phase_pipeline() -> None:
    print("\n[6] through finalize (section tag, --skip-render refusal)")
    import _e2e_harness as H
    import _state
    tmp, out = H.new_build(2)
    try:
        meta_p = _p.meta_json(out)
        meta = json.loads(meta_p.read_text(encoding="utf-8"))
        meta["slides"][0]["section_tag"] = "The market"
        meta["slides"][1]["section_tag"] = "(none)"
        meta_p.write_text(json.dumps(meta, indent=2), encoding="utf-8")
        for n in (1, 2):
            H.write_option(out, n)
        for run in (1, 2):
            r = H.finalize(out, "--skip-render")
            assert r.returncode == 0, f"finalize run {run} failed ({r.returncode}):\n{r.stdout[-3000:]}\n{r.stderr[-2000:]}"
            tags = {}
            for n in (1, 2):
                p = out / f"slide_{n:02d}" / "option_A.pptx"          # the finished slide
                assert p.exists(), sorted(str(x) for x in out.rglob("*.pptx"))
                s = Presentation(str(p)).slides[0]
                tags[n] = [sh for sh in s.shapes if sh.name == CS.SECTION_TAG_NAME]
            assert len(tags[1]) == 1 and tags[1][0].text_frame.text == "The market", tags
            assert len(tags[2]) == 0, tags
        ok("finalize draws one tag on slide 1 (still one after a second run), none for (none)")
        _state.record_picks(out, {_p.slide_key(1): "A", _p.slide_key(2): "A"})
        r = H.finalize(out, "--skip-render")
        assert r.returncode == 18 and "--skip-render refused" in (r.stdout + r.stderr), r.stdout[-1500:] + r.stderr[-1500:]
        ok("--skip-render is refused once the user has picked")
    finally:
        H.cleanup(tmp)


def main() -> int:
    os.environ.setdefault("SLIDE_LAB_NO_OPEN", "1")
    phase_roles()
    phase_section_tag()
    phase_palette()
    phase_checks()
    phase_fit()
    phase_pipeline()
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
