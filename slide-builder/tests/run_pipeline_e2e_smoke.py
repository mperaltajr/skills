#!/usr/bin/env python3
"""End-to-end smoke: the real scripts, run in order, on the bundled template.

Every other smoke calls functions with hand-built state. This runs prep, a
worker-shaped option script per slide, finalize, the review page, record_picks,
the final check and compile as separate processes, the way a build does, and
checks that what each stage records is what actually happened.

The order it drives is the owner's decision (2026-09-30):
  REVIEW.html -> record_picks.py -> (translate picks) -> finalize ->
  build_review.py --final (FINAL-CHECK.html) -> compile --final-token

Cases (each one shipped, or could have shipped, a wrong deck before):
  1   clean build: every stage records what happened, compile succeeds,
      check_done passes only after a vision pass
  1b  partial picks and tampered picks are refused by record_picks
  1c  a deck failing its integrity check never replaces the good one
  1d  a finished file changed after the final check refuses compile
  1e  the review page's JavaScript check code equals Python's
  2   compile before finalize / before the final check is refused
  3   a crashing option script is recorded as blocked, by name
  4   finalize --slide 1 must not wipe slide 2's block
  4b  a rebuild moves the old option files aside
  5   a picked but untranslated sketch points at the translator; an unpicked
      one is simply not converted

Run:  py -3 slide-builder/tests/run_pipeline_e2e_smoke.py   (about two minutes)
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))

import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402

CRASHING_SCRIPT = "# Slide 2 option A — crashes on purpose\nraise RuntimeError('boom')\n"


def _review(out: Path) -> str:
    r = H.run("build_review.py", "--out", out)
    assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
    return _state.read_state(out)["review"]["token"]


def _line(token: str, picks) -> str:
    """What the page's Build button puts on the picks line."""
    body = _state.canonical_picks(picks)
    return f"PICKS {body} CHECK {_state.approval_check(token, body)}"


def _record(out: Path, picks) -> "subprocess.CompletedProcess":
    token = _review(out)
    return H.run("record_picks.py", "--out", out, "--approved", _line(token, picks))


def _final(out: Path):
    r = H.run("build_review.py", "--out", out, "--final")
    tok = (_state.read_state(out).get("final_check") or {}).get("token")
    return r, tok


def _compile(out: Path, token, *extra):
    args = ["--out", out, *extra]
    if token:
        args += ["--final-token", token]
    return H.run("compile_picks.py", *args)


def main() -> int:
    print("[1] clean build, in the owner's order")
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1)
        H.write_option(out, 2)
        both = {"slide_01": "A", "slide_02": "A"}

        print("[2] compile before finalize / before the final check is refused")
        r = _record(out, both)
        assert r.returncode == 0, r.stdout
        r = _compile(out, None)
        assert r.returncode == 5 and "final look" in r.stdout, r.stdout[-600:]
        r, _ = _final(out)
        assert r.returncode == 5, "a final check before finalize must refuse"
        print("    ok: no final check yet, compile refuses; final check refuses unfinalized")

        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        assert _state.read_state(out)["review"].get("picks") == both, (
            "a re-finalize must keep the approved picks")
        r, tok = _final(out)
        assert r.returncode == 0 and tok, r.stdout + r.stderr
        assert (out / "FINAL-CHECK.html").exists()
        r = _compile(out, "0000000000000000")
        assert r.returncode == 5 and "does not match" in r.stdout, r.stdout[-600:]
        r = _compile(out, tok)
        assert r.returncode == 0, r.stdout[-2000:] + r.stderr[-2000:]
        comp = _state.read_state(out)["stages"]["compile"]
        assert comp["kind"] == "picks" and comp["slides"] == 2, comp
        print("    ok: finalize -> final check -> compile with its token")

        print("[1a] vision pass + check_done over the real compiled deck")
        deck = _state.compiled_deck(out)
        r = H.run("check_done.py", "--out", out)
        assert r.returncode == 1 and "no vision pass" in r.stdout, r.stdout
        r = H.run("record_vision_qc.py", "--out", out, "--deck", deck,
                  "--slides-reviewed", 2, "--criticals", 0, "--majors", 0,
                  "--advisories", 1)
        assert r.returncode == 0, r.stdout + r.stderr
        r = H.run("check_done.py", "--out", out)
        assert r.returncode == 0 and "DELIVERABLE (final deck)" in r.stdout, r.stdout
        print("    ok: refused before the vision pass, deliverable after")

        print("[1a2] overrides used in the build are shown at delivery")
        _state.record_override(out, "assume_gated", "test")
        r = H.run("check_done.py", "--out", out)
        assert r.returncode == 0 and "passed over" in r.stdout and "assume_gated" in r.stdout, r.stdout
        print("    ok")

        print("[1a3] a single-slide rebuild after compile makes the deck stale")
        r = H.run("build_deck.py", "--brief", tmp / "brief.md", "--template", H.TEMPLATE,
                  "--out", out, "--pattern", "direct", "--slide", "2")
        assert r.returncode == 0, r.stderr[-800:]
        r = H.run("check_done.py", "--out", out)
        assert r.returncode == 1 and "rebuilt or inserted after" in r.stdout, (
            "same brief, so the content hash alone missed this; the old deck must "
            "not stay DELIVERABLE: " + r.stdout)
        print("    ok")
        # put slide 2 back so the rest of this build can continue
        H.write_option(out, 2)
        assert H.finalize(out, "--slide", "2").returncode == 0

        print("[1b] partial and tampered picks are refused")
        r = _record(out, {"slide_01": "A"})
        assert r.returncode == 5 and "no pick" in r.stdout, r.stdout
        token = _review(out)
        good = _line(token, both)
        tampered = good.replace("slide_02=A", "slide_02=B")
        r = H.run("record_picks.py", "--out", out, "--approved", tampered)
        assert r.returncode == 5 and "check code" in r.stdout, r.stdout
        print("    ok: missing slide and edited pick both refused")

        print("[1c] a deck failing its integrity check never replaces the good one")
        assert _record(out, both).returncode == 0
        r, tok = _final(out)
        assert r.returncode == 0
        good_digest = _state.file_digest(deck)
        import contextlib, io, os
        import compile_picks
        _orig = compile_picks._report_integrity
        compile_picks._report_integrity = lambda p: 6
        argv = sys.argv
        sys.argv = ["compile_picks.py", "--out", str(out), "--final-token", tok]
        os.environ["SLIDE_LAB_OPTIONS_PER_SLIDE"] = "1"
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                rc = compile_picks.main()
        finally:
            compile_picks._report_integrity = _orig
            sys.argv = argv
        assert rc == 6, rc
        assert _state.file_digest(deck) == good_digest, "good deck was replaced"
        assert (out / "final_deck.REJECTED.pptx").exists()
        print("    ok: exit 6, previous deck untouched, rejected file set aside")

        print("[1d] a finished file changed after the final check refuses compile")
        themed = out / "slide_01" / "option_A.pptx"
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(str(themed))
        prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        prs.save(str(themed))
        r = _compile(out, tok)
        assert r.returncode == 5 and "changed after the final" in r.stdout, r.stdout[-800:]
        print("    ok: the user approved different bytes; refused")

        print("[1e] the page's JavaScript check code equals Python's")
        try:
            from playwright.sync_api import sync_playwright
        except ImportError:
            print("    skipped: playwright not installed")
        else:
            token = _review(out)
            with sync_playwright() as pw:
                b = pw.chromium.launch()
                pg = b.new_page()
                pg.goto((out / "REVIEW.html").resolve().as_uri())
                js = pg.evaluate("() => approvalCheck(canonicalPicks("
                                 "{slide_02: 'A', slide_01: 'A'}))")
                js_all = pg.evaluate("() => approvalCheck('ALL')")
                b.close()
            assert js == _state.approval_check(token, "slide_01=A;slide_02=A"), js
            assert js_all == _state.approval_check(token, "ALL"), js_all
            print("    ok: identical for a pick list and for ALL")

            print("[1f] the page's Build button: refuses a pending slide, then emits a")
            print("     line record_picks accepts; a rebuilt slide loses only its own pick")
            with sync_playwright() as pw:
                b = pw.chromium.launch()
                pg = b.new_page()
                pg.goto((out / "REVIEW.html").resolve().as_uri())
                pg.evaluate("() => { localStorage.clear(); window.__copied = null;"
                            " navigator.clipboard.writeText = t => {"
                            " window.__copied = t; return Promise.resolve(); }; }")
                pg.click('#card-slide_01 .option[data-letter="A"] img, '
                         '#card-slide_01 .option[data-letter="A"]')
                pg.evaluate("() => buildDeck()")
                assert pg.evaluate("() => window.__copied") is None, (
                    "Build copied a command with slide 2 still pending")
                pg.click('#card-slide_02 .option[data-letter="A"] img, '
                         '#card-slide_02 .option[data-letter="A"]')
                pg.evaluate("() => buildDeck()")
                pg.wait_for_function("() => window.__copied")
                cmd = pg.evaluate("() => window.__copied")
                # A rebuilt slide 2 gets a new stamp: its pick must fall away,
                # slide 1's must stay.
                pg.evaluate("() => { window.__STAMPS__['slide_02'] = 'rebuilt';"
                            " STAMPS['slide_02'] = 'rebuilt'; }")
                kept = pg.evaluate("() => [pickForSlide('slide_01'), pickForSlide('slide_02')]")
                b.close()
            import re as _re
            # The page copies a plain message now; its Picks line is what
            # record_picks.py receives (check-coded, verified there).
            m = _re.search(r"^(Picks: .+\(check [0-9a-f]{8}\))$", cmd, _re.M)
            assert m and "py -3" not in cmd, cmd
            r = H.run("record_picks.py", "--out", out, "--approved", m.group(1))
            assert r.returncode == 0, r.stdout
            assert kept == ["A", None], kept
            print("    ok")
    finally:
        H.cleanup(tmp)

    print("[3] a crashing option script is recorded as blocked, by name")
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1)
        H.write_option(out, 2, body=CRASHING_SCRIPT)
        r = H.finalize(out)
        assert r.returncode == 15, f"expected exit 15, got {r.returncode}\n{r.stdout[-1500:]}"
        by = _state.read_state(out)["qc"]["by_option"]
        assert by["slide_02/A"]["blocks"] > 0, by
        assert "did not build" in " ".join(by["slide_02/A"]["reasons"]), by
        assert _record(out, {"slide_01": "A", "slide_02": "A"}).returncode == 0
        r, _ = _final(out)
        assert r.returncode == 5 and "slide_02/A" in r.stderr, r.stderr[-800:]
        print("    ok: exit 15, slide_02/A blocked, the final check refuses to show it")

        print("[4] finalize --slide 1 must not wipe slide 2's block")
        r = H.finalize(out, "--slide", "1")
        assert r.returncode == 0, r.stdout[-1500:]
        by = _state.read_state(out)["qc"]["by_option"]
        assert by.get("slide_02/A", {}).get("blocks", 0) > 0, by
        r, _ = _final(out)
        assert r.returncode == 5, "slide 2 is still blocked"
        print("    ok: slide 2 still blocked")
    finally:
        H.cleanup(tmp)

    print("[4b] a rebuild moves the old option files aside")
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1)
        H.write_option(out, 2)
        assert H.finalize(out).returncode == 0
        r = H.run("build_deck.py", "--brief", tmp / "brief.md", "--template", H.TEMPLATE,
                  "--out", out, "--pattern", "direct", "--confirm-template", "--slide", "2")
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        left = sorted(p.name for p in (out / "slide_02").glob("option_*"))
        assert left == [], f"old option files left in place: {left}"
        assert list((out / "slide_02" / "_prev").glob("*/option_A.py"))
        r = H.finalize(out, "--slide", "2")
        assert r.returncode == 11, f"expected 11, got {r.returncode}"
        print("    ok: old files in _prev/, finalize waits for the new design")
    finally:
        H.cleanup(tmp)

    print("[5] untranslated sketches: unpicked ones wait, a picked one needs the translator")
    tmp, out = H.new_build(2)
    saved = os.environ.get("SLIDE_LAB_TRANSLATOR")
    os.environ["SLIDE_LAB_TRANSLATOR"] = "agent"
    try:
        H.write_option(out, 1)
        (out / "slide_02" / "option_A.html").write_text("<html></html>", encoding="utf-8")
        r = H.finalize(out)
        assert r.returncode == 0, (
            "before any pick, a sketch is reviewed as a sketch, not 'missing' "
            f"(exit {r.returncode})\n{r.stdout[-800:]}")
        assert _record(out, {"slide_01": "A", "slide_02": "A"}).returncode == 0
        r = H.finalize(out)
        assert r.returncode == 11, r.returncode
        msg = r.stdout + r.stderr
        assert "slide-builder-translator" in msg and "slide_02/option_A" in msg, msg[-1200:]
        print("    ok (agent mode): waits before picks; after the pick, exit 11 names the translator")
    finally:
        if saved is None:
            os.environ.pop("SLIDE_LAB_TRANSLATOR", None)
        else:
            os.environ["SLIDE_LAB_TRANSLATOR"] = saved
        H.cleanup(tmp)

    print("[5b] script mode: recording the picks translates the picked sketch on the spot")
    tmp, out = H.new_build(2)
    saved = os.environ.get("SLIDE_LAB_TRANSLATOR")
    os.environ["SLIDE_LAB_TRANSLATOR"] = "script"
    try:
        H.write_option(out, 1)
        (out / "slide_02" / "option_A.html").write_text(
            '<html><body style="margin:0"><div class="slide-canvas" style="position:relative;'
            'width:1280px;height:720px"><div data-shape-id="box" style="position:absolute;'
            'left:100px;top:200px;width:300px;height:120px;background:#1f4e79"></div>'
            '<p data-shape-id="note" style="position:absolute;left:100px;top:360px;margin:0;'
            'font:16px Arial">A short note under the box.</p></div></body></html>',
            encoding="utf-8")
        r = _record(out, {"slide_01": "A", "slide_02": "A"})
        assert r.returncode == 0, r.stdout[-800:]
        sd = out / "slide_02"
        for f in ("option_A_native.py", "option_A_native.plan.json", "option_A_translation_report.json"):
            assert (sd / f).exists(), f"record_picks did not translate: {f} missing"
        assert "translated 1 picked sketch" in r.stdout, r.stdout[-800:]
        r = H.finalize(out)
        assert r.returncode == 0, f"finalize after script translation: exit {r.returncode}\n{r.stdout[-800:]}"
        print("    ok: native script, plan and report written at pick time; finalize builds it")
    finally:
        if saved is None:
            os.environ.pop("SLIDE_LAB_TRANSLATOR", None)
        else:
            os.environ["SLIDE_LAB_TRANSLATOR"] = saved
        H.cleanup(tmp)

    print("[6] a pinned supplied page is carried through, and gates compile")
    import tempfile
    tmp = Path(tempfile.mkdtemp(prefix="slidelab_e2e_pin_"))
    try:
        brief = tmp / "brief.md"
        slides = H.SLIDE.format(n=1).replace(
            "**Slide type:** Content",
            "**Slide type:** Content\n**Pinned source page:** client_onepager.pptx slide 1")
        brief.write_text(H.BRIEF.format(slides=slides), encoding="utf-8")
        assert H.run("seal_brief.py", "--brief", brief, "--accepted", "test brief").returncode == 0
        out = tmp / "out"
        r = H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE,
                  "--out", out, "--pattern", "direct")
        assert r.returncode == 0, r.stderr[-800:]
        import json as _json
        meta = _json.loads((out / "_meta.json").read_text(encoding="utf-8"))
        assert meta["slides"][0].get("pinned_source_page", "").startswith("client_onepager"), \
            meta["slides"][0]
        prompt = (out / "slide_01" / "_prompt.md").read_text(encoding="utf-8")
        assert "option A reproduces that page" in prompt, "the worker is not told to reproduce it"
        H.write_option(out, 1)
        assert H.finalize(out).returncode == 0
        html = _review(out) and (out / "REVIEW.html").read_text(encoding="utf-8")
        assert "reproduces the page you supplied" in html
        assert _record(out, {"slide_01": "A"}).returncode == 0
        assert H.finalize(out).returncode == 0
        r, tok = _final(out)
        assert r.returncode == 0, r.stderr
        r = _compile(out, tok)
        assert r.returncode == 5 and "never reconciled" in r.stdout, r.stdout[-600:]
        _state.record_source_ledger(out, unresolved=0, keep_source=1)
        r = _compile(out, tok)
        assert r.returncode == 0, r.stdout[-1200:]
        print("    ok: in _meta, in the prompt, on the review page; compile waits for the ledger")
    finally:
        H.cleanup(tmp)

    print("[7] Leave out: the user's choice, not a compile flag")
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1)
        H.write_option(out, 2)
        assert H.finalize(out).returncode == 0
        r = _record(out, {"slide_01": "-", "slide_02": "-"})
        assert r.returncode == 5 and "nothing to build" in r.stdout, r.stdout
        assert _record(out, {"slide_01": "A", "slide_02": "-"}).returncode == 0
        assert H.finalize(out).returncode == 0
        r, tok = _final(out)
        assert r.returncode == 0, r.stderr
        r = _compile(out, tok, "--drop", "1")
        assert r.returncode == 5 and "retired" in r.stdout, "--drop must not work any more"
        r = _compile(out, tok)
        assert r.returncode == 0, r.stdout[-800:]
        assert _state.read_state(out)["stages"]["compile"]["slides"] == 1
        print("    ok: left-out slide omitted; --drop refused; all-left-out refused")
    finally:
        H.cleanup(tmp)

    print("[8] --insert moves the per-option records with the folders")
    tmp, out = H.new_build(2)
    try:
        H.write_option(out, 1)
        H.write_option(out, 2, body=CRASHING_SCRIPT)       # slide 2 blocked
        assert H.finalize(out).returncode == 15
        brief = tmp / "brief.md"
        brief.write_text(H.BRIEF.format(slides="".join(
            H.SLIDE.format(n=i) for i in (1, 2, 3))), encoding="utf-8")
        assert H.run("seal_brief.py", "--brief", brief, "--accepted", "test brief").returncode == 0
        r = H.run("build_deck.py", "--brief", brief, "--template", H.TEMPLATE,
                  "--out", out, "--pattern", "direct", "--insert", "2")
        assert r.returncode == 0, r.stderr[-800:]
        by = _state.read_state(out)["qc"]["by_option"]
        assert by.get("slide_03/A", {}).get("blocks", 0) > 0, (
            f"the blocked slide moved to 3 but its record did not: {by}")
        assert "slide_02/A" not in by, f"the new slide 2 inherited a record: {by}"
        print("    ok: the block moved to slide 3; the new slide 2 has no record")
    finally:
        H.cleanup(tmp)

    print("[9] the printed steps for a slide rebuild stay scoped to that slide")
    import build_deck
    lines = "\n".join(build_deck.review_sequence_lines("OUT", "TPL", 4))
    assert lines.count("--slide 4") == 2, (
        "both finalize steps must be scoped to the rebuilt slide; a full finalize "
        "breaks the 6b flow:\n" + lines)
    print("    ok")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
