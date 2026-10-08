#!/usr/bin/env python3
"""Smoke test: a failed LibreOffice render is retried once, and its error is kept.

Renders fail now and then (a non-zero exit or no PDF) and a second try
usually works; the session had to rerun them by hand. finalize also cut the
error to 50 characters, so the cause was lost (2026-10-08).

  - a renderer that fails on its first call: one automatic retry succeeds,
    the second try renders from a local copy with a fresh profile
  - a deck under a OneDrive folder is rendered from a local copy
  - when both tries fail, the error carries the full output of both
  - finalize prints the whole error, not 50 characters of it
  - two renders at the same time both finish (renders are not serialized)

Needs LibreOffice, like the other render smokes.
Run:  py -3 slide-builder/tests/run_render_retry_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
import threading
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

import render_slides as R  # noqa: E402

LONG_ERR = ("simulated renderer failure: " + "the full reason is this long line of text " * 4).strip()


def _deck(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "Render retry test"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _input_of(cmd: list) -> str:
    return cmd[-1]


def main() -> int:
    real = R._run_soffice
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        deck = _deck(td / "plain" / "deck.pptx")

        print("[1] fails once, then the automatic retry renders it")
        calls = []

        def fail_once(cmd):
            calls.append(list(cmd))
            if len(calls) == 1:
                return subprocess.CompletedProcess(cmd, 1, "", LONG_ERR)
            return real(cmd)
        R._run_soffice = fail_once
        try:
            R.render_libre(deck, td / "png1", 40)
        finally:
            R._run_soffice = real
        assert (td / "png1" / "slide_01.png").exists(), "the retry did not render"
        assert len(calls) == 2, len(calls)
        assert _input_of(calls[0]) == str(deck), calls[0]
        assert "try2" in _input_of(calls[1]) and _input_of(calls[1]) != str(deck), calls[1]
        assert calls[0][1] != calls[1][1], "the retry reused the first profile"
        print("    ok: 2 calls; the second from a local copy with a fresh profile")

        print("[2] a deck under a OneDrive folder renders from a local copy")
        synced = _deck(td / "OneDrive - Test Org" / "deck.pptx")
        seen = []

        def record(cmd):
            seen.append(list(cmd))
            return real(cmd)
        R._run_soffice = record
        try:
            R.render_libre(synced, td / "png2", 40)
        finally:
            R._run_soffice = real
        assert (td / "png2" / "slide_01.png").exists()
        assert "onedrive" not in _input_of(seen[0]).lower(), seen[0]
        print("    ok: rendered from", Path(_input_of(seen[0])).parent.name, "copy")

        print("[3] both tries fail: the full error of both is kept")

        n_fail = []

        def always_fail(cmd):
            n_fail.append(1)
            return subprocess.CompletedProcess(cmd, 3, "out-%d" % len(n_fail),
                                               LONG_ERR + " #%d" % len(n_fail))
        R._run_soffice = always_fail
        try:
            R.render_libre(deck, td / "png3", 40)
            raise AssertionError("a failing render raised nothing")
        except RuntimeError as exc:
            msg = str(exc)
        finally:
            R._run_soffice = real
        assert LONG_ERR + " #1" in msg and LONG_ERR + " #2" in msg, msg
        assert "try 1" in msg and "try 2" in msg and "exit code 3" in msg, msg
        print(f"    ok: {len(msg)} characters of error kept")

        print("[4] finalize prints the whole error")
        import finalize_deck as fd
        flag = fd._fail_flag("render: RuntimeError: " + msg)
        assert LONG_ERR + " #2" in flag, flag
        print("    ok")

        print("[5] two renders at once both finish")
        decks = [_deck(td / f"par{k}" / "deck.pptx") for k in range(2)]
        errors = []

        def go(k):
            try:
                R.render_libre(decks[k], td / f"pngp{k}", 40)
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)
        th = [threading.Thread(target=go, args=(k,)) for k in range(2)]
        [t.start() for t in th]
        [t.join(300) for t in th]
        assert not errors, errors
        assert all((td / f"pngp{k}" / "slide_01.png").exists() for k in range(2))
        print("    ok")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
