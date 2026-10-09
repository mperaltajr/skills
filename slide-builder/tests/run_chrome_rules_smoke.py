#!/usr/bin/env python3
"""Smoke test: the takeaway, footnotes and source line sit in one place on every slide.

Owner's rule (2026-10-08): Slide Lab places these three itself, whatever the
designer drew:
  takeaway   directly under the title, at the title's text left edge, the
             template's Subtitle-slot size (else 16 pt), not bold, the
             template's main text color
  source     at the template's source position, title's text left edge, 9 pt
  footnotes  directly above the source line, stacked upward, same edge, 9 pt

A made-up deck on the bundled test template, through the real pipeline: two
direct-path pages and one sketch page whose designers put the takeaway at
different places, sizes and weights, footnotes mid-page and the source line
at different x positions and sizes. In the compiled deck:
  - every slide's takeaway, source line and footnotes have the same left
    edge (the title's text edge), top and size
  - footnotes stack directly above the source line, in order
  - the takeaway is regular weight in the theme's text color
The registration mock slide, made from a copy of the same template, shows the
same three positions and sizes.

Needs Playwright's Chromium and LibreOffice, like the other render smokes.
Run:  py -3 slide-builder/tests/run_chrome_rules_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402
import _chrome_schema as CS  # noqa: E402

EMU = 9525

DIRECT = '''# Slide {n} option A -- pattern: text column (test harness)
# Variant: plain
# Brief fidelity check: test fixture
import sys
from pathlib import Path
sys.path.insert(0, r"{skill}")
from twins.helpers import new_slide, add_text
from pptx.dml.color import RGBColor


def build():
    prs, slide = new_slide()
    add_text(slide, "subtitle", "Takeaway {n}.", {sx}, {sy}, 500, 40,
             font_size_pt={spt}, bold={sbold}, color=RGBColor(0xC0, 0x30, 0x30))
    add_text(slide, "body", "Point one\\nPoint two", 80, 260, 600, 120, font_size_pt=14)
{notes}
    add_text(slide, "source", "Source: made-up figures for page {n}", {cx}, {cy}, 520, 20,
             font_size_pt={cpt})
    return prs


if __name__ == "__main__":
    build().save(str(Path(__file__).resolve().parent / "option_A.pptx"))
'''

NOTE = '    add_text(slide, "footnote-{k}", "{k}. A note the designer put mid-page ({k})", {x}, {y}, 400, 24, font_size_pt={pt})\n'

SKETCH = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0;padding:0}
.slide-canvas{position:relative;width:1280px;height:720px;background:#fff;font-family:Arial}
p,h1{margin:0;position:absolute}
</style></head><body><div class="slide-canvas">
<h1 data-template-field="title" style="left:48px;top:60px;font-size:28px">Claim number 3 is stated plainly here</h1>
<p data-template-field="subtitle" style="left:420px;top:170px;font-size:22px;font-weight:700;color:#2a6">Takeaway 3, drawn big and bold.</p>
<p data-shape-id="body" style="left:100px;top:260px;font-size:16px">Body text for the sketch page.</p>
<p data-shape-id="footnote-1" style="left:640px;top:420px;font-size:15px">1. A footnote drawn mid-page</p>
<p data-shape-id="footnote-2" style="left:200px;top:470px;font-size:13px">2. A second footnote somewhere else</p>
<p data-template-field="footer" style="left:520px;top:684px;font-size:14px">Source: made-up figures for page 3</p>
</div></body></html>"""


def _review(out: Path) -> str:
    r = H.run("build_review.py", "--out", out)
    assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
    return _state.read_state(out)["review"]["token"]


def _line(token: str, picks) -> str:
    body = _state.canonical_picks(picks)
    return f"PICKS {body} CHECK {_state.approval_check(token, body)}"


def _size(sh):
    from pptx.oxml.ns import qn
    sizes = {int(r.get("sz")) / 100 for r in sh._element.iter(qn("a:rPr")) if r.get("sz")}
    return sizes.pop() if len(sizes) == 1 else tuple(sorted(sizes))


def _bold(sh):
    from pptx.oxml.ns import qn
    return {r.get("b") for r in sh._element.iter(qn("a:rPr"))}


def _text_left(sh) -> int:
    from twins.chrome_rules import _inset
    return int(sh.left) + _inset(sh, "lIns")


def _roles(slide):
    from twins.chrome_rules import _ph_type
    take, src, notes = None, None, []
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False) or not sh.text_frame.text.strip():
            continue
        name, t = (sh.name or "").lower(), _ph_type(sh)
        if t == 4 or (t is None and name.startswith("subtitle")):
            take = sh
        elif t == 15 or (t is None and name.startswith("source")):
            src = sh
        elif name.startswith("footnote"):
            notes.append(sh)
    return take, src, sorted(notes, key=lambda s: int(s.top))


def _check_slide(slide, label: str, takeaway_pt: float) -> dict:
    from twins.chrome_rules import _title_geometry
    take, src, notes = _roles(slide)
    assert take is not None and src is not None, f"{label}: takeaway or source missing"
    left, _right, title_bottom = _title_geometry(slide, None)
    for sh in [take, src, *notes]:
        assert abs(_text_left(sh) - left) <= EMU, \
            f"{label}: {sh.name} text starts at {_text_left(sh) / EMU:.1f}px, title at {left / EMU:.1f}px"
    assert _size(take) == takeaway_pt, f"{label}: takeaway size {_size(take)}"
    assert _bold(take) <= {"0", None}, f"{label}: takeaway bold {_bold(take)}"
    assert _size(src) == CS.CHROME_NOTE_FONT_PT, f"{label}: source size {_size(src)}"
    for nt in notes:
        assert _size(nt) == CS.CHROME_NOTE_FONT_PT, f"{label}: {nt.name} size {_size(nt)}"
    # footnotes stacked directly above the source line, in order
    edge = int(src.top)
    for nt in reversed(notes):
        gap = (edge - (int(nt.top) + int(nt.height))) / EMU
        assert 0 <= gap <= CS.CHROME_NOTE_GAP_PX + 0.5, f"{label}: {nt.name} {gap:.1f}px above"
        edge = int(nt.top)
    names = [nt.name for nt in notes]
    assert names == sorted(names), f"{label}: footnotes out of order {names}"
    return {"take": (int(take.left), int(take.top), _size(take)),
            "src": (int(src.left), int(src.top), _size(src)),
            "note_left": {int(n.left) for n in notes}, "title_bottom": title_bottom,
            "take_color": [c.get("val") for c in take._element.iter(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}schemeClr")]}


def main() -> int:
    from pptx import Presentation
    print("[1] three pages drawn three ways, through the pipeline")
    tmp, out = H.new_build(3)
    saved = os.environ.get("SLIDE_LAB_TRANSLATOR")
    os.environ["SLIDE_LAB_TRANSLATOR"] = "script"
    try:
        H.write_option(out, 1, body=DIRECT.format(
            n=1, skill=str(SKILL), sx=300, sy=140, spt=20, sbold=True,
            notes=NOTE.format(k=1, x=500, y=400, pt=14), cx=300, cy=650, cpt=12))
        H.write_option(out, 2, body=DIRECT.format(
            n=2, skill=str(SKILL), sx=64, sy=180, spt=12, sbold=False,
            notes=NOTE.format(k=1, x=100, y=300, pt=11) + NOTE.format(k=2, x=150, y=340, pt=16),
            cx=700, cy=600, cpt=8))
        (out / "slide_03" / "option_A.html").write_text(SKETCH, encoding="utf-8")
        picks = {"slide_01": "A", "slide_02": "A", "slide_03": "A"}
        r = H.run("record_picks.py", "--out", out, "--approved", _line(_review(out), picks))
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-1500:]
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-2500:] + r.stderr[-2500:]
        r = H.run("build_review.py", "--out", out, "--final")
        tok = (_state.read_state(out).get("final_check") or {}).get("token")
        assert r.returncode == 0 and tok, r.stdout[-1500:] + r.stderr[-1500:]
        r = H.run("compile_picks.py", "--out", out, "--final-token", tok)
        assert r.returncode == 0, r.stdout[-2500:] + r.stderr[-2500:]
        deck = Presentation(str(_state.compiled_deck(out)))
        from twins.chrome_rules import template_takeaway_pt
        from _chrome_schema import load_chrome_yml
        import _paths as _p
        t_pt = template_takeaway_pt(Presentation(str(H.TEMPLATE)), load_chrome_yml(_p.chrome_yml(H.TEMPLATE)))
        print(f"    ok: compiled; the template's takeaway size is {t_pt:g} pt")

        print("[2] every slide: same left edge, top and size for the three lines")
        got = [_check_slide(s, f"slide {i}", t_pt) for i, s in enumerate(deck.slides, 1)]
        for g in got[1:]:
            assert g["take"] == got[0]["take"], (g["take"], got[0]["take"])
            assert g["src"] == got[0]["src"], (g["src"], got[0]["src"])
        lefts = set().union(*[g["note_left"] for g in got])
        assert len(lefts) == 1 and lefts == {got[0]["src"][0]}, lefts
        assert got[0]["take"][1] >= got[0]["title_bottom"], "takeaway is not under the title"
        assert all("tx1" in g["take_color"] for g in got), [g["take_color"] for g in got]
        print(f"    ok: takeaway at ({got[0]['take'][0] / EMU:.0f}, {got[0]['take'][1] / EMU:.0f}) px "
              f"{got[0]['take'][2]:g} pt; source at ({got[0]['src'][0] / EMU:.0f}, "
              f"{got[0]['src'][1] / EMU:.0f}) px {got[0]['src'][2]:g} pt on all 3 slides")
    finally:
        if saved is None:
            os.environ.pop("SLIDE_LAB_TRANSLATOR", None)
        else:
            os.environ["SLIDE_LAB_TRANSLATOR"] = saved
        H.cleanup(tmp)

    print("[3] the registration mock slide shows the same")
    import _paths as _p
    import register_template as RT
    td = Path(tempfile.mkdtemp(prefix="slidelab_chrome_mock_"))
    try:
        tpl = td / H.TEMPLATE.name
        shutil.copy2(H.TEMPLATE, tpl)
        shutil.copytree(_p.template_sidecar_dir(H.TEMPLATE), _p.template_sidecar_dir(tpl))
        tj = _p.theme_json(tpl)
        data = json.loads(tj.read_text(encoding="utf-8"))
        data["default_content_layout"] = "body_canonical_light"
        tj.write_text(json.dumps(data, indent=2), encoding="utf-8")
        fails, infos = RT._render_mock_page_selftest(tpl)
        assert not fails, fails
        mock = Presentation(str(_p.selftest_pptx(tpl)))
        m = _check_slide(mock.slides[0], "mock slide", t_pt)
        assert m["take"] == got[0]["take"], (m["take"], got[0]["take"])
        assert m["src"] == got[0]["src"], (m["src"], got[0]["src"])
        print("    ok: mock slide's takeaway, footnote and source match the built slides")
    finally:
        shutil.rmtree(td, ignore_errors=True)

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
