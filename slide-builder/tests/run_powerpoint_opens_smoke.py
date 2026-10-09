#!/usr/bin/env python3
"""Smoke test: PowerPoint's own opens-cleanly check runs before a deck is
called done (owner's decision, 2026-10-09).

compile_picks opens the saved deck in PowerPoint (read-only, no window, via
slide-qc/scripts/ppt_safe.py) before it replaces the previous deck: no repair
question and the same slide count, or the deck is refused.

  [1] a clean fictional deck opens: ok
  [2] a deck PowerPoint refuses (an incomplete <p:style>, the defect that
      shipped once) is caught, with PowerPoint's own words, and compile's
      check returns 6 (refused)
  [3] a deck with duplicate shape ids: the offline check refuses it; PowerPoint
      itself opens it (it repairs this one silently), which is why both
      checks run
  [4] a silent repair that loses slides (stand-in PowerPoint reporting one
      slide fewer) is caught by the slide-count comparison
  [5] without PowerPoint (SLIDE_LAB_NO_POWERPOINT=1, a Mac) the check says
      it did not run and does not block

Fictional files only; the user's open PowerPoint presentations are recorded
before and after (attaching only) and must be unchanged.
Run:  py -3 slide-builder/tests/run_powerpoint_opens_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import contextlib
import io
import os
import sys
import tempfile
import types
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
QC = SKILL.parent / "slide-qc" / "scripts"
for p in (str(SKILL), str(SKILL / "scripts"), str(QC)):
    sys.path.insert(0, p)

import ppt_safe  # noqa: E402
import render_slides as RS  # noqa: E402
from _ppt_snapshot import snapshot  # noqa: E402


def _deck(path: Path, kind: str = "good") -> Path:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
        s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(1), Inches(3), Inches(2))
        s.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1)).text_frame.text = \
            f"Fictional {i + 1}"
    if kind == "style":
        st = prs.slides[1].shapes[0]._element.find(qn("p:style"))
        st.remove(st.find(qn("a:effectRef")))
    if kind == "dupids":
        for s in prs.slides:
            for c in s.shapes._spTree.iter(qn("p:cNvPr")):
                c.set("id", "2")
    prs.save(str(path))
    return path


def main() -> int:
    if not RS.powerpoint_available():
        print("  skipped: PowerPoint can't be driven on this computer")
        print("SMOKE PASSED.")
        return 0
    import compile_picks as CP
    before = snapshot()
    with tempfile.TemporaryDirectory() as tdn:
        td = Path(tdn)
        print("[1] a clean deck opens cleanly")
        ok, why = CP.powerpoint_opens(_deck(td / "good.pptx"))
        assert ok is True and not why, why
        print("    ok")

        print("[2] a deck PowerPoint refuses is caught")
        bad = _deck(td / "style.pptx", "style")
        ok, why = CP.powerpoint_opens(bad)
        assert ok is False and why, why
        with contextlib.redirect_stdout(io.StringIO()) as buf:
            rc = CP._report_powerpoint_opens(bad)
        assert rc == 6 and "REFUSED" in buf.getvalue(), buf.getvalue()
        print(f"    ok: refused, PowerPoint says: {why[:70]}")

        print("[3] duplicate shape ids: the offline check refuses, PowerPoint repairs")
        dup = _deck(td / "dupids.pptx", "dupids")
        problems = CP.assert_package_integrity(dup)
        assert any("used more than once" in p for p in problems), problems
        ok, why = CP.powerpoint_opens(dup)
        print(f"    ok: offline check refused it; PowerPoint alone says "
              f"{'opens' if ok else 'refuses: ' + why[:60]}")

        print("[4] a silent repair that loses a slide is caught")

        class _Pres:
            Slides = types.SimpleNamespace(Count=2)

        @contextlib.contextmanager
        def fake_session():
            yield object(), False

        @contextlib.contextmanager
        def fake_opened(app, p):
            yield _Pres()
        real = ppt_safe.session, ppt_safe.opened
        ppt_safe.session, ppt_safe.opened = fake_session, fake_opened
        try:
            res = ppt_safe.check_opens(td / "good.pptx", expected_slides=3)
        finally:
            ppt_safe.session, ppt_safe.opened = real
        assert not res["ok"] and "repaired" in res["error"], res
        print(f"    ok: {res['error']}")

        print("[5] no PowerPoint here: not run, not blocking")
        os.environ["SLIDE_LAB_NO_POWERPOINT"] = "1"
        try:
            ok, why = CP.powerpoint_opens(bad)
            with contextlib.redirect_stdout(io.StringIO()):
                rc = CP._report_powerpoint_opens(bad)
        finally:
            os.environ.pop("SLIDE_LAB_NO_POWERPOINT", None)
        assert ok is None and rc == 0 and "not installed" in why, (ok, rc, why)
        print("    ok")

    after = snapshot()
    assert before[1] == after[1], "the user's open presentations changed"
    if before[0]:
        assert after[0], "PowerPoint was running before and is not now"
    print(f"  ok: the user's PowerPoint: {len(before[1])} presentation(s) open before, "
          f"the same {len(after[1])} after")
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
