#!/usr/bin/env python3
"""Smoke test: the pipeline must not ship a deck PowerPoint refuses to open.

Background. An 87-slide deck was delivered twice showing "PowerPoint found a
problem with content... Repair", then "could not open the file". python-pptx read
it. LibreOffice rendered every slide. check_pptx_hygiene reported zero violations.
Nothing in the pipeline looked at the structures PowerPoint actually validates.

Three of them are known to produce that symptom, and this test covers each:
  1. <p:style> missing a child. The element takes exactly lnRef, fillRef,
     effectRef, fontRef, in order. Deleting effectRef is the obvious-looking way
     to stop LibreOffice drawing a preset shadow, and it invalidates the file.
  2. Duplicate shape ids, which every graft produces unless it renumbers.
  3. A graphicFrame or picture whose r:id points at a relationship the slide
     part does not have.

Also asserts the two graft sites renumber and that both gates run the checks —
a fix at one site lets the other keep shipping the defect.

Run:  py -3 slide-builder/tests/run_openability_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import inspect
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from twins.composer import reassign_shape_ids  # noqa: E402
from pptx_openability import check_openability  # noqa: E402

_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')

_STYLE_OK = (
    f'<p:style {_NS}>'
    '<a:lnRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:lnRef>'
    '<a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef>'
    '<a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef>'
    '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>'
    '</p:style>'
)
_STYLE_NO_EFFECTREF = _STYLE_OK.replace(
    '<a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef>', '')


def _ids(slide) -> list[str]:
    return [c.get("id") for c in slide.shapes._spTree.iter(qn("p:cNvPr"))]


def _blank_slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def main() -> int:
    print("1. a <p:style> with effectRef deleted is caught")
    prs, slide = _blank_slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box._element.append(etree.fromstring(_STYLE_NO_EFFECTREF))
    found = check_openability(prs)
    assert any("effectRef" in v["issue"] for v in found), \
        f"the incomplete <p:style> was not reported: {found}"
    assert all(v["severity"] == "Critical" for v in found), \
        "an unopenable deck must not be reported as advisory"
    print(f"    ok: {found[0]['issue'][:78]}...")

    print("2. a complete <p:style> with effectRef idx=0 is accepted")
    prs, slide = _blank_slide()
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box._element.append(etree.fromstring(_STYLE_OK))
    assert check_openability(prs) == [], \
        "the sanctioned no-shadow form (effectRef idx=0) was rejected"
    print("    ok: the sanctioned no-shadow form passes")

    print("3. the graft's id collision is caught, and reassign resolves it")
    prs, slide = _blank_slide()
    for i in range(4):
        slide.shapes.add_textbox(Inches(0.5 + i), Inches(1), Inches(1), Inches(1))
    for cNvPr in list(slide.shapes._spTree.iter(qn("p:cNvPr")))[1:]:
        cNvPr.set("id", "2")   # what four deepcopied source trees leave behind
    assert check_openability(prs), "a real id collision was not reported"
    changed = reassign_shape_ids(slide)
    after = _ids(slide)
    assert changed == 3, f"expected 3 shapes moved, got {changed}"
    assert len(set(after)) == len(after), f"ids still collide: {after}"
    assert after[0] == "1", "the root group's id must not be reassigned"
    assert check_openability(prs) == []
    print(f"    ok: {after}")

    print("4. reassign is idempotent and leaves unique ids where they were")
    snapshot = list(after)
    assert reassign_shape_ids(slide) == 0, "a second pass moved already-unique shapes"
    assert _ids(slide) == snapshot, "a second pass renumbered unique ids"
    print("    ok: no churn on a clean slide")

    print("5. a dangling relationship is caught")
    prs, slide = _blank_slide()
    pic_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    blip = etree.fromstring(
        f'<a:blip {_NS} xmlns:r="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships" r:embed="rId9999"/>')
    pic_box._element.append(blip)
    found = check_openability(prs)
    assert any("rId9999" in v["issue"] for v in found), \
        f"a reference to a relationship that does not exist was not reported: {found}"
    print("    ok: the missing relationship is named")

    print("6. both graft sites renumber")
    import finalize_deck
    import compile_picks

    assert "reassign_shape_ids" in inspect.getsource(finalize_deck.graft_and_theme), (
        "finalize_deck's graft no longer renumbers shape ids; every themed option "
        "would ship a deck PowerPoint refuses to open")
    assert "reassign_shape_ids" in inspect.getsource(
        compile_picks.copy_picked_slide_into), \
        "compile_picks's graft no longer renumbers shape ids"
    print("    ok: finalize_deck and compile_picks both renumber")

    print("7. an embedded chart is blocked, not merely warned about")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    prs, slide = _blank_slide()
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1.0, 2.0))
    slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(1), Inches(1),
                           Inches(3), Inches(3), data)
    tmp = Path(__file__).resolve().parent / "_tmp_embedded_chart.pptx"
    try:
        prs.save(str(tmp))
        res = finalize_deck._check_no_embedded_objects(tmp)
        assert not res["pass"], "an embedded chart was accepted"
        assert "proportion bar" in res["detail"], \
            f"the refusal does not name the sanctioned alternative: {res['detail']}"
    finally:
        tmp.unlink(missing_ok=True)
    print(f"    ok: {res['detail'][:76]}...")

    print("8. both gates run the checks")
    assert "check_openability" in inspect.getsource(
        compile_picks.assert_package_integrity), (
        "the compile gate no longer checks openability; a post-compile hand-edit "
        "could ship a deck that will not open")
    import check_pptx_hygiene
    assert "check_openability" in inspect.getsource(check_pptx_hygiene.run_all_checks), \
        "the QC pre-pass no longer checks openability"
    print("    ok: compile's refusal gate and the QC pre-pass both check")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
