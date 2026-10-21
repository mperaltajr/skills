#!/usr/bin/env python3
"""Smoke test for the context design (owner's decisions, 2026-10-14).

  [1] add_ruling.py: a paraphrase is refused; the user's real words are
      stored with the chat's reading on its own line; a newer ruling on the
      same topic strikes the older (both shown); a repeat says REPEAT;
      compaction summaries, tool results and agent threads do not count.
  [2] slide ids: added once, kept, never reused; the seal survives them (and a
      brief sealed the old way still builds once it has its ids).
  [3] prep: the slide id reaches _meta.json and REVIEW-style labels; every
      _context.md carries the pack (under 2 KB, the ruling verbatim); a
      hand-written "Direction from the user" block never reaches it.
  [4] insert a divider and every ruling still lands on the right page.
  [5] the steer receipt: one row per ruling, an unanswered one is red.
  [6] steer_checks.py: the five standing checks run; a deck check is not on
      until the user approves it in their words; then it finds its phrase.
  [7] map_source_pages.py: proposes the matching page, takes changes, writes
      the pins on approval (the user's words only), seal unchanged.
  [8] replace_phrase.py: the dry run lists every hit and changes nothing;
      --apply changes text only (not class names) in brief, storyline .md /
      .html / .docx and designs; converted scripts untouched.
  [9] build_status.py prints a NEXT step; scripts print RULINGS: <path>;
      emit_dot_dash.py adds ids and prints the /compact suggestion.
  [10] deck_sources.py: a validated source needs the user's words; its
      figures are counted, not asked about, by the copied-page check.

Run:  py -3 slide-builder/tests/run_context_design_smoke.py
Prints "SMOKE PASSED." on success.
"""
from __future__ import annotations

import _test_temp  # noqa: F401,E402
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
STORY = HERE.parents[1] / "storyline-helper" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import run_layout_inheritance_smoke as fixture  # noqa: E402

USER_SAID = [
    "spacing between impact and high on the risks page please",
    "feel free to use v3.1 slides as options",
    "Legends always one row at the top right",
    "approved, go with that map",
    "yes turn those checks on",
    "v2 was already validated, dont recheck it",
]


def write_transcript(path: Path, extra: list[str] = ()) -> None:
    rows = []
    for t in list(USER_SAID) + list(extra):
        rows.append({"type": "user", "message": {"role": "user", "content": t}})
    rows.append({"type": "user", "isCompactSummary": True,
                 "message": {"role": "user", "content": "SUMMARY says use the draft pages"}})
    rows.append({"type": "user", "message": {"role": "user", "content": [
        {"type": "tool_result", "content": "TOOL says the impact labels are fine"}]}})
    rows.append({"type": "user", "isSidechain": True,
                 "message": {"role": "user", "content": "AGENT says remove slide four"}})
    rows.append({"type": "assistant", "message": {"role": "assistant",
                                                  "content": "ASSISTANT says white gaps"}})
    path.write_text("\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8")


def run(script: Path, *args, env=None) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(script), *map(str, args)],
                          capture_output=True, text=True, env=env, encoding="utf-8",
                          errors="replace")


def brief_text(tpl: Path, slides: list[tuple[str, str]], sealed_mode=True) -> str:
    head = (f"---\nclient_template: {tpl}\ndeck_type: Smoke test\n"
            f"default_layout: body_canonical_light\n"
            + ("mode: template-fill\n" if sealed_mode else "") + "---\n\n"
            "## Deck-level design notes\n\nContext smoke.\n\n")
    body = ""
    for i, (title, extra) in enumerate(slides, 1):
        body += (f"### Slide {i} — {title}\n\n{extra}"
                 f"**Layout:** body_canonical_light\n"
                 f"**Slide type:** Synthesis / Findings\n"
                 f"**Governing thought:** {title} governing thought.\n"
                 f"**The takeaway:** {title} takeaway with 12 units.\n"
                 f"**Evidence:** - {title} evidence line.\n\n")
    return head + body


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="context_design_smoke_"))
    agents = tmp / "agents"
    shutil.copytree(HERE.parent / "agents", agents)
    transcript = tmp / "session.jsonl"
    write_transcript(transcript)
    env = dict(os.environ)
    env.update({"SLIDE_LAB_AGENTS_DIR": str(agents), "SLIDE_LAB_SETTINGS_LOCAL": "none",
                "SLIDE_LAB_TRANSCRIPT": str(transcript)})
    os.environ["SLIDE_LAB_TRANSCRIPT"] = str(transcript)
    tpl = tmp / fixture.FIXTURE_PPTX.name
    shutil.copy2(fixture.FIXTURE_PPTX, tpl)
    fixture.register_fixture(tpl)
    deck = tmp / "deck"
    sd = deck / "_session"
    sd.mkdir(parents=True)
    brief = sd / "narrative-brief-smoke.md"
    brief.write_text(brief_text(tpl, [("Revenue grows fast", ""), ("The biggest risks have answers", ""),
                                      ("Leaders built big businesses", "")]), encoding="utf-8")

    print("[2] slide ids")
    import slide_ids
    r = run(SCRIPTS / "slide_ids.py", "--brief", brief, env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    ids = slide_ids.ids_in_brief(brief.read_text(encoding="utf-8"))
    assert ids == {1: "revenue", 2: "risks", 3: "leaders"}, ids
    assert "RULINGS:" in r.stdout, r.stdout
    before = brief.read_text(encoding="utf-8")
    r = run(SCRIPTS / "slide_ids.py", "--brief", brief, env=env)
    assert brief.read_text(encoding="utf-8") == before, "second run changed the brief"
    # An id once used is never handed out again: drop "risks", add a new risks slide.
    reg = json.loads((sd / "slide_ids.json").read_text(encoding="utf-8"))
    assert set(reg["ids"]) == {"revenue", "risks", "leaders"}, reg
    probe = tmp / "probe" / "_session"
    probe.mkdir(parents=True)
    pb = probe / "narrative-brief-p.md"
    pb.write_text(brief_text(tpl, [("Risks again", "")]), encoding="utf-8")
    (probe / "slide_ids.json").write_text(json.dumps({"ids": {"risks": {}}}), encoding="utf-8")
    slide_ids.assign(pb)
    assert slide_ids.ids_in_brief(pb.read_text(encoding="utf-8"))[1] != "risks"
    import seal_brief
    import build_deck
    t0 = before.replace("\n\n**Slide id:** revenue", "").replace(
        "\n\n**Slide id:** risks", "").replace("\n\n**Slide id:** leaders", "")
    fm0, b0 = build_deck.extract_front_matter(t0)
    fm1, b1 = build_deck.extract_front_matter(before)
    assert seal_brief.brief_fingerprint(fm0, b0) == seal_brief.brief_fingerprint(fm1, b1), \
        "the seal counts the slide id lines"
    assert seal_brief.brief_fingerprint(fm0, b0, legacy=True) == \
        seal_brief.brief_fingerprint(fm1, b1, legacy=True), "an old seal breaks on ids"
    print("    ok: ids revenue/risks/leaders, kept on a rerun, never reused, seal unchanged")

    print("[1] add_ruling.py")
    ar = SCRIPTS / "add_ruling.py"
    r = run(ar, "--brief", brief, "--id", "risks", "--topic", "impact spacing",
            "--quote", "white gaps between the impact labels and the rows", env=env)
    assert r.returncode == 4 and "REFUSED" in r.stdout, r.stdout
    for fake in ("SUMMARY says use the draft pages", "TOOL says the impact labels are fine",
                 "AGENT says remove slide four", "ASSISTANT says white gaps"):
        r = run(ar, "--brief", brief, "--id", "deck", "--topic", "x", "--quote", fake, env=env)
        assert r.returncode == 4, f"{fake!r} counted as the user's words:\n{r.stdout}"
    r = run(ar, "--brief", brief, "--id", "nosuch", "--topic", "x", "--quote",
            "feel free to use v3.1 slides", env=env)
    assert r.returncode == 3, r.stdout
    r = run(ar, "--brief", brief, "--id", "risks", "--topic", "impact spacing",
            "--quote", "spacing between impact and high",
            "--reading", "a gap between the High label column and the risk rows", env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    r = run(ar, "--brief", brief, "--id", "deck", "--topic", "source pages",
            "--quote", "feel free to use v3.1 slides as options", env=env)
    assert r.returncode == 0, r.stdout
    r = run(ar, "--brief", brief, "--id", "risks", "--topic", "Impact spacing!",
            "--quote", "spacing between impact ... risks page", env=env)
    assert r.returncode == 0 and "REPEAT" in r.stdout and "replaces R1" in r.stdout, r.stdout
    md = (sd / "rulings.md").read_text(encoding="utf-8")
    assert "~~R1" in md and "(replaced by R3)" in md, md
    assert "**Chat's reading:** a gap between" in md, md
    assert '"spacing between impact and high"' in md, md
    print("    ok: paraphrase / summary / tool / agent / assistant text refused; stored verbatim; "
          "R3 strikes R1; REPEAT printed")

    print("[3] prep: ids in _meta, the pack in every _context.md, no chat direction")
    out = deck
    (out / "slide_02").mkdir(parents=True, exist_ok=True)
    (out / "slide_02" / "_prior_feedback.md").write_text(
        "# Direction from the user (13 Oct): EDIT the kept design\n- add white space between "
        "the High / Medium / Low impact labels and the rows\n", encoding="utf-8")
    r = run(SCRIPTS / "build_deck.py", "--brief", brief, "--template", tpl, "--out", out,
            "--confirm-template", env=env)
    assert r.returncode == 0, f"prep failed ({r.returncode}):\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}"
    assert "RULINGS:" in r.stdout, "build_deck did not print RULINGS"
    assert "Direction from the user" in r.stdout and "add_ruling.py" in r.stdout, r.stdout[-2000:]
    meta = json.loads((out / "_meta.json").read_text(encoding="utf-8"))
    assert [s.get("slide_id") for s in meta["slides"]] == ["revenue", "risks", "leaders"], meta
    ctx2 = (out / "slide_02" / "_context.md").read_text(encoding="utf-8")
    assert "Direction from the user" not in ctx2, "chat direction reached the designer"
    import _context_pack
    a, b = ctx2.index(_context_pack.START), ctx2.index(_context_pack.END)
    pack = ctx2[a:b + len(_context_pack.END)]
    assert len(pack.encode("utf-8")) <= _context_pack.MAX_BYTES, len(pack.encode("utf-8"))
    assert 'R3 (Impact spacing!): the user said "spacing between impact ... risks page"' in pack, pack
    assert "R1 " not in pack, "a replaced ruling reached the designer"
    assert "feel free to use v3.1 slides as options" in pack, pack
    ctx1 = (out / "slide_01" / "_context.md").read_text(encoding="utf-8")
    assert "spacing between impact" not in ctx1, "slide 2's ruling reached slide 1"
    assert "feel free to use v3.1" in ctx1, "deck ruling missing on slide 1"
    print(f"    ok: _meta ids; pack {len(pack.encode('utf-8'))} bytes; slide ruling on its "
          "slide only; deck ruling on all; chat direction left out with a warning")

    print("[5] steer receipt")
    import build_review
    slide = {"n": 2, "prompt_path": out / "slide_02" / "_prompt.md"}
    rows = _context_pack.receipt(out, 2)
    assert len(rows) == 1 and not rows[0]["ok"], rows
    (out / "slide_02" / "_steer_answers.txt").write_text(
        "R3: left a 12 px white gap between the impact column and the rows\n", encoding="utf-8")
    rows = _context_pack.receipt(out, 2)
    assert rows[0]["ok"] and "12 px" in rows[0]["answer"], rows
    r = run(ar, "--out", out, "--id", "risks", "--topic", "legend",
            "--quote", "Legends always one row at the top right", env=env)
    assert r.returncode == 0 and "refreshed the rulings pack" in r.stdout, r.stdout
    assert "Legends always one row" in (out / "slide_02" / "_context.md").read_text(encoding="utf-8")
    import time
    time.sleep(0.05)
    rows = _context_pack.receipt(out, 2)
    assert len(rows) == 2 and rows[0]["ok"] and not rows[1]["ok"], rows
    page = build_review.render_steer_receipt(slide)
    assert "Not answered" in page and "Legends always one row" in page and "12 px" in page, page
    print("    ok: answered ruling green, the new unanswered one red on the card")

    print("[4] insert a divider: rulings stay on their page")
    brief.write_text(brief.read_text(encoding="utf-8").replace(
        "### Slide 2 — The biggest", "### Slide 2 — Section two\n\n**Slide type:** Section divider\n"
        "**Layout:** body_canonical_light\n**Governing thought:** Section two.\n\n"
        "### Slide 3 — The biggest").replace("### Slide 3 — Leaders", "### Slide 4 — Leaders"),
        encoding="utf-8")
    slide_ids.assign(brief)
    r = run(SCRIPTS / "build_deck.py", "--brief", brief, "--template", tpl, "--out", out,
            "--insert", "2", "--confirm-template", env=env)
    assert r.returncode == 0, f"insert failed:\n{r.stdout[-1200:]}\n{r.stderr[-1200:]}"
    meta = json.loads((out / "_meta.json").read_text(encoding="utf-8"))
    by_n = {s["n"]: s.get("slide_id") for s in meta["slides"]}
    assert by_n[3] == "risks" and by_n[4] == "leaders" and by_n[2].startswith("divider"), by_n
    assert "spacing between impact" in (out / "slide_03" / "_context.md").read_text(encoding="utf-8")
    assert "spacing between impact" not in (out / "slide_02" / "_context.md").read_text(encoding="utf-8")
    assert [r_["n"] for r_ in _context_pack.receipt(out, 3)] == [3, 4]
    assert _context_pack.receipt(out, 2) == []
    r = run(SCRIPTS / "slide_ids.py", "--brief", brief, "--check", env=env)
    assert "Slide 3 (risks)" in r.stdout, r.stdout
    print(f"    ok: {by_n}; the risks rulings moved with the page to slide 3")

    print("[6] steer_checks.py")
    sc = SCRIPTS / "steer_checks.py"
    (out / "slide_01" / "option_A.html").write_text(
        '<html><body><div data-template-field="title">Revenue grows fast</div>'
        '<div data-template-field="subtitle">Revenue grows fast takeaway with 12 units</div>'
        '<div class="card">The window is now for buyers.</div>'
        '<div class="chart-title">Revenue, [$B]</div>'
        '<div class="so-what band">Revenue grows fast takeaway with 12 units</div>'
        '<div class="footnote">Globex buys GPUs</div><p class="card">Initech too</p>'
        '</body></html>', encoding="utf-8")
    import steer_checks
    found = steer_checks.run_slide(out, 1)
    kinds = {f.split(": ")[1] for f in found}
    assert {"trailing_periods", "unit_no_brackets", "one_summary_line"} <= kinds, found
    assert not any("banned_phrases" in f for f in found), "a deck check ran before approval"
    r = run(sc, "propose", "--out", out, "--kind", "banned_phrases",
            "--param", "phrases=the window is now", "--because", "R5", env=env)
    assert r.returncode == 0 and "P1" in r.stdout, r.stdout
    r = run(sc, "propose", "--out", out, "--kind", "footnote_only", "--param", "terms=Globex|Initech", env=env)
    r = run(sc, "propose", "--out", out, "--kind", "preferred_terms", "--param", "terms=buyers:customers", env=env)
    assert not any("banned_phrases" in f for f in steer_checks.run_slide(out, 1))
    r = run(sc, "approve", "--out", out, "--proposal", "P1,P2,P3",
            "--approved-in-chat", "sure, turn them on", env=env)
    assert r.returncode == 4, r.stdout
    r = run(sc, "approve", "--out", out, "--proposal", "P1,P2,P3",
            "--approved-in-chat", "yes turn those checks on", env=env)
    assert r.returncode == 0, r.stdout
    found = steer_checks.run_slide(out, 1)
    assert any("banned_phrases" in f for f in found), found
    assert any("preferred_terms" in f and "customers" in f for f in found), found
    ft = [f for f in found if "footnote_only" in f]
    assert len(ft) == 1 and "Initech" in ft[0], ft
    meta = json.loads((out / "_meta.json").read_text(encoding="utf-8"))
    meta["slides"][2]["pinned_source_page"] = "x.pptx slide 1"
    (out / "_meta.json").write_text(json.dumps(meta), encoding="utf-8")
    assert any("pinned_option_a" in f for f in steer_checks.run_slide(out, 3))
    meta["slides"][2].pop("pinned_source_page")
    (out / "_meta.json").write_text(json.dumps(meta), encoding="utf-8")
    print(f"    ok: {len(found)} findings on slide 1; deck checks only after the user's words")

    print("[7] map_source_pages.py")
    from pptx import Presentation
    from pptx.util import Inches
    src = tmp / "v2_deck.pptx"
    prs = Presentation()
    texts = ["Quarterly weather in Lisbon and tram schedules",
             "The biggest risks have answers: weaker demand, owners keep GPUs longer; each risk has a named response",
             "Leaders built big businesses at large cloud companies"]
    for t in texts:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = t
    prs.save(str(src))
    mb = sd / "narrative-brief-smoke.md"
    fm_b, body_b = build_deck.extract_front_matter(mb.read_text(encoding="utf-8"))
    fp_before = seal_brief.brief_fingerprint(fm_b, body_b)
    ms = SCRIPTS / "map_source_pages.py"
    r = run(ms, "propose", "--brief", mb, "--from", src, "--no-pictures", env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    pm = json.loads((sd / "page_map.json").read_text(encoding="utf-8"))
    prop = {s["id"]: (s["proposal"] or {}).get("page") for s in pm["slides"]}
    assert prop.get("risks") == 2 and prop.get("leaders") == 3, prop
    assert not (sd / "PAGE-MAP.html").exists(), "the map has no page of its own"
    import map_source_pages
    sec = map_source_pages.render_section(sd)
    assert "Pages you gave us" in sec and "Slide 3 (risks)" in sec and "not used" in sec, sec
    assert "Not approved yet" in sec, sec
    import build_review as _br
    assert "Pages you gave us" in _br._page_map_section({"brief": str(mb), "out": str(out)})
    r = run(ms, "set", "--brief", mb, "leaders=new", "revenue=1", env=env)
    assert r.returncode == 0, r.stdout
    r = run(ms, "approve", "--brief", mb, "--approved-in-chat", "looks fine to me", env=env)
    assert r.returncode == 4, r.stdout
    r = run(ms, "approve", "--brief", mb, "--approved-in-chat", "approved, go with that map", env=env)
    assert r.returncode == 0, r.stdout
    t = mb.read_text(encoding="utf-8")
    pins = {b_["id"]: build_deck.extract_field(t[b_["header_end"]:b_["end"]], ("pinned source page",))
            for b_ in slide_ids.blocks(t)}
    assert pins["risks"].endswith("slide 2") and pins["revenue"].endswith("slide 1"), pins
    assert not pins["leaders"], pins
    fm_a, body_a = build_deck.extract_front_matter(t)
    assert seal_brief.brief_fingerprint(fm_a, body_a) == fp_before, "pins changed the seal"
    r = run(ms, "set", "--brief", mb, "revenue=new", env=env)
    r = run(ms, "approve", "--brief", mb, "--approved-in-chat", "approved, go with that map", env=env)
    t = mb.read_text(encoding="utf-8")
    assert t.count("**Pinned source page:**") == 1, t
    print(f"    ok: proposed {prop}; changes applied; pins written only with the user's words; "
          "a pin removed for 'new'; seal unchanged")

    print("[8] replace_phrase.py")
    (deck / "dot-dash-smoke.md").write_text("Revenue grows for buyers\n", encoding="utf-8")
    (deck / "dot-dash-smoke.html").write_text(
        '<p class="buyers">Buyers and buyers choose</p>', encoding="utf-8")
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    p.add_run("Big buy")
    p.add_run("ers here")
    d.save(str(deck / "dot-dash-smoke.docx"))
    (out / "slide_01" / "option_A_native.py").write_text("# buyers\n", encoding="utf-8")
    rp = SCRIPTS / "replace_phrase.py"
    r = run(rp, "--out", out, "--find", "buyers", "--replace", "customers", env=env)
    assert r.returncode == 0 and "DRY RUN" in r.stdout, r.stdout
    assert "buyers" in (deck / "dot-dash-smoke.md").read_text(encoding="utf-8")
    r = run(rp, "--out", out, "--find", "buyers", "--replace", "customers", "--apply", env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    h = (deck / "dot-dash-smoke.html").read_text(encoding="utf-8")
    assert 'class="buyers"' in h and "Buyers and customers" in h, h
    assert (deck / "dot-dash-smoke.md").read_text(encoding="utf-8") == "Revenue grows for customers\n"
    assert "customers" in (out / "slide_01" / "option_A.html").read_text(encoding="utf-8")
    assert Document(str(deck / "dot-dash-smoke.docx")).paragraphs[0].text == "Big customers here"
    assert (out / "slide_01" / "option_A_native.py").read_text(encoding="utf-8") == "# buyers\n"
    assert "--edited-by-parent" in r.stdout, r.stdout
    print("    ok: dry run lists, --apply changes text (across Word runs), not tags or converted scripts")

    print("[9] build_status.py, RULINGS lines, emit_dot_dash")
    r = run(SCRIPTS / "build_status.py", "--out", out, env=env)
    assert r.returncode == 0 and "NEXT:" in r.stdout and "Rulings:" in r.stdout, r.stdout
    assert "Repeated rulings" in r.stdout, r.stdout
    r = run(SCRIPTS / "deck_sources.py", "list", "--out", out, env=env)
    assert "RULINGS:" in r.stdout, r.stdout
    eb = tmp / "emit" / "_session" / "narrative-brief-emit.md"
    eb.parent.mkdir(parents=True)
    eb.write_text(brief_text(tpl, [("Revenue grows fast", ""), ("Risks have answers", "")]),
                  encoding="utf-8")
    r = run(STORY / "emit_dot_dash.py", eb, env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "/compact keep the path to" in r.stdout and "saves tokens" in r.stdout, r.stdout
    assert "RULINGS:" in r.stdout, r.stdout
    assert len(slide_ids.ids_in_brief(eb.read_text(encoding="utf-8"))) == 2
    import _rulings
    hint = "\n".join(_rulings.compact_hint(sd, "picks recorded"))
    assert "page map" in hint and "rulings.md" in hint, hint
    print("    ok: status NEXT; RULINGS printed; emit added ids and suggested /compact")

    print("[10] deck_sources.py")
    ds = SCRIPTS / "deck_sources.py"
    r = run(ds, "add", "--out", out, "--name", "Investor deck v2", "--path", src,
            "--validated", "--approved-in-chat", "it is validated", env=env)
    assert r.returncode == 4, r.stdout
    r = run(ds, "add", "--out", out, "--name", "Investor deck v2", "--path", src, "--slides", "risks",
            "--validated", "--approved-in-chat", "v2 was already validated", env=env)
    assert r.returncode == 0, r.stdout
    r = run(ds, "trail", "--out", out, "--figure", "120%", "--where", "model tab 30_Revenue F12",
            "--slides", "risks", env=env)
    assert r.returncode == 0, r.stdout
    ctx3 = (out / "slide_03" / "_context.md").read_text(encoding="utf-8")
    assert "validated by the user, do not re-check" in ctx3 and "120%: model tab" in ctx3, ctx3
    import deck_sources
    assert deck_sources.is_validated(sd, src)
    import source_ledger
    assert source_ledger._validated_source(out, src)
    print("    ok: validated only with the user's words; pack shows status and trail")

    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
