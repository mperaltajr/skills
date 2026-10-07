#!/usr/bin/env python3
"""Smoke test: arrowheads that run into a box are flagged, before anyone looks.

On cycle and flow diagrams arrows came out ending inside the boxes; only the
final visual check caught it, and the defect was already in the design sketch
(2026-10-06). scripts/arrow_ends.py is the deterministic check. Everything
here is made up (a fictional loop diagram, never client content).

Drawn slides (python-pptx), boxes 200x100 px:
  1. a straight arrow whose head is inside a box: flagged, "inside"
  2. a head 2 px short of the box: flagged
  3. a head 10 px short: clean
  4. a curved freeform arrow (cubic Bezier) ending inside a box: flagged
  5. a plain line (no arrowhead) touching a box, and a panel behind the whole
     diagram: neither counts
  6. a clean four-box loop: no findings; the CLI exits 0, a bad one exits 1
     and --report writes the findings into the translation report
The design (HTML), through translate_html.py:
  7. a fictional loop sketch whose SVG arrows (arrow markers, one curved)
     touch two boxes records two MAJOR_ARROW_END_AT_BOX warnings; the same
     sketch with clearance records none

Run:  py -3 slide-builder/tests/run_arrow_ends_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SCRIPTS))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu  # noqa: E402

import arrow_ends as AE  # noqa: E402

PX = 9525


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * PX), Emu(720 * PX)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _box(slide, name, x, y, w=200, h=100):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x * PX), Emu(y * PX),
                               Emu(w * PX), Emu(h * PX))
    s.name = name
    return s


def _arrow_head(shape, tail=True, head=False):
    ln = shape.line._get_or_add_ln()
    if head:
        etree.SubElement(ln, qn("a:headEnd"), type="triangle")
    if tail:
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle")


def _arrow(slide, name, x1, y1, x2, y2, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1 * PX), Emu(y1 * PX),
                                   Emu(x2 * PX), Emu(y2 * PX))
    c.name = name
    if head:
        _arrow_head(c)
    return c


def _curve(slide, name, p0, c1, c2, p1):
    """A freeform cubic Bezier arrow from p0 to p1 (px), arrowhead at p1."""
    xs = [p[0] for p in (p0, c1, c2, p1)]
    ys = [p[1] for p in (p0, c1, c2, p1)]
    x0, y0 = min(xs), min(ys)
    fb = slide.shapes.build_freeform(Emu(p0[0] * PX), Emu(p0[1] * PX), scale=1.0)
    fb.add_line_segments([(Emu(p1[0] * PX), Emu(p1[1] * PX))], close=False)
    s = fb.convert_to_shape()
    s.name = name
    s.left, s.top = Emu(x0 * PX), Emu(y0 * PX)
    s.width, s.height = Emu((max(xs) - x0) * PX), Emu((max(ys) - y0) * PX)
    path = s._element.find(".//" + qn("a:path"))
    for c in list(path):
        path.remove(c)
    path.set("w", str((max(xs) - x0) * PX))
    path.set("h", str((max(ys) - y0) * PX))
    mv = etree.SubElement(path, qn("a:moveTo"))
    etree.SubElement(mv, qn("a:pt"), x=str((p0[0] - x0) * PX), y=str((p0[1] - y0) * PX))
    cb = etree.SubElement(path, qn("a:cubicBezTo"))
    for p in (c1, c2, p1):
        etree.SubElement(cb, qn("a:pt"), x=str((p[0] - x0) * PX), y=str((p[1] - y0) * PX))
    s.fill.background()
    _arrow_head(s)
    return s


def _hits(slide):
    arrows, boxes = AE.read_slide(slide)
    return AE.find_hits(arrows, boxes)


LOOP_HTML = """<!doctype html><html><head><style>
*{box-sizing:border-box} html,body{margin:0;padding:0}
.slide-canvas{width:1280px;height:720px;position:relative;overflow:hidden;background:#FFFFFF;font-family:Arial}
.node{position:absolute;width:220px;height:90px;background:#DCE6F2;border-radius:8px;
      font-size:16px;display:flex;align-items:center;justify-content:center}
</style></head><body><div class="slide-canvas">
<div data-template-field="title" style="position:absolute;left:60px;top:30px;font-size:28px">A made-up loop</div>
<div class="node" data-shape-id="n1" style="left:530px;top:140px">Step one</div>
<div class="node" data-shape-id="n2" style="left:900px;top:330px">Step two</div>
<div class="node" data-shape-id="n3" style="left:530px;top:520px">Step three</div>
<div class="node" data-shape-id="n4" style="left:160px;top:330px">Step four</div>
<svg style="position:absolute;left:0;top:0" width="1280" height="720">
<defs><marker id="m" markerWidth="10" markerHeight="10" refX="9" refY="5" orient="auto">
<path d="M0,0 L10,5 L0,10 z" fill="#333"/></marker></defs>
<line x1="750" y1="185" x2="{x2}" y2="330" stroke="#333" stroke-width="2" marker-end="url(#m)"/>
<path d="M1010,420 C1010,560 900,565 {px},565" fill="none" stroke="#333" stroke-width="2" marker-end="url(#m)"/>
<line x1="530" y1="565" x2="300" y2="440" stroke="#333" stroke-width="2" marker-end="url(#m)"/>
</svg></div></body></html>"""


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="slidelab_arrow_ends_"))
    try:
        print("[1-3] straight arrows: inside / 2 px short / 10 px short")
        for label, end_x, want in (("inside", 520, True), ("2px", 498, True),
                                   ("10px", 490, False)):
            prs, s = _deck()
            _box(s, "source box", 100, 300)
            _box(s, "target box", 500, 300)
            _arrow(s, f"arrow {label}", 300, 350, end_x, 350)
            hits = _hits(s)
            assert bool(hits) is want, (label, hits)
            if hits:
                assert hits[0]["box"] == "target box" and hits[0]["end"] == "end", hits
                assert hits[0]["inside"] is (label == "inside"), hits
        print("    ok")

        print("[4] a curved freeform arrow ending inside a box")
        prs, s = _deck()
        _box(s, "top", 540, 100)
        _box(s, "right", 900, 300)
        _curve(s, "curved arrow", (740, 150), (900, 150), (1000, 200), (1000, 320))
        hits = _hits(s)
        assert len(hits) == 1 and hits[0]["box"] == "right" and hits[0]["inside"], hits
        print(f"    ok: {AE.describe(hits[0])[:80]}...")

        print("[5] a plain line touching a box, and a panel behind it all, do not count")
        prs, s = _deck()
        _box(s, "panel", 60, 200, 940, 300).fill.solid()   # 30% of the slide: not a backdrop
        _box(s, "a", 100, 300)
        _box(s, "b", 500, 300)
        _arrow(s, "plain line", 300, 350, 500, 350, head=False)
        _arrow(s, "clear arrow", 300, 380, 488, 380)
        assert _hits(s) == [], _hits(s)
        print("    ok")

        print("[6] a clean loop: 0 findings; CLI exit codes and --report")
        prs, s = _deck()
        for n, (x, y) in {"n1": (540, 120), "n2": (900, 310), "n3": (540, 500),
                          "n4": (180, 310)}.items():
            _box(s, n, x, y)
        _arrow(s, "1-2", 745, 170, 990, 300)
        _curve(s, "2-3", (1000, 420), (1000, 540), (900, 550), (752, 550))
        _arrow(s, "3-4", 540, 550, 380, 418)
        _arrow(s, "4-1", 280, 300, 530, 170)
        assert _hits(s) == [], _hits(s)
        good = tmp / "good.pptx"
        prs.save(str(good))
        r = subprocess.run([sys.executable, str(SCRIPTS / "arrow_ends.py"), str(good)],
                           capture_output=True, text=True)
        assert r.returncode == 0, r.stdout + r.stderr
        prs, s = _deck()
        _box(s, "target", 500, 300)
        _arrow(s, "bad", 300, 350, 510, 350)
        bad = tmp / "bad.pptx"
        prs.save(str(bad))
        rep = tmp / "option_B_translation_report.json"
        rep.write_text(json.dumps({"warnings": [{"code": "OTHER", "detail": "kept"}]}),
                       encoding="utf-8")
        r = subprocess.run([sys.executable, str(SCRIPTS / "arrow_ends.py"), str(bad),
                            "--report", str(rep)], capture_output=True, text=True)
        assert r.returncode == 1 and "inside" in r.stdout, r.stdout + r.stderr
        w = json.loads(rep.read_text(encoding="utf-8"))["warnings"]
        assert [x["code"] for x in w] == ["OTHER", AE.CODE], w
        print("    ok")

        print("[7] the design: a loop sketch with arrows touching boxes is flagged")
        import translate_html as T
        from playwright.sync_api import sync_playwright
        from _browser import launch
        results = {}
        for label, x2, px in (("touching", 900, 750), ("clear", 890, 760)):
            html = tmp / f"loop_{label}.html"
            html.write_text(LOOP_HTML.replace("{x2}", str(x2)).replace("{px}", str(px)),
                            encoding="utf-8")
            with sync_playwright() as pw:
                br = launch(pw)
                page = br.new_page(viewport={"width": 1280, "height": 720})
                data = T.extract(page, html, True)
                br.close()
            plan, _f = T.plan_from_extract(data, True)
            results[label] = [w for w in plan.warnings if w["code"] == AE.CODE]
        assert len(results["touching"]) == 2, results["touching"]
        assert results["clear"] == [], results["clear"]
        print(f"    ok: {results['touching'][0]['detail'][:90]}...")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
