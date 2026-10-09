#!/usr/bin/env python3
"""Smoke test: many one-slide decks render as ONE deck and each page goes back
to the slide it came from (owner's decision, 2026-10-09).

  [1] mapping: six fictional one-slide decks, each a different solid color
      and a different label, one of them hidden; every picture has its own
      slide's color and every PDF page its own slide's label, on each
      renderer installed (PowerPoint, LibreOffice)
  [2] one render per template: the six render in ONE render call; a deck on a
      different template (4:3) is rendered in a deck of its own
  [3] a slide that fails to render is reported for that slide only: a
      renderer that refuses any deck holding the bad slide; the deck is
      halved until the bad slide stands alone, the other slides still render
      with the right colors
  [4] the same with real PowerPoint and no LibreOffice: a deck PowerPoint
      refuses to open (incomplete <p:style>) fails alone
  [5] the converter's self-check (translate_alarm.render_full) uses one render
      for a round, and returns the failing slide's reason
  [6] finalize's render_all marks only the failing option as not rendered

Fictional files only; the user's open PowerPoint presentations are recorded
before and after (attaching only) and must be unchanged.
Run:  py -3 slide-builder/tests/run_one_deck_render_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import os
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
QC = SKILL.parent / "slide-qc" / "scripts"
for p in (str(SKILL), str(SKILL / "scripts"), str(QC)):
    sys.path.insert(0, p)

import one_deck  # noqa: E402
import render_slides as RS  # noqa: E402
from _ppt_snapshot import snapshot  # noqa: E402

COLORS = ["C00000", "00A000", "0000C0", "C0C000", "00C0C0", "C000C0"]


def _deck(path: Path, color: str, label: str, hidden=False, four_three=False,
          broken=False) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(10 if four_three else 13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = RGBColor.from_string(color)
    r.line.fill.background()
    tb = s.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    tb.text_frame.text = label
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    if hidden:
        s._element.set("show", "0")
    if broken:
        st = r._element.find(qn("p:style"))
        st.remove(st.find(qn("a:effectRef")))     # PowerPoint refuses this file
    prs.save(str(path))
    return path


def _color(png: Path) -> tuple:
    from PIL import Image
    with Image.open(png) as im:
        return im.convert("RGB").getpixel((im.width - 20, 20))


def _close(rgb, hexc) -> bool:
    want = tuple(int(hexc[i:i + 2], 16) for i in (0, 2, 4))
    return all(abs(a - b) <= 12 for a, b in zip(rgb, want))


def _engines() -> list[str]:
    out = []
    if RS.powerpoint_available():
        out.append("powerpoint")
    if RS.libreoffice_available():
        out.append("libreoffice")
    return out


def main() -> int:
    before = snapshot()
    engines = _engines()
    assert engines, "no renderer on this computer"
    with tempfile.TemporaryDirectory() as tdn:
        td = Path(tdn)
        srcs = [_deck(td / f"s{i}.pptx", c, f"Label {i} {c}", hidden=(i == 2))
                for i, c in enumerate(COLORS)]

        print("[1] each page goes back to its own slide")
        for eng in engines:
            pngs = [td / eng / f"p{i}.png" for i in range(len(srcs))]
            calls = []
            real = RS.render

            def counting(*a, **k):
                calls.append(a[0])
                return real(*a, **k)
            RS.render = counting
            try:
                errs = one_deck.render_pngs(srcs, pngs, 40, renderer=eng)
            finally:
                RS.render = real
            assert not errs, errs
            for i, c in enumerate(COLORS):
                assert _close(_color(pngs[i]), c), (eng, i, _color(pngs[i]), c)
            assert len(calls) == 1, f"{eng}: {len(calls)} renders for one template"
            pdf_calls = []
            real_pdf = RS.to_pdf

            def counting_pdf(*a, **k):
                pdf_calls.append(a[0])
                return real_pdf(*a, **k)
            RS.to_pdf = counting_pdf
            try:
                got = one_deck.pdf_pages(srcs, td / f"pdf_{eng}", renderer=eng)
            finally:
                RS.to_pdf = real_pdf
            assert len(pdf_calls) == 1, f"{eng}: {len(pdf_calls)} PDF renders for one template"
            import pypdfium2 as pdfium
            for i, g in enumerate(got):
                assert isinstance(g, tuple), g
                doc = pdfium.PdfDocument(str(g[0]))
                page = doc[g[1]]
                text = page.get_textpage().get_text_range()
                page.close()
                doc.close()
                assert f"Label {i} {COLORS[i]}" in text, (eng, i, text)
            print(f"    ok ({eng}): 6 slides, one render, colors and labels match, "
                  f"the hidden slide included")

        print("[2] a deck on another template renders in a deck of its own")
        other = _deck(td / "other.pptx", "808080", "Other template", four_three=True)
        mixed = srcs[:3] + [other] + srcs[3:]
        cols = COLORS[:3] + ["808080"] + COLORS[3:]
        assert one_deck.signature(other) != one_deck.signature(srcs[0])
        calls = []
        real = RS.render

        def counting2(*a, **k):
            calls.append(a[0])
            return real(*a, **k)
        RS.render = counting2
        try:
            pngs = [td / "mixed" / f"p{i}.png" for i in range(len(mixed))]
            errs = one_deck.render_pngs(mixed, pngs, 40, renderer=engines[0])
        finally:
            RS.render = real
        assert not errs and len(calls) == 2, (errs, calls)
        for i, c in enumerate(cols):
            assert _close(_color(pngs[i]), c), (i, _color(pngs[i]), c)
        print("    ok: two renders (one per template), every picture right")

        print("[3] a slide that fails is reported for that slide only")
        bad_label = "Label 4 "
        real = RS.render

        def refuses_bad(deck, out, dpi, renderer=None, quiet=False):
            with zipfile.ZipFile(deck) as z:
                if any(bad_label.encode() in z.read(n) for n in z.namelist()
                       if n.startswith("ppt/slides/slide")):
                    raise RuntimeError("simulated: the renderer refused this deck")
            return real(deck, out, dpi, renderer=renderer, quiet=quiet)
        RS.render = refuses_bad
        try:
            pngs = [td / "iso" / f"p{i}.png" for i in range(len(srcs))]
            errs = one_deck.render_pngs(srcs, pngs, 40, renderer=engines[0])
        finally:
            RS.render = real
        assert list(errs) == [4] and "simulated" in errs[4], errs
        for i, c in enumerate(COLORS):
            if i != 4:
                assert _close(_color(pngs[i]), c), (i, _color(pngs[i]), c)
        assert not pngs[4].exists()
        print("    ok: only slide 5 of 6 failed, with its own reason; the rest rendered")

        if RS.powerpoint_available():
            print("[4] real PowerPoint, no LibreOffice: the deck it refuses fails alone")
            broken = _deck(td / "broken.pptx", "404040", "Broken", broken=True)
            set4 = srcs[:2] + [broken] + srcs[2:4]
            cols4 = COLORS[:2] + ["404040"] + COLORS[2:4]
            os.environ["SLIDE_LAB_NO_LIBREOFFICE"] = "1"
            try:
                pngs = [td / "ppt_iso" / f"p{i}.png" for i in range(len(set4))]
                errs = one_deck.render_pngs(set4, pngs, 40, renderer="powerpoint")
            finally:
                os.environ.pop("SLIDE_LAB_NO_LIBREOFFICE", None)
            assert list(errs) == [2], errs
            assert "could not open" in errs[2].lower() or "powerpoint" in errs[2].lower(), errs[2]
            for i, c in enumerate(cols4):
                if i != 2:
                    assert _close(_color(pngs[i]), c), (i, _color(pngs[i]), c)
            print(f"    ok: PowerPoint refused only that slide ({errs[2].splitlines()[-1][:80]})")

        print("[5] the converter's self-check renders a round as one deck")
        import translate_alarm as A
        calls = []
        real_pdf = RS.to_pdf

        def counting_pdf(*a, **k):
            calls.append(a[0])
            return real_pdf(*a, **k)
        RS.to_pdf = counting_pdf
        RS_render = RS.render
        try:
            out = A.render_full(srcs, engine=engines[0])
        finally:
            RS.to_pdf = real_pdf
        assert len(calls) == 1 and all(o is not None for o in out), (calls, out)
        for i, c in enumerate(COLORS):
            img = out[i][0]
            assert _close(tuple(int(v) for v in img[20, 1260]), c), (i, img[20, 1260], c)
        assert not A.LAST_RENDER_ERRORS
        print("    ok: 6 slides, one render, each image its own slide")

        print("[6] finalize marks only the failing option")
        import finalize_deck as fd

        class St:
            def __init__(self, src, png):
                self.themed_pptx_path, self.themed_png_path = src, png
                self.rendered, self.error = None, ""
        sts = [St(s, td / "fin" / f"o{i}.png") for i, s in enumerate(srcs)]
        RS.render = refuses_bad
        try:
            fd.render_all(sts)
        finally:
            RS.render = RS_render
        assert [s.rendered for s in sts] == [True, True, True, True, False, True], \
            [s.rendered for s in sts]
        assert "simulated" in sts[4].error
        print("    ok")

    after = snapshot()
    assert before[1] == after[1], "the user's open presentations changed"
    print(f"  ok: the user's PowerPoint: {len(before[1])} presentation(s) open before, "
          f"the same {len(after[1])} after")
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
