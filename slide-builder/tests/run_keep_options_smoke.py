#!/usr/bin/env python3
"""Smoke test: keep some options of a slide ("B and C, not A"; item 5 of the
2026-10-06 session report).

The pick line names several letters per slide ("Picks: 1BC 2A", canonical
"slide_01=BC;slide_02=A"). That is recorded as the user's choice and compiled
as a labeled all-options deck of exactly those options. Before, the only way
was to move option A's files aside and finalize with --allow-missing, which
logged an override on every run and read like a defect at delivery.

  1. the plain and canonical forms parse; canonical_picks writes lists the
     way the page does
  2. REVIEW.html: "All options in one deck" shows a keep / leave-out switch per
     option; leaving one out copies "Picks: 1BC 2A (check ...)" (Playwright,
     launched through scripts/_browser.py)
  3. record_picks records ["B", "C"] with all_options set; finalize needs no
     --allow-missing and logs no override; compile ships the 3 kept options,
     labeled
  4. "1AC 2-" leaves slide 2 out; a letter not on disk or given twice is refused
  5. an identical override is logged once, with a count

Run:  py -3 slide-builder/tests/run_keep_options_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent / "scripts"))
import _e2e_harness as H  # noqa: E402
import _state  # noqa: E402
import record_picks  # noqa: E402


def _line(out: Path, plain: str, body: str) -> str:
    tok = _state.read_state(out)["review"]["token"]
    return f"Picks: {plain} (check {_state.approval_check(tok, body)})"


def _page_line(out: Path) -> str | None:
    """Drive REVIEW.html: enter all-options mode, leave slide 1's option A out,
    copy the message. None when Playwright is not installed."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None
    import _browser
    with sync_playwright() as pw:
        b = _browser.launch(pw)
        try:
            pg = b.new_page()
            pg.on("dialog", lambda d: d.accept())
            pg.goto((out / "REVIEW.html").resolve().as_uri())
            pg.evaluate("() => { localStorage.clear(); window.__copied = null;"
                        " navigator.clipboard.writeText = t => {"
                        " window.__copied = t; return Promise.resolve(); }; }")
            assert not pg.is_visible('.keep-toggle[data-slide="slide_01"][data-letter="A"]'), \
                "the switch shows outside all-options mode"
            pg.click("#btn-all")
            assert pg.evaluate("() => window.__copied") is None, "copied on the first click"
            assert pg.is_visible('.keep-toggle[data-slide="slide_01"][data-letter="A"]')
            pg.click('.keep-toggle[data-slide="slide_01"][data-letter="A"]')
            label = pg.text_content('.keep-toggle[data-slide="slide_01"][data-letter="A"]')
            assert "Left out" in label, label
            pg.click("#btn-all")
            pg.wait_for_function("() => window.__copied")
            cmd = pg.evaluate("() => window.__copied")
            # leave out B and C too: one option per slide is a normal deck
            pg.evaluate("() => { window.__copied = null; }")
            pg.click('.keep-toggle[data-slide="slide_01"][data-letter="B"]')
            pg.click('.keep-toggle[data-slide="slide_01"][data-letter="C"]')
            pg.click("#btn-all")
            assert pg.evaluate("() => window.__copied") is None, \
                "copied an all-options message with no slide keeping two options"
        finally:
            b.close()
    return cmd


def main() -> int:
    print("[1] the pick forms parse")
    assert record_picks._plain_body("1BC 2A") == "slide_01=BC;slide_02=A"
    assert record_picks._plain_body("2a, 1cb") == "slide_01=BC;slide_02=A"
    assert record_picks._plain_body("1B 2-") == "slide_01=B;slide_02=-"
    assert record_picks._plain_body("1BX").startswith("BAD:")
    assert _state.canonical_picks({"slide_02": ["A"], "slide_01": ["C", "B"]}) == \
        "slide_01=BC;slide_02=A"
    assert _state.canonical_picks({"slide_01": "A"}) == "slide_01=A"
    print("    ok")

    tmp, out = H.new_build(2)
    try:
        for L in "ABC":
            H.write_option(out, 1, L)
        H.write_option(out, 2, "A")
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        assert H.run("build_review.py", "--out", out).returncode == 0
        body = "slide_01=BC;slide_02=A"
        expected = _line(out, "1BC 2A", body)

        print("[2] REVIEW.html: the keep / leave-out switch writes the multi-letter line")
        cmd = _page_line(out)
        if cmd is None:
            print("    skipped: playwright not installed")
            line = expected
        else:
            m = re.search(r"^(Picks: .+\(check [0-9a-f]{8}\))$", cmd, re.M)
            assert m, cmd
            line = m.group(1)
            assert line == expected, (line, expected)
            assert "the options I kept" in cmd and "py -3" not in cmd, cmd
            print(f"    ok: {line.split('(')[0].strip()}")

        print("[3] recorded as the user's choice; finalize and compile need no override")
        r = H.run("record_picks.py", "--out", out, "--approved", line)
        assert r.returncode == 0, r.stdout
        assert "keep several options" in r.stdout, r.stdout
        st = _state.read_state(out)
        assert st["review"]["picks"] == {"slide_01": ["B", "C"], "slide_02": ["A"]}, st["review"]
        assert st["review"]["all_options"] is True
        r = H.finalize(out)
        assert r.returncode == 0, r.stdout[-1500:] + r.stderr[-800:]
        assert not _state.read_state(out).get("overrides"), _state.read_state(out).get("overrides")
        assert H.run("build_review.py", "--out", out, "--final").returncode == 0
        final_page = (out / "FINAL-CHECK.html").read_text(encoding="utf-8")
        assert "the 3 options you kept" in final_page
        assert "--all-variations --badge" in final_page or "Deck: all options, labeled" in final_page
        tok = _state.read_state(out)["final_check"]["token"]
        r = H.run("compile_picks.py", "--out", out, "--final-token", tok)
        assert r.returncode == 5 and "--all-variations" in r.stdout, r.stdout
        r = H.run("compile_picks.py", "--out", out, "--final-token", tok,
                  "--all-variations", "--badge")
        assert r.returncode == 0, r.stdout[-1500:]
        from pptx import Presentation
        deck = Presentation(str(out / "final_deck_all_variations.pptx"))
        badges = [sh.text_frame.text for s in deck.slides for sh in s.shapes
                  if sh.name == "chrome-option-badge"]
        assert len(deck.slides) == 3 and badges == ["Option B", "Option C", "Option A"], badges
        assert _state.read_state(out)["stages"]["compile"]["options"] == \
            ["slide_01/B", "slide_01/C", "slide_02/A"]
        assert not _state.read_state(out).get("overrides")
        print("    ok: 3 labeled slides (1B, 1C, 2A), no --allow-missing, no override")

        print("[4] leave a slide out; a letter not on disk or twice is refused")
        assert H.run("build_review.py", "--out", out).returncode == 0
        r = H.run("record_picks.py", "--out", out, "--approved",
                  _line(out, "1AC 2-", "slide_01=AC;slide_02=-"))
        assert r.returncode == 0, r.stdout
        assert _state.read_state(out)["review"]["picks"] == {"slide_01": ["A", "C"],
                                                             "slide_02": "-"}
        r = H.run("record_picks.py", "--out", out, "--approved",
                  _line(out, "1BD 2A", "slide_01=BD;slide_02=A"))
        assert r.returncode == 5 and "no option D" in r.stdout, r.stdout
        tok = _state.read_state(out)["review"]["token"]
        r = H.run("record_picks.py", "--out", out, "--approved",
                  f"PICKS slide_01=BB;slide_02=A CHECK "
                  f"{_state.approval_check(tok, 'slide_01=BB;slide_02=A')}")
        assert r.returncode == 5 and "twice" in r.stdout, r.stdout
        print("    ok")

        print("[5] an identical override is logged once, with a count")
        _state.record_override(out, "allow_missing", "finalize ran with 1 option(s) missing")
        _state.record_override(out, "allow_missing", "finalize ran with 1 option(s) missing")
        _state.record_override(out, "allow_missing", "finalize ran with 2 option(s) missing")
        ov = _state.read_state(out)["overrides"]
        assert len(ov) == 2 and ov[0]["count"] == 2 and "count" not in ov[1], ov
        print("    ok")
    finally:
        H.cleanup(tmp)
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
