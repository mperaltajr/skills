#!/usr/bin/env python3
"""Smoke test: the fixes from the 2026-10-01 "agents are slow" session report.

That session lost two full rounds of pages and ran 9-20 minutes a page:
  1. a 960x540 template registered and passed its self-test; every page was
     measured for 1280x720 and the round was thrown away
     -> registration refuses it (exit 4) and rescale_template.py resizes a
        16:9 template cleanly (positions, sizes and text all scale)
  2. the body area started under the title, ignoring a subtitle below it
     -> body_top is below whichever sits lower, plus 12px
  3. every revision round built three options per page after the user had
     already said what to change
     -> a --slide rebuild builds options_per_slide_revision (1); --options N
        overrides; a full build still uses options_per_slide
  4. renders written straight into OneDrive timed out and were retried
     -> render_html.py screenshots to a local temp file and copies it in
  5. designers had no picture of the client's deck and came back generic
     -> style references (add_style_refs.py) are listed in every _context.md

Run:  py -3 slide-builder/tests/run_speed_fixes_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SKILL / "scripts"))

import _e2e_harness as H  # noqa: E402
import _paths as _p  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

import rescale_template as R  # noqa: E402


def _small_template(path: Path) -> tuple[int, int, int]:
    """A 960x540 (10 x 5.625 in) deck. Returns the master title box's left and
    width and the master title text size (hundredths of a point) before scaling."""
    from pptx.oxml.ns import qn
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    master = prs.slide_masters[0]
    title = next(s for s in master.placeholders if int(s.placeholder_format.type) == 1)
    sz = master.element.find(".//" + qn("p:titleStyle") + "//" + qn("a:defRPr"))
    prs.save(str(path))
    return int(title.left), int(title.width), int(sz.get("sz"))


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="slidelab_speed_"))
    try:
        print("[1] a 960x540 template is refused at registration and rescales cleanly")
        small = tmp / "Small.pptx"
        left0, width0, sz0 = _small_template(small)
        r = H.run("register_template.py", "propose", small)
        assert r.returncode == 4, f"expected exit 4, got {r.returncode}: {r.stdout[-600:]}"
        assert "rescale_template.py" in r.stdout and "1280 x 720" in r.stdout, r.stdout[-600:]
        out = R.rescale(small)
        assert out and out.exists(), out
        big = Presentation(str(out))
        assert (int(big.slide_width), int(big.slide_height)) == (R.TARGET_W, R.TARGET_H)
        from pptx.oxml.ns import qn
        title = next(s for s in big.slide_masters[0].placeholders
                     if int(s.placeholder_format.type) == 1)
        f = R.TARGET_W / 9144000
        assert abs(int(title.left) - left0 * f) < 2 and abs(int(title.width) - width0 * f) < 2
        sz = int(big.slide_masters[0].element.find(
            ".//" + qn("p:titleStyle") + "//" + qn("a:defRPr")).get("sz"))
        assert abs(sz - sz0 * f) <= 1, (sz0, sz)
        assert R.rescale(out) is None, "an already-sized template should be left alone"
        sq = tmp / "Square.pptx"
        prs = Presentation()                            # python-pptx default is 4:3
        prs.save(str(sq))
        try:
            R.rescale(sq)
            raise AssertionError("a 4:3 template must be refused, not stretched")
        except ValueError:
            pass
        print("    ok: exit 4 with the fix; 960 -> 1280 scales the master's boxes and title size; 4:3 refused")

        print("[2] the body area starts below the subtitle when it sits lower")
        import register_template as RT
        prs = Presentation(str(H.TEMPLATE))
        hits = 0
        for lay in prs.slide_layouts:
            phs = {int(p.placeholder_format.type): p for p in lay.placeholders
                   if p.placeholder_format is not None}
            if 4 in phs and (1 in phs or 3 in phs):
                sub = phs[4]
                try:
                    sub_bottom = RT._emu_to_px(int(sub.top) + int(sub.height))
                except Exception:
                    continue
                fields = RT._extract_body_zone_for_canonical(lay)
                tit = phs.get(1) or phs.get(3)
                tit_bottom = RT._emu_to_px(int(tit.top) + int(tit.height))
                if sub_bottom > tit_bottom:
                    assert fields["body_top_y_px"] == int(sub_bottom + 12), (lay.name, fields)
                    hits += 1
        assert hits, "fixture has no layout with a subtitle below its title"
        print(f"    ok: {hits} layout(s) start the body 12px under the subtitle")

        print("[3] a rebuild builds one option; --options overrides; a full build keeps the setting")
        btmp, bout = H.new_build(2)
        try:
            base_env = {k: v for k, v in H.env().items() if k != "SLIDE_LAB_OPTIONS_PER_SLIDE"}

            def rebuild(*extra):
                return subprocess.run(
                    [sys.executable, str(H.SCRIPTS / "build_deck.py"), "--template", str(H.TEMPLATE),
                     "--out", str(bout), "--slide", "1", "--confirm-template", *extra],
                    capture_output=True, text=True, env=base_env, encoding="utf-8", errors="replace")

            def letters():
                meta = json.loads((bout / "_meta.json").read_text(encoding="utf-8"))
                return next(s for s in meta["slides"] if s["n"] == 1)["options"]

            r = rebuild("--options", "2")
            assert r.returncode == 0, r.stdout[-800:] + r.stderr[-800:]
            assert letters() == ["A", "B"], letters()
            r = rebuild()
            assert r.returncode == 0, r.stdout[-800:] + r.stderr[-800:]
            assert letters() == ["A"] * _p.options_per_slide_revision() or \
                letters() == list("ABC")[:_p.options_per_slide_revision()], letters()
            prompt = (bout / "slide_01" / "_prompt.md").read_text(encoding="utf-8")
            n_rev = _p.options_per_slide_revision()
            assert f"Produce the **{n_rev} option(s)**" in prompt, "the prompt must ask for the revision count"
            r = rebuild("--options", "5")
            assert r.returncode == 1 and "--options must be" in r.stderr, r.stderr[-400:]
        finally:
            H.cleanup(btmp)
        print("    ok: --options 2 -> A,B; plain rebuild -> the revision count; --options 5 refused")

        print("[4] render_html.py renders through a local temp file")
        html = tmp / "page.html"
        html.write_text('<html><body style="margin:0"><div style="width:1280px;height:720px;'
                        'background:#1f4e79"></div></body></html>', encoding="utf-8")
        png = tmp / "deep" / "page.png"
        r = subprocess.run([sys.executable, str(H.SCRIPTS / "render_html.py"), str(html), str(png)],
                           capture_output=True, text=True)
        assert r.returncode == 0 and png.exists(), r.stderr[-400:]
        import render_html
        assert "tempfile" in Path(render_html.__file__).read_text(encoding="utf-8")
        print("    ok: rendered, written via a temp file")

        print("[5] style references reach the designer's context")
        import build_deck as B
        tpl = tmp / "Client.pptx"
        shutil.copy2(H.TEMPLATE, tpl)
        none_block = B._style_refs_block(tpl)
        assert "No client style references" in none_block and "add_style_refs.py" in none_block
        refs = _p.style_refs_dir(tpl)
        refs.mkdir(parents=True)
        shutil.copy2(png, refs / "client_p04.png")
        block = B._style_refs_block(tpl)
        assert str(refs / "client_p04.png") in block and "Do NOT copy" in block, block
        print("    ok: listed with match-the-look, don't-copy instructions; a hint when absent")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("SMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
