#!/usr/bin/env python3
"""Smoke test: the template's own sample text never ships, and QC catches it if it does.

A layout's footer slot held bracketed sample text ("<Customize with ...>").
When the design supplied no footer text, finalize cloned the footer with that
line and it appeared on every slide; only a human eye on the final check
caught it, the automated check could not see it (2026-10-08).

On a COPY of the bundled test template whose footer slot carries made-up
sample text:
  - finishing a slide with no footer text leaves no sample line on it
  - finishing a slide with a real source line keeps the real text
  - check_pptx_hygiene flags a deck that still shows such a line as Critical
    (angle-bracket lines and "Click to edit" prompts), and leaves a sentence
    with a "<" inside it alone

Run:  py -3 slide-builder/tests/run_sample_text_smoke.py
Prints "SMOKE PASSED." on success; raises AssertionError otherwise.
"""
from __future__ import annotations

import contextlib
import io
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(SKILL / "scripts"))
sys.path.insert(0, str(SKILL.parent / "slide-qc" / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import finalize_deck as fd  # noqa: E402
import _paths as _p  # noqa: E402
from _chrome_schema import load_chrome_yml  # noqa: E402
from twins.composer import clone_missing_chrome_placeholders  # noqa: E402
import check_pptx_hygiene as HY  # noqa: E402

FIXTURE = HERE / "fixtures" / "layout_diverse_template.pptx"
LAYOUT = "body_canonical_light"
SAMPLE = "<Customize with the project name>"
REAL = "Source: made-up figures for the test"


def _with_sample_footer(src: Path, dst: Path) -> None:
    prs = Presentation(str(src))
    lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
    n = 0
    for ph in lay.placeholders:
        if int(ph.placeholder_format.type) == 15:
            ph.text_frame.text = SAMPLE
            n += 1
    assert n == 1, f"expected one footer slot on {LAYOUT}, found {n}"
    prs.save(str(dst))


def _finish(template: Path, footer: str | None):
    prs = Presentation(str(template))
    lay = next(l for l in prs.slide_layouts if l.name == LAYOUT)
    chrome = load_chrome_yml(_p.chrome_yml(FIXTURE)).layouts[LAYOUT]
    slide = prs.slides.add_slide(lay)
    clone_missing_chrome_placeholders(slide, lay)
    src = Presentation().slides.add_slide(Presentation().slide_layouts[6])
    fields = {"title": "A test title"}
    if footer:
        fields["footer"] = footer
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        fd._apply_body_canonical_finishing(slide, prs, chrome, src, 2,
                                           template_fields_override=fields)
    return prs, slide, buf.getvalue()


def _all_text(slide) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes
                     if getattr(sh, "has_text_frame", False))


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="slidelab_sample_text_"))
    try:
        tpl = tmp / "with_sample.pptx"
        _with_sample_footer(FIXTURE, tpl)

        print("[1] no footer text supplied: the sample line is cleared")
        prs, slide, out = _finish(tpl, None)
        assert SAMPLE not in _all_text(slide), _all_text(slide)
        assert "sample text" in out, out
        print("    ok: " + out.strip().splitlines()[-1].strip())

        print("[2] a real source line is kept")
        _prs, slide2, _out = _finish(tpl, REAL)
        txt = _all_text(slide2)
        assert REAL in txt and SAMPLE not in txt, txt
        print("    ok")

        print("[3] the hygiene check flags sample text left on a slide")
        bad = Presentation()
        s = bad.slides.add_slide(bad.slide_layouts[6])
        for i, t in enumerate((SAMPLE, "Click to edit Master text styles",
                               "Costs fell <5% of plan in the quarter")):
            tb = s.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(6), Inches(0.6))
            tb.text_frame.text = t
        good = Presentation()
        g = good.slides.add_slide(good.slide_layouts[6])
        g.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(0.6)).text_frame.text = \
            "Costs fell <5% of plan in the quarter"
        bad_p, good_p = tmp / "bad.pptx", tmp / "good.pptx"
        bad.save(str(bad_p))
        good.save(str(good_p))
        hits = [v for v in HY.run_all_checks(bad_p)["violations"]
                if v["category"] == "template-sample-text"]
        assert len(hits) == 1 and hits[0]["severity"] == "Critical", hits
        assert SAMPLE[:20] in hits[0]["issue"] and "Click to edit" in hits[0]["issue"], hits
        assert "<5%" not in hits[0]["issue"], hits
        clean = [v for v in HY.run_all_checks(good_p)["violations"]
                 if v["category"] == "template-sample-text"]
        assert not clean, clean
        print("    ok: Critical on the sample lines; a sentence with '<' is left alone")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nSMOKE PASSED.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
