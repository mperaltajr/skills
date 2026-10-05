#!/usr/bin/env python3
"""source_check.py — what is on the finished slides that is not in the brief.

Lists, per picked slide:
  - numbers that do not appear anywhere in the brief (page numbers and
    small list numerals are ignored);
  - ALL-CAPS labels whose words do not appear in the brief (designers have
    invented eyebrows such as "DECISION TODAY");
and, once for the deck, figures written two ways ($1.2B on one slide,
$1,200M on another).

FINAL-CHECK.html shows the result at the top so the person approving the
deck sees exactly what to verify. It never blocks: a derived number (a total,
a ratio) can be right without appearing in the brief. The point is that
nobody has to hunt for it.

  py -3 scripts/source_check.py --out <out_dir>     (prints the same list)
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

NUM_RE = re.compile(r"\$?\d[\d,]*(?:\.\d+)?\s?(?:%|[KMB]\b|bn\b|m\b|pp\b|x\b)?", re.I)
CAPS_RE = re.compile(r"\b[A-Z][A-Z&'\-]{1,}(?:\s+[A-Z][A-Z&'\-]{1,})+\b")
MONEY_RE = re.compile(r"\$(\d[\d,]*(?:\.\d+)?)\s?([KMB])\b", re.I)


def _norm_num(s: str) -> str:
    return re.sub(r"[\s$,]", "", s).lower().rstrip(".")


def _brief_text(out_dir: Path) -> str:
    try:
        meta = json.loads((out_dir / "_meta.json").read_text(encoding="utf-8"))
        bp = (meta.get("deck") or {}).get("brief") or meta.get("brief")
        if bp and Path(bp).exists():
            return Path(bp).read_text(encoding="utf-8")
    except Exception:
        pass
    hits = sorted((out_dir / "_session").glob("narrative-brief-*.md"))
    return hits[0].read_text(encoding="utf-8") if hits else ""


def _slide_text(pptx: Path) -> str:
    from pptx import Presentation
    out = []

    def walk(shapes):
        for sh in shapes:
            if getattr(sh, "shape_type", None) == 6:  # group
                walk(sh.shapes)
            elif getattr(sh, "has_text_frame", False):
                out.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for c in row.cells:
                        out.append(c.text)
    prs = Presentation(str(pptx))
    for s in prs.slides:
        walk(s.shapes)
    return "\n".join(out)


def _money_values(text: str) -> dict[float, set[str]]:
    scale = {"k": 1e3, "m": 1e6, "b": 1e9}
    vals: dict[float, set[str]] = {}
    for m in MONEY_RE.finditer(text):
        v = round(float(m.group(1).replace(",", "")) * scale[m.group(2).lower()], 0)
        vals.setdefault(v, set()).add(m.group(2).upper())
    return vals


def check(out_dir: Path, ships: list[tuple[str, Path]]) -> dict:
    brief = _brief_text(out_dir)
    brief_nums = {_norm_num(m.group(0)) for m in NUM_RE.finditer(brief)}
    brief_plain = {re.sub(r"[^\d.]", "", n) for n in brief_nums}
    brief_lower = brief.lower()
    per_slide: dict[str, dict] = {}
    deck_money: dict[float, set[str]] = {}
    for key, pptx in ships:
        try:
            text = _slide_text(pptx)
        except Exception:
            continue
        nums, labels = [], []
        for m in NUM_RE.finditer(text):
            raw = m.group(0).strip()
            n = _norm_num(raw)
            digits = re.sub(r"[^\d.]", "", n)
            if not digits or (re.fullmatch(r"\d{1,2}", n) and int(n) <= 20):
                continue  # page numbers, list numerals
            if n in brief_nums or digits in brief_plain:
                continue
            if raw not in nums:
                nums.append(raw)
        # Chart axis ticks (200, 400, 600 ...) are drawn, not claimed: drop
        # round numbers when a slide carries four or more of them.
        def _round(x):
            d = re.sub(r"[^\d]", "", x)
            return bool(d) and len(d.rstrip("0")) <= 2 and len(d) >= 2 and not re.search(r"[%$KMB]", x, re.I)
        if sum(1 for x in nums if _round(x)) >= 4:
            nums = [x for x in nums if not _round(x)]
        for m in CAPS_RE.finditer(text):
            words = [w for w in re.split(r"[\s&'\-]+", m.group(0).lower()) if len(w) > 2]
            if words and not all(w in brief_lower for w in words):
                if m.group(0) not in labels:
                    labels.append(m.group(0))
        for v, units in _money_values(text).items():
            deck_money.setdefault(v, set()).update(units)
        if nums or labels:
            per_slide[key] = {"numbers": nums[:12], "labels": labels[:12]}
    mixed = []
    # $1,200M vs $1.2B: the same scale of figure written in two units.
    units_by_scale: dict[str, set[str]] = {}
    for v, units in deck_money.items():
        if v >= 1e9:
            units_by_scale.setdefault("billions", set()).update(units)
    if len(units_by_scale.get("billions", set())) > 1:
        mixed.append("Figures of a billion or more are written both in $B and in $M "
                     "across the deck; pick one.")
    return {"slides": per_slide, "mixed_units": mixed}


def html_block(result: dict) -> str:
    import html as _h
    rows = []
    for key, r in sorted(result["slides"].items()):
        bits = []
        if r["numbers"]:
            bits.append("numbers: " + ", ".join(_h.escape(x) for x in r["numbers"]))
        if r["labels"]:
            bits.append("labels: " + ", ".join(_h.escape(x) for x in r["labels"]))
        rows.append(f"<tr><td>{_h.escape(key.replace('_', ' ').title())}</td><td>{'; '.join(bits)}</td></tr>")
    for m in result["mixed_units"]:
        rows.append(f"<tr><td>Whole deck</td><td>{_h.escape(m)}</td></tr>")
    if not rows:
        return ('<div class="srccheck ok">Every number and label on these slides '
                'appears in the brief.</div>')
    return ('<div class="srccheck"><b>Check these before you build.</b> They are on the '
            'slides but not in the brief. A total or a ratio can be right; an invented '
            'number or label is not.<table>' + "".join(rows) + "</table></div>")


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, type=Path)
    a = ap.parse_args(argv)
    import _state
    import _paths as _p
    picks = (_state.read_state(a.out).get("review") or {}).get("picks") or {}
    ships = [(k, a.out / k / _p.option_pptx_name(v)) for k, v in sorted(picks.items())
             if isinstance(v, str) and v != "-"]
    res = check(a.out, ships)
    for k, r in sorted(res["slides"].items()):
        print(k, r)
    for m in res["mixed_units"]:
        print("deck:", m)
    if not res["slides"] and not res["mixed_units"]:
        print("Every number and label on the picked slides appears in the brief.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
