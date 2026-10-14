#!/usr/bin/env python3
"""source_ledger.py — the figures a COPIED page brings onto a shipping slide.

The problem this solves. When a user hands over a page to be reproduced (a
one-pager, a page of an earlier deck), replicating it faithfully also
replicates its numbers. In a real build that shipped "$300B / 6-12 months / 3x"
when the brief already carried the verified "~$450B / sold out / ~4x". So the
fix for "you dropped my page" silently becomes "you shipped my stale numbers".

Why a ledger and not a differ. A differ would compare each figure against the
brief and pass or fail. It cannot: the brief has no typed numeric field, only
prose, so deciding that "6-12 months" and "sold out" denote the same quantity
is a judgment the machine does not own. A differ that reported "0 conflicts"
would be asserting it. The machine records only what it knows: *this figure
came from the copied page onto the slide about to ship, and a human has or has
not answered for it.*

How it works since 2026-10-13 (owner's decision, the lighter check, option 2):

  * One file, one SECTION PER PINNED SLIDE ("pages" keyed slide_NN): the
    source page and slide, when it was made, the fingerprint of the picked,
    finished design it was made from, its rows and its unreadable surfaces.
    A build replaces only its own section. Before, every build overwrote the
    whole file, so with 15 pinned slides only the last one's rows existed and
    the gates read one count: a silent false pass.
  * Rows only for figures that are on BOTH the copied page and the picked
    design's text (a figure the redesign dropped is not asked about). Slide
    numbers, Source/Note lines, footers and dates are skipped. Figures that
    also appear in the brief word for word are COUNTED ("12 match the brief"),
    not listed, and never resolved automatically.
  * finalize runs it for every pinned slide once the user has picked, and
    FINAL-CHECK.html shows all open rows in one table (slide, figure on the
    copied page, closest brief line, keep or replace). The user's approval of
    that final check (Build it, or "build it" in chat with it open) answers
    every row it showed: compile records them as kept on the final check.
  * compile and check_done check EVERY pinned slide that ships: no section,
    a section made from an older design than the pick, or an open row stops
    them, naming the slides and the one command.
  * Pins are never dropped automatically, and a redesign that keeps none of
    the copied figures simply has no rows.

A row can still be answered in the file by hand (`resolution`):
    bind_from_brief   the value comes from the brief, not the copied page
    keep_source       the human states this value is still correct
    replace_with      the human supplies a value (`replacement`)
`null` is not a resolution.

Run:
  py -3 scripts/source_ledger.py build  --out <out> --all-picked
      (every pinned slide that ships, from _meta.json and the recorded picks;
       a pin that cannot be read or found is an error. finalize runs this.)
  py -3 scripts/source_ledger.py build  --out <out> --deck <page> --slide N [--target-slide T]
      (one slide by hand: the page is a .pptx, a PDF with a text layer, or a
       scanned PDF / picture, which is checked by eye only)
  py -3 scripts/source_ledger.py status --out <out>
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _extract  # noqa: E402
import _paths as _p  # noqa: E402
import _state  # noqa: E402

LEDGER_NAME = "source_ledger.json"
SCHEMA = 2
RESOLUTIONS = ("bind_from_brief", "keep_source", "replace_with")
VISUAL_SUFFIXES = frozenset({".pdf", ".png", ".jpg", ".jpeg", ".gif", ".bmp",
                             ".tif", ".tiff", ".webp"})
PPTX_SUFFIXES = frozenset({".pptx", ".pptm", ".potx"})
SCRIPTS = Path(__file__).resolve().parent


class PinError(Exception):
    """A pin that cannot be read, or names a page that cannot be found."""


def _now() -> str:
    return _dt.datetime.now().isoformat(timespec="seconds")


def ledger_path(out_dir: Path) -> Path:
    return Path(out_dir) / LEDGER_NAME


# ---------------------------------------------------------------- the file
def _old_to_pages(old: dict) -> dict:
    """The one-page format (before 2026-10-13) read as one section. It has no
    design fingerprint, so the gates treat it as older than the pick."""
    rows = old.get("rows") or []
    n = None
    for r in rows + (old.get("unreachable") or []):
        if isinstance(r.get("slide"), int):
            n = r["slide"]
            break
    if n is None:
        n = int(old.get("source_slide") or 1)
    sec = {k: v for k, v in old.items() if k not in ("_how_to",)}
    sec.setdefault("source_kind", "pptx")
    sec.update({"slide": n, "fingerprint": "", "old_format": True,
                "brief_matches": 0})
    for i, r in enumerate(rows):
        r.setdefault("id", f"{_p.slide_key(n)}:old{i}")
        r.setdefault("figure", r.get("source_text", ""))
    return {_p.slide_key(n): sec}


def load_ledger(out_dir: Path) -> dict:
    """{"schema": 2, "pages": {slide_NN: section}}; {} when there is none."""
    p = ledger_path(out_dir)
    if not p.exists():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if isinstance(data, dict) and isinstance(data.get("pages"), dict):
        return data
    if isinstance(data, dict) and ("rows" in data or "source_page" in data):
        return {"schema": SCHEMA, "pages": _old_to_pages(data)}
    return {}


_load = load_ledger          # older name


def _open(r: dict) -> bool:
    if r.get("resolution") not in RESOLUTIONS:
        return True
    return (r.get("resolution") == "replace_with"
            and not str(r.get("replacement") or "").strip())


def _all_rows(ledger: dict) -> list:
    if "pages" in ledger:
        return [r for sec in ledger["pages"].values() for r in (sec.get("rows") or [])]
    return list(ledger.get("rows") or [])


def tally(ledger: dict) -> tuple[int, int, int]:
    """(unresolved, keep_source, total) over every section's rows."""
    rows = _all_rows(ledger)
    return (sum(1 for r in rows if _open(r)),
            sum(1 for r in rows if r.get("resolution") == "keep_source"),
            len(rows))


def _slide_counts(sec: dict) -> dict:
    rows = sec.get("rows") or []
    return {"rows": len(rows), "unresolved": sum(1 for r in rows if _open(r)),
            "kept": sum(1 for r in rows if r.get("resolution") == "keep_source"),
            "brief_matches": int(sec.get("brief_matches") or 0),
            "unreachable": len(sec.get("unreachable") or []),
            "visual_only": sec.get("source_kind") == "visual",
            "source": f"{Path(str(sec.get('source_page', ''))).name} slide "
                      f"{sec.get('source_slide')}",
            "at": sec.get("generated_at", "")}


def save_ledger(out_dir: Path, ledger: dict) -> None:
    """Write the file, the brief_qc.json banner/table data, and the per-slide
    counts in the build state."""
    ledger["schema"] = SCHEMA
    ledger.setdefault("_how_to", (
        "One section per pinned slide. Rows are figures that came from the copied "
        "page onto the picked design and are not in the brief word for word. The user "
        "answers them on FINAL-CHECK.html (keep, or replace through a fix); compile "
        "records the final check's approval on each row. A row can also be answered "
        "here: resolution bind_from_brief, keep_source, or replace_with plus "
        "replacement. null is not a resolution."))
    ledger_path(out_dir).write_text(json.dumps(ledger, indent=2, ensure_ascii=False),
                                    encoding="utf-8")
    _write_brief_qc(out_dir, ledger)
    by = {sid: _slide_counts(sec) for sid, sec in sorted(ledger["pages"].items())}
    unresolved, keeps, _ = tally(ledger)
    _state.record_source_ledger(
        out_dir, unresolved, keeps, sum(c["unreachable"] for c in by.values()),
        visual_only=bool(by) and all(c["visual_only"] for c in by.values()), by_slide=by)


def _write_brief_qc(out_dir: Path, ledger: dict) -> None:
    """The ledger in the brief_qc.json shape REVIEW.html reads (summary,
    blocking, warnings), plus the rows FINAL-CHECK.html shows as a table.
    Deliberately never a green "0 conflicts": silence must not read as a pass."""
    pages = ledger.get("pages") or {}
    blocking, warnings, table, slides = [], [], [], []
    for sid, sec in sorted(pages.items()):
        n = int(sec.get("slide") or sid.split("_")[-1])
        src = f"{Path(str(sec.get('source_page', ''))).name} slide {sec.get('source_slide')}"
        c = _slide_counts(sec)
        slides.append({"slide": n, "source": src, **c})
        for r in sec.get("rows") or []:
            if _open(r):
                blocking.append(
                    f"slide {n}: {r.get('figure')!r} came from the copied page ({src}) "
                    "and is not answered yet: keep or replace it on FINAL-CHECK.html")
                table.append({"id": r.get("id"), "slide": n, "figure": r.get("figure"),
                              "source_text": r.get("source_text", ""), "source": src,
                              "brief_value": r.get("brief_value", "")})
        for u in sec.get("unreachable") or []:
            warnings.append(f"slide {n}, {u.get('addr')}: {u.get('why')} — NOT checked; "
                            "the vision pass is responsible for this surface")
    unresolved, keeps, total = tally(ledger)
    visual = [s for s in slides if s["visual_only"]]
    summary = (f"Copied pages: {len(slides)} slide(s). {total} figure(s) came from the "
               f"copied page onto the picked design, {total - unresolved} answered "
               f"({keeps} kept), {unresolved} still to answer on the final check; "
               f"{sum(s['brief_matches'] for s in slides)} more match the brief (counted, "
               f"not listed); {sum(s['unreachable'] for s in slides)} surface(s) unreadable "
               "and referred to vision QC.")
    if visual:
        summary += (f" {len(visual)} copied page(s) are a PDF or picture with nothing "
                    "machine-readable (slide " + ", ".join(str(s["slide"]) for s in visual)
                    + "): checked by eye only; take every figure from the brief or confirm "
                    "it with the user.")
    _p.brief_qc_json(out_dir).write_text(json.dumps(
        {"summary": summary, "blocking": blocking, "warnings": warnings,
         "copied_page_rows": table, "copied_page_slides": slides}, indent=2,
        ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------- figures
_MONTHS = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_DATE_RES = [
    re.compile(r"\b\d{4}-\d{1,2}-\d{1,2}\b"),
    re.compile(r"\b\d{1,2}/\d{1,2}(?:/\d{2,4})?\b"),
    re.compile(rf"\b{_MONTHS}\s+\d{{1,2}}(?:st|nd|rd|th)?(?:,?\s+\d{{4}})?\b", re.I),
    re.compile(rf"\b\d{{1,2}}(?:st|nd|rd|th)?\s+{_MONTHS}(?:\s+\d{{4}})?\b", re.I),
    re.compile(rf"\b{_MONTHS}\s+\d{{4}}\b", re.I),
]
_NUM = r"\d[\d,]*(?:\.\d+)?"
_FIG_RE = re.compile(
    r"(?<![A-Za-z0-9.,])"                 # not glued to a word: Q3, FY26, H1, v2
    r"(?:[~≈]\s?)?(?:[$€£]\s?)?" + _NUM +
    r"(?:\s?(?:-|–|to)\s?(?:[$€£]\s?)?" + _NUM + r")?"
    r"(?:\s?(?:%|percent|x|bn|billion|mn|million|trillion|thousand|b|m|k|t|gw|mw|kw)"
    r"(?![A-Za-z]))?",
    re.I)
_YEAR_RE = re.compile(r"^(?:19|20)\d\d(?:-(?:19|20)?\d\d)?$")
_SKIP_NAME_RE = re.compile(r"slide\s*number|page\s*number|\bfooter|\bdate\b", re.I)
_SKIP_TEXT_RE = re.compile(r"^\s*(?:sources?|notes?|footnotes?)\s*[:.]", re.I)
_UNIT_WORDS = (("billion", "b"), ("bn", "b"), ("million", "m"), ("mn", "m"),
               ("trillion", "t"), ("thousand", "k"), ("percent", "%"))


def norm_figure(tok: str) -> str:
    t = tok.lower().replace("–", "-").replace("≈", "").replace("~", "")
    t = re.sub(r"\s+", "", t)
    t = re.sub(r"(?<=\d)to(?=[$€£\d])", "-", t)
    t = re.sub(r"(?<=\d),(?=\d)", "", t)
    for w, u in _UNIT_WORDS:
        if t.endswith(w):
            t = t[: -len(w)] + u
    t = re.sub(r"(\d)\.0+(?!\d)", r"\1", t)
    return t


def figures(text: str) -> list[tuple[str, str]]:
    """(as written, normalized) for each figure in the text: amounts,
    percents, multiples, ranges and plain numbers. Dates and bare years are
    left out, and so are numbers glued to letters (Q3, FY26, v2)."""
    t = text or ""
    for rx in _DATE_RES:
        t = rx.sub(" ", t)
    out, seen = [], set()
    for m in _FIG_RE.finditer(t):
        raw = m.group(0).strip()
        n = norm_figure(raw)
        if not any(ch.isdigit() for ch in n) or _YEAR_RE.match(n) or n in seen:
            continue
        seen.add(n)
        out.append((raw, n))
    return out


def _skip_surface(s: dict) -> bool:
    """Slide numbers, footers, dates, Source/Note lines."""
    text = (s.get("text") or "").strip()
    if _SKIP_NAME_RE.search(s.get("name") or ""):
        return True
    if _SKIP_TEXT_RE.match(text):
        return True
    return bool(re.fullmatch(r"\d{1,3}", text))


# ---------------------------------------------------------------- the pin
_PIN_RE = re.compile(r"^(?P<path>.+?)[\s,;:]+(?:slide|page|p\.?)\s*#?\s*(?P<n>\d+)\s*\.?$", re.I)


def parse_pin(pin: str, bases) -> tuple[Path, int]:
    """'<path> slide N' (either slash style, quotes or backticks allowed) to
    (an existing file, N). A relative path is looked up next to the brief,
    then in the build folder. A picture needs no slide number."""
    raw = (pin or "").strip().strip("`").strip()
    m = _PIN_RE.match(raw)
    path_s, n = (m.group("path"), int(m.group("n"))) if m else (raw, None)
    path_s = path_s.strip().strip("`\"'").strip()
    if os.sep == "/":
        path_s = path_s.replace("\\", "/")
    cand = Path(path_s)
    tries = [cand] if cand.is_absolute() else [Path(b) / cand for b in bases if b] + [cand]
    found = next((t for t in tries if t.is_file()), None)
    if found is None:
        raise PinError(f"the pinned page {path_s!r} was not found (looked in: "
                       + "; ".join(str(t.parent) for t in tries) + ")")
    if n is None:
        if found.suffix.lower() in VISUAL_SUFFIXES - {".pdf"}:
            n = 1
        else:
            raise PinError(f"the pin {pin!r} does not say which slide or page; write it "
                           "as '<path to the file> slide N' in the brief")
    if n < 1:
        raise PinError(f"the pin {pin!r} names slide {n}; slides count from 1")
    return found.resolve(), n


def _meta(out_dir: Path) -> dict:
    try:
        return json.loads(_p.meta_json(out_dir).read_text(encoding="utf-8"))
    except Exception:
        return {}


def pinned_slides(out_dir: Path, meta: dict | None = None) -> dict:
    """{slide number: pin text} from _meta.json."""
    meta = _meta(out_dir) if meta is None else meta
    out = {}
    for s in meta.get("slides") or []:
        pin = (s.get("pinned_source_page") or "").strip()
        if pin and isinstance(s.get("n"), int):
            out[s["n"]] = pin
    return out


def _bases(out_dir: Path, meta: dict) -> list:
    b = []
    if meta.get("brief"):
        b.append(Path(meta["brief"]).parent)
    b.append(Path(out_dir))
    return b


def shipping_letters(state: dict, n: int) -> list | None:
    """The letters that ship for slide n: [] when left out, None when nothing
    has been picked yet."""
    picks = (state.get("review") or {}).get("picks")
    if not picks:
        return None
    v = picks.get(_p.slide_key(n))
    if v in (None, "", "-", []):
        return []
    return sorted(L for L in (v if isinstance(v, list) else [v]) if L and L != "-")


def design_fingerprint(out_dir: Path, n: int, letters: list) -> str | None:
    """'A:<md5>' of the finished option file(s) that ship; None when one is
    not finished yet. The same digests the final check is bound to."""
    parts = []
    for L in letters:
        f = Path(out_dir) / _p.slide_key(n) / _p.option_pptx_name(L)
        d = _state.file_digest(f) if f.exists() else ""
        if not d:
            return None
        parts.append(f"{L}:{d}")
    return ";".join(parts)


# ---------------------------------------------------------------- reading
def _pdf_pages(path: Path) -> int | None:
    """Page count of a PDF, or None when no PDF reader is installed."""
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(path))
        try:
            return len(doc)
        finally:
            doc.close()
    except ImportError:
        pass
    except Exception:
        return None
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(path)).pages)
    except Exception:
        return None


def _pdf_text_page(path: Path, n: int) -> tuple[list[str], int] | None:
    """(text lines, number of pictures) on page n of a PDF, from its text
    layer. None when no PDF reader is installed or the page can't be read."""
    try:
        import pypdfium2 as pdfium
        import pypdfium2.raw as pdfium_c
    except ImportError:
        pdfium = None
    if pdfium is not None:
        try:
            doc = pdfium.PdfDocument(str(path))
            try:
                page = doc[n - 1]
                tp = page.get_textpage()
                text = tp.get_text_range() or ""
                try:
                    pictures = sum(1 for _ in page.get_objects(
                        filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE], max_depth=3))
                except Exception:
                    pictures = 0
                return [ln.strip() for ln in text.replace("\r", "\n").split("\n")
                        if ln.strip()], pictures
            finally:
                doc.close()
        except Exception:
            pass
    try:
        from pypdf import PdfReader
        page = PdfReader(str(path)).pages[n - 1]
        text = page.extract_text() or ""
        try:
            pictures = len(page.images)
        except Exception:
            pictures = 0
        return [ln.strip() for ln in text.splitlines() if ln.strip()], pictures
    except Exception:
        return None


def read_source(deck: Path, n: int, target: int) -> tuple[str, list, list]:
    """(kind, surfaces, unreachable) of page n of the copied file. kind is
    "pptx", "pdf_text" or "visual" (nothing machine-readable)."""
    deck = Path(deck)
    suffix = deck.suffix.lower()
    if suffix in PPTX_SUFFIXES:
        from pptx import Presentation
        slides = list(Presentation(str(deck)).slides)
        if not (1 <= n <= len(slides)):
            raise PinError(f"slide {n} is out of range ({deck.name} has {len(slides)})")
        surfaces, unreachable = _extract.walk_slide(slides[n - 1], n)
        for u in unreachable:
            u["slide"] = target
        return "pptx", surfaces, unreachable
    if suffix not in VISUAL_SUFFIXES:
        raise PinError(f"{deck.name}: a copied page must be a PowerPoint file, a PDF or a "
                       f"picture ({', '.join(sorted(VISUAL_SUFFIXES))}). Save it as a PDF "
                       "and run this again.")
    is_pdf = suffix == ".pdf"
    why_pdf = ""
    if is_pdf:
        pages = _pdf_pages(deck)
        if pages is not None and not (1 <= n <= pages):
            raise PinError(f"page {n} is out of range (the PDF has {pages} page(s))")
        got = _pdf_text_page(deck, n)
        if got is not None and any(ch.isalnum() for ln in got[0] for ch in ln):
            lines, pictures = got
            surfaces = [{"slide": n, "addr": f"p{n}/line{k + 1}", "kind": "pdf_text",
                         "name": "", "text": ln} for k, ln in enumerate(lines)]
            if surfaces and re.fullmatch(r"\d{1,3}", surfaces[-1]["text"]):
                surfaces[-1]["name"] = "page number"
            unreachable = []
            if pictures:
                unreachable.append({
                    "slide": target, "addr": f"p{n}/pictures", "kind": "picture",
                    "why": (f"{pictures} picture(s) on PDF page {n}: any figure inside "
                            "them is not read by machine; checked by eye only (slide-qc "
                            "vision pass)")})
            return "pdf_text", surfaces, unreachable
        why_pdf = ("no text layer: a scan" if got is not None
                   else "its text could not be read")
    elif n != 1:
        raise PinError(f"slide {n}: a picture has one page; use slide 1")
    what = f"PDF page {n} ({why_pdf})" if is_pdf else "picture"
    return "visual", [], [{
        "slide": target, "addr": f"whole page ({what})", "kind": "pdf" if is_pdf else "picture",
        "why": (f"{what}: its figures are not read by machine; checked by eye only "
                "(slide-qc vision pass), no figure on it was machine-checked")}]


def design_surfaces(out_dir: Path, n: int, letters: list) -> list:
    """Every readable text surface of the finished option file(s) that ship."""
    from pptx import Presentation
    out = []
    for L in letters:
        f = Path(out_dir) / _p.slide_key(n) / _p.option_pptx_name(L)
        prs = Presentation(str(f))
        for slide in list(prs.slides)[:1]:
            s, _ = _extract.walk_slide(slide, n)
            for x in s:
                x["option"] = L
            out.extend(s)
    return out


def _brief_parts(meta: dict, n: int) -> tuple[str, list]:
    """(the whole brief's text without headers and pin lines, this slide's
    lines that carry a figure)."""
    bp = Path(meta.get("brief") or "")
    try:
        text = bp.read_text(encoding="utf-8") if bp.is_file() else ""
    except OSError:
        text = ""
    lines, body, blocks = [], text, []
    try:
        import build_deck
        _, body = build_deck.extract_front_matter(text)   # no seal or dates
        blocks = build_deck.split_slide_blocks(body)
    except Exception:
        pass
    keep = [ln for ln in body.splitlines()
            if not re.match(r"^\s*#+\s*slide\s+\d+", ln, re.I)
            and "pinned source page" not in ln.lower()]
    whole = "\n".join(keep)
    try:
        for sn, _t, block in blocks:
            if sn == n:
                lines = [ln.strip(" -*\t") for ln in block.splitlines()
                         if figures(ln) and "pinned source page" not in ln.lower()]
    except Exception:
        lines = []
    return whole, lines


def _closest(source_text: str, lines: list) -> str:
    """The brief line for this slide that shares the most words with the
    copied text (shown beside the row so the user can compare)."""
    if not lines:
        return ""
    words = set(re.findall(r"[a-z]{3,}", (source_text or "").lower()))
    best = max(lines, key=lambda ln: len(words & set(re.findall(r"[a-z]{3,}", ln.lower()))))
    best = re.sub(r"\*\*([^*]+)\*\*", r"\1", best).strip()
    return best[:160] + ("..." if len(best) > 160 else "")


def _validated_source(out_dir: Path, deck: Path) -> dict | None:
    """The deck_sources.py entry when the user marked this copied file as
    validated, else None. Never raises."""
    try:
        import _rulings
        import deck_sources
        return deck_sources.is_validated(_rulings.session_dir(out=out_dir), deck)
    except Exception:
        return None


# ---------------------------------------------------------------- building
def build_section(out_dir: Path, n: int, deck: Path, src_n: int, letters: list,
                  meta: dict, prior: dict | None = None) -> dict:
    """One pinned slide's section, made from the picked, finished design."""
    fp = design_fingerprint(out_dir, n, letters)
    if fp is None:
        raise PinError(f"slide {n}'s picked design (option {', '.join(letters)}) is not "
                       "finished yet; run finalize_deck.py first")
    kind, surfaces, unreachable = read_source(deck, src_n, n)
    sid = _p.slide_key(n)
    design = [s for s in design_surfaces(out_dir, n, letters) if not _skip_surface(s)]
    on_design: dict = {}
    for s in design:
        for _raw, f in figures(s["text"]):
            on_design.setdefault(f, s["text"])
    whole, brief_lines = _brief_parts(meta, n)
    in_brief = {f for _r, f in figures(whole)}
    vsrc = _validated_source(out_dir, deck)
    prior_rows = {r.get("id"): r for r in ((prior or {}).get("rows") or [])}
    rows, matched, dropped, seen, from_validated = [], [], 0, set(), []
    for s in surfaces:
        if _skip_surface(s):
            continue
        for raw, f in figures(s["text"]):
            if f in seen:
                continue
            seen.add(f)
            if f not in on_design:
                dropped += 1              # the picked design does not carry it
                continue
            if f in in_brief:
                matched.append(raw)       # counted, never listed or auto-resolved
                continue
            if vsrc:
                # The user said this source is validated (deck_sources.py, in
                # their words): counted, not asked about again (2026-10-14).
                from_validated.append(raw)
                continue
            row = {"id": f"{sid}:{f}", "slide": n, "figure": raw, "addr": s["addr"],
                   "kind": s.get("kind", ""), "shape_name": s.get("name", ""),
                   "source_text": s["text"], "design_text": on_design[f],
                   "brief_value": _closest(s["text"], brief_lines),
                   "resolution": None, "replacement": None, "note": ""}
            old = prior_rows.get(row["id"])
            # A human's answer stands while the figure, the copied text and the
            # text it landed in are unchanged.
            if old and not _open(old) and all(old.get(k) == row[k] for k in
                                               ("source_text", "design_text")):
                for k in ("resolution", "replacement", "note", "resolved_at"):
                    if k in old:
                        row[k] = old[k]
            rows.append(row)
    return {"slide": n, "source_page": str(Path(deck).resolve()), "source_slide": src_n,
            "source_kind": kind, "generated_at": _now(),
            "design": {"letters": list(letters), "fingerprint": fp}, "fingerprint": fp,
            "rows": rows, "brief_matches": len(matched), "brief_matched": matched,
            "validated_source": (vsrc or {}).get("name", ""),
            "validated_matches": len(from_validated), "validated_matched": from_validated,
            "not_on_design": dropped, "unreachable": unreachable}


def build_slides(out_dir: Path, only=None, strict: bool = True) -> tuple[list, list]:
    """Build the section of every pinned slide that ships (or of `only`).
    Returns (built slide numbers, [(slide, error)]). A slide left out of the
    deck needs no section."""
    out_dir = Path(out_dir)
    meta = _meta(out_dir)
    state = _state.read_state(out_dir)
    pins = pinned_slides(out_dir, meta)
    ledger = load_ledger(out_dir) or {"schema": SCHEMA, "pages": {}}
    built, errors = [], []
    for n, pin in sorted(pins.items()):
        if only is not None and n not in only:
            continue
        letters = shipping_letters(state, n)
        if letters is None:
            errors.append((n, "nothing has been picked yet; the copied-page check runs on "
                              "the picked design, after the picks"))
            continue
        if not letters:
            continue
        try:
            deck, src_n = parse_pin(pin, _bases(out_dir, meta))
            sid = _p.slide_key(n)
            ledger["pages"][sid] = build_section(out_dir, n, deck, src_n, letters, meta,
                                                 ledger["pages"].get(sid))
            built.append(n)
        except PinError as exc:
            errors.append((n, str(exc)))
        except Exception as exc:                      # a page python-pptx cannot open
            errors.append((n, f"{type(exc).__name__}: {exc}"))
    if built:
        save_ledger(out_dir, ledger)
    return built, errors


def stale_slides(out_dir: Path, state: dict | None = None) -> list:
    """Pinned slides that ship and have no section, or one made from another
    design than the one picked now."""
    out_dir = Path(out_dir)
    state = _state.read_state(out_dir) if state is None else state
    pages = (load_ledger(out_dir) or {}).get("pages") or {}
    out = []
    for n in sorted(pinned_slides(out_dir)):
        letters = shipping_letters(state, n)
        if not letters:
            continue
        sec = pages.get(_p.slide_key(n))
        if not sec or sec.get("fingerprint") != design_fingerprint(out_dir, n, letters):
            out.append(n)
    return out


def auto_after_finalize(out_dir: Path, slide: int | None = None) -> list[str]:
    """finalize's hook: once the user has picked, (re)build the section of
    every pinned slide whose section is missing or older than its pick.
    Returns plain lines to print. Never raises."""
    try:
        state = _state.read_state(out_dir)
        if not (state.get("review") or {}).get("picks"):
            return []
        todo = [n for n in stale_slides(out_dir, state) if slide is None or n == slide]
        if not todo:
            return []
        built, errors = build_slides(out_dir, only=set(todo))
    except Exception as exc:
        return [f"  ERROR: the copied-page check could not run ({type(exc).__name__}: "
                f"{exc}). Run: py -3 \"{SCRIPTS / 'source_ledger.py'}\" build --out "
                f"\"{out_dir}\" --all-picked"]
    lines = []
    if built:
        led = load_ledger(out_dir)
        for n in built:
            c = _slide_counts(led["pages"][_p.slide_key(n)])
            lines.append(f"  copied page, slide {n}: {c['unresolved']} figure(s) from "
                         f"{c['source']} to answer on the final check, "
                         f"{c['brief_matches']} match the brief (counted)"
                         + (", checked by eye only" if c["visual_only"] else ""))
    for n, why in errors:
        lines.append(f"  ERROR: copied page, slide {n}: {why}. Fix the pin in the brief "
                     "(or the file), then run: py -3 "
                     f"\"{SCRIPTS / 'source_ledger.py'}\" build --out \"{out_dir}\" "
                     "--all-picked. Compile refuses until it has run.")
    return lines


# ---------------------------------------------------------------- the gates
def gate_problems(out_dir: Path, state: dict | None = None) -> list[tuple[int, str, str]]:
    """[(slide, what, plain reason)] for every pinned slide that ships and is
    not settled, and every section with an open row. what is "missing",
    "stale" or "open". Empty when nothing is pinned."""
    out_dir = Path(out_dir)
    state = _state.read_state(out_dir) if state is None else state
    meta = _meta(out_dir)
    pins = pinned_slides(out_dir, meta)
    pages = (load_ledger(out_dir) or {}).get("pages") or {}
    if not pins and not pages:
        return []
    probs, seen = [], set()
    for n, pin in sorted(pins.items()):
        letters = shipping_letters(state, n)
        if letters == []:
            continue
        sid = _p.slide_key(n)
        seen.add(sid)
        sec = pages.get(sid)
        if not sec:
            probs.append((n, "missing", "no copied-page check yet"))
            continue
        if letters is not None and sec.get("fingerprint") != design_fingerprint(
                out_dir, n, letters):
            probs.append((n, "stale", "its check was made from an older design than the "
                                      "one picked now"))
            continue
        k = sum(1 for r in sec.get("rows") or [] if _open(r))
        if k:
            probs.append((n, "open", f"{k} figure(s) from the copied page not answered yet"))
    for sid, sec in sorted(pages.items()):
        if sid in seen:
            continue
        k = sum(1 for r in sec.get("rows") or [] if _open(r))
        if k:
            n = int(sec.get("slide") or sid.split("_")[-1])
            probs.append((n, "open", f"{k} figure(s) from the copied page not answered yet"))
    return probs


def gate_message(out_dir: Path, probs: list) -> str:
    """One refusal that names the slides and gives the one command."""
    names = "; ".join(f"slide {n} ({why})" for n, _w, why in probs[:12])
    more = f" and {len(probs) - 12} more" if len(probs) > 12 else ""
    cmd_ledger = (f'py -3 "{SCRIPTS / "source_ledger.py"}" build --out "{out_dir}" '
                  "--all-picked")
    cmd_final = f'py -3 "{SCRIPTS / "build_review.py"}" --out "{out_dir}" --final --open'
    if any(w in ("missing", "stale") for _n, w, _r in probs):
        todo = (f"Run {cmd_ledger}, then {cmd_final}, and let the user answer the "
                "copied-page table (keep or replace each figure).")
    else:
        todo = (f"Run {cmd_final} and let the user answer the copied-page table (keep "
                "or replace each figure); their approval of that final check answers it.")
    return (f"{len(probs)} slide(s) that copy a supplied page are not checked: {names}"
            f"{more}. {todo} Never remove a pin to get past this.")


def resolve_by_final_check(out_dir: Path, state: dict | None = None) -> int:
    """The user approved the final check: every open row its table showed is
    answered "kept" (owner's decision C, 2026-10-13: the final-check approval
    counts). Only rows of sections whose design is still the one shown.
    Returns how many rows were answered."""
    state = _state.read_state(out_dir) if state is None else state
    fc = state.get("final_check") or {}
    shown = fc.get("ledger") or {}
    ids = set(shown.get("rows") or [])
    fps = shown.get("fingerprints") or {}
    ledger = load_ledger(out_dir)
    if not ids or not ledger:
        return 0
    k = 0
    for sid, sec in (ledger.get("pages") or {}).items():
        if fps.get(sid) != sec.get("fingerprint"):
            continue
        for r in sec.get("rows") or []:
            if _open(r) and r.get("id") in ids:
                r.update({"resolution": "keep_source", "resolved_at": _now(),
                          "note": "kept: the user approved FINAL-CHECK.html with this "
                                  f"figure in its table (final check {fc.get('token')})"})
                k += 1
    if k:
        save_ledger(out_dir, ledger)
    return k


def final_check_rows(out_dir: Path) -> tuple[list, list, dict]:
    """(open rows, per-slide summaries, the record to bind to the final
    check) for FINAL-CHECK.html."""
    led = load_ledger(out_dir)
    if not led:
        return [], [], {}
    _write_brief_qc(out_dir, led)
    try:
        qc = json.loads(_p.brief_qc_json(out_dir).read_text(encoding="utf-8"))
    except Exception:
        qc = {}
    rows = qc.get("copied_page_rows") or []
    slides = qc.get("copied_page_slides") or []
    bind = {"rows": sorted(r["id"] for r in rows if r.get("id")),
            "fingerprints": {sid: sec.get("fingerprint", "")
                             for sid, sec in sorted((led.get("pages") or {}).items())}}
    return rows, slides, bind


# ---------------------------------------------------------------- commands
def _print_section(sec: dict) -> None:
    c = _slide_counts(sec)
    n = sec.get("slide")
    if c["visual_only"]:
        print(f"     slide {n}: {c['source']} is a PDF page or picture with nothing "
              "machine-readable: CHECKED BY EYE ONLY (the slide-qc vision pass covers it, "
              "and check_done.py says so at delivery). Take every figure from the brief, "
              "or confirm it with the user.")
        return
    print(f"     slide {n}: {c['rows']} figure(s) from {c['source']} came through onto the "
          f"picked design, {c['unresolved']} to answer; {c['brief_matches']} match the "
          f"brief (counted, not listed); {sec.get('not_on_design', 0)} not on the design.")
    if sec.get("validated_matches"):
        print(f"     slide {n}: {sec['validated_matches']} more come from "
              f"\"{sec.get('validated_source')}\", which the user validated: counted, not "
              "asked about again.")
    if c["unreachable"]:
        print(f"       {c['unreachable']} surface(s) could NOT be read (picture / SmartArt / "
              "embedded object); the slide-qc vision pass is responsible for them.")


def cmd_build(args) -> int:
    out_dir = Path(args.out)
    if not out_dir.exists():
        print(f"[error] out dir not found: {out_dir}", file=sys.stderr)
        return 2
    if args.all_picked:
        pins = pinned_slides(out_dir)
        if not pins:
            print(f"[info] no slide in {out_dir / '_meta.json'} copies a supplied page.")
            return 0
        built, errors = build_slides(out_dir)
        led = load_ledger(out_dir)
        print(f"[ok] {LEDGER_NAME}: {len(built)} of {len(pins)} pinned slide(s) checked")
        for n in built:
            _print_section(led["pages"][_p.slide_key(n)])
        for n, why in errors:
            print(f"[error] slide {n}: {why}", file=sys.stderr)
        if errors:
            print("Fix each pin in the brief (`**Pinned source page:** <path> slide N`) or "
                  "the file, prep that slide again if the brief changed, and run this again. "
                  "Never remove a pin to get past the check.", file=sys.stderr)
            return 2
        print("     The user answers these on FINAL-CHECK.html (build_review.py --final).")
        return 0
    if not args.deck:
        print("[error] give --all-picked, or --deck <page> --slide N", file=sys.stderr)
        return 2
    deck = Path(args.deck)
    if not deck.exists():
        print(f"[error] supplied page not found: {deck}", file=sys.stderr)
        return 2
    target = args.target_slide or args.slide
    state = _state.read_state(out_dir)
    letters = shipping_letters(state, target)
    if letters is None:
        letters = [L for L in _p._ALL_OPTION_LETTERS
                   if (out_dir / _p.slide_key(target) / _p.option_pptx_name(L)).exists()][:1]
        if not letters:
            print(f"[error] slide {target} has no finished design yet. The copied-page "
                  "check runs on the picked, finished design: after the picks, finalize "
                  "runs it for every pinned slide.", file=sys.stderr)
            return 2
        print(f"[info] nothing picked yet: checking slide {target} against option "
              f"{letters[0]} (the faithful copy). The pick replaces it.")
    if not letters:
        print(f"[info] slide {target} is left out of the deck; nothing to check.")
        return 0
    ledger = load_ledger(out_dir) or {"schema": SCHEMA, "pages": {}}
    sid = _p.slide_key(target)
    try:
        ledger["pages"][sid] = build_section(out_dir, target, deck, args.slide, letters,
                                             _meta(out_dir), ledger["pages"].get(sid))
    except PinError as exc:
        print(f"[error] {exc}", file=sys.stderr)
        return 2
    save_ledger(out_dir, ledger)
    print(f"[ok] wrote {ledger_path(out_dir)} (section {sid}; the other slides' sections "
          "are kept)")
    _print_section(ledger["pages"][sid])
    print("     compile_picks.py refuses until every row is answered (on FINAL-CHECK.html).")
    return 0


def cmd_status(args) -> int:
    out_dir = Path(args.out)
    ledger = load_ledger(out_dir)
    if not ledger:
        print(f"[info] no {LEDGER_NAME} in {out_dir}: no copied page checked yet.")
        return 0
    save_ledger(out_dir, ledger)
    unresolved, keeps, total = tally(ledger)
    print(f"[ok] {total} figure(s): {total - unresolved} answered, {unresolved} not "
          f"answered, {keeps} kept from the copied page.")
    for sec in sorted(ledger["pages"].values(), key=lambda s: int(s.get("slide") or 0)):
        _print_section(sec)
        for r in sec.get("rows") or []:
            if _open(r):
                print(f"       NOT ANSWERED {r.get('figure')!r} in {r.get('source_text')!r}")
    probs = gate_problems(out_dir)
    if probs:
        print("     " + gate_message(out_dir, probs))
    return 1 if (unresolved or probs) else 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Check the figures a copied page brings "
                                             "onto the shipping slides.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build", help="Make the section of each pinned slide.")
    b.add_argument("--out", required=True)
    b.add_argument("--all-picked", action="store_true",
                   help="Every pinned slide that ships, from _meta.json and the picks.")
    b.add_argument("--deck", default=None,
                   help="One copied page by hand: a .pptx; a PDF (a scan is checked by "
                        "eye only); or a picture (checked by eye only).")
    b.add_argument("--slide", type=int, default=1, help="Which slide or page of that file.")
    b.add_argument("--target-slide", type=int, default=None,
                   help="Slide number in the deck being built (defaults to --slide).")
    b.set_defaults(func=cmd_build)
    s = sub.add_parser("status", help="Re-count the ledger and record it.")
    s.add_argument("--out", required=True)
    s.set_defaults(func=cmd_status)
    args = ap.parse_args(argv)
    try:
        import _rulings
        _rulings.announce(out=args.out)
    except Exception:
        pass
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
