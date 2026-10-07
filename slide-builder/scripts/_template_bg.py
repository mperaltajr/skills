"""_template_bg.py — what a template layout really looks like behind the body.

Two facts the designers and the translator need and used not to have:

  background   the layout's own background: the layout's <p:bg>, else the
               master's, else PowerPoint's default (bg1). Sketches were drawn
               on a hard-coded white canvas, so on a template whose master is
               a light gray-blue a pale panel was approved on white and then
               vanished in the deck (2026-10-06).
  margins      the left and right edges of the layout's text area (title and
               content placeholders). Designers were given the body's top and
               bottom only, so panels ran past the template's right margin.

Read straight from the template, so a template registered before these were
recorded works without registering it again. register_template.py records
the background per layout in chrome.yml (background_hex); build prep prefers
that and falls back to this lookup.
"""
from __future__ import annotations

import colorsys
from functools import lru_cache
from pathlib import Path

from pptx.oxml.ns import qn

EMU_PER_PX = 9525

# Roles of a color map (bg1 -> lt1 ...) and the scheme slots they name.
_DEFAULT_CLRMAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
_PRESET = {"white": "FFFFFF", "black": "000000"}


def _theme_colors(master) -> dict[str, str]:
    """{scheme slot: hex} from the master's theme (dk1, lt1, accent1, ...)."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    out: dict[str, str] = {}
    try:
        theme_part = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme_part.blob)
    except Exception:
        return out
    cs = root.find(".//" + qn("a:clrScheme"))
    if cs is None:
        return out
    for slot in cs:
        name = etree.QName(slot).localname
        for c in slot:
            tag = etree.QName(c).localname
            if tag == "srgbClr" and c.get("val"):
                out[name] = c.get("val").upper()
            elif tag == "sysClr":
                out[name] = (c.get("lastClr") or ("FFFFFF" if c.get("val") == "window"
                                                  else "000000")).upper()
    return out


def _clr_map(layout) -> dict[str, str]:
    m = dict(_DEFAULT_CLRMAP)
    try:
        cm = layout.slide_master.element.find(qn("p:clrMap"))
        if cm is not None:
            m.update({k: v for k, v in cm.attrib.items()})
        ov = layout.element.find(qn("p:clrMapOvr"))
        oc = ov.find(qn("a:overrideClrMapping")) if ov is not None else None
        if oc is not None:
            m.update({k: v for k, v in oc.attrib.items()})
    except Exception:
        pass
    return m


def _apply_mods(hex_: str, el) -> str:
    """Apply lumMod / lumOff / tint / shade children of a color element."""
    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
    for m in el:
        if not isinstance(m.tag, str):
            continue
        tag = m.tag.split("}")[-1]
        try:
            v = int(m.get("val")) / 100000
        except (TypeError, ValueError):
            continue
        if tag in ("lumMod", "lumOff"):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            l = l * v if tag == "lumMod" else l + v
            r, g, b = colorsys.hls_to_rgb(h, min(1.0, max(0.0, l)), s)
        elif tag == "tint":
            r, g, b = (c + (1 - c) * (1 - v) for c in (r, g, b))
        elif tag == "shade":
            r, g, b = (c * v for c in (r, g, b))
    return "".join(f"{int(round(min(1, max(0, c)) * 255)):02X}" for c in (r, g, b))


def _resolve_color(el, layout) -> str | None:
    """Hex of a DrawingML color element (srgbClr / schemeClr / sysClr / prstClr)."""
    if el is None:
        return None
    tag = el.tag.split("}")[-1]
    base = None
    if tag == "srgbClr":
        base = (el.get("val") or "").upper() or None
    elif tag == "sysClr":
        base = (el.get("lastClr") or "").upper() or None
    elif tag == "prstClr":
        base = _PRESET.get(el.get("val") or "")
    elif tag == "schemeClr":
        val = el.get("val") or ""
        slot = _clr_map(layout).get(val, val)
        base = _theme_colors(layout.slide_master).get(slot)
    if not base or len(base) != 6:
        return None
    return _apply_mods(base, el)


def _first_color(fill_el, layout) -> str | None:
    for c in fill_el.iter():
        if isinstance(c.tag, str) and c.tag.split("}")[-1] in ("srgbClr", "schemeClr", "sysClr", "prstClr"):
            return _resolve_color(c, layout)
    return None


def _bg_of(element, layout) -> dict | None:
    """The background an element (layout or master) sets, or None if it sets none."""
    csld = element.find(qn("p:cSld"))
    bg = csld.find(qn("p:bg")) if csld is not None else None
    if bg is None:
        return None
    pr = bg.find(qn("p:bgPr"))
    if pr is not None:
        for f in pr:
            if not isinstance(f.tag, str):
                continue
            kind = f.tag.split("}")[-1]
            if kind == "solidFill":
                return {"hex": _first_color(f, layout), "kind": "solid"}
            if kind == "gradFill":
                return {"hex": _first_color(f, layout), "kind": "gradient"}
            if kind == "blipFill":
                return {"hex": None, "kind": "picture"}
            if kind == "pattFill":
                return {"hex": _first_color(f, layout), "kind": "pattern"}
            if kind == "noFill":
                return {"hex": "FFFFFF", "kind": "solid"}
        return None
    ref = bg.find(qn("p:bgRef"))
    if ref is not None:
        # A theme background style tinted with this color. Styles 1001 / 1002
        # are solid in nearly every theme; 1003 is usually a gradient or picture.
        try:
            idx = int(ref.get("idx") or 0)
        except ValueError:
            idx = 0
        color = next((c for c in ref if isinstance(c.tag, str) and c.tag.split("}")[-1] in
                      ("srgbClr", "schemeClr", "sysClr", "prstClr")), None)
        return {"hex": _resolve_color(color, layout),
                "kind": "solid" if idx in (0, 1001, 1002) else "theme-style"}
    return None


def layout_background(layout) -> dict:
    """{'hex': 'RRGGBB' or None, 'kind': solid|gradient|picture|pattern|theme-style,
    'from': layout|master|default} for one python-pptx slide layout. `hex` is
    the color to draw a sketch on: exact for a solid background, the first
    color of a gradient (approximate; kind says so), None for a picture."""
    for element, origin in ((layout.element, "layout"), (layout.slide_master.element, "master")):
        try:
            got = _bg_of(element, layout)
        except Exception:
            got = None
        if got is not None:
            got["from"] = origin
            return got
    white = _theme_colors(layout.slide_master).get(_clr_map(layout).get("bg1", "lt1"))
    return {"hex": white or "FFFFFF", "kind": "solid", "from": "default"}


_TEXT_AREA_TYPES = {1, 2, 3, 7}   # TITLE, BODY, CENTER_TITLE, OBJECT


def layout_margins(layout, slide_w_emu: int | None = None) -> tuple[int, int] | None:
    """(left, right) px of the layout's text area: the leftmost and rightmost
    edge of its title and content placeholders. None when it has neither."""
    lefts, rights = [], []
    for ph in layout.placeholders:
        try:
            t = int(ph.placeholder_format.type)
            if t not in _TEXT_AREA_TYPES or ph.width is None or ph.left is None:
                continue
            left, w = int(ph.left), int(ph.width)
        except Exception:
            continue
        if slide_w_emu and w < 0.25 * slide_w_emu:
            continue    # a small label or number slot, not the text area
        lefts.append(left)
        rights.append(left + w)
    if not lefts:
        return None
    return round(min(lefts) / EMU_PER_PX), round(max(rights) / EMU_PER_PX)


@lru_cache(maxsize=8)
def _open(path: str, mtime: float):
    from pptx import Presentation
    return Presentation(path)


def find_layout(template_path, layout_name: str):
    p = Path(template_path)
    prs = _open(str(p), p.stat().st_mtime)
    for master in prs.slide_masters:
        for lay in master.slide_layouts:
            if (lay.name or "").strip() == (layout_name or "").strip():
                return prs, lay
    return prs, None


def template_layout_facts(template_path, layout_name: str) -> dict:
    """{'background': {...}, 'margins': (left, right) | None} for a layout of a
    template file; empty dict when the file or layout can't be read."""
    try:
        prs, lay = find_layout(template_path, layout_name)
    except Exception:
        return {}
    if lay is None:
        return {}
    return {"background": layout_background(lay),
            "margins": layout_margins(lay, int(prs.slide_width))}


def facts_for(template_path, layout_name: str, layout_chrome=None) -> dict:
    """{'bg_hex', 'bg_kind', 'left', 'right'} for a slide's layout: what the
    registration recorded in chrome.yml when it did, else read from the
    template now (so a template registered before these were recorded needs
    no re-registration). Values are None when unknown."""
    lc = layout_chrome
    out = {"bg_hex": getattr(lc, "background_hex", None),
           "bg_kind": getattr(lc, "background_kind", None),
           "left": getattr(lc, "body_left_px", None),
           "right": getattr(lc, "body_right_px", None)}
    if out["bg_kind"] is None or out["left"] is None:
        live = template_layout_facts(template_path, layout_name) if template_path else {}
        bg = live.get("background") or {}
        if out["bg_kind"] is None:
            out["bg_hex"], out["bg_kind"] = bg.get("hex"), bg.get("kind")
        if out["left"] is None and live.get("margins"):
            out["left"], out["right"] = live["margins"]
    return out


def color_distance(a: str, b: str) -> float:
    """Plain RGB distance between two 'RRGGBB' colors (0 = same, 441 = black/white)."""
    try:
        pa = [int(a.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
        pb = [int(b.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)]
    except (ValueError, AttributeError):
        return 999.0
    return sum((x - y) ** 2 for x, y in zip(pa, pb)) ** 0.5
