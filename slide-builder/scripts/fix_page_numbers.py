#!/usr/bin/env python3
"""fix_page_numbers.py: turn typed page numbers in a finished deck into real
slide-number fields.

Owner's rule (2026-10-12): page numbers are never hardcoded. Decks built
before that date carry a typed "5" in every slide-number box, so moving,
adding or deleting a slide in PowerPoint leaves every number after it wrong.
This tool fixes such a deck: in every page-number box (the slide-number
placeholder, or a box named like a page number) a typed bare number becomes a
slide-number field that keeps the same formatting, with the slide's current
number as its stored text. A box that already holds a field keeps it; only its
stored number is set to the slide's position. Nothing else on any slide is
touched.

  py -3 scripts/fix_page_numbers.py "<deck.pptx>"                  writes "<deck>_page_fields.pptx"
  py -3 scripts/fix_page_numbers.py "<deck.pptx>" --out "<new.pptx>"
  py -3 scripts/fix_page_numbers.py "<deck.pptx>" --in-place       overwrites the input
  py -3 scripts/fix_page_numbers.py "<deck.pptx>" --check          reports only, writes nothing

The input is never overwritten unless --in-place is given. Prints one line per
change. Exit code 0 when the deck was written (or --check found nothing to
fix), 1 when --check found typed page numbers, 2 on a usage or file error.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
for _p in (str(HERE), str(SKILL)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from pptx import Presentation  # noqa: E402

from twins.composer import (  # noqa: E402
    has_slidenum_field, is_page_number_shape, slide_number_offset, stamp_page_number,
)


def fix_presentation(prs) -> list[dict]:
    """Fix every page-number box in `prs` in memory. Returns one record per
    box changed: {"slide", "shape", "before", "after", "action"}."""
    changes: list[dict] = []
    offset = slide_number_offset(prs)
    for pos, slide in enumerate(prs.slides, start=1):
        number = pos + offset
        for shape in slide.shapes:
            if not (getattr(shape, "has_text_frame", False) and is_page_number_shape(shape)):
                continue
            before = (shape.text_frame.text or "").strip()
            had_field = has_slidenum_field(shape)
            action = stamp_page_number(shape, number)
            if not action:
                continue
            after = (shape.text_frame.text or "").strip()
            if action == "kept" and had_field and before == after:
                continue  # already a field with the right number
            changes.append({"slide": pos, "shape": shape.name, "before": before,
                            "after": after, "action": action})
    return changes


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--in-place", action="store_true")
    ap.add_argument("--check", action="store_true",
                    help="report what would change; write nothing")
    a = ap.parse_args(argv)

    src = a.pptx
    if not src.is_file() or src.suffix.lower() != ".pptx":
        print(f"ERROR: not a .pptx file: {src}")
        return 2
    if a.in_place and a.out:
        print("ERROR: give --out or --in-place, not both.")
        return 2
    out = src if a.in_place else (a.out or src.with_name(src.stem + "_page_fields.pptx"))
    if not a.in_place and out.resolve() == src.resolve():
        print("ERROR: --out is the input file; use --in-place to overwrite it.")
        return 2

    prs = Presentation(str(src))
    changes = fix_presentation(prs)
    for c in changes:
        verb = {"converted": "typed number -> field", "kept": "field kept, stored number set",
                "added": "field added"}[c["action"]]
        print(f"  slide {c['slide']:>3}  {c['shape']!r}: {c['before']!r} -> {c['after']!r} ({verb})")
    n_conv = sum(1 for c in changes if c["action"] == "converted")
    print(f"{len(changes)} page-number box(es) changed ({n_conv} typed number(s) made into "
          f"fields) on {len(prs.slides)} slide(s).")
    if a.check:
        return 1 if n_conv else 0
    try:
        prs.save(str(out))
    except (PermissionError, OSError) as exc:
        print(f"ERROR: could not write {out}: {type(exc).__name__}: {exc} (is it open in PowerPoint?)")
        return 2
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
