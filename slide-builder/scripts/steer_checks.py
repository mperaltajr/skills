#!/usr/bin/env python3
"""steer_checks.py: the user's repeated steers, checked by script.

On two decks, 9 of the 12 steers the user had to repeat about rules (not
content) could be checked by a script (owner's decision, 2026-10-14). This
file holds those 9 checks and runs them on every slide's designs: REVIEW.html
shows what they find on each card, before the user looks.

Five are the owner's standing rules and always run:
  trailing_periods   no text item ends with a final period
  legend_top_right   a chart's legend is one row at its top right
                     (checked on finished slides; sketches are measured at finalize)
  unit_no_brackets   "Chart title, Unit", the unit plain, never [M GPUs] or (M GPUs)
  one_summary_line   one summary line, the takeaway under the headline; no second
                     copy of it and no so-what band repeating it
  pinned_option_a    every slide with a pinned source page has an option A
Four are deck rules that need the deck's own words, so they run only once the
user approved them for this deck:
  no_word_on_main    a word kept off the main deck ("scenario"), allowed on
                     appendix slides
  banned_phrases     phrases never used ("the window is now")
  preferred_terms    a term always written another way ("buyers" -> "customers",
                     "Project X" -> "Acme")
  footnote_only      names allowed only in a footnote or source line ("Globex")

A new check is never turned on by the chat. When the user repeats a ruling
(add_ruling.py says REPEAT), the chat PROPOSES one and the user approves it in
their own words (checked against the transcript):

  py -3 scripts/steer_checks.py kinds
  py -3 scripts/steer_checks.py propose --out <deck> --kind banned_phrases \\
        --param "phrases=the window is now" --because R3,R9
  py -3 scripts/steer_checks.py approve --out <deck> --proposal P1 \\
        --approved-in-chat "<the user's words>"
  py -3 scripts/steer_checks.py propose --out <deck> --from-deck <earlier deck folder>
  py -3 scripts/steer_checks.py run --out <deck> [--slide N]
  py -3 scripts/steer_checks.py list --out <deck>

A ruling no kind can check is a code change: say so, and write the change only
after the user approves it.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import _rulings  # noqa: E402

CONFIG = "steer_checks.json"
ALWAYS = ("trailing_periods", "legend_top_right", "unit_no_brackets", "one_summary_line",
          "pinned_option_a")
KINDS = {
    "trailing_periods": ("No text item ends with a final period.", []),
    "legend_top_right": ("A chart's legend is one row at the top right of the chart.", []),
    "unit_no_brackets": ("Chart header 'Title, Unit' with the unit written plainly.", []),
    "one_summary_line": ("One summary line (the takeaway); no second copy, no so-what band.", []),
    "pinned_option_a": ("Every slide with a pinned source page has an option A.", []),
    "no_word_on_main": ("A word kept off main-deck slides (appendix slides allowed).",
                        ["words=<word>[|<word>...]", "appendix_from=<slide number> (optional)"]),
    "banned_phrases": ("Phrases never used anywhere in the deck.", ["phrases=<phrase>[|...]"]),
    "preferred_terms": ("A term always written another way.",
                        ["terms=<wrong>:<right>[|<wrong>:<right>...]"]),
    "footnote_only": ("Names allowed only in a footnote or source line.",
                      ["terms=<name>[|<name>...]"]),
}


def config_path(sd: Path) -> Path:
    return Path(sd) / CONFIG


def load_config(sd: Path) -> dict:
    try:
        cfg = json.loads(config_path(sd).read_text(encoding="utf-8"))
    except Exception:
        cfg = {}
    cfg.setdefault("schema", 1)
    cfg.setdefault("checks", [])
    cfg.setdefault("proposals", [])
    return cfg


def save_config(sd: Path, cfg: dict) -> None:
    Path(sd).mkdir(parents=True, exist_ok=True)
    config_path(sd).write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")


def _split(v: str) -> list[str]:
    return [x.strip() for x in (v or "").split("|") if x.strip()]


# ------------------------------------------------------------------ reading designs

class _Blocks(HTMLParser):
    """Text blocks of a sketch: (role names of the element and its parents, text)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack: list[tuple[str, str]] = []
        self.out: list[list] = []
        self.skip = 0

    VOID = {"br", "img", "hr", "input", "meta", "link", "source", "wbr", "area", "col"}

    def handle_starttag(self, tag, attrs):
        if tag in self.VOID:
            if tag == "br" and self.out:
                self.out[-1][1] += " "
            return
        a = dict(attrs)
        names = " ".join(x for x in (a.get("id"), a.get("class"), a.get("data-template-field"),
                                     a.get("data-shape-id"), a.get("data-role")) if x)
        if tag in ("style", "script", "head", "title"):
            self.skip += 1
        self.stack.append((tag, names))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i][0] == tag:
                if tag in ("style", "script", "head", "title"):
                    self.skip = max(0, self.skip - 1)
                del self.stack[i:]
                return

    def handle_data(self, data):
        if self.skip or not data.strip():
            return
        names = " ".join(n for _t, n in self.stack if n)
        key = id(self.stack[-1]) if self.stack else 0
        if self.out and self.out[-1][2] == key:
            self.out[-1][1] += data
        else:
            self.out.append([names, data, key])


def html_blocks(path: Path) -> list[tuple[str, str]]:
    p = _Blocks()
    try:
        p.feed(path.read_text(encoding="utf-8", errors="replace"))
    except Exception:
        return []
    return [(n.lower(), " ".join(t.split())) for n, t, _k in p.out if t.strip()]


def pptx_blocks(path: Path):
    """(blocks, slide object) of a finished option, or ([], None)."""
    try:
        from pptx import Presentation
        import chart_rules
        prs = Presentation(str(path))
        slide = prs.slides[0]
        blocks = []
        for name, tf in chart_rules.iter_text_frames(slide.shapes):
            for para in tf.paragraphs:
                t = " ".join("".join(r.text for r in para.runs).split())
                if t:
                    blocks.append((name.lower(), t))
        return blocks, slide
    except Exception:
        return [], None


def _is_note(names: str) -> bool:
    return bool(re.search(r"footnote|source|footer|chrome-", names))


def _is_takeaway(names: str) -> bool:
    return bool(re.search(r"(^|[\s/])subtitle(\b|-)", names))


# ------------------------------------------------------------------ the checks

def _word_re(term: str) -> re.Pattern:
    flags = 0 if any(c.isupper() for c in term) else re.I
    return re.compile(rf"(?<![A-Za-z0-9]){re.escape(term)}s?(?![A-Za-z0-9])", flags)


def check_option(kind: str, params: dict, blocks: list, slide_obj, meta_slide: dict,
                 n: int) -> list[str]:
    """Findings (plain sentences) of one check on one design."""
    out: list[str] = []
    if kind == "trailing_periods":
        import chart_rules
        hits = [t for names, t in blocks if chart_rules.ends_with_final_period(t)]
        if hits:
            out.append(f"{len(hits)} text item(s) end with a final period, e.g. "
                       f"\"...{hits[0][-40:]}\"")
    elif kind in ("legend_top_right", "unit_no_brackets"):
        if slide_obj is not None:
            import chart_rules
            for p in chart_rules.chart_problems(slide_obj):
                if kind == "legend_top_right" and p["kind"] == "legend":
                    out.append(p["detail"])
                if kind == "unit_no_brackets" and p["kind"] == "unit_brackets":
                    out.append(p["detail"])
        elif kind == "unit_no_brackets":
            import chart_rules
            for names, t in blocks:
                if "chart-title" in names or "chart-header" in names:
                    b = chart_rules.bracketed_unit(t)
                    if b:
                        out.append(f"chart header \"{t[:60]}\" has its unit in brackets ({b}); "
                                   "write \"Title, Unit\"")
    elif kind == "one_summary_line":
        take = " ".join((meta_slide.get("subtitle") or next(
            (t for nm, t in blocks if _is_takeaway(nm)), "")).split()).lower().rstrip(".")
        for names, t in blocks:
            if _is_takeaway(names) or _is_note(names):
                continue
            tl = t.lower().rstrip(".")
            if take and len(take) > 20 and tl == take:
                out.append("the takeaway appears a second time on the page: "
                           f"\"{t[:60]}\"")
                break
            if re.search(r"so-?what|summary-line|takeaway-band|bottom-band|band-so", names) \
                    and len(t) > 30:
                out.append(f"a second summary line (\"{t[:60]}\"); one summary line only, "
                           "the takeaway under the headline")
                break
    elif kind == "no_word_on_main":
        tag = (meta_slide.get("section_tag") or "").lower()
        start = int(params.get("appendix_from") or 0)
        if "appendix" in tag or (start and n >= start):
            return out
        for w in _split(params.get("words", "")):
            rx = _word_re(w)
            hit = next((t for _nm, t in blocks if rx.search(t)), "")
            if hit:
                out.append(f"\"{w}\" on a main-deck slide: \"{hit[:60]}\"")
    elif kind == "banned_phrases":
        for ph in _split(params.get("phrases", "")):
            rx = re.compile(re.escape(ph), re.I)
            hit = next((t for _nm, t in blocks if rx.search(t)), "")
            if hit:
                out.append(f"\"{ph}\" is not used in this deck: \"{hit[:60]}\"")
    elif kind == "preferred_terms":
        for pair in _split(params.get("terms", "")):
            wrong, _, right = pair.partition(":")
            if not wrong.strip():
                continue
            rx = _word_re(wrong.strip())
            hit = next((t for _nm, t in blocks if rx.search(t)), "")
            if hit:
                out.append(f"\"{wrong.strip()}\" should be \"{right.strip()}\": \"{hit[:60]}\"")
    elif kind == "footnote_only":
        for term in _split(params.get("terms", "")):
            rx = _word_re(term)
            hit = next((t for nm, t in blocks if rx.search(t) and not _is_note(nm)), "")
            if hit:
                out.append(f"\"{term}\" only in a footnote or source line, not in "
                           f"\"{hit[:60]}\"")
    return out


def checks_in_force(sd: Path) -> list[dict]:
    cfg = load_config(sd)
    got = [{"id": k, "kind": k, "params": {}, "standing": True} for k in ALWAYS]
    return got + [c for c in cfg["checks"] if c.get("kind") in KINDS]


def _meta(out: Path) -> dict:
    try:
        return json.loads((out / "_meta.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def _designs(slide_dir: Path) -> list[tuple[str, Path | None, Path | None]]:
    """(letter, sketch html, finished pptx) of every option on disk."""
    letters = sorted({m.group(1) for f in slide_dir.glob("option_*")
                      for m in [re.match(r"option_([A-F])[._]", f.name)] if m})
    out = []
    for L in letters:
        h = slide_dir / f"option_{L}.html"
        x = slide_dir / f"option_{L}.pptx"
        out.append((L, h if h.exists() else None, x if x.exists() else None))
    return out


def run_slide(out: Path, n: int, meta: dict | None = None) -> list[str]:
    """Findings for slide n, one line each, 'option B: <check>: <finding>'."""
    out = Path(out)
    meta = meta if meta is not None else _meta(out)
    ms = next((s for s in meta.get("slides") or [] if int(s.get("n") or 0) == n), {})
    sd = _rulings.session_dir(out=out)
    sdir = out / f"slide_{n:02d}"
    found: list[str] = []
    checks = checks_in_force(sd)
    for c in checks:
        if c["kind"] == "pinned_option_a":
            if (ms.get("pinned_source_page") or "").strip() and sdir.is_dir() and not any(
                    (sdir / f"option_A.{ext}").exists() for ext in ("html", "py")):
                found.append("pinned_option_a: the slide has a pinned source page but no "
                             "option A (option A reproduces that page)")
    for L, h, x in _designs(sdir) if sdir.is_dir() else []:
        blocks, slide_obj = pptx_blocks(x) if x else ([], None)
        if not blocks and h:
            blocks = html_blocks(h)
        for c in checks:
            if c["kind"] == "pinned_option_a":
                continue
            for f in check_option(c["kind"], c.get("params") or {}, blocks, slide_obj, ms, n):
                found.append(f"option {L}: {c['kind']}: {f}")
    return found


# ------------------------------------------------------------------ commands

def _params(items: list[str]) -> dict:
    out = {}
    for it in items:
        k, _, v = it.partition("=")
        if k.strip():
            out[k.strip()] = v.strip()
    return out


def cmd_kinds(_a) -> int:
    for k, (what, params) in KINDS.items():
        on = " (always on)" if k in ALWAYS else ""
        print(f"{k}{on}: {what}")
        for p in params:
            print(f"    --param \"{p}\"")
    return 0


def cmd_propose(a) -> int:
    sd = _rulings.session_dir(out=a.out)
    cfg = load_config(sd)
    new = []
    if a.from_deck:
        other = load_config(_rulings.session_dir(out=a.from_deck))
        for c in other["checks"]:
            new.append({"kind": c["kind"], "params": c.get("params") or {},
                        "because": [f"approved for {a.from_deck.name}"]})
        if not new:
            print(f"No approved deck checks in {a.from_deck}.")
            return 0
    else:
        if a.kind not in KINDS:
            print(f"ERROR: no check kind {a.kind!r}. The kinds: {', '.join(KINDS)}. A "
                  "ruling none of them can check needs a code change: describe it to the "
                  "user and write it only after they approve.")
            return 2
        if a.kind in ALWAYS:
            print(f"{a.kind} already runs on every deck; nothing to propose.")
            return 0
        params = _params(a.param)
        needed = [p.split("=")[0] for p in KINDS[a.kind][1] if "optional" not in p]
        missing = [p for p in needed if not params.get(p)]
        if missing:
            print(f"ERROR: {a.kind} needs --param for: {', '.join(missing)} "
                  f"({'; '.join(KINDS[a.kind][1])}).")
            return 2
        new.append({"kind": a.kind, "params": params,
                    "because": [x.strip() for x in (a.because or "").split(",") if x.strip()]})
    base = 1 + max([int(p["id"][1:]) for p in cfg["proposals"]] or [0])
    for i, p in enumerate(new):
        p["id"] = f"P{base + i}"
        p["at"] = datetime.now().isoformat(timespec="seconds")
        cfg["proposals"].append(p)
    save_config(sd, cfg)
    print("Proposed (NOT on yet). Ask the user, in one message:")
    for p in new:
        ps = "; ".join(f"{k} = {v}" for k, v in p["params"].items())
        print(f"  {p['id']}: an automatic check, {p['kind']} ({KINDS[p['kind']][0]}) "
              f"with {ps}" + (f", because of {', '.join(p['because'])}" if p['because'] else ""))
    print("  \"Turn these checks on for this deck? They run on every design before you see it.\"")
    print("When the user says yes, run with their words exactly as typed:")
    ids = ",".join(p["id"] for p in new)
    print(f"  py -3 scripts/steer_checks.py approve --out \"{a.out}\" --proposal {ids} "
          "--approved-in-chat \"<their words>\"")
    return 0


def cmd_approve(a) -> int:
    sd = _rulings.session_dir(out=a.out)
    cfg = load_config(sd)
    want = [x.strip() for x in a.proposal.split(",") if x.strip()]
    props = {p["id"]: p for p in cfg["proposals"]}
    bad = [w for w in want if w not in props]
    if bad:
        print(f"REFUSED: no proposal {', '.join(bad)} (open: {', '.join(props) or 'none'}).")
        return 3
    ok, why, where = _rulings.check_users_words(a.approved_in_chat or "", a.transcript)
    if not ok:
        print(f"REFUSED: {why}")
        print("Nothing was turned on. The user approves checks in their own words.")
        return 4
    n0 = 1 + max([int(c["id"][1:]) for c in cfg["checks"]] or [0])
    for i, w in enumerate(want):
        p = props[w]
        cfg["checks"].append({"id": f"C{n0 + i}", "kind": p["kind"], "params": p["params"],
                              "because": p.get("because") or [],
                              "approved": {"words": a.approved_in_chat.strip(),
                                           "at": datetime.now().isoformat(timespec="seconds"),
                                           "where": f"{Path(where[0]).name}:{where[1]}"}})
    cfg["proposals"] = [p for p in cfg["proposals"] if p["id"] not in want]
    save_config(sd, cfg)
    print(f"[ok] {len(want)} check(s) on for this deck. They run on every design: "
          "build_review.py shows what they find on each slide's card, and "
          f"py -3 scripts/steer_checks.py run --out \"{a.out}\" prints it.")
    return 0


def cmd_list(a) -> int:
    sd = _rulings.session_dir(out=a.out)
    cfg = load_config(sd)
    for k in ALWAYS:
        print(f"  {k} (always on)")
    for c in cfg["checks"]:
        print(f"  {c['id']} {c['kind']} {c.get('params')} approved: "
              f"\"{(c.get('approved') or {}).get('words', '')[:60]}\"")
    for p in cfg["proposals"]:
        print(f"  {p['id']} PROPOSED, not on: {p['kind']} {p.get('params')}")
    return 0


def cmd_run(a) -> int:
    meta = _meta(a.out)
    ns = [a.slide] if a.slide else [int(s["n"]) for s in meta.get("slides") or []]
    total = 0
    for n in ns:
        f = run_slide(a.out, n, meta)
        total += len(f)
        for line in f:
            print(f"slide {n}: {line}")
    print(f"{total} finding(s) on {len(ns)} slide(s).")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="The user's repeated steers, checked by script.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("kinds")
    p = sub.add_parser("propose")
    p.add_argument("--out", required=True, type=Path)
    p.add_argument("--kind", default="")
    p.add_argument("--param", action="append", default=[])
    p.add_argument("--because", default="")
    p.add_argument("--from-deck", type=Path, default=None, dest="from_deck")
    p = sub.add_parser("approve")
    p.add_argument("--out", required=True, type=Path)
    p.add_argument("--proposal", required=True)
    p.add_argument("--approved-in-chat", dest="approved_in_chat", required=True)
    p.add_argument("--transcript", default=None, help=argparse.SUPPRESS)
    for name in ("list", "run"):
        p = sub.add_parser(name)
        p.add_argument("--out", required=True, type=Path)
        if name == "run":
            p.add_argument("--slide", type=int, default=None)
    a = ap.parse_args(argv)
    if getattr(a, "out", None):
        _rulings.announce(out=a.out)
    return {"kinds": cmd_kinds, "propose": cmd_propose, "approve": cmd_approve,
            "list": cmd_list, "run": cmd_run}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main())
