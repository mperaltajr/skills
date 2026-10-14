#!/usr/bin/env python3
"""clone_pages.py: pages that repeat are designed once, then copied by script.

Owner's decision 2026-10-14 (speed verdict Q3, items 1 to 3):

  Section dividers  Only the FIRST divider in the brief goes to a designer.
                    Its design keeps the per-page words (section number, name,
                    one-line description, the section list) as named settings
                    in a PAGE VALUES block at the top of its file. This script
                    writes the file of every other divider from that design,
                    with the divider's own values from the brief. A comment on
                    any divider goes to the first one in edit mode; the copies
                    are then made again from it.
  Cover, contents,  Designed once per template. The first deck designs them
  executive summary (same PAGE VALUES block); when that deck compiles, the
                    picked design is kept for the template. A later deck on
                    the same template sends no designer for them: this script
                    fills the kept design from the new storyline, word for
                    word (contents page ranges come from the sections).
  Supplied pages    A page pinned to a slide of a PowerPoint file is copied as
  (.pptx pins)      is (twins.composer.copy_shape_with_parts). A worker in EDIT
                    MODE then changes only the text the brief changes, in the
                    copy's TEXT_CHANGES block. It never redraws the page.

Content pages are never built by script (verdict Q3, "What not to do").

Every copy is an ordinary option file: finalize builds, themes, fit-checks and
renders it like any page, it is on REVIEW.html and FINAL-CHECK.html, and
slide-qc reads its picture like any page.

Commands (finalize runs `refresh` itself; the others are for a person):
  py -3 scripts/clone_pages.py refresh --out <build> [--slide N]
  py -3 scripts/clone_pages.py status  --out <build>
  py -3 scripts/clone_pages.py save-masters --out <build>
Exit: 0 ok | 1 a copy could not be made (the reason is printed) | 2 bad usage
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import os
import pprint
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
for _d in (str(HERE), str(SKILL)):
    if _d not in sys.path:
        sys.path.insert(0, _d)

import _paths as _p  # noqa: E402

# ---------------------------------------------------------------- the kinds
DIVIDER, COVER, CONTENTS, SUMMARY, PINNED = (
    "divider", "cover", "contents", "exec_summary", "pinned")
TEMPLATE_KINDS = (COVER, CONTENTS, SUMMARY)
LABEL = {DIVIDER: "section divider", COVER: "cover", CONTENTS: "contents page",
         SUMMARY: "executive summary", PINNED: "copied page"}
# The names a design of each kind must keep in its PAGE VALUES block.
REQUIRED = {
    DIVIDER: ("SECTION_NUMBER", "SECTION_TITLE", "SECTION_DESCRIPTION", "SECTIONS"),
    COVER: ("COVER_TITLE", "COVER_SUBTITLE", "COVER_PRESENTER", "COVER_DATE", "COVER_CLIENT"),
    CONTENTS: ("CONTENTS",),
    SUMMARY: ("SUMMARY_ROWS",),
}

BLOCK_BEGIN = "# --- PAGE VALUES"
BLOCK_BEGIN_LINE = ("# --- PAGE VALUES (clone_pages.py fills these for each copy; "
                    "keep the names) ---")
BLOCK_END = "# --- END PAGE VALUES ---"
COPY_MARK = "# PAGE COPY:"
PENDING_MARK = "# COPY_EDIT_PENDING:"
CHANGES_BEGIN = "# --- TEXT CHANGES"
CHANGES_END = "# --- END TEXT CHANGES ---"
RECORD_NAME = "_page_copy.json"
EDIT_REQUEST = "_edit_request.md"
PREV_DIR = "_prev_copy"
MASTERS_ENV = "SLIDE_LAB_PAGE_MASTERS_DIR"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def _sid(n: int) -> str:
    return _p.slide_key(n)


# ---------------------------------------------------------------- the brief
def _norm_type(a: str) -> str:
    return re.sub(r"\s*/\s*", "/", (a or "").strip().lower())


def page_kind(slide: dict) -> str | None:
    """divider / cover / contents / exec_summary from the slide type, else None."""
    a = _norm_type(slide.get("archetype") or "")
    if not a:
        return None
    if "divider" in a:
        return DIVIDER
    if a in ("cover/title", "cover", "title", "title page"):
        return COVER
    if "agenda" in a or "contents" in a:
        return CONTENTS
    if "executive summary" in a:
        return SUMMARY
    return None


_BULLET_RE = re.compile(r"^(\s*)[-*•]\s+(.*\S)\s*$")
_LABEL_RE = re.compile(r"^\*\*(.+?)\*\*\s*(?:[—–:]|-(?=\s))?\s*(.*)$")
_NONE_RE = re.compile(r"^\s*(?:none|n/?a|no line|-|—)\b", re.I)


def _clean(s: str) -> str:
    """The brief's words as written: spaces tidied, one closing period dropped
    (slide text carries no closing periods)."""
    s = " ".join((s or "").split())
    if s.endswith(".") and not s.endswith(".."):
        s = s[:-1]
    return s


def bullets(text: str) -> list[dict]:
    """[{label, text, points}] from a brief block: top-level bullets, a leading
    **bold label** split off, indented bullets kept as points, a wrapped line
    joined to the bullet above it."""
    items: list[dict] = []
    for line in (text or "").splitlines():
        m = _BULLET_RE.match(line)
        if not m:
            if items and line.strip() and not line.lstrip().startswith(("#", "**", "|")):
                last = items[-1]
                if last["points"]:
                    last["points"][-1] = _clean(last["points"][-1] + " " + line.strip())
                else:
                    last["text"] = _clean(last["text"] + " " + line.strip())
            continue
        indent = len(m.group(1).replace("\t", "    "))
        body = m.group(2)
        if indent >= 2 and items:
            items[-1]["points"].append(_clean(body))
            continue
        lm = _LABEL_RE.match(body)
        label, rest = (lm.group(1), lm.group(2)) if lm else ("", body)
        items.append({"label": _clean(label.strip().rstrip(":").strip()),
                      "text": _clean(rest), "points": []})
    return items


def _section_description(d: dict) -> str:
    """A divider's one line: **Section description:**, else its takeaway, else
    its first bullet (a bold label and a leading repeat of the title dropped)."""
    own = _clean(d.get("section_description") or "")
    if own:
        return own
    tk = (d.get("so_what") or "").strip()
    if tk and not _NONE_RE.match(tk):
        return _clean(tk)
    items = bullets(d.get("evidence_content") or "")
    if not items:
        return ""
    t = items[0]["text"]
    title = (d.get("title") or "").strip()
    if title and t.casefold().startswith(title.casefold()):
        t = t[len(title):].lstrip(" :—–-")
    return _clean(t)


def sections(slides: list) -> list[dict]:
    """The deck's sections, one per section divider, in brief order: number,
    name, one-line description, first and last page."""
    ordered = sorted(slides, key=lambda s: s["slide_n"])
    total = max((s["slide_n"] for s in ordered), default=0)
    divs = [s for s in ordered if page_kind(s) == DIVIDER]
    out = []
    for i, d in enumerate(divs):
        last = divs[i + 1]["slide_n"] - 1 if i + 1 < len(divs) else total
        out.append({"number": i + 1, "name": _clean(d.get("title") or ""),
                    "description": _section_description(d),
                    "first_page": d["slide_n"], "last_page": last})
    return out


_PAGES_RE = re.compile(r"^pages?\s+(\d+)\s*(?:to|-|–|—)\s*(\d+)\s*[:,.—–-]?\s*(.*)$",
                       re.I)


def _contents(slide: dict, secs: list) -> tuple[list, list]:
    items = bullets(slide.get("evidence_content") or "")
    warns: list[str] = []
    if not items:
        return [dict(s) for s in secs], warns
    rows = []
    for i, it in enumerate(items):
        label, text = it["label"], it["text"]
        if not label:
            label, text = text, ""
        m = re.match(r"^(\d{1,2})[.)]?\s+(.*)$", label)
        num, name = (int(m.group(1)), m.group(2).strip()) if m else (i + 1, label)
        first = last = None
        pm = _PAGES_RE.match(text)
        if pm:
            first, last, text = int(pm.group(1)), int(pm.group(2)), pm.group(3)
        if len(secs) == len(items):
            sf, sl = secs[i]["first_page"], secs[i]["last_page"]
            if first is not None and (first, last) != (sf, sl):
                warns.append(f"slide {slide['slide_n']} (contents) says pages {first} to "
                             f"{last} for {name!r}; the sections run {sf} to {sl}. The page "
                             "uses the sections' pages. Fix the storyline if it is wrong.")
            first, last = sf, sl
        rows.append({"number": num, "name": name, "description": _clean(text),
                     "first_page": first, "last_page": last})
    return rows, warns


def _cover(slide: dict) -> dict:
    keys = {"title": "COVER_TITLE", "subtitle": "COVER_SUBTITLE",
            "presenter": "COVER_PRESENTER", "date": "COVER_DATE",
            "audience / client": "COVER_CLIENT", "audience/client": "COVER_CLIENT",
            "client": "COVER_CLIENT", "audience": "COVER_CLIENT"}
    vals = {k: "" for k in REQUIRED[COVER]}
    for it in bullets(slide.get("evidence_content") or ""):
        k = keys.get(it["label"].strip().lower())
        if k and not vals[k]:
            vals[k] = it["text"]
    if not vals["COVER_TITLE"]:
        vals["COVER_TITLE"] = _clean(slide.get("title") or "")
    return vals


def values_for(kind: str, slide: dict, slides: list, secs: list | None = None) -> tuple[dict, list]:
    """The PAGE VALUES of one page, from the brief. Returns (values, warnings)."""
    secs = sections(slides) if secs is None else secs
    if kind == DIVIDER:
        me = next((s for s in secs if s["first_page"] == slide["slide_n"]), None)
        return ({"SECTION_NUMBER": me["number"] if me else 1,
                 "SECTION_TITLE": me["name"] if me else _clean(slide.get("title") or ""),
                 "SECTION_DESCRIPTION": me["description"] if me else "",
                 "SECTIONS": [s["name"] for s in secs]}, [])
    if kind == COVER:
        return _cover(slide), []
    if kind == CONTENTS:
        rows, warns = _contents(slide, secs)
        return {"CONTENTS": rows}, warns
    if kind == SUMMARY:
        return {"SUMMARY_ROWS": bullets(slide.get("evidence_content") or "")}, []
    return {}, []


# ------------------------------------------------------- per-template masters
def masters_dir(template) -> Path:
    """Where the designs kept for this template live."""
    template = Path(template)
    env = os.environ.get(MASTERS_ENV)
    if env:
        return Path(env) / re.sub(r"[^A-Za-z0-9_.-]+", "_", template.stem)
    return _p.template_sidecar_dir(template) / "page_masters"


def template_master(template, kind: str) -> Path | None:
    """The kept design of this kind for the template, when it is usable."""
    if not template or kind not in TEMPLATE_KINDS:
        return None
    f = masters_dir(template) / f"{kind}.py"
    if not f.is_file():
        return None
    _vals, _span, err = parse_block(f.read_text(encoding="utf-8"), kind)
    return None if err else f


# ---------------------------------------------------------------- the plan
def plan(slides: list, *, template=None, bases=(), design_fresh=(),
         keep_masters=()) -> tuple[dict, list]:
    """{slide_n: page_copy entry} for every page made or copied by script, and
    warnings. design_fresh: slides that get a designer even when the template
    has a kept design (a rebuild). keep_masters: slides an earlier prep of this
    build already gave a designer (kept that way on a single-slide prep)."""
    out: dict[int, dict] = {}
    warns: list[str] = []
    for s in slides:
        pin = (s.get("pinned_source_page") or "").strip()
        if not pin:
            continue
        try:
            import source_ledger
            path, k = source_ledger.parse_pin(pin, bases)
        except Exception:
            continue                     # source_ledger reports a bad pin itself
        if path.suffix.lower() == ".pptx":
            out[s["slide_n"]] = {"kind": PINNED, "role": "pinned", "source": str(path),
                                 "source_slide": k}
    secs = sections(slides)
    groups: dict[str, list] = {}
    for s in sorted(slides, key=lambda x: x["slide_n"]):
        k = page_kind(s)
        if k and s["slide_n"] not in out:
            groups.setdefault(k, []).append(s)
    fresh = set(design_fresh) | set(keep_masters)
    for kind, members in groups.items():
        if kind == DIVIDER:
            if len(members) < 2:
                continue
            m = members[0]
            mv, _ = values_for(kind, m, slides, secs)
            out[m["slide_n"]] = {"kind": kind, "role": "master", "values": mv,
                                 "copies": [c["slide_n"] for c in members[1:]]}
            for c in members[1:]:
                cv, _ = values_for(kind, c, slides, secs)
                out[c["slide_n"]] = {"kind": kind, "role": "copy", "master": m["slide_n"],
                                     "values": cv}
            continue
        s = members[0]
        v, w = values_for(kind, s, slides, secs)
        warns += w
        kept = template_master(template, kind)
        if kept and s["slide_n"] not in fresh:
            out[s["slide_n"]] = {"kind": kind, "role": "copy", "master": "template",
                                 "master_file": str(kept), "values": v}
        else:
            out[s["slide_n"]] = {"kind": kind, "role": "master", "values": v, "copies": []}
    return out, warns


def is_deck_copy(pc: dict | None) -> bool:
    return bool(pc) and pc.get("role") == "copy" and isinstance(pc.get("master"), int)


def no_worker_slides(meta: dict) -> list[int]:
    """Slides no designer is sent for: copies (of a page in this deck or of a
    design kept for the template). Pinned copies get a worker in EDIT MODE."""
    return sorted(s["n"] for s in meta.get("slides") or []
                  if (s.get("page_copy") or {}).get("role") == "copy")


def edit_mode_slides(meta: dict) -> list[int]:
    """Copied supplied pages: their worker runs in EDIT MODE on _edit_request.md."""
    return sorted(s["n"] for s in meta.get("slides") or []
                  if (s.get("page_copy") or {}).get("role") == "pinned")


def apply_to_meta(meta_path: Path, copy_plan: dict) -> None:
    """Record the plan in _meta.json: page_copy per slide, and the option
    letters of each copy (the master's for a deck copy, A for the others)."""
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    by_n = {s.get("n"): s for s in meta.get("slides") or []}
    for n, s in by_n.items():
        pc = copy_plan.get(n)
        if pc:
            s["page_copy"] = pc
        else:
            s.pop("page_copy", None)
    for n, s in by_n.items():
        pc = s.get("page_copy") or {}
        if is_deck_copy(pc) and pc["master"] in by_n:
            s["options"] = list(by_n[pc["master"]].get("options") or _p.option_letters())
        elif pc.get("role") in ("pinned",) or (pc.get("role") == "copy"):
            s["options"] = ["A"]
    meta_path.write_text(json.dumps(meta, indent=2), encoding="utf-8")


# ----------------------------------------------------- the PAGE VALUES block
def _literal(v) -> str:
    return pprint.pformat(v, width=96, sort_dicts=False)


def render_block(values: dict) -> str:
    lines = [BLOCK_BEGIN_LINE]
    for k, v in values.items():
        lit = _literal(v)
        lines.append(f"{k} = " + lit.replace("\n", "\n" + " " * (len(k) + 3)))
    lines.append(BLOCK_END)
    return "\n".join(lines) + "\n"


def parse_block(text: str, kind: str) -> tuple[dict | None, tuple[int, int] | None, str]:
    """(values, (start, end) of the block in text, error). The block is the
    lines from '# --- PAGE VALUES' to '# --- END PAGE VALUES ---'; between them
    only NAME = <plain value> lines, with every name the kind needs."""
    lines = text.splitlines(keepends=True)
    start = next((i for i, ln in enumerate(lines) if ln.startswith(BLOCK_BEGIN)), None)
    if start is None:
        return None, None, "there is no PAGE VALUES block"
    end = next((i for i in range(start + 1, len(lines)) if lines[i].startswith(BLOCK_END)), None)
    if end is None:
        return None, None, f"the PAGE VALUES block has no closing line ({BLOCK_END})"
    body = "".join(lines[start + 1:end])
    try:
        tree = ast.parse(body)
    except SyntaxError as exc:
        return None, None, f"the PAGE VALUES block is not valid Python ({exc.msg})"
    vals: dict = {}
    for node in tree.body:
        if not (isinstance(node, ast.Assign) and len(node.targets) == 1
                and isinstance(node.targets[0], ast.Name)):
            return None, None, ("the PAGE VALUES block holds something other than "
                                "NAME = value lines; move it below the block")
        try:
            vals[node.targets[0].id] = ast.literal_eval(node.value)
        except ValueError:
            return None, None, (f"{node.targets[0].id} in the PAGE VALUES block is not a "
                                "plain value (text, a number, a list or a dict)")
    need = REQUIRED.get(kind, ())
    missing = [k for k in need if k not in vals]
    if missing:
        return None, None, "the PAGE VALUES block lacks " + ", ".join(missing)
    extra = [k for k in vals if k not in need]
    if extra:
        return None, None, ("the PAGE VALUES block holds names a copy cannot fill ("
                            + ", ".join(extra) + "); move them below the block")
    a = sum(len(ln) for ln in lines[:start])
    b = sum(len(ln) for ln in lines[:end + 1])
    return vals, (a, b), ""


def _type_problems(got: dict, want: dict) -> list[str]:
    """Names whose value changed kind (a number made text, a list made a dict):
    the copies are filled with the brief's kinds, so the design must use them."""
    probs = []
    for k, w in want.items():
        g = got.get(k)
        if type(g) is not type(w):
            probs.append(f"{k} must stay {type(w).__name__} ({_literal(w)[:60]}), "
                         f"it is {type(g).__name__}")
        elif isinstance(w, list) and w and g and type(g[0]) is not type(w[0]):
            probs.append(f"{k} must stay a list of {type(w[0]).__name__}")
    return probs


def fill(master_text: str, kind: str, values: dict, slide_n: int, letter: str,
         from_what: str) -> str:
    """The master's design with this page's values (raises ValueError when
    the master has no usable block)."""
    _vals, span, err = parse_block(master_text, kind)
    if err:
        raise ValueError(err)
    text = master_text[:span[0]] + render_block(values) + master_text[span[1]:]
    lines = text.splitlines(keepends=True)
    lines = [ln for ln in lines if not ln.startswith(COPY_MARK)]
    if lines:
        lines[0] = re.sub(r"^# Slide \d+ option [A-F]", f"# Slide {slide_n} option {letter}",
                          lines[0])
    nl = "\r\n" if "\r\n" in master_text else "\n"
    mark = (f"{COPY_MARK} made by clone_pages.py from {from_what}. Do not edit this file: "
            f"edit {from_what} (every copy is made again from it).{nl}")
    lines.insert(1 if lines else 0, mark)
    return "".join(lines)


# ---------------------------------------------------------------- refresh
def _meta(out: Path) -> dict:
    try:
        return json.loads(_p.meta_json(out).read_text(encoding="utf-8"))
    except Exception:
        return {}


def _record(sd: Path) -> dict:
    try:
        return json.loads((sd / RECORD_NAME).read_text(encoding="utf-8"))
    except Exception:
        return {}


def _write_record(sd: Path, rec: dict) -> None:
    (sd / RECORD_NAME).write_text(json.dumps(rec, indent=2), encoding="utf-8")


def _design_file(sd: Path, letter: str) -> Path:
    return sd / _p.option_py_name(letter)


def _drop_letter(sd: Path, letter: str) -> None:
    for name in (_p.option_py_name(letter), _p.option_pptx_name(letter),
                 _p.option_png_name(letter), _p.option_qc_json_name(letter)):
        try:
            (sd / name).unlink()
        except OSError:
            pass


def refresh(out: Path, only: int | None = None) -> dict:
    """Make every copy current from its master. `only`: one slide (a copy, or
    a master and its copies). Returns what happened:
      written  [(slide, letter)]       files written or made again
      waiting  [(slide, letter, why)]  the master's file is not there yet
      errors   [(slide, why)]          the master's design cannot be copied
      pending  [slide]                 a copied supplied page waits on its edit
      notes    [text]
    A file is only written when its text changes (its date is part of the
    pick stamp)."""
    out = Path(out)
    meta = _meta(out)
    by_n = {s.get("n"): s for s in meta.get("slides") or [] if isinstance(s.get("n"), int)}
    rep: dict = {"written": [], "waiting": [], "errors": [], "pending": [], "notes": [],
                 "same": 0}
    for n in sorted(by_n):
        pc = by_n[n].get("page_copy") or {}
        if not pc:
            continue
        role, kind = pc.get("role"), pc.get("kind")
        sd = _p.slide_dir(out, n)
        if role == "pinned":
            if only not in (None, n):
                continue
            f = _design_file(sd, "A")
            if f.is_file() and PENDING_MARK in f.read_text(encoding="utf-8", errors="replace"):
                rep["pending"].append(n)
            continue
        if role == "master":
            if only not in (None, n):
                continue
            _check_master(out, n, by_n[n], rep)
            continue
        if role != "copy":
            continue
        master = pc.get("master")
        if only is not None and only != n and only != master:
            continue
        rec = _record(sd)
        if rec.get("detached"):
            continue
        if master == "template":
            mf = Path(pc.get("master_file") or "")
            sources = {"A": mf}
            from_what = f"the {LABEL[kind]} design kept for this template ({mf})"
        else:
            msd = _p.slide_dir(out, master)
            letters = by_n.get(master, {}).get("options") or list(_p.option_letters())
            sources = {L: _design_file(msd, L) for L in letters}
            from_what = f"slide {master} option {{L}}"
        letters_rec = rec.setdefault("letters", {})
        for L, mf in sources.items():
            what = from_what.replace("{L}", L)
            if not mf.is_file():
                if (mf.parent / f"option_{L}.html").exists():
                    rep["errors"].append((n, f"{what} is a sketch (HTML); a copy needs the "
                                          "direct-path script. Prep this build again: the "
                                          "first divider is set to the direct path."))
                else:
                    rep["waiting"].append((n, L, f"made from {what} once that file is written"))
                continue
            mtext = mf.read_text(encoding="utf-8")
            if mtext.splitlines()[:1] and mtext.splitlines()[0].startswith("# SKELETON_REJECTED:"):
                rep["errors"].append((n, f"{what} was turned down by its designer "
                                      "(SKELETON_REJECTED), so there is no design to copy"))
                continue
            try:
                new = fill(mtext, kind, pc.get("values") or {}, n, L, what)
            except ValueError as exc:
                rep["errors"].append((n, f"{what} cannot be copied: {exc}. Send its "
                                      f"designer back to add the block (its _prompt.md, "
                                      "section 0, has it word for word)"))
                continue
            target = _design_file(sd, L)
            if target.is_file():
                cur = target.read_text(encoding="utf-8", errors="replace")
                if cur == new:
                    rep["same"] += 1
                    continue
                written = (letters_rec.get(L) or {}).get("written_sha")
                if written and _sha(cur) != written:
                    if master == "template":
                        # Edited in place from the user's comment: this page is
                        # its own design now, and the next deck may keep it.
                        rec["detached"] = True
                        rec["detached_at"] = _now()
                        _write_record(sd, rec)
                        rep["notes"].append(
                            f"slide {n}: its {LABEL[kind]} was edited in place, so it is "
                            "no longer filled from the kept design; when this deck "
                            "compiles, the edited design is kept for the template.")
                        break
                    prev = sd / PREV_DIR
                    prev.mkdir(exist_ok=True)
                    shutil.copy2(target, prev / target.name)
                    rep["notes"].append(
                        f"WARNING: slide {n} option {L} was changed by hand, but it is a "
                        f"copy of slide {master}: it was made again from slide {master} "
                        f"(the changed file is kept at {prev / target.name}). Put the "
                        f"change on slide {master}; every copy follows it.")
            target.write_text(new, encoding="utf-8", newline="")
            letters_rec[L] = {"written_sha": _sha(new), "master_sha": _sha(mtext),
                              "master": str(mf), "at": _now()}
            rep["written"].append((n, L))
        else:
            # Letters the master no longer has go (they were made here).
            for f in sorted(sd.glob("option_*.py")):
                L = f.stem.split("_")[1] if "_" in f.stem else ""
                if f.stem.endswith("_native") or L in sources or L not in letters_rec:
                    continue
                _drop_letter(sd, L)
                letters_rec.pop(L, None)
                rep["notes"].append(f"slide {n}: option {L} removed (slide {master} has "
                                    "no such design any more).")
            rec.update({"kind": kind, "master": master, "updated_at": _now()})
            _write_record(sd, rec)
    return rep


def _check_master(out: Path, n: int, entry: dict, rep: dict) -> None:
    """A master's PAGE VALUES must hold the brief's words, in the kinds the
    copies use. Words that differ are set back to the brief's (the block is
    data, the design is the code below it); a changed kind is an error."""
    pc = entry.get("page_copy") or {}
    kind, want = pc.get("kind"), pc.get("values") or {}
    sd = _p.slide_dir(out, n)
    for L in entry.get("options") or list(_p.option_letters()):
        f = _design_file(sd, L)
        if not f.is_file():
            continue
        text = f.read_text(encoding="utf-8")
        if text.splitlines()[:1] and text.splitlines()[0].startswith("# SKELETON_REJECTED:"):
            continue
        got, span, err = parse_block(text, kind)
        if err:
            if pc.get("copies"):
                rep["errors"].append((n, f"slide {n} option {L} ({LABEL[kind]}) cannot be "
                                      f"copied to slides {_nums(pc['copies'])}: {err}. Send "
                                      "its designer back with its _prompt.md (section 0 has "
                                      "the block word for word)"))
            else:
                rep["notes"].append(f"slide {n} option {L}: {err}, so this {LABEL[kind]} "
                                    "cannot be kept for later decks on this template.")
            continue
        probs = _type_problems(got, want)
        if probs and not pc.get("copies"):
            rep["notes"].append(f"slide {n} option {L}: " + "; ".join(probs) + f", so this "
                                f"{LABEL[kind]} cannot be kept for later decks as it is.")
            continue
        if probs:
            rep["errors"].append((n, f"slide {n} option {L}: " + "; ".join(probs)
                                  + ". Copies are filled with the brief's values, so the "
                                  "design must format them in its code. Send its designer "
                                  "back with this line"))
            continue
        if got != want:
            new = text[:span[0]] + render_block(want) + text[span[1]:]
            f.write_text(new, encoding="utf-8", newline="")
            rep["notes"].append(f"slide {n} option {L}: page values set back to the brief's "
                                "words (the design had changed them).")


def _nums(ns) -> str:
    return ", ".join(str(x) for x in ns)


def report_lines(rep: dict, out: Path, template=None, only: int | None = None) -> list[str]:
    """Plain lines for finalize's output."""
    lines = []
    w = rep.get("written") or []
    if w:
        lines.append("  made " + str(len(w)) + " copied page file(s): "
                     + ", ".join(f"slide {n} option {L}" for n, L in w))
    if rep.get("same"):
        lines.append(f"  {rep['same']} copied page file(s) already current")
    for n, L, why in rep.get("waiting") or []:
        lines.append(f"  slide {n} option {L}: {why}")
    lines += ["  " + t for t in rep.get("notes") or []]
    if only is not None and w:
        others = sorted({n for n, _L in w if n != only})
        if others and template:
            lines.append("  The copies changed. Finalize them too:")
            for n in others:
                lines.append(f"    py -3 scripts/finalize_deck.py --out \"{out}\" "
                             f"--template \"{template}\" --slide {n}")
    return lines


def problem_lines(rep: dict, out: Path) -> list[str]:
    lines = []
    for n, why in rep.get("errors") or []:
        lines.append(f"  slide {n}: {why}.")
    for n in rep.get("pending") or []:
        lines.append(f"  slide {n}: the supplied page is copied, and its text is not "
                     "updated yet. Dispatch slide-builder-worker in EDIT MODE on "
                     f"{_p.slide_dir(out, n) / EDIT_REQUEST} (it changes only the text "
                     "the brief changes), then run finalize again.")
    return lines


# ------------------------------------------------------- picks and comments
def _letters(v) -> str:
    return "".join(sorted(v)) if isinstance(v, (list, tuple)) else str(v or "")


def check_copy_picks(meta: dict, picks: dict) -> str:
    """'' when every copy of a page in this deck has the same pick as its
    master (or is left out); else the plain reason."""
    for s in meta.get("slides") or []:
        pc = s.get("page_copy") or {}
        if not is_deck_copy(pc):
            continue
        n, m = s["n"], pc["master"]
        cv, mv = picks.get(_sid(n)), picks.get(_sid(m))
        if cv in (None, "-", "") or mv in (None, "-", ""):
            continue
        if _letters(cv) != _letters(mv):
            return (f"slide {n} is a copy of slide {m}'s design, so it takes the same "
                    f"design: slide {m} is {_letters(mv)}, slide {n} is {_letters(cv)}. Ask "
                    f"the user to pick {_letters(mv)} on slide {n} too (or pick again on "
                    f"slide {m}), then click Build my deck again on REVIEW.html.")
    return ""


def route_copy_feedback(meta: dict, fb: dict, picks: dict) -> tuple[dict, list, str]:
    """A comment on a copy goes to its master (edit mode there; the copies are
    made again after it). Returns (feedback, notes, refusal)."""
    fb = {k: dict(v) for k, v in (fb or {}).items()}
    notes = []
    for s in meta.get("slides") or []:
        pc = s.get("page_copy") or {}
        if not is_deck_copy(pc):
            continue
        n, m = s["n"], pc["master"]
        sid, msid = _sid(n), _sid(m)
        if sid not in fb or picks.get(sid) in (None, "-"):
            continue
        if picks.get(msid) in (None, "-"):
            return fb, [], (f"there is a comment on slide {n}, a copy of slide {m}'s design, "
                            f"and slide {m} is left out. A comment on a copy is applied to "
                            f"slide {m} and every copy is made again from it, so slide {m} "
                            "must be in the deck. Ask the user to keep it (or to leave out "
                            "the copies too), then click Build my deck again.")
        mf = fb.setdefault(msid, {})
        for k, v in fb.pop(sid).items():
            add = f"(from slide {n}, a copy of this design) {v}"
            mf[k] = (mf[k] + "; " + add) if mf.get(k) else add
        notes.append(f"comment on slide {n} goes to slide {m}: slide {n} is a copy of "
                     f"slide {m}'s design, so the edit is made there and every copy is made "
                     "again from it at finalize.")
    return fb, notes, ""


# ----------------------------------------------------- keep for the template
def save_template_masters(out: Path, picks: dict | None = None) -> list[str]:
    """When a deck compiles: keep its picked cover, contents and executive
    summary designs for the template, so the next deck on it fills them by
    script. Only a direct-path design with a usable PAGE VALUES block is kept."""
    out = Path(out)
    meta = _meta(out)
    if picks is None:
        try:
            import _state
            picks = ((_state.read_state(out).get("review") or {}).get("picks")) or {}
        except Exception:
            picks = {}
    template = meta.get("template")
    if not template:
        return []
    lines = []
    for s in meta.get("slides") or []:
        pc = s.get("page_copy") or {}
        kind = pc.get("kind")
        if kind not in TEMPLATE_KINDS:
            continue
        n = s["n"]
        sd = _p.slide_dir(out, n)
        if pc.get("role") == "copy" and not _record(sd).get("detached"):
            continue                       # filled from the kept design, unchanged
        v = picks.get(_sid(n))
        L = (v[0] if isinstance(v, list) and v else v) if v not in (None, "-", "") else None
        if not L:
            continue
        f = _design_file(sd, L)
        if not f.is_file() or (sd / f"option_{L}.html").exists():
            continue
        text = f.read_text(encoding="utf-8")
        _vals, _span, err = parse_block(text, kind)
        if err:
            lines.append(f"  slide {n}: the {LABEL[kind]} is not kept for later decks ({err}).")
            continue
        text = "".join(ln for ln in text.splitlines(keepends=True)
                       if not ln.startswith(COPY_MARK))
        dst_dir = masters_dir(template)
        dst = dst_dir / f"{kind}.py"
        if dst.is_file() and dst.read_text(encoding="utf-8") == text:
            continue
        dst_dir.mkdir(parents=True, exist_ok=True)
        dst.write_text(text, encoding="utf-8", newline="")
        (dst_dir / f"{kind}.json").write_text(json.dumps(
            {"kind": kind, "deck": str(out), "slide": n, "letter": L, "saved_at": _now()},
            indent=2), encoding="utf-8")
        lines.append(f"  kept the {LABEL[kind]} design of slide {n} for this template: the "
                     "next deck on it fills it from its own storyline, with no designer. "
                     f"Kept at {dst}")
    return lines


# -------------------------------------------------- prep: prompts and notes
def master_prompt_section(n: int, pc: dict, plan_all: dict) -> str:
    kind, vals = pc["kind"], pc["values"]
    copies = pc.get("copies") or []
    if kind == DIVIDER:
        why = (f"This is the deck's first section divider. Your design is copied to the "
               f"other dividers (slides {_nums(copies)}) by `scripts/clone_pages.py`, which "
               "changes ONLY the values below. No other designer works on a divider.")
        others = [plan_all[c]["values"] for c in copies if c in plan_all]
        longest_t = max([vals["SECTION_TITLE"]] + [o["SECTION_TITLE"] for o in others], key=len)
        longest_d = max([vals["SECTION_DESCRIPTION"]] + [o["SECTION_DESCRIPTION"] for o in others],
                        key=len)
        fit = (f"Every copy must fit: the longest section name is {longest_t!r} and the "
               f"longest description is {longest_d!r}. Section numbers run 1 to "
               f"{len(vals['SECTIONS'])}; draw the number from SECTION_NUMBER (for example "
               "f\"{SECTION_NUMBER:02d}\"). If the layout prints a fixed number of its own "
               "(such as \"01\"), cover or replace it here: every copy carries the fix.")
    else:
        why = (f"This page is the {LABEL[kind]} design for this template. When the deck "
               "compiles, the picked design is kept; later decks on this template fill it "
               "from their own storyline with `scripts/clone_pages.py` and send no "
               "designer, so the lists below may be shorter or longer in another deck.")
        fit = ("Draw every list in a loop over its items, and make the design hold fewer or "
               "more items (about half to one and a half times as many as here); a filled "
               "page is fit-checked like any page.")
    fence = "```"
    return "\n".join([
        "## 0. Page values (this page is designed once, then copied by script)",
        "",
        why,
        "",
        "Put this block at the top of each option script, right after the three "
        "comment lines, exactly as written (same names, same markers, these values):",
        "",
        fence + "python",
        render_block(vals).rstrip("\n"),
        fence,
        "",
        "- Every word that differs from one copy to the next comes from these names. "
        "Never type those words anywhere else in the script.",
        "- Put nothing else between the two marker lines, and keep each value's kind "
        "(a number stays a number, a list stays a list). Format in the code, not in the "
        "block.",
        "- Draw lists in a loop over the list (rows, sections, points).",
        f"- {fit}",
        "",
        "---",
        "",
    ])


def copy_prompt_banner(n: int, pc: dict) -> str:
    kind = pc["kind"]
    if pc.get("master") == "template":
        src = f"the {LABEL[kind]} design kept for this template ({pc.get('master_file')})"
        comment = ("A comment on this page is applied to it in edit mode, like any picked "
                   "page.")
    else:
        src = f"slide {pc['master']}'s design"
        comment = (f"A comment on this page goes to slide {pc['master']} in edit mode, "
                   "and every copy is made again from it.")
    return "\n".join([
        f"# NO WORKER FOR THIS SLIDE: slide {n} is a copy of {src}",
        "",
        f"Do not design this slide. `scripts/clone_pages.py` (finalize runs it) writes "
        f"its option file from {src}, with this page's own words from the brief. It is "
        "rendered, fit-checked and read like any page. " + comment,
        "If you are a slide-builder-worker dispatched on this file: write nothing and "
        "report \"page copy, no work\".",
        "",
        "---",
        "",
    ])


def pinned_prompt_banner(n: int, pc: dict) -> str:
    return "\n".join([
        f"# EDIT MODE ONLY: slide {n} copies {Path(pc['source']).name} slide "
        f"{pc['source_slide']}",
        "",
        "Prep copied the supplied page onto this slide as is (option_A.py). Do not design "
        f"options. Work from `{EDIT_REQUEST}` in this folder: change only the text the brief "
        "changes, in option_A.py's TEXT_CHANGES block.",
        "",
        "---",
        "",
    ])


def _page_text(source: str, k: int) -> list[str]:
    from pptx import Presentation
    lines = []
    try:
        s = Presentation(source).slides[k - 1]
    except Exception as exc:
        return [f"(could not read the page: {type(exc).__name__}: {exc})"]

    def walk(shapes):
        for sh in shapes:
            if getattr(sh, "shape_type", None) == 6 and hasattr(sh, "shapes"):
                yield from walk(sh.shapes)
            else:
                yield sh
    for sh in walk(s.shapes):
        if getattr(sh, "has_text_frame", False):
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    lines.append(f"- [{sh.name}] {t}")
        if getattr(sh, "has_table", False):
            for row in sh.table.rows:
                for c in row.cells:
                    t = c.text.strip()
                    if t:
                        lines.append(f"- [{sh.name} table] {t}")
    return lines


def write_pinned_copy(sd: Path, n: int, pc: dict) -> list[Path]:
    """option_A.py (the supplied page copied as is) and _edit_request.md."""
    src, k = pc["source"], int(pc["source_slide"])
    name = Path(src).name
    script = "\n".join([
        f"# Slide {n} option A - pattern: copied page ({name} slide {k})",
        "# Variant: the supplied page copied as is; only the text the brief changes is rewritten",
        "# Brief fidelity check: the copied page's own words, changed only where TEXT_CHANGES says",
        f"{PENDING_MARK} the worker in EDIT MODE fills TEXT_CHANGES ({EDIT_REQUEST}), "
        "then deletes this line",
        "import sys",
        "from pathlib import Path",
        f"sys.path.insert(0, r\"{SKILL}\")",
        f"sys.path.insert(0, r\"{HERE}\")",
        "from clone_pages import build_copied_page",
        "",
        f"SOURCE_DECK = r\"{src}\"",
        f"SOURCE_SLIDE = {k}",
        "",
        "# --- TEXT CHANGES (old text exactly as on the copied page -> new text from the brief) ---",
        "TEXT_CHANGES = {",
        "}",
        CHANGES_END,
        "",
        "",
        "def build():",
        "    return build_copied_page(SOURCE_DECK, SOURCE_SLIDE, TEXT_CHANGES)",
        "",
        "",
        'if __name__ == "__main__":',
        "    prs = build()",
        '    prs.save(str(Path(__file__).resolve().parent / "option_A.pptx"))',
        "",
    ])
    f = _design_file(sd, "A")
    f.write_text(script, encoding="utf-8")
    req = sd / EDIT_REQUEST
    req.write_text("\n".join([
        f"# Edit request: {_sid(n)}, option A (COPIED PAGE, EDIT IN PLACE)",
        "",
        f"Prep copied slide {k} of {src} onto this page as is: {f} copies its shapes, "
        "pictures and charts. Change ONLY the text the brief changes:",
        "",
        "1. Read _context.md and _prompt.md for the brief of this page.",
        "2. In option_A.py, fill TEXT_CHANGES: each entry is \"old text exactly as on the "
        "copied page\": \"new text from the brief\". The old text is a whole paragraph or a "
        "piece of one (a run). An old text that is not on the page stops the build, so copy "
        "it from the list below.",
        "3. Do not add, move, resize, restyle or remove shapes. If the brief needs a new "
        "element or a different layout, stop and say so: that page needs a design, not a copy.",
        f"4. Delete the line starting \"{PENDING_MARK}\", run the script, then "
        "scripts/check_fit.py on the option_A.pptx it saved.",
        "5. Rename this file to _edit_request.applied.md.",
        "",
        "The text on the copied page now:",
    ] + _page_text(src, k)) + "\n", encoding="utf-8")
    return [f, req]


def build_copied_page(source: str, slide_no: int, changes: dict):
    """Run by a copied page's option_A.py: the supplied slide's shapes on a
    fresh slide, with their pictures and charts, then the text changes."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from twins.helpers import new_slide
    from twins.composer import copy_shape_with_parts
    src = Presentation(source)
    if not 1 <= slide_no <= len(src.slides):
        raise ValueError(f"{Path(source).name} has {len(src.slides)} slide(s), not {slide_no}")
    s = src.slides[slide_no - 1]
    prs, slide = new_slide()
    if (src.slide_width, src.slide_height) != (prs.slide_width, prs.slide_height):
        sys.stderr.write(f"  WARN: {Path(source).name} pages are a different size from "
                         "this deck's; the copy keeps the shapes where they were. Check "
                         "the picture.\n")
    tree = slide.shapes._spTree
    memo: dict = {}
    for sh in s.shapes:
        el = copy_shape_with_parts(sh.element, s.part, slide.part, memo)
        ph = el.find(".//" + qn("p:ph"))
        if ph is not None:
            kind = (ph.get("type") or "body").lower()
            if kind in ("dt", "ftr", "sldnum"):
                continue                       # the layout draws its own
            if not (getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()) \
                    and not getattr(sh, "has_table", False) and kind not in ("pic",):
                continue                       # an empty placeholder
            cnv = el.find(".//" + qn("p:cNvPr"))
            if kind in ("title", "ctrtitle"):
                if cnv is not None:
                    cnv.set("name", "title")
            elif kind == "subtitle":
                if cnv is not None:
                    cnv.set("name", "subtitle")
            else:
                # A body placeholder becomes a plain shape where it was drawn,
                # so finishing keeps its content.
                _materialize(el, sh)
                ph.getparent().remove(ph)
        tree.append(el)
    _apply_changes(slide, changes or {})
    return prs


def _materialize(el, sh) -> None:
    from pptx.oxml.ns import qn
    sppr = el.find(qn("p:spPr"))
    if sppr is None or sppr.find(qn("a:xfrm")) is not None:
        return
    from lxml import etree
    xfrm = etree.SubElement(sppr, qn("a:xfrm"))
    sppr.remove(xfrm)
    sppr.insert(0, xfrm)
    off = etree.SubElement(xfrm, qn("a:off"))
    off.set("x", str(int(sh.left or 0)))
    off.set("y", str(int(sh.top or 0)))
    ext = etree.SubElement(xfrm, qn("a:ext"))
    ext.set("cx", str(int(sh.width or 0)))
    ext.set("cy", str(int(sh.height or 0)))


def _apply_changes(slide, changes: dict) -> None:
    left = dict(changes)

    def frames(shapes):
        for sh in shapes:
            if getattr(sh, "shape_type", None) == 6 and hasattr(sh, "shapes"):
                yield from frames(sh.shapes)
                continue
            if getattr(sh, "has_text_frame", False):
                yield sh.text_frame
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for c in row.cells:
                        yield c.text_frame
    for tf in frames(slide.shapes):
        for p in tf.paragraphs:
            whole = "".join(r.text for r in p.runs)
            if whole.strip() in changes and p.runs:
                p.runs[0].text = changes[whole.strip()]
                for r in p.runs[1:]:
                    r.text = ""
                left.pop(whole.strip(), None)
                continue
            for r in p.runs:
                for old, new in changes.items():
                    if old and old in r.text:
                        r.text = r.text.replace(old, new)
                        left.pop(old, None)
    if left:
        raise ValueError("TEXT_CHANGES names text that is not on the copied page: "
                         + "; ".join(repr(k) for k in list(left)[:5])
                         + ". Copy the old text exactly from _edit_request.md.")


def prep_lines(copy_plan: dict, out: Path, template) -> list[str]:
    """What prep says about these pages (also written into dispatch_plan.md)."""
    if not copy_plan:
        return []
    lines = ["", "Pages made by script (owner's decision 2026-10-14; no designer for copies):"]
    no_worker = []
    for n in sorted(copy_plan):
        pc = copy_plan[n]
        kind, role = pc["kind"], pc["role"]
        if role == "master" and kind == DIVIDER:
            lines.append(f"  slide {n}: the first section divider. Its designer keeps the page "
                         f"values at the top of the file (_prompt.md section 0); slides "
                         f"{_nums(pc['copies'])} are copies of it.")
        elif role == "master":
            lines.append(f"  slide {n}: the {LABEL[kind]} is designed here once (its _prompt.md "
                         "section 0); when the deck compiles, the pick is kept for this "
                         "template and later decks fill it by script.")
        elif role == "copy" and pc.get("master") == "template":
            no_worker.append(n)
            lines.append(f"  slide {n}: the {LABEL[kind]} is filled from the design kept for "
                         f"this template ({pc['master_file']}), word for word from the brief. "
                         "To design it fresh instead: py -3 scripts/build_deck.py --slide "
                         f"{n} --out \"{out}\" --template \"{template}\" (a rebuild always "
                         "sends a designer).")
        elif role == "copy":
            no_worker.append(n)
        elif role == "pinned":
            lines.append(f"  slide {n}: copies {Path(pc['source']).name} slide "
                         f"{pc['source_slide']} as is. Send its worker in EDIT MODE on "
                         f"{_p.slide_dir(out, n) / EDIT_REQUEST} (it changes only the text "
                         "the brief changes), not a designer.")
    if no_worker:
        lines.append(f"  Do NOT send a worker for slide(s) {_nums(no_worker)}: finalize_deck.py "
                     "writes them (clone_pages.py), and they are rendered, fit-checked, "
                     "reviewed and QC-read like any page.")
    return lines


# ---------------------------------------------------------------- CLI
def _cmd_refresh(args) -> int:
    rep = refresh(args.out, only=args.slide)
    meta = _meta(args.out)
    for ln in report_lines(rep, args.out, meta.get("template"), args.slide):
        print(ln)
    probs = problem_lines(rep, args.out)
    for ln in probs:
        print(ln)
    if not (rep["written"] or rep["waiting"] or rep["notes"] or probs or rep["same"]):
        print("  no copied pages in this build")
    return 1 if rep["errors"] else 0


def _cmd_status(args) -> int:
    meta = _meta(args.out)
    any_ = False
    for s in meta.get("slides") or []:
        pc = s.get("page_copy")
        if not pc:
            continue
        any_ = True
        m = pc.get("master")
        how = {"master": "designed here once", "pinned":
               f"copy of {Path(pc.get('source', '')).name} slide {pc.get('source_slide')}",
               "copy": ("filled from the design kept for this template" if m == "template"
                        else f"copy of slide {m}")}.get(pc.get("role"), pc.get("role"))
        print(f"  slide {s['n']}: {LABEL.get(pc.get('kind'), pc.get('kind'))}, {how}")
    if not any_:
        print("  no copied pages in this build")
    return 0


def _cmd_save(args) -> int:
    lines = save_template_masters(args.out)
    for ln in lines or ["  nothing new to keep for this template"]:
        print(ln)
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Pages designed once and copied by script.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("refresh", help="Make every copied page current (finalize runs it).")
    r.add_argument("--out", required=True, type=Path)
    r.add_argument("--slide", type=int, default=None)
    s = sub.add_parser("status", help="Which pages are copies, and of what.")
    s.add_argument("--out", required=True, type=Path)
    k = sub.add_parser("save-masters", help="Keep this deck's picked cover, contents and "
                                            "executive summary for its template.")
    k.add_argument("--out", required=True, type=Path)
    args = ap.parse_args(argv)
    if not _p.meta_json(args.out).exists():
        print(f"ERROR: no _meta.json in {args.out}; run build_deck.py first.")
        return 2
    return {"refresh": _cmd_refresh, "status": _cmd_status, "save-masters": _cmd_save}[args.cmd](args)


if __name__ == "__main__":
    sys.exit(main())
