#!/usr/bin/env python3
"""chart_rules.py: chart header, legend place and no final periods
(three owner-approved rules, 2026-10-09).

  1. Legend: every chart's legend is one row at the top right of the chart,
     on the chart's header line, right-aligned to the chart's right edge. A
     chart with one series has no legend.
  2. Chart header: every chart has a header line "Chart title, [Unit]" at the
     chart's top left: the title bold, the unit regular in square brackets
     after a comma, an optional superscript footnote mark after the title.
     Axis titles where an axis carries a unit.
  3. No final periods: no text item ends with a period (titles, subtitles,
     takeaways, labels, bullets, panels, callouts, table cells, chart labels,
     footnotes, sources). Periods BETWEEN sentences stay.

What this file does:

  strip_final_periods(slide)  finalize calls it on every option after the
      slide is built and before it is rendered: removes one final "." from
      the end of every paragraph in every text box, group and table cell, and
      returns one record per change (slide, shape, before / after tail).
      Kept: an ellipsis ("..." or the one-character ellipsis), an abbreviation
      (Inc., Ltd., Co., Corp., etc., e.g., i.e., vs., No., St., U.S., Jr.,
      Sr., Mr., Mrs., Ms., Dr.; any single capital letter such as "J."; any
      dotted abbreviation such as "U.K." or "a.m."), and a list number on its
      own ("3."). Text inside native chart objects is not touched.
  leftover_periods(slide)     the same test without changing anything.
  chart_problems(slide)       charts without a header line, without a unit
      (when they have a value axis), or with a legend that is not one row at
      the top right.
  deck_findings(prs)          slide-qc's hygiene findings for all three
      (layout Majors, fixed automatically; keys final_period, chart_header,
      legend_position in qc_fix_policy.py).

How a drawn chart is found (Slide Lab draws charts from shapes): by shape
name. Every chart piece is named chart-...: marks chart-bar-..., chart-line-
..., chart-dot-...; value-axis ticks chart-ylab-... (or chart-ytick-...,
chart-axis-y0 ...); gridlines chart-grid-...; the header chart-title (or
chart-title-<id> when a slide has two charts; chart-header is accepted too);
the legend chart-legend (or chart-legend-<series> plus chart-legend-swatch-
<series>); axis titles chart-axis-title-x / chart-axis-title-y. A slide
whose pieces are not named this way (an older or external deck) has no chart
this check can see, so it is not flagged: the vision pass still looks.
Pieces are grouped into charts by position (pieces within about a third of
an inch of each other belong to the same chart), so two charts side by side
are checked one by one.

Native PowerPoint charts (an external deck) are read from the chart object:
the header is the chart's own title or a text box just above it; the unit is
in the header, the value-axis title, or a $ or % in the tick-label format;
the legend must be at the top or top right (position "top" or "corner", or a
manual layout in the top-right quarter). A legend at the right middle is
flagged: it is not on the header line.

Run:  py -3 chart_rules.py <deck.pptx>      prints the findings per slide
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

EMU_PER_IN = 914400

# ---------------------------------------------------------------------------
# Rule 3: final periods
# ---------------------------------------------------------------------------
ABBREVIATIONS = frozenset({
    "inc.", "ltd.", "co.", "corp.", "etc.", "e.g.", "i.e.", "vs.", "no.", "st.",
    "u.s.", "jr.", "sr.", "mr.", "mrs.", "ms.", "dr.",
})
_SINGLE_CAP = re.compile(r"^[A-Z]\.$")
_DOTTED_ABBR = re.compile(r"^(?:[A-Za-z]\.){2,}$")
_LIST_NUMBER = re.compile(r"^\(?\d{1,3}\.$")
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def keeps_final_period(text: str) -> bool:
    """True when `text` ends in a period that must stay (or no period)."""
    t = (text or "").rstrip()
    if not t.endswith("."):
        return True
    if t.endswith("..") or t.endswith("…."):
        return True
    last = t.split()[-1] if t.split() else t
    # "(Inc." or "Smith & Co." : look at the word itself
    word = last.lstrip("([{\"'“‘")
    if word.lower() in ABBREVIATIONS:
        return True
    if _SINGLE_CAP.match(word) or _DOTTED_ABBR.match(word):
        return True
    if _LIST_NUMBER.match(t):
        return True
    return False


def ends_with_final_period(text: str) -> bool:
    t = (text or "").rstrip()
    return t.endswith(".") and not keeps_final_period(t)


def _para_parts(p):
    """The a:r and a:fld children of a paragraph, in order, as (element,
    text). a:fld is a field (slide number, date): text PowerPoint fills in."""
    out = []
    for child in p._p:
        tag = child.tag
        if tag in (_A + "r", _A + "fld"):
            t = child.find(_A + "t")
            out.append((child, (t.text or "") if t is not None else ""))
    return out


def _para_text(p) -> str:
    return "".join(t for _, t in _para_parts(p))


def _tail(text: str, n: int = 40) -> str:
    t = (text or "").replace("\v", " ").replace("\n", " ").rstrip()
    return t if len(t) <= n else "..." + t[-n:]


def iter_text_frames(shapes, prefix: str = ""):
    """(shape name, text frame) for every text box, every shape inside a
    group, and every table cell. Native charts are skipped."""
    for sh in shapes:
        name = prefix + (getattr(sh, "name", "") or "")
        try:
            if sh.shape_type == 6:  # group
                yield from iter_text_frames(sh.shapes, prefix=name + "/")
                continue
        except Exception:
            pass
        if getattr(sh, "has_table", False):
            try:
                for ri, row in enumerate(sh.table.rows):
                    for ci, cell in enumerate(row.cells):
                        yield f"{name} cell r{ri + 1}c{ci + 1}", cell.text_frame
            except Exception:
                pass
            continue
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            yield name, sh.text_frame


def _scan(slide, fix: bool, slide_n=None) -> list[dict]:
    changes: list[dict] = []
    for name, tf in iter_text_frames(slide.shapes):
        for p in tf.paragraphs:
            parts = _para_parts(p)
            full = "".join(t for _, t in parts)
            if not ends_with_final_period(full):
                continue
            last = next(((el, t) for el, t in reversed(parts) if t.strip()), None)
            if last is None or last[0].tag != _A + "r":
                continue  # ends in a field: PowerPoint's text, leave it
            el, t = last
            cut = len(t.rstrip()) - 1
            if cut < 0 or t[cut] != ".":
                continue
            new_t = t[:cut] + t[cut + 1:]
            after = full[: len(full) - len(t)] + new_t
            rec = {"slide": slide_n, "shape": name or "<unnamed>",
                   "before": _tail(full), "after": _tail(after)}
            if fix:
                el.find(_A + "t").text = new_t
            changes.append(rec)
    return changes


def strip_final_periods(slide, slide_n=None) -> list[dict]:
    """Remove one final period from every paragraph that ends with one (the
    kept cases aside). Returns one record per change."""
    return _scan(slide, fix=True, slide_n=slide_n)


def leftover_periods(slide, slide_n=None) -> list[dict]:
    """The paragraphs strip_final_periods would change; changes nothing."""
    return _scan(slide, fix=False, slide_n=slide_n)


# ---------------------------------------------------------------------------
# Rules 1 and 2: chart header and legend
# ---------------------------------------------------------------------------
_TOKEN_SPLIT = re.compile(r"[-_ ]+")
LEGEND_TOKENS = {"legend", "leg", "key"}
HEADER_TOKENS = {"title", "header", "heading"}
AXIS_SIDE_TOKENS = {"x", "y", "xaxis", "yaxis"}
MARK_TOKENS = {"bar", "bars", "col", "column", "columns", "line", "dot", "dots", "point",
               "pt", "bubble", "area", "seg", "segment", "series", "slice", "wedge",
               "marker", "stack"}
GRID_TOKENS = {"grid", "gridline", "gridlines"}
PANEL_TOKENS = {"panel", "bg", "background", "card", "frame", "divider", "note", "notes",
                "callout", "box", "subhead", "caption"}
VALUE_TICK_RE = re.compile(
    r"(^|[-_ ])(ylab|ylabel|ytick|yticks|yval|y[-_]?tick)([-_ \d]|$)|(^|[-_ ])axis[-_ ]?y[-_ ]?\d", re.I)
UNIT_TOKENS = {"unit", "units"}
# A unit in the header: "[$B]", "($B)", "$", "%", or ", <unit>" after the title
UNIT_IN_TEXT = re.compile(r"\[[^\]]+\]|\([^)]+\)|[$%€£¥]|,\s*\S")
HEADER_TEXT = re.compile(r"^[^\n]{2,160},\s*\[[^\]]+\]\s*$")
GROUP_GAP_IN = 0.35
HEADER_REACH_IN = 3.5
ONE_ROW_TOL_IN = 0.15


def _tokens(name: str) -> list[str]:
    return [t for t in _TOKEN_SPLIT.split((name or "").strip().lower()) if t]


def _role(name: str) -> str:
    """header | axis_title | unit | legend | tick | mark | grid | other | ''
    ('' = not a chart piece)."""
    low = (name or "").strip().lower()
    toks = _tokens(low)
    if not toks:
        return ""
    chartish = toks[0] in ("chart", "charts")
    if any(t in LEGEND_TOKENS for t in toks) and (chartish or toks[0] in LEGEND_TOKENS):
        return "legend"
    if not chartish:
        return ""
    if any(t in HEADER_TOKENS for t in toks):
        return "axis_title" if any(t in AXIS_SIDE_TOKENS for t in toks) else "header"
    if any(t in UNIT_TOKENS for t in toks):
        return "unit"
    if VALUE_TICK_RE.search(low):
        return "tick"
    if "label" in toks or "labels" in toks:
        return "other"   # chart-grid-label, chart-map-label: text, not marks
    if any(t in MARK_TOKENS for t in toks):
        return "mark"
    if any(t in GRID_TOKENS for t in toks):
        return "grid"
    return "other"


def _walk(shapes, off=(0, 0, 1.0, 1.0)):
    """Every shape with its slide-space box (x0, y0, x1, y1) in EMU. Group
    children are mapped through the group's child coordinate space."""
    ox, oy, sx, sy = off
    for sh in shapes:
        try:
            x0 = ox + (sh.left or 0) * sx
            y0 = oy + (sh.top or 0) * sy
            x1 = x0 + (sh.width or 0) * sx
            y1 = y0 + (sh.height or 0) * sy
        except Exception:
            continue
        try:
            is_group = sh.shape_type == 6
        except Exception:
            is_group = False
        if is_group:
            try:
                xfrm = sh._element.grpSpPr.find(_A + "xfrm")
                ch_off = xfrm.find(_A + "chOff")
                ch_ext = xfrm.find(_A + "chExt")
                cx, cy = int(ch_off.get("x")), int(ch_off.get("y"))
                cw, chh = int(ch_ext.get("cx")) or 1, int(ch_ext.get("cy")) or 1
                gsx = (x1 - x0) / cw if cw else sx
                gsy = (y1 - y0) / chh if chh else sy
                yield from _walk(sh.shapes, (x0 - cx * gsx, y0 - cy * gsy, gsx, gsy))
            except Exception:
                yield from _walk(sh.shapes, off)
            continue
        yield sh, (x0, y0, x1, y1)


def _text_of(sh) -> str:
    try:
        return sh.text_frame.text.strip() if sh.has_text_frame else ""
    except Exception:
        return ""


def _union(boxes):
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def _near(a, b, gx, gy=None):
    gy = gx if gy is None else gy
    return not (a[2] + gx < b[0] or b[2] + gx < a[0] or a[3] + gy < b[1] or b[3] + gy < a[1])


def _pieces_near(p, q, gap) -> bool:
    """Do two chart pieces belong to the same chart? Marks and gridlines:
    within `gap`. Value-axis ticks stand in a column, often more than `gap`
    apart: two ticks join when they share a column and are within an inch of
    each other; a tick joins the marks it stands beside (up to an inch to
    their left or right, overlapping them in height)."""
    ticks = (p[2] == "tick") + (q[2] == "tick")
    if ticks == 2:
        return _near(p[1], q[1], 0.05 * EMU_PER_IN, 1.0 * EMU_PER_IN)
    if ticks == 1:
        return _near(p[1], q[1], 1.0 * EMU_PER_IN, gap)
    return _near(p[1], q[1], gap)


def _cluster(items, gap):
    """Group (shape, box, role) items that belong to the same chart."""
    groups: list[list] = []
    for it in items:
        hit = [g for g in groups if any(_pieces_near(it, o, gap) for o in g)]
        merged = [it]
        for g in hit:
            merged.extend(g)
            groups.remove(g)
        groups.append(merged)
    return groups


def _is_title_placeholder(sh) -> bool:
    try:
        return bool(sh.is_placeholder and sh.placeholder_format.type in (1, 3, 4))
    except Exception:
        return False


def _drawn_chart_problems(slide) -> list[dict]:
    pieces = []
    for sh, box in _walk(slide.shapes):
        if getattr(sh, "has_chart", False) and sh.has_chart:
            continue
        role = _role(getattr(sh, "name", "") or "")
        if role:
            pieces.append((sh, box, role))
    core = [p for p in pieces if p[2] in ("mark", "grid", "tick")]
    if not core:
        return []
    gap = GROUP_GAP_IN * EMU_PER_IN
    charts = []
    for g in _cluster(core, gap):
        n_marks = sum(1 for p in g if p[2] == "mark")
        n_ticks = sum(1 for p in g if p[2] == "tick")
        if n_marks >= 2 or n_ticks >= 3:
            charts.append(g)
    if not charts:
        return []
    out: list[dict] = []
    legends = [p for p in pieces if p[2] == "legend"]
    headers = [p for p in pieces if p[2] == "header"]
    units = [p for p in pieces if p[2] in ("unit", "axis_title")]
    # fallback header: a text box written as "Title, [Unit]"
    for sh, box in _walk(slide.shapes):
        if (not _role(getattr(sh, "name", "") or "") and not _is_title_placeholder(sh)
                and HEADER_TEXT.match(_text_of(sh))):
            headers.append((sh, box, "header"))

    def _dist(box, cb):
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        dx = max(cb[0] - cx, 0, cx - cb[2])
        dy = max(cb[1] - cy, 0, cy - cb[3])
        return (dx * dx + dy * dy) ** 0.5

    # Widen each chart sideways with its other named pieces in the same band
    # (data labels, line-end labels), so a chart drawn with only its value
    # ticks named still gets its real width. Only sideways: growing up or
    # down would swallow the chart above or below. Panels and notes are left
    # out: a background panel would swallow everything.
    others = [p for p in pieces if p[2] == "other"
              and not any(t in PANEL_TOKENS for t in _tokens(getattr(p[0], "name", "") or ""))]
    chart_boxes = []
    for g in charts:
        core_box = _union([p[1] for p in g])
        box = core_box
        for _ in range(4):
            add = [o[1] for o in others
                   if o[1][1] >= core_box[1] - gap and o[1][3] <= core_box[3] + gap
                   and _near(o[1], box, gap)]
            new = _union([box] + add) if add else box
            new = (new[0], core_box[1], new[2], core_box[3])
            if new == box:
                break
            box = new
        chart_boxes.append((box, core_box, g))
    # each legend piece goes to the nearest chart (on a one-chart slide, every
    # legend piece is that chart's)
    legend_of = {i: [] for i in range(len(chart_boxes))}
    for lp in legends:
        best = min(range(len(chart_boxes)), key=lambda i: _dist(lp[1], chart_boxes[i][0]))
        if (len(chart_boxes) == 1
                or _dist(lp[1], chart_boxes[best][0]) <= HEADER_REACH_IN * EMU_PER_IN):
            legend_of[best].append(lp)
    reach = HEADER_REACH_IN * EMU_PER_IN
    for i, (cb, plot, g) in enumerate(chart_boxes):
        label = (f"chart near {cb[0] / EMU_PER_IN:.1f} in, {cb[1] / EMU_PER_IN:.1f} in"
                 if len(chart_boxes) > 1 else "the chart")
        mid_y = (plot[1] + plot[3]) / 2

        def _above_and_over(box):
            overlap = min(box[2], cb[2]) - max(box[0], cb[0])
            cy = (box[1] + box[3]) / 2
            return overlap > 0 and cy < mid_y and box[3] >= plot[1] - reach

        hdr = [h for h in headers if _above_and_over(h[1])]
        hdr_text = " ".join(_text_of(h[0]) for h in hdr)
        if not hdr:
            out.append({"kind": "header", "chart": label,
                        "detail": f"{label} has no header line (a shape named chart-title "
                                  f"reading \"Chart title, [Unit]\" at its top left)"})
        has_value_axis = sum(1 for p in g if p[2] == "tick") >= 2
        if has_value_axis:
            unit_ok = bool(hdr_text and UNIT_IN_TEXT.search(hdr_text))
            if not unit_ok:
                unit_ok = any(_text_of(u[0]) and _dist(u[1], cb) <= reach for u in units)
            if not unit_ok:
                ticks = [_text_of(p[0]) for p in g if p[2] == "tick"]
                unit_ok = any(re.search(r"[$%€£¥]|\d\s*[KMBkmb]\b", t) for t in ticks)
            if not unit_ok:
                out.append({"kind": "unit", "chart": label,
                            "detail": f"{label} has a value axis but no unit: write it in the "
                                      f"header as \", [Unit]\" or in the axis title"})
        lg = legend_of[i]
        if lg:
            lb = _union([p[1] for p in lg])
            lcx, lcy = (lb[0] + lb[2]) / 2, (lb[1] + lb[3]) / 2
            mid_x = (cb[0] + cb[2]) / 2
            if not (lcy < mid_y and lcx > mid_x):
                out.append({"kind": "legend", "chart": label,
                            "detail": f"{label}: the legend is not at the top right (it sits "
                                      f"{'below' if lcy >= mid_y else 'above'} the middle and "
                                      f"{'left of' if lcx <= mid_x else 'right of'} the middle; "
                                      f"put it on the header line, right-aligned)"})
            else:
                texted = [p for p in lg if _text_of(p[0])]
                ys = [(p[1][1] + p[1][3]) / 2 for p in texted]
                if len(ys) >= 2 and max(ys) - min(ys) > ONE_ROW_TOL_IN * EMU_PER_IN:
                    out.append({"kind": "legend", "chart": label,
                                "detail": f"{label}: the legend is stacked, not one row; put the "
                                          f"entries side by side on the header line"})
    return out


def _native_chart_problems(slide) -> list[dict]:
    from pptx.enum.chart import XL_LEGEND_POSITION
    out: list[dict] = []
    frames = []
    texts = []
    for sh, box in _walk(slide.shapes):
        if getattr(sh, "has_chart", False) and sh.has_chart:
            frames.append((sh, box))
        elif _text_of(sh) and not _is_title_placeholder(sh):
            texts.append((sh, box))
    for sh, box in frames:
        label = f"chart \"{sh.name}\""
        try:
            chart = sh.chart
        except Exception:
            continue
        title_text = ""
        try:
            if chart.has_title:
                title_text = chart.chart_title.text_frame.text.strip()
        except Exception:
            title_text = ""
        if not title_text:
            h = box[3] - box[1]
            for tsh, tb in texts:
                overlap = min(tb[2], box[2]) - max(tb[0], box[0])
                if (overlap > 0 and tb[3] <= box[1] + 0.25 * h
                        and tb[3] >= box[1] - 1.2 * EMU_PER_IN
                        and (tb[1] + tb[3]) / 2 < box[1] + 0.25 * h):
                    if _role(getattr(tsh, "name", "") or "") == "legend":
                        continue
                    title_text = _text_of(tsh)
                    break
        if not title_text:
            out.append({"kind": "header", "chart": label,
                        "detail": f"{label} has no header line (a title \"Chart title, [Unit]\" "
                                  f"at its top left)"})
        value_axis = None
        try:
            value_axis = chart.value_axis
        except Exception:
            value_axis = None   # pie / doughnut: no value axis
        if value_axis is not None:
            unit_ok = bool(title_text and UNIT_IN_TEXT.search(title_text))
            if not unit_ok:
                try:
                    unit_ok = bool(value_axis.has_title
                                   and value_axis.axis_title.text_frame.text.strip())
                except Exception:
                    pass
            if not unit_ok:
                try:
                    fmt = value_axis.tick_labels.number_format or ""
                    unit_ok = bool(re.search(r"[$%€£¥]", fmt))
                except Exception:
                    pass
            if not unit_ok:
                out.append({"kind": "unit", "chart": label,
                            "detail": f"{label} has a value axis but no unit: write it in the "
                                      f"header as \", [Unit]\" or in the axis title"})
        try:
            has_legend = chart.has_legend
        except Exception:
            has_legend = False
        if has_legend:
            ok = False
            try:
                leg = chart._chartSpace.chart.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/chart}legend")
                ml = leg.find(".//{http://schemas.openxmlformats.org/drawingml/2006/chart}manualLayout")
                if ml is not None:
                    c = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
                    x = float(ml.find(c + "x").get("val"))
                    y = float(ml.find(c + "y").get("val"))
                    w = ml.find(c + "w")
                    hh = ml.find(c + "h")
                    cx = x + (float(w.get("val")) / 2 if w is not None else 0)
                    cy = y + (float(hh.get("val")) / 2 if hh is not None else 0)
                    ok = cx > 0.5 and cy < 0.5
                else:
                    ok = chart.legend.position in (XL_LEGEND_POSITION.TOP,
                                                   XL_LEGEND_POSITION.CORNER)
            except Exception:
                ok = True   # cannot tell: do not flag
            if not ok:
                try:
                    pos = str(chart.legend.position).split(".")[-1].split(" ")[0].lower()
                except Exception:
                    pos = "?"
                out.append({"kind": "legend", "chart": label,
                            "detail": f"{label}: the legend is at the {pos}, not one row at the "
                                      f"top right on the header line"})
    return out


def chart_problems(slide) -> list[dict]:
    """[{kind: header | unit | legend, chart, detail}] for the slide."""
    out: list[dict] = []
    try:
        out += _drawn_chart_problems(slide)
    except Exception:
        pass
    try:
        out += _native_chart_problems(slide)
    except Exception:
        pass
    return out


# ---------------------------------------------------------------------------
# slide-qc findings
# ---------------------------------------------------------------------------
def slide_findings(slide, slide_num: int) -> list[dict]:
    out: list[dict] = []
    left = leftover_periods(slide, slide_n=slide_num)
    if left:
        shown = "; ".join(f"{c['shape']}: '{c['before']}'" for c in left[:3])
        more = "" if len(left) <= 3 else f" (+{len(left) - 3} more)"
        out.append({
            "slide": slide_num, "severity": "Major", "category": "final_period",
            "auto_fix": True, "count": len(left),
            "issue": (f"Slide {slide_num}: {len(left)} text item(s) end with a final period "
                      f"({shown}{more}). Remove the final period; periods between "
                      f"sentences stay."),
        })
    for p in chart_problems(slide):
        cat = "legend_position" if p["kind"] == "legend" else "chart_header"
        out.append({
            "slide": slide_num, "severity": "Major", "category": cat, "auto_fix": True,
            "kind": p["kind"],
            "issue": f"Slide {slide_num}: {p['detail']}.",
        })
    return out


def deck_findings(prs) -> list[dict]:
    out: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        out += slide_findings(slide, i)
    return out


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if not argv:
        print(__doc__)
        return 2
    from pptx import Presentation
    prs = Presentation(str(Path(argv[0])))
    findings = deck_findings(prs)
    for f in findings:
        print(f"[{f['category']}] {f['issue']}")
    counts: dict[str, int] = {}
    for f in findings:
        counts[f["category"]] = counts.get(f["category"], 0) + 1
    n_periods = sum(f.get("count", 0) for f in findings if f["category"] == "final_period")
    print(f"\nfinal_period: {counts.get('final_period', 0)} slide(s), {n_periods} text item(s); "
          f"chart_header: {counts.get('chart_header', 0)}; "
          f"legend_position: {counts.get('legend_position', 0)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
