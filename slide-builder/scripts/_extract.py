"""_extract.py — one full-fidelity walk over a slide's visible text.

Both adopt_deck (option 6b) and refresh_deck (option 6c) used to walk only
`slide.shapes` and keep shapes where `has_text_frame` was true. That silently
dropped:

  * every TABLE cell        (a GraphicFrame has no text frame)
  * every shape inside a GROUP (the walk never recursed)
  * every chart category and series value

Measured on a 5-shape probe slide carrying 8 figures, the old walk recovered 2
of 8. This walk recovers 8. In practice that is the difference between
"refresh my PMO deck" seeing the numbers in your table and silently ignoring
them, and between a replicated page's figures being reviewable or invisible.

Anything genuinely unreadable from the slide XML (pictures, SmartArt, OLE and
think-cell objects) is NOT skipped silently. It is returned in `unreachable` so
callers can state how large the unchecked surface is instead of assuming it is
empty. Silence has been the recurring failure mode in this pipeline; an explicit
"4 surfaces could not be read" is the whole point of that second list.

Addresses are stable strings so they can live in a JSON spec/ledger:
    s3/sh2            top-level text frame
    s3/sh2/g0/g1      nested inside group(s)
    s3/sh4/r1c2       table cell, row 1 col 2
    s3/sh5/cat0       chart category
    s3/sh5/ser0pt2    chart series point
"""
from __future__ import annotations

import re

from pptx.enum.shapes import MSO_SHAPE_TYPE

# A "figure" for ledger purposes: currency, percent, multiple, plain number,
# or a duration/quantity word that behaves like one. Deliberately broad on the
# read side; the ledger decides what to prompt on.
_NUMERAL_RE = re.compile(
    r"(\$\s?\d|\d+\s?%|\b\d+\s?[xX]\b|\b\d[\d,\.]*\b|"
    r"\b(?:sold out|none|nil|all|every|half|double|triple)\b)",
    re.IGNORECASE,
)


def has_figure(text: str) -> bool:
    """True when a surface carries something a reviewer would call a figure."""
    return bool(_NUMERAL_RE.search(text or ""))


def _txt(shape) -> str:
    try:
        return (shape.text_frame.text or "").strip()
    except Exception:
        return ""


def _name(shape) -> str:
    try:
        return shape.name or ""
    except Exception:
        return ""


def _walk(shapes, slide_no: int, prefix: str,
          surfaces: list, unreachable: list, top: bool = True) -> None:
    for idx, shape in enumerate(shapes):
        addr = f"{prefix}/sh{idx}" if top else f"{prefix}/g{idx}"
        try:
            st = shape.shape_type
        except Exception:
            st = None

        # 1. Groups: recurse. The old walk stopped here, so anything a designer
        #    grouped (which on a client one-pager is most of the callouts) was
        #    invisible. finalize_deck already recurses groups elsewhere.
        if st == MSO_SHAPE_TYPE.GROUP:
            try:
                _walk(shape.shapes, slide_no, addr, surfaces, unreachable, top=False)
                continue
            except Exception:
                unreachable.append({"slide": slide_no, "addr": addr, "kind": "group",
                                    "name": _name(shape), "why": "group could not be walked"})
                continue

        # 2. Tables: every cell is its own editable text frame.
        if getattr(shape, "has_table", False):
            try:
                for r, row in enumerate(shape.table.rows):
                    for c, cell in enumerate(row.cells):
                        t = (cell.text_frame.text or "").strip()
                        if t:
                            surfaces.append({
                                "slide": slide_no, "addr": f"{addr}/r{r}c{c}",
                                "kind": "table_cell", "name": _name(shape),
                                "text": t, "writable": True})
                continue
            except Exception:
                unreachable.append({"slide": slide_no, "addr": addr, "kind": "table",
                                    "name": _name(shape), "why": "table could not be read"})
                continue

        # 3. Charts: categories and plotted values are READ-ONLY here. The
        #    numbers live in an embedded workbook, not the slide part, so they
        #    are reportable but not safely rewritable from this walk.
        if getattr(shape, "has_chart", False):
            try:
                chart = shape.chart
                for pi, plot in enumerate(chart.plots):
                    for ci, cat in enumerate(plot.categories):
                        if str(cat).strip():
                            surfaces.append({
                                "slide": slide_no, "addr": f"{addr}/cat{ci}",
                                "kind": "chart_category", "name": _name(shape),
                                "text": str(cat).strip(), "writable": False})
                    for si, ser in enumerate(plot.series):
                        for vi, val in enumerate(ser.values):
                            if val is not None:
                                surfaces.append({
                                    "slide": slide_no, "addr": f"{addr}/ser{si}pt{vi}",
                                    "kind": "chart_value", "name": _name(shape),
                                    "text": str(val), "writable": False})
                continue
            except Exception:
                unreachable.append({"slide": slide_no, "addr": addr, "kind": "chart",
                                    "name": _name(shape), "why": "chart data could not be read"})
                continue

        # 4. Plain text frames.
        if getattr(shape, "has_text_frame", False):
            t = _txt(shape)
            if t:
                surfaces.append({"slide": slide_no, "addr": addr, "kind": "text",
                                 "name": _name(shape), "text": t, "writable": True})
            continue

        # 5. Everything else is text we CANNOT read from the XML. Report it
        #    rather than pretending the surface was checked.
        if st == MSO_SHAPE_TYPE.PICTURE:
            why = "picture: any numerals are baked into the image"
        elif st == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
            why = "embedded object (e.g. think-cell / Excel): not readable as text"
        else:
            why = "no readable text (SmartArt / media / placeholder graphic)"
        unreachable.append({"slide": slide_no, "addr": addr,
                            "kind": str(st), "name": _name(shape), "why": why})


def walk_slide(slide, slide_no: int = 1) -> tuple[list[dict], list[dict]]:
    """Return (surfaces, unreachable) for ONE slide."""
    surfaces: list[dict] = []
    unreachable: list[dict] = []
    _walk(slide.shapes, slide_no, f"s{slide_no}", surfaces, unreachable)
    return surfaces, unreachable


def resolve_text_frame(slide, addr: str):
    """Resolve an address produced by `walk_slide` back to a WRITABLE text frame.

    Returns None when the address no longer resolves (the deck's shapes moved,
    or the address points at a read-only surface such as chart data). Callers
    treat None as "skip", and the drift guard on text still applies on top.
    """
    parts = (addr or "").split("/")
    if len(parts) < 2 or not parts[0].startswith("s"):
        return None
    cur = None
    shapes = slide.shapes
    for tok in parts[1:]:
        if tok.startswith("sh") or (tok.startswith("g") and tok[1:].isdigit()):
            try:
                i = int(tok[2:] if tok.startswith("sh") else tok[1:])
                seq = list(shapes)
            except Exception:
                return None
            if not (0 <= i < len(seq)):
                return None
            cur = seq[i]
            try:
                shapes = cur.shapes          # only valid for a group
            except Exception:
                shapes = []
        elif tok.startswith("r") and "c" in tok:
            try:
                r, c = tok[1:].split("c", 1)
                return cur.table.rows[int(r)].cells[int(c)].text_frame
            except Exception:
                return None
        else:
            return None                      # chart/other: not writable
    try:
        return cur.text_frame if (cur is not None and cur.has_text_frame) else None
    except Exception:
        return None


def walk_deck(prs) -> tuple[list[dict], list[dict]]:
    """Return (surfaces, unreachable) for every slide in a presentation."""
    surfaces: list[dict] = []
    unreachable: list[dict] = []
    for n, slide in enumerate(prs.slides, start=1):
        s, u = walk_slide(slide, n)
        surfaces.extend(s)
        unreachable.extend(u)
    return surfaces, unreachable
