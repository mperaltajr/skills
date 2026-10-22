#!/usr/bin/env python3
"""translate_html.py — turn a picked HTML design into editable native PowerPoint.

This does the job the slide-builder-translator agent used to do for every
option, as a script. An audit of 228 real agent translations found the job is
almost entirely mechanical (read each element's position and style from the
browser, map it to a shape) and that the agents did it inconsistently: six
different letter-spacing factors, opposite baseline nudges, improvised XML
edits that made PowerPoint refuse a deck, and the same five corrections
rediscovered by eye every session. Here each rule is written down once.

What it reads: the design opened in the same headless Chromium that drew it,
so flex, grid and calc() are already resolved into boxes.
What it writes, next to the HTML, keeping the contract finalize already reads:
  option_X_native.py           runs the plan; finalize executes it unchanged
  option_X_native.plan.json    every shape and text box, decided
  option_X_translation_report.json   counts, warnings, fallbacks (fixed schema)

What it does NOT draw: anything it cannot translate faithfully (curved SVG
paths, rotated or skewed elements, images, vertical text). Those are listed
in the report as fallback elements and the script is marked FALLBACK_PENDING;
finalize refuses it until the slide-builder-translator agent has drawn just
those elements (see the agent's "Fallback mode").

Converting again after the design changed (refresh_stale, which
record_picks.py and finalize_deck.py call) keeps what the agent drew only when
every element it draws is unchanged; otherwise the script waits on the agent
with the old drawing pasted in to adapt. A hand run over a script the agent
has drawn in is refused unless --force.

Run:
  py -3 scripts/translate_html.py --out <out_dir> --slide 4 --letter B
  py -3 scripts/translate_html.py --html <file.html> --emit <dir>   (standalone)
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL))
sys.path.insert(0, str(HERE))

TRANSLATOR_VERSION = "script-1"
CANVAS_W, CANVAS_H = 1280, 720

# ---------------------------------------------------------------------------
# Browser side: read the laid-out page
# ---------------------------------------------------------------------------

EXTRACT_JS = r"""
(opts) => {
  const canvas = document.querySelector('.slide-canvas') || document.querySelector('.slide') || document.body;
  const cr = canvas.getBoundingClientRect();
  const ox = cr.x, oy = cr.y;
  const R = (r) => ({x: r.x - ox, y: r.y - oy, w: r.width, h: r.height});
  // Where Chrome drew each visible character: [left, top, right, bottom]. The
  // alarm compares these with where LibreOffice put the same characters.
  const glyphBoxes = (nodes) => {
    const out = [];
    const rg = document.createRange();
    for (const n of nodes) {
      const t = n.textContent;
      for (let i = 0; i < t.length; i++) {
        if (/\s/.test(t[i])) continue;
        const j = (t.codePointAt(i) > 0xffff) ? i + 2 : i + 1;
        rg.setStart(n, i); rg.setEnd(n, j);
        const q = rg.getBoundingClientRect();
        out.push([+(q.left - ox).toFixed(1), +(q.top - oy).toFixed(1), +(q.right - ox).toFixed(1), +(q.bottom - oy).toFixed(1)]);
        if (j === i + 2) i++;
      }
    }
    return out;
  };
  const glyphCount = (t) => { let n = 0; for (let i = 0; i < t.length; i++) { if (/\s/.test(t[i])) continue; if (t.codePointAt(i) > 0xffff) i++; n++; } return n; };
  const transparent = (c) => !c || c === 'transparent' || /rgba\([^)]*,\s*0\)$/.test(c);
  const out = {boxes: [], texts: [], svgs: [], icons: [], pictures: [], fallbacks: [], fields: {}, notes: [],
    canvas: {w: cr.width, h: cr.height}};
  let order = 0;

  // 1. Materialize ::before / ::after as real spans, so they lay out and can be
  //    measured like any element (a DOM walk cannot see pseudo-elements; ~5% of
  //    designs draw bullets and accent bars with them).
  const style = document.createElement('style');
  style.textContent = '[data-pm-b]::before{content:none!important}[data-pm-a]::after{content:none!important}';
  document.head.appendChild(style);
  for (const el of Array.from(canvas.querySelectorAll('*'))) {
    if (el.closest('svg')) continue;
    for (const [p, attr, first] of [['::before', 'data-pm-b', true], ['::after', 'data-pm-a', false]]) {
      const ps = getComputedStyle(el, p);
      if (!ps || ps.content === 'none' || ps.content === 'normal' || ps.display === 'none') continue;
      const span = document.createElement('span');
      for (let i = 0; i < ps.length; i++) {
        const name = ps[i];
        if (name === 'content') continue;
        span.style.setProperty(name, ps.getPropertyValue(name));
      }
      let txt = ps.content;
      txt = /^["'].*["']$/.test(txt) ? txt.slice(1, -1) : '';
      span.textContent = txt.replace(/\\([0-9a-fA-F]{2,6})\s?/g, (m, h) => String.fromCodePoint(parseInt(h, 16)));
      span.setAttribute('data-pseudo', p);
      el.setAttribute(attr, '');
      if (first) el.insertBefore(span, el.firstChild); else el.appendChild(span);
    }
  }

  const visible = (cs) => cs.display !== 'none' && cs.visibility !== 'hidden' && parseFloat(cs.opacity) > 0.01;
  const effOpacity = (el) => { let o = 1; for (let a = el; a && a !== document.documentElement; a = a.parentElement) o *= parseFloat(getComputedStyle(a).opacity || 1); return o; };
  const isRotated = (cs) => { const t = cs.transform; if (!t || t === 'none') return false; const m = t.match(/matrix\(([^)]+)\)/); if (!m) return true; const v = m[1].split(',').map(parseFloat); return Math.abs(v[1]) > 1e-3 || Math.abs(v[2]) > 1e-3; };
  const skip = new Set();
  const markSubtree = (el) => { skip.add(el); el.querySelectorAll('*').forEach(e => skip.add(e)); };

  // 1b. Tables and lists (2026-10-08). Each one is marked, so every box and
  //     text read below says which table cell or list item it belongs to. The
  //     planner builds one native PowerPoint table (or one text box with real
  //     bullets) from them, and keeps the separate shapes as the fallback.
  const sidesOf = (s) => ['Top','Right','Bottom','Left'].map(n => ({w: parseFloat(s['border'+n+'Width']) || 0, c: s['border'+n+'Color'], s: s['border'+n+'Style']}));
  const radiusOf = (s) => Math.max(...[s.borderTopLeftRadius, s.borderTopRightRadius, s.borderBottomRightRadius, s.borderBottomLeftRadius].map(v => parseFloat(v) || 0));
  const inField = (el) => !!el.closest('[data-template-field]');
  out.tables = []; out.lists = [];
  const TROOT = 'table, [data-native="table"]';
  for (const t of canvas.querySelectorAll(TROOT)) {
    if ((t.parentElement && t.parentElement.closest(TROOT)) || inField(t)) continue;
    const tcs = getComputedStyle(t);
    if (!visible(tcs)) continue;
    const k = out.tables.length;
    t.setAttribute('data-pm-tbl', k);
    const tag = t.tagName.toLowerCase();
    const rec = {k, ...R(t.getBoundingClientRect()), tag, id: t.getAttribute('data-shape-id') || '',
      collapse: tcs.borderCollapse === 'collapse', spacing: tcs.borderSpacing, sides: sidesOf(tcs),
      radius: radiusOf(tcs), opacity: effOpacity(t), cells: []};
    const rows = [];
    if (tag === 'table') {
      for (const r of t.rows) rows.push([r, Array.from(r.cells)]);
    } else {
      for (const c of t.children) {
        if (c.getAttribute('role') === 'row' || c.hasAttribute('data-row')) rows.push([c, Array.from(c.children)]);
        else rows.push([null, [c]]);
      }
    }
    for (const [row, cells] of rows) {
      if (row) row.setAttribute('data-pm-part', '');
      const rcs = row ? getComputedStyle(row) : null;
      const sec = row && row.parentElement !== t ? row.parentElement : null;
      if (sec) sec.setAttribute('data-pm-part', '');
      const scs = sec ? getComputedStyle(sec) : null;
      for (const c of cells) {
        const ccs = getComputedStyle(c);
        if (!visible(ccs)) continue;
        const i = rec.cells.length;
        c.setAttribute('data-pm-cell', k + ':' + i);
        c.setAttribute('data-pm-part', '');
        rec.cells.push({i, ...R(c.getBoundingClientRect()),
          bg: [tcs.backgroundColor, scs ? scs.backgroundColor : '', rcs ? rcs.backgroundColor : '', ccs.backgroundColor],
          bgimg: [tcs.backgroundImage, scs ? scs.backgroundImage : 'none', rcs ? rcs.backgroundImage : 'none', ccs.backgroundImage].join(' '),
          sides: sidesOf(ccs), rowSides: rcs ? sidesOf(rcs) : null, rowRect: row ? R(row.getBoundingClientRect()) : null,
          secSides: scs ? sidesOf(scs) : null, secRect: sec ? R(sec.getBoundingClientRect()) : null,
          radius: Math.max(radiusOf(ccs), rcs ? radiusOf(rcs) : 0), opacity: effOpacity(c),
          id: c.getAttribute('data-shape-id') || ''});
      }
    }
    out.tables.push(rec);
  }

  // A list's marker: where Chrome drew it and what it is.
  const firstChar = (el) => {
    const tw = document.createTreeWalker(el, NodeFilter.SHOW_TEXT);
    for (let n = tw.nextNode(); n; n = tw.nextNode()) {
      const p = n.parentElement;
      if (p.closest('[data-pm-item]') !== el || p.closest('[data-pm-marker]')) continue;
      const pcs = getComputedStyle(p);
      if (!visible(pcs)) continue;
      const m = n.textContent.search(/\S/);
      if (m < 0) continue;
      const rg = document.createRange(); rg.setStart(n, m); rg.setEnd(n, m + 1);
      const q = rg.getBoundingClientRect();
      return {x: q.left - ox, y: q.top - oy, h: q.height};
    }
    return null;
  };
  const markerOf = (c, k, i, ccs) => {
    const tag = k + ':' + i;
    if (ccs.display === 'list-item' && ccs.listStyleType !== 'none' && ccs.listStyleImage === 'none') {
      const ms = getComputedStyle(c, '::marker');
      const a = firstChar(c);
      if (!a) return {type: 'none'};
      const pos = ccs.listStylePosition;
      const old = c.style.listStylePosition;
      c.style.listStylePosition = pos === 'inside' ? 'outside' : 'inside';
      const b = firstChar(c);
      c.style.listStylePosition = old;
      if (!b) return {type: 'none'};
      // outside: switching to inside pushes the text right by the marker's width
      const mw = Math.abs(b.x - a.x);
      const mx = pos === 'inside' ? b.x : a.x - mw;
      return {type: 'native', style: ccs.listStyleType, color: ms.color || ccs.color,
        fs: parseFloat(ms.fontSize) || parseFloat(ccs.fontSize), ff: ms.fontFamily || ccs.fontFamily,
        fw: ms.fontWeight, x: mx, y: a.y, w: mw, h: a.h, value: (c.tagName === 'LI' ? c.value : null)};
    }
    const f = c.firstElementChild;
    if (!f) return {type: 'none'};
    const fr = f.getBoundingClientRect();
    const fcs = getComputedStyle(f);
    const txt = f.textContent.trim();
    const isPseudo = f.getAttribute('data-pseudo') === '::before';
    const smallEmpty = !txt && fr.width > 0 && fr.width <= 24 && fr.height <= 24;
    const icon = f.hasAttribute('data-icon-name') || f.tagName.toLowerCase() === 'svg';
    // a short label or glyph set before the text in its own element ("01", a check)
    f.setAttribute('data-pm-marker', tag);
    const tc = firstChar(c);
    const before = tc && fr.right - ox <= tc.x + 1 && txt.length <= 3;
    if (!(isPseudo || smallEmpty || (icon && fr.width <= 32) || (before && txt))) {
      f.removeAttribute('data-pm-marker');
      return {type: 'none'};
    }
    return {type: icon ? 'icon' : (txt ? 'glyph' : 'box'), char: txt, color: txt ? fcs.color : fcs.backgroundColor,
      fs: parseFloat(fcs.fontSize), ff: fcs.fontFamily, fw: fcs.fontWeight, ...R(fr), radius: radiusOf(fcs),
      bg: fcs.backgroundColor, bgimg: fcs.backgroundImage, sides: sidesOf(fcs), inline: fcs.display === 'inline'};
  };
  const LROOT = 'ul, ol, [data-native="list"]';
  for (const L of canvas.querySelectorAll(LROOT)) {
    if ((L.parentElement && L.parentElement.closest(LROOT)) || inField(L) || L.closest('[data-pm-tbl]')) continue;
    if (!visible(getComputedStyle(L))) continue;
    const k = out.lists.length;
    L.setAttribute('data-pm-list', k);
    const rec = {k, id: L.getAttribute('data-shape-id') || '', tag: L.tagName.toLowerCase(), ...R(L.getBoundingClientRect()), items: []};
    const walk = (list, level) => {
      const html = list.tagName === 'UL' || list.tagName === 'OL';
      for (const c of Array.from(list.children)) {
        if (html && c.tagName !== 'LI') continue;
        const ccs = getComputedStyle(c);
        if (!visible(ccs)) continue;
        const i = rec.items.length;
        c.setAttribute('data-pm-item', k + ':' + i);
        const it = {i, level, ordered: list.tagName === 'OL', ...R(c.getBoundingClientRect())};
        rec.items.push(it);
        it.marker = markerOf(c, k, i, ccs);
        for (const n of c.querySelectorAll(LROOT)) {
          if (n.parentElement.closest(LROOT) === list && n.parentElement.closest('[data-pm-item]') === c) walk(n, level + 1);
        }
      }
    };
    walk(L, 0);
    out.lists.push(rec);
  }
  const grp = (el) => {
    const g = (sel, a) => { const e = el.closest(sel); return e ? e.getAttribute(a) : null; };
    return {tbl: g('[data-pm-tbl]', 'data-pm-tbl'), cell: g('[data-pm-cell]', 'data-pm-cell'),
      list: g('[data-pm-list]', 'data-pm-list'), item: g('[data-pm-item]', 'data-pm-item'),
      marker: g('[data-pm-marker]', 'data-pm-marker'),
      part: el.hasAttribute('data-pm-part') || el.hasAttribute('data-pm-tbl')};
  };

  // 2. Template fields: the graft owns title / footer / page number. The
  //    takeaway is drawn as a body shape named "subtitle" when the layout has no
  //    subtitle placeholder; otherwise it goes to the placeholder too.
  for (const el of canvas.querySelectorAll('[data-template-field]')) {
    const name = el.getAttribute('data-template-field');
    out.fields[name] = el.innerText.replace(/\s+/g, ' ').trim();
    if (!(name === 'subtitle' && opts.subtitleAsShape)) markSubtree(el);
  }

  const all = [canvas, ...canvas.querySelectorAll('*')];
  for (const el of all) {
    if (skip.has(el)) continue;
    const tag = el.tagName.toLowerCase();
    if (el.closest('svg') && tag !== 'svg') continue;
    const cs = getComputedStyle(el);
    if (!visible(cs)) { if (cs.display === 'none') markSubtree(el); continue; }
    const r = el.getBoundingClientRect();
    const sid = el.getAttribute('data-shape-id') || el.getAttribute('data-template-field') || '';

    // A library icon (drawn into the page by icon_svg.draw_icons): one icon
    // step, inserted as the real vector icon. It used to vanish silently.
    if (el.hasAttribute('data-icon-name')) {
      out.icons.push({name: (el.getAttribute('data-icon-name') || '').trim(), ...R(r), color: cs.color,
        drawn: el.getAttribute('data-icon-drawn') || '', id: sid, g: grp(el), order: order++});
      markSubtree(el); continue;
    }

    if (isRotated(cs)) {
      out.fallbacks.push({kind: 'rotated', reason: 'rotated or skewed element', id: sid, ...R(r), g: grp(el), html: el.outerHTML, order: order++});
      markSubtree(el); continue;
    }
    if (cs.writingMode && cs.writingMode !== 'horizontal-tb') {
      out.fallbacks.push({kind: 'vertical-text', reason: 'vertical writing mode', id: sid, ...R(r), g: grp(el), html: el.outerHTML, order: order++});
      markSubtree(el); continue;
    }
    if (tag === 'img') {
      // A picture from a file next to the design, cropped and shaped as the
      // browser drew it, is a native picture (2026-10-13); anything the
      // picture tool can't carry (a web address, a filter, see-through, a
      // border, clipped by its container, inside a table or list) stays
      // with the agent.
      const src = el.currentSrc || el.src || '';
      let why = '';
      if (!/^file:/i.test(src)) why = /^data:/i.test(src) ? 'embedded picture data' : 'picture from a web address';
      else if (!el.complete || !el.naturalWidth || !el.naturalHeight) why = 'picture did not load';
      else if (cs.filter && cs.filter !== 'none') why = 'filtered picture';
      else if (effOpacity(el) < 0.999) why = 'see-through picture';
      else if (sidesOf(cs).some(s => s.w > 0 && s.s !== 'none' && s.s !== 'hidden' && !transparent(s.c))) why = 'picture with a border';
      else if (el.closest('[data-pm-tbl], [data-pm-list]')) why = 'picture inside a table or list';
      else if (cs.clipPath && cs.clipPath !== 'none') why = 'clipped picture';
      else {
        for (let a = el.parentElement; a && a !== canvas.parentElement; a = a.parentElement) {
          const acs = getComputedStyle(a);
          if (acs.overflowX !== 'visible' || acs.overflowY !== 'visible' || (acs.clipPath && acs.clipPath !== 'none')) {
            const ar = a.getBoundingClientRect();
            if (r.left < ar.left - 0.5 || r.top < ar.top - 0.5 || r.right > ar.right + 0.5 || r.bottom > ar.bottom + 0.5 || (radiusOf(acs) > 0.5 && a !== canvas)) { why = 'picture clipped by its container'; break; }
          }
        }
      }
      if (!why) {
        out.pictures.push({...R(r), src, nw: el.naturalWidth, nh: el.naturalHeight, fit: cs.objectFit,
          pos: cs.objectPosition, radii: [cs.borderTopLeftRadius, cs.borderTopRightRadius, cs.borderBottomRightRadius, cs.borderBottomLeftRadius],
          shadow: cs.boxShadow !== 'none', id: sid, g: grp(el), html: el.outerHTML, order: order++});
      } else {
        out.fallbacks.push({kind: tag, reason: 'img element (' + why + ')', id: sid, ...R(r), g: grp(el), html: el.outerHTML, order: order++});
      }
      markSubtree(el); continue;
    }
    if (tag === 'canvas' || tag === 'video' || tag === 'iframe' || tag === 'picture') {
      out.fallbacks.push({kind: tag, reason: tag + ' element', id: sid, ...R(r), g: grp(el), html: el.outerHTML, order: order++});
      markSubtree(el); continue;
    }
    if (tag === 'svg') {
      { const sv = readSvg(el, sid); sv.g = grp(el); out.svgs.push(sv); }
      continue;
    }

    // Box
    const sides = ['Top','Right','Bottom','Left'].map(s => ({w: parseFloat(cs['border'+s+'Width']) || 0, c: cs['border'+s+'Color'], s: cs['border'+s+'Style']}));
    const liveSides = sides.filter(s => s.w > 0 && s.s !== 'none' && s.s !== 'hidden' && !transparent(s.c));
    const inline = cs.display === 'inline';
    const bgimg = cs.backgroundImage;
    if (bgimg && bgimg !== 'none' && /url\(/.test(bgimg)) {
      out.fallbacks.push({kind: 'background-image', reason: 'background image', id: sid, ...R(r), g: grp(el),
        html: el.cloneNode(false).outerHTML + ' ' + bgimg + ' ' + cs.backgroundSize + ' ' + cs.backgroundPosition, order: order++});
    }
    const hasBg = !transparent(cs.backgroundColor) || (bgimg && /gradient/.test(bgimg));
    if ((hasBg || liveSides.length) && r.width >= 0 && r.height >= 0 && el !== canvas || (el === canvas && hasBg)) {
      const rects = inline ? Array.from(el.getClientRects()) : [r];
      for (const q of rects) {
        out.boxes.push({...R(q), bg: cs.backgroundColor, bgimg, sides,
          radii: [cs.borderTopLeftRadius, cs.borderTopRightRadius, cs.borderBottomRightRadius, cs.borderBottomLeftRadius],
          opacity: effOpacity(el), shadow: cs.boxShadow !== 'none', filter: cs.filter !== 'none',
          id: sid, tag, canvas: el === canvas, g: grp(el), order: order++});
      }
    }

    // Text block: an element that is not inline, plus its inline descendants.
    if (!inline) {
      const runs = [];
      const ownRects = [];
      const textNodes = [];
      const collect = (n) => {
        for (const c of n.childNodes) {
          if (c.nodeType === 3) {
            if (!c.textContent.length) continue;
            const pcs = getComputedStyle(c.parentElement);
            const rg = document.createRange(); rg.selectNodeContents(c);
            const rs = Array.from(rg.getClientRects()).filter(q => q.width > 0.3 && q.height > 0.3);
            rs.forEach(q => ownRects.push(q));
            textNodes.push(c);
            runs.push({t: c.textContent, color: pcs.color, fw: pcs.fontWeight, fi: pcs.fontStyle,
              fs: parseFloat(pcs.fontSize), ff: pcs.fontFamily, ls: pcs.letterSpacing, tt: pcs.textTransform,
              td: pcs.textDecorationLine, op: effOpacity(c.parentElement), ws: pcs.whiteSpace,
              mk: !!c.parentElement.closest('[data-pm-marker]'), gc: glyphCount(c.textContent),
              first: rs.length ? {x: rs[0].left - ox, y: rs[0].top - oy, r: rs[0].right - ox, h: rs[0].height} : null,
              last: rs.length ? {x: rs[rs.length-1].left - ox, r: rs[rs.length-1].right - ox, y: rs[rs.length-1].top - oy, h: rs[rs.length-1].height} : null});
          } else if (c.nodeType === 1) {
            if (skip.has(c)) continue;
            const ccs = getComputedStyle(c);
            if (!visible(ccs)) continue;
            if (c.tagName.toLowerCase() === 'br') { runs.push({br: true}); continue; }
            if (ccs.display === 'inline') collect(c);
          }
        }
      };
      collect(el);
      const joined = runs.filter(x => !x.br).map(x => x.t).join('');
      if (joined.trim().length && ownRects.length) {
        let x0 = 1e9, y0 = 1e9, x1 = -1e9, y1 = -1e9;
        const lineTops = [];
        for (const q of ownRects.slice().sort((a, b) => a.top - b.top)) {
          x0 = Math.min(x0, q.left); y0 = Math.min(y0, q.top); x1 = Math.max(x1, q.right); y1 = Math.max(y1, q.bottom);
          if (!lineTops.length || q.top - lineTops[lineTops.length - 1] > q.height * 0.5) lineTops.push(q.top);
        }
        const firstLineH = Math.max(...ownRects.filter(q => Math.abs(q.top - y0) < 2).map(q => q.height));
        const pl = parseFloat(cs.paddingLeft) + parseFloat(cs.borderLeftWidth), pr = parseFloat(cs.paddingRight) + parseFloat(cs.borderRightWidth);
        const pt = parseFloat(cs.paddingTop) + parseFloat(cs.borderTopWidth), pb = parseFloat(cs.paddingBottom) + parseFloat(cs.borderBottomWidth);
        // Hidden behind something painted later, or clipped away by an ancestor?
        // Only something that paints over the text counts: a transparent
        // container or the empty middle of an SVG ring on top does not hide
        // it. Every line has to be hidden, not just the middle point.
        const paints = (e) => {
          const tag = e.tagName.toLowerCase();
          if (tag === 'img' || tag === 'canvas' || tag === 'video') return true;
          if (e.closest('svg')) return tag !== 'svg' && tag !== 'g';
          const s = getComputedStyle(e);
          if (s.backgroundImage && s.backgroundImage !== 'none') return true;
          const m = s.backgroundColor.match(/rgba?\(([^)]+)\)/);
          if (!m) return false;
          const v = m[1].split(',').map(parseFloat);
          return (v.length < 4 || v[3] > 0.9) && effOpacity(e) > 0.9;
        };
        const hiddenAt = (px, py) => {
          for (const e of document.elementsFromPoint(px, py)) {
            if (e === el || el.contains(e) || e.contains(el)) return false;
            if (paints(e)) return true;
          }
          return false;
        };
        const samples = ownRects.filter(q => q.width > 0 && q.height > 0)
          .map(q => [q.left + q.width / 2, q.top + q.height / 2]);
        const covered = samples.length > 0 && samples.every(([px, py]) => hiddenAt(px, py));
        let clipped = false;
        for (let a = el.parentElement; a && a !== canvas.parentElement; a = a.parentElement) {
          const acs = getComputedStyle(a);
          if (acs.overflow === 'hidden' || acs.overflowX === 'hidden' || acs.overflowY === 'hidden') {
            const ar = a.getBoundingClientRect();
            if (x1 <= ar.left || x0 >= ar.right || y1 <= ar.top || y0 >= ar.bottom) { clipped = true; break; }
          }
        }
        out.texts.push({box: R(r), content: {x: r.x - ox + pl, y: r.y - oy + pt, w: r.width - pl - pr, h: r.height - pt - pb},
          ink: {x: x0 - ox, y: y0 - oy, w: x1 - x0, h: y1 - y0}, lines: lineTops.length, firstLineH,
          glyphs: glyphBoxes(textNodes),
          runs, align: cs.textAlign, lh: cs.lineHeight, fs: parseFloat(cs.fontSize), ws: cs.whiteSpace,
          display: cs.display, jc: cs.justifyContent, covered, clipped,
          listItem: cs.display === 'list-item' && cs.listStyleType !== 'none', g: grp(el),
          id: sid, tf: el.getAttribute('data-template-field') || (el.closest('[data-template-field]') ? el.closest('[data-template-field]').getAttribute('data-template-field') : null),
          order: order++});
      }
    }
  }

  // SVG: map the primitives these designs use (line, polyline, polygon, rect,
  // circle, ellipse, text, straight-segment paths). Anything else, or anything
  // decorated in a way shapes cannot carry, sends the whole SVG to the agent.
  // How many screen px one unit of an svg element's own drawing is (its
  // viewBox and transforms): line widths are in those units.
  function ctmScale(el) { try { const m = el.getScreenCTM(); return m ? Math.sqrt(Math.abs(m.a * m.d - m.b * m.c)) || 1 : 1; } catch (e) { return 1; } }
  // A marker that is one plain triangle pointing along the line, in the
  // line's color: {len, wid, ext (px past the line's end to the tip), fill}.
  function arrowHead(el, cs, which, sw) {
    const ref = el.getAttribute('marker-' + which) || (which === 'end' ? cs.markerEnd : cs.markerStart) || '';
    const m = /url\(\s*["']?#([^"')]+)["']?\s*\)/.exec(ref);
    const mk = m ? document.getElementById(m[1]) : null;
    if (!mk || mk.tagName.toLowerCase() !== 'marker') return null;
    const kids = Array.from(mk.children).filter(k => !['title', 'desc'].includes(k.tagName.toLowerCase()));
    if (kids.length !== 1) return null;
    const k = kids[0], kt = k.tagName.toLowerCase();
    if (k.getAttribute('transform')) return null;
    let v = [];
    if (kt === 'polygon') {
      const nums = (k.getAttribute('points') || '').trim().split(/[\s,]+/).map(parseFloat);
      for (let i = 0; i + 1 < nums.length; i += 2) v.push([nums[i], nums[i + 1]]);
    } else if (kt === 'path') {
      const d = k.getAttribute('d') || '';
      if (/[^MLZmlz\d\s,.\-e]/.test(d) || /[ml]/.test(d.replace(/e-?\d/g, ''))) return null;
      const nums = (d.match(/-?[\d.]+(?:e-?\d+)?/g) || []).map(parseFloat);
      for (let i = 0; i + 1 < nums.length; i += 2) v.push([nums[i], nums[i + 1]]);
      if (v.length === 4 && Math.abs(v[3][0] - v[0][0]) < 1e-6 && Math.abs(v[3][1] - v[0][1]) < 1e-6) v.pop();
    } else return null;
    if (v.length !== 3) return null;
    const orient = (mk.getAttribute('orient') || '0').trim();
    if (orient !== 'auto' && orient !== 'auto-start-reverse') return null;
    const units = mk.getAttribute('markerUnits') || 'strokeWidth';
    let scale = (units === 'strokeWidth' ? sw : 1) * ctmScale(el);
    const vb = mk.viewBox && mk.viewBox.baseVal;
    if (vb && vb.width > 0 && vb.height > 0) {
      const mw = parseFloat(mk.getAttribute('markerWidth') || 3), mh = parseFloat(mk.getAttribute('markerHeight') || 3);
      if (Math.abs(mw / vb.width - mh / vb.height) > 1e-3) return null;
      scale *= mw / vb.width;
    }
    // the tip: the corner farthest along x from the other two, which share an x
    let best = null;
    for (let i = 0; i < 3; i++) {
      const o = v.filter((_, j) => j !== i);
      const bx = (o[0][0] + o[1][0]) / 2, len = Math.abs(v[i][0] - bx), wid = Math.abs(o[0][1] - o[1][1]);
      if (!best || len > best.len) best = {tip: v[i], bx, len, wid, spread: Math.abs(o[0][0] - o[1][0]), my: (o[0][1] + o[1][1]) / 2};
    }
    if (!best || best.len <= 0 || best.wid <= 0 || best.spread > 0.25 * best.len || Math.abs(best.tip[1] - best.my) > 0.25 * best.wid) return null;
    const dir = Math.sign(best.tip[0] - best.bx);
    const outward = which === 'end' ? dir > 0 : (orient === 'auto-start-reverse' ? dir > 0 : dir < 0);
    if (!outward) return null;
    const refX = parseFloat(mk.getAttribute('refX') || 0);
    const ext = (which === 'start' && orient === 'auto' ? refX - best.tip[0] : best.tip[0] - refX) * scale;
    return {len: best.len * scale, wid: best.wid * scale, ext, fill: getComputedStyle(k).fill};
  }

  function readSvg(svg, sid) {
    const r = svg.getBoundingClientRect();
    const res = {id: sid, ...R(r), order: order++, items: [], unsupported: [], arrows: [], rects: [],
                 html: svg.outerHTML};
    const pt = (el, x, y) => { const m = el.getScreenCTM(); const p = new DOMPoint(x, y).matrixTransform(m); return [p.x - ox, p.y - oy]; };
    for (const el of svg.querySelectorAll('*')) {
      const tag = el.tagName.toLowerCase();
      if (['defs','title','desc','g','style','tspan','lineargradient','radialgradient','stop','clippath','mask','marker','symbol'].includes(tag)) {
        // a marker is judged where it is used (arrowHead below)
        if (['lineargradient','radialgradient','clippath','mask','symbol'].includes(tag)) res.unsupported.push(tag);
        continue;
      }
      if (el.closest('marker')) continue;      // an arrowhead's own drawing
      const cs = getComputedStyle(el);
      if (!visible(cs)) continue;
      if (!el.closest('defs, marker, symbol, clipPath, mask')) {
        // where each drawn piece is, so the self-check can leave out exactly
        // what the agent will draw when the whole drawing goes to it
        const q = el.getBoundingClientRect();
        if (q.width > 0 || q.height > 0) res.rects.push(R(q));
      }
      const mEnd = !!el.getAttribute('marker-end') || cs.markerEnd !== 'none';
      const mStart = !!el.getAttribute('marker-start') || cs.markerStart !== 'none';
      let heads = null;
      if (mEnd || mStart) {
        // Where the arrow's ends are, for the arrow-end check (scripts/arrow_ends.py).
        try {
          if (!el.closest('defs, marker, symbol') && typeof el.getTotalLength === 'function') {
            const L = el.getTotalLength(), pts = [];
            for (let k = 0; k <= 8; k++) { const q = el.getPointAtLength(L * k / 8); pts.push(pt(el, q.x, q.y)); }
            res.arrows.push({pts, start: mStart, end: mEnd, id: el.getAttribute('data-shape-id') || el.id || ''});
          }
        } catch (e) {}
        // A straight line with a plain triangle arrowhead is drawn natively
        // (a line with a PowerPoint arrowhead, 2026-10-13); any other marker
        // sends the drawing to the agent.
        const sw0 = parseFloat(cs.strokeWidth) || 1;
        const hs = {start: mStart ? arrowHead(el, cs, 'start', sw0) : null, end: mEnd ? arrowHead(el, cs, 'end', sw0) : null};
        if (['line', 'path', 'polyline'].includes(tag) && (!mStart || hs.start) && (!mEnd || hs.end)) heads = hs;
        else { res.unsupported.push('arrow marker'); continue; }
      }
      if (cs.strokeDasharray && cs.strokeDasharray !== 'none' && (tag === 'circle' || tag === 'ellipse')) { res.unsupported.push('dashed ring'); continue; }
      const st = {fill: cs.fill, fillOp: parseFloat(cs.fillOpacity) * parseFloat(cs.opacity), stroke: cs.stroke,
                  sw: (parseFloat(cs.strokeWidth) || 0) * ctmScale(el), dash: cs.strokeDasharray && cs.strokeDasharray !== 'none',
                  def: !!el.closest('defs, marker, symbol')};
      if (/url\(/.test(st.fill) || /url\(/.test(st.stroke)) { res.unsupported.push('gradient paint'); continue; }
      const A = (n) => parseFloat(el.getAttribute(n) || 0);
      const n0 = res.items.length;
      if (tag === 'line') res.items.push({k: 'poly', closed: false, pts: [pt(el, A('x1'), A('y1')), pt(el, A('x2'), A('y2'))], ...st});
      else if (tag === 'polyline' || tag === 'polygon') {
        const nums = (el.getAttribute('points') || '').trim().split(/[\s,]+/).map(parseFloat);
        const pts = []; for (let i = 0; i + 1 < nums.length; i += 2) pts.push(pt(el, nums[i], nums[i+1]));
        res.items.push({k: 'poly', closed: tag === 'polygon', pts, ...st});
      } else if (tag === 'rect') {
        const x = A('x'), y = A('y'), w = A('width'), h = A('height');
        if (A('rx') > 0 || A('ry') > 0) { res.unsupported.push('rounded svg rect'); continue; }
        res.items.push({k: 'poly', closed: true, pts: [pt(el,x,y), pt(el,x+w,y), pt(el,x+w,y+h), pt(el,x,y+h)], ...st});
      } else if (tag === 'circle' || tag === 'ellipse') {
        const cx = A('cx'), cy = A('cy'), rx = tag === 'circle' ? A('r') : A('rx'), ry = tag === 'circle' ? A('r') : A('ry');
        const a = pt(el, cx - rx, cy - ry), b = pt(el, cx + rx, cy + ry);
        res.items.push({k: 'oval', x: Math.min(a[0], b[0]), y: Math.min(a[1], b[1]), w: Math.abs(b[0]-a[0]), h: Math.abs(b[1]-a[1]), ...st});
      } else if (tag === 'path') {
        const d = el.getAttribute('d') || '';
        if (/[CcQqSsTtAa]/.test(d)) { res.unsupported.push('curved path'); continue; }
        const toks = d.match(/[MmLlHhVvZz]|-?[\d.]+(?:e-?\d+)?/g) || [];
        let i = 0, cmd = 'M', x = 0, y = 0, sx = 0, sy = 0; const subs = []; let cur = null;
        const num = () => parseFloat(toks[i++]);
        while (i < toks.length) {
          if (/[A-Za-z]/.test(toks[i])) cmd = toks[i++];
          if (cmd === 'Z' || cmd === 'z') { if (cur) { cur.closed = true; } x = sx; y = sy; cmd = 'M'; continue; }
          if (cmd === 'M' || cmd === 'm') { const nx = num(), ny = num(); x = cmd === 'm' ? x + nx : nx; y = cmd === 'm' ? y + ny : ny; sx = x; sy = y; cur = {pts: [[x, y]], closed: false}; subs.push(cur); cmd = cmd === 'm' ? 'l' : 'L'; continue; }
          if (cmd === 'L' || cmd === 'l') { const nx = num(), ny = num(); x = cmd === 'l' ? x + nx : nx; y = cmd === 'l' ? y + ny : ny; }
          else if (cmd === 'H' || cmd === 'h') { const nx = num(); x = cmd === 'h' ? x + nx : nx; }
          else if (cmd === 'V' || cmd === 'v') { const ny = num(); y = cmd === 'v' ? y + ny : ny; }
          else { i++; continue; }
          if (cur) cur.pts.push([x, y]);
        }
        for (const s of subs) res.items.push({k: 'poly', closed: s.closed, pts: s.pts.map(p => pt(el, p[0], p[1])), ...st});
      }
      if (heads) {
        const got = res.items.slice(n0);
        if (got.length !== 1 || got[0].k !== 'poly' || got[0].closed || got[0].pts.length !== 2) {
          res.items.length = n0; res.unsupported.push('arrow marker'); continue;
        }
        got[0].heads = heads;
        continue;
      }
      if (tag === 'text') {
        const bb = el.getBoundingClientRect();
        res.items.push({k: 'text', t: el.textContent.replace(/\s+/g, ' ').trim(), x: bb.left - ox, y: bb.top - oy, w: bb.width, h: bb.height,
          fs: parseFloat(cs.fontSize), ff: cs.fontFamily, fw: cs.fontWeight, color: cs.fill, anchor: cs.textAnchor || el.getAttribute('text-anchor') || 'start'});
      } else if (!['line', 'polyline', 'polygon', 'rect', 'circle', 'ellipse', 'path'].includes(tag)) res.unsupported.push(tag);
    }
    return res;
  }

  out.pseudo_count = canvas.querySelectorAll('[data-pseudo]').length;
  out.visible_text = canvas.innerText;
  return out;
}
"""

# ---------------------------------------------------------------------------
# Python side: decide what to draw
# ---------------------------------------------------------------------------

_RGB = re.compile(r"rgba?\(([^)]+)\)")


def parse_color(s):
    """CSS color -> (hex, alpha), or None for transparent/none."""
    if not s or s in ("transparent", "none", "currentcolor"):
        return None
    m = _RGB.match(s.strip())
    if m:
        parts = [p for p in re.split(r"[,\s/]+", m.group(1).strip()) if p]
        try:
            r, g, b = (int(round(float(p))) for p in parts[:3])
        except ValueError:
            return None
        a = float(parts[3]) if len(parts) > 3 else 1.0
        if a <= 0.001:
            return None
        return ("%02X%02X%02X" % (r, g, b), a)
    m = re.match(r"#([0-9a-fA-F]{6})$", s.strip())
    if m:
        return (m.group(1).upper(), 1.0)
    m = re.match(r"#([0-9a-fA-F]{3})$", s.strip())
    if m:
        return ("".join(c * 2 for c in m.group(1)).upper(), 1.0)
    return None


def _px(v, ref: float = 0.0) -> float:
    if v in (None, "", "normal", "auto"):
        return 0.0
    v = str(v).strip()
    if v.endswith("%"):
        return float(v[:-1]) / 100 * ref
    try:
        return float(v.replace("px", "").split()[0])
    except ValueError:
        return 0.0


def _gradient_mid(bgimg):
    cols = re.findall(r"rgba?\([^)]+\)|#[0-9a-fA-F]{3,6}\b", bgimg or "")
    return parse_color(cols[len(cols) // 2]) if cols else None


def _split_top(s: str) -> list[str]:
    """Split on commas that are not inside parentheses."""
    out, depth, cur = [], 0, ""
    for ch in s:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            out.append(cur.strip())
            cur = ""
        else:
            cur += ch
    if cur.strip():
        out.append(cur.strip())
    return out


_SIDES = {"to top": 0, "to right": 90, "to bottom": 180, "to left": 270,
          "to top right": 45, "to right top": 45, "to bottom right": 135, "to right bottom": 135,
          "to bottom left": 225, "to left bottom": 225, "to top left": 315, "to left top": 315}


def _linear_gradient(bgimg: str, w: float, h: float) -> dict | None:
    """A single CSS linear-gradient as {"angle": css degrees, "stops": [[0..1,
    hex, alpha], ...]}; None for anything else (radial, conic, repeating,
    several layers), which is flattened to one color instead."""
    layers = _split_top(bgimg or "")
    if len(layers) != 1:
        return None
    m = re.match(r"^linear-gradient\((.*)\)$", layers[0].strip(), re.S)
    if not m:
        return None
    parts = _split_top(m.group(1))
    angle = 180.0
    first = parts[0].strip().lower()
    if first.startswith("to "):
        if first not in _SIDES:
            return None
        angle = float(_SIDES[first])
        if " " in first[3:] and w > 0 and h > 0:
            # corner directions follow the box's own diagonal
            # (the gradient runs perpendicular to the other diagonal)
            a = math.degrees(math.atan2(h, w))
            angle = {45: a, 135: 180 - a, 225: 180 + a, 315: 360 - a}[_SIDES[first]] % 360
        parts = parts[1:]
    elif re.match(r"^-?[\d.]+(deg|turn|rad|grad)$", first):
        v, unit = re.match(r"^(-?[\d.]+)(\w+)$", first).groups()
        angle = float(v) * {"deg": 1, "turn": 360, "rad": 57.29578, "grad": 0.9}[unit]
        parts = parts[1:]
    stops = []
    length = abs(w * math.sin(math.radians(angle))) + abs(h * math.cos(math.radians(angle))) \
        if (w and h) else 0
    for p in parts:
        cm = re.match(r"^(rgba?\([^)]*\)|#[0-9a-fA-F]{3,8}|[a-z]+)\s*(.*)$", p.strip())
        if not cm:
            return None
        col = parse_color(cm.group(1))
        if col is None:
            col = (None, 0.0)              # transparent: fades its neighbor's color
        pos = None
        rest = cm.group(2).strip().split()
        if rest:
            r = rest[0]
            if r.endswith("%"):
                pos = float(r[:-1]) / 100
            elif r.endswith("px") and length:
                pos = float(r[:-2]) / length
        stops.append([pos, col[0], col[1]])
    if len(stops) < 2:
        return None
    # missing positions: first 0, last 1, the rest spread evenly between known ones
    if stops[0][0] is None:
        stops[0][0] = 0.0
    if stops[-1][0] is None:
        stops[-1][0] = 1.0
    i = 0
    while i < len(stops):
        if stops[i][0] is None:
            j = i
            while stops[j][0] is None:
                j += 1
            a, b = stops[i - 1][0], stops[j][0]
            for k in range(i, j):
                stops[k][0] = a + (b - a) * (k - i + 1) / (j - i + 1)
            i = j
        i += 1
    for k, s_ in enumerate(stops):
        s_[0] = min(1.0, max(0.0, s_[0]))
        if s_[1] is None:
            near = [q for q in stops[k + 1:] + stops[:k][::-1] if q[1]]
            if not near:
                return None
            s_[1] = near[0][1]
    return {"angle": angle % 360, "stops": stops}


GENERIC_FAMILIES = {"sans-serif": "Arial", "serif": "Times New Roman", "monospace": "Consolas",
                    "system-ui": "Segoe UI", "-apple-system": "Segoe UI", "ui-sans-serif": "Arial"}


def _family(ff: str) -> str:
    """The font Chrome actually drew with: the first family in the stack that
    is installed. Naming an uninstalled first choice ("Helvetica Neue")
    made PowerPoint and LibreOffice substitute a different, wider font."""
    stack = [f.strip().strip('"\'') for f in (ff or "").split(",") if f.strip()]
    if not stack:
        return ""
    fams = _family_index()
    installed = {f for (f, _s) in fams} | {f"{f} {s}" for (f, s) in fams}
    for f in stack:
        if f.lower() in GENERIC_FAMILIES:
            return GENERIC_FAMILIES[f.lower()]
        if f.lower() in installed:
            return f
    return stack[0]


def _bold(fw) -> bool:
    try:
        return int(fw) >= 600
    except (TypeError, ValueError):
        return str(fw) in ("bold", "bolder")


class Plan:
    def __init__(self):
        self.ops: list[dict] = []
        self.warnings: list[dict] = []
        self.fallbacks: list[dict] = []
        self.kill = {"gradients_flattened": 0, "shadows_stripped": 0, "filters_dropped": 0,
                     "opacity_on_text_normalized": 0, "text_decorations_stripped": 0}
        self.counts = {"shapes": 0, "texts": 0, "svg_items": 0, "pseudo_elements": 0,
                       "boxes_grown_for_clearance": 0, "wrap_adjusted": 0,
                       "covered_text_skipped": 0, "triangles": 0,
                       "self_check_handed_over": 0}
        self.self_check = {"ran": False, "fired": []}
        # The template's own background on this slide's layout ('RRGGBB'),
        # when known (a build's slide); None when translating a loose file.
        self.template_bg: str | None = None
        # The design's canvas color when it was left to the template (not drawn).
        self.canvas_bg: str | None = None
        # Tables and lists drawn natively, each with its fallback versions
        # (see _plan_native_groups).
        self.groups: dict = {}

    def warn(self, code, detail):
        self.warnings.append({"code": code, "detail": detail})


# --- boxes -----------------------------------------------------------------

# A canvas this close (RGB distance) to the template's background IS the
# template's background.
CANVAS_MATCH = 3.0
# A fill this close to what is behind it does not show (PALE_FILL_ON_BACKGROUND).
PALE_FILL_LIMIT = 10.0
PALE_FILL_MIN_AREA = 0.02 * 1280 * 720     # only panels, not dots and rules


def _color_distance(a: str, b: str) -> float:
    from _template_bg import color_distance
    return color_distance(a, b)


def _mix(fg: str, bg: str, alpha: float) -> str:
    f = [int(fg[i:i + 2], 16) for i in (0, 2, 4)]
    g = [int(bg[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(a * alpha + b * (1 - alpha)):02X}" for a, b in zip(f, g))


def _check_pale_fills(plan: "Plan") -> None:
    """Warn when a large fill nearly matches what is behind it: the slide's
    background (the template's, when known) or the panel it sits on. A pale
    panel designed and approved on white vanished on a gray-blue master
    (2026-10-06). Advisory: the review page shows it."""
    designed_bg = plan.canvas_bg or "FFFFFF"          # what the designer saw
    slide_bg = plan.template_bg or designed_bg        # what the deck will show
    filled = [o for o in plan.ops if o.get("op") == "shape" and o.get("fill")
              and o.get("_box") and o.get("_filled")]

    def _against(o, bg):
        x, y, w, h = o["_box"]
        under = bg
        for u in reversed([u for u in filled if u["_order"] < o["_order"]]):
            ux, uy, uw, uh = u["_box"]
            if ux - 1 <= x and uy - 1 <= y and x + w <= ux + uw + 1 and y + h <= uy + uh + 1:
                under = _mix(u["fill"], bg, float(u.get("alpha") or 1.0))
                break
        return _mix(o["fill"], under, float(o.get("alpha") or 1.0)), under

    for o in filled:
        x, y, w, h = o["_box"]
        if w * h < PALE_FILL_MIN_AREA:
            continue
        seen, under = _against(o, slide_bg)
        line = o.get("line") or {}
        if line.get("color") and line.get("w", 0) >= 0.5 and \
                _color_distance(line["color"], under) > PALE_FILL_LIMIT:
            continue          # an outline shows the panel's edge
        d = _color_distance(seen, under)
        # Exactly the color behind it, as designed, is deliberate (a white card
        # holding a one-sided accent bar, a mask): on 99 past designs that was
        # 89 of 94 hits. It is a defect only when the design showed it and the
        # deck's background swallows it, or when a tint nearly matches.
        d_design = _color_distance(*_against(o, designed_bg))
        if (1.0 <= d <= PALE_FILL_LIMIT) or (d < 1.0 and d_design >= 1.0):
            what = ("the template's background" if under == plan.template_bg
                    else "the slide's background" if under == slide_bg else "the panel behind it")
            plan.warn("PALE_FILL_ON_BACKGROUND",
                      f"{o.get('name') or 'a panel'} ({w:.0f}x{h:.0f} px) is #{seen}, "
                      f"{d:.0f} away from {what} (#{under}): it will barely show. Give it "
                      "a visibly different fill or an outline")


def _check_arrow_ends(plan: "Plan", data: dict) -> None:
    """The design's arrows (SVG lines and paths with arrow markers) must stop
    clear of every box (scripts/arrow_ends.py). On cycle diagrams the sketch's
    arrowheads touched the boxes and the converter copied that (2026-10-06)."""
    import arrow_ends as AE
    arrows = []
    for s in data.get("svgs") or []:
        for a in s.get("arrows") or []:
            pts = [tuple(p) for p in a.get("pts") or []]
            if len(pts) < 2:
                continue
            heads = ([(pts[0][0], pts[0][1], "start")] if a.get("start") else []) + \
                    ([(pts[-1][0], pts[-1][1], "end")] if a.get("end") else [])
            arrows.append({"name": a.get("id") or s.get("id") or "", "points": pts,
                           "heads": heads})
    if not arrows:
        return
    boxes = []
    for b in data.get("boxes") or []:
        if b.get("canvas"):
            continue
        live = any(sd["w"] > 0 and sd["s"] not in ("none", "hidden") and parse_color(sd["c"])
                   for sd in b.get("sides") or [])
        if parse_color(b.get("bg")) or live:
            boxes.append({"name": b.get("id") or b.get("tag") or "", "box": (b["x"], b["y"], b["w"], b["h"])})
    for s in data.get("svgs") or []:
        for it in s.get("items") or []:
            if it.get("def") or not (parse_color(it.get("fill")) or parse_color(it.get("stroke"))):
                continue
            if it["k"] == "oval":
                boxes.append({"name": s.get("id") or "svg shape",
                              "box": (it["x"], it["y"], it["w"], it["h"])})
            elif it["k"] == "poly" and it.get("closed") and len(it.get("pts") or []) >= 3:
                xs, ys = [p[0] for p in it["pts"]], [p[1] for p in it["pts"]]
                boxes.append({"name": s.get("id") or "svg shape",
                              "box": (min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys))})
    for h in AE.find_hits(arrows, boxes):
        plan.warnings.append({"code": AE.CODE, "source": "design",
                              "detail": "design: " + AE.describe(h)})


def _plan_box(plan: Plan, b: dict, order_key) -> None:
    x, y, w, h = b["x"], b["y"], b["w"], b["h"]
    if b.get("shadow"):
        plan.kill["shadows_stripped"] += 1
    if b.get("filter"):
        plan.kill["filters_dropped"] += 1
    sides = b["sides"]
    live = [s for s in sides if s["w"] > 0 and s["s"] not in ("none", "hidden") and parse_color(s["c"])]

    # CSS triangle: a zero-size box drawn entirely by borders, one side colored
    # and its two neighbors transparent. The point is opposite the colored side.
    zero_content = (abs(w - (sides[1]["w"] + sides[3]["w"])) < 1.0
                    and abs(h - (sides[0]["w"] + sides[2]["w"])) < 1.0)
    if zero_content and len(live) == 1 and parse_color(b.get("bg")) is None:
        side = sides.index(live[0])
        neighbors = [sides[(side + 1) % 4], sides[(side + 3) % 4]]
        if all(n["w"] > 0 for n in neighbors):
            col = parse_color(live[0]["c"])
            rot = {0: 180, 1: 270, 2: 0, 3: 90}[side]   # colored top -> points down
            sw, sh, sx, sy = w, h, x, y
            if rot in (90, 270):
                # The triangle shape points up; turned sideways its width and
                # height swap, so draw it swapped and keep the same center.
                sw, sh = h, w
                sx, sy = x + (w - sw) / 2, y + (h - sh) / 2
            plan.ops.append({"op": "shape", "kind": "triangle", "x": sx, "y": sy, "w": sw, "h": sh,
                             "rotation": rot, "fill": col[0], "alpha": col[1] * b["opacity"],
                             "name": b.get("id") or None, "_order": order_key})
            plan.counts["triangles"] += 1
            return

    if b.get("canvas"):
        col = parse_color(b["bg"])
        if not col or col[0] == "FFFFFF":
            plan.canvas_bg = "FFFFFF"
            return
        if plan.template_bg and _color_distance(col[0], plan.template_bg) <= CANVAS_MATCH:
            # The design is drawn on the template's own background, which the
            # slide already has. A full-slide rectangle in that color would only
            # cover the master's own artwork (2026-10-06).
            plan.canvas_bg = col[0]
            plan.counts["canvas_left_to_template"] = 1
            return
    col = parse_color(b["bg"])
    flattened = False
    gradient = None
    if b.get("bgimg") and "gradient" in b["bgimg"]:
        # A gradient paints over the background color. A single linear one
        # is drawn natively; radial, conic or layered ones become one color.
        gradient = _linear_gradient(b["bgimg"], w, h)
        col = _gradient_mid(b["bgimg"]) or col
        if gradient is None:
            plan.kill["gradients_flattened"] += 1
            flattened = True
    if w < 0.5 or h < 0.5:
        return
    uniform = len(live) == 4 and len({(round(s["w"], 1), s["c"], s["s"]) for s in live}) == 1
    radii = [_px(r, min(w, h)) for r in b["radii"]]
    rad = max(radii)
    if rad > 0.5 and len({round(r) for r in radii}) > 1:
        plan.warn("MIXED_CORNER_RADII", f"{b.get('id') or b['tag']}: drawn with one radius")
    if col or uniform:
        if rad >= min(w, h) / 2 - 0.5 and abs(w - h) < 1.0:
            kind, ratio = "oval", None
        elif rad > 0.5:
            kind, ratio = "rounded", min(0.5, rad / max(1.0, min(w, h)))
        else:
            kind, ratio = "rect", None
        line = None
        if uniform:
            lc = parse_color(live[0]["c"])
            line = {"color": lc[0], "w": live[0]["w"], "dash": live[0]["s"]}
        bw = live[0]["w"] if uniform else 0
        plan.ops.append({"op": "shape", "kind": kind, "radius_ratio": ratio,
                         "x": x + bw / 2, "y": y + bw / 2, "w": w - bw, "h": h - bw,
                         "fill": col[0] if col else None,
                         # a native gradient's stops carry their own alpha; the
                         # middle stop's (col) used to fade all of them again
                         "alpha": (1.0 if gradient else (col[1] if col else 1.0)) * b["opacity"],
                         "line": line, "name": b.get("id") or None, "_order": order_key,
                         "_box": (x, y, w, h), "_filled": bool(col), "_flattened": flattened})
        if gradient:
            plan.ops[-1]["gradient"] = gradient
        plan.counts["shapes"] += 1
    if live and not uniform:
        # One-sided borders: python-pptx outlines are all-or-nothing, so each
        # bordered side is its own thin bar at that edge.
        for i, s in enumerate(sides):
            c = parse_color(s["c"]) if s["w"] > 0 and s["s"] not in ("none", "hidden") else None
            if not c:
                continue
            bw = s["w"]
            geo = [(x, y, w, bw), (x + w - bw, y, bw, h), (x, y + h - bw, w, bw), (x, y, bw, h)][i]
            plan.ops.append({"op": "shape", "kind": "rect", "x": geo[0], "y": geo[1], "w": geo[2],
                             "h": geo[3], "fill": c[0], "alpha": c[1] * b["opacity"],
                             "line": None, "_order": order_key + 0.001 * (i + 1)})
            plan.counts["shapes"] += 1
            if s["s"] in ("dashed", "dotted"):
                plan.warn("DASHED_EDGE_AS_SOLID", f"{b.get('id') or b['tag']}: one-sided {s['s']} border drawn solid")


# --- text ------------------------------------------------------------------

_FONT_CACHE: dict = {}
_FAMILIES: dict | None = None


def _family_index() -> dict:
    """{(family, style): path} from the fonts' own name tables, cached on disk.

    The shared font index is keyed by FILE name (arial.ttf), so looking up the
    CSS family ("Arial") found nothing, and the wrap check silently did nothing.
    """
    global _FAMILIES
    if _FAMILIES is not None:
        return _FAMILIES
    import _scratch
    cache = _scratch.cache_dir() / "slidelab_font_families.json"
    try:
        _FAMILIES = {tuple(k.split("|", 1)): v for k, v in
                     json.loads(cache.read_text(encoding="utf-8")).items()}
        return _FAMILIES
    except Exception:
        pass
    _FAMILIES = {}
    try:
        from PIL import ImageFont
        from _chrome_schema import _font_index
        for path in set(_font_index().values()):
            try:
                fam, sty = ImageFont.truetype(path, 10).getname()
            except Exception:
                continue
            _FAMILIES.setdefault(((fam or "").lower(), (sty or "").lower()), path)
        cache.write_text(json.dumps({f"{k[0]}|{k[1]}": v for k, v in _FAMILIES.items()}),
                         encoding="utf-8")
    except Exception:
        pass
    return _FAMILIES


_FACES = None


def _face_table() -> dict:
    """{family (lowercase): [face, ...]} for every installed font, cached on disk.

    A face: path, weight, width, italic, and its PowerPoint name: the legacy
    family (name id 1) plus whether that is the family's Bold style (id 2).
    Chrome groups faces by the typographic family (id 16, else id 1) and picks
    by weight; PowerPoint knows each face by its legacy family + bold on/off
    ("Brand Sans Medium", or "Brand Sans" + bold for the Bold face).
    """
    global _FACES
    if _FACES is not None:
        return _FACES
    import _scratch
    cache = _scratch.cache_dir() / "slidelab_font_faces.json"
    try:
        _FACES = json.loads(cache.read_text(encoding="utf-8"))
        return _FACES
    except Exception:
        pass
    _FACES = {}
    try:
        from fontTools.ttLib import TTFont
        from _chrome_schema import _font_index
        for path in sorted(set(_font_index().values())):
            try:
                f = TTFont(path, lazy=True, fontNumber=0)
                n = f["name"]
                fam1, sty2 = n.getDebugName(1) or "", n.getDebugName(2) or ""
                typo = n.getDebugName(16) or fam1
                os2 = f["OS/2"]
                face = {"path": path, "weight": int(os2.usWeightClass), "width": int(os2.usWidthClass),
                        "italic": bool(os2.fsSelection & 1) or "italic" in sty2.lower(),
                        "ppt_name": fam1, "ppt_bold": "bold" in sty2.lower()}
            except Exception:
                continue
            for key in {typo.lower(), fam1.lower()}:
                _FACES.setdefault(key, []).append(face)
        cache.write_text(json.dumps(_FACES), encoding="utf-8")
    except Exception:
        pass
    return _FACES


def _pick_weight(faces: list[dict], w: int) -> dict:
    """CSS font matching by weight: 400-500 look up to 500 then down; above
    500 look up then down; below 400 look down then up."""
    ws = sorted({f["weight"] for f in faces})
    if w in ws:
        best = w
    elif 400 <= w <= 500:
        up = [x for x in ws if w < x <= 500]
        down = [x for x in ws if x < w]
        best = up[0] if up else (down[-1] if down else min(x for x in ws if x > w))
    elif w > 500:
        up = [x for x in ws if x > w]
        best = up[0] if up else max(x for x in ws if x < w)
    else:
        down = [x for x in ws if x < w]
        best = down[-1] if down else min(x for x in ws if x > w)
    return next(f for f in faces if f["weight"] == best)


def _face(ff: str, fw, italic: bool = False) -> tuple[str, bool]:
    """(font name, bold) as PowerPoint needs it to draw the face Chrome drew.

    Weight 600 and up is always the base family with bold on. Naming the
    heavy face instead (Arial at 800 is "Arial Black", a brand family at 600
    its own "Semibold" face) did not survive: finalize's theme pass swaps
    every font name that is not the theme font for the theme font, and the
    weight went with the name, so the text shipped regular, in PowerPoint
    too (2026-10-06). extract() draws 600+ at 700 before measuring, so the
    self-check compares against the face that actually ships.

    Below 600, Chrome picks a face by weight from everything installed under
    the family (a brand family at 500 is often its own "Medium" face), and
    that face is named the way PowerPoint knows it.
    """
    fam = _family(ff)
    try:
        w = int(fw)
    except (TypeError, ValueError):
        w = 700 if str(fw) in ("bold", "bolder") else 400
    if w >= 600:
        return fam, True
    faces = [f for f in _face_table().get(fam.lower(), []) if f["italic"] == italic] or \
        _face_table().get(fam.lower(), [])
    normal = [f for f in faces if f["width"] == 5] or faces
    if not normal:
        return fam, w >= 600
    face = _pick_weight(normal, w)
    return face["ppt_name"] or fam, face["ppt_bold"]


class _Measure:
    """Text width from the font file's own advance widths: exact at any size
    (PIL rounds each letter to whole pixels at small sizes, so 13px and 13.3px
    measured the same)."""

    def __init__(self, metrics, size_px: float):
        self.cmap, self.adv, self.upm, self.default = metrics
        self.k = size_px / self.upm

    def getlength(self, s: str) -> float:
        return sum(self.adv.get(self.cmap.get(ord(ch)), self.default) for ch in s) * self.k


_METRICS: dict = {}


def _font(family: str, bold: bool, size_px: float):
    """A measurer for this face and size, or None when the font file can't be found."""
    key = (family.lower(), bold, round(size_px, 2))
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    fams = _family_index()
    fam = family.lower()
    path = _face_path(family, bold) or next((p for (f, s), p in fams.items() if f == fam), None)
    font = None
    if path:
        if path not in _METRICS:
            try:
                from fontTools.ttLib import TTFont
                f = TTFont(path, lazy=True, fontNumber=0)
                hmtx = f["hmtx"].metrics
                adv = {g: m[0] for g, m in hmtx.items()}
                _METRICS[path] = (f.getBestCmap() or {}, adv, f["head"].unitsPerEm,
                                  adv.get(".notdef", f["head"].unitsPerEm // 2))
            except Exception:
                _METRICS[path] = None
        if _METRICS[path]:
            font = _Measure(_METRICS[path], max(1.0, float(size_px)))
    _FONT_CACHE[key] = font
    return font


_DESC_CACHE: dict = {}


def _face_path(face: str, bold: bool):
    """Font file for a PowerPoint font name + bold flag."""
    faces = [f for f in _face_table().get(face.lower(), []) if f["ppt_name"].lower() == face.lower()
             and not f["italic"]]
    hit = next((f for f in faces if f["ppt_bold"] == bold), None) or (faces[0] if faces else None)
    return hit["path"] if hit else None


def vertical_metrics(face: str, bold: bool) -> tuple[float, float, float]:
    """(ascent, descent, normal line height) in em, as Chrome sets this face.

    Chrome on Windows takes ascent and descent from the font's Windows
    metrics (its typographic ones when the font says USE_TYPO_METRICS), and
    line-height "normal" from the hhea ascent + descent + line gap (measured
    on Arial, Arial Black, Calibri, Georgia, Segoe UI, Consolas, Verdana and
    Times New Roman, 2026-10-09). Arial's values when the file can't be found.
    """
    key = (face.lower(), bold)
    if key not in _DESC_CACHE:
        got = (0.905, 0.212, 1.149)
        path = _face_path(face, bold)
        if path:
            try:
                from fontTools.ttLib import TTFont
                f = TTFont(path, lazy=True)
                upm = f["head"].unitsPerEm
                o, hh = f["OS/2"], f["hhea"]
                if o.fsSelection & 128:
                    asc, desc = o.sTypoAscender / upm, abs(o.sTypoDescender) / upm
                else:
                    asc, desc = o.usWinAscent / upm, o.usWinDescent / upm
                normal = (hh.ascent + abs(hh.descent) + hh.lineGap) / upm
                if asc > 0 and desc >= 0:
                    got = (asc, desc, normal if normal > 0.5 else asc + desc)
            except Exception:
                pass
        _DESC_CACHE[key] = got
    return _DESC_CACHE[key]


# Where each renderer puts the first baseline of a text box (top inset 0,
# anchored at the top) whose paragraph has exact line spacing L ("spcPts") and
# whose largest letters are S, below the box top. Measured from the glyph
# origins in each program's PDF and checked against the pixels, on 8 fonts
# (Arial, Arial Black, Calibri, Georgia, Segoe UI, Consolas, Verdana, Times
# New Roman), 9 to 72 pt, L from 0.69 S to 2.3 S (2026-10-09):
#   PowerPoint   rounds L to whole points, for the line pitch too; then
#                L > floor(1.2 S) pt (past its "single" line):  0.75 L, for
#                every font;  otherwise  max(0.75 L, L - D), where D is the
#                font's share of a 1.2 S line below the baseline:
#                D = 1.2 S x descent / (ascent + descent)
#   LibreOffice  L - 0.2 S, for every font and every L; pitch L, unrounded
# Chrome puts it at the line's top + half the leading + the font's ascent.
# So the converter writes line spacing in whole points (both programs then set
# the same pitch) and places each box for PowerPoint, the program the deck is
# opened in. A one-line text box gets L = 0.8 S, the one spacing at which both
# programs put the baseline in the same place (0.6 S below the top; letters
# are not cut off); its line spacing is not visible otherwise. Text on more
# lines keeps the design's pitch, and LibreOffice draws it
# (L - 0.2 S) - PowerPoint's baseline lower, which the self-check expects.
PPT_SINGLE = 1.2
TIGHT_SINGLE = 0.8


def line_spacing_pt(lh_px: float, size_pt: float, single: bool) -> int:
    """The exact line spacing to write, in whole points."""
    if single:
        return max(1, int(round(TIGHT_SINGLE * size_pt)))
    return max(1, int(round(lh_px * 0.75)))


def ppt_first_baseline(L: float, S: float, asc: float, desc: float) -> float:
    """px below the box top; L (the spacing written) and S in px."""
    import math
    Lpt = round(L * 0.75)
    Lr = Lpt / 0.75
    if Lpt > math.floor(PPT_SINGLE * S * 0.75 + 1e-6):
        return 0.75 * Lr
    return max(0.75 * Lr, Lr - PPT_SINGLE * S * desc / (asc + desc))


def lo_first_baseline(L: float, S: float) -> float:
    return L - 0.2 * S


def place_first_line(ink_top: float, first_h: float, run: dict | None, lh_px: float,
                     single: bool) -> tuple[float, float, float]:
    """(box top, line spacing to write in px, how much lower LibreOffice draws
    the text than PowerPoint) for a text whose first line's text box (Chrome)
    starts at ink_top and is first_h tall, set in run's face and size."""
    asc, desc, _n = vertical_metrics(run["font"], run["bold"]) if run else (0.905, 0.212, 1.149)
    base = ink_top + first_h * asc / (asc + desc)      # Chrome's first baseline
    S_pt = _snapped_px(run["size_px"]) * 0.75 if run else lh_px * 0.75 / 1.2
    S = S_pt / 0.75
    L = line_spacing_pt(lh_px, S_pt, single) / 0.75
    b = ppt_first_baseline(L, S, asc, desc)
    return base - b, L, lo_first_baseline(L, S) - b


_CAP_CACHE: dict = {}
PAD_STEP_EMU = 360          # the top inset is written in whole steps of this


def cap_height(face: str, bold: bool) -> float:
    """Capital-letter height in em: the font's OS/2 sCapHeight, else Arial's 0.716."""
    key = (face.lower(), bold)
    if key not in _CAP_CACHE:
        got = 0.716
        path = _face_path(face, bold)
        if path:
            try:
                from fontTools.ttLib import TTFont
                f = TTFont(path, lazy=True)
                cap = getattr(f["OS/2"], "sCapHeight", 0) or 0
                if cap > 0:
                    got = cap / f["head"].unitsPerEm
            except Exception:
                pass
        _CAP_CACHE[key] = got
    return _CAP_CACHE[key]


def top_pad_px(run: dict | None, L: float) -> float:
    """How far the capitals of a top-anchored text box (top inset 0, exact
    line spacing L px) reach above the box in PowerPoint, rounded up to whole
    PAD_STEP_EMU; 0 when they don't.

    One-line boxes get spacing 0.8 x size (TIGHT_SINGLE), which puts the
    first baseline 0.6 x size below the box top, so capitals (~0.72 x size)
    stuck out above the box and its selection frame looked wrong
    (2026-10-13). The box is then drawn this much higher and taller with
    this much top inset (html_emit._text): the text itself stays put."""
    if not run:
        return 0.0
    asc, desc, _n = vertical_metrics(run["font"], run["bold"])
    S = _snapped_px(run["size_px"])
    k = cap_height(run["font"], run["bold"]) * S - ppt_first_baseline(L, S, asc, desc)
    if k <= 0:
        return 0.0
    return math.ceil(k * 9525 / PAD_STEP_EMU - 1e-9) * PAD_STEP_EMU / 9525


CAPS_ABOVE_LIMIT_PX = 1.0


def caps_above_frame(slide) -> list[dict]:
    """Advisory check on a built slide: top-anchored text boxes with exact
    line spacing whose capitals reach more than CAPS_ABOVE_LIMIT_PX above the
    box (where PowerPoint draws its selection frame). Each finding: {"shape",
    "text", "above_px"}. Fired on 427 boxes of the 10/12 deck; boxes drawn
    with top_pad_px don't fire."""
    from pptx.oxml.ns import qn
    out = []

    def walk(shapes):
        for sh in shapes:
            if getattr(sh, "shape_type", None) == 6:
                yield from walk(sh.shapes)
            else:
                yield sh
    for sh in walk(slide.shapes):
        if not getattr(sh, "has_text_frame", False) or not sh.text_frame.text.strip():
            continue
        body = sh._element.find(".//" + qn("a:bodyPr"))
        if body is None or (body.get("anchor") or "t") != "t":
            continue
        tins = int(body.get("tIns") if body.get("tIns") is not None else 45720) / 9525
        p0 = sh.text_frame.paragraphs[0]
        spc = p0._p.find(qn("a:pPr") + "/" + qn("a:lnSpc") + "/" + qn("a:spcPts"))
        runs = [r for r in p0.runs if r.text.strip() and r.font.size]
        if spc is None or not runs:
            continue
        big = max(runs, key=lambda r: r.font.size)
        S_px = big.font.size.pt / 0.75
        L_px = int(spc.get("val")) / 100 / 0.75
        face = big.font.name or "Arial"
        bold = bool(big.font.bold)
        asc, desc, _n = vertical_metrics(face, bold)
        above = cap_height(face, bold) * S_px - (tins + ppt_first_baseline(L_px, S_px, asc, desc))
        if above > CAPS_ABOVE_LIMIT_PX:
            out.append({"shape": sh.name, "text": sh.text_frame.text[:40], "above_px": round(above, 1)})
    return out


def _snapped_px(size_px: float) -> float:
    from twins.helpers import snap_font_pt
    return snap_font_pt(size_px * 0.75) / 0.75


def _wrapped_lines(paragraphs, width_px: float) -> int | None:
    """Greedy word-wrap line count, measured with the real font files (as
    PowerPoint will set the text, which is slightly wider than Chrome). None if
    a font is missing. A line may end at a space, or inside a word after a
    hyphen or dash (_chrome_schema.hyphen_pieces), as PowerPoint and Chrome
    both do."""
    from _chrome_schema import hyphen_pieces
    total = 0
    for para in paragraphs:
        words, widths = [], []
        for r in para:
            # measured at the size PowerPoint will use (snapped to the allowed list)
            f = _font(r.get("font") or "", r.get("bold", False), _snapped_px(r["size_px"]))
            if f is None:
                return None
            for wd in re.split(r"(\s+)", r["t"]):
                for piece in ([wd] if not wd or wd.isspace() else hyphen_pieces(wd)):
                    if piece:
                        words.append(piece)
                        widths.append(f.getlength(piece)
                                      + (r.get("letter_spacing_px") or 0) * len(piece))
        lines, cur = 1, 0.0
        for wd, wl in zip(words, widths):
            if wd.isspace():
                cur += wl
                continue
            if cur > 0 and cur + wl > width_px:
                lines += 1
                cur = wl
            else:
                cur += wl
        total += lines
    return total


def _plan_text(plan: Plan, t: dict, order_key, name=None) -> None:
    from twins.helpers import upper_preserving_terms
    if t.get("covered"):
        plan.counts["covered_text_skipped"] += 1
        plan.warn("TEXT_HIDDEN_IN_DESIGN", f"{t.get('id') or 'text'}: covered by a later shape in the design; not drawn")
        return
    if t.get("clipped"):
        plan.warn("TEXT_HIDDEN_IN_DESIGN", f"{t.get('id') or 'text'}: clipped away by its container; not drawn")
        return
    pre = t["ws"] in ("pre", "pre-wrap", "pre-line", "break-spaces")
    paras: list[list[dict]] = [[]]
    prev_last = None
    for r in t["runs"]:
        if r.get("br"):
            paras.append([])
            prev_last = None
            continue
        txt = r["t"]
        if not pre:
            txt = re.sub(r"\s+", " ", txt)
        # Side-by-side pieces with a visible gap but no space between them
        # ("$1.2B" and "to" set as separate spans with a margin): keep the gap.
        if prev_last and r.get("first") and paras[-1] and not paras[-1][-1]["t"].endswith(" ") \
                and not txt.startswith(" ") \
                and r["first"]["y"] < prev_last["y"] + prev_last.get("h", r["first"]["h"]) \
                and prev_last["y"] < r["first"]["y"] + r["first"]["h"] \
                and r["first"]["x"] - prev_last["r"] > 0.2 * r["fs"]:
            # Same line = the pieces overlap vertically (a 32px figure and a
            # 16px "to" share a baseline, not a top). Spaces sized to the gap.
            gap = r["first"]["x"] - prev_last["r"]
            txt = " " * max(1, round(gap / (0.28 * r["fs"]))) + txt
        if r.get("last"):
            prev_last = r["last"]
        if r["tt"] == "uppercase":
            txt = upper_preserving_terms(txt)          # "TaaS" keeps its casing
        elif r["tt"] == "lowercase":
            txt = txt.lower()
        elif r["tt"] == "capitalize":
            txt = txt.title()
        if r.get("td") and r["td"] != "none":
            plan.kill["text_decorations_stripped"] += 1
        col = parse_color(r["color"])
        if (col and col[1] < 0.999) or r.get("op", 1) < 0.999:
            plan.kill["opacity_on_text_normalized"] += 1
        face, bold = _face(r["ff"], r["fw"])
        run = {"t": txt, "font": face, "size_px": r["fs"], "bold": bold,
               "italic": r["fi"] == "italic", "color": col[0] if col else None,
               "letter_spacing_px": _px(r["ls"]) if r["ls"] != "normal" else 0}
        parts = txt.split("\n") if pre else [txt]
        for k, part in enumerate(parts):
            if k:
                paras.append([])
            paras[-1].append(dict(run, t=part))
    for p in paras:
        if p:
            p[0]["t"] = p[0]["t"].lstrip()
            p[-1]["t"] = p[-1]["t"].rstrip()
    paras = [p for p in paras if "".join(x["t"] for x in p).strip()] or []
    if not paras:
        return
    if t.get("listItem"):
        paras[0].insert(0, dict(paras[0][0], t="• "))

    fs = t["fs"]
    r0 = next((r for p in paras for r in p if r["t"].strip()), None)
    if t["lh"] not in ("normal", None):
        lh = _px(t["lh"])
    else:
        # "normal" is the face's own line height (1.15 em for Arial, 1.33 for
        # Segoe UI), not a fixed 1.15
        lh = fs * (vertical_metrics(r0["font"], r0["bold"])[2] if r0 else 1.15)
    lines = max(1, t["lines"])
    ink, cont = t["ink"], t["content"]
    align = t["align"] if t["align"] in ("left", "center", "right", "justify") else \
        {"start": "left", "end": "right"}.get(t["align"], "left")
    if t["display"] in ("flex", "inline-flex", "grid") and t.get("jc") == "center":
        align = "center"
    x, w = cont["x"], cont["w"]
    if w < 2:
        x, w = ink["x"], ink["w"]
    elif align in ("left", "justify") and ink["w"] > 0:
        right = cont["x"] + cont["w"]
        x = ink["x"]
        w = max(right - x, ink["w"])
    elif align == "center" and ink["w"] > 0:
        cx = ink["x"] + ink["w"] / 2
        half = max(ink["w"] / 2, min(cx - cont["x"], cont["x"] + cont["w"] - cx))
        x, w = cx - half, 2 * half
    elif align == "right" and ink["w"] > 0:
        right = max(ink["x"] + ink["w"], cont["x"] + cont["w"])
        w = max(right - cont["x"], ink["w"])
        x = right - w
    single = lines == 1 and len(paras) == 1
    wrap = not (single or t["ws"] == "nowrap")
    design_px = max((r["size_px"] for p in paras for r in p if r["t"].strip()), default=fs)
    cased = any(r.get("tt") == "uppercase" and upper_preserving_terms(r["t"]) != r["t"].upper()
                for r in t["runs"] if not r.get("br"))

    # Will PowerPoint wrap it into more lines than Chrome did? Measured with
    # the real font files. Widen a little; if that is not enough, keep the
    # design's size and say so (WRAP_RISK, shown on the review page). The
    # font used to be stepped down one size silently: 12 pt body text shipped
    # at 11 pt on a misprediction, and at 11 pt the predictor still saw the
    # extra line (2026-10-13).
    if wrap:
        predicted = _wrapped_lines(paras, w)
        if predicted is not None and predicted > lines:
            grown = False
            for extra in (0.02, 0.04, 0.06):
                nw = w * (1 + extra)
                if _wrapped_lines(paras, nw) <= lines:
                    if align == "center":
                        x -= (nw - w) / 2
                    elif align == "right":
                        x -= nw - w
                    w, grown = nw, True
                    break
            if grown:
                plan.counts["wrap_adjusted"] += 1
                plan.warn("WRAP_ADJUSTED", f"{t.get('id') or 'text'}: would wrap to "
                          f"{predicted} lines in PowerPoint vs {lines} in the design; widened")
            else:
                plan.counts["wrap_risk"] = plan.counts.get("wrap_risk", 0) + 1
                plan.warn("WRAP_RISK", f"{t.get('id') or 'text'}: may wrap to {predicted} lines "
                          f"in PowerPoint vs {lines} in the design, even 6% wider; kept at the "
                          "design's size. Check this text in the finished slide; if it runs onto "
                          "an extra line, shorten it or give its box more room in the design")
        elif predicted is not None and predicted < lines:
            # Fits on fewer lines than the design (Chrome broke at a width the
            # PowerPoint text just clears): narrow the box toward the design's
            # widest line until the breaks match, so the card keeps its shape.
            nw = w
            while nw - 2 >= ink["w"]:
                nw -= 2
                if _wrapped_lines(paras, nw) == lines:
                    if align == "center":
                        x += (w - nw) / 2
                    elif align == "right":
                        x += w - nw
                    w = nw
                    plan.counts["wrap_adjusted"] += 1
                    plan.warn("WRAP_ADJUSTED", f"{t.get('id') or 'text'}: would fit on "
                              f"{predicted} lines in PowerPoint vs {lines} in the design; "
                              "narrowed to keep the design's line breaks")
                    break

    # First baseline, placed where PowerPoint draws it (place_first_line).
    # Chrome's first baseline is the first line's text-box top (ink top) plus
    # the face's ascent; the first line's text-box height (firstLineH) is
    # ascent + descent, so the ascent's share of it comes from the font.
    flh = t.get("firstLineH") or fs * 1.15
    big = max((r for r in paras[0] if r["t"].strip()), key=lambda r: r["size_px"], default=r0)
    y, L, dy_lo = place_first_line(ink["y"], flh, big, lh, single)
    h = max(lh * lines, ink["h"]) + 2

    plan.ops.append({"op": "text", "x": x, "y": y, "w": w, "h": h, "align": align,
                     "line_height_px": L, "_lh_design": lh, "wrap": wrap, "paragraphs": paras,
                     "name": name or (t.get("id") or None), "_order": order_key,
                     "_ink": (ink["x"], ink["y"], ink["w"], ink["h"]), "_single": single,
                     "_lines": lines, "_id": t.get("id") or "", "_glyphs": t.get("glyphs") or [],
                     "_design_px": design_px, "_cased": cased, "_dy_lo": dy_lo})
    pad = top_pad_px(big, L)
    if pad:
        plan.ops[-1]["pad_top_px"] = pad
    plan.counts["texts"] += 1


# --- pictures --------------------------------------------------------------

PICTURE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}


def _file_path(src: str) -> Path | None:
    """The file a file:// address points at, or None."""
    from urllib.parse import unquote, urlparse
    from urllib.request import url2pathname
    try:
        u = urlparse(src)
        if u.scheme.lower() != "file":
            return None
        if u.netloc:
            return Path(url2pathname(f"//{u.netloc}{u.path}"))
        return Path(url2pathname(u.path))
    except Exception:
        return None


def _object_offset(tok: str, room: float) -> float:
    """One object-position value as px from the box's left/top edge."""
    tok = (tok or "").strip().lower()
    named = {"left": 0.0, "top": 0.0, "center": 0.5, "right": 1.0, "bottom": 1.0}
    if tok in named:
        return named[tok] * room
    if tok.endswith("%"):
        return float(tok[:-1]) / 100 * room
    return _px(tok)


def _plan_picture(plan: "Plan", pc: dict) -> None:
    """A picture from a file, drawn as a native picture: the part of it the
    browser showed (object-fit and object-position as a crop) and its shape
    (round corners, or an oval for 50%). Before 2026-10-13 every picture went
    to the agent, which drew exactly this by hand."""
    x, y, w, h = pc["x"], pc["y"], pc["w"], pc["h"]

    def agent(why):
        plan.fallbacks.append({"kind": "img", "id": pc.get("id") or "", "x": x, "y": y, "w": w,
                               "h": h, "reason": f"img element ({why})", "_order": pc["order"],
                               "_html": pc.get("html")})
    if w < 1 or h < 1:
        return
    path = _file_path(pc.get("src") or "")
    if path is None or not path.is_file():
        return agent("picture file not found")
    if path.suffix.lower() not in PICTURE_TYPES:
        return agent(f"{path.suffix or 'this kind of'} picture can't be placed natively")
    nw, nh = float(pc["nw"]), float(pc["nh"])
    fit = (pc.get("fit") or "fill").lower()
    if fit == "cover":
        k = max(w / nw, h / nh)
        dw, dh = nw * k, nh * k
    elif fit in ("contain", "scale-down"):
        k = min(w / nw, h / nh)
        if fit == "scale-down":
            k = min(1.0, k)
        dw, dh = nw * k, nh * k
    elif fit == "none":
        dw, dh = nw, nh
    else:                                 # fill: stretched to the box
        dw, dh = w, h
    toks = (pc.get("pos") or "50% 50%").split()
    if len(toks) == 1:
        toks.append("50%")
    ox = _object_offset(toks[0], w - dw)
    oy = _object_offset(toks[1], h - dh)
    # the picture as drawn, cut to the box
    px0, py0, px1, py1 = x + ox, y + oy, x + ox + dw, y + oy + dh
    vx0, vy0, vx1, vy1 = max(x, px0), max(y, py0), min(x + w, px1), min(y + h, py1)
    if vx1 - vx0 < 1 or vy1 - vy0 < 1:
        return
    crop = [(vx0 - px0) / dw, (vy0 - py0) / dh, (px1 - vx1) / dw, (py1 - vy1) / dh]
    radii = [_px(r_, min(w, h)) for r_ in pc.get("radii") or []]
    rad = max(radii or [0.0])
    geom, ratio = "rect", None
    if rad > 0.5:
        if len({round(r_) for r_ in radii}) > 1:
            return agent("picture with mixed corner radii")
        if (vx1 - vx0) < w - 0.5 or (vy1 - vy0) < h - 0.5:
            return agent("rounded picture that doesn't fill its box")
        if rad >= min(w, h) / 2 - 0.5:
            geom = "ellipse"
        else:
            geom, ratio = "roundRect", min(0.5, rad / max(1.0, min(w, h)))
    if pc.get("shadow"):
        plan.kill["shadows_stripped"] += 1
    plan.ops.append({"op": "picture", "path": str(path), "x": vx0, "y": vy0, "w": vx1 - vx0,
                     "h": vy1 - vy0, "crop": [round(c, 6) for c in crop], "geom": geom,
                     "radius_ratio": ratio, "name": pc.get("id") or None,
                     "_order": pc["order"], "_box": (vx0, vy0, vx1 - vx0, vy1 - vy0)})
    plan.counts["pictures"] = plan.counts.get("pictures", 0) + 1


# --- svg -------------------------------------------------------------------

def _plan_icon(plan: Plan, ic: dict) -> None:
    """One library icon: inserted at the same spot and color as the sketch
    drew it (the picture fitted into its box, keeping its shape, as the
    browser did). The name goes through icon_names, the one lookup: an old
    name becomes its plain name; an unknown name keeps its step as a labeled
    placeholder, and its ICON_UNKNOWN warning (from icon_svg.draw_icons)
    blocks the review card."""
    import icon_svg
    import icon_names
    name = ic.get("name") or ""
    known = icon_names.lookup(name)
    if known is not None and known["name"] != name.strip():
        name = known["name"]
    x, y, w, h = ic["x"], ic["y"], ic["w"], ic["h"]
    if w < 1 or h < 1:
        plan.warn("ICON_NO_SIZE", f'icon "{name}" has no size on the page; not inserted')
        return
    svg = icon_svg.svg_for(name) if ic.get("drawn") == "icon" else None
    m = re.search(r'viewBox="0 0 (\d+) (\d+)"', svg or "")
    if m:
        vw, vh = int(m.group(1)), int(m.group(2))
        k = min(w / vw, h / vh)
        fw, fh = vw * k, vh * k
        x, y, w, h = x + (w - fw) / 2, y + (h - fh) / 2, fw, fh
    col = parse_color(ic.get("color"))
    plan.ops.append({"op": "icon", "name": name, "x": x, "y": y, "w": w, "h": h,
                     "color": col[0] if col else "000000",
                     "placeholder": not bool(m), "_order": ic["order"]})
    plan.counts["icons"] = plan.counts.get("icons", 0) + 1


# PowerPoint's arrowhead sizes, as multiples of the line's width.
_HEAD_SIZES = (("sm", 2.0), ("med", 3.0), ("lg", 5.0))
HEAD_COLOR_LIMIT = 30.0     # RGB distance: an arrowhead in another color stays the agent's


def _head_size(px: float, line_w: float) -> str:
    k = px / max(0.5, line_w)
    return min(_HEAD_SIZES, key=lambda s_: abs(math.log(k / s_[1])) if k > 0 else 9)[0]


def _head_fits(hd: dict, line_w: float) -> bool:
    """Is a PowerPoint arrowhead close to this size (small 2x the line's
    width, large 5x)? A 16 px head on a 2 px line (8x) came out half size."""
    lw = max(0.5, line_w)
    return all(1.6 <= hd[k] / lw <= 6.0 for k in ("len", "wid"))


def _arrow_heads_ok(s: dict) -> bool:
    for it in s.get("items") or []:
        for hd in (it.get("heads") or {}).values():
            if not hd:
                continue
            fill, stroke = parse_color(hd.get("fill")), parse_color(it.get("stroke"))
            if not fill or not stroke or _color_distance(fill[0], stroke[0]) > HEAD_COLOR_LIMIT:
                return False
    return True


def _plan_svg(plan: Plan, s: dict) -> None:
    if not s["unsupported"] and not _arrow_heads_ok(s):
        s = dict(s, unsupported=["arrowhead in another color than its line"])
    if s["unsupported"]:
        plan.fallbacks.append({"kind": "svg", "id": s.get("id") or "", "x": s["x"], "y": s["y"],
                               "w": s["w"], "h": s["h"],
                               "reason": "svg with " + ", ".join(sorted(set(s["unsupported"]))),
                               "_order": s["order"], "_rects": s.get("rects") or None,
                               "_html": s.get("html")})
        return
    for i, it in enumerate(s["items"]):
        k = s["order"] + 0.0001 * (i + 1)
        fill = parse_color(it.get("fill"))
        stroke = parse_color(it.get("stroke"))
        line = {"color": stroke[0], "w": it["sw"], "dash": "dashed" if it.get("dash") else None} \
            if stroke and it.get("sw", 0) > 0 else None
        if it["k"] == "poly":
            if len(it["pts"]) == 2 and not it["closed"]:
                pts = [list(it["pts"][0]), list(it["pts"][1])]
                heads = it.get("heads") or {}
                if line and any(heads.values()):
                    # PowerPoint puts the arrow's tip on the line's end: move
                    # each end out to where the design's tip is
                    dx, dy = pts[1][0] - pts[0][0], pts[1][1] - pts[0][1]
                    n = math.hypot(dx, dy) or 1.0
                    ux, uy = dx / n, dy / n
                    line = dict(line)
                    own = []          # arrowheads PowerPoint can't size: drawn as triangles
                    for which, key, j, sgn in (("start", "head", 0, -1), ("end", "tail", 1, 1)):
                        hd = heads.get(which)
                        if not hd:
                            continue
                        e = float(hd.get("ext") or 0)
                        tip = [pts[j][0] + sgn * ux * e, pts[j][1] + sgn * uy * e]
                        if _head_fits(hd, line["w"]):
                            pts[j] = tip
                            line[key] = {"type": "triangle", "w": _head_size(hd["wid"], line["w"]),
                                         "len": _head_size(hd["len"], line["w"])}
                        else:
                            bx, by = tip[0] - sgn * ux * hd["len"], tip[1] - sgn * uy * hd["len"]
                            hw = hd["wid"] / 2
                            own.append([tip, [bx - uy * hw, by + ux * hw], [bx + uy * hw, by - ux * hw]])
                    plan.counts["native_arrowheads"] = plan.counts.get("native_arrowheads", 0) + 1
                else:
                    own = []
                name = (s.get("id") or None) if len(s["items"]) == 1 else None
                plan.ops.append({"op": "connector", "points": pts, "line": line, "_order": k,
                                 "name": name})
                for n_, tri in enumerate(own):
                    plan.ops.append({"op": "freeform", "points": tri, "closed": True,
                                     "fill": line["color"], "alpha": 1.0, "line": None,
                                     "_order": k + 0.00001 * (n_ + 1),
                                     "name": f"{name}-head-{n_ + 1}" if name else None})
            else:
                plan.ops.append({"op": "freeform", "points": it["pts"], "closed": it["closed"],
                                 "fill": fill[0] if fill and it["closed"] else None,
                                 "alpha": (fill[1] if fill else 1) * it.get("fillOp", 1),
                                 "line": line, "_order": k})
        elif it["k"] == "oval":
            plan.ops.append({"op": "shape", "kind": "oval", "x": it["x"], "y": it["y"],
                             "w": it["w"], "h": it["h"], "fill": fill[0] if fill else None,
                             "alpha": (fill[1] if fill else 1) * it.get("fillOp", 1),
                             "line": line, "_order": k})
        elif it["k"] == "text" and it["t"]:
            col = parse_color(it["color"])
            align = {"middle": "center", "end": "right"}.get(it["anchor"], "left")
            run = {"t": it["t"], "font": _face(it["ff"], it["fw"])[0], "size_px": it["fs"],
                   "bold": _face(it["ff"], it["fw"])[1], "italic": False,
                   "color": col[0] if col else None}
            # the text's box is Chrome's (ascent + descent tall): one line
            ty, L, dy_lo = place_first_line(it["y"], it["h"], run, it["h"], True)
            plan.ops.append({"op": "text", "x": it["x"], "y": ty, "w": it["w"] + 4, "h": it["h"],
                             "align": align, "line_height_px": L, "_lh_design": it["h"],
                             "wrap": False, "paragraphs": [[run]], "_order": k, "_dy_lo": dy_lo})
            pad = top_pad_px(run, L)
            if pad:
                plan.ops[-1]["pad_top_px"] = pad
        plan.counts["svg_items"] += 1


# --- clearance ---------------------------------------------------------------

# Clearance growth (owner's rule, 2026-10-08): a box grows only into empty
# space, by at most a few px, and keeps this gap to the next shape.
CLEARANCE_GAP_PX = 2.0


def _op_rect(o: dict) -> tuple[float, float, float, float] | None:
    """(x0, y0, x1, y1) of an op as drawn: a box's outer edge, a text's ink
    when known, a line's points."""
    if o.get("_box"):
        x, y, w, h = o["_box"]
    elif o.get("op") == "text" and o.get("_ink"):
        x, y, w, h = o["_ink"]
    elif all(k in o for k in ("x", "y", "w", "h")):
        x, y, w, h = o["x"], o["y"], o["w"], o["h"]
    elif o.get("points"):
        xs = [p[0] for p in o["points"]]
        ys = [p[1] for p in o["points"]]
        w_line = ((o.get("line") or {}).get("w") or 1) / 2
        return min(xs) - w_line, min(ys) - w_line, max(xs) + w_line, max(ys) + w_line
    else:
        return None
    return float(x), float(y), float(x + w), float(y + h)


def _free_space(plan: Plan, b: dict, side: str) -> float:
    """Empty px past box b's bottom ('bottom') or right ('right') edge before
    the next thing drawn there: a shape, line or text beyond the edge, the
    edge of a shape that holds b, the canvas edge. A shape that crosses the
    edge leaves no room. Text that starts inside b and runs past its edge is
    overhang, not a neighbor."""
    bx0, by0, bx1, by1 = _op_rect(b)
    if side == "bottom":
        lim = CANVAS_H - by1
        lo, hi, edge, a0, a1 = bx0, bx1, by1, 1, 3     # axis: y; span: x
    else:
        lim = CANVAS_W - bx1
        lo, hi, edge, a0, a1 = by0, by1, bx1, 0, 2     # axis: x; span: y
    # elements left to the agent are drawn there too
    others = list(plan.ops) + [dict(f, op="fallback") for f in plan.fallbacks
                               if f.get("w") and f.get("h")]
    for o in others:
        if o is b:
            continue
        r = _op_rect(o)
        if r is None:
            continue
        s0, s1 = (r[0], r[2]) if side == "bottom" else (r[1], r[3])
        if s1 <= lo + 0.5 or s0 >= hi - 0.5:
            continue                                   # not beside this edge
        start, end = r[a0], r[a1]
        if start >= edge - 0.5:
            lim = min(lim, start - edge)               # below / right of the edge
        elif end > edge + 0.5:
            if o.get("op") == "text":
                continue                               # overhanging text
            holds = (r[0] <= bx0 + 0.5 and r[1] <= by0 + 0.5 and r[2] >= bx1 - 0.5
                     and r[3] >= by1 - 0.5)
            lim = min(lim, end - edge) if holds else 0.0
    return max(0.0, lim)


def _apply_clearance(plan: Plan) -> None:
    """Text sitting right on the edge of the filled box behind it gets room.
    LibreOffice and PowerPoint can set text a few pixels off Chrome, enough to
    push the last line past a dark band's edge. Deliberately overhanging text
    (big numerals breaking out of a card) is left alone.

    Runs after the self-check, which compares the design as drawn (a grown box
    read as drawn wrong and went to the agent, 2026-10-08). The box grows only
    into empty space and keeps CLEARANCE_GAP_PX to the next shape: thin
    stacked boxes used to grow into each other. A box with no room keeps its
    design size."""
    boxes = [o for o in plan.ops if o["op"] == "shape" and o.get("_filled") and o.get("_box")]
    for t in (o for o in plan.ops if o["op"] == "text" and o.get("_ink")):
        ix, iy, iw, ih = t["_ink"]
        holders = [b for b in boxes if b["_order"] < t["_order"]
                   and b["_box"][0] - 1 <= ix and ix + iw <= b["_box"][0] + b["_box"][2] + 1
                   and b["_box"][1] - 1 <= iy and iy + ih <= b["_box"][1] + b["_box"][3] + 1]
        if not holders:
            continue
        b = min(holders, key=lambda q: q["_box"][2] * q["_box"][3])
        bx, by, bw, bh = b["_box"]
        gap_bottom = (by + bh) - (iy + ih)
        if 0 <= gap_bottom < 4:
            grow = min(6 - gap_bottom, _free_space(plan, b, "bottom") - CLEARANCE_GAP_PX)
            if grow >= 0.5:
                b["h"] += grow
                b["_box"] = (bx, by, bw, bh + grow)
                bh += grow
                plan.counts["boxes_grown_for_clearance"] += 1
            else:
                plan.counts["boxes_kept_no_room"] = plan.counts.get("boxes_kept_no_room", 0) + 1
        gap_right = (bx + bw) - (ix + iw)
        if t.get("_single") and t["align"] == "left" and 0 <= gap_right < 4:
            grow = min(6 - gap_right, _free_space(plan, b, "right") - CLEARANCE_GAP_PX)
            if grow >= 0.5:
                b["w"] += grow
                b["_box"] = (bx, by, bw + grow, bh)
                plan.counts["boxes_grown_for_clearance"] += 1
            else:
                plan.counts["boxes_kept_no_room"] = plan.counts.get("boxes_kept_no_room", 0) + 1


# ---------------------------------------------------------------------------
# Native tables and lists (owner's decision, 2026-10-08)
# ---------------------------------------------------------------------------
# An HTML <table> (or a grid marked data-native="table") becomes ONE native
# PowerPoint table, and a <ul>/<ol> (or a column marked data-native="list")
# becomes ONE text box with real bullets, instead of a rectangle and a text
# box per cell or per line. Each is a "group" with a chain of versions, best
# first; the self-check moves a group one step down its chain when the
# version on the slide renders differently from the design:
#   table: native table -> separate shapes (as before)
#   list:  one box with bullets -> one box, markers drawn as shapes
#          -> one text box per item (as before)
# Text is never lost on the way: every version holds all of it.

GRID_TOL = 1.5          # px: cell edges this close are the same grid line
TABLE_MAX_RADIUS = 4.0  # px: rounder corners than this can't be a native table

_BULLET_CHARS = {"disc": "•", "circle": "◦", "square": "▪"}
_AUTONUM = {"decimal": "arabicPeriod", "lower-alpha": "alphaLcPeriod",
            "lower-latin": "alphaLcPeriod", "upper-alpha": "alphaUcPeriod",
            "upper-latin": "alphaUcPeriod", "lower-roman": "romanLcPeriod",
            "upper-roman": "romanUcPeriod"}


def _g(rec: dict) -> dict:
    return rec.get("g") or rec.get("_g") or {}


def _cluster(vals: list[float], tol: float = GRID_TOL) -> list[float]:
    out: list[list[float]] = []
    for v in sorted(vals):
        if out and v - out[-1][-1] <= tol:
            out[-1].append(v)
        else:
            out.append([v])
    return [sum(c) / len(c) for c in out]


def _nearest(lines: list[float], v: float) -> tuple[int, float]:
    j = min(range(len(lines)), key=lambda k: abs(lines[k] - v))
    return j, abs(lines[j] - v)


def _live_side(s):
    """A CSS border side as {"color", "w", "dash"}, "hidden", or None."""
    if not s or s.get("w", 0) <= 0 or s.get("s") in (None, "none"):
        return None
    if s["s"] == "hidden":
        return "hidden"
    c = parse_color(s.get("c"))
    if not c:
        return None
    color = c[0] if c[1] >= 0.999 else _mix(c[0], "FFFFFF", c[1])
    return {"color": color, "w": float(s["w"]),
            "dash": s["s"] if s["s"] in ("dashed", "dotted") else None}


def _compose_fill(layers: list[str]):
    """Backgrounds from the bottom (table) to the top (cell) as one fill."""
    col = None
    for layer in layers:
        c = parse_color(layer)
        if not c:
            continue
        if c[1] >= 0.999 or col is None:
            col = (c[0], c[1])
        else:
            col = (_mix(c[0], col[0], c[1]), col[1])
    return col


def _scratch_text(t: dict):
    """Plan one text block on its own; the text op, or None (nothing drawn)."""
    sp = Plan()
    _plan_text(sp, t, t["order"])
    return sp.ops[0] if sp.ops else None


def _resolve_edge(cands):
    if not cands or any(c == "hidden" for c, _p in cands):
        return None
    c, _p = max(cands, key=lambda q: (round(q[0]["w"], 1), q[1]))
    return dict(c)


def _table_group(plan: "Plan", tb: dict, data: dict):
    """(table op, the ops it replaces) or (None, reason it stays shapes)."""
    k = str(tb["k"])
    name = tb.get("id") or f"table-{tb['k'] + 1}"
    # What sits on top of a cell (a pill, a numbered badge, an icon, a small
    # drawing) stays its own shape, drawn over the native table, exactly as
    # before; the text inside such a badge goes with it.
    decos = [b for b in data.get("boxes") or [] if _g(b).get("tbl") == k and not _g(b).get("part")]

    def _inside(t, b):
        ix, iy, iw, ih = t["ink"]["x"], t["ink"]["y"], t["ink"]["w"], t["ink"]["h"]
        return (b["order"] <= t["order"] and b["x"] - 1 <= ix and b["y"] - 1 <= iy
                and ix + iw <= b["x"] + b["w"] + 1 and iy + ih <= b["y"] + b["h"] + 1)

    texts = [t for t in data.get("texts") or [] if _g(t).get("tbl") == k]
    deco_texts = [t for t in texts if any(_inside(t, b) for b in decos)]
    texts = [t for t in texts if not any(t is d for d in deco_texts)]

    def _overlaps(t, b):
        ix, iy, iw, ih = t["ink"]["x"], t["ink"]["y"], t["ink"]["w"], t["ink"]["h"]
        return ix < b["x"] + b["w"] - 1 and b["x"] < ix + iw - 1 and \
            iy < b["y"] + b["h"] - 1 and b["y"] < iy + ih - 1
    if any(_overlaps(t, b) for t in texts for b in decos):
        return None, "a cell's text runs through its own badge or band"
    others =(data.get("svgs") or []) + (data.get("icons") or [])
    keep = {x["order"] for x in decos + deco_texts + others if _g(x).get("tbl") == k}
    alt = [o for o in plan.ops if _g(o).get("tbl") == k and o.get("_src") not in keep]
    if any(not _g(t).get("cell") for t in texts):
        return None, "text in the table sits outside its cells (a caption)"
    cells = tb.get("cells") or []
    if not cells:
        return None, "it has no cells"
    if tb.get("opacity", 1) < 0.99 or any(c.get("opacity", 1) < 0.99 for c in cells):
        return None, "it is see-through"
    if tb.get("radius", 0) > TABLE_MAX_RADIUS or any(c.get("radius", 0) > TABLE_MAX_RADIUS for c in cells):
        return None, "it has rounded corners"
    if any("gradient" in (c.get("bgimg") or "") or "url(" in (c.get("bgimg") or "") for c in cells):
        return None, "a cell has a gradient or picture background"
    if not alt:
        return None, "nothing in it is drawn"

    fills = {c["i"]: _compose_fill(c["bg"]) for c in cells}
    has_lines = any(_live_side(s) not in (None, "hidden") for c in cells for s in c["sides"])

    def lines_and_gaps(pairs):
        lines = _cluster([v for p in pairs for v in p])
        covered = [False] * max(0, len(lines) - 1)
        for lo, hi in pairs:
            a, _ = _nearest(lines, lo)
            b, _ = _nearest(lines, hi)
            for j in range(a, b):
                covered[j] = True
        gaps = [j for j, cv in enumerate(covered) if not cv]
        return lines, gaps

    xs, gx = lines_and_gaps([(c["x"], c["x"] + c["w"]) for c in cells])
    ys, gy = lines_and_gaps([(c["y"], c["y"] + c["h"]) for c in cells])
    tol = GRID_TOL
    for lines, gaps in ((xs, gx), (ys, gy)):
        if not gaps:
            continue
        widest = max(lines[j + 1] - lines[j] for j in gaps)
        if widest > GRID_TOL and (any(fills.values()) or has_lines):
            return None, "its cells are spaced apart"
        tol = max(tol, widest + GRID_TOL)
        for j in sorted(gaps, reverse=True):
            del lines[j + 1]
    if len(xs) < 2 or len(ys) < 2:
        return None, "its cells don't form a grid"
    nr, nc = len(ys) - 1, len(xs) - 1
    own = [[None] * nc for _ in range(nr)]
    span = {}
    for c in cells:
        (c0, e0), (c1, e1) = _nearest(xs, c["x"]), _nearest(xs, c["x"] + c["w"])
        (r0, e2), (r1, e3) = _nearest(ys, c["y"]), _nearest(ys, c["y"] + c["h"])
        if max(e0, e1, e2, e3) > tol or c1 <= c0 or r1 <= r0:
            return None, "its cells don't form a grid"
        for r in range(r0, r1):
            for cc in range(c0, c1):
                if own[r][cc] is not None:
                    return None, "its cells overlap"
                own[r][cc] = c["i"]
        span[c["i"]] = (r0, r1, c0, c1)

    # Borders, one per grid edge. In CSS's collapsed model the widest border
    # wins and a cell's own beats its row's, then its group's, then the table's.
    html_sep = tb.get("tag") == "table" and not tb.get("collapse")
    Hs: dict = {}
    Vs: dict = {}

    def add(d, key, side, prio):
        ls = _live_side(side)
        if ls:
            d.setdefault(key, []).append((ls, prio))

    by_i = {c["i"]: c for c in cells}
    for c in cells:
        r0, r1, c0, c1 = span[c["i"]]
        T_, R_, B_, L_ = c["sides"]
        for cc in range(c0, c1):
            add(Hs, (r0, cc), T_, 3)
            add(Hs, (r1, cc), B_, 3)
        for rr in range(r0, r1):
            add(Vs, (c0, rr), L_, 3)
            add(Vs, (c1, rr), R_, 3)
        if html_sep:
            continue
        for sides, rect, prio in ((c.get("rowSides"), c.get("rowRect"), 2),
                                  (c.get("secSides"), c.get("secRect"), 1)):
            if not sides or not rect:
                continue
            if abs(rect["y"] - c["y"]) <= tol:
                for cc in range(c0, c1):
                    add(Hs, (r0, cc), sides[0], prio)
            if abs(rect["y"] + rect["h"] - c["y"] - c["h"]) <= tol:
                for cc in range(c0, c1):
                    add(Hs, (r1, cc), sides[2], prio)
            if abs(rect["x"] - c["x"]) <= tol:
                for rr in range(r0, r1):
                    add(Vs, (c0, rr), sides[3], prio)
            if abs(rect["x"] + rect["w"] - c["x"] - c["w"]) <= tol:
                for rr in range(r0, r1):
                    add(Vs, (c1, rr), sides[1], prio)
    ts = tb.get("sides") or [None] * 4
    for cc in range(nc):
        add(Hs, (0, cc), ts[0], 0)
        add(Hs, (nr, cc), ts[2], 0)
    for rr in range(nr):
        add(Vs, (0, rr), ts[3], 0)
        add(Vs, (nc, rr), ts[1], 0)

    table_fill = _compose_fill([cells[0]["bg"][0]])
    slots = []
    for r in range(nr):
        row = []
        for cc in range(nc):
            o = own[r][cc]
            f = fills.get(o) if o is not None else table_fill
            row.append({
                "fill": f[0] if f else None, "alpha": f[1] if f else 1.0,
                "L": _resolve_edge(Vs.get((cc, r))) if cc == 0 or own[r][cc - 1] != o or o is None else None,
                "R": _resolve_edge(Vs.get((cc + 1, r))) if cc == nc - 1 or own[r][cc + 1] != o or o is None else None,
                "T": _resolve_edge(Hs.get((r, cc))) if r == 0 or own[r - 1][cc] != o or o is None else None,
                "B": _resolve_edge(Hs.get((r + 1, cc))) if r == nr - 1 or own[r + 1][cc] != o or o is None else None})
        slots.append(row)

    # Cell text: planned exactly as a text box would be, then turned into the
    # cell's inner margins, so the letters land where Chrome put them.
    texts_by = {}
    for t in texts:
        texts_by.setdefault(_g(t)["cell"], []).append(t)
    out_cells, cmp_ops = [], []
    for c in cells:
        r0, r1, c0, c1 = span[c["i"]]
        X0, X1, Y0, Y1 = xs[c0], xs[c1], ys[r0], ys[r1]
        ts_ = sorted((t for t in texts_by.get(f"{k}:{c['i']}", [])
                      if not (t.get("covered") or t.get("clipped"))), key=lambda t: t["order"])
        blocks = [o for o in (_scratch_text(t) for t in ts_) if o is not None]
        rec = {"r": r0, "c": c0, "rs": r1 - r0, "cs": c1 - c0, "paragraphs": [],
               "mar": [0, 0, 0, 0], "align": "left", "line_height_px": 0}
        label = f"{name} r{r0 + 1}c{c0 + 1}"
        if blocks:
            for p_, q_ in zip(blocks, blocks[1:]):
                if q_["_ink"][1] < p_["_ink"][1] + p_["_ink"][3] - 2:
                    return None, "a cell holds pieces of text side by side"
            aligns = {o["align"] for o in blocks}
            if len(aligns) > 1:
                return None, "a cell mixes text alignments"
            align = aligns.pop()
            wrap = any(o["wrap"] for o in blocks)
            x = min(o["x"] for o in blocks)
            w = max(o["x"] + o["w"] for o in blocks) - x
            if not wrap:
                # one line each: give the text all the room the cell has, so a
                # renderer that sets the font a hair wider doesn't break a line
                if align == "right":
                    mar_l, mar_r = 0.0, X1 - (x + w)
                elif align == "center":
                    cx = x + w / 2
                    half = min(cx - X0, X1 - cx)
                    mar_l, mar_r = cx - half - X0, X1 - cx - half
                else:
                    mar_l, mar_r = x - X0, 0.0
            else:
                mar_l, mar_r = x - X0, X1 - (x + w)
            mar_t = blocks[0]["y"] - Y0
            if min(mar_l, mar_r) < -2:
                return None, "text runs past the edge of its cell"
            if mar_t < -3:
                return None, "text sits too close to the top of its cell"
            mar_l, mar_r, mar_t = max(0.0, mar_l), max(0.0, mar_r), max(0.0, mar_t)
            paragraphs, props, prev = [], [], None
            for o in blocks:
                spc = 0.0 if prev is None else \
                    o["y"] - (prev["y"] + max(1, prev.get("_lines") or 1) * prev["line_height_px"])
                if spc < -2:
                    return None, "a cell's lines of text overlap"
                ml = (o["x"] - x) if align in ("left", "justify") else 0.0
                mr = (x + w - o["x"] - o["w"]) if (align == "right" or (wrap and align != "center")) else 0.0
                avail = X1 - X0 - mar_l - mar_r - ml - mr
                pred = _wrapped_lines(o["paragraphs"], avail)
                want = max(1, o.get("_lines") or 1) if o["wrap"] else len(o["paragraphs"])
                if pred is not None and pred > want:
                    return None, "cell text would break onto more lines than in the design"
                for pi, para in enumerate(o["paragraphs"]):
                    paragraphs.append(para)
                    props.append({"line_height_px": o["line_height_px"], "marL_px": ml,
                                  "marR_px": mr, "indent_px": 0.0,
                                  "space_before_px": max(0.0, spc) if pi == 0 else 0.0})
                prev = o
            total = blocks[-1]["y"] + max(1, blocks[-1].get("_lines") or 1) * \
                blocks[-1]["line_height_px"] - blocks[0]["y"]
            if mar_t + total > (Y1 - Y0) + 1.0:
                return None, "cell text is taller than its row"
            rec.update(paragraphs=paragraphs, mar=[mar_l, mar_t, mar_r, 0.0], align=align,
                       line_height_px=blocks[0]["line_height_px"])
            if len(blocks) > 1:
                rec["para"] = props
            for n_, o in enumerate(blocks):
                cmp_ops.append(dict(o, name=label + (f" text {n_ + 1}" if len(blocks) > 1 else "")))
        out_cells.append(rec)
        f = fills.get(c["i"])
        cmp_ops.append({"op": "shape", "kind": "rect", "x": X0, "y": Y0, "w": X1 - X0,
                        "h": Y1 - Y0, "fill": f[0] if f else None, "name": label + " cell",
                        "_filled": bool(f), "_cell": True})
    x0, y0 = xs[0], ys[0]
    op = {"op": "table", "x": x0, "y": y0, "w": xs[-1] - x0, "h": ys[-1] - y0,
          "cols": [xs[j + 1] - xs[j] for j in range(nc)],
          "rows": [ys[j + 1] - ys[j] for j in range(nr)],
          "slots": slots, "cells": out_cells, "name": name,
          "_order": min(o["_order"] for o in alt), "_box": (x0, y0, xs[-1] - x0, ys[-1] - y0),
          "_filled": any(fills.values()), "_cmp": cmp_ops}
    return op, alt


def _strip_marker_runs(t: dict) -> dict:
    """Text block t without the runs of its marker (an inline ::before glyph),
    with the glyph boxes and ink trimmed to match."""
    if not any(r.get("mk") for r in t["runs"] if not r.get("br")):
        return t
    glyphs = t.get("glyphs") or []
    keep_runs, keep_g, gi = [], [], 0
    for r in t["runs"]:
        if r.get("br"):
            keep_runs.append(r)
            continue
        n = int(r.get("gc") or 0)
        if not r.get("mk"):
            keep_runs.append(r)
            keep_g += glyphs[gi:gi + n]
        gi += n
    q = dict(t, runs=keep_runs, glyphs=keep_g)
    if keep_g:
        x0 = min(g_[0] for g_ in keep_g)
        x1 = max(g_[2] for g_ in keep_g)
        q["ink"] = dict(t["ink"], x=x0, w=x1 - x0)
    q["listItem"] = False
    return q


_GLYPH_BOX: dict = {}


def _glyph_box(ch: str, face: str = "Arial"):
    """(left bearing, ink width) of one glyph in em, or None."""
    key = (ch, face)
    if key not in _GLYPH_BOX:
        val = None
        path = _face_path(face, False)
        if path:
            try:
                from fontTools.ttLib import TTFont
                f = TTFont(path, lazy=True)
                gname = (f.getBestCmap() or {}).get(ord(ch))
                if gname and "glyf" in f:
                    gl = f["glyf"][gname]
                    upm = f["head"].unitsPerEm
                    if getattr(gl, "numberOfContours", 0):
                        val = (gl.xMin / upm, (gl.xMax - gl.xMin) / upm)
            except Exception:
                val = None
        _GLYPH_BOX[key] = val
    return _GLYPH_BOX[key]


_SHOT_CACHE: dict = {}


def _ink_box(data: dict, x0: float, y0: float, x1: float, y1: float, bg=None,
             near: tuple | None = None):
    """(x, y, w, h) of the marks inside a region of the design's picture, or None.

    bg: the background color there ('RRGGBB'), when known; else the median of
    the region's border. near: (x, y) on the slide; only the connected mark
    closest to it counts (a letter's smoothed edge beside a bullet is not the
    bullet)."""
    import io
    import numpy as np
    from PIL import Image
    shot = data.get("_shot")
    if not shot:
        return None
    if _SHOT_CACHE.get("key") is not shot:          # one design at a time
        _SHOT_CACHE.update(key=shot, arr=np.asarray(
            Image.open(io.BytesIO(shot)).convert("RGB"), dtype=np.float32))
    arr = _SHOT_CACHE["arr"]
    h, w = arr.shape[:2]
    x0, y0 = max(0, int(math.floor(x0))), max(0, int(math.floor(y0)))
    x1, y1 = min(w, int(math.ceil(x1))), min(h, int(math.ceil(y1)))
    if x1 - x0 < 2 or y1 - y0 < 2:
        return None
    crop = arr[y0:y1, x0:x1]
    if bg is not None:
        bgv = np.array([int(bg[k:k + 2], 16) for k in (0, 2, 4)], dtype=np.float32)
    else:
        ring = np.concatenate([crop[0], crop[-1], crop[:, 0], crop[:, -1]])
        bgv = np.median(ring, axis=0)
    mask = np.linalg.norm(crop - bgv, axis=2) > 60
    if near is not None and mask.any():
        mask = _blob_nearest(mask, near[0] - x0, near[1] - y0)
    ys, xs = np.nonzero(mask)
    if not len(xs):
        return None
    return (float(x0 + xs.min()), float(y0 + ys.min()),
            float(xs.max() - xs.min() + 1), float(ys.max() - ys.min() + 1))


def _blob_nearest(mask, cx: float, cy: float):
    """The connected part of mask (8-neighbors) with the pixel nearest (cx, cy)."""
    import numpy as np
    ys, xs = np.nonzero(mask)
    k = int(np.argmin((xs - cx) ** 2 + (ys - cy) ** 2))
    h, w = mask.shape
    keep = np.zeros_like(mask)
    todo = [(int(ys[k]), int(xs[k]))]
    keep[todo[0]] = True
    while todo:
        y, x = todo.pop()
        for dy in (-1, 0, 1):
            for dx in (-1, 0, 1):
                ny, nx = y + dy, x + dx
                if 0 <= ny < h and 0 <= nx < w and mask[ny, nx] and not keep[ny, nx]:
                    keep[ny, nx] = True
                    todo.append((ny, nx))
    return keep


# A disc, circle or square bullet read off the design's picture must be
# between these shares of the text size; outside, the reading is wrong (a
# strip of the page beside a dark panel read as one 1,200% bullet,
# 2026-10-13) and the browser's own marker size and place are used instead.
BULLET_MIN_PCT, BULLET_MAX_PCT = 60.0, 200.0


def _fill_under(data: dict | None, px: float, py: float, before) -> dict | None:
    """The smallest filled box of the design under point (px, py), painted
    before `before` (an order number), or None."""
    best = None
    for b in (data or {}).get("boxes") or []:
        if b.get("order") is not None and before is not None and b["order"] >= before:
            continue
        if not (b["x"] <= px <= b["x"] + b["w"] and b["y"] <= py <= b["y"] + b["h"]):
            continue
        if not (parse_color(b.get("bg")) or (b.get("bgimg") and "gradient" in b["bgimg"])):
            continue
        if best is None or b["w"] * b["h"] < best["w"] * best["h"]:
            best = b
    return best


def _bg_hex(box: dict | None, data: dict | None) -> str:
    """The color behind a list marker: its filled box's, else the slide's."""
    for b in ([box] if box else []) + [b for b in (data or {}).get("boxes") or [] if b.get("canvas")]:
        col = parse_color(b.get("bg")) or (_gradient_mid(b.get("bgimg")) if b.get("bgimg") else None)
        if col:
            return col[0] if col[1] >= 0.999 else _mix(col[0], "FFFFFF", col[1])
    return "FFFFFF"


def _bullet_for(it: dict, op: dict, marker_texts: dict, data: dict | None = None,
                warn: list | None = None):
    """The PowerPoint bullet for one list item: (spec, bullet x, region the
    marker covers in the design), "keep" when the marker has to stay a shape,
    or None when the item has no marker."""
    m = it.get("marker") or {}
    kind = m.get("type")
    runs = [r for p in op["paragraphs"] for r in p if r["t"].strip()]
    text_px = runs[0]["size_px"] if runs else 16.0
    if kind == "native":
        col = parse_color(m.get("color"))
        style = m.get("style") or ""
        pct = 100.0 * float(m.get("fs") or text_px) / text_px
        # Where the marker's ink is, read off the design's own picture: the
        # browser does not report where it drew a list marker. Read only the
        # marker's own space first, against the color actually behind it
        # (its filled box's), keeping the mark nearest the marker; widen
        # toward the left only inside that box. The old wide strip crossed a
        # dark panel's edge and read the page beside it as the bullet, and
        # it caught the first letter's smoothed edge ("A": 404%).
        fs = float(m.get("fs") or text_px)
        right = m["x"] + m["w"]
        cy = m["y"] + m["h"] / 2
        holder = _fill_under(data, m["x"] + m["w"] / 2, cy, op.get("_order"))
        bg = _bg_hex(holder, data)
        lo = (holder["x"] + 1) if holder else 0.0
        ink = None
        if data is not None:
            ink = _ink_box(data, max(lo, m["x"]), m["y"] - 2, right - 2, m["y"] + m["h"] + 2,
                           bg=bg, near=(m["x"] + m["w"] / 2, cy))
            if ink is None:
                ink = _ink_box(data, max(lo, right - max(m["w"], 3 * fs)), m["y"] - 2, right - 2,
                               m["y"] + m["h"] + 2, bg=bg, near=(right - fs / 2, cy))
        why = None
        lbox = it.get("_L")
        if style in _AUTONUM:
            spec = {"autonum": _AUTONUM[style], "start": int(m.get("value") or 1)}
            bx = ink[0] if ink else m["x"]
        elif style in _BULLET_CHARS:
            spec = {"char": _BULLET_CHARS[style], "font": "Arial"}
            gb = _glyph_box(spec["char"])
            bx = m["x"]
            if ink and gb:
                size_px = ink[2] / gb[1]
                got = 100.0 * size_px / text_px
                if not BULLET_MIN_PCT <= got <= BULLET_MAX_PCT:
                    why = f"read as {got:.0f}% of the text size"
                else:
                    pct = got
                    bx = ink[0] - gb[0] * size_px
        else:
            return "unsupported"
        if why is None and ink:
            if lbox and bx < lbox[0] - 1:
                why = f"read {lbox[0] - bx:.0f} px left of the list"
            elif holder and not (holder["x"] - 1 <= ink[0] and ink[0] + ink[2] <= holder["x"] + holder["w"] + 1):
                why = "read outside the panel the list sits on"
        if why:
            # the browser's own marker place and size instead
            bx, pct = m["x"], 100.0 * fs / text_px
            ink = None
            if warn is not None:
                warn.append(f"item {it.get('i', 0) + 1}: the bullet {why}; used the browser's marker "
                            f"position and size instead")
        spec.update(color=col[0] if col else None, pct=pct)
        region = (ink[0], ink[1], ink[2], ink[3]) if ink else (m["x"], m["y"], max(1.0, m["w"]), m["h"])
        return spec, bx, region
    if kind == "glyph":
        c = (m.get("char") or "").strip()
        boxed = parse_color(m.get("bg")) or any(_live_side(s) not in (None, "hidden")
                                                 for s in m.get("sides") or [])
        if boxed:
            return "keep"
        tb = marker_texts.get(f"{it['_k']}:{it['i']}")
        col = parse_color(m.get("color"))
        pct = 100.0 * float(m.get("fs") or text_px) / text_px
        num = re.fullmatch(r"(\d{1,2})([.)])", c)
        if num and not c.startswith("0"):
            spec = {"autonum": "arabicPeriod" if num.group(2) == "." else "arabicParenR",
                    "start": int(num.group(1))}
        elif len(c) == 1 and not c.isalnum():
            face, _b = _face(m.get("ff") or "", m.get("fw"))
            spec = {"char": c, "font": face or "Arial"}
        else:
            return "keep"
        spec.update(color=col[0] if col else None, pct=pct)
        mx = tb["ink"]["x"] if tb else m["x"]
        if m.get("inline") and it.get("_mk_x") is not None:
            mx = it["_mk_x"]
        return spec, mx, (m["x"], m["y"], max(1.0, m["w"]), max(1.0, m["h"]))
    if kind == "box":
        col = parse_color(m.get("bg"))
        bgimg = m.get("bgimg") or "none"
        w, h = float(m.get("w") or 0), float(m.get("h") or 0)
        if not col:
            if any(_live_side(s) not in (None, "hidden") for s in m.get("sides") or []):
                return "keep"
            return None                    # an empty ::before: no marker at all
        if bgimg != "none" or w < 1 or h < 1:
            return "keep"
        ratio = w / h
        if 0.7 <= ratio <= 1.43:
            ch = "●" if (m.get("radius") or 0) >= min(w, h) / 2 - 0.5 else "■"
        elif ratio >= 2.5:
            ch = "–"
        else:
            return "keep"
        gb = _glyph_box(ch)
        if not gb:
            return "keep"
        lsb, ink_w = gb
        size_px = w / ink_w                          # bullet font size that draws it w wide
        pct = 100.0 * size_px / text_px
        if not 25 <= pct <= 400:
            return "keep"
        spec = {"char": ch, "font": "Arial", "color": col[0], "pct": pct}
        return spec, m["x"] - lsb * size_px, (m["x"], m["y"], w, h)
    if kind == "icon":
        return "keep"
    return None


def _list_op(name: str, blocks: list, order) -> dict | None:
    """One text box for a whole list. blocks: (item, first block of the item,
    text op, bullet (spec, x, region) or None), in reading order."""
    x0 = min([b[2]["x"] for b in blocks] + [b[3][1] for b in blocks if b[3]])
    right = max(b[2]["x"] + b[2]["w"] for b in blocks)
    wrap = any(b[2]["wrap"] for b in blocks)
    paragraphs, props = [], []
    prev = None
    for it, first, op, bul in blocks:
        spc = 0.0 if prev is None else \
            op["y"] - (prev["y"] + max(1, prev.get("_lines") or 1) * prev["line_height_px"])
        if spc < -2:
            return None
        for pi, para in enumerate(op["paragraphs"]):
            pr = {"lvl": int(it.get("level") or 0), "marL_px": op["x"] - x0,
                  "marR_px": (right - (op["x"] + op["w"])) if wrap else 0.0,
                  "indent_px": 0.0, "space_before_px": max(0.0, spc) if pi == 0 else 0.0,
                  "line_height_px": op["line_height_px"]}
            if pi == 0 and first and bul:
                pr["indent_px"] = bul[1] - op["x"]
                pr["bullet"] = bul[0]
            props.append(pr)
            paragraphs.append(para)
        prev = op
    first_op, last_op = blocks[0][2], blocks[-1][2]
    inks = [b[2]["_ink"] for b in blocks]
    ix0, iy0 = min(i[0] for i in inks), min(i[1] for i in inks)
    ix1, iy1 = max(i[0] + i[2] for i in inks), max(i[1] + i[3] for i in inks)
    bottom = last_op["y"] + max(1, last_op.get("_lines") or 1) * last_op["line_height_px"]
    return {"op": "text", "x": x0, "y": first_op["y"], "w": right - x0,
            "h": bottom - first_op["y"] + 2, "align": "left",
            "line_height_px": first_op["line_height_px"], "wrap": wrap,
            "paragraphs": paragraphs, "para": props, "name": name, "_order": order,
            "_ink": (ix0, iy0, ix1 - ix0, iy1 - iy0), "_single": False,
            "_lines": sum(max(1, b[2].get("_lines") or 1) for b in blocks), "_id": name,
            "_glyphs": [q for b in blocks for q in (b[2].get("_glyphs") or [])],
            "_design_px": max(b[2].get("_design_px") or 0 for b in blocks),
            "_cased": any(b[2].get("_cased") for b in blocks), "_list": True,
            "_dy_lo": first_op.get("_dy_lo") or 0.0,
            **({"pad_top_px": first_op["pad_top_px"]} if first_op.get("pad_top_px") else {})}


def _list_group(plan: "Plan", L: dict, data: dict):
    """(chain of versions, best first) or (None, reason it stays as before)."""
    k = str(L["k"])
    name = L.get("id") or f"list-{L['k'] + 1}"
    items = L.get("items") or []
    if len(items) < 2:
        return None, None
    alt = [o for o in plan.ops if _g(o).get("list") == k
           and (o["op"] == "text" or _g(o).get("marker"))]
    if not alt:
        return None, None
    if any(_g(b).get("list") == k and _g(b).get("item") and not _g(b).get("marker")
           for b in data.get("boxes") or []):
        return None, "its items are drawn as panels or have their own borders"
    if any(_g(f).get("list") == k for f in (data.get("fallbacks") or []) + (data.get("pictures") or [])):
        return None, "it holds a picture or an element left to the agent"
    texts = [t for t in data.get("texts") or [] if _g(t).get("list") == k]
    if any(not _g(t).get("item") for t in texts):
        return None, "it holds text outside its items"
    by_item, marker_texts = {}, {}
    for t in texts:
        if _g(t).get("marker"):
            marker_texts[_g(t)["marker"]] = t
        else:
            by_item.setdefault(_g(t)["item"], []).append(t)
    plain_blocks, bullet_blocks = [], []
    bullets_ok = True
    for it in items:
        it = dict(it, _k=k, _L=(L["x"], L["y"], L["w"], L["h"]))
        ts = sorted((t for t in by_item.get(f"{k}:{it['i']}", [])
                     if not (t.get("covered") or t.get("clipped"))), key=lambda t: t["order"])
        first = True
        for t in ts:
            op_plain = _scratch_text(dict(t, listItem=False) if (it.get("marker") or {}).get("type") == "native" else t)
            stripped = _strip_marker_runs(t)
            if stripped is not t:
                mk_runs = [r for r in t["runs"] if r.get("mk") and r.get("first")]
                if mk_runs:
                    it["_mk_x"] = mk_runs[0]["first"]["x"]
            op_bul = _scratch_text(dict(stripped, listItem=False))
            if op_plain is None or op_bul is None:
                continue
            notes = []
            bul = _bullet_for(it, op_bul, marker_texts, data, notes) if first else None
            for n_ in notes:
                plan.warn("BULLET_FALLBACK", f"{name}: {n_}")
            if bul in ("keep", "unsupported"):
                bullets_ok = False
                if bul == "unsupported":
                    return None, "its marker style has no PowerPoint bullet"
                bul = None
            plain_blocks.append((it, first, op_plain, None))
            bullet_blocks.append((it, first, op_bul, bul))
            first = False
    if len(plain_blocks) < 2:
        return None, None
    for blocks in (plain_blocks,):
        for a, b in zip(blocks, blocks[1:]):
            if b[2]["_ink"][1] < a[2]["_ink"][1] + a[2]["_ink"][3] - 2:
                return None, "its items sit side by side"
    if any(b[2]["align"] not in ("left", "justify") for b in plain_blocks):
        return None, "its text is not left-aligned"
    order = min(o["_order"] for o in alt)
    markers = [o for o in alt if _g(o).get("marker")]
    inline_marks = any(r.get("mk") for t in texts for r in t["runs"] if not r.get("br"))
    plain = _list_op(name, plain_blocks, order)
    chain = []
    if bullets_ok and any(b[3] for b in bullet_blocks):
        native = _list_op(name, bullet_blocks, order)
        if native is not None:
            regions = [b[3][2] for b in bullet_blocks if b[3]]
            native["_cmp"] = [native] + [
                {"op": "shape", "kind": "rect", "x": rx, "y": ry, "w": rw, "h": rh,
                 "name": f"{name} marker {n + 1}", "_filled": True}
                for n, (rx, ry, rw, rh) in enumerate(regions)]
            chain.append([native])
    elif bullets_ok and plain is not None and not markers and not inline_marks:
        chain.append([plain])                 # no markers at all: one box, no bullets
    if plain is not None and (markers or inline_marks) and \
            not (chain and chain[0][0] is plain):
        chain.append([plain] + markers)
    if not chain:
        return None, "its items could not be set as one text box"
    chain.append(alt)
    return chain, None


def _register_group(plan: "Plan", kind: str, name: str, chain: list) -> None:
    gid = len(plan.groups)
    for v, ops in enumerate(chain):
        for o in ops:
            o["_gid"], o["_gv"] = gid, v
    plan.groups[gid] = {"kind": kind, "name": name, "chain": chain, "cur": 0}
    drop = {id(o) for o in chain[-1]}
    plan.ops = [o for o in plan.ops if id(o) not in drop] + list(chain[0])


def _plan_native_groups(plan: "Plan", data: dict) -> None:
    for tb in data.get("tables") or []:
        op, alt_or_why = _table_group(plan, tb, data)
        name = tb.get("id") or f"table-{tb['k'] + 1}"
        if op is None:
            if any(_g(o).get("tbl") == str(tb["k"]) for o in plan.ops):
                plan.warn("TABLE_KEPT_AS_SHAPES",
                          f"{name}: not drawn as a native table because {alt_or_why}; "
                          "drawn as separate shapes as before")
            continue
        _register_group(plan, "table", name, [[op], alt_or_why])
    for L in data.get("lists") or []:
        chain, why = _list_group(plan, L, data)
        name = L.get("id") or f"list-{L['k'] + 1}"
        if chain is None:
            if why:
                plan.warn("LIST_KEPT_AS_SEPARATE_BOXES",
                          f"{name}: not set as one text box because {why}; each item "
                          "drawn as its own text box as before")
            continue
        _register_group(plan, "list", name, chain)
    plan.ops.sort(key=lambda o: o["_order"])


_STEP_WARN = {
    "table": ("TABLE_NOT_NATIVE", "the native table rendered differently from the design "
              "({why}); drawn as separate shapes as before"),
    "list1": ("LIST_MARKERS_AS_SHAPES", "the bullets rendered differently from the design's "
              "markers ({why}); the text stays one box and the markers are drawn as shapes"),
    "list2": ("LIST_NOT_SINGLE_BOX", "the one-box list rendered differently from the design "
              "({why}); each item drawn as its own text box as before"),
}


def _step_group(plan: "Plan", gid: int, why: str) -> list[dict]:
    """Move a group one version down its chain; return the ops swapped in."""
    grp = plan.groups[gid]
    cur_ids = {id(o) for o in grp["chain"][grp["cur"]]}
    plan.ops = [o for o in plan.ops if id(o) not in cur_ids]
    grp["cur"] += 1
    new = grp["chain"][grp["cur"]]
    plan.ops += list(new)
    plan.ops.sort(key=lambda o: o["_order"])
    last = grp["cur"] == len(grp["chain"]) - 1
    key = "table" if grp["kind"] == "table" else ("list2" if last else "list1")
    code, text = _STEP_WARN[key]
    plan.warn(code, f"{grp['name']}: " + text.format(why=why))
    return list(new)


def native_counts(plan: "Plan") -> dict:
    n_tables = sum(1 for g_ in plan.groups.values() if g_["kind"] == "table" and g_["cur"] == 0)
    n_lists = sum(1 for g_ in plan.groups.values()
                  if g_["kind"] == "list" and g_["cur"] < len(g_["chain"]) - 1)
    return {"native_tables": n_tables, "single_box_lists": n_lists}


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------

HEAVY_TO_BOLD_JS = r"""() => {
  let n = 0;
  for (const el of document.querySelectorAll('*')) {
    const w = parseInt(getComputedStyle(el).fontWeight, 10);
    if (w >= 600 && w !== 700) { el.style.setProperty('font-weight', '700', 'important'); n++; }
  }
  return n;
}"""


def extract(page, html: Path, subtitle_as_shape: bool) -> dict:
    page.goto(html.resolve().as_uri())
    page.wait_for_load_state("load")
    try:
        page.evaluate("document.fonts.ready")
    except Exception:
        pass
    page.wait_for_timeout(100)
    # Library icons are drawn exactly as render_html drew them for the review,
    # so the boxes and the self-check's reference image include them.
    import icon_svg
    icon_warnings = icon_svg.draw_icons(page)
    # PowerPoint gets bold on/off only (see _face): draw every weight from 600
    # up at 700, so the boxes, wraps and the self-check's reference image are
    # those of the face that ships, not of a heavier face the deck cannot keep.
    page.evaluate(HEAVY_TO_BOLD_JS)
    # The design as the alarm compares it: without the fields the template
    # draws (title, footer, page number), since the script doesn't draw them.
    page.evaluate("""(sas) => { for (const el of document.querySelectorAll('[data-template-field]')) {
        if (sas && el.getAttribute('data-template-field') === 'subtitle') continue;
        el.dataset.alarmHidden = el.style.visibility; el.style.visibility = 'hidden'; } }""",
                  subtitle_as_shape)
    cx, cy = page.evaluate("""() => { const c = document.querySelector('.slide-canvas') || document.querySelector('.slide') || document.body;
        const r = c.getBoundingClientRect(); return [r.x, r.y]; }""")
    shot = page.screenshot(clip={"x": cx, "y": cy, "width": CANVAS_W, "height": CANVAS_H})
    page.evaluate("""() => { for (const el of document.querySelectorAll('[data-alarm-hidden]')) {
        el.style.visibility = el.dataset.alarmHidden; delete el.dataset.alarmHidden; } }""")
    data = page.evaluate(EXTRACT_JS, {"subtitleAsShape": subtitle_as_shape})
    data["_shot"] = shot
    data["icon_warnings"] = icon_warnings
    # @font-face rules that pointed an installed font at a file, removed by
    # the browser (scripts/_browser.py) before anything was measured.
    try:
        data["font_face_removed"] = sorted(set(
            page.evaluate("() => window.__slidelabFontFaceRemoved || []") or []))
    except Exception:
        data["font_face_removed"] = []
    return data


def plan_from_extract(data: dict, subtitle_as_shape: bool,
                      template_bg: str | None = None) -> tuple[Plan, dict]:
    """template_bg: the slide layout's own background ('RRGGBB'), when the
    slide belongs to a build. A canvas of that color is left to the template
    instead of drawn as a full-slide rectangle over the master's artwork."""
    plan = Plan()
    plan.template_bg = (template_bg or "").upper().lstrip("#") or None
    for fam in data.get("font_face_removed") or []:
        plan.warn("FONT_FACE_REMOVED",
                  f"the design pointed the installed font {fam!r} at a font file "
                  f"(@font-face); the rule was ignored, so the design was read in the "
                  f"installed font, the one the finished slide uses")
    fields = {k: v for k, v in (data.get("fields") or {}).items()
              if not (k == "subtitle" and subtitle_as_shape)}
    # Template fields the graft fills are title, subtitle, footer and page
    # number. Anything else marked data-template-field is not drawn: say so
    # rather than drop it silently. The section tag is drawn by finalize from
    # the brief (owner's decision, 2026-10-09), so the design's copy is left out.
    for k, v in (data.get("fields") or {}).items():
        if k in ("title", "subtitle", "footer", "page_number"):
            continue
        low = k.lower().replace("_", "-")
        if "section" in low or low in ("kicker", "eyebrow", "tag"):
            plan.warn("SECTION_TAG_FIELD",
                      f"the design marks {v!r} as template field {k!r}: Slide Lab draws the "
                      "section tag itself from the brief, above the title, so this copy is "
                      "not drawn. Do not draw or mark the section tag in the design")
        else:
            plan.warn("UNKNOWN_TEMPLATE_FIELD",
                      f"the design marks {v!r} as template field {k!r}, which no template "
                      "slot takes (only title, subtitle, footer, page_number): it is not drawn")
    cv = data.get("canvas") or {}
    if abs((cv.get("w") or CANVAS_W) - CANVAS_W) > 1 or abs((cv.get("h") or CANVAS_H) - CANVAS_H) > 1:
        # Not on the fixed 1280x720 canvas the design spec requires (an older
        # 1600x900 design, say). Every coordinate would be off; the agent
        # translates the whole slide instead of the script guessing a rescale.
        plan.fallbacks.append({"kind": "canvas", "id": "", "x": 0, "y": 0, "w": CANVAS_W,
                               "h": CANVAS_H, "reason": f"design is {cv.get('w'):.0f}x{cv.get('h'):.0f}, "
                                                        "not the 1280x720 canvas; translate the whole slide"})
        return plan, fields
    items = [(b["order"], "box", b) for b in data["boxes"]]
    items += [(t["order"], "text", t) for t in data["texts"]]
    items += [(s["order"], "svg", s) for s in data["svgs"]]
    items += [(ic["order"], "icon", ic) for ic in data.get("icons") or []]
    items += [(pc["order"], "picture", pc) for pc in data.get("pictures") or []]
    for w in data.get("icon_warnings") or []:
        plan.warn(w["code"], w["detail"])
    for f in data["fallbacks"]:
        plan.fallbacks.append(dict({k: f.get(k) for k in ("kind", "id", "x", "y", "w", "h", "reason")},
                                   _order=f.get("order"), _html=f.get("html")))
    for _, kind, it in sorted(items, key=lambda z: z[0]):
        n0 = len(plan.ops)
        _plan_item(plan, kind, it, subtitle_as_shape)
        for o in plan.ops[n0:]:
            o["_g"] = it.get("g") or {}
            o["_src"] = it.get("order")
    # Clearance growth (_apply_clearance) is NOT applied here: it runs after
    # the self-check, in translate_many, so the check compares the design as
    # drawn.
    plan.ops.sort(key=lambda o: o["_order"])
    plan.counts["pseudo_elements"] = int(data.get("pseudo_count") or 0)
    _check_pale_fills(plan)
    _check_arrow_ends(plan, data)
    _plan_native_groups(plan, data)
    return plan, fields


def _plan_item(plan: Plan, kind: str, it: dict, subtitle_as_shape: bool) -> None:
    if kind == "box":
        _plan_box(plan, it, it["order"])
    elif kind == "text":
        name = "subtitle" if it.get("tf") == "subtitle" and subtitle_as_shape else None
        _plan_text(plan, it, it["order"], name=name)
    elif kind == "icon":
        _plan_icon(plan, it)
    elif kind == "picture":
        _plan_picture(plan, it)
    else:
        _plan_svg(plan, it)


def _public_plan(plan: Plan) -> dict:
    out = {"version": TRANSLATOR_VERSION,
           "ops": [{k: v for k, v in o.items() if not k.startswith("_")} for o in plan.ops]}
    if plan.fallbacks:
        _fallback_stacking(plan)
        out["fallbacks"] = [{k: f.get(k) for k in ("id", "draw_under", "draw_after_n")}
                            for f in plan.fallbacks]
    return out


def _fallback_stacking(plan: Plan) -> None:
    """Where each element left to the agent sits in the design's stacking
    order, so html_emit.restack_fallbacks can put the agent's shapes there.
    The agent's code runs after the plan is drawn, so everything it drew
    landed on top: an arrow drawn over its label (2026-10-12).

    Every entry gets an id (the design's, else "fallback-N"), "draw_after_n"
    (how many plan steps come before it) and "draw_under" (the name of the
    first plan shape drawn after it, when it has one). Both are None for an
    element on top of everything."""
    ops = plan.ops
    for n, f in enumerate(plan.fallbacks, 1):
        if not f.get("id"):
            f["id"] = f"fallback-{n}"
        order = f.get("_order")
        if order is None:
            f["draw_after_n"] = f["draw_under"] = None
            continue
        later = [i for i, o in enumerate(ops) if o.get("_order") is not None and o["_order"] > order]
        if not later:
            f["draw_after_n"] = f["draw_under"] = None
            continue
        f["draw_after_n"] = later[0]
        f["draw_under"] = ops[later[0]].get("name") or None


NATIVE_TEMPLATE = '''# {name} -- translated by translate_html.py ({version}), deterministic.
# The plan is in {plan_name}; twins/html_emit.py draws it.
{fields_block}
{fallback_line}import json
import sys
from pathlib import Path

sys.path.insert(0, r"{skill}")
from twins.html_emit import build, restack_fallbacks  # noqa: E402

HERE = Path(__file__).resolve().parent
PLAN = json.loads((HERE / "{plan_name}").read_text(encoding="utf-8"))


def build_slide():
    prs, slide = build(PLAN)
    # --- FALLBACK ELEMENTS: drawn by slide-builder-translator (agent) ---
    # (none)
    # --- END FALLBACK ELEMENTS ---
    # Puts the agent's shapes where the design stacks them (by their names).
    restack_fallbacks(slide, PLAN)
    return prs


if __name__ == "__main__":
    build_slide().save(str(Path(__file__).resolve().with_suffix(".pptx")))
'''


def _fields_block(fields: dict) -> str:
    lines = ["# __template_fields__ = {"]
    for k in ("title", "subtitle", "footer", "page_number"):
        if k in fields and fields[k]:
            lines.append(f"#     {k!r}: {json.dumps(fields[k], ensure_ascii=False)},")
    lines.append("# }")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Converting a design again (item 8 step 5 and item 7, 2026-10-13). The agent's
# drawing (the code between the FALLBACK markers) is kept only when every
# element it draws is unchanged: the same elements, each with the same
# fingerprint (its HTML, its box, the bytes of any picture it shows), and the
# kept code still runs and names a shape after each element. Anything else and
# the slide waits on the agent again, with its previous drawing pasted into the
# block as comments to adapt. All or nothing: no element is kept on its own.
# The old script, plan and report are copied to slide_NN/_prev_translation/
# <time>/ first, every time (not slide_NN/_prev/, which means "an earlier
# design" to apply_qc_fix).
# ---------------------------------------------------------------------------

PREV_TRANSLATION_DIR = "_prev_translation"
_BLOCK_RE = re.compile(r"^[ \t]*# --- FALLBACK ELEMENTS[^\n]*\n(.*?)^[ \t]*# --- END FALLBACK ELEMENTS ---",
                       re.M | re.S)
PREV_OPEN = ("    # --- PREVIOUS DRAWING: the design changed after the agent drew this. Adapt it "
             "to the new design above this section, then delete the section ---\n")
PREV_CLOSE = "    # --- END PREVIOUS DRAWING ---\n"
_PREV_RE = re.compile(r"^[ \t]*# --- PREVIOUS DRAWING[^\n]*\n(.*?)^[ \t]*# --- END PREVIOUS DRAWING ---[^\n]*\n",
                      re.M | re.S)
PENDING_PREV_NOTE = (". The design changed after the agent drew it: its previous drawing is "
                     "pasted in the fallback block as comments; adapt it")


def _referenced_files(src: str, html: Path) -> list[Path]:
    """Local files an element's HTML points at (src, href, url(...))."""
    from urllib.parse import unquote
    refs = re.findall(r"""(?:src|href)\s*=\s*["']([^"']+)["']""", src or "", re.I)
    refs += re.findall(r"""url\(\s*["']?([^"')]+)["']?\s*\)""", src or "", re.I)
    out = []
    for ref in refs:
        ref = ref.strip().replace("&amp;", "&")
        if not ref or ref.startswith("#") or re.match(r"(?:https?|data):", ref, re.I):
            continue
        path = (_file_path(ref) if ref.lower().startswith("file:")
                else Path(html).parent / unquote(ref.split("?")[0]))
        try:
            if path is not None and path.is_file():
                out.append(path)
        except OSError:
            pass
    return out


def _fallback_fp(f: dict, html: Path) -> str:
    """One element left to the agent, as a fingerprint: what it is in the
    design (its HTML; for an element the self-check handed over, the plan
    step the design gave it), its box to the pixel, and the bytes of every
    picture it shows. A design-wide element (kind "canvas") is the whole file."""
    import hashlib
    h = hashlib.sha256(f"{f.get('kind')}|{f.get('id')}|".encode())
    if f.get("kind") != "self-check":
        # (a self-check entry's box is where the render differed: not used)
        h.update(",".join(str(round(float(f.get(k) or 0))) for k in ("x", "y", "w", "h")).encode())
    src = f.get("_html")
    if f.get("kind") == "canvas" or not src:
        h.update(_sha256(html).encode())
    else:
        h.update(re.sub(r"\s+", " ", src).strip().encode("utf-8", "replace"))
        for ref in _referenced_files(src, html):
            try:
                h.update(hashlib.sha256(ref.read_bytes()).hexdigest().encode())
            except OSError:
                pass
    return h.hexdigest()[:20]


def _drawing_lines(block: str) -> list[str]:
    """The agent's code in a fallback block (not the placeholder, not a
    pasted previous drawing)."""
    block = _PREV_RE.sub("", block or "")
    return [ln for ln in block.splitlines() if ln.strip() and not ln.strip().startswith("#")]


def _old_translation(emit_dir: Path, letter: str) -> dict | None:
    """What the last conversion of this option left: its script, plan, report,
    the agent's drawing in it (if any) and any previous drawing pasted in."""
    native = Path(emit_dir) / f"option_{letter}_native.py"
    if not native.exists():
        return None
    text = native.read_text(encoding="utf-8", errors="replace")
    old = {"text": text, "plan": None, "report": {}}
    for k, name in (("plan", f"option_{letter}_native.plan.json"),
                    ("report", f"option_{letter}_translation_report.json")):
        try:
            old[k] = json.loads((Path(emit_dir) / name).read_text(encoding="utf-8"))
        except Exception:
            pass
    m = _BLOCK_RE.search(text)
    block = m.group(1) if m else ""
    done = re.search(r"^# FALLBACK_DONE:.*$", text, re.M)
    old["block"] = block
    old["done_line"] = done.group(0) if done else ""
    old["drawn"] = bool(done) and bool(_drawing_lines(block))
    pm = _PREV_RE.search(block)
    old["prev_section"] = pm.group(0) if pm else ""
    return old


def _snapshot_translation(emit_dir: Path, letter: str) -> Path | None:
    """Copy the option's script, plan and report to _prev_translation/<time>/."""
    import shutil
    from datetime import datetime
    names = [f"option_{letter}_native.py", f"option_{letter}_native.plan.json",
             f"option_{letter}_translation_report.json"]
    have = [Path(emit_dir) / n for n in names if (Path(emit_dir) / n).exists()]
    if not have:
        return None
    base = Path(emit_dir) / PREV_TRANSLATION_DIR
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    dest, k = base / stamp, 1
    while (dest / names[0]).exists():
        dest, k = base / f"{stamp}-{k}", k + 1
    try:
        dest.mkdir(parents=True, exist_ok=True)
        for f in have:
            shutil.copy2(f, dest / f.name)
    except OSError:
        return None
    return dest


def _as_previous(block: str) -> str:
    """The agent's old code as a commented section to adapt."""
    lines = []
    for ln in (block or "").splitlines():
        if ln.strip() == "# (none)":
            continue
        lines.append("    # | " + (ln[4:] if ln.startswith("    ") else ln.lstrip()))
    return PREV_OPEN + "\n".join(lines) + ("\n" if lines else "") + PREV_CLOSE


def _kept_drawing_check(native: Path, fallbacks: list[dict]) -> str:
    """Run the script with the kept drawing. '' when it runs and (with more
    than one element) names a shape after every element; else why not."""
    import subprocess
    try:
        r = subprocess.run([sys.executable, str(native)], cwd=str(native.parent),
                           capture_output=True, text=True, timeout=300)
    except Exception as exc:
        return f"it could not be run ({type(exc).__name__})"
    if r.returncode != 0:
        tail = [ln for ln in (r.stderr or r.stdout or "").strip().splitlines() if ln.strip()]
        return "it no longer runs (" + (tail[-1][:200] if tail else f"exit {r.returncode}") + ")"
    if len(fallbacks) > 1:
        try:
            from pptx import Presentation
            names = [sh.name for sh in
                     Presentation(str(native.with_suffix(".pptx"))).slides[0].shapes]
        except Exception as exc:
            return f"its slide could not be read ({type(exc).__name__})"
        miss = [f["id"] for f in fallbacks
                if not any(n == f["id"] or n.startswith(f["id"] + "-") for n in names)]
        if miss:
            return "no shape is named after " + ", ".join(miss)
    return ""


def _drawing_decision(old: dict, plan: "Plan") -> str:
    """'' when the agent's old drawing can be kept for this plan, else why not."""
    old_fp = {f.get("id"): f.get("fp") for f in (old["report"] or {}).get("fallback") or []}
    new_fp = {f["id"]: f["fp"] for f in plan.fallbacks}
    if set(old_fp) != set(new_fp):
        new, gone = sorted(set(new_fp) - set(old_fp)), sorted(set(old_fp) - set(new_fp))
        return "; ".join(x for x in (
            ("new element(s) for the agent: " + ", ".join(new)) if new else "",
            ("no longer left to the agent: " + ", ".join(gone)) if gone else "") if x)
    if not all(old_fp.values()):
        return "the earlier conversion did not record what each element looked like"
    changed = sorted(k for k in new_fp if new_fp[k] != old_fp[k])
    return ("changed in the design: " + ", ".join(changed)) if changed else ""


def write_outputs(html: Path, emit_dir: Path, letter: str, plan: Plan, fields: dict,
                  visible_text: str) -> dict:
    stem = f"option_{letter}_native"
    plan_name = f"{stem}.plan.json"
    native = emit_dir / f"{stem}.py"
    old = _old_translation(emit_dir, letter)
    snap = _snapshot_translation(emit_dir, letter) if old else None
    if plan.fallbacks:
        _fallback_stacking(plan)
        for f in plan.fallbacks:
            f["fp"] = _fallback_fp(f, html)
    pub = _public_plan(plan)
    (emit_dir / plan_name).write_text(json.dumps(pub, indent=1), encoding="utf-8")
    id_lines = ""
    for f in plan.fallbacks:
        where = (f"stacked under {f['draw_under']!r}" if f.get("draw_under")
                 else f"stacked after plan step {f['draw_after_n']}"
                 if f.get("draw_after_n") is not None else "on top of everything")
        id_lines += (f"#   {f['id']}: name its shapes {f['id']}, {f['id']}-1, {f['id']}-2 ...; "
                     f"{where} (set for you)\n")
    pending = (f"# FALLBACK_PENDING: {len(plan.fallbacks)} element(s) need the "
               f"slide-builder-translator agent; see option_{letter}_translation_report.json")

    def _write(fb: str, block: str | None = None) -> None:
        text = NATIVE_TEMPLATE.format(
            name=f"{stem}.py", version=TRANSLATOR_VERSION, plan_name=plan_name,
            fields_block=_fields_block(fields), fallback_line=fb, skill=str(SKILL))
        if block is not None:
            text = text.replace("    # (none)\n", block, 1)
        native.write_text(text, encoding="utf-8")

    def _write_with_previous(section: str) -> None:
        _write(pending + PENDING_PREV_NOTE + "\n" + id_lines, "    # (none)\n" + section)

    kept = False
    re_info = None
    if old is None:
        _write(pending + "\n" + id_lines if plan.fallbacks else "")
    else:
        try:
            body_same = old["plan"] == json.loads(json.dumps(pub))
        except Exception:
            body_same = False
        fields_same = (old["report"] or {}).get("template_fields") == fields
        # a: converts the same; b: only the title / takeaway / footer text the
        # template draws changed; c: the body changed; d: c where the agent drew
        case = "a" if body_same and fields_same else "b" if body_same else "c"
        re_info = {"case": case, "drawing": None, "why": "",
                   "snapshot": str(snap) if snap else ""}
        if old["drawn"] and plan.fallbacks:
            if case == "c":
                re_info["case"] = "d"
            why = _drawing_decision(old, plan)
            if not why:
                # The drawing as the agent wrote it, under the new header and
                # plan; the template's restack_fallbacks call stays after it.
                _write(old["done_line"] + "\n" + id_lines, old["block"])
                why = _kept_drawing_check(native, plan.fallbacks)
                kept = not why
            re_info.update(drawing="kept" if kept else "pending", why=why)
            if not kept:
                _write_with_previous(_as_previous(old["block"]))
        elif old["drawn"]:
            re_info.update(drawing="not needed",
                           why="the script now draws everything the agent drew")
            _write("")
        elif plan.fallbacks and old["prev_section"]:
            # Still waiting on the agent from an earlier change: the pasted
            # previous drawing stays with it.
            re_info.update(drawing="pending", carried=True,
                           why="the agent had not adapted its previous drawing yet")
            _write_with_previous(old["prev_section"])
        else:
            _write(pending + "\n" + id_lines if plan.fallbacks else "")
    plan.counts.update(native_counts(plan))
    try:
        from twins.html_emit import build as _build
        _prs, _sl = _build(_public_plan(plan))
        for hit in caps_above_frame(_sl):
            plan.warn("TEXT_ABOVE_FRAME", f"{hit['shape']}: capitals reach {hit['above_px']} px "
                      "above the text box (its frame when selected); advisory")
    except Exception:
        pass
    report = {
        "translator": TRANSLATOR_VERSION,
        "html": str(html),
        "css_kill_list_applied": plan.kill,
        "warnings": plan.warnings,
        "counts": plan.counts,
        "fallback": _public_fallbacks(plan),
        "self_check": plan.self_check,
        "needs_agent": bool(plan.fallbacks) and not kept,
        "html_sha256": _sha256(html),
        "template_fields": fields,
        "visible_text_chars": len(re.sub(r"\s+", "", visible_text or "")),
    }
    if re_info:
        report["retranslation"] = re_info
    (emit_dir / f"option_{letter}_translation_report.json").write_text(
        json.dumps(report, indent=1), encoding="utf-8")
    return report


def _public_fallbacks(plan: "Plan") -> list[dict]:
    return [{k: v for k, v in f.items() if not k.startswith("_")} for f in plan.fallbacks]


def _sha256(path: Path) -> str:
    import hashlib
    try:
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()
    except OSError:
        return ""


def fallback_pending(native_py: Path) -> bool:
    """True while the translator agent still has elements to draw."""
    try:
        return "# FALLBACK_PENDING:" in Path(native_py).read_text(encoding="utf-8")
    except OSError:
        return False


def needs_conversion(html: Path, emit_dir: Path, letter: str) -> bool:
    """Does this design need converting (again)?

    Yes when there is no option_X_native.py yet, or the design changed after
    it was written (the report records the design's fingerprint; an older
    report without one falls back to file times). No otherwise: converting
    again rewrites option_X_native.py from scratch, which erased what the
    translator agent had drawn in fallback mode and marked it pending again,
    so a QC fix or redesign looped on "run the agent" for ever (2026-10-06).
    """
    native = Path(emit_dir) / f"option_{letter}_native.py"
    if not native.exists():
        return True
    try:
        rep = json.loads((Path(emit_dir) / f"option_{letter}_translation_report.json")
                         .read_text(encoding="utf-8"))
    except Exception:
        rep = {}
    sha = rep.get("html_sha256") if isinstance(rep, dict) else None
    if sha:
        return sha != _sha256(html)
    return Path(html).stat().st_mtime > native.stat().st_mtime


def refresh_line(rep: dict) -> str:
    """One plain line for a converted option: what happened and why."""
    h = Path(rep["html"])
    m = re.match(r"slide_(\d+)", h.parent.name)
    who = (f"slide {int(m.group(1))} option {h.stem.split('_')[-1]}" if m
           else f"{h.parent.name}/{h.name}")
    ri = rep.get("retranslation")
    n_fb = len(rep.get("fallback") or [])
    wait = (f"; {n_fb} element(s) wait on the translator agent (FALLBACK MODE)"
            if rep.get("needs_agent") else "")
    if not ri:
        return f"[ok] {who}: converted{wait}"
    d, why = ri.get("drawing"), ri.get("why") or ""
    if d == "kept":
        head = {"a": "the design file changed but converts the same",
                "b": "only the title or takeaway text changed; script header updated",
                "d": "the design changed; converted again"}.get(ri["case"], "converted again")
        return (f"[ok] {who}: {head}; the translator agent's drawing is kept (every element "
                "it draws is unchanged and it still runs)")
    if d == "pending" and ri.get("carried"):
        return (f"[!!] {who}: the design changed again before the translator agent adapted "
                "its previous drawing; converted again, and it still waits on the agent "
                "(FALLBACK MODE) with that drawing pasted in")
    if d == "pending":
        return (f"[!!] {who}: the design changed where the translator agent drew ({why}); "
                "converted again, and it waits on the agent (FALLBACK MODE) with its previous "
                "drawing pasted in to adapt")
    if d == "not needed":
        return (f"[ok] {who}: the design changed; converted again, and the agent's earlier "
                "drawing is no longer needed (the script draws it all now)" + wait)
    head = {"a": "the design file changed but converts the same",
            "b": "only the title or takeaway text changed; converted again",
            "c": "the design changed; converted again"}.get(ri["case"], "converted again")
    return f"[ok] {who}: {head}{wait}"


def refresh_stale(out: Path, items: list[tuple[int, str]], quiet: bool = False) -> list[Path]:
    """Bring the picked sketch options up to date with their designs.

    items: (slide number, letter). An option whose design (option_X.html) is
    the one it was converted from is kept as it is. One that changed, or was
    never converted, is converted again in one call (one browser, one
    self-check pass), and write_outputs decides what happens to what the
    translator agent drew:
      a  converts the same: script kept, fingerprint updated
      b  only the template's text (title, takeaway, footer) changed: the
         script's __template_fields__ header is rewritten, drawing kept
      c  the body changed and the agent drew nothing: replaced
      d  the body changed and the agent drew: its drawing is kept only when
         every element it draws is unchanged and the script still runs; else
         the option waits on the agent with the old drawing pasted in
    The old script is copied to slide_NN/_prev_translation/<time>/ first.
    Prints one line per option (quiet=True: only the ones converted).
    Returns the option_X_native.py files that wait on the translator agent:
    the ones just converted with elements left over, and the ones kept whose
    agent part is not drawn yet.
    """
    jobs, waiting = [], []
    for n, letter in items:
        sd = out / f"slide_{n:02d}"
        html = sd / f"option_{letter}.html"
        if not html.exists():
            continue
        native = sd / f"option_{letter}_native.py"
        if needs_conversion(html, sd, letter):
            jobs.append(job_for(out, n, letter))
            continue
        if fallback_pending(native):
            waiting.append(native)
            if not quiet:
                print(f"[..] slide {n} option {letter}: converted from the current design; "
                      "still waits on the translator agent (FALLBACK MODE)")
        elif not quiet:
            print(f"[ok] slide {n} option {letter}: already converted from the current "
                  "design; kept as is (nothing the translator agent drew is redone)")
    if jobs:
        for r in translate_many(jobs):
            print(refresh_line(r))
            if r["needs_agent"]:
                h = Path(r["html"])
                waiting.append(h.with_name(h.stem + "_native.py"))
    return waiting


def convert_picked(out: Path, items: list[tuple[int, str]]) -> list[Path]:
    """The picked sketch options brought up to date (refresh_stale)."""
    return refresh_stale(out, items)


def drawn_by_agent(native_py: Path) -> bool:
    """True when the translator agent has drawn in this script (FALLBACK_DONE
    and code between the markers)."""
    try:
        text = Path(native_py).read_text(encoding="utf-8")
    except OSError:
        return False
    m = _BLOCK_RE.search(text)
    return bool(re.search(r"^# FALLBACK_DONE:", text, re.M)) and bool(
        _drawing_lines(m.group(1) if m else ""))


def _subtitle_as_shape(out: Path, slide_n: int) -> bool:
    """True when this slide's layout has no subtitle placeholder (the takeaway is
    then drawn as a body shape named 'subtitle', which finalize recognizes)."""
    try:
        import _paths as _p
        from _chrome_schema import load_chrome_yml
        meta = json.loads(_p.meta_json(out).read_text(encoding="utf-8"))
        slide = next(s for s in meta["slides"] if s.get("n") == slide_n)
        spec = load_chrome_yml(_p.chrome_yml(Path(meta["template"])))
        lc = spec.layouts.get(slide.get("layout") or "")
        return lc is None or getattr(lc, "subtitle_placeholder_idx", None) is None
    except Exception:
        return True


_TEMPLATES: dict = {}   # template opened once per run, not once per slide


def _title_text_span(out: Path, slide_n: int) -> tuple[float, float] | None:
    """(left, width) in px of the title's TEXT on this slide's layout: the
    title box (chrome.yml's title_box_x_px / title_box_width_px, else the
    layout's title placeholder) less its inner left and right margins. The
    takeaway drawn as a shape has no inner margin, so this is where its text
    must sit to line up with the title's. None when it can't be read."""
    try:
        import _paths as _p
        from _chrome_schema import load_chrome_yml
        from pptx import Presentation
        meta = json.loads(_p.meta_json(out).read_text(encoding="utf-8"))
        slide = next(s for s in meta["slides"] if s.get("n") == slide_n)
        layout_name = slide.get("layout") or ""
        tpl = Path(meta["template"])
        spec = load_chrome_yml(_p.chrome_yml(tpl))
        lc = spec.layouts.get(layout_name)
        key = (str(tpl), tpl.stat().st_mtime)
        prs = _TEMPLATES.get(key)
        if prs is None:
            prs = _TEMPLATES[key] = Presentation(str(tpl))
        layout = next(l for l in prs.slide_layouts if l.name == layout_name)
        t_idx = getattr(lc, "title_placeholder_idx", None) if lc else None
        ph = next((p for p in layout.placeholders if t_idx is not None
                   and p.placeholder_format.idx == t_idx), None) or next(
            p for p in layout.placeholders if int(p.placeholder_format.type or 0) in (1, 3))
        x = getattr(lc, "title_box_x_px", None) if lc else None
        w = getattr(lc, "title_box_width_px", None) if lc else None
        if x is None or w is None:
            x, w = int(ph.left) / 9525, int(ph.width) / 9525
        from pptx.oxml.ns import qn

        def _ins(attr):
            for el in (ph._element, getattr(ph, "_base_placeholder", None)):
                el = getattr(el, "_element", el)
                bp = el.find(".//" + qn("a:bodyPr")) if el is not None else None
                if bp is not None and bp.get(attr) is not None:
                    return int(bp.get(attr)) / 9525
            return 91440 / 9525          # PowerPoint's default inner margin
        lin, rin = _ins("lIns"), _ins("rIns")
        return float(x) + lin, float(w) - lin - rin
    except Exception:
        return None


def _template_bg(out: Path, slide_n: int) -> str | None:
    """This slide's layout background ('RRGGBB'): chrome.yml's record, else
    read from the template. None when unknown or a picture."""
    try:
        import _paths as _p
        import _template_bg as TB
        from _chrome_schema import load_chrome_yml
        meta = json.loads(_p.meta_json(out).read_text(encoding="utf-8"))
        slide = next(s for s in meta["slides"] if s.get("n") == slide_n)
        tpl = Path(meta["template"])
        lc = None
        try:
            lc = load_chrome_yml(_p.chrome_yml(tpl)).layouts.get(slide.get("layout") or "")
        except Exception:
            pass
        f = TB.facts_for(tpl, slide.get("layout") or "", lc)
        # For a gradient or pattern this is its first color, the one prep gave
        # the designer for the canvas: a canvas of that color is still the
        # template's background and is left to it.
        return f["bg_hex"] if f.get("bg_hex") and f.get("bg_kind") != "picture" else None
    except Exception:
        return None


def job_for(out: Path, slide_n: int, letter: str) -> tuple:
    """The translate_many job for one picked option of a build."""
    sd = out / f"slide_{slide_n:02d}"
    sas = _subtitle_as_shape(out, slide_n)
    return (sd / f"option_{letter}.html", sd, letter, sas,
            _title_text_span(out, slide_n) if sas else None, _template_bg(out, slide_n))


def snap_subtitle(plan: "Plan", span: tuple[float, float] | None) -> None:
    """Line the takeaway drawn as a shape up with the title: its text starts
    where the title's text starts and spans the same width. The design's
    height position is kept. Workers are not told the title's left edge, so
    the sketch's x was a guess, and the line came out indented under the
    title (2026-10-06). Runs after the self-check, which compares against the
    design as drawn."""
    if not span:
        return
    op = next((o for o in plan.ops if o.get("op") == "text" and o.get("name") == "subtitle"),
              None)
    if op is None:
        return
    x, w = span
    if w <= 0 or (abs(op["x"] - x) < 0.5 and abs(op["w"] - w) < 0.5):
        return
    lines = op.get("_lines")
    predicted = _wrapped_lines(op["paragraphs"], w)
    if lines and predicted and predicted > lines:
        op["h"] = op["h"] * predicted / lines
    op["x"], op["w"] = x, w
    plan.counts["subtitle_aligned_to_title"] = 1


def font_ratio_message(font: str, ratio: float, installed: bool | None = None) -> str:
    """The RENDER_FONT_RATIO warning. ratio: finished text width / sketch text
    width, shared by every text in `font` (font may end in ' bold').

    Installed font: the sketch was drawn in a different version of it than the
    finished slide uses (the cause on 2026-10-08 was a designer's font-file
    rule; the old wording blamed the check's renderer and called it "not a
    defect", which sent the session after the wrong cause). Not installed: the
    check had to use a stand-in, so text size was not judged for that font."""
    fam = re.sub(r"\s+bold$", "", font or "", flags=re.I).strip()
    if installed is None:
        try:
            from _chrome_schema import installed_families
            installed = fam.lower() in installed_families()
        except Exception:
            installed = True
    pct = abs(ratio - 1)
    more = "wider" if ratio >= 1 else "narrower"
    if installed:
        return (f"the design was drawn in a different version of {font} than the finished "
                f"slide uses: the finished text runs {pct:.0%} {more} than the approved "
                f"sketch. Check every {font} text box for overflow; this is a design issue "
                f"to fix, not a renderer quirk")
    return (f"this computer lacks {font}, so the check used a stand-in {pct:.0%} {more}; "
            f"text size was not judged for this font")


def _render_plans(plans: list["Plan"], engine: str = "auto") -> list:
    """(image, characters) per plan, from one render pass (translate_alarm)."""
    import _scratch
    import translate_alarm as A
    from twins.html_emit import build
    with _scratch.temp_dir("alarm_") as td:
        paths = []
        for k, plan in enumerate(plans):
            prs, _s = build(_public_plan(plan))
            if plan.canvas_bg and plan.canvas_bg != "FFFFFF":
                # The canvas was left to the template's background; the
                # check's blank slide gets that color instead, so it
                # compares like with like.
                from pptx.dml.color import RGBColor
                _s.background.fill.solid()
                _s.background.fill.fore_color.rgb = RGBColor.from_string(plan.canvas_bg)
            paths.append(Path(td) / f"t{k:03d}.pptx")
            prs.save(str(paths[-1]))
        return A.render_full(paths, engine=engine)


def _bullet_rows(plan: "Plan") -> list[dict]:
    """Self-check rule for one-box lists: a bullet bigger than
    BULLET_MAX_PCT of its text, or starting outside the filled panel its text
    sits on, is wrong whatever the render looks like (the render is compared
    with the same design the bullet was misread from). The list steps down to
    markers drawn as shapes, like any other list that came out different."""
    import translate_alarm as A
    fills = [o for o in plan.ops if o.get("op") == "shape" and o.get("_filled") and o.get("_box")]
    rows = []
    for i, o in enumerate(plan.ops):
        if o.get("op") != "text" or not o.get("para") or not o.get("_ink"):
            continue
        ix, iy, iw, ih = o["_ink"]
        holders = [b for b in fills if b.get("_order", 0) < o.get("_order", 0)
                   and b["_box"][0] - 1 <= ix and ix + iw <= b["_box"][0] + b["_box"][2] + 1
                   and b["_box"][1] - 1 <= iy and iy + ih <= b["_box"][1] + b["_box"][3] + 1]
        holder = min(holders, key=lambda b: b["_box"][2] * b["_box"][3]) if holders else None
        bad = None
        for pr in o["para"]:
            b = pr.get("bullet")
            if not b:
                continue
            left = o["x"] + (pr.get("marL_px") or 0) + (pr.get("indent_px") or 0)
            if not b.get("autonum") and float(b.get("pct") or 100) > BULLET_MAX_PCT:
                bad = f"a bullet is {float(b['pct']):.0f}% of its text"
            elif holder and not (holder["_box"][0] - 1 <= left <= holder["_box"][0] + holder["_box"][2]):
                bad = "a bullet starts outside the panel its text sits on"
            if bad:
                break
        if bad:
            r = A.region(o)
            if r:
                rows.append({"i": i, "name": o.get("name") or f"text#{i}", "op": "text", "box": r,
                             "fired": ["bullet"], "detail": bad})
    return rows


# Pictures, drawings, rotated and vertical elements the agent draws: their
# areas (x0, y0, x1, y1), a little wider than the element so its softened
# edge is left out too. Elements handed over by the self-check itself are not
# here: they are out of the plan, nothing is judged against them.
AGENT_MASK_PAD = 2.0


def _agent_masks(plan: "Plan") -> list[tuple]:
    out = []
    for f in plan.fallbacks:
        if f.get("kind") == "self-check":
            continue
        rects = f.get("_rects") or [{"x": f.get("x"), "y": f.get("y"), "w": f.get("w"),
                                     "h": f.get("h")}]
        for q in rects:
            try:
                x, y, w, h = (float(q[k]) for k in ("x", "y", "w", "h"))
            except (TypeError, ValueError, KeyError):
                continue
            p = AGENT_MASK_PAD
            out.append((x - p, y - p, x + w + p, y + h + p))
    return out


SELF_CHECK_ROUNDS = 4    # a group steps down at most twice, plus a last look


def self_check(items: list[tuple["Plan", dict]], engine: str = "auto",
               hand_over: bool = True) -> None:
    """Render every translated slide and hand any element that came out
    looking different from the design to the agent (scripts/translate_alarm.py).

    A native table or a one-box list that comes out different is not handed
    over: it steps down to its next version (see _plan_native_groups) and
    that slide is rendered again, judging only what was swapped in. Text is
    never dropped on the way.

    items: (plan, extracted data) pairs; plans are changed in place.
    engine: "auto" (the default renderer: PowerPoint on Windows when it is
    installed, else LibreOffice; slide-builder/settings.json "renderer"),
    "libreoffice" or "powerpoint". Each round renders all its slides as ONE
    deck (translate_alarm.render_full); a later round renders only the slides
    that changed. hand_over=False (the owner's converter check) keeps fired elements
    in the plan and only records them.
    A slide that does not render gets a SELF_CHECK_SKIPPED warning (with the
    reason) and keeps all its elements; the other slides are still checked.
    """
    import translate_alarm as A
    judge: dict[int, set | None] = {k: None for k in range(len(items))}
    for plan, _ in items:
        plan.self_check = {"ran": False, "fired": []}
        plan.counts["self_check_handed_over"] = 0
    todo = list(range(len(items)))
    for rnd in range(SELF_CHECK_ROUNDS):
        if not todo:
            break
        try:
            rendered = _render_plans([items[k][0] for k in todo], engine)
        except Exception as exc:
            for k in todo:
                items[k][0].warn("SELF_CHECK_SKIPPED",
                                 f"could not render to compare ({type(exc).__name__})")
            return
        nxt = []
        why_not = dict(A.LAST_RENDER_ERRORS)
        for j, (k, got) in enumerate(zip(todo, rendered)):
            plan, data = items[k]
            if not got:
                reason = (why_not.get(j) or "").strip().splitlines()
                plan.warn("SELF_CHECK_SKIPPED", "the slide did not render"
                          + (f" ({reason[-1][:200]})" if reason else ""))
                continue
            native, chars = got
            plan._last_render = got
            rows = [r for r in A.compare(plan.ops, A.load_rgb(data["_shot"]), native, chars,
                                         engine=A.LAST_RENDER_ENGINES.get(j),
                                         masks=_agent_masks(plan))
                    if r["fired"]]
            rows += _bullet_rows(plan)
            if rnd == 0:
                for font, ratio in (getattr(A.compare, "font_ratios", None) or {}).items():
                    plan.warn("RENDER_FONT_RATIO", font_ratio_message(font, ratio))
            plan.self_check["ran"] = True
            if judge[k] is not None:
                rows = [r for r in rows if id(plan.ops[r["i"]]) in judge[k]]
            by_op: dict[int, list] = {}
            for r in rows:
                by_op.setdefault(r["i"], []).append(r)
            stepped, swapped_in, handed = set(), [], []
            last_round = rnd == SELF_CHECK_ROUNDS - 1
            for i, rs in sorted(by_op.items()):
                o = plan.ops[i]
                gid = o.get("_gid")
                grp = plan.groups.get(gid) if gid is not None else None
                if grp and grp["cur"] < len(grp["chain"]) - 1 and not last_round:
                    if gid not in stepped:
                        stepped.add(gid)
                        why = ", ".join(sorted({w for r in rs for w in r["fired"]}))
                        swapped_in.append((gid, why))
                    continue
                handed.append((o, rs))
            for o, rs in handed:
                plan.self_check["fired"] += [{"element": r["name"], "what": r["fired"]} for r in rs]
            if hand_over:
                for o, rs in handed:
                    x0 = min(r["box"][0] for r in rs)
                    y0 = min(r["box"][1] for r in rs)
                    x1 = max(r["box"][2] for r in rs)
                    y1 = max(r["box"][3] for r in rs)
                    plan.fallbacks.append({
                        "kind": "self-check", "id": o.get("name") or "",
                        "x": x0, "y": y0, "w": x1 - x0, "h": y1 - y0,
                        "reason": "came out looking different from the design ("
                                  + ", ".join(sorted({w for r in rs for w in r["fired"]})) + ")",
                        "_order": o.get("_order"),
                        # what the element is, for _fallback_fp: the plan step
                        # the design gave it (not where the render put it)
                        "_html": json.dumps({k: v for k, v in o.items()
                                             if not k.startswith("_")}, sort_keys=True,
                                            default=str)})
                gone = {id(o) for o, _ in handed}
                plan.ops = [o for o in plan.ops if id(o) not in gone]
                plan.counts["self_check_handed_over"] += len(handed)
            if swapped_in:
                new_ids = set()
                for gid, why in swapped_in:
                    new_ids |= {id(o) for o in _step_group(plan, gid, why)}
                judge[k] = new_ids
                nxt.append(k)
        todo = nxt


def translate_many(jobs: list[tuple], check: bool = True) -> list[dict]:
    """jobs: (html, emit_dir, letter, subtitle_as_shape[, title_text_span[, template_bg]]);
    job_for() builds one for a build's slide. One browser for all, then one
    render of every slide, as one deck, for the self-check."""
    from playwright.sync_api import sync_playwright
    done = []
    from _browser import scratch_env, render_deadline, render_slot, render_timeout, step
    # Time limit (scripts/_browser.py): 120 s (SLIDE_LAB_RENDER_TIMEOUT) plus
    # 15 s for each slide after the first, 30 s to read any one slide. Past it,
    # only this run's own browser is ended and the run exits 4: re-run.
    # The browser part holds one of the computer's render slots, taken BEFORE
    # the time limit starts, so waiting for it is free (slide-qc/scripts/
    # render_slot.py). The self-check render below takes its own slot.
    limit = render_timeout() + 15 * max(0, len(jobs) - 1)
    with render_slot(f"translate_html.py ({len(jobs)} design(s))"), \
            render_deadline("translate_html.py", seconds=limit), \
            scratch_env(), sync_playwright() as pw:
        from _browser import launch
        browser = launch(pw)
        page = browser.new_page(viewport={"width": CANVAS_W, "height": CANVAS_H})
        for html, emit_dir, letter, sas, *rest in jobs:
            with step(f"reading the design {Path(html).parent.name}/{Path(html).name}"):
                data = extract(page, html, sas)
            plan, fields = plan_from_extract(data, sas, rest[1] if len(rest) > 1 else None)
            done.append((html, emit_dir, letter, plan, fields, data, rest[0] if rest else None))
        browser.close()
    if check and done:
        self_check([(d[3], d[5]) for d in done])
    for d in done:
        snap_subtitle(d[3], d[6])
        # after the check, as snap_subtitle: grown boxes are not "drawn wrong"
        _apply_clearance(d[3])
    return [write_outputs(html, emit_dir, letter, plan, fields, data.get("visible_text", ""))
            for html, emit_dir, letter, plan, fields, data, _span in done]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Translate picked HTML designs to native PowerPoint.")
    ap.add_argument("--out", type=Path, help="Build output dir.")
    ap.add_argument("--slide", type=int, help="Slide number.")
    ap.add_argument("--letter", help="Option letter.")
    ap.add_argument("--html", type=Path, help="Standalone: an HTML file to translate.")
    ap.add_argument("--emit", type=Path, help="Standalone: where to write the outputs.")
    ap.add_argument("--no-self-check", action="store_true",
                    help="Skip rendering and comparing with the design (debugging only).")
    ap.add_argument("--force", action="store_true",
                    help="Convert again even though the translator agent has drawn in this "
                         "option's script (its drawing is kept only if every element it draws "
                         "is unchanged; else it is pasted back as the previous drawing).")
    args = ap.parse_args(argv)
    if args.html:
        emit = args.emit or args.html.parent
        emit.mkdir(parents=True, exist_ok=True)   # a manual run used to crash here
        letter = re.match(r"option_([A-Z])", args.html.name)
        jobs = [(args.html, emit, letter.group(1) if letter else "A", True)]
    elif args.out and args.slide and args.letter:
        sd = args.out / f"slide_{args.slide:02d}"
        html = sd / f"option_{args.letter}.html"
        if not html.exists():
            print(f"ERROR: no design at {html}")
            return 2
        jobs = [job_for(args.out, args.slide, args.letter)]
    else:
        ap.error("give --out/--slide/--letter, or --html")
    # A hand run over a script the agent has drawn in (2026-10-13): it used to
    # rewrite the script from scratch and the drawing was lost for good.
    html0, emit0, letter0 = jobs[0][0], jobs[0][1], jobs[0][2]
    native0 = Path(emit0) / f"option_{letter0}_native.py"
    if not args.force and drawn_by_agent(native0):
        same = not needs_conversion(Path(html0), Path(emit0), letter0)
        print(f"REFUSED: {native0} has the translator agent's drawing in it (FALLBACK_DONE). "
              + ("It is already converted from the current design: nothing to do."
                 if same else
                 "The design changed since. Bring it up to date with "
                 + (f"py -3 scripts/record_picks.py --out \"{args.out}\" --edits-done "
                    "(it converts what changed and keeps the drawing when the elements it "
                    "draws are unchanged), or finalize_deck.py, which does the same."
                    if args.out else "--force.")))
        print("  Add --force to convert anyway: the drawing is kept only if every element it "
              "draws is unchanged; else the script waits on the agent with the old drawing "
              "pasted in. The old script is copied to _prev_translation\\ first.")
        return 5
    for rep in translate_many(jobs, check=not args.no_self_check):
        c = rep["counts"]
        print(f"[ok] {Path(rep['html']).parent.name}/{Path(rep['html']).name}: "
              f"{c['shapes']} shapes, {c['texts']} text boxes, {c['svg_items']} drawing items"
              + (f"; {len(rep['fallback'])} element(s) need the agent" if rep["needs_agent"]
                 else ""))
        if rep.get("retranslation"):
            print("  " + refresh_line(rep))
    return 0


if __name__ == "__main__":
    sys.exit(main())
