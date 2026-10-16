#!/usr/bin/env python3
"""fit_boxes.py — fit each plain text box to its text, with proof that no text moved.

Why (owner's request, 2026-10-16): "many times the text boxes are too large
that it is difficult to click on." A label in a full-width card came out as a
full-width box, so clicking anywhere on the card picked the label.

How:
  1. PowerPoint measures the real text of every box that may be fitted
     (read-only, no window, on a local copy; slide-qc/scripts/ppt_safe.py,
     so the user's PowerPoint and decks are never touched). It also says what
     the box inherits from its layout and master: anchor, word wrap, margins,
     shrink-on-overflow, fill and outline.
  2. First its left, right and bottom margins (insets) become 0, the margin
     turned into box geometry so the text area, the text and its line
     breaks stay exactly where they were (owner, 2026-10-16: "Can all just
     be 0?"; a right margin of 0.54" used to force wrapping becomes a
     narrower box). The top margin stays (the item-15 top padding).
     Then each box shrinks around its text, with the anchored edge (top, middle,
     bottom) and the aligned side (left, center, right) kept where they are:
       - a box with word wrap on keeps its height to the text plus its
         bottom margin plus 2 pt; its width narrows only while the spare room
         stays at least 1 character or 10% of the width (whichever is larger),
         so a word typed later does not wrap at once
       - a one-line box (word wrap off) gets the text's width plus 1 character
     Only the custom text boxes the designers and the converter added are
     fitted. Anything from the template is never touched, not even its
     height (owner's decision, 2026-10-16: "Everything pre-template does not
     need to adjust. It's just these custom text boxes that cause issues").
  3. Proof, inside this script: PowerPoint measures the fitted copy again.
     Any box whose text moved more than 0.6 pt or whose line count changed
     goes back to its old size. Then the pages with fitted boxes are drawn
     before and after and compared pixel by pixel: a page with more than 30
     changed pixels gets every box on it back. This repeats up to 3 rounds.
     The fitted deck is kept only when the last round shows 0 moved boxes and
     0 changed pages; otherwise the deck stays exactly as it was.

Never fitted (the box keeps its size):
  - any placeholder or slot from the template, layout or master: title,
    subtitle, body, page number, date, footer, any layout slot; and Slide
    Lab's fixed slots (takeaway, footnote, source, section tag)
  - a box holding a field: page number, date, footer (owner: "those stay
    constant"; a page number fitted to "9" stacks as "1" over "0" once a page
    is inserted before it)
  - shrink-text-on-overflow boxes, justified text
  - a box with a visible fill or outline, its own or from the layout, master
    or theme style (the box is part of the picture)
  - rotated, flipped, vertical or multi-column text; shapes that are not
    plain rectangles; hidden shapes
  - shapes a connector is attached to; connectors and pictures
  - anything in a table, chart, SmartArt or group

Without PowerPoint (a Mac, or Windows without it) the step is skipped with a
note. It never guesses from font files: that is how text rewraps.

Library:
  fit_deck(path, out=None, pages=None, log=print) -> record dict
      fits `path` (in place, or into `out`); pages = 1-based page numbers to
      consider (None = every page). The record says what happened:
      status fitted | unchanged | skipped | error, plus counts.
  summary_lines(record) -> plain-English lines for COMPILED.md and check_done

Command line (decks Slide Lab did not build: fit only the pages Slide Lab
changed):
  py -3 fit_boxes.py "<deck.pptx>" --pages 3,7-9 [--out "<new.pptx>"]
  py -3 fit_boxes.py "<deck.pptx>" --all-pages   [--out "<new.pptx>"]
The original is never overwritten; the default output is <name>_fitted.pptx.

SLIDE_LAB_NO_BOX_FIT=1 turns the step off (recorded as skipped, with that
reason, and shown at delivery).
"""
from __future__ import annotations

import argparse
import contextlib
import os
import shutil
import sys
import time
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
for _sp in (str(SKILL), str(HERE), str(SKILL.parent / "slide-qc" / "scripts")):
    if _sp not in sys.path:
        sys.path.insert(0, _sp)

PT = 12700                    # EMU per point
BOUNDS_TOL_PT = 0.6           # text may move at most this much
PIXEL_DELTA = 40              # a pixel counts as changed above this (0-255)
PIXEL_PAGE_LIMIT = 30         # a page counts as changed above this many pixels
MAX_ROUNDS = 3
HEIGHT_PAD_PT = 2.0           # room kept below (or above) the text
MIN_GAIN_PT = 1.0             # smaller savings are not worth a change
RENDER_WIDTH_PX = 1920

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_FILLS = ("solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
_FIELD_PH_TYPES = (13, 15, 16)          # slide number, footer, date
_CHROME_NAMES = ("section-tag",)
_CHROME_PREFIXES = ("chrome-",)

# PowerPoint constants (Office type library)
_MSO_TRUE = -1
_ANCHOR = {1: "t", 3: "m", 4: "b"}      # top, middle, bottom (baselines: skipped)
_ALIGN = {1: "l", 2: "ctr", 3: "r"}     # left, center, right; others: see below
_ALIGN_JUSTIFIED = (4, 5, 6, 7)         # justify, distribute, Thai, justify low
_AUTOSIZE_SHRINK_TEXT = 2               # msoAutoSizeTextToFitShape
_ORIENT_HORIZONTAL = 1
_RECTANGLE = 1                          # msoShapeRectangle
_NOT_PRIMITIVE = 138                    # msoShapeNotPrimitive
_PLACEHOLDER = 14                       # msoPlaceholder

NO_PPT_NOTE = ("PowerPoint is not on this computer (a Mac, or Windows without "
               "PowerPoint), so text boxes were not fitted to their text. "
               "Nothing else changes; fit them later on Windows with: "
               "py -3 slide-builder/scripts/fit_boxes.py \"<deck.pptx>\" --all-pages")


# ---------------------------------------------------------------------------
# Which boxes may be fitted (file side; PowerPoint adds its own checks)
# ---------------------------------------------------------------------------
def _ph_chain(shape):
    """The layout and master placeholders this placeholder inherits from."""
    out = []
    cur = shape
    for _ in range(3):
        try:
            cur = cur._base_placeholder
        except Exception:
            break
        if cur is None:
            break
        out.append(cur)
    return out


def _bodypr(el):
    return el.find(P + "txBody/" + A + "bodyPr")


def _fill_state(el, which: str):
    """'on' / 'off' / None (not said here) for the fill or the outline."""
    sp = el.find(P + "spPr")
    if sp is None:
        return None
    if which == "fill":
        if sp.find(A + "noFill") is not None:
            return "off"
        if any(sp.find(A + t) is not None for t in _FILLS):
            return "on"
        return None
    ln = sp.find(A + "ln")
    if ln is None:
        return None
    if ln.find(A + "noFill") is not None:
        return "off"
    if any(ln.find(A + t) is not None for t in ("solidFill", "gradFill", "pattFill")):
        return "on"
    return None


def _style_ref_on(el, tag: str) -> bool:
    st = el.find(P + "style")
    if st is None:
        return False
    ref = st.find(A + tag)
    return ref is not None and ref.get("idx") not in (None, "0")


def visible_fill_or_outline(shape) -> bool:
    """True when the box has a fill or outline you can see, its own or one it
    inherits (layout, master, theme style). The fit's file-side check;
    PowerPoint's own answer is checked too when it measures."""
    chain = [shape] + (_ph_chain(shape) if getattr(shape, "is_placeholder", False) else [])
    for which, ref in (("fill", "fillRef"), ("line", "lnRef")):
        state = None
        for s in chain:
            state = _fill_state(s._element, which)
            if state is not None:
                break
        if state == "on":
            return True
        if state is None and _style_ref_on(shape._element, ref):
            return True
    return False


def _shrink_on_overflow(shape) -> bool:
    chain = [shape] + (_ph_chain(shape) if getattr(shape, "is_placeholder", False) else [])
    for s in chain:
        bp = _bodypr(s._element)
        if bp is None:
            continue
        if bp.find(A + "normAutofit") is not None:
            return True
        if bp.find(A + "noAutofit") is not None or bp.find(A + "spAutoFit") is not None:
            return False
    return False


def _is_field_box(shape) -> bool:
    el = shape._element
    if el.find(".//" + A + "fld") is not None:
        return True
    try:
        if shape.is_placeholder and int(shape.placeholder_format.type) in _FIELD_PH_TYPES:
            return True
    except Exception:
        pass
    try:
        from twins.composer import is_page_number_shape
        return bool(is_page_number_shape(shape))
    except Exception:
        return False


def _is_fixed_slot(shape, slide_h: int) -> bool:
    """Slide Lab's chrome slots: takeaway, footnote, source, section tag."""
    name = (getattr(shape, "name", "") or "").strip().lower()
    if name in _CHROME_NAMES or name.startswith(_CHROME_PREFIXES):
        return True
    try:
        import _chrome_schema as CS
        return bool(CS.is_page_takeaway(shape, slide_h) or CS.note_role(shape, slide_h))
    except Exception:
        return False


def _glued_ids(slide) -> set:
    ids = set()
    for tag in ("stCxn", "endCxn"):
        for c in slide._element.iter(A + tag):
            if c.get("id"):
                ids.add(int(c.get("id")))
    return ids


def plan(prs, pages=None):
    """(candidates, skipped) for the pages asked for.
    candidates: {(page, shape_id): {"name": str}}, only the custom text boxes
    the designers and the converter added (never a placeholder)
    skipped: Counter of reason -> number of text boxes kept as they are."""
    cands, skipped = {}, Counter()
    slide_h = int(prs.slide_height or 0)
    for i, slide in enumerate(prs.slides, 1):
        if pages is not None and i not in pages:
            continue
        glued = _glued_ids(slide)
        for sh in slide.shapes:
            el = sh._element
            tag = el.tag
            if tag == P + "grpSp":
                if any((t.text or "").strip() for t in el.iter(A + "t")):
                    skipped["inside a group"] += 1
                continue
            if tag == P + "graphicFrame":
                skipped["table, chart or SmartArt"] += 1
                continue
            if tag != P + "sp":
                continue                          # pictures, connectors
            if not getattr(sh, "has_text_frame", False) or not (sh.text_frame.text or "").strip():
                continue
            if _is_field_box(sh):
                skipped["holds a field (page number, date, footer)"] += 1
                continue
            if getattr(sh, "is_placeholder", False):
                skipped["template slot (placeholder)"] += 1
                continue
            if _is_fixed_slot(sh, slide_h):
                skipped["Slide Lab fixed slot (takeaway, footnote, source, section tag)"] += 1
                continue
            if sh.shape_id in glued:
                skipped["a connector is attached"] += 1
                continue
            if _shrink_on_overflow(sh):
                skipped["shrinks text on overflow"] += 1
                continue
            if visible_fill_or_outline(sh):
                skipped["visible fill or outline"] += 1
                continue
            sp = el.find(P + "spPr")
            geom = sp.find(A + "prstGeom") if sp is not None else None
            if (sp is not None and sp.find(A + "custGeom") is not None) or (
                    geom is not None and geom.get("prst") not in (None, "rect")):
                skipped["not a plain rectangle"] += 1
                continue
            xf = sp.find(A + "xfrm") if sp is not None else None
            if xf is not None and (xf.get("rot") not in (None, "0") or xf.get("flipH") in ("1", "true")
                                   or xf.get("flipV") in ("1", "true")):
                skipped["rotated or flipped"] += 1
                continue
            bp = _bodypr(el)
            if bp is not None and ((bp.get("vert") not in (None, "horz"))
                                   or int(bp.get("numCol") or 1) > 1):
                skipped["vertical or multi-column text"] += 1
                continue
            cands[(i, int(sh.shape_id))] = {"name": sh.name}
    return cands, skipped


# ---------------------------------------------------------------------------
# PowerPoint: measure and draw
# ---------------------------------------------------------------------------
def _call_count(tr, attr):
    """Count of a text range's Lines / Paragraphs / Runs.

    On the classic TextFrame.TextRange these are METHODS: Lines().Count. The
    first version of this check read the newer TextRange2 the wrong way and
    got no line count for any box (so rewraps went unseen); through pywin32,
    TextRange2.Lines() returns the TEXT, not the lines. So: the classic
    method first, the TextRange2 property second, and None (the box is not
    fitted) when neither gives a number."""
    try:
        return int(getattr(tr, attr)().Count)
    except Exception:
        pass
    try:
        v = getattr(tr, attr)
        if not isinstance(v, str):
            return int(v.Count)
    except Exception:
        pass
    return None


def _item(tr, attr, i):
    """The i-th paragraph or run of a classic TextRange (a method call)."""
    return getattr(tr, attr)(i, 1)


_BUSY = (-2147418111, -2147417846)      # call rejected / retry later: PowerPoint is busy
_BUSY_TRIES = 40


def _is_busy(exc) -> bool:
    try:
        return int(exc.args[0]) in _BUSY
    except Exception:
        return False


def _retry(fn):
    """fn(), tried again while PowerPoint answers 'busy' (it does while
    another program, or the user, is working in it). Up to about 20 s."""
    for i in range(_BUSY_TRIES):
        try:
            return fn()
        except Exception as exc:  # noqa: BLE001
            if not _is_busy(exc) or i == _BUSY_TRIES - 1:
                raise
            time.sleep(0.5)


@contextlib.contextmanager
def _open(app, pptx: Path):
    """ppt_safe.opened (read-only, no window, closes only our file), with the
    open tried again while PowerPoint is busy."""
    import ppt_safe
    for i in range(_BUSY_TRIES):
        cm = ppt_safe.opened(app, Path(pptx))
        try:
            pres = cm.__enter__()
            break
        except Exception as exc:  # noqa: BLE001
            if not _is_busy(exc) or i == _BUSY_TRIES - 1:
                raise
            time.sleep(0.5)
    try:
        yield pres
    finally:
        cm.__exit__(None, None, None)


class PowerPointEngine:
    """Measures and draws through ONE ppt_safe session (opened read-only,
    without a window; only our own copies are ever opened)."""

    def __init__(self, app):
        self.app = app

    def _shapes_by_id(self, pres, page):
        sl = pres.Slides(page)
        out = {}
        for j in range(1, sl.Shapes.Count + 1):
            shp = sl.Shapes(j)
            try:
                out[int(shp.Id)] = shp
            except Exception:
                pass
        return out

    def measure(self, pptx: Path, keys) -> dict:
        """{key: facts} for every (page, shape_id) key. A box PowerPoint
        cannot read fully gets {"skip": reason}."""
        res = {}
        by_page = {}
        for k in keys:
            by_page.setdefault(k[0], []).append(k)
        with _open(self.app, Path(pptx)) as pres:
            for page, ks in sorted(by_page.items()):
                shapes = _retry(lambda: self._shapes_by_id(pres, page))
                for k in ks:
                    shp = shapes.get(k[1])
                    if shp is None:
                        res[k] = {"skip": "PowerPoint could not find the box"}
                        continue
                    try:
                        res[k] = _retry(lambda: self._facts(shp))
                    except Exception as exc:  # noqa: BLE001
                        res[k] = {"skip": f"PowerPoint could not read it ({type(exc).__name__})"}
        return res

    @staticmethod
    def _facts(shp) -> dict:
        f = {"L": float(shp.Left), "T": float(shp.Top),
             "W": float(shp.Width), "H": float(shp.Height)}
        if int(shp.Visible) != _MSO_TRUE:
            return {"skip": "hidden"}
        if float(shp.Rotation or 0) != 0 or int(shp.HorizontalFlip) == _MSO_TRUE \
                or int(shp.VerticalFlip) == _MSO_TRUE:
            return {"skip": "rotated or flipped"}
        try:
            if int(shp.Fill.Visible) == _MSO_TRUE or int(shp.Line.Visible) == _MSO_TRUE:
                return {"skip": "visible fill or outline"}
        except Exception:
            return {"skip": "visible fill or outline"}
        try:
            ast = int(shp.AutoShapeType)
            # a placeholder reports "not primitive" (138); its geometry was
            # checked in the file
            if ast != _RECTANGLE and not (int(shp.Type) == _PLACEHOLDER and ast == _NOT_PRIMITIVE):
                return {"skip": "not a plain rectangle"}
        except Exception:
            pass
        tf = shp.TextFrame2
        if int(tf.AutoSize) == _AUTOSIZE_SHRINK_TEXT:
            return {"skip": "shrinks text on overflow"}
        if int(tf.Orientation) != _ORIENT_HORIZONTAL:
            return {"skip": "vertical or multi-column text"}
        try:
            if int(tf.Column.Number) > 1:
                return {"skip": "vertical or multi-column text"}
        except Exception:
            pass
        anchor = _ANCHOR.get(int(tf.VerticalAnchor))
        if anchor is None:
            return {"skip": "text anchored on its baseline"}
        tr = tf.TextRange                         # TextRange2: bounds
        tr1 = shp.TextFrame.TextRange             # classic: lines, paragraphs, runs
        lines = _call_count(tr1, "Lines") or _call_count(tr, "Lines")
        if not lines:
            return {"skip": "PowerPoint could not count its lines"}
        aligns = set()
        n_par = _call_count(tr1, "Paragraphs") or 0
        if not n_par:
            return {"skip": "PowerPoint could not read its paragraphs"}
        for i in range(1, n_par + 1):
            par = _item(tr1, "Paragraphs", i)
            if not (par.Text or "").strip():
                continue
            a = int(par.ParagraphFormat.Alignment)
            if a in _ALIGN_JUSTIFIED:
                return {"skip": "justified text"}
            aligns.add(_ALIGN.get(a, "mixed"))
        align = aligns.pop() if len(aligns) == 1 else "mixed"
        try:
            if int(tf.HorizontalAnchor) == 2:        # text block centered in the box
                align = "ctr" if align in ("l", "ctr", "r") else "mixed"
        except Exception:
            pass
        em = 0.0
        n_runs = _call_count(tr1, "Runs") or 0
        for i in range(1, n_runs + 1):
            try:
                em = max(em, float(_item(tr1, "Runs", i).Font.Size))
            except Exception:
                pass
        f.update({
            "bl": float(tr.BoundLeft), "bt": float(tr.BoundTop),
            "bw": float(tr.BoundWidth), "bh": float(tr.BoundHeight),
            "lines": lines, "anchor": anchor, "align": align,
            "wrap": int(tf.WordWrap) == _MSO_TRUE,
            "mL": float(tf.MarginLeft), "mR": float(tf.MarginRight),
            "mT": float(tf.MarginTop), "mB": float(tf.MarginBottom),
            "em": em if 1 <= em <= 400 else 12.0})
        if f["bw"] <= 0 or f["bh"] <= 0:
            return {"skip": "PowerPoint found no text to measure"}
        return f

    def render(self, pptx: Path, pages, out_dir: Path) -> dict:
        """{page: png} for the pages asked for."""
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        res = {}
        with _open(self.app, Path(pptx)) as pres:
            try:
                ratio = float(pres.PageSetup.SlideHeight) / float(pres.PageSetup.SlideWidth)
            except Exception:
                ratio = 9 / 16
            h = int(round(RENDER_WIDTH_PX * ratio))
            for p in sorted(pages):
                png = out_dir / f"page_{p:03d}.png"
                _retry(lambda: pres.Slides(p).Export(str(png.resolve()), "PNG",
                                                     RENDER_WIDTH_PX, h))
                res[p] = png
        return res


@contextlib.contextmanager
def _powerpoint_engine():
    import ppt_safe
    with ppt_safe.session() as (app, _started):
        yield PowerPointEngine(app)


def powerpoint_here() -> bool:
    try:
        from render_slides import powerpoint_available
        return bool(powerpoint_available())
    except Exception:
        return False


# ---------------------------------------------------------------------------
# The fit itself (points; pure)
# ---------------------------------------------------------------------------
def compute_fit(m: dict):
    """(L, T, W, H, zero_insets) in points, or None when the box already fits.

    First the left, right and bottom insets become part of the box (owner,
    2026-10-16: "Why does Right margin have 0.54" when everything else is 0?
    Can all just be 0?"): the left edge moves right by the left inset, the
    width loses both side insets, the height loses the bottom inset. The
    text area is exactly what it was, so the text stays where it is and
    wraps at the same words (a right inset used to force wrapping becomes a
    narrower box). The top inset stays: it is the item-15 top padding that
    keeps capitals inside the selection frame. Then the box is fitted."""
    L0, T0, W0, H0 = m["L"], m["T"], m["W"], m["H"]
    bl, bt, bw, bh = m["bl"], m["bt"], m["bw"], m["bh"]
    mL, mR, mT, mB = m["mL"], m["mR"], m["mT"], m["mB"]
    zero = any(v > 0.01 for v in (mL, mR, mB))
    if zero and (W0 - mL - mR > 1 and H0 - mB > 1):
        L, T, W, H = L0 + mL, T0, W0 - mL - mR, H0 - mB
        mL = mR = mB = 0.0
    else:
        zero = False
        L, T, W, H = L0, T0, W0, H0
    pad = HEIGHT_PAD_PT
    nl, nt, nw, nh = L, T, W, H
    if m["anchor"] == "t":
        nh = (bt + bh + mB + pad) - T
    elif m["anchor"] == "b":
        nt = bt - mT - pad
        nh = (T + H) - nt
    else:
        c = T + H / 2
        half = max(c - (bt - mT - pad), (bt + bh + mB + pad) - c)
        nh, nt = 2 * half, c - half
    if nh > H - MIN_GAIN_PT:
        nt, nh = T, H
    if m["align"] in ("l", "ctr", "r"):
        em = m["em"]
        if m["align"] == "l":
            base = (bl + bw + mR) - L
        elif m["align"] == "r":
            base = (L + W) - (bl - mL)
        else:
            c = L + W / 2
            base = 2 * max(c - (bl - mL), (bl + bw + mR) - c)
        # wrap on: spare >= 1 character and >= 10% of the new width
        spare = max(em, base / 9.0) if m["wrap"] else em
        w = base + spare
        if w <= W - MIN_GAIN_PT:
            nw = w
            if m["align"] == "l":
                nl = L
            elif m["align"] == "r":
                nl = (L + W) - w
            else:
                nl = (L + W / 2) - w / 2
    if not zero and (nl, nt, nw, nh) == (L, T, W, H):
        return None
    return nl, nt, nw, nh, zero


def _apply(src: Path, dst: Path, fits: dict, meas: dict) -> None:
    """Write the deck with the fitted geometry for the keys in `fits`."""
    from pptx import Presentation
    prs = Presentation(str(src))
    slides = list(prs.slides)
    for (page, sid), (nl, nt, nw, nh, zero) in fits.items():
        slide = slides[page - 1]
        sh = next((s for s in slide.shapes if int(s.shape_id) == sid), None)
        if sh is None:
            continue
        m = meas[(page, sid)]
        x, y, cx, cy = int(sh.left or 0), int(sh.top or 0), int(sh.width or 0), int(sh.height or 0)
        # Changes are applied to the file's own numbers, so an edge that does
        # not move keeps its exact value.
        nx = x + int(round((nl - m["L"]) * PT))
        ny = y + int(round((nt - m["T"]) * PT))
        ncx = cx + int(round((nw - m["W"]) * PT))
        ncy = cy + int(round((nh - m["H"]) * PT))
        sh.left, sh.top, sh.width, sh.height = nx, ny, ncx, ncy
        if zero:
            bp = _bodypr(sh._element)
            if bp is not None:
                for att in ("lIns", "rIns", "bIns"):      # tIns (top padding) stays
                    bp.set(att, "0")
    prs.save(str(dst))


def _moved(a: dict, b: dict) -> str | None:
    if "skip" in b:
        return "rewrapped"
    if a["lines"] != b["lines"]:
        return "rewrapped"
    for k in ("bl", "bt", "bw", "bh"):
        if abs(a[k] - b[k]) > BOUNDS_TOL_PT:
            return "moved"
    return None


def _changed_pixels(png_a: Path, png_b: Path) -> int:
    from PIL import Image
    import numpy as np
    a = Image.open(png_a).convert("RGB")
    b = Image.open(png_b).convert("RGB")
    if a.size != b.size:
        return a.size[0] * a.size[1]
    d = np.abs(np.asarray(a, dtype=np.int16) - np.asarray(b, dtype=np.int16)).max(axis=2)
    return int((d > PIXEL_DELTA).sum())


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------
def _record(status, **kw) -> dict:
    rec = {"status": status, "engine": kw.pop("engine", None), "fitted": 0,
           "skipped_by_rule": {}, "reverted_boxes": 0,
           "reverted_pages": [], "moved": 0, "pixel_pages_changed": 0,
           "rounds": 0, "pages": "all", "note": "", "seconds": 0.0}
    rec.update(kw)
    return rec


def fit_deck(path, out=None, pages=None, log=print, engine_factory=None) -> dict:
    """Fit the text boxes of `path`; write the result to `out` (default: in
    place). Writes only when the proof passed and something was fitted."""
    t0 = time.monotonic()
    path = Path(path)
    out = Path(out) if out else path
    pages_set = set(pages) if pages is not None else None
    pages_label = "all" if pages_set is None else sorted(pages_set)
    if os.environ.get("SLIDE_LAB_NO_BOX_FIT", "").strip().lower() in ("1", "true", "yes"):
        return _record("skipped", pages=pages_label,
                       note="turned off for this run (SLIDE_LAB_NO_BOX_FIT=1)")
    if engine_factory is None:
        if not powerpoint_here():
            return _record("skipped", pages=pages_label, note=NO_PPT_NOTE)
        engine_factory = _powerpoint_engine
        engine_name = "PowerPoint"
    else:
        engine_name = getattr(engine_factory, "engine_name", "test engine")

    from pptx import Presentation
    import _scratch
    with _scratch.temp_dir("boxfit_") as td:
        td = Path(td)
        orig = td / "before.pptx"
        shutil.copyfile(path, orig)
        prs = Presentation(str(orig))
        if pages_set is not None:
            pages_set = {p for p in pages_set if 1 <= p <= len(prs.slides)}
        cands, skipped = plan(prs, pages_set)
        del prs
        rec = _record("unchanged", engine=engine_name, pages=pages_label)
        with engine_factory() as eng:
            meas_a = eng.measure(orig, list(cands))
            fits = {}
            for k, m in meas_a.items():
                if "skip" in m:
                    skipped[m["skip"]] += 1
                    continue
                g = compute_fit(m)
                if g is None:
                    skipped["already fits its text"] += 1
                else:
                    fits[k] = g
            rec["skipped_by_rule"] = dict(sorted(skipped.items()))
            if not fits:
                rec["note"] = "no box was larger than its text"
                rec["seconds"] = round(time.monotonic() - t0, 1)
                return rec
            before_png = eng.render(orig, {k[0] for k in fits}, td / "before_png")
            reverted, rev_pages = set(), set()
            accepted = None
            for rnd in range(1, MAX_ROUNDS + 1):
                rec["rounds"] = rnd
                live = {k: g for k, g in fits.items() if k not in reverted}
                if not live:
                    break
                trial = td / f"round{rnd}.pptx"
                _apply(orig, trial, live, meas_a)
                meas_b = eng.measure(trial, list(live))
                bad = {k for k in live if _moved(meas_a[k], meas_b.get(k, {"skip": "?"}))}
                changed = []
                if not bad:
                    after_png = eng.render(trial, {k[0] for k in live}, td / f"round{rnd}_png")
                    for p, png in sorted(after_png.items()):
                        if _changed_pixels(before_png[p], png) > PIXEL_PAGE_LIMIT:
                            changed.append(p)
                log(f"round {rnd}: {len(live)} box(es) fitted, {len(bad)} moved or "
                    f"rewrapped, {len(changed)} page(s) look different")
                if not bad and not changed:
                    accepted = (trial, live)
                    break
                reverted |= bad
                for p in changed:
                    rev_pages.add(p)
                    reverted |= {k for k in live if k[0] == p}
            rec["reverted_boxes"] = len(reverted)
            rec["reverted_pages"] = sorted(rev_pages)
        if accepted is None:
            rec["note"] = (f"the proof did not pass within {MAX_ROUNDS} rounds, so every "
                           "box was kept as it was")
            rec["seconds"] = round(time.monotonic() - t0, 1)
            return rec
        trial, live = accepted
        if out.resolve() != path.resolve():
            out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(trial, out)
        rec.update(status="fitted", fitted=len(live),
                   insets_zeroed=sum(1 for g in live.values() if g[4]),
                   moved=0, pixel_pages_changed=0,
                   seconds=round(time.monotonic() - t0, 1))
        return rec


def summary_lines(rec: dict | None) -> list[str]:
    """Plain-English lines for COMPILED.md, the compile output and check_done."""
    if not rec:
        return ["- not run (this deck was compiled before the text-box fit step, 2026-10-16)"]
    st = rec.get("status")
    if st == "skipped":
        return [f"- skipped: {rec.get('note')}"]
    if st == "error":
        return [f"- FAILED: {rec.get('note')}. The deck itself is fine but its text boxes "
                "were not fitted; recompile to try again."]
    lines = []
    pages = rec.get("pages")
    where = "every page" if pages == "all" else f"page(s) {', '.join(map(str, pages or []))}"
    if st == "fitted":
        lines.append(f"- {rec.get('fitted', 0)} custom text box(es) fitted to their text on "
                     f"{where}, measured by {rec.get('engine')} (nothing from the template "
                     "is changed: placeholders, page numbers and fixed slots keep their size)")
        if rec.get("insets_zeroed"):
            lines.append(f"- {rec['insets_zeroed']} of them had their left, right and bottom "
                         "margins set to 0 (the margin became part of the box, so the text "
                         "and its line breaks stay; the top margin is kept)")
        lines.append(f"- proof: 0 boxes moved or rewrapped (0.6 pt, line count), 0 pages look "
                     f"different (pixel check), {rec.get('rounds', 0)} round(s), "
                     f"{rec.get('seconds', 0)} s")
    else:
        lines.append(f"- no box changed on {where}: {rec.get('note') or 'nothing to fit'}")
    if rec.get("reverted_boxes"):
        rp = rec.get("reverted_pages") or []
        lines.append(f"- {rec['reverted_boxes']} box(es) put back by the proof"
                     + (f" (every box on page(s) {', '.join(map(str, rp))}, which looked different)"
                        if rp else ""))
    sk = rec.get("skipped_by_rule") or {}
    if sk:
        lines.append("- kept as they are: " + "; ".join(f"{n} {why}" for why, n in sk.items()))
    return lines


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------
def _parse_pages(text: str) -> set:
    out = set()
    for part in (text or "").replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            out |= set(range(int(a), int(b) + 1))
        else:
            out.add(int(part))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Fit text boxes to their text (PowerPoint measures, then proves no text moved).")
    ap.add_argument("deck", type=Path)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--pages", help="only these pages, e.g. 3,7-9 (the pages Slide Lab changed)")
    g.add_argument("--all-pages", action="store_true", help="every page (a deck Slide Lab built)")
    ap.add_argument("--out", type=Path, default=None,
                    help="where to write the fitted deck (default: <name>_fitted.pptx beside it)")
    args = ap.parse_args(argv)
    deck = args.deck
    if not deck.exists():
        print(f"ERROR: no deck at {deck}")
        return 2
    out = args.out or deck.with_name(f"{deck.stem}_fitted{deck.suffix}")
    if out.resolve() == deck.resolve():
        print("ERROR: --out is the original deck. The original is never overwritten; "
              "pass a different --out.")
        return 2
    try:
        pages = None if args.all_pages else _parse_pages(args.pages)
    except ValueError:
        print(f"ERROR: could not read --pages {args.pages!r}; write it like 3,7-9")
        return 2
    rec = fit_deck(deck, out=out, pages=pages, log=lambda s: print("  " + s))
    for ln in summary_lines(rec):
        print(ln)
    if rec["status"] == "fitted":
        print(f"Fitted deck: {out}")
        print("Next: run slide-qc on the fitted deck before sending it.")
        return 0
    print(f"No fitted deck was written; {deck.name} is unchanged.")
    print("Next: send the original as it is, or run this on Windows with PowerPoint.")
    return 0 if rec["status"] in ("unchanged", "skipped") else 1


if __name__ == "__main__":
    sys.exit(main())
