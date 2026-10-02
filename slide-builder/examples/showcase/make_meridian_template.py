#!/usr/bin/env python3
"""Build the fictional "Meridian" PowerPoint template used for the
Slide Lab showcase examples.

Meridian is invented. It exists so the shareable examples never show a real
client's brand. The template is a plain 16:9 (13.333 x 7.5 in) file with six
layouts, a navy and amber palette, and Arial throughout, so it renders the same
on any PC or Mac.

  py -3 make_meridian_template.py "<output folder>"
Writes "<output folder>/Meridian Template.pptx".
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = "0B3C49"     # main brand color (dk2)
AMBER = "F2A541"    # highlight color (accent1)
TEAL = "2A7F8E"     # secondary (accent2)
SLATE = "5B6770"    # muted text (accent3)
MIST = "E9EEF0"     # light fill (lt2)
INK = "1F2933"      # body text (dk1)
FONT = "Arial"

W, H = 12192000, 6858000
SCALE_X = W / 9144000          # python-pptx's default template is 4:3

KEEP = {                       # default layout name -> Meridian name
    "Title Slide": "Cover",
    "Title and Content": "Title and Content",
    "Section Header": "Section Divider",
    "Two Content": "Two Content",
    "Title Only": "Title Only",
    "Blank": "Blank",
}

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def _scale_x(root) -> None:
    for off in root.iter(qn("a:off")):
        off.set("x", str(int(int(off.get("x")) * SCALE_X)))
    for ext in root.iter(qn("a:ext")):
        if ext.get("cx") is not None:
            ext.set("cx", str(int(int(ext.get("cx")) * SCALE_X)))


def _theme(prs) -> None:
    master = prs.slide_masters[0]
    part = next(r.target_part for r in master.part.rels.values()
                if r.reltype.endswith("/theme"))
    root = etree.fromstring(part.blob)
    root.set("name", "Meridian")
    cs = root.find(".//a:clrScheme", NS)
    cs.set("name", "Meridian")
    colors = {"dk1": INK, "lt1": "FFFFFF", "dk2": NAVY, "lt2": MIST,
              "accent1": AMBER, "accent2": TEAL, "accent3": SLATE,
              "accent4": "8FB8C2", "accent5": "C9792B", "accent6": "A3ADB4",
              "hlink": TEAL, "folHlink": SLATE}
    for slot, hexv in colors.items():
        el = cs.find(f"a:{slot}", NS)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, qn("a:srgbClr")).set("val", hexv)
    fs = root.find(".//a:fontScheme", NS)
    fs.set("name", "Meridian")
    for tag in ("a:majorFont", "a:minorFont"):
        fs.find(tag, NS).find("a:latin", NS).set("typeface", FONT)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                standalone=True)


def _set_xfrm(shape, x, y, w, h) -> None:
    shape.left, shape.top, shape.width, shape.height = (
        Inches(x), Inches(y), Inches(w), Inches(h))


def _top_left(ph) -> None:
    """Text sits at the top-left of the box, under the amber bar."""
    body = ph._element.find(qn("p:txBody"))
    body.find(qn("a:bodyPr")).set("anchor", "t")
    lst = body.find(qn("a:lstStyle"))
    lvl = lst.find(qn("a:lvl1pPr"))
    if lvl is None:
        lvl = etree.SubElement(lst, qn("a:lvl1pPr"))
    lvl.set("algn", "l")
    lvl.set("marL", "0")
    lvl.set("indent", "0")


def _text_styles(master) -> None:
    tx = master._element.find(qn("p:txStyles"))
    title = tx.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    title.set("algn", "l")
    rpr = title.find(qn("a:defRPr"))
    rpr.set("sz", "2800")
    rpr.set("b", "1")
    for child in list(rpr):
        if child.tag == qn("a:solidFill"):
            rpr.remove(child)
    fill = etree.Element(qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr")).set("val", "tx2")
    rpr.insert(0, fill)
    sizes = ["1600", "1400", "1200", "1200", "1200"]
    body = tx.find(qn("p:bodyStyle"))
    for i, sz in enumerate(sizes, start=1):
        lvl = body.find(qn(f"a:lvl{i}pPr"))
        if lvl is not None:
            lvl.find(qn("a:defRPr")).set("sz", sz)


def _move_into(target_tree, shape) -> None:
    """Move a scratch-slide shape onto a master/layout: drop its theme style
    (explicit fill and line are set, and the style adds a shadow) and give it
    an id no other shape in the target uses (PowerPoint refuses duplicates)."""
    el = shape._element
    el.getparent().remove(el)
    style = el.find(qn("p:style"))
    if style is not None:
        el.remove(style)
    used = [int(c.get("id")) for c in target_tree.iter(qn("p:cNvPr"))]
    el.find(".//" + qn("p:cNvPr")).set("id", str(max(used + [1]) + 1))
    target_tree.append(el)


def _decorate(prs, master, layouts) -> None:
    """Shapes drawn on a scratch slide, then moved onto the master/layouts."""
    scratch = prs.slides.add_slide(layouts["Blank"])
    sh = scratch.shapes

    bar = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.32), Inches(0.7), Inches(0.07))
    # Brand line top right: the bottom strip belongs to Slide Lab's footnote
    # and source lines, which landed on top of a bottom-left brand line.
    brand = sh.add_textbox(Inches(6.88), Inches(0.2), Inches(6.0), Inches(0.3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(AMBER)
    bar.line.fill.background()
    p = brand.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r2 = p.add_run()
    r2.text = "Illustrative example, not real data   "
    r1 = p.add_run()
    r1.text = "MERIDIAN"
    for r, bold, hexv in ((r1, True, NAVY), (r2, False, SLATE)):
        r.font.size, r.font.bold, r.font.name = Pt(9), bold, FONT
        r.font.color.rgb = RGBColor.from_string(hexv)
    tree = master.shapes._spTree
    for s in (bar, brand):
        _move_into(tree, s)

    # Cover and section divider: navy background, white title, amber bar.
    for name in ("Cover", "Section Divider"):
        lay = layouts[name]
        lay._element.set("showMasterSp", "0")
        csld = lay._element.find(qn("p:cSld"))
        bg = etree.SubElement(csld, qn("p:bg"))
        csld.remove(bg)
        csld.insert(0, bg)
        bgpr = etree.SubElement(bg, qn("p:bgPr"))
        f = etree.SubElement(bgpr, qn("a:solidFill"))
        etree.SubElement(f, qn("a:srgbClr")).set("val", NAVY)
        etree.SubElement(bgpr, qn("a:effectLst"))
        for ph in lay.placeholders:
            if ph.placeholder_format.idx in (10, 11, 12):   # date/footer/number
                continue
            lst = ph._element.find(qn("p:txBody")).find(qn("a:lstStyle"))
            lvl = lst.find(qn("a:lvl1pPr"))
            if lvl is None:
                lvl = etree.SubElement(lst, qn("a:lvl1pPr"))
            # One defRPr per level: PowerPoint repairs a file with two.
            rpr = lvl.find(qn("a:defRPr"))
            if rpr is None:
                rpr = etree.SubElement(lvl, qn("a:defRPr"))
            for old in rpr.findall(qn("a:solidFill")):
                rpr.remove(old)
            is_title = ph.placeholder_format.type in (1, 3)
            rpr.set("sz", "3600" if is_title else "1600")
            fill = etree.Element(qn("a:solidFill"))
            rpr.insert(0, fill)
            etree.SubElement(fill, qn("a:srgbClr")).set(
                "val", "FFFFFF" if is_title else "C9D3D8")
        b2 = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.5), Inches(1.0), Inches(0.09))
        b2.fill.solid()
        b2.fill.fore_color.rgb = RGBColor.from_string(AMBER)
        b2.line.fill.background()
        b2.shadow.inherit = False
        _move_into(lay.shapes._spTree, b2)

    # drop the scratch slide
    sld_id = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)


def build(out_dir: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    master = prs.slide_masters[0]
    _scale_x(master._element)
    for lay in list(prs.slide_layouts):
        if lay.name not in KEEP:
            prs.slide_layouts.remove(lay)
            continue
        _scale_x(lay._element)
        lay._element.find(qn("p:cSld")).set("name", KEEP[lay.name])
    layouts = {l.name: l for l in prs.slide_layouts}
    _theme(prs)
    _text_styles(master)
    # The default layouts override sizes and set ALL CAPS on the section
    # header; drop that so every layout follows the master's text styles.
    for lay in layouts.values():
        for rpr in lay._element.iter(qn("a:defRPr")):
            for attr in ("sz", "cap"):
                rpr.attrib.pop(attr, None)

    # Content geometry: title band at the top, body below, footer strip.
    for ph in master.placeholders:
        t = ph.placeholder_format.type
        if t == 1:
            _set_xfrm(ph, 0.45, 0.45, 12.43, 0.95)
        elif t == 2:
            _set_xfrm(ph, 0.45, 1.6, 12.43, 5.1)
    for name in ("Title and Content", "Two Content", "Title Only"):
        for ph in layouts[name].placeholders:
            idx = ph.placeholder_format.idx
            if idx == 0:
                _set_xfrm(ph, 0.45, 0.45, 12.43, 0.95)
            elif name == "Title and Content" and idx == 1:
                _set_xfrm(ph, 0.45, 1.6, 12.43, 5.1)
            elif name == "Two Content" and idx == 1:
                _set_xfrm(ph, 0.45, 1.6, 6.05, 5.1)
            elif name == "Two Content" and idx == 2:
                _set_xfrm(ph, 6.83, 1.6, 6.05, 5.1)
    for name in ("Cover", "Section Divider"):
        for ph in layouts[name].placeholders:
            idx = ph.placeholder_format.idx
            if idx == 0:
                _set_xfrm(ph, 0.8, 2.8, 11.0, 1.5)
            elif idx == 1:
                _set_xfrm(ph, 0.8, 4.4, 11.0, 1.0)
            if idx in (0, 1):
                _top_left(ph)
    for lay in (master, *layouts.values()):
        for ph in lay.placeholders:
            if ph.placeholder_format.type == 13 or ph.placeholder_format.idx == 12:
                _set_xfrm(ph, 11.4, 6.98, 1.48, 0.3)

    _decorate(prs, master, layouts)
    prs.core_properties.title = "Meridian Template"
    prs.core_properties.author = "Slide Lab showcase"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "Meridian Template.pptx"
    prs.save(str(out))
    return out


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    print(build(Path(sys.argv[1])))
