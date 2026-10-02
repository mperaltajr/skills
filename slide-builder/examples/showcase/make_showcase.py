#!/usr/bin/env python3
"""Turn one real Meridian showcase run into the three shareable example files.

  py -3 make_showcase.py --session "<showcase session folder>" --dest "<folder>"

Writes to <dest>:
  Slide-Lab-Example-Review.html     the first-round review page exactly as Slide
                                    Lab wrote it (3 options per slide), images
                                    embedded, local paths removed, plus a banner
                                    saying how the example was made
  Slide-Lab-Example-Storyline.html  the storyline page Slide Lab wrote
  Slide-Lab-Example-Deck.pptx       the compiled deck, with speaker notes saying
                                    what each slide shows and how it was made

Nothing in the session folder changes: the review page is rebuilt in a temporary
copy. Slides redesigned after QC have their first-round options in
slide_NN/_prev/<stamp>/; the copy restores those so the page shows round one.

Run after a showcase run and before sharing. RUN.md in this folder lists the
whole sequence. The honesty note comes from HOW_MADE below: update it with the
times and edits measured in the run (see _session/timing.log).
"""
from __future__ import annotations

import argparse
import base64
import io
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parents[1] / "scripts"

HOW_MADE = [
    "Made on 2 October 2026 in one real Slide Lab run, on a fictional template "
    "(Meridian) with invented numbers. Nothing here is client material.",
    "Storyline: written from a finished data sheet, 2 minutes. A real storyline "
    "conversation usually takes 10 to 30 minutes.",
    "Designs: 10 slides with 3 options each, all designed at the same time, "
    "13 minutes.",
    "After the picks: converting them to editable PowerPoint took 25 seconds "
    "(no agent needed, including the growth loop); finalize and compile took "
    "under 2 minutes.",
    "Quality check: 1 Major finding (the slide 5 chart axis started at $300M, "
    "which overstated the gap). Slide 5 was redesigned in 2 minutes and the "
    "deck recompiled; 6 Advisory items remain (mostly titles that wrap).",
    "Hand edits, all disclosed: slide 6's footer placeholder was replaced with "
    "the deck's source line before review; slide 8's takeaway was shortened "
    "from 131 to 98 characters before the final check. Everything else is the "
    "pipeline's own output.",
]

SLIDE_NOTES = {
    1: "Built on the template's own cover layout. Covers are never numbered.",
    2: "Executive summary as a dense question-and-answer table, the ask highlighted.",
    3: "Bar chart with one highlight (the 2030 bar) and its callout drawn on the chart.",
    4: "Scored table: five markets against two cut-offs, passing rows shaded, small "
       "bars showing how far each market is from the cut-off.",
    5: "Line chart with the crossover marked. Redesigned after the quality check so "
       "the axis starts at zero and the drawn gap matches the 1.5x on the slide.",
    6: "Numbered process with icons: Slide Lab's default of making the structure visible.",
    7: "Timeline with the two go/no-go gates as diamonds and what happens after each.",
    8: "Growth loop diagram, converted to editable PowerPoint shapes by the script.",
    9: "Waterfall that builds the $12M total from its five parts.",
    10: "Decision table with owner and date, the decision being asked for highlighted.",
}


def _banner() -> str:
    items = "".join(f"<li>{t}</li>" for t in HOW_MADE)
    return (
        '<div style="font-family:Arial,sans-serif;background:#FFF7E8;border:1px solid '
        '#F2A541;border-radius:6px;padding:16px 20px;margin:16px;color:#1F2933;'
        'font-size:14px;line-height:1.5">'
        '<div style="font-size:16px;font-weight:bold;margin-bottom:6px">'
        'Example: the review page Slide Lab wrote for this run</div>'
        '<p style="margin:0 0 8px">This is where you choose. For each slide, '
        '<b>Pick</b> one of three designs, <b>Replace these</b> to get a new design, '
        'or <b>Leave out</b>. The one-click feedback buttons cover common asks. In '
        'this shared copy your clicks are saved only in your own browser, and '
        '<b>Build my deck</b> copies a command that works only inside a real '
        'Slide Lab session. After this page, Slide Lab shows every pick finished '
        'on the template for a final look, then builds the deck.</p>'
        '<div style="font-weight:bold;margin-top:8px">How this example was made</div>'
        f'<ul style="margin:4px 0 0 18px;padding:0">{items}</ul></div>')


def _data_uri(path: Path) -> str:
    from PIL import Image
    im = Image.open(path).convert("RGB")
    if im.width > 1280:
        im = im.resize((1280, round(im.height * 1280 / im.width)))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=85, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def _file_url_to_path(url: str) -> Path:
    from urllib.parse import unquote
    p = unquote(url[len("file:///"):])
    return Path(p)


def build_review(session: Path, dest: Path) -> Path:
    with tempfile.TemporaryDirectory(prefix="showcase_") as td:
        tmp = Path(td) / "run"
        shutil.copytree(session, tmp, ignore=shutil.ignore_patterns(
            "_qc", "_qc*", "final_pngs", "_render_tmp", "_raw", "*.pptx.bak"))
        # Round one for slides redesigned after QC.
        for prev in sorted(tmp.glob("slide_*/_prev/*")):
            slide_dir = prev.parent.parent
            for f in slide_dir.glob("option_*"):
                f.unlink()
            for f in prev.iterdir():
                shutil.copy2(f, slide_dir / f.name)
        # Round one, before any pick: a sketch option shows as its sketch, not
        # as the finished slide it became after it was picked and converted.
        for html_opt in tmp.glob("slide_*/option_?.html"):
            stem = html_opt.stem
            for suffix in (".png", ".pptx", ".qc.json", "_native.py",
                           "_native.plan.json", "_translation_report.json"):
                f = html_opt.with_name(stem + suffix)
                if f.exists():
                    f.unlink()
        r = subprocess.run([sys.executable, str(SCRIPTS / "build_review.py"), "--out", str(tmp)],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        if r.returncode != 0:
            raise SystemExit("build_review failed:\n" + r.stdout[-1500:] + r.stderr[-1500:])
        html = (tmp / "REVIEW.html").read_text(encoding="utf-8")

        def _inline(m):
            p = _file_url_to_path(m.group(1))
            return f'src="{_data_uri(p)}"' if p.exists() else m.group(0)
        html = re.sub(r'src="(file:///[^"]+\.png)"', _inline, html)
        # No local paths in a shared file: the run folder (short and long forms),
        # then any remaining user folder.
        for base in {tmp, tmp.resolve(), session, session.resolve()}:
            for form in (str(base), str(base).replace("\\", "\\\\")):
                html = html.replace(form, "&lt;your session folder&gt;")
        html = re.sub(r"([A-Za-z]:(?:\\\\|\\)Users(?:\\\\|\\))[^\\<>\"']+",
                      lambda m: m.group(1) + "&lt;you&gt;", html)
        html = re.sub(r"(<body[^>]*>)", r"\1" + _banner().replace("\\", "\\\\"), html, count=1)
    out = dest / "Slide-Lab-Example-Review.html"
    out.write_text(html, encoding="utf-8")
    return out


def build_deck(session: Path, dest: Path) -> Path:
    from pptx import Presentation
    prs = Presentation(str(session / "final_deck.pptx"))
    cp = prs.core_properties
    cp.title = "Slide Lab example: Meridian Southeast Asia entry"
    cp.author = cp.last_modified_by = "Slide Lab example"
    cp.subject = cp.keywords = cp.comments = ""
    for i, slide in enumerate(prs.slides, start=1):
        text = SLIDE_NOTES.get(i, "")
        if i == 1:
            text += "\n\nHow this example was made:\n" + "\n".join(f"- {t}" for t in HOW_MADE)
        slide.notes_slide.notes_text_frame.text = text
    out = dest / "Slide-Lab-Example-Deck.pptx"
    prs.save(str(out))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--session", required=True, type=Path)
    ap.add_argument("--dest", required=True, type=Path)
    a = ap.parse_args(argv)
    a.dest.mkdir(parents=True, exist_ok=True)
    made = [build_review(a.session, a.dest)]
    dd = sorted(a.session.glob("dot-dash-*.html"))
    if dd:
        shutil.copy2(dd[0], a.dest / "Slide-Lab-Example-Storyline.html")
        made.append(a.dest / "Slide-Lab-Example-Storyline.html")
    made.append(build_deck(a.session, a.dest))
    leaks = []
    for f in made:
        if f.suffix == ".html":
            t = f.read_text(encoding="utf-8")
            for bad in ("m.a.peralta", "file:///", "— [add", "[add source here"):
                if bad in t:
                    leaks.append(f"{f.name}: contains {bad!r}")
    for f in made:
        print(f"  {f}  ({f.stat().st_size // 1024} KB)")
    if leaks:
        print("CHECK FAILED:\n  " + "\n  ".join(leaks))
        return 1
    print("[ok] no local paths in the shared files")
    return 0


if __name__ == "__main__":
    sys.exit(main())
